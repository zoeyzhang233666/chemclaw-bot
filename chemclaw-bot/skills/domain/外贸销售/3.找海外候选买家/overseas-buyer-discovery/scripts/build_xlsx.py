# -*- coding: utf-8 -*-
"""Build the Indonesian heavy soda ash lead workbook (Excel).
Input : leads_merged.json  (list of lead dicts), market_facts.json (optional)
Output: 印尼重质纯碱_海外候选买家线索工作簿.xlsx
"""
import json, os, sys, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

HERE = os.path.dirname(os.path.abspath(__file__))
LEADS = os.path.join(HERE, "leads_merged.json")
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "印尼重质纯碱_海外候选买家线索工作簿.xlsx")
TODAY = datetime.date(2026, 10, 19)

C_INC = "1F4E79"; C_ACC = "2E75B6"; C_LT = "DDEBF7"; C_A = "C6EFCE"; C_B = "FFEB9C"; C_C = "F2F2F2"; C_HEAD = "1F3864"
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()

def header(ws, row, cols, widths, fill=C_INC):
    for i, (c, w) in enumerate(zip(cols, widths), start=1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font = Font(bold=True, color="FFFFFF", size=10, name="Microsoft YaHei")
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BOX
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 34

def title_block(ws, text, sub, ncols):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=text)
    c.font = Font(bold=True, size=15, color="FFFFFF", name="Microsoft YaHei")
    c.fill = PatternFill("solid", fgColor=C_HEAD)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 34
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    c2 = ws.cell(row=2, column=1, value=sub)
    c2.font = Font(size=9, color="595959", name="Microsoft YaHei")
    c2.fill = PatternFill("solid", fgColor=C_LT)
    c2.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
    ws.row_dimensions[2].height = 30

leads = json.load(open(LEADS, encoding="utf-8"))
try:
    facts = json.load(open(os.path.join(HERE, "market_facts.json"), encoding="utf-8"))
except Exception:
    facts = []

def tname(t):
    return {"A": "A 有直接贸易证据", "B": "B 有产品或业务相关证据", "C": "C 仅根据行业关系推断"}.get(t, t)

# ---------------------------------------------------------------- 0 说明
ws = wb.active
ws.title = "00_说明与证据规则"
title_block(ws, "印度尼西亚 · 重质纯碱(Heavy/Dense Soda Ash, Na₂CO₃ 99.2%) 海外候选买家线索工作簿",
            "用途：外贸销售候选客户发现与证据核验。所有条目均为“候选”，不代表已确认买家，也不代表存在与贵司的采购关系。生成日期 %s。" % TODAY.isoformat(), 8)
rows = [
    ("字段", "含义"),
    ("公司", "经公开来源核实的公司法定名称（印尼多为 PT xxx）；集团/上市简称括注"),
    ("类型", "终端生产企业＝自己用该产品做生产；直接进口商＝自行报关进口并转售；经销商＝批发分销（是否自行进口未确认）；待确认"),
    ("产品相关性", "该产品在其工艺/配方中的位置与用量逻辑"),
    ("证据等级", "A 有直接贸易证据 / B 有产品或业务相关证据 / C 仅根据行业关系推断"),
    ("A 的判定", "公开来源直接出现该企业就该产品（或明确同类品）的采购、进口、消耗记录，并能给出产品、方向、期间"),
    ("B 的判定", "公开来源能证明企业真实存在且其业务对该产品有明确产品或工艺关联（如产品目录列明、工艺必然消耗），但无交易记录"),
    ("C 的判定", "仅凭“属于某行业”推断，未取得任何产品或贸易层面证据——只能用于电话核实，不得写入报价或承诺"),
    ("可信度", "高＝一手来源（公司年报/官网/政府数据）；中＝行业目录、第三方数据库、报道；低＝单一弱来源或未核实"),
    ("使用红线", "① 不得把 C 级写成客户；② 不得虚构采购记录、联系人姓名或个人手机/邮箱；③ 官网与贸易数据须在联系前复核"),
    ("更新规则", "任何一条证据超过 18 个月，联系前应重新打开链接确认；年报类证据以最近一期为准"),
]
start = 4
for i, (a, b) in enumerate(rows):
    r = start + i
    ws.cell(row=r, column=1, value=a).font = Font(bold=True, name="Microsoft YaHei", size=10)
    ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor=C_LT)
    ws.cell(row=r, column=1).border = BOX
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
    c = ws.cell(row=r, column=2, value=b)
    c.font = Font(name="Microsoft YaHei", size=10)
    c.alignment = Alignment(wrap_text=True, vertical="center")
    c.border = BOX
    ws.row_dimensions[r].height = 30
r = start + len(rows) + 1
ws.cell(row=r, column=1, value="本工作簿工作表").font = Font(bold=True, size=11, name="Microsoft YaHei")
sheets_desc = [
    ("01_候选企业总表", "全部候选企业（可按类型/证据等级/地区/业务领域筛选），含证据摘要与链接"),
    ("02_证据明细", "逐条证据：等级、来源类型、期间、原文要点、链接"),
    ("03_公开联系方式", "仅公开企业级渠道（官网、总机、企业邮箱、官网表单），不含个人联系方式"),
    ("04_后市核验清单", "每条线索仍需验证的关键问题、核验渠道与优先级"),
    ("05_行业与需求地图", "纯碱下游用途、用量强度与印尼对应产业集群"),
    ("06_进口与需求数据", "印尼 HS 2836.20 进口金额/数量/来源国（UN Comtrade 实测）"),
    ("07_排除与同业", "同行/未来本土生产商/竞品供应渠道，避免误当客户"),
    ("08_话术与询盘草稿", "按客户类型与证据等级给出首触点话术（英文，仅供改写）"),
]
for i, (s, d) in enumerate(sheets_desc, start=1):
    rr = r + i
    ws.cell(row=rr, column=1, value=s).font = Font(name="Microsoft YaHei", size=10, bold=True)
    ws.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=8)
    ws.cell(row=rr, column=2, value=d).font = Font(name="Microsoft YaHei", size=10)
for col, w in zip("ABCDEFGH", [18, 16, 14, 14, 14, 14, 14, 24]):
    ws.column_dimensions[col].width = w

# ---------------------------------------------------------------- 1 总表
ws = wb.create_sheet("01_候选企业总表")
cols = ["编号", "公司名称", "别名/上市代码", "客户类型", "买方角色", "地区（省/市）", "主营业务",
        "产品相关性", "证据等级", "证据要点", "证据期间", "用量线索", "官网/主要来源", "公开联系渠道",
        "可信度", "关联原因分析", "首要核验问题", "跟进优先级"]
widths = [7, 34, 22, 16, 24, 22, 40, 40, 18, 60, 18, 22, 40, 40, 8, 46, 40, 10]
title_block(ws, "01 候选企业总表 — 印尼重质纯碱潜在终端工厂 / 进口商 / 经销商", "共 %d 条。A 级＝有公开的直接采购/消耗/进口证据；B 级＝产品与业务强相关；C 级＝仅行业推断。勾选/筛选后请先做电话或邮件核实。" % len(leads), len(cols))
header(ws, 3, cols, widths)

def prio(l):
    tier = l.get("evidence_tier", "C"); conf = l.get("confidence", "低")
    vol = l.get("volume_hint", "")
    big = any(k in vol for k in ["万吨", "10,000", "万"])
    if tier == "A" and big: return "P1 立即"
    if tier == "A": return "P2 本周"
    if tier == "B" and conf == "高": return "P2 本周"
    if tier == "B": return "P3 两周内"
    return "P4 观察"

for idx, l in enumerate(leads, start=1):
    r = 3 + idx
    tier = l.get("evidence_tier", "C")
    vals = [
        l.get("id", idx), l.get("company", ""), l.get("alt_names", ""), l.get("type", ""), l.get("buyer_role", ""),
        l.get("region", ""), l.get("main_business", ""), l.get("product_relation", ""), tier,
        l.get("evidence_summary", ""), l.get("evidence_period", ""), l.get("volume_hint", ""),
        (l.get("evidence_urls") or [""])[0], "；".join(l.get("contact") or []), l.get("confidence", ""),
        l.get("why_related", ""), "；".join((l.get("open_questions") or [])[:2]), prio(l),
    ]
    for ci, v in enumerate(vals, start=1):
        c = ws.cell(row=r, column=ci, value=v)
        c.font = Font(name="Microsoft YaHei", size=9)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.border = BOX
    ws.cell(row=r, column=9).alignment = Alignment(horizontal="center", vertical="center")
    fill = {"A": C_A, "B": C_B, "C": C_C}[tier]
    ws.cell(row=r, column=9).fill = PatternFill("solid", fgColor=fill)
    ws.cell(row=r, column=2).font = Font(name="Microsoft YaHei", size=9, bold=True)
    ws.row_dimensions[r].height = 78

ws.freeze_panes = "C4"
ws.auto_filter.ref = "A3:%s%d" % (get_column_letter(len(cols)), 3 + len(leads))
dv = DataValidation(type="list", formula1='"A 有直接贸易证据,B 有产品或业务相关证据,C 仅根据行业关系推断"', allow_blank=True)
ws.add_data_validation(dv)
dv.add("I4:I%d" % (3 + len(leads) + 50))

# ---------------------------------------------------------------- 2 证据明细
ws = wb.create_sheet("02_证据明细")
cols = ["证据ID", "公司名称", "证据等级", "证据类型", "来源名称", "来源链接", "覆盖期间", "原文要点（摘录/转述）", "可作为什么用", "局限与反证"]
widths = [9, 32, 18, 18, 22, 46, 16, 70, 34, 40]
title_block(ws, "02 证据明细 — 每条候选背后的可追溯出处", "所有链接均于任务期间实际抓取；若日后失效，请以公司名称+关键词重新检索并在“局限与反证”列记录差异。", len(cols))
header(ws, 3, cols, widths)
r = 4
for l in leads:
    for j, ev in enumerate(l.get("evidence_list") or []):
        vals = ["E%s-%d" % (l.get("id", ""), j + 1), l.get("company", ""), ev.get("tier", l.get("evidence_tier", "")),
                ev.get("type", ""), ev.get("source_name", ""), ev.get("url", ""), ev.get("period", ""),
                ev.get("quote", ""), ev.get("usable_as", ""), ev.get("caveat", "")]
        for ci, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=ci, value=v)
            c.font = Font(name="Microsoft YaHei", size=9)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = BOX
        ws.row_dimensions[r].height = 64
        r += 1
ws.freeze_panes = "C4"
ws.auto_filter.ref = "A3:J%d" % max(4, r - 1)

# ---------------------------------------------------------------- 3 联系方式
ws = wb.create_sheet("03_公开联系方式")
cols = ["公司名称", "类型", "官网", "企业总机/公开电话", "企业邮箱/联系表单", "地址", "社交/其他官方渠道", "联系注意", "核验日期"]
widths = [34, 16, 40, 26, 40, 46, 30, 40, 12]
title_block(ws, "03 公开联系方式 — 仅企业级公开渠道", "不记录任何人名、个人手机或个人邮箱。上市公司优先走投资者关系/公司秘书渠道；询价前先确认收件方是否为采购职能。", len(cols))
header(ws, 3, cols, widths)
r = 4
for l in leads:
    ct = l.get("contact_struct") or {}
    vals = [l.get("company", ""), l.get("type", ""), ct.get("website", (l.get("evidence_urls") or [""])[0]),
            ct.get("phone", ""), ct.get("email", ""), ct.get("address", ""), ct.get("social", ""),
            ct.get("note", ""), ct.get("verified_on", str(TODAY))]
    for ci, v in enumerate(vals, start=1):
        c = ws.cell(row=r, column=ci, value=v)
        c.font = Font(name="Microsoft YaHei", size=9)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.border = BOX
    ws.row_dimensions[r].height = 58
    r += 1
ws.freeze_panes = "B4"
ws.auto_filter.ref = "A3:I%d" % max(4, r - 1)

# ---------------------------------------------------------------- 4 后验清单
ws = wb.create_sheet("04_后市核验清单")
cols = ["公司名称", "本期证据等级", "待验证问题", "为什么必须验证", "建议核验渠道", "优先级"]
widths = [32, 18, 60, 44, 40, 10]
title_block(ws, "04 后市核验清单 — 每条线索的“下一步该查什么”", "按优先级排序：P1＝A 级大用量客户（先确认供应结构与决策链）；P2＝B 级高可信（先确认是否已从中国采购）；P4＝C 级（先确认是否真的用该产品）。", len(cols))
header(ws, 3, cols, widths)
r = 4
for l in sorted(leads, key=lambda x: ({'A': 0, 'B': 1, 'C': 2}.get(x.get("evidence_tier", "C"), 3), x.get("id", 0))):
    for q in (l.get("open_questions") or ["（补充待验证问题）"]):
        vals = [l.get("company", ""), tname(l.get("evidence_tier", "C")), q, l.get("why_related", "")[:90],
                l.get("verify_channel", "官网/年报/电话询证"), prio(l)]
        for ci, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=ci, value=v)
            c.font = Font(name="Microsoft YaHei", size=9)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = BOX
        ws.row_dimensions[r].height = 46
        r += 1
ws.freeze_panes = "B4"
ws.auto_filter.ref = "A3:F%d" % max(4, r - 1)

# ---------------------------------------------------------------- 5 行业地图
ws = wb.create_sheet("05_行业与需求地图")
cols = ["下游行业", "产品在该行业的作用", "工艺/配方位置", "单客户用量强度", "印尼对应集群/省市", "本次线索数", "说明"]
widths = [22, 40, 34, 18, 34, 10, 50]
title_block(ws, "05 行业与需求地图 — 重质纯碱在印尼的下游去向", "用量强度：★★★★★＝单客户万吨级/年；★★★★＝千吨级；★★★＝百吨级；★★＝十吨级；★＝零星。", len(cols))
header(ws, 3, cols, widths)
lane_map = [
    ("浮法玻璃/容器玻璃/汽车玻璃", "助熔剂：降低石英砂熔融温度与黏度，构成 Na₂O 组分", "配料仓→混合→熔窑（1550–1600℃）", "★★★★★", "Jawa Barat (Cikampek/Bekasi), Jawa Timur, DKI Jakarta", "glass"),
    ("日用玻璃/玻璃砖/餐具", "同上，钠钙玻璃配方 10–15% Na₂O", "配料→熔窑→成型", "★★★–★★★★", "Jawa Timur, Banten", "glass"),
    ("洗衣粉/洗涤剂粉体", "碳酸钠提供碱性、软化水硬度、作喷雾干燥载体", "配料→喷粉/附聚→包装", "★★★★", "Jawa Barat, Banten, DKI Jakarta", "detergent"),
    ("液体洗涤剂/工业清洗", "碱性调节与缓冲，皂化油脂", "配料混合", "★★", "Jawa Barat, DKI Jakarta", "detergent"),
    ("水玻璃/硅酸钠", "与石英砂高温熔融生成 Na₂SiO₃（硅酸钠）", "熔窑 1200–1400℃", "★★★–★★★★", "Jawa Timur, Jawa Tengah", "silicate"),
    ("纸浆造纸/化学品回收", "碱回收与漂白段碱度调节、硅回收", "蒸煮/碱回收/漂白", "★★★–★★★★", "Riau, Sumatera Selatan, Jawa Timur", "pulp_paper"),
    ("印染/纺织湿加工", "退浆、精练、染色浴碱剂", "染整工序", "★★–★★★", "Jawa Barat (Bandung/Majalaya/Cimahi), Pekalongan", "textile_ceramic"),
    ("陶瓷釉料/熔块", "釉料助熔与熔块配方（并非所有工艺都需要）", "釉料/熔块熔制", "★★–★★★", "Jawa Barat, Jawa Timur", "textile_ceramic"),
    ("水处理/锅炉", "pH 调节与硬度沉淀", "加药系统", "★★", "全国", "detergent"),
    ("化工分销/进口贸易", "分装、复配、转售给上述终端", "仓储/分装", "—（渠道）", "DKI Jakarta, Banten, Jawa Timur", "distributor"),
]
for i, (a, b, c_, d, e, lane) in enumerate(lane_map, start=4):
    cnt = sum(1 for l in leads if l.get("lane", "") == lane)
    vals = [a, b, c_, d, e, cnt, "本次线索见 01 表 lane=%s" % lane]
    for ci, v in enumerate(vals, start=1):
        cell = ws.cell(row=i, column=ci, value=v)
        cell.font = Font(name="Microsoft YaHei", size=9)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BOX
    ws.row_dimensions[i].height = 44
ws.freeze_panes = "B4"

# ---------------------------------------------------------------- 6 进口数据
ws = wb.create_sheet("06_进口与需求数据")
title_block(ws, "06 印尼 HS 2836.20（碳酸钠）进口规模 — UN Comtrade 实测序列", "来源：UN Comtrade 公开 API（reporter=Indonesia 360, cmdCode=283620, flow=M）。数量列为 netWgt 折算；个别年份进口方未报重量，显示为“未报”。", 9)
cols = ["年份", "进口额 (百万美元)", "进口量 (千吨)", "均价估算 (美元/吨)", "第一大来源国", "第二来源国", "第三来源国", "对华进口 (百万美元)", "中国出口至印尼 (千吨)"]
widths = [10, 20, 18, 18, 18, 18, 18, 20, 20]
header(ws, 3, cols, widths)
try:
    series = json.load(open(os.path.join(HERE, "comtrade_idn_283620.json"), encoding="utf-8"))
except Exception:
    series = []
try:
    chn = json.load(open(os.path.join(HERE, "china_to_idn_283620.json"), encoding="utf-8"))
except Exception:
    chn = {}
for i, s in enumerate(series, start=4):
    q = s.get("total_qty_kt") or None
    price = round(s["total_value_usd"] * 1e6 / (q * 1e3), 1) if q else "—"
    p = s.get("partners", []) + [{}, {}, {}]
    chn_v, chn_q = chn.get(str(s["year"]), [None, None])
    vals = [s["year"], s["total_value_usd"], (q if q else "未报"), price,
            "%s %s百万美元" % (p[0].get("name", ""), p[0].get("value_usd_m", "")),
            "%s %s百万美元" % (p[1].get("name", ""), p[1].get("value_usd_m", "")),
            "%s %s百万美元" % (p[2].get("name", ""), p[2].get("value_usd_m", "")),
            chn_v if chn_v is not None else "", chn_q if chn_q is not None else ""]
    for ci, v in enumerate(vals, start=1):
        c = ws.cell(row=i, column=ci, value=v)
        c.font = Font(name="Microsoft YaHei", size=9)
        c.border = BOX
        c.alignment = Alignment(horizontal="center" if ci <= 4 else "left", vertical="center", wrap_text=True)
r = 4 + len(series) + 1
for line in [
    "读数提示：",
    "· 印尼长期无本土纯碱产能，供应基本依赖进口（后续新建项目见 07 表），因此“进口来源国结构”直接决定终端工厂的替代空间。",
    "· 美国（ANSAC trona 天然碱）与土耳其（Ciner）合计占比高，中国货以重质/轻质定向报价切入，价格与批次一致性是关键卖点。",
    "· 2025 年中国→印尼出口量明显抬升（见“中国出口至印尼”列），说明中国货源已在这条链上被接受，可据此定位正在切换供应源的进口商。",
]:
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
    c = ws.cell(row=r, column=1, value=line)
    c.font = Font(name="Microsoft YaHei", size=9, bold=line.endswith("："))
    c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[r].height = 26
    r += 1

# ---------------------------------------------------------------- 7 排除
ws = wb.create_sheet("07_排除与同业")
cols = ["公司/机构", "判定", "理由", "证据", "处理建议"]
widths = [36, 18, 56, 46, 40]
title_block(ws, "07 排除与同业名单 — peer gate", "这些公司看似相关，但不是买家（或是未来竞争者）。写入名单前先过这一关，避免销售资源浪费与商业信息外泄。", len(cols))
header(ws, 3, cols, widths)
try:
    peers = json.load(open(os.path.join(HERE, "peer_gate.json"), encoding="utf-8"))
except Exception:
    peers = []
for i, p in enumerate(peers, start=4):
    vals = [p.get("company", ""), p.get("verdict", ""), p.get("why", ""), p.get("evidence", ""), p.get("action", "")]
    for ci, v in enumerate(vals, start=1):
        c = ws.cell(row=i, column=ci, value=v)
        c.font = Font(name="Microsoft YaHei", size=9)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.border = BOX
    ws.row_dimensions[i].height = 56
ws.freeze_panes = "B4"

# ---------------------------------------------------------------- 8 话术
ws = wb.create_sheet("08_话术与询盘草稿")
cols = ["适用对象", "适用证据等级", "主题/切入角度", "正文草稿（英文，需按实际规格改写）", "禁止表述", "跟进动作"]
widths = [26, 18, 32, 90, 36, 40]
title_block(ws, "08 话术与询盘草稿 — 按客户类型分层的首触点", "原则：先给“可验证的专业信息”，再要“一个低门槛的确认问题”；A 级客户可直接引用其公开披露的用量与来源结构，B/C 级不要假设对方在用重质纯碱。", len(cols))
header(ws, 3, cols, widths)
scripts = [
 ("A 级终端工厂（玻璃）", "A", "供应来源结构复核 + 批次一致性",
  "Subject: Dense soda ash supply option for your glass furnaces\n\nDear [Company] Procurement team,\n\nYour most recent annual & sustainability report lists soda ash consumption of about [X] t per year. We are a dense soda ash producer in [Province], China (Na2CO3 99.2% min, bulk density 0.9-1.1 g/cm3, Fe 0.003% max, 50 kg PP bags / 1.25 t jumbo bags), with ISO 9001 and COA per lot.\n\nWe would like to understand how your current supply is structured (origin country, annual volume, contract length) so that we can quote on a like-for-like basis. Would you be open to a 20-minute call, or shall we send a COA and a price indication for 500 t / 1,000 t?\n\nBest regards,\n[Name] | [Company] | [Email] | [Phone]",
  "不要写“我们了解贵司正在从某国采购”；不要在首封邮件承诺价格有效期超过 15 天",
  "3 个工作日未回→电话总机转采购部；2 周→寄样+COA"),
 ("A/B 级进口商与分销商", "A/B", "补足现有产品线的重质纯碱",
  "Subject: Dense soda ash (Na2CO3 99.2%) - CIF [Jakarta/Surabaya] price for Q1\n\nDear [Company],\n\nWe noticed your catalogue covers soda ash dense/light for glass and detergent customers. We produce dense soda ash in [Province], China, with 20,000 t/month of export capacity and full COA/TDS support.\n\nCan you confirm: (1) annual volume for dense grade, (2) preferred pack (50 kg / 1,000 kg jumbo), (3) target CIF port?\n\nPrice indication: USD [xxx]/t CIF [port] for 500 t, valid in principle for [15] days and subject to final confirmation.\n\nBest regards,\n[Name] | [Company]",
  "不要宣称“最低价”；不要对未核实的进口商报 FOB 最终价",
  "附 TDS+COA 样本；询问是否需要 SGS/BV 检验"),
 ("B 级终端（洗涤剂/水处理/造纸等）", "B", "工艺碱度与成本核算",
  "Subject: Soda ash (dense/light) for your [detergent/paper/water treatment] production\n\nDear [Company],\n\nWe supply soda ash for [application]. Based on your product range, dense grade is usually preferred for [reason: bulk density in slurry/silo handling] and light grade for [reason].\n\nWhich grade and purity do you require, and what is your annual consumption? We can send a COA and a small sample at your cost for trial.\n\nBest regards,\n[Name] | [Company]",
  "不要说“我们确定贵司使用重质纯碱”",
  "样品试用→小批量试单→年度框架"),
 ("C 级（仅行业推断）", "C", "需求确认型开场",
  "Subject: Quick question - soda ash usage in your glass/detergent line\n\nDear [Company],\n\nWe are a soda ash producer in China. Before sending any details, could you confirm whether your plant currently purchases soda ash (dense or light) and roughly how much per year?\n\nIf yes, I will send a COA and price indication right away. If you source through a distributor, we would be glad to know who, so we do not duplicate efforts.\n\nBest regards,\n[Name] | [Company]",
  "不要发送报价单、不要承诺账期",
  "电话优先；如确认在用，升级为 B/A 并补证据"),
]
for i, s in enumerate(scripts, start=4):
    for ci, v in enumerate(s, start=1):
        c = ws.cell(row=i, column=ci, value=v)
        c.font = Font(name="Microsoft YaHei", size=9)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.border = BOX
    ws.row_dimensions[i].height = 190

for sheet in wb.worksheets:
    sheet.sheet_view.showGridLines = False

wb.save(OUT)
print("saved", OUT, "leads:", len(leads))
