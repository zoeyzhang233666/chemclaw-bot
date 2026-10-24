# -*- coding: utf-8 -*-
"""
selfcheck.py — 一键自检

    python scripts/selfcheck.py [--skip-xlsx]

检查内容：
  1. 目录结构完整性（必需文件是否都在）
  2. 引擎可运行 + 关键恒等式（K1–K7）
  3. 关键数值与已知案例一致（回归基线）
  4. 缺失数据机制（该留空的确实留空了，没有用 0 顶替）
  5. 退税未进入任何术语成本（V9）
  6. HTML 生成成功且含关键片段
  7. Excel 生成成功、公式数量正常（可选：需 openpyxl）
"""
from __future__ import annotations

import json
import os
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)

from calc_engine import QuoteInput, compute, load_input  # noqa: E402

PASS, FAIL = [], []


def chk(cond, name, detail=""):
    if cond:
        PASS.append(name)
        print(f"  [PASS] {name}" + (f" — {detail}" if detail else ""))
    else:
        FAIL.append(f"{name}" + (f" — {detail}" if detail else ""))
        print(f"  [FAIL] {name}" + (f" — {detail}" if detail else ""))


def approx(a, b, tol=1e-4):
    return a is not None and b is not None and abs(a - b) <= tol


# ---------------------------------------------------------------- 1. 结构
def check_layout():
    print("\n[1] 目录结构与必需文件")
    need = [
        "SKILL.md", "README.md", "INSTALL.md", "VERSION", "CHANGELOG.md", "PORTABLE_PROMPT.md",
        "references/incoterms.md", "references/formulas.md", "references/parameters.md",
        "references/validation.md", "references/sensitivity.md", "references/qa.md",
        "references/portability.md", "references/citations.md",
        "scripts/calc_engine.py", "scripts/build_html.py", "scripts/build_xlsx.py",
        "scripts/selfcheck.py", "templates/calculator.template.html",
        "products/soda-ash.json", "products/_template.json",
        "examples/soda-ash-qingdao-jakarta/input.json",
        "examples/soda-ash-qingdao-jakarta/案例说明.md",
    ]
    missing = [p for p in need if not os.path.exists(os.path.join(ROOT, p))]
    chk(not missing, "必需文件齐全", f"缺 {len(missing)} 个: {missing[:5]}" if missing else f"{len(need)} 个文件")


# ---------------------------------------------------------------- 2/3. 引擎
def check_engine():
    print("\n[2] 计算引擎与恒等式")
    inp_path = os.path.join(ROOT, "examples", "soda-ash-qingdao-jakarta", "input.json")
    if not os.path.exists(inp_path):
        chk(False, "示例输入存在", inp_path)
        return None
    inp = load_input(inp_path)
    res = compute(inp)

    ids = res.get("identity_checks", {})

    def collect_ok(node, path=""):
        """identity_checks 是嵌套 dict，叶子是 {'ok': bool, ...}"""
        out = []
        if isinstance(node, dict) and "ok" in node and isinstance(node["ok"], bool):
            out.append((path, node["ok"]))
        elif isinstance(node, dict):
            for k, v in node.items():
                out.extend(collect_ok(v, f"{path}.{k}" if path else k))
        return out

    leaves = collect_ok(ids)
    bad = [p for p, ok in leaves if not ok]
    chk(not bad, "K 系列恒等式全部成立", f"{len(leaves)} 项" if not bad else f"失败: {bad[:4]}")

    print("\n[3] 数值回归基线")
    q = res["quotes"]
    base = {
        "FOB": (1740.0, 1891.30, 266.38),
        "CFR": (2095.0, 2277.17, 320.73),
        "CIF": (2101.9364, 2284.71, 321.79),
    }
    for term, (cost, price, fx) in base.items():
        got = q.get(term, {})
        chk(approx(got.get("cost_cny_per_t"), cost, 1e-3),
            f"{term} 成本 = {cost}", f"got {got.get('cost_cny_per_t')}")
        chk(approx(got.get("price_cny_per_t"), price, 0.01),
            f"{term} 报价 = {price}", f"got {got.get('price_cny_per_t')}")
        chk(approx(got.get("price_fx_per_t"), fx, 0.01),
            f"{term} 外币价 = {fx}", f"got {got.get('price_fx_per_t')}")

    # 保险费自洽解
    seg = res["segments"]
    s6 = seg["S6"]["sum_cny_per_t"]
    chk(approx(s6, 6.9364, 1e-3), "保险费为自洽解 6.9364", f"got {s6}")
    chk(approx(s6, 2101.9364 * 1.10 * 0.003, 1e-4), "S6 = 1.10r × C_CIF", f"got {s6}")

    # 毛利两口径恒等式（用未取整的 price_raw 校验，避免 2 位小数取整干扰）
    fob = q["FOB"]
    p_raw = fob.get("price_raw_cny_per_t") or fob["price_cny_per_t"]
    m = fob["gp_on_raw_cny_per_t"] / p_raw
    k = fob["gp_on_raw_cny_per_t"] / fob["cost_cny_per_t"]
    chk(approx(m, 0.08, 1e-6), "毛利率 = 8%（按未取整价格）", f"got {m:.6f}")
    chk(approx(k, 0.086957, 1e-5), "加成率 ≈ 8.6957%", f"got {k:.6f}")
    chk(approx((1 + k) * (1 - m), 1.0, 1e-6), "(1+k)(1−m) = 1", f"got {(1 + k) * (1 - m):.9f}")
    # 展示用的取整价毛利率应贴近 8%（允许取整误差）
    m_disp = fob["gp_cny_per_t"] / fob["price_cny_per_t"]
    chk(approx(m_disp, 0.08, 1e-4), "取整后展示毛利率仍 ≈ 8%", f"got {m_disp:.6f}")

    # 风险/费用边界
    S = {k: v["sum_cny_per_t"] for k, v in seg.items()}
    chk(approx(q["FOB"]["cost_cny_per_t"], sum(S[s] for s in ["S1", "S2", "S3", "S4", "S7", "S8"]), 1e-6),
        "FOB = S1+S2+S3+S4+S7+S8")
    chk(approx(q["CFR"]["cost_cny_per_t"], q["FOB"]["cost_cny_per_t"] + S["S5"], 1e-6),
        "CFR = FOB + S5")
    chk(approx(q["CIF"]["cost_cny_per_t"], q["CFR"]["cost_cny_per_t"] + S["S6"], 1e-6),
        "CIF = CFR + S6")
    chk(S["S9"] is None, "F/C 组不含 S9（S9 为 None，未被计入 FOB/CFR/CIF）", f"S9={S['S9']}")

    print("\n[4] 缺失数据机制")
    # 构造一个缺海运费的输入
    d = json.load(open(inp_path, encoding="utf-8"))
    d2 = dict(d)
    d2["freight_usd_per_container"] = None
    tmp = os.path.join(tempfile.gettempdir(), "_selfcheck_missing.json")
    json.dump(d2, open(tmp, "w", encoding="utf-8"), ensure_ascii=False)
    try:
        res2 = compute(load_input(tmp))
        cfr_cost = res2["quotes"]["CFR"]["cost_cny_per_t"]
        chk(cfr_cost is None, "缺海运费时 CFR 成本为 None（不是 0）", f"got {cfr_cost}")
        cif_price = res2["quotes"]["CIF"]["price_cny_per_t"]
        chk(cif_price is None, "缺海运费时 CIF 报价为 None（不是 0）", f"got {cif_price}")
        ms = " ".join(json.dumps(m, ensure_ascii=False) for m in res2.get("missing", []))
        chk("运费" in ms or "freight" in ms.lower(), "缺失清单里点名了海运费")
        chk(approx(res2["quotes"]["FOB"]["cost_cny_per_t"], 1740.0, 1e-6),
            "FOB 不受海运费缺失影响（仍可算）", f"got {res2['quotes']['FOB']['cost_cny_per_t']}")
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)

    print("\n[5] 退税未进入成本（V9）")
    refund = res.get("refund", {})
    chk(refund.get("enabled") is True, "退税已计算并单列", f"{refund.get('refund_cny_per_t')} ¥/t")
    # 退税金额不得出现在任何术语成本里：术语成本必须严格等于其组成段之和
    term_segs = {
        "FOB": ["S1", "S2", "S3", "S4", "S7", "S8"],
        "CFR": ["S1", "S2", "S3", "S4", "S5", "S7", "S8"],
        "CIF": ["S1", "S2", "S3", "S4", "S5", "S6", "S7", "S8"],
    }
    polluted = []
    for t, segs in term_segs.items():
        c = q.get(t, {}).get("cost_cny_per_t")
        ssum = sum(S[s] for s in segs)
        if c is not None and not approx(c, ssum, 1e-6):
            polluted.append(f"{t}: {c} != {ssum}")
    chk(not polluted, "退税未抵减任何术语成本（成本 = 组成段之和）",
        "; ".join(polluted) if polluted else "")
    return res


# ---------------------------------------------------------------- 6. HTML
def check_html(res):
    print("\n[6] HTML 生成")
    tpl = os.path.join(ROOT, "templates", "calculator.template.html")
    if not os.path.exists(tpl):
        chk(False, "HTML 模板存在")
        return
    src = open(tpl, encoding="utf-8").read()
    chk("__DATA__" in src, "模板含 __DATA__ 占位符")

    import subprocess
    out = os.path.join(tempfile.gettempdir(), "_selfcheck.html")
    inp = os.path.join(ROOT, "examples", "soda-ash-qingdao-jakarta", "input.json")
    r = subprocess.run([sys.executable, os.path.join(HERE, "build_html.py"), inp, "-o", out],
                       capture_output=True, text=True, encoding="utf-8", errors="replace")
    if r.returncode != 0:
        chk(False, "build_html.py 运行成功", (r.stderr or "")[-300:])
        return
    chk(True, "build_html.py 运行成功")
    html = open(out, encoding="utf-8").read()
    chk(len(html) > 50000, "HTML 体积正常", f"{len(html) // 1024} KB")
    chk("__DATA__" not in html, "占位符已被替换")
    for token in ["FOB", "CFR", "CIF", "Incoterms", "SIMULATED" if "SIMULATED" in html else "待补数据"]:
        chk(token in html, f"HTML 含关键内容: {token}")
    chk('charset="utf-8"' in html.lower().replace('"', '"') or "charset=utf-8" in html.lower(),
        "声明 UTF-8 编码")
    chk("http://" not in html.replace("http://www.w3.org", "") and "https://" not in html.replace("https://www.w3.org", ""),
        "无外部网络依赖（离线可用）")
    os.remove(out)


# ---------------------------------------------------------------- 7. Excel
def check_xlsx():
    print("\n[7] Excel 生成")
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        print("  [SKIP] 未安装 openpyxl，跳过 Excel 检查")
        return
    import subprocess
    out = os.path.join(tempfile.gettempdir(), "_selfcheck.xlsx")
    inp = os.path.join(ROOT, "examples", "soda-ash-qingdao-jakarta", "input.json")
    r = subprocess.run([sys.executable, os.path.join(HERE, "build_xlsx.py"), inp, "-o", out],
                       capture_output=True, text=True, encoding="utf-8", errors="replace")
    if r.returncode != 0:
        chk(False, "build_xlsx.py 运行成功", (r.stderr or "")[-400:])
        return
    chk(True, "build_xlsx.py 运行成功")
    from openpyxl import load_workbook
    wb = load_workbook(out)
    chk(len(wb.sheetnames) == 12, "工作簿含 12 张表", f"{len(wb.sheetnames)}")
    formulas = sum(1 for ws in wb.worksheets for row in ws.iter_rows() for c in row
                   if isinstance(c.value, str) and c.value.startswith("="))
    chk(formulas > 200, "公式单元格数量正常（>200）", f"{formulas} 个")
    chk("03 费用明细" in wb.sheetnames, "含唯一录入区 03 费用明细")
    chk("11 公式与校验" in wb.sheetnames, "含校验表 11")
    os.remove(out)


def main(argv=None):
    argv = argv or sys.argv[1:]
    print("=" * 66)
    print("  算出口报价 · 自检 (selfcheck)")
    print("  root:", ROOT)
    print("=" * 66)
    check_layout()
    res = check_engine()
    if res and "--skip-xlsx" not in argv:
        check_html(res)
        check_xlsx()
    elif "--skip-xlsx" not in argv:
        check_html(None)
    print("\n" + "=" * 66)
    print(f"  PASS {len(PASS)} / FAIL {len(FAIL)}")
    if FAIL:
        print("  失败项：")
        for f in FAIL:
            print("   -", f)
    else:
        print("  全部通过 ✓")
    print("=" * 66)
    return 0 if not FAIL else 1


if __name__ == "__main__":
    sys.exit(main())
