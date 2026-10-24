# -*- coding: utf-8 -*-
"""生成 6 页 Excel 核对工作簿：单证索引 / 跨文件字段矩阵 / 计算核验 / 差异定位 / 问题清单 / 修正记录。"""
import io
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

STATUS_CN = {"OK": "一致", "DIFF": "不一致", "CALC": "计算不符", "MISS": "缺失",
             "NA": "不适用", "UNK": "无法判断", "MANUAL": "需人工确认"}
FILL = {"DIFF": "FDE8E6", "CALC": "FDE8E6", "MISS": "FDF3DC", "NA": "F2F4F7",
        "UNK": "F3ECF8", "OK": "E9F6EE", "MANUAL": "E4EDF6"}

NAVY = "1F3A5F"
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
HEAD_FONT = Font(color="FFFFFF", bold=True, size=10.5)
THIN = Side(style="thin", color="D8DEE7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _sheet(wb, title, widths, header, rows, freeze="A2", status_col=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in range(1, len(header) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = CENTER
        cell.border = BORDER
    for row in rows:
        ws.append(row)
    for r in range(2, ws.max_row + 1):
        for c in range(1, len(header) + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            cell.alignment = WRAP
            cell.font = Font(size=10)
        if status_col:
            sv = ws.cell(row=r, column=status_col).value
            key = None
            for k, cn in STATUS_CN.items():
                if sv == cn:
                    key = k
            if key and FILL.get(key):
                ws.cell(row=r, column=status_col).fill = PatternFill("solid", fgColor=FILL[key])
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 28
    return ws


def build(result, out_path):
    wb = Workbook()
    wb.remove(wb.active)
    docs = result["docs"]
    live = [d for d in docs if not d.get("excluded")]
    doc_ids = [d["id"] for d in live]
    dn = {d["id"]: (d.get("type_cn") or d["type"]) for d in docs}

    # ---- 1. 单证索引 ----
    rows = []
    for d in docs:
        rows.append([
            d["id"], d.get("type", ""), d.get("type_cn", ""), d["filename"],
            d.get("pages", ""), len(d.get("fields") or {}),
            "参考文件（不参与核对）" if d.get("excluded") else "本票业务单证",
            d.get("exclude_reason", "") if d.get("excluded") else "",
        ])
    _sheet(wb, "单证索引", [10, 10, 20, 44, 8, 12, 22, 40],
           ["编号", "类型代码", "单证类型", "文件名", "页数", "已抽取字段数", "归类", "备注"],
           rows)

    # ---- 2. 跨文件字段矩阵 ----
    header = ["字段", "英文名", "分组", "字段性质"] + \
             ["%s\n%s" % (dn.get(i, i), i) for i in doc_ids]
    rows = []
    for row in result["matrix"]:
        line = [row["cn"], row.get("en", ""), row.get("group", ""),
                {"fact": "事实字段", "role": "角色字段", "calc": "派生计算"}.get(row.get("kind"), row.get("kind", ""))]
        for i in doc_ids:
            c = row["cells"].get(i) or {}
            st = c.get("status", "NA")
            v = c.get("value", "")
            line.append("%s：%s" % (STATUS_CN.get(st, st), v) if v else STATUS_CN.get(st, st))
        rows.append(line)
    _sheet(wb, "跨文件字段矩阵", [18, 20, 10, 11] + [26] * len(doc_ids), header, rows)

    # ---- 3. 计算核验 ----
    rows = []
    for c in result["checks"]:
        rows.append([c["id"], c["name"], c["detail"],
                     "通过" if c["ok"] else "不符",
                     ", ".join(dn.get(d, d) for d in dict.fromkeys(c.get("docs") or []))])
    _sheet(wb, "计算核验", [8, 34, 56, 10, 26],
           ["编号", "核验项", "计算说明", "结论", "涉及单证"], rows, status_col=4)

    # ---- 4. 差异定位 ----
    text_by = {d["id"]: d.get("text", "") for d in docs}

    def ctx(did, val):
        t = text_by.get(did, "")
        if not t or not val:
            return ""
        for n in (str(val), str(val)[:20], str(val)[:12]):
            if len(n) < 3:
                continue
            i = t.find(n)
            if i >= 0:
                a = max(0, i - 40)
                return t[a:i + len(n) + 40].replace("\n", " ⏎ ")
        return ""

    rows = []
    for it in result["issues"]:
        rows.append([
            it.get("id", ""), it.get("severity", ""), it.get("field_cn", ""),
            dn.get(it.get("doc"), it.get("doc", "")), it.get("status_cn", ""),
            it.get("value", ""), it.get("page") or "", it.get("status", ""),
            STATUS_CN.get(it.get("status", ""), it.get("status", "")),
            it.get("note", ""), ctx(it.get("doc"), it.get("value")),
        ])
    _sheet(wb, "差异定位", [7, 7, 16, 14, 10, 24, 7, 10, 10, 30, 52],
           ["编号", "级别", "字段", "单证", "页面", "当前值", "页码", "状态码", "状态",
            "系统判断", "原文上下文（定位用）"], rows, status_col=9)

    # ---- 5. 问题清单 ----
    rows = []
    for it in result["issues"]:
        peers = it.get("peers") or {}
        ref = "；".join("%s=%s" % (dn.get(k, k), v) for k, v in list(peers.items())[:3])
        rows.append([
            it.get("id", ""), it.get("severity", ""),
            {"S1": "阻断交单", "S2": "需修正", "S3": "建议核对"}.get(it.get("severity"), ""),
            it.get("field_cn", ""), dn.get(it.get("doc"), it.get("doc", "")),
            it.get("value", ""), ref, it.get("advice", "") or it.get("note", ""),
            ("第 %s 页" % it["page"]) if it.get("page") else "",
        ])
    _sheet(wb, "问题清单", [7, 7, 12, 16, 14, 22, 40, 56, 10],
           ["编号", "级别", "处理要求", "字段", "单证", "当前值", "其他单证取值（参考）",
            "问题说明与修改建议", "定位"], rows)

    # ---- 6. 修正记录 ----
    _sheet(wb, "修正记录", [6, 16, 14, 26, 26, 12, 20, 30],
           ["#", "字段", "单证", "原值", "修正值", "类型", "时间", "备注"],
           [], freeze="A2")
    ws = wb["修正记录"]
    ws.cell(row=2, column=1, value="（由检查页面「导出修正记录」回填；此处留空表示本次未做人工修正）")
    ws.cell(row=2, column=1).font = Font(italic=True, color="66718A", size=10)

    wb.save(out_path)
    return out_path


def main():
    base = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    src = sys.argv[1] if len(sys.argv) > 1 else os.path.join(base, "out", "check_result.json")
    dst = sys.argv[2] if len(sys.argv) > 2 else os.path.join(base, "out", "出口单证检查结果.xlsx")
    with io.open(src, encoding="utf-8") as f:
        result = json.load(f)
    build(result, dst)
    print("已生成: %s" % dst)


if __name__ == "__main__":
    main()
