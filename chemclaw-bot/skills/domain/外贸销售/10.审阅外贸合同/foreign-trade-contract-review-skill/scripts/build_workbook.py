"""Build the Excel review workbook (审阅清单).

Sheets:
    00 使用说明         Cover, scope, disclaimer, how to use
    01 审阅总览         KPI dashboard + chart data
    02 条款索引         Every article / sub-clause, with risk mapping
    03 问题清单         Main deliverable: one row per issue, verbatim quote
    04 确认事项         Exploded checklist (one row per confirmation item)
    05 缺失条款         Clauses absent from the contract
    06 原文摘录         Full contract text, block by block (citable)
    07 法律引用         Legal references with verification status
    08 依据与边界       Methodology, verification register, disclaimer

Usage:
    python build_workbook.py <contract_parsed.json> <data_meta.json> <issues-*.json> --out review.xlsx
"""
from __future__ import annotations

import argparse
import io
import json
import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# ---------------------------------------------------------------- palette
INK = "1F2937"
INK2 = "4B5563"
LINE = "D8DEE9"
HEAD_BG = "101828"
HEAD_FG = "FFFFFF"
ZEBRA = "F7F9FC"
BAND = "EEF2F9"

SEV_STYLE = {
    "high": ("DC2626", "FEE2E2"),
    "medium": ("D97706", "FEF3C7"),
    "low": ("059669", "D1FAE5"),
}

STATUS_LIST = ["待处理", "处理中", "需法务确认", "已解决", "不适用"]
STATUS_STYLE = {
    "待处理": ("9A3412", "FFF7ED"),
    "处理中": ("1E40AF", "EFF6FF"),
    "需法务确认": ("6D28D9", "F5F3FF"),
    "已解决": ("166534", "F0FDF4"),
    "不适用": ("475569", "F1F5F9"),
}

CAT_LABEL = {
    "payment": "付款",
    "delivery": "交付与运输",
    "quality": "质量与检验",
    "claim": "索赔与责任",
    "compliance": "合规与单证",
    "dispute": "争议解决与效力",
    "missing": "条款缺失",
}
TYPE_LABEL = {
    "ambiguous": "表述不清",
    "unilateral": "单方有利",
    "conflict": "内部冲突",
    "missing": "条款缺失",
    "unbalanced": "权利义务失衡",
    "unenforceable": "可能不可执行",
    "commercial": "商业不合理",
}
SEV_LABEL = {"high": "高风险", "medium": "中风险", "low": "低风险"}

thin = Side(style="thin", color=LINE)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


# ---------------------------------------------------------------- helpers
def load_issues(paths):
    out = []
    for p in paths:
        out.extend(json.load(io.open(p, encoding="utf-8")).get("issues", []))
    out.sort(key=lambda i: i.get("seq", 999))
    return out


def title_block(ws, title_cn, title_en, ncols, subtitle=""):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=title_cn)
    c.font = Font(name="微软雅黑", size=15, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=HEAD_BG)
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[1].height = 34

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    c = ws.cell(row=2, column=1, value=title_en)
    c.font = Font(name="Arial", size=9.5, color="9FB0CC", italic=True)
    c.fill = PatternFill("solid", fgColor=HEAD_BG)
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[2].height = 20

    r = 3
    if subtitle:
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=ncols)
        c = ws.cell(row=3, column=1, value=subtitle)
        c.font = Font(name="微软雅黑", size=9, color="92400E")
        c.fill = PatternFill("solid", fgColor="FFF8E6")
        c.alignment = Alignment(vertical="center", horizontal="left", indent=1, wrap_text=False)
        ws.row_dimensions[3].height = 22
        r = 4
    return r


def header_row(ws, row, headers, widths=None, height=30):
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = Font(name="微软雅黑", size=9.5, bold=True, color=HEAD_FG)
        c.fill = PatternFill("solid", fgColor="334155")
        c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[row].height = height
    if widths:
        for j, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w


def body_cell(c, wrap=True, size=9.5, bold=False, color=INK, valign="top", halign="left"):
    c.font = Font(name="微软雅黑", size=size, bold=bold, color=color)
    c.alignment = Alignment(vertical=valign, horizontal=halign, wrap_text=wrap)
    c.border = BORDER


def set_row_height(ws, row, text, width_chars, min_h=15, line_h=13.5):
    lines = 1
    for seg in str(text or "").split("\n"):
        lines += max(0, len(seg) * 2 // max(width_chars, 1))
    ws.row_dimensions[row].height = max(min_h, min(lines, 14) * line_h)


# ---------------------------------------------------------------- sheet: cover
def sheet_cover(ws, meta, issues):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 96

    title_block(
        ws,
        f"外贸合同审阅清单  |  {meta['title']}",
        "Foreign Trade Contract Review Checklist",
        3,
        "⚠ 模拟合同 SIMULATED — 本合同为教学/测试用虚构文本，不构成真实商业文件。",
    )

    rows = [
        ("合同编号 Contract No.", meta["contract_no"]),
        ("合同日期 Date", meta["contract_date"]),
        ("我方身份 Our role", f"{meta['our_role_cn']}（seller）— 中国出口方"),
        ("交易对手 Counterparty", "买方（buyer）— 印度尼西亚"),
        ("产品 Product", f"{meta['product_cn']}  |  {meta['product']}"),
        ("HS 编码", meta["hs_code"]),
        ("贸易术语 Trade term", meta["trade_term"]),
        ("合同金额 Amount", meta["amount"]),
        ("数量 Quantity", meta["quantity"]),
        ("审阅日期 Review date", meta["review_date"]),
        ("审阅人 Reviewer", meta["reviewer"]),
        ("问题总数 Issues", f"{len(issues)} 项"),
    ]
    r = 5
    for k, v in rows:
        c = ws.cell(row=r, column=2, value=k)
        body_cell(c, bold=True, color=INK2, size=9.5)
        c.fill = PatternFill("solid", fgColor=BAND)
        c = ws.cell(row=r, column=3, value=v)
        body_cell(c)
        r += 1

    r += 1
    c = ws.cell(row=r, column=2, value="使用说明 / How to use")
    c.font = Font(name="微软雅黑", size=11, bold=True, color=HEAD_BG)
    r += 1

    steps = [
        ("01 审阅总览", "先看风险分布与高优先级事项，确定审阅重点。"),
        ("02 条款索引", "合同每一章、每一条的位置与对应风险，用于快速定位。"),
        ("03 问题清单", "逐条问题：原文引用、问题类型、潜在后果、确认事项、建议措辞。"
                     "「负责人 / 处理状态 / 备注」三列为填写区，已设置下拉与筛选。"),
        ("04 确认事项", "把所有「需要进一步确认的事项」拆成可勾选的待办，便于分工跟进。"),
        ("05 缺失条款", "合同中未出现、但外贸履约通常需要约定的条款。缺失 ≠ 违法，但意味着无约定可依。"),
        ("06 原文摘录", "合同全文按条款块编号摘录，供引用与比对。"),
        ("07 法律引用", "本报告引用的规则与来源，并标注核实状态（已核实 / 待核实）。"),
        ("08 依据与边界", "审阅方法、核实登记，以及业务审阅与法律意见的边界声明。"),
    ]
    for k, v in steps:
        c = ws.cell(row=r, column=2, value=k)
        body_cell(c, bold=True, size=9.5, color="1D4ED8")
        c = ws.cell(row=r, column=3, value=v)
        body_cell(c)
        ws.row_dimensions[r].height = 30
        r += 1

    r += 1
    warn = (
        "边界声明：本清单为业务风险提醒，不构成法律意见。清单仅基于所提供合同文本及附件作出，"
        "未核验交易对手资信、货物实际品质或目的国监管现状；涉及印度尼西亚法的事项一律标注为待当地律师确认。"
        "任何法律适用结论须经具备中国法与印度尼西亚法执业资格的律师复核，方可用于谈判或争议处理。"
    )
    ws.merge_cells(start_row=r, start_column=2, end_row=r + 3, end_column=3)
    c = ws.cell(row=r, column=1 + 1, value=warn)
    c.font = Font(name="微软雅黑", size=9.5, color="7C2D12")
    c.fill = PatternFill("solid", fgColor="FEF3C7")
    c.alignment = Alignment(vertical="top", horizontal="left", wrap_text=True, indent=1)
    c.border = BORDER
    return r


# ---------------------------------------------------------------- sheet: overview
def sheet_overview(ws, meta, issues):
    ws.sheet_view.showGridLines = False
    ncols = 5
    title_block(ws, "审阅总览 | 风险分布", "Review Overview — risk distribution", ncols)
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 46

    counts = {"high": 0, "medium": 0, "low": 0}
    for i in issues:
        counts[i["severity"]] = counts.get(i["severity"], 0) + 1

    r = 5
    c = ws.cell(row=r, column=2, value="按风险等级")
    c.font = Font(name="微软雅黑", size=10, bold=True)

    r = 6
    for h, col in (("等级", 2), ("数量", 3), ("占比", 4)):
        c = ws.cell(row=r, column=col, value=h)
        body_cell(c, bold=True, halign="center", size=9)
        c.fill = PatternFill("solid", fgColor=BAND)

    data_start = r + 1
    for k in ("high", "medium", "low"):
        r += 1
        fg, bg = SEV_STYLE[k]
        c = ws.cell(row=r, column=2, value=SEV_LABEL[k])
        body_cell(c, bold=True, color=fg)
        c.fill = PatternFill("solid", fgColor=bg)
        c = ws.cell(row=r, column=3, value=counts[k])
        body_cell(c, halign="center")
        c = ws.cell(row=r, column=4, value=counts[k] / max(len(issues), 1))
        body_cell(c, halign="center")
        c.number_format = "0%"
    data_end = r

    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = "风险等级分布 / Issues by severity"
    chart.y_axis.title = None
    chart.x_axis.title = None
    data = Reference(ws, min_col=3, min_row=data_start - 1, max_row=data_end)
    cats = Reference(ws, min_col=2, min_row=data_start, max_row=data_end)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height = 6.5
    chart.width = 14
    chart.legend = None
    ws.add_chart(chart, "F5")

    # category table
    r = data_end + 3
    c = ws.cell(row=r, column=2, value="按主题")
    c.font = Font(name="微软雅黑", size=10, bold=True)

    r += 1
    for h, col in (("主题", 2), ("数量", 3), ("高风险", 4), ("说明", 5)):
        c = ws.cell(row=r, column=col, value=h)
        body_cell(c, bold=True, halign="center", size=9)
        c.fill = PatternFill("solid", fgColor=BAND)

    cat_note = {
        "payment": "货款能否收回、付款条件是否可执行",
        "delivery": "交货、装运、风险与运输责任",
        "quality": "质量标准是否可客观判定、检验权是否对等",
        "claim": "索赔程序、举证与责任限额",
        "compliance": "目的国准入、单证与产品合规",
        "dispute": "适用法律、仲裁与执行",
        "missing": "应有未有的条款",
    }
    for key, label in CAT_LABEL.items():
        n = sum(1 for i in issues if i["category"] == key)
        if not n:
            continue
        r += 1
        h = sum(1 for i in issues if i["category"] == key and i["severity"] == "high")
        body_cell(ws.cell(row=r, column=2, value=label), bold=True)
        body_cell(ws.cell(row=r, column=3, value=n), halign="center")
        body_cell(ws.cell(row=r, column=4, value=h), halign="center",
                  color=SEV_STYLE["high"][0], bold=True)
        body_cell(ws.cell(row=r, column=5, value=cat_note.get(key, "")), color=INK2)

    # high priority short list
    r += 3
    c = ws.cell(row=r, column=2, value="高优先级事项（高风险 + 条款缺失）")
    c.font = Font(name="微软雅黑", size=10, bold=True)
    r += 1
    for h, col in (("编号", 2), ("条款", 3), ("问题", 5)):
        c = ws.cell(row=r, column=col, value=h)
        body_cell(c, bold=True, halign="center", size=9)
        c.fill = PatternFill("solid", fgColor=BAND)
    for i in issues:
        if i["severity"] != "high" and i["category"] != "missing":
            continue
        r += 1
        body_cell(ws.cell(row=r, column=2, value=i["id"]), halign="center", bold=True,
                  color=SEV_STYLE[i["severity"]][0])
        body_cell(ws.cell(row=r, column=3, value=i["clause_title"]), size=8.5, color=INK2)
        c = ws.cell(row=r, column=5, value=i["title"])
        body_cell(c)
        ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=5)
        set_row_height(ws, r, i["title"], 46)
    return r


# ---------------------------------------------------------------- sheet: clause index
def sheet_index(ws, parsed, issues):
    ws.sheet_view.showGridLines = False
    headers = ["条款位置", "层级", "章节 / 条款标题", "块编号", "涉及风险", "风险等级", "条数"]
    widths = [30, 10, 58, 46, 26, 12, 8]
    title_block(ws, "条款索引 | 合同结构与风险映射", "Clause Index — structure and risk mapping", len(headers))
    ws.freeze_panes = "A5"
    header_row(ws, 4, headers, widths)

    # map block -> issues
    by_block = {}
    for i in issues:
        for a in i.get("anchors", []):
            by_block.setdefault(a["block"], []).append(i)

    r = 5
    for art in parsed["index"]:
        sev = [i["severity"] for b in art["blocks"] for i in by_block.get(b, [])]
        ids = sorted({i["id"] for b in art["blocks"] for i in by_block.get(b, [])})
        worst = "high" if "high" in sev else ("medium" if "medium" in sev else ("low" if "low" in sev else ""))
        vals = [art["title"].split("—")[0].strip(), "章", art["title"], art["id"],
                ", ".join(ids), SEV_LABEL.get(worst, ""), len(ids) or ""]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=v)
            body_cell(c, bold=(j == 3), size=9)
            c.fill = PatternFill("solid", fgColor=ZEBRA)
        if worst:
            c = ws.cell(row=r, column=6)
            c.font = Font(name="微软雅黑", size=9, bold=True, color=SEV_STYLE[worst][0])
            c.fill = PatternFill("solid", fgColor=SEV_STYLE[worst][1])
        ws.row_dimensions[r].height = 20
        r += 1

        for cl in art["clauses"]:
            sev = [i["severity"] for b in cl["blocks"] for i in by_block.get(b, [])]
            ids = sorted({i["id"] for b in cl["blocks"] for i in by_block.get(b, [])})
            worst = "high" if "high" in sev else ("medium" if "medium" in sev else ("low" if "low" in sev else ""))
            vals = ["", "条", cl["title"], cl["id"], ", ".join(ids),
                    SEV_LABEL.get(worst, ""), len(ids) or ""]
            for j, v in enumerate(vals, start=1):
                c = ws.cell(row=r, column=j, value=v)
                body_cell(c, size=9)
            if worst:
                c = ws.cell(row=r, column=6)
                c.font = Font(name="微软雅黑", size=9, bold=True, color=SEV_STYLE[worst][0])
                c.fill = PatternFill("solid", fgColor=SEV_STYLE[worst][1])
            ws.row_dimensions[r].height = 18
            r += 1
    ws.auto_filter.ref = f"A4:G{r - 1}"
    return r


# ---------------------------------------------------------------- sheet: issues
def sheet_issues(ws, issues):
    ws.sheet_view.showGridLines = False
    headers = [
        "编号", "问题分类", "风险等级", "问题类型", "条款",
        "合同原文位置", "合同原文引用",
        "问题说明", "潜在商务后果", "需要确认事项", "建议修改措辞",
        "负责人", "处理状态", "备注",
    ]
    widths = [8, 12, 9, 11, 20, 26, 52, 52, 48, 42, 48, 11, 12, 26]
    title_block(
        ws, "问题清单 | 合同原文 · 问题类型 · 商务后果 · 确认事项",
        "Issue Register — quote, issue type, consequence, confirmations",
        len(headers),
    )
    ws.freeze_panes = "F5"
    header_row(ws, 4, headers, widths, height=34)

    r = 5
    for i in issues:
        anchors = i.get("anchors", [])
        locs = "\n".join(a.get("clause", "") for a in anchors) or "（条款缺失）"
        quotes = "\n\n".join(
            f"[{a.get('clause','')}] \"{a.get('quote','')}\"" for a in anchors
        ) or "（该事项为条款缺失，合同中无对应原文）"
        confirms = "\n".join(f"{n}. {c}" for n, c in enumerate(i.get("confirm", []), 1))

        vals = [
            i["id"], CAT_LABEL.get(i["category"], i["category"]),
            SEV_LABEL.get(i["severity"], i["severity"]),
            TYPE_LABEL.get(i["issue_type"], i["issue_type"]),
            i.get("clause_title", ""), locs, quotes,
            i.get("finding", ""), i.get("consequence", ""), confirms,
            i.get("suggested_wording", ""),
            i.get("owner", ""), i.get("status", "待处理"), i.get("note", ""),
        ]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=v)
            body_cell(c, size=9)
        # severity colour
        c = ws.cell(row=r, column=3)
        fg, bg = SEV_STYLE[i["severity"]]
        c.font = Font(name="微软雅黑", size=9, bold=True, color=fg)
        c.fill = PatternFill("solid", fgColor=bg)
        c.alignment = Alignment(horizontal="center", vertical="center")
        # id
        ws.cell(row=r, column=1).font = Font(name="Consolas", size=9, bold=True, color=INK2)
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="center", vertical="top")
        # quote block: serif italic feel
        q = ws.cell(row=r, column=7)
        q.font = Font(name="Georgia", size=9, italic=True, color="1E293B")
        # status styling
        st = i.get("status", "待处理")
        sc = ws.cell(row=r, column=13)
        sfg, sbg = STATUS_STYLE.get(st, STATUS_STYLE["待处理"])
        sc.font = Font(name="微软雅黑", size=9, bold=True, color=sfg)
        sc.fill = PatternFill("solid", fgColor=sbg)
        sc.alignment = Alignment(horizontal="center", vertical="center")

        # height driven by the longest wrapped column
        est = max(
            len(quotes) / 46, len(i.get("finding", "")) / 46,
            len(i.get("consequence", "")) / 42, len(confirms) / 38,
            len(i.get("suggested_wording", "")) / 42, 1,
        )
        ws.row_dimensions[r].height = min(max(est * 13.5, 40), 320)
        r += 1

    last = r - 1
    ws.auto_filter.ref = f"A4:N{last}"

    # data validation on status
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUS_LIST) + '"', allow_blank=True)
    dv.error = "请从下拉列表中选择处理状态"
    dv.prompt = "选择处理状态"
    ws.add_data_validation(dv)
    dv.add(f"M5:M{max(last, 5) + 200}")
    return r


# ---------------------------------------------------------------- sheet: confirmations
def sheet_confirm(ws, issues):
    ws.sheet_view.showGridLines = False
    headers = ["序号", "关联编号", "条款", "需要确认的事项", "建议负责部门", "确认结果", "完成", "备注"]
    widths = [7, 10, 20, 76, 14, 30, 8, 24]
    title_block(
        ws, "确认事项清单 | 可直接分工跟进",
        "Confirmation Checklist — assignable action items",
        len(headers),
        "提示：本表把「问题清单」中的确认事项逐条拆开，便于按部门分工。“建议负责部门”为常见分工建议，需按贵司实际调整。",
    )
    ws.freeze_panes = "A5"
    header_row(ws, 4, headers, widths)

    dept_guess = {
        "owner": "业务/销售",
        "payment": "财务/单证",
        "delivery": "物流/单证",
        "quality": "品控/技术",
        "claim": "法务/业务",
        "compliance": "单证/合规",
        "dispute": "法务",
        "missing": "法务/业务",
    }

    r = 5
    n = 0
    for i in issues:
        for c in i.get("confirm", []):
            n += 1
            vals = [n, i["id"], i.get("clause_title", ""), c, dept_guess.get(i["category"], ""), "", "☐", ""]
            for j, v in enumerate(vals, start=1):
                cell = ws.cell(row=r, column=j, value=v)
                body_cell(cell, size=9)
            ws.cell(row=r, column=1).alignment = Alignment(horizontal="center", vertical="top")
            ws.cell(row=r, column=2).alignment = Alignment(horizontal="center", vertical="top")
            ws.cell(row=r, column=2).font = Font(name="Consolas", size=9, bold=True, color=INK2)
            ws.cell(row=r, column=7).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row=r, column=7).font = Font(name="微软雅黑", size=12, color=INK2)
            set_row_height(ws, r, c, 76)
            r += 1

    last = r - 1
    ws.auto_filter.ref = f"A4:H{last}"
    dv = DataValidation(type="list", formula1='"☐,☑,不适用"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"G5:G{last}")
    return r


# ---------------------------------------------------------------- sheet: missing
def sheet_missing(ws, issues):
    ws.sheet_view.showGridLines = False
    headers = ["编号", "主题", "缺失/薄弱条款", "合同中的相关原文", "为什么要约定", "潜在后果", "建议补充内容"]
    widths = [8, 12, 40, 50, 44, 48, 50]
    title_block(
        ws, "缺失与薄弱条款清单",
        "Missing / Weak Clause Register",
        len(headers),
        "说明：标为「缺失」表示合同文本中未出现该主题的约定，不代表违法。未约定的事项在争议中将回到法律补充规则，结果不可控。",
    )
    ws.freeze_panes = "A5"
    header_row(ws, 4, headers, widths)

    r = 5
    for i in issues:
        if i["category"] != "missing" and not i["title"].startswith("【条款缺失】"):
            continue
        quotes = "\n\n".join(
            f"[{a.get('clause','')}] \"{a.get('quote','')}\"" for a in i.get("anchors", [])
        ) or "（合同全文未出现相关约定）"
        vals = [
            i["id"], CAT_LABEL.get(i["category"], ""), i["title"],
            quotes, i.get("finding", ""), i.get("consequence", ""),
            i.get("suggested_wording", ""),
        ]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=v)
            body_cell(c, size=9)
        ws.cell(row=r, column=1).font = Font(name="Consolas", size=9, bold=True, color=INK2)
        ws.cell(row=r, column=4).font = Font(name="Georgia", size=9, italic=True, color="1E293B")
        est = max(len(quotes) / 44, len(i.get("finding", "")) / 40,
                  len(i.get("consequence", "")) / 44, len(i.get("suggested_wording", "")) / 46, 1)
        ws.row_dimensions[r].height = min(max(est * 13.5, 46), 300)
        r += 1
    ws.auto_filter.ref = f"A4:G{r - 1}"
    return r


# ---------------------------------------------------------------- sheet: source text
def sheet_source(ws, parsed):
    ws.sheet_view.showGridLines = False
    headers = ["块编号", "类型", "所属条款", "原文"]
    widths = [46, 12, 34, 108]
    title_block(
        ws, "合同原文摘录 | 按块编号",
        "Contract Text Extract — by block id",
        len(headers),
        "说明：块编号与 HTML 报告中的定位锚点一一对应，可用于核对任何一条风险提示的原文依据。",
    )
    ws.freeze_panes = "A5"
    header_row(ws, 4, headers, widths)

    kind_cn = {
        "h1": "章", "annex-title": "附件", "h2": "条", "p": "段",
        "exec-col": "签署栏", "exec-line": "签署行", "table": "规格表",
    }
    r = 5
    for b in parsed["blocks"]:
        if b["kind"] == "table":
            text = "\n".join(" | ".join(row) for row in b["rows"])
        elif b["kind"] == "exec-col":
            text = "\n".join(ln["text"] for ln in b.get("lines", []))
        else:
            text = b.get("text", "")
        vals = [b["id"], kind_cn.get(b["kind"], b["kind"]), b.get("article", ""), text]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=v)
            body_cell(c, size=9)
        ws.cell(row=r, column=1).font = Font(name="Consolas", size=8.5, color=INK2)
        if b["kind"] in ("h1", "annex-title"):
            for j in range(1, 5):
                cc = ws.cell(row=r, column=j)
                cc.font = Font(name="微软雅黑", size=9, bold=True, color="FFFFFF")
                cc.fill = PatternFill("solid", fgColor="475569")
        ws.row_dimensions[r].height = min(max(len(text) / 100 * 14, 16), 120)
        r += 1
    ws.auto_filter.ref = f"A4:D{r - 1}"
    return r


# ---------------------------------------------------------------- sheet: legal refs
def sheet_legal(ws, issues):
    ws.sheet_view.showGridLines = False
    headers = ["关联编号", "条款", "规则 / 依据", "说明", "核实状态"]
    widths = [10, 20, 50, 76, 30]
    title_block(
        ws, "法律与规则引用 | 含核实状态",
        "Legal & Rule References — with verification status",
        len(headers),
        "说明：凡涉及印度尼西亚法的事项，均需由印尼执业律师复核。「待核实」表示本报告未取得一手文本，不得作为确定性依据引用。",
    )
    ws.freeze_panes = "A5"
    header_row(ws, 4, headers, widths)

    verified_markers = ("UNCITRAL", "lartas", "Permendag", "纽约公约", "纽约公约", "SEMA", "PERMA",
                        "Nine AM", "UCP", "Incoterms")
    r = 5
    for i in issues:
        for ref in i.get("legal_refs", []):
            label = ref.get("label", "")
            status = "已核实来源" if any(m in label for m in verified_markers) else "需律师复核"
            vals = [i["id"], i.get("clause_title", ""), label, ref.get("note", ""), status]
            for j, v in enumerate(vals, start=1):
                c = ws.cell(row=r, column=j, value=v)
                body_cell(c, size=9)
            ws.cell(row=r, column=1).font = Font(name="Consolas", size=9, bold=True, color=INK2)
            c = ws.cell(row=r, column=5)
            if status == "已核实来源":
                c.fill = PatternFill("solid", fgColor="D1FAE5")
                c.font = Font(name="微软雅黑", size=9, color="065F46", bold=True)
            else:
                c.fill = PatternFill("solid", fgColor="FEF3C7")
                c.font = Font(name="微软雅黑", size=9, color="92400E", bold=True)
            c.alignment = Alignment(horizontal="center", vertical="center")
            set_row_height(ws, r, ref.get("note", ""), 76)
            r += 1
    ws.auto_filter.ref = f"A4:E{r - 1}"
    return r


# ---------------------------------------------------------------- sheet: method
def sheet_method(ws, meta, issues):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 92

    title_block(ws, "审阅依据与边界", "Methodology, verification register and limitations", 3)

    r = 5
    blocks = [
        ("一、审阅范围", meta["review_scope"]),
        ("二、审阅方法",
         "1) 先识别合同主体、产品、交易条件与条款结构，形成条款索引；\n"
         "2) 逐条比对合同内部一致性（同一事项在多处出现时是否冲突）；\n"
         "3) 将我方身份（卖方/买方）代入，判断该条款使谁承担义务与风险；\n"
         "4) 对缺失的常见条款逐项检核，明确标注为「缺失」而非推定其存在；\n"
         "5) 对涉及法律适用的事项，检索可核实的规则来源并标注核实状态；\n"
         "6) 每条问题必须锚定合同原文位置，并给出可执行的确认事项。"),
        ("三、本报告不做的事",
         "· 不对合同效力、条款有效性作出确定性法律结论；\n"
         "· 不推定合同中不存在的内容；\n"
         "· 不将正常商业安排判定为违法（例如较高的违约金、特定仲裁地，本身均属商业谈判结果）；\n"
         "· 不核验交易对手资信、货物实测品质、目的国监管现状；\n"
         "· 不提供税务、外汇、海关归类结论。"),
        ("四、需要特别提示的一点",
         "本样例中，印度尼西亚并非《联合国国际货物销售合同公约》（CISG）缔约国"
         "（据 UNCITRAL 官方缔约国表），因此 CISG 本不适用于本合同，"
         "第 15.1 条的「排除 CISG」表述属于巩固性安排，真正起作用的是「适用中国法」与「仲裁地/执行」的选择。"),
        ("五、免责声明", meta["disclaimer"]),
    ]

    for head, body in blocks:
        c = ws.cell(row=r, column=2, value=head)
        c.font = Font(name="微软雅黑", size=10.5, bold=True, color=HEAD_BG)
        c.alignment = Alignment(vertical="top")
        r += 1
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        c = ws.cell(row=r, column=2, value=body)
        c.font = Font(name="微软雅黑", size=9.5, color="374151")
        c.alignment = Alignment(vertical="top", wrap_text=True, indent=1)
        c.border = BORDER
        lines = body.count("\n") + 1 + len(body) // 90
        ws.row_dimensions[r].height = max(30, lines * 15)
        r += 2

    # verification register
    c = ws.cell(row=r, column=2, value="六、核实登记")
    c.font = Font(name="微软雅黑", size=10.5, bold=True, color=HEAD_BG)
    r += 1
    for j, h in enumerate(["事项", "结论", "来源"], start=2):
        cc = ws.cell(row=r, column=j, value=h)
        body_cell(cc, bold=True, size=9, halign="center")
        cc.fill = PatternFill("solid", fgColor=BAND)
    r += 1

    reg = [
        ("中国为 CISG 缔约国", "已核实：1986-12-11 批准，1988-01-01 生效，附第 1(1)(b) 条保留声明",
         "UNCITRAL CISG Status 官方页面"),
        ("印度尼西亚 CISG 缔约国身份", "已核实：未出现在 97 个缔约方名单中（非缔约国）",
         "UNCITRAL CISG Status 官方页面"),
        ("CIF 保险最低要求", "已核实：ICC Institute Cargo Clauses (C)；ICC 官方说明 CIP 已提高为 (A)，CIF 仍为 (C)",
         "ICC Digital Library（Incoterms 2020 规则文本）"),
        ("CIF 风险转移点", "已核实：货物在装运港装上船时风险转移（与 FOB 相同），但卖方仍须订舱至目的港",
         "ICC Digital Library（Incoterms 2020 规则文本）"),
        ("印尼合同语言要求（现行立场）", "已核实：SEMA No. 3/2023 规定缺印尼语译本本身不得作为撤销理由，恶意除外",
         "印尼最高法院 SEMA 3/2023；Nine AM 案 Putusan 601 K/Pdt/2015"),
        ("印中纽约公约关系", "已核实：两国均为缔约国（印尼 1981-10-07、中国 1987-01-22），均作互惠+商事保留",
         "newyorkconvention.org 缔约国页"),
        ("外国仲裁裁决在印尼执行程序", "已核实：PERMA No. 3/2023（2023-10-17 生效）",
         "印尼最高法院 JDIH"),
        ("纯碱 HS 2836 是否在印尼限制进口清单", "已核实（负向）：Permendag 22/2025 与 20/2025 封闭清单均未列入 2836",
         "peraturan.bpk.go.id"),
        ("纯碱是否需强制 SNI", "未核实：未查得针对 2836.20 的强制 SNI，但未能取得官方强制名录，不得据此认定“不存在”",
         "—"),
        ("印尼对华纯碱反倾销措施", "未核实：KADI 公开页面未见相关案件，属证据缺失而非证明不存在",
         "KADI 官网（未取得完整措施清单）"),
        ("ACFTA 3.0 生效状态", "未核实：截至检索时点未见生效确认；现行 Form E 仍为可操作路径",
         "越南工贸部公告 / 泰媒报道（二手）"),
        ("中国—东盟协定税率与 MFN 税率差额", "未核实：需按印尼海关税则具体税目核算", "—"),
    ]
    for a, b, c_, in reg:
        body_cell(ws.cell(row=r, column=2, value=a), size=9, bold=True)
        cc = ws.cell(row=r, column=3, value=b)
        body_cell(cc, size=9)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=3)
        # put source on next line inside the same cell to keep 3 columns
        cc.value = f"{b}\n来源：{c_}"
        ws.row_dimensions[r].height = max(28, (len(b) // 88 + 2) * 13.5)
        r += 1
    return r


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("parsed")
    ap.add_argument("meta")
    ap.add_argument("issues", nargs="+")
    ap.add_argument("--out", required=True)
    ap.add_argument("--stamp", default="")
    args = ap.parse_args()

    parsed = json.load(io.open(args.parsed, encoding="utf-8"))
    meta = json.load(io.open(args.meta, encoding="utf-8"))["meta"]
    issues = load_issues(args.issues)

    wb = Workbook()
    ws = wb.active
    ws.title = "00 使用说明"
    sheet_cover(ws, meta, issues)

    sheet_overview(wb.create_sheet("01 审阅总览"), meta, issues)
    sheet_index(wb.create_sheet("02 条款索引"), parsed, issues)
    sheet_issues(wb.create_sheet("03 问题清单"), issues)
    sheet_confirm(wb.create_sheet("04 确认事项"), issues)
    sheet_missing(wb.create_sheet("05 缺失条款"), issues)
    sheet_source(wb.create_sheet("06 原文摘录"), parsed)
    sheet_legal(wb.create_sheet("07 法律引用"), issues)
    sheet_method(wb.create_sheet("08 依据与边界"), meta, issues)

    wb.properties.title = "外贸合同审阅清单 — 模拟合同"
    wb.properties.subject = meta["contract_no"]
    wb.properties.creator = "Foreign Trade Contract Review Skill"
    wb.properties.description = "SIMULATED — business risk flags, not legal advice."

    wb.save(args.out)
    print(f"wrote {args.out}  ({len(issues)} issues, {len(wb.sheetnames)} sheets: {', '.join(wb.sheetnames)})")


if __name__ == "__main__":
    main()
