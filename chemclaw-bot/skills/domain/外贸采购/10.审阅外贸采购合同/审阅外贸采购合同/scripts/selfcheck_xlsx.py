"""Excel 工作簿结构自检：读取生成的 xlsx，逐表核对行数与关键字段。"""
import sys

from openpyxl import load_workbook

p = sys.argv[1]
sys.stdout.reconfigure(encoding="utf-8")
wb = load_workbook(p)
print("sheets:", wb.sheetnames)
for ws in wb.worksheets:
    print(f"  {ws.title}: dims={ws.dimensions} max_row={ws.max_row} max_col={ws.max_column} freeze={ws.freeze_panes} filter={ws.auto_filter.ref}")

# 02 总表：检查关键列非空
ws = wb["02 风险事项总表"]
hdr = [c.value for c in ws[3]]
print("\n02 headers:", hdr)
idx = {h: i + 1 for i, h in enumerate(hdr)}
empty = []
for r in range(4, ws.max_row + 1):
    for key in ("风险编号", "风险类型", "严重度", "风险标题", "合同原文（逐字摘录）", "合同写了什么（中文解释）", "对中国买方的商务后果", "需确认事项（分条）"):
        if not ws.cell(row=r, column=idx[key]).value:
            empty.append((r, key))
print("02 empty required cells:", empty[:10], "count", len(empty))

# 06 中英对照：英文列非空
ws6 = wb["06 中英对照问题清单"]
h6 = [c.value for c in ws6[3]]
i6 = {h: i + 1 for i, h in enumerate(h6)}
miss = [r for r in range(4, ws6.max_row + 1) if not ws6.cell(row=r, column=i6["Question to Supplier (English)"]).value]
print("06 rows:", ws6.max_row - 3, "missing EN:", len(miss))

# 04 法规
ws4 = wb["04 法规核对"]
print("04 law rows:", ws4.max_row - 3)
for r in range(4, ws4.max_row + 1):
    print("   -", ws4.cell(row=r, column=2).value, "|", ws4.cell(row=r, column=4).value)

# 08 统计
ws8 = wb["08 统计"]
for r in range(3, ws8.max_row + 1):
    a, b = ws8.cell(row=r, column=1).value, ws8.cell(row=r, column=2).value
    if a:
        print("08:", a, "=", (str(b)[:60] if b is not None else ""))
