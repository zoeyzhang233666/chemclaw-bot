"""把案例 JSON 渲染为「逐条对照」的双语合同 DOCX。

用途：当用户没有真实合同 PDF 时，产出一份**明确标注「模拟合同」**的完整合同中英对照本，
作为审查对象与示例附件。真实合同请直接使用原件，不要用本脚本改写正文。

版式（逐条对照，条款序号相同）：
    P2  第三条  3.1 Specifications
    EN │ 3.1 Ethylene glycol (...)...          ← 英文原文（逐字）
    ZH │ 3.1 乙二醇（...）...                    ← 中文对照
    风险标签：R08（由 review.json 自动标注，可选 --with-tags）

用法:
    python build_contract_docx.py --dir <案例目录> --out <合同.docx> [--with-tags]
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from common import load_case

LATIN = "Times New Roman"
CJK = "微软雅黑"
INK = "0F3557"


def set_font(run, name=LATIN, size=10.5, bold=False, color=None, cjk=CJK, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), cjk)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)


def shade(cell, fill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def para_shade(p, fill: str):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def add_page_field(paragraph):
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    paragraph._p.append(fld)


def build(case_dir: Path, out: Path, with_tags: bool = False) -> dict:
    case = load_case(case_dir)
    c = case["contract"]
    meta = c.get("contract_meta", {})
    bilingual = c.get("language", "en-zh") == "en-zh"
    tags: dict[str, list[str]] = {}
    if with_tags:
        for f in case["review"].get("findings", []):
            for cid in f["clause_ids"]:
                tags.setdefault(cid, []).append(f["id"])

    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
    sec.left_margin, sec.right_margin = Cm(2.4), Cm(2.2)

    style = doc.styles["Normal"]
    style.font.name = LATIN
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), CJK)
    style.paragraph_format.line_spacing = 1.32
    style.paragraph_format.space_after = Pt(3)

    # 页脚
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("GY-PO-2026-MEG-017 · 模拟合同（Simulated Contract） · 第 ")
    set_font(r, size=8, color="94A3B8")
    add_page_field(fp)
    r = fp.add_run(" 页")
    set_font(r, size=8, color="94A3B8")

    # 模拟标记
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para_shade(p, "FEF3C7")
    r = p.add_run("【模拟合同 SIMULATED CONTRACT】本文件中英对照本为外贸采购合同审查 Skill 的演示材料："
                  "双方名称、地址、金额与条款均为编写，不对应任何真实交易或主体，不得用于真实签约或对外引用。")
    set_font(r, size=8.5, bold=True, color="78350F")

    # 标题
    for txt, size, bold, color in (
        ("PURCHASE CONTRACT", 20, True, INK),
        (meta.get("title_en", "Purchase Contract for Industrial Grade Monoethylene Glycol (MEG)"), 10.5, False, "334155"),
        (meta.get("title", ""), 11, True, INK),
        (f"合同编号 Contract No.: {meta.get('contract_no', '')}", 9.5, False, "475569"),
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(txt)
        set_font(r, size=size, bold=bold, color=color)

    doc.add_paragraph()

    clauses = {x["id"]: x for x in c["clauses"]}

    def emit_bilingual(idx_no, en_text, zh_text, kind=""):
        """no | EN | ZH 三列单行表格（逐条对照）。"""
        t = doc.add_table(rows=1, cols=3)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        t.style = "Table Grid"
        w = (Cm(2.2), Cm(8.1), Cm(6.3))
        cells = t.rows[0].cells
        for i, wd in enumerate(w):
            cells[i].width = wd
        set_font(cells[0].paragraphs[0].add_run(str(idx_no)), size=8, bold=True, color=INK)
        shade(cells[0], "F1F5F9")
        pe = cells[1].paragraphs[0]
        set_font(pe.add_run(en_text), size=10, name=LATIN)
        pz = cells[2].paragraphs[0]
        set_font(pz.add_run(zh_text), size=9.2, color="1F2937")
        for i in (1, 2):
            cells[i].paragraphs[0].paragraph_format.space_after = Pt(0)
        if kind == "table_row":
            shade(cells[1], "FAFAF9")
            shade(cells[2], "FAFAF9")

    for s in c["sections"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        para_shade(p, "E8EFF5")
        r = p.add_run(f'{s["no"]}　{s["title"]}')
        set_font(r, size=11.5, bold=True, color=INK)
        r = p.add_run(f"　　[P{s['page']}]")
        set_font(r, size=8, color="94A3B8")
        for cid in s["clause_ids"]:
            cl = clauses[cid]
            no = cl["no"].split(" ")[0] if cl["no"] else cid
            emit_bilingual(no, cl["text"], cl.get("zh", ""), cl.get("kind", ""))
            if with_tags and tags.get(cid):
                tp = doc.add_paragraph()
                tp.paragraph_format.left_indent = Cm(2.3)
                tp.paragraph_format.space_after = Pt(2)
                r = tp.add_run("审查标注：" + "、".join(tags[cid]))
                set_font(r, size=7.5, color="B91C1C")

    # 签署
    doc.add_paragraph()
    sig = doc.add_table(rows=1, cols=2)
    sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    left, right = sig.rows[0].cells
    for cell, who, name, addr in (
        (left, "SELLER 卖方", "GULF PETROCHEMICAL TRADING FZE", "Office 1203, Tower B, Jebel Ali Free Zone, Dubai, U.A.E."),
        (right, "BUYER 买方", "ZHEJIANG HONYU NEW MATERIALS CO., LTD. 浙江宏宇新材料有限公司", "No. 88 Huancheng East Road, Hangzhou, Zhejiang, China"),
    ):
        set_font(cell.paragraphs[0].add_run(who), size=9, bold=True, color=INK)
        set_font(cell.add_paragraph().add_run(name), size=9)
        set_font(cell.add_paragraph().add_run(addr), size=8, color="475569")
        for line in ("Authorised Signature 授权签字：__________________",
                     "Name / Title 姓名职务：__________________",
                     "Date 日期：__________________"):
            set_font(cell.add_paragraph().add_run(line), size=9)
        shade(cell, "F8FAFC")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run("注：本文件为模拟合同，仅用于外贸采购合同审查方法演示与 Skill 测试；"
                  "其中刻意保留的条款问题（如预付款无担保、异议期过短、尾款触发条件冲突、"
                  "仲裁地矛盾等）用于考察审查能力，不代表任何真实交易的约定。")
    set_font(r, size=8, color="64748B")

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return {"out": str(out), "bytes": out.stat().st_size, "tables": len(doc.tables)}


def main() -> None:
    ap = argparse.ArgumentParser(description="由案例 JSON 生成中英对照合同 DOCX（模拟合同）")
    ap.add_argument("--dir", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--with-tags", action="store_true", help="在条款下标注关联风险编号")
    a = ap.parse_args()
    print(json.dumps({"ok": True, **build(Path(a.dir), Path(a.out), a.with_tags)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
