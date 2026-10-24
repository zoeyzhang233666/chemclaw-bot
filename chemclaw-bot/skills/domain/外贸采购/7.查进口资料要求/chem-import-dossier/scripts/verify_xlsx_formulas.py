#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
verify_xlsx_formulas.py — 在没有 Excel/LibreOffice 的环境里验证工作簿公式

做法：读取公式文本 → 解析跨表引用与本地引用 → 用真实单元格值求值
覆盖本技能用到的那几类公式（四则运算、COUNTIF/COUNTIFS/COUNTA/SUM、
IFERROR、TEXT、& 连接、$ 绝对引用、跨表 '名称'!A1 引用）。
对无法解析的公式直接报 FAIL，不会静默放过。

用法： python scripts/verify_xlsx_formulas.py examples/h2o2-kr-30/dossier.xlsx
"""
from __future__ import annotations

import argparse
import os
import re
import sys

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

# 只匹配：'表名'!A1 / 表名!A1 / A1 / A1:B2（可带 $）
# 关键：
#  - 表名后必须有 ! 才算跨表引用（否则会把 COUNTIF 当成表名）
#  - 单元格引用两侧不能紧邻标识符字符（避免把 "P0" 这类字符串当引用）
#  - 字符串字面量在替换前会被整体遮蔽
_C = r"(?<![A-Za-z0-9_\u4e00-\u9fff\"'])"          # 左边界
_R = r"(?![A-Za-z0-9_\u4e00-\u9fff\"])"             # 右边界
REF = re.compile(
    _C + r"(?:(?:'([^']+)'|([A-Za-z0-9_\u4e00-\u9fff]{1,40}))!)?"
    r"\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?" + _R
)
STR = re.compile(r'"[^"]*"')

# 合法的公式起始：= 后紧跟函数名( / 引用 / 数字 / 括号 / 引号
FORMULA_START = re.compile(
    r"^=(?:[A-Z][A-Z0-9.]*\(|[A-Z]{1,3}\$?\d|\$|[0-9]|\(|\"|')"
)


def is_formula(v):
    """区分真公式与「以 = 开头的说明文字」（如 '= USD × 汇率'）。"""
    return isinstance(v, str) and bool(FORMULA_START.match(v))


class Evaluator:
    def __init__(self, path):
        self.path = path
        self.wb = load_workbook(path)
        self.cache = {}
        self.errors = []
        self.detail = []

    # ---------------------------------------------------------------- 取值
    def val(self, sheet, col, row):
        key = (sheet, col, row)
        if key in self.cache:
            return self.cache[key]
        ws = self.wb[sheet]
        cell = ws.cell(row=row, column=col)
        v = cell.value
        if is_formula(v):
            v = self.eval_formula(sheet, v[1:])
        self.cache[key] = v
        return v

    def range_vals(self, sheet, c1, r1, c2, r2):
        out = []
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                out.append(self.val(sheet, c, r))
        return out

    # ---------------------------------------------------------------- 解析引用
    def _resolve(self, cur_sheet, m):
        sname = m.group(1) or m.group(2) or cur_sheet
        c1 = column_index_from_string(m.group(3))
        r1 = int(m.group(4))
        if m.group(5):
            return sname, c1, r1, column_index_from_string(m.group(5)), int(m.group(6))
        return sname, c1, r1, c1, r1

    def _sub_refs(self, formula, cur_sheet):
        """把公式里的引用替换成数值/列表，返回可 eval 的 Python 表达式。

        先把字符串字面量（"P0"、"齐备" 等）整体遮蔽，避免被误认成单元格引用。
        """
        missing = []
        literals = []

        def stash(m):
            literals.append(m.group(0))
            return f"\x00{len(literals) - 1}\x00"

        masked = STR.sub(stash, formula)

        def repl(m):
            sname, c1, r1, c2, r2 = self._resolve(cur_sheet, m)
            if sname not in self.wb.sheetnames:
                missing.append(sname)
                return "None"
            if c2 == c1 and r2 == r1:
                return repr(self.val(sname, c1, r1))
            return repr(self.range_vals(sname, c1, r1, c2, r2))

        out = REF.sub(repl, masked)

        def unstash(s):
            return re.sub(r"\x00(\d+)\x00", lambda m: literals[int(m.group(1))], s)

        return unstash(out), missing

    # ---------------------------------------------------------------- 函数
    @staticmethod
    def _flat(args):
        out = []
        for a in args:
            if isinstance(a, list):
                out.extend(a)
            else:
                out.append(a)
        return out

    @staticmethod
    def _countif(vals, crit):
        def match(v):
            if isinstance(crit, str):
                op = re.match(r"^(<=|>=|<>|<|>|=)?(.*)$", crit)
                o, target = op.group(1) or "=", op.group(2)
                try:
                    t = float(target)
                    vv = float(v)
                    return {"=": vv == t, "<>": vv != t, "<": vv < t,
                            ">": vv > t, "<=": vv <= t, ">=": vv >= t}[o]
                except (TypeError, ValueError):
                    sv = "" if v is None else str(v)
                    if o == "=":
                        return sv == target
                    if o == "<>":
                        return sv != target
                    return False
            return v == crit
        return sum(1 for v in vals if match(v))

    def eval_formula(self, sheet, f):
        expr, missing = self._sub_refs(f, sheet)
        if missing:
            self.errors.append(f"{sheet}: 公式引用了不存在的表 {missing} → {f}")
            return None
        env = {
            "COUNTIF": lambda vals, crit: self._countif(self._flat([vals]), crit),
            "COUNTIFS": self._countifs,
            "COUNTA": lambda vals: sum(1 for v in self._flat([vals]) if v not in (None, "")),
            "SUM": lambda *a: sum(v for v in self._flat(list(a)) if isinstance(v, (int, float))),
            "IFERROR": lambda v, d: d if v in (None, "") else v,
            "TEXT": lambda v, fmt: self._text(v, fmt),
            "IF": lambda c, a, b: a if c else b,
            "ROUND": lambda v, n=2: round(float(v), int(n)),
        }
        try:
            py = expr.replace("<>", "!=")
            # 只有顶层出现 & 时才按连接处理（引用已替换为字面量/列表）
            depth = 0
            top_amp = False
            in_str = False
            for ch in py:
                if ch == '"':
                    in_str = not in_str
                elif not in_str:
                    if ch == "(":
                        depth += 1
                    elif ch == ")":
                        depth -= 1
                    elif ch == "&" and depth == 0:
                        top_amp = True
            if top_amp:
                parts = self._split_top(py, "&")
                vals = [eval(p, {"__builtins__": {}}, dict(env)) for p in parts]
                return "".join("" if v is None else str(v) for v in vals)
            return eval(py, {"__builtins__": {}}, dict(env))
        except Exception as e:  # noqa: BLE001
            self.errors.append(f"{sheet}: 无法求值 {f} → {e}")
            return None

    @staticmethod
    def _split_top(s, sep):
        out, buf, depth, in_str = [], "", 0, False
        for ch in s:
            if ch == '"':
                in_str = not in_str
            if not in_str:
                if ch == "(":
                    depth += 1
                elif ch == ")":
                    depth -= 1
                elif ch == sep and depth == 0:
                    out.append(buf)
                    buf = ""
                    continue
            buf += ch
        out.append(buf)
        return [x.strip() for x in out]

    @staticmethod
    def _text(v, fmt):
        try:
            if "%" in str(fmt):
                return f"{float(v):.1%}"
            return str(v)
        except (TypeError, ValueError):
            return str(v)

    def _countifs(self, *a):
        """COUNTIFS(range1, crit1, range2, crit2, ...)"""
        ranges = [(a[i], a[i + 1]) for i in range(0, len(a), 2)]
        first = ranges[0][0]
        n = len(first) if isinstance(first, list) else 1
        hits = 0
        for i in range(n):
            ok = True
            for rng, crit in ranges:
                vals = [rng[i]] if isinstance(rng, list) else [rng]
                if self._countif(vals, crit) != 1:
                    ok = False
                    break
            if ok:
                hits += 1
        return hits

    # ---------------------------------------------------------------- 报告
    def run(self):
        n = 0
        for s in self.wb.sheetnames:
            ws = self.wb[s]
            for row in ws.iter_rows():
                for c in row:
                    if is_formula(c.value):
                        n += 1
                        v = self.eval_formula(s, c.value[1:])
                        self.detail.append((s, c.coordinate, c.value, v))
        return n


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("-v", "--verbose", action="store_true")
    a = ap.parse_args()
    if not os.path.exists(a.xlsx):
        print(f"找不到 {a.xlsx}")
        return 1
    ev = Evaluator(a.xlsx)
    n = ev.run()
    print(f"公式总数 {n}")
    if a.verbose:
        for s, coord, f, v in ev.detail:
            print(f"  {s}!{coord:6s} {f[:70]:70s} = {v}")
    # 关键指标断言
    checks = []
    import openpyxl
    wb = load_workbook(a.xlsx)
    dd = wb["11 进度看板"]
    dash = {}
    for r in range(1, 40):
        k = dd.cell(row=r, column=2).value
        v = dd.cell(row=r, column=3).value
        if k:
            if is_formula(v):
                dash[k] = ev.val("11 进度看板", 3, r)
            else:
                dash[k] = v
    tx = wb["10 税费测算"]
    for r in range(20, 45):
        k = tx.cell(row=r, column=1).value
        if k and is_formula(tx.cell(row=r, column=2).value):
            dash["TAX:" + k] = ev.val("10 税费测算", 2, r)

    print("\n--- 关键指标（求值结果）---")
    for k, v in dash.items():
        if k.startswith(("单证", "P0", "缺口", "待确认", "监管", "TAX:", "齐备", "有但", "缺失", "待定", "过期", "不适用")):
            print(f"  {k:36s} = {v}")

    # 断言
    ok = True
    tax_checks = {
        "TAX:CIF 金额（CNY）": (24200 * 7.118),
        "TAX:关税额（CNY）": (24200 * 7.118 * 0.055),
        "TAX:增值税计税价格（CNY）": (24200 * 7.118 * 1.055),
        "TAX:增值税额（CNY）": (24200 * 7.118 * 1.055 * 0.13),
    }
    for k, expect in tax_checks.items():
        got = dash.get(k)
        if got is None:
            print(f"  ✗ 缺少 {k}")
            ok = False
        elif abs(float(got) - expect) > 1.0:
            print(f"  ✗ {k} 期望 {expect:,.2f} 实得 {got}")
            ok = False
        else:
            print(f"  ✓ {k} = {got:,.2f}")
    if ev.errors:
        print("\n--- 错误 ---")
        for e in ev.errors:
            print("  ✗", e)
        ok = False
    print("\n" + ("全部公式求值通过" if ok else "存在公式问题"))
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
