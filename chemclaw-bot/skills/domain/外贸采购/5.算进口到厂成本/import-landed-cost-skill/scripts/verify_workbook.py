# -*- coding: utf-8 -*-
"""
verify_workbook.py — Excel 工作簿 ←→ Python 引擎 逐项对账

步骤：
  1. 重新生成工作簿（可选 --no-build 跳过）
  2. 用 LibreOffice 重算，导出带缓存值的副本
  3. 用 Python 引擎按同一输入计算
  4. 比对：关税计税价格 / 关税 / 增值税 / 现金支出 / 最终成本 / 元/吨 / 敏感性
"""
from __future__ import annotations

_CHEMCLAW_ROOT = next(parent for parent in __import__('pathlib').Path(__file__).resolve().parents if (parent / 'runtime/deliver.mjs').is_file())

import argparse
import json
import os
import subprocess
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from openpyxl import load_workbook  # noqa: E402

from landed_cost_engine import LandedCostInput, land_cost  # noqa: E402

NODE = __import__('sys').executable
CLI = (str(_CHEMCLAW_ROOT / 'runtime/office-cli.py'))
XL = os.path.abspath(os.path.join(HERE, "..", "assets", "import_landed_cost_calculator.xlsx"))
RECALC = os.path.join(HERE, "..", "build", "recalc.xlsx")

# Excel 页 = 引擎输入（与 build_workbook 默认值保持一致）
COMMON = dict(
    currency="USD", incoterm_version="2020", price_unit="MT", quantity_unit="MT", packing_included=True,
    product="工业级乙二醇", hs_code="29053100", origin_country="SA",
    destination_port="宁波港", final_address="江苏省苏州市",
    fx_market_rate=6.6807, fx_customs_rate=6.72,
    packing_extra=0.0, freight_includes_dest_thc=False,
    insurance_rate_pct=0.55, insurance_amount=None,
    port_charges_cny=600, customs_broker_cny=500, inspection_cny=0,
    warehouse_cny=3000, warehouse_confirmed_none=False, inland_truck_cny=7812.5, financing_cny=0,
    tariff_mfn_rate=0.055, tariff_interim_rate=None, tariff_agreement_rate=None,
    tariff_general_rate=0.30, applied_rate_type="MFN", origin_eligible=False,
    anti_dumping_rate=0.0, vat_rate=0.13, vat_deductible=True, consumption_tax_rate=0.0,
    duty_paid_by="buyer", quantity=62.5,
)
CASES = {
    "FOB": dict(COMMON, incoterm="FOB", incoterm_place="朱拜勒", unit_price=640.00,
                ocean_freight_per_mt=45.00, insurance_included=False),
    "CFR": dict(COMMON, incoterm="CFR", incoterm_place="宁波", unit_price=685.00,
                ocean_freight_per_mt=None, insurance_included=False),
    "CIF": dict(COMMON, incoterm="CIF", incoterm_place="宁波", unit_price=685.05,
                ocean_freight_per_mt=None, insurance_included=True),
}
CELLS = {"customs_value": "C41", "duty": "C42", "vat": "C51", "credit": "C52",
         "cash": "C69", "final": "C71", "per_ton": "C72", "cif_cny": "D35", "status": "C75"}


def recalc() -> None:
    os.makedirs(os.path.dirname(RECALC), exist_ok=True)
    if os.path.exists(RECALC):
        os.remove(RECALC)
    cmd = [NODE, CLI, "recalculate", "--input", XL, "--output", RECALC]
    print(" ".join(f'"{c}"' if " " in c else c for c in cmd[:2]), "...")
    p = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="replace")
    if p.returncode != 0:
        raise SystemExit(f"recalculate 失败：{p.stderr[:800]}")
    print("recalculated ->", RECALC)


CELLMAP = {"customs_value": ("cv_cny", "C"), "duty": ("duty", "C"), "vat": ("vat_cash", "C"),
           "credit": ("vat_credit", "C"), "cash": ("cash", "C"), "final": ("final_tot", "C"),
           "per_ton": ("final_pt", "C"), "cif_cny": ("cif_total", "D"), "status": ("status", "C")}


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--no-build", action="store_true")
    a = ap.parse_args()
    if not a.no_build:
        subprocess.run([sys.executable, os.path.join(HERE, "build_workbook.py")], check=True)
    recalc()

    rowmap = json.load(open(os.path.join(HERE, "..", "build", "rowmap.json"), encoding="utf-8"))
    wb = load_workbook(RECALC, data_only=True)
    rows, fails = [], []
    for term, cfg in CASES.items():
        ws = wb[term]
        rm = rowmap[term]
        res = land_cost(LandedCostInput.from_dict(cfg))
        L, S = res["lines"], res["summary"]
        exp = {"customs_value": L["customs_value_cny"], "duty": L["duty_cny"],
               "vat": L["import_vat_cash_cny"], "credit": L["import_vat_input_credit_cny"],
               "cash": S["cash_out_total_cny"], "final": S["final_cost_total_cny"],
               "per_ton": S["landed_cost_per_ton_cny"], "cif_cny": S["cif_equivalent_total_cny"]}
        for key, (rk, col) in CELLMAP.items():
            cell = f"{col}{rm[rk]}"
            got = ws[cell].value
            if key == "status":
                ok = got == res["status"]
                rows.append({"case": term, "item": "状态", "cell": cell, "excel": got,
                             "engine": res["status"], "ok": ok})
            else:
                e = exp[key]
                ok = (isinstance(got, (int, float)) and e is not None and abs(got - e) < 0.02)
                rows.append({"case": term, "item": key, "cell": cell, "excel": got,
                             "engine": None if e is None else round(e, 4), "ok": bool(ok)})
            if not rows[-1]["ok"]:
                fails.append(rows[-1])

    # 敏感性：按标签定位行，再逐情景对账
    ws = wb["敏感性分析"]
    label_row = {}
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=2).value
        if isinstance(v, str):
            label_row[v.strip()] = r
    # 敏感性页一律基于 FOB 页 —— 对账也统一用 FOB 配置
    cfg = CASES["FOB"]
    for term in ("FOB",):
        checks = []
        for k in (-0.03, 0.0, 0.03):
            d2 = dict(cfg)
            d2["fx_market_rate"] = cfg["fx_market_rate"] * (1 + k)
            d2["fx_customs_rate"] = cfg["fx_customs_rate"] * (1 + k)
            checks.append((f"汇率 {k:+.0%}", land_cost(LandedCostInput.from_dict(d2))
                           ["summary"]["landed_cost_per_ton_cny"]))
        if term == "FOB":
            for k in (-0.20, -0.10, 0.0, 0.10, 0.20):
                d2 = dict(cfg)
                d2["ocean_freight_per_mt"] = cfg["ocean_freight_per_mt"] * (1 + k)
                checks.append((f"海运费 {k:+.0%}", land_cost(LandedCostInput.from_dict(d2))
                               ["summary"]["landed_cost_per_ton_cny"]))
        for k in (0.5, 0.75, 1.0, 1.5, 2.0):
            d2 = dict(cfg)
            d2["quantity"] = cfg["quantity"] * k
            checks.append((f"数量 ×{k:.2f}".rstrip('0').rstrip('.'), land_cost(LandedCostInput.from_dict(d2))
                           ["summary"]["landed_cost_per_ton_cny"]))
        for label, e in checks:
            r = label_row.get(label)
            if r is None:
                rows.append({"case": f"{term}/敏感性", "item": label, "excel": None,
                             "engine": None if e is None else round(e, 4), "ok": False})
                fails.append(rows[-1])
                continue
            got = ws[f"D{r}"].value
            ok = isinstance(got, (int, float)) and e is not None and abs(got - e) < 0.05
            rows.append({"case": f"{term}/敏感性", "item": label, "cell": f"D{r}", "excel": got,
                         "engine": None if e is None else round(e, 4), "ok": bool(ok)})
            if not ok:
                fails.append(rows[-1])

    passed = sum(1 for r in rows if r["ok"])
    print(f"\n对账项：{len(rows)}  通过：{passed}  失败：{len(fails)}")
    for f in fails:
        print("  ✗", f)
    out = os.path.join(HERE, "..", "build", "workbook_reconciliation.json")
    with open(out, "w", encoding="utf-8") as f:
        json.dump({"total": len(rows), "passed": passed, "rows": rows}, f, ensure_ascii=False, indent=2)
    print("detail ->", os.path.abspath(out))
    return 0 if not fails else 1


if __name__ == "__main__":
    sys.exit(main())
