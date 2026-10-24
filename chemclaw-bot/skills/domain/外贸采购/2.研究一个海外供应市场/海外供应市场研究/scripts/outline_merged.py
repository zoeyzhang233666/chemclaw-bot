#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
outline_merged.py —— 给 Excel 中被合并的单元格区域补外边框

原因：openpyxl 合并单元格后，只有左上角单元格保留边框，打开后会出现"断线"。
本脚本在 build_excel.py 生成之后运行（或作为函数被 import）。
"""
from openpyxl.styles import Border, Side
from openpyxl.utils import get_column_letter

THIN = Side(style="thin", color="BFBFBF")

def outline(ws, r0, r1, c0, c1):
    for r in range(r0, r1 + 1):
        for c in range(c0, c1 + 1):
            cell = ws.cell(row=r, column=c)
            b = cell.border
            cell.border = Border(
                left=THIN if c == c0 else (b.left if b and b.left else None),
                right=THIN if c == c1 else (b.right if b and b.right else None),
                top=THIN if r == r0 else (b.top if b and b.top else None),
                bottom=THIN if r == r1 else (b.bottom if b and b.bottom else None),
            )

def outline_all(ws):
    for rng in list(ws.merged_cells.ranges):
        outline(ws, rng.min_row, rng.max_row, rng.min_col, rng.max_col)
