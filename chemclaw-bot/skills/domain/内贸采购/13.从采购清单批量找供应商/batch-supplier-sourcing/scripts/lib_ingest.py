# -*- coding: utf-8 -*-
"""batch_sourcing.ingest — 采购清单解析：Excel/CSV/JSON → 原始行 → 标准字段 → 化学品映射。"""
from __future__ import annotations

import csv
import io
import os
import re

from lib_common import MaterialsDict, is_form_variant, cas_norm, cas_valid, clean_text, fmt_num, norm_key, normalize_unit, to_number

READERS = ("xlsx", "xls", "xlsm", "csv", "json")


# ------------------------------------------------------------------ 原始读取
def _read_xlsx(path, sheet=None, header_row=None):
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = [[c for c in r] for r in ws.iter_rows(values_only=True)]
    wb.close()
    return rows, ws.title, header_row


def _read_csv(path, encoding=None):
    raw = open(path, "rb").read()
    for enc in (encoding, "utf-8-sig", "gb18030", "utf-8"):
        if not enc:
            continue
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = raw.decode("utf-8", errors="replace")
    dialect = csv.Sniffer().sniff(text[:2000], delimiters=",\t;") if text.count(",") > 10 or "\t" in text else csv.excel
    rows = [r for r in csv.reader(io.StringIO(text), dialect)]
    return rows, "CSV", None


def _read_json(path):
    import json as _j

    data = _j.load(open(path, "r", encoding="utf-8-sig"))
    rows = data.get("rows") if isinstance(data, dict) else data
    headers = data.get("fields") if isinstance(data, dict) else None
    if headers is None:
        headers = list(rows[0].keys())
        body = [[r.get(h) for h in headers] for r in rows]
    else:
        body = [[r.get(h) for h in headers] for r in rows]
    return [headers] + body, "JSON", 1


def read_rows(path, sheet=None, header_row=None):
    ext = os.path.splitext(path)[1].lower().lstrip(".")
    if ext in ("xlsx", "xlsm", "xls"):
        return _read_xlsx(path, sheet, header_row)
    if ext == "csv":
        return _read_csv(path)
    if ext == "json":
        return _read_json(path)
    raise ValueError(f"不支持的清单格式: .{ext}（支持 xlsx/csv/json；pdf/docx 请先用 tools/ocr_to_csv.py 转成 CSV）")


# ------------------------------------------------------------------ 表头映射
def map_headers(header_cells, field_map):
    fields = field_map["fields"]
    used = set()
    col_map = {}
    for idx, raw in enumerate(header_cells):
        k = norm_key(raw)
        if not k:
            continue
        hit = None
        # 1) 精确别名
        for f in fields:
            if f["key"] in used and f["key"] != "spec":
                continue
            if any(norm_key(a) == k for a in f.get("aliases", [])):
                hit = f
                break
        # 2) 包含匹配（长别名优先）
        if hit is None:
            cands = []
            for f in fields:
                if f["key"] in used and f["key"] != "spec":
                    continue
                for a in f.get("aliases", []):
                    ak = norm_key(a)
                    if len(ak) >= 2 and ak in k:
                        cands.append((len(ak), f))
            if cands:
                cands.sort(key=lambda x: -x[0])
                hit = cands[0][1]
        # 3) 正则 hint
        if hit is None:
            for f in fields:
                if f.get("hints") and f["key"] not in used:
                    if any(re.search(h, k) for h in f["hints"]):
                        hit = f
                        break
        if hit is not None:
            col_map[idx] = hit["key"]
            used.add(hit["key"])
        else:
            col_map[idx] = None
    return col_map


def _is_skip_row(cells, skip_tokens):
    text = " ".join(clean_text(c) for c in cells)
    if not text.strip():
        return True
    return any(t in text for t in skip_tokens) and len([c for c in cells if clean_text(c)]) <= 6


# ------------------------------------------------------------------ 化学品识别
_SPEC_IN_NAME = re.compile(r"(\d{1,2}(?:\.\d)?\s*%|无水|二水|六水|十水|片状|球状|固体|液体|粉状|颗粒|结晶|溶液|重质|轻质|工业级|食品级|饲料级|聚酯级|纤维级|电子级)")


def guess_form_from_name(raw_name):
    m = _SPEC_IN_NAME.findall(clean_text(raw_name))
    return m or []


def parse_purity(raw):
    s = clean_text(raw)
    if not s:
        return None, ""
    m = re.search(r"(\d{1,3}(?:\.\d+)?)\s*%", s)
    if not m:
        return None, s
    return float(m.group(1)), s[m.end():].strip(" ;；,，")


def normalize_list(project, mdict: MaterialsDict, fm, rules):
    """核心入口：list_spec + rows → 标准化物料表 + ambiguities + parse_notes。"""
    list_spec = project["list"]
    rows, sheet_name, _ = read_rows(list_spec["path"], list_spec.get("sheet"), list_spec.get("header_row"))
    header_row_no = int(list_spec.get("header_row") or 0)
    if header_row_no <= 0:
        # 自动定位表头：第一个非空且命中 raw_name 的行
        for i, r in enumerate(rows[:10]):
            cm = map_headers(r, fm)
            if "raw_name" in cm.values() or "combination" in cm.values():
                header_row_no = i + 1
                break
        if header_row_no <= 0:
            header_row_no = 1
    header = rows[header_row_no - 1]
    col_map = map_headers(header, fm)
    skip_tokens = rules.get("header_skip", [])

    parse_notes = []
    unmapped = [clean_text(h) for i, h in enumerate(header) if col_map.get(i) is None and clean_text(h)]
    if unmapped:
        parse_notes.append({"level": "info", "msg": f"未映射列（未参与标准化）: {unmapped}"})

    items, ambiguities = [], []
    seen_codes = {}
    for r_i, raw_row in enumerate(rows[header_row_no:], start=header_row_no + 1):
        if _is_skip_row(raw_row, skip_tokens):
            continue
        rec = {}
        for c_i, key in col_map.items():
            if key and c_i < len(raw_row):
                rec.setdefault(key, clean_text(raw_row[c_i]))
        if not any(rec.values()):
            continue
        raw_name = rec.get("raw_name", "") or rec.get("combination", "")
        if not raw_name:
            parse_notes.append({"level": "warn", "msg": f"第 {r_i} 行缺少物料名称，已跳过"})
            continue
        name_before = raw_name
        codem = re.search(r"([A-Z]{1,4}-?\d{2,}|M\d{3,})", raw_name)
        combination_cas = None
        if rec.get("combination"):
            cm = re.search(r"\b(\d{2,7}-\d{2}-\d)\b", rec["combination"])
            if cm:
                combination_cas = cm.group(1)
                raw_name = rec["combination"].replace(combination_cas, "").strip(" ()（）、,")
        spec_from_code = None
        if codem:
            spec_from_code = "清单内编码"
        item = _build_item(project, mdict, rules, rec, raw_name, combination_cas, r_i, name_before)
        item["source_row"] = r_i
        item["code"] = item.get("code") or f"P{r_i - header_row_no:02d}"
        # 合并单元格/重复行：同一物料同一形态再次出现 → 标记“疑似重复”
        key = (item.get("name_key"), item.get("form"))
        if item.get("name_key") and key in seen_codes:
            item["dup_of"] = seen_codes[key]
            item["issues"].append(f"与第 {seen_codes[key]} 行同一物料/形态，疑似重复行或不同批次，需确认合并与否")
        seen_codes.setdefault(key, r_i)
        items.append(item)
        AMB_KEYS = ("数量缺失", "数量为 0", "计量单位", "CAS 校验失败", "未提供 CAS", "计数单位", "不在字典形态清单")
        if item["status"] != "ok" or item.get("dup_of") or any(k in " ".join(item["issues"]) for k in AMB_KEYS):
            for a in _ambiguity_for(item, rec):
                a["id"] = f"A-{len(ambiguities) + 1:03d}"
                ambiguities.append(a)
    return {"items": items, "ambiguities": ambiguities, "parse_notes": parse_notes,
            "sheet": sheet_name, "header_row": header_row_no,
            "header": [clean_text(h) for h in header],
            "col_map": {clean_text(header[i]) if i < len(header) else str(i): v for i, v in col_map.items() if v}}


def _build_item(project, mdict, rules, rec, raw_name, combination_cas, r_i, name_before):
    code = rec.get("code") or f"M{r_i:03d}"
    item = {
        "code": rec.get("code") or "",
        "raw_name": name_before,
        "name": "",
        "en": "",
        "aliases": [],
        "cas": "",
        "cas_raw": rec.get("cas", "") or combination_cas or "",
        "form": rec.get("form", ""),
        "spec": rec.get("spec", ""),
        "purity": None,
        "unit": "",
        "unit_raw": rec.get("unit", ""),
        "qty": None,
        "qty_annual": None,
        "packing": rec.get("packing", ""),
        "delivery": rec.get("delivery", ""),
        "application": rec.get("application", ""),
        "supplier": rec.get("supplier", ""),
        "hazard": "",
        "hazard_ref": "",
        "status": "ok",
        "issues": [],
        "name_key": "",
        "notes": rec.get("notes", ""),
    }
    # 名称/CAS
    cas_in = cas_norm(item["cas_raw"]) if item["cas_raw"] else ""
    if cas_in and not cas_valid(cas_in):
        item["issues"].append("CAS 校验失败（校验位不符或格式错误）")
    name_hits = mdict.match_name(raw_name)
    cas_hits = mdict.lookup_cas(cas_in) if cas_in else []
    if len(name_hits) > 1:
        base = [h for h in name_hits if "(" not in h["name"] and "（" not in h["name"]]
        if base and all(is_form_variant(base[0], h) for h in name_hits):
            name_hits = base[:1]
        elif base:
            name_hits = base[:1]
            item["issues"].append("名称可对应多个形态：" + "、".join(h["name"] for h in name_hits))
        else:
            item["status"] = "ambiguous"
            item["issues"].append(f"名称可对应 {len(name_hits)} 种物料：" + "、".join(h["name"] for h in name_hits))
            item["candidates"] = [h["name"] for h in name_hits]
    if not name_hits and cas_hits:
        name_hits = cas_hits
        item["issues"].append("按 CAS 反查得到名称，原名称未在字典命中，需确认是否为同一物质")
        item["status"] = "ambiguous"
    if name_hits:
        m = name_hits[0]
        item["name"] = m["name"]
        item["en"] = m.get("en", "")
        item["cas"] = m.get("cas", "")
        item["hazard"] = m.get("hazard", "")
        item["hazard_ref"] = m.get("hazard_ref", "")
        item["aliases"] = list(m.get("aliases") or [])
        item["family"] = m.get("families", [])
        item["dict_key"] = m["key"]
        item["name_key"] = m["key"]
        if cas_in and m.get("cas") and cas_norm(m["cas"]) != cas_in:
            item["status"] = "conflict"
            item["issues"].append(f"清单 CAS {cas_in} 与字典 CAS {m['cas']} 不一致")
        if not item["form"]:
            forms = m.get("forms") or []
            g = guess_form_from_name(raw_name)
            matched = [f for f in forms if f and any(f in raw_name or g_ in f for g_ in g)]
            item["form"] = matched[0] if matched else (g[0] if g and g[0] in forms else "")
            if g and not matched:
                item["issues"].append(f"名称内自带的形态/浓度『{'/'.join(g)}』不在字典形态清单（{ '、'.join(forms) }）内，需确认口径")
                if item["status"] == "ok":
                    item["status"] = "ambiguous"
    else:
        item["status"] = "unrecognized"
        item["name"] = raw_name
        item["issues"].append("化学品名称未在标准化字典命中；CAS 未提供或无法交叉验证")
    if not item["cas"]:
        item["cas"] = cas_in
    if not cas_in and item["cas"]:
        item["issues"].append("清单未提供 CAS，已按名称字典补齐，建议供应商确认")

    # 规格 & 纯度
    purity_val, purity_rest = parse_purity(item["spec"])
    if rec.get("purity"):
        purity_val2, _ = parse_purity(rec["purity"])
        purity_val = purity_val2 if purity_val is None else purity_val
    item["purity"] = purity_val
    if purity_val is not None and not item["spec"]:
        item["spec"] = purity_rest
    # 名称里的浓度 → form
    if purity_val is None:
        nm = re.search(r"(\d{1,2}(?:\.\d)?)\s*%", raw_name)
        if nm and not item["form"]:
            item["form"] = f"{nm.group(1)}%"
    if not item["spec"] and not item["purity"]:
        item["issues"].append("未提供执行标准/质量指标，需在询价中澄清（不阻塞）")

    # 数量 & 单位
    q = to_number(rec.get("quantity"))
    item["qty"] = q
    unit_c, factor, cat, ambig = normalize_unit(item["unit_raw"], rules)
    item["unit"] = unit_c
    if q is None:
        item["status"] = "blocked" if item["status"] == "ok" else item["status"]
        item["issues"].append("数量缺失或不可解析")
    elif q <= 0:
        item["issues"].append("数量为 0 或负数，疑似占位/未填")
        if item["status"] == "ok":
            item["status"] = "blocked"
    elif unit_c == "t" and factor:
        item["qty"] = round(q * factor, 3)
    elif ambig and cat not in ("count",):
        item["status"] = "blocked" if item["status"] == "ok" else item["status"]
        item["issues"].append(f"计量单位『{item['unit_raw']}』无法标准化（{cat}）")
    elif cat == "count":
        item["issues"].append(f"计量单位『{item['unit_raw']}』为件/桶等计数单位，需换算为净重吨再比价")
    item["unit"] = unit_c or item["unit_raw"]
    ann = to_number(rec.get("annual_usage"))
    mon = to_number(rec.get("monthly_usage"))
    item["qty_annual"] = ann if ann is not None else (round(mon * 12, 3) if mon is not None else None)
    if item["qty_annual"] is None and q is not None and cat == "mass":
        item["qty_annual"] = None
    item["qty_monthly"] = mon
    item["usage_basis"] = "年用量" if ann is not None else ("月用量×12" if mon is not None else "未披露")
    if item["status"] not in ("ok",):
        item["issues"].append("状态 = " + item["status"])
    return item


def _ambiguity_for(item, rec):
    level = {"blocked": "P0", "conflict": "P0", "unrecognized": "P0", "ambiguous": "P1"}.get(item["status"], "P1")
    issues = item["issues"]
    qs = []
    if item["status"] == "ambiguous":
        if "candidates" in item:
            qs.append(f"『{item['raw_name']}』可对应 " + "、".join(item["candidates"]) +
                      "（CAS 相同但规格不同）。请按最终物料拆行，并写明浓度/形态与执行标准")
        else:
            qs.append("名称内规格与字典形态不一致，请确认浓度/形态口径")
    elif item["status"] == "unrecognized":
        qs.append(f"『{item['raw_name']}』未在字典命中，请补充标准化学名、CAS、执行标准或 MSDS 后确认")
    elif item["status"] == "conflict":
        qs.append("清单 CAS 与名称对应的物质不一致，请确认哪一栏为准；建议以 MSDS 第 3 节为准")
    for msg in issues:
        if msg.startswith("计量单位") and not any(qs):
            qs.append(f"计量单位『{item['unit_raw']}』需确认：请改为吨/千克/立方米，或给出单位净重以便折算")
        elif msg.startswith("数量缺失") or msg.startswith("数量为 0"):
            qs.append("数量缺失或为 0，请补齐实际需求量（数量/单位必须可折算成吨或立方米）")
        elif msg.startswith("不在字典形态清单"):
            qs.append(msg)
    if item.get("dup_of"):
        qs.append(f"与第 {item['dup_of']} 行物料相同（同 CAS 同形态），是否合并为一行统计总量？")
    missing_cas = any(m.startswith("清单未提供 CAS") for m in issues)
    if missing_cas:
        qs.append(f"清单未填 CAS（已按名称补 {item.get('cas') or '—'}），请确认是否与该物质一致，特别是其浓度/形态")
    if item["status"] in ("ambiguous", "unrecognized", "conflict", "blocked") and not any(
            k in " ".join(qs) for k in ("CAS", "物质", "字典", "数量", "重复")):
        qs.append("；".join(issues))
    if not qs:
        qs.append("；".join(issues) if issues else "待确认")
    if any(m.startswith("CAS 校验失败") for m in issues) and item["status"] != "conflict":
        level = "P0"
    if item.get("dup_of"):
        level = "P2" if level == "P1" else level
    return [{
        "row": item["raw_name"],
        "code": item.get("code", ""),
        "source_row": item.get("source_row"),
        "level": level,
        "type": item["status"],
        "question": qs[0],
        "extra": qs[1:],
        "detail": "；".join(issues),
        "owner": "采购方",
        "status": "待确认",
    }]
