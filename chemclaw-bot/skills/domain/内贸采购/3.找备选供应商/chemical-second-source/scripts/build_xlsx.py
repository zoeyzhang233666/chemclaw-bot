#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Render the 10-sheet Excel workbook from data.json.

Usage:
  python build_xlsx.py data.json out.xlsx
"""
from __future__ import annotations

import argparse
import math
import os
from collections import Counter

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.formatting.rule import ColorScaleRule, DataBarRule, FormulaRule
    from openpyxl.chart import BarChart, PieChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError as e:  # pragma: no cover
    raise SystemExit("缺少 openpyxl：python -m pip install openpyxl") from e

import common

NAVY = "0B2545"
NAVY2 = "123B63"
BLUE = "1F5EFF"
TEAL = "0E7490"
AMBER = "B45309"
RED = "B42318"
GREEN = "067647"
VIOLET = "6D28D9"
LIGHT = "F5F8FC"
LIGHT2 = "FBFDFF"
LINE = "D9E2EC"
WHITE = "FFFFFF"
INK = "1A2B3C"
MUTED = "5A6B7D"
FONT = "微软雅黑"

thin = Side(style="thin", color=LINE)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
RISK_LEVELS = ["低", "中", "中高", "高"]


def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)


def font(size=10, bold=False, color=INK, italic=False, underline=None):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic, underline=underline)


def align(h="left", v="center", wrap=True, indent=0):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap, indent=indent)


def est_lines(text, width):
    if text is None:
        return 1
    units = sum(2 if ord(ch) > 127 else 1 for ch in str(text))
    return max(1, math.ceil(units / max(width * 1.02, 4)))


def fit_row(ws, row, specs, min_h=18, max_h=150, line_h=13.2):
    lines = 1
    for text, width in specs:
        lines = max(lines, est_lines(text, width))
    ws.row_dimensions[row].height = min(max_h, max(min_h, lines * line_h + 5))


def section_bar(ws, row, ncols, text, color=NAVY, size=11):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1, value=text)
    c.font = font(size, True, WHITE)
    c.fill = fill(color)
    c.alignment = align("left", "center", False, 1)
    ws.row_dimensions[row].height = 26


def title_block(ws, ncols, title, subtitle, tag="内部采购评估 · 公开信息研究"):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=title)
    c.font = font(20, True, WHITE)
    c.fill = fill(NAVY)
    c.alignment = align("left", "center", False, 1)
    ws.row_dimensions[1].height = 44
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    c = ws.cell(row=2, column=1, value=subtitle)
    c.font = font(10, False, "C9D9EA")
    c.fill = fill(NAVY2)
    c.alignment = align("left", "center", False, 1)
    ws.row_dimensions[2].height = 24
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=ncols)
    c = ws.cell(row=3, column=1, value=tag)
    c.font = font(9, False, MUTED, italic=True)
    c.fill = fill(LIGHT)
    c.alignment = align("left", "center", False, 1)
    ws.row_dimensions[3].height = 20


def header_row(ws, row, headers, widths, height=42, fill_color=NAVY):
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = font(9, True, WHITE)
        c.fill = fill(fill_color)
        c.alignment = align("center", "center", True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = height


def table_rows(ws, start_row, rows, widths, zebra=True, align_map=None):
    for r, row in enumerate(rows, start=start_row):
        for ci, val in enumerate(row, start=1):
            c = ws.cell(row=r, column=ci, value=val)
            c.font = font(9)
            c.border = BORDER
            if align_map and ci in align_map:
                c.alignment = align(align_map[ci], "top" if align_map[ci] == "left" else "center", True)
            else:
                c.alignment = align("left" if ci <= 2 else "center", "top", True)
            if zebra and (r - start_row) % 2 == 1:
                c.fill = fill(LIGHT2)


def style_link(cell, url):
    cell.hyperlink = url
    cell.font = font(8.5, False, BLUE, underline="single")


def page_setup(ws, title_rows="1:3"):
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = title_rows
    ws.oddFooter.center.text = "第二供应源备选池 · 第 &P / &N 页"
    ws.oddFooter.center.size = 8
    ws.oddFooter.center.color = "808080"


def build(data_path, out_path):
    raw = common.load(data_path)
    D = common.normalize(raw)
    M = D["meta"]
    base = D["baseline"]
    cands = D["candidates"]
    specs = D["spec_items"]
    status = D["spec_status"]
    same = D["same_source"]
    risks = D["risks"]
    actions = D["actions"]
    sources = D["sources"]
    src = {s["id"]: s for s in sources}
    product_short = M.get("product_short") or M.get("product")
    unit = M.get("unit") or "万吨/年"
    buyer_short = M.get("buyer_short") or str(M.get("buyer_site") or "").split("（")[0].split("(")[0].rstrip("市") or "收货地"
    qshort = M.get("qualification_short") or "准入"
    qlabel = M.get("qualification_label") or "准入 / 资质"
    grade_label = M.get("grade_label") or "等级 / 品牌"
    fy = f'{M.get("fin_year")} ' if M.get("fin_year") else ""
    oy = f'{M.get("output_year")} ' if M.get("output_year") else ""

    wb = Workbook()
    wb.properties.title = f"{M.get('product')} 第二供应源备选池 · 对比工作簿"
    wb.properties.subject = f"{M.get('buyer_site')} · {M.get('volume')} · 备选供应商研究"
    wb.properties.creator = "chemical-second-source skill"
    wb.properties.keywords = "化工品, 供应商, 第二供应源, 采购, 比选"

    # ------------------------------------------------------------------
    # 00 说明
    # ------------------------------------------------------------------
    ws = wb.active
    ws.title = "00_说明"
    ws.sheet_properties.tabColor = NAVY
    ws.sheet_view.showGridLines = False
    for col, w in {"A": 18, "B": 24, "C": 24, "D": 24, "E": 24, "F": 24}.items():
        ws.column_dimensions[col].width = w
    ws.merge_cells("A1:F1")
    c = ws["A1"]
    c.value = f"{M.get('product')} 第二供应源备选池"
    c.font = font(22, True, WHITE)
    c.fill = fill(NAVY)
    c.alignment = align("left", "center", False, 1)
    ws.row_dimensions[1].height = 50

    ws.merge_cells("A2:F2")
    c = ws["A2"]
    c.value = f"现有供应商：{base.get('name')}（{base.get('process')} {base.get('capacity')} {unit}）｜收货地：{M.get('buyer_site')}｜采购量：{M.get('volume')}｜规格：{M.get('spec')}"
    c.font = font(10, False, "D6E4F2")
    c.fill = fill(NAVY2)
    c.alignment = align("left", "center", False, 1)
    ws.row_dimensions[2].height = 32

    ws.merge_cells("A3:F3")
    c = ws["A3"]
    c.value = f"数据截止 {M.get('asof')} · 公开信息研究，不构成供应商替换或授标建议 · “待核实”项须以证照、COA、现场审核为准"
    c.font = font(9, False, MUTED, italic=True)
    c.fill = fill(LIGHT)
    c.alignment = align("left", "center", False, 1)
    ws.row_dimensions[3].height = 22

    section_bar(ws, 5, 6, "一、场景参数（可在询价单中直接复用）")
    params_rows = [
        ("现有供应商", f"{base.get('name')}（{base.get('listed','')} {base.get('code','')}）｜{base.get('process')} {base.get('capacity')} {unit}｜{base.get('city','')} {base.get('site','')}｜{qshort} {base.get('brand','')}｜距收货地约 {base.get('distance')} km"),
        ("采购产品", f"{M.get('product')}；用途：{M.get('use_case') or '待填'}；不接受替代等级，除非试样与工况验证通过"),
        ("质量规格", M.get("spec") or "待填"),
        ("采购量", f"{M.get('volume')}（年约 {round(float(M.get('volume_t') or 0) * (1 if M.get('volume_period') == '年' else 12)):,} 吨）"),
        ("收货地", M.get("buyer_site") or "待填"),
        ("目标", "建立 1 家主供 + 1 家备供的第二供应源；先验证，不先替换"),
    ]
    r = 6
    for k, v in params_rows:
        ws.cell(row=r, column=1, value=k).font = font(10, True, NAVY)
        ws.cell(row=r, column=1).fill = fill(LIGHT)
        ws.cell(row=r, column=1).border = BORDER
        ws.cell(row=r, column=1).alignment = align("left", "center", True, 1)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        c = ws.cell(row=r, column=2, value=v)
        c.font = font(9.5)
        c.border = BORDER
        c.alignment = align("left", "center", True, 1)
        ws.row_dimensions[r].height = 30 if len(str(v)) < 70 else 44
        r += 1

    r += 1
    section_bar(ws, r, 6, "二、研究方法（可复核、可复用）")
    r += 1
    method = M.get("method_steps") or []
    if not method:
        method = [
            {"label": "规格锚定", "text": "把现有供应商的牌号、COA、内控指标拆成可检测指标；先定义“兼容”，再找厂。"},
            {"label": "厂家身份", "text": "以统一社会信用代码/股票代码/集团股权穿透确认“生产企业”而非贸易商。"},
            {"label": "产能与区位", "text": "公开产能、投产时间、等级比例、到收货地运距/运输方式；估算运费区间。"},
            {"label": "资质与风险", "text": "行业准入、体系认证、环保/安全/财务公开记录；区分“已公开”与“待核实”。"},
            {"label": "试样与审厂", "text": "新装置、新工艺、非准入品牌一律首样；首样 20–50 kg × 3 批，按规格逐项验证。"},
        ]
    for s in method:
        ws.cell(row=r, column=1, value=str(s.get("label") or "")).font = font(9.5, True, BLUE)
        ws.cell(row=r, column=1).border = BORDER
        ws.cell(row=r, column=1).alignment = align("left", "top", True, 1)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        c = ws.cell(row=r, column=2, value=str(s.get("text") or ""))
        c.font = font(9.5)
        c.border = BORDER
        c.alignment = align("left", "top", True, 1)
        ws.row_dimensions[r].height = 30
        r += 1

    r += 1
    section_bar(ws, r, 6, "三、关键判断（中性，不替代采购决策）", color=TEAL)
    r += 1
    js = M.get("key_judgements") or []
    if not js:
        js = [{"label": "池子结构", "text": f"共 {len(cands)} 家备选；其中 {sum(1 for x in cands if x.get('qualified'))} 家具备{qshort}；{sum(1 for x in cands if x.get('new_plant'))} 家为新装置。"},
              {"label": "下一步", "text": "先按规格兼容性与试样结果缩小范围，再做现场审核与商务谈判；不以排序替代决策。"}]
    for s in js:
        ws.cell(row=r, column=1, value=str(s.get("label") or "")).font = font(9.5, True, TEAL)
        ws.cell(row=r, column=1).border = BORDER
        ws.cell(row=r, column=1).alignment = align("left", "top", True, 1)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        c = ws.cell(row=r, column=2, value=str(s.get("text") or ""))
        c.font = font(9.5)
        c.border = BORDER
        c.alignment = align("left", "top", True, 1)
        ws.row_dimensions[r].height = 38
        r += 1

    r += 1
    section_bar(ws, r, 6, "四、工作簿导航", color=BLUE)
    r += 1
    nav = [
        ("01_备选池", f"{len(cands)} 家备选的关键字段、风险摘要、证据链接；适合筛选与排序。"),
        ("02_图表", "产能、运费、工艺结构可视化；图表数据可编辑。"),
        ("03_横向对比", "现有 + 备选逐项对比；运距/运费/产能/风险用色阶标记相对位置。"),
        ("04_试样计划", f"{len(specs)} 项指标 × {len(cands)} 家的试样矩阵与状态跟踪。"),
        ("05_同源关系", "股权、集团、实控人与同源判断，避免“假第二源”。"),
        ("06_风险台账", f"{len(risks)} 条公开风险事实、影响、来源与核实状态。"),
        ("07_行动清单", f"{M.get('action_window')} 双源推进动作、责任与产出。"),
        ("08_数据来源", f"{len(sources)} 条来源，A/B/C 证据等级与日期。"),
        ("09_尽调清单", f"{len(cands)} 家 × {len(M.get('dd_items') or [])} 项尽调状态矩阵。"),
    ]
    for k, v in nav:
        ws.cell(row=r, column=1, value=k).font = font(9, True, NAVY)
        ws.cell(row=r, column=1).border = BORDER
        ws.cell(row=r, column=1).alignment = align("left", "center", True, 1)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        c = ws.cell(row=r, column=2, value=v)
        c.font = font(9)
        c.border = BORDER
        c.alignment = align("left", "center", True, 1)
        ws.row_dimensions[r].height = 26
        r += 1

    r += 1
    section_bar(ws, r, 6, "五、证据等级、风险口径与免责", color=AMBER)
    r += 1
    legend = [
        ("证据等级", "A = 上市公司公告/政府网站；B = 行业协会/券商研报/主流媒体；C = 企业官网/百科。C 级需二次核验。"),
        ("风险等级", "低 / 中 / 中高 / 高：综合财务、装置、物流、资质、同源、规模六项的研究判断；1–5 分仅用于排序。"),
        ("试样状态", "COA / 试样 / 审核；在 04 表与 09 表可跟踪。"),
        ("免责", M.get("disclaimer") or "本工作簿为采购研究辅助材料，不构成供应商选择、替换或授标建议。"),
    ]
    for k, v in legend:
        ws.cell(row=r, column=1, value=k).font = font(9, True, AMBER)
        ws.cell(row=r, column=1).border = BORDER
        ws.cell(row=r, column=1).alignment = align("left", "top", True, 1)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        c = ws.cell(row=r, column=2, value=v)
        c.font = font(9)
        c.border = BORDER
        c.alignment = align("left", "top", True, 1)
        ws.row_dimensions[r].height = 26 if len(str(v)) < 80 else 58
        r += 1
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"A1:F{r-1}"

    # ------------------------------------------------------------------
    # 01 备选池
    # ------------------------------------------------------------------
    ws = wb.create_sheet("01_备选池")
    ws.sheet_properties.tabColor = BLUE
    widths = [7, 30, 8, 8, 28, 10, 16, 8, 8, 8, 11, 40, 20, 9, 12, 26, 22, 10, 14, 12, 12, 12, 30, 10, 8, 22, 8, 7, 44, 46, 34, 40, 7]
    headers = ["编号", "企业名称", "类型", "省", "市/园区", f"距{buyer_short} km", "运输方式", "运费低\n元/吨", "运费高\n元/吨", "工艺",
               f"{product_short}产能\n{unit}", "产能备注", grade_label, qshort, "上市/代码", "集团", "实控人", "股权性质",
               "员工（公开）", f"{fy}营收\n亿元", f"{fy}净利\n亿元", f"{oy}产量", "主营产品", "规格\n兼容", "试样", "同源关系",
               "风险\n等级", "风险分", "风险摘要", f"与{base.get('short','现有')}差异", "需人工核实", "首条证据链接", "证据\n等级"]
    title_block(ws, len(widths), f"01 · 备选供应商池（{len(cands)} 家）",
                "按“身份—区位—产能—资质—风险—差异”组织；运费为估算区间，不是承运商报价。")
    header_row(ws, 4, headers, widths, height=40)
    rows = []
    for c in [base] + cands:
        ev = src.get((c.get("evidence") or [None])[0]) if c.get("evidence") else None
        rows.append([
            c.get("id"), c.get("name"), "现有" if c is base else "备选", c.get("province"), f'{c.get("city","")} · {c.get("site","")}',
            c.get("distance"), c.get("transport"), c.get("freight_low"), c.get("freight_high"), c.get("process"),
            c.get("capacity"), c.get("capacity_note"), f'{c.get("grade","")} / {c.get("brand","")}',
            "是" if c.get("qualified") else "否", f'{c.get("listed","")} {c.get("code","")}'.strip(),
            c.get("group"), c.get("controller"), c.get("ownership"), c.get("employees"), c.get("revenue"), c.get("profit"),
            c.get("output"), c.get("products"), c.get("spec_fit"), c.get("sample"), c.get("same_source"),
            c.get("risk_level"), c.get("risk_score"), c.get("risk_summary"), c.get("diff"), "；".join(c.get("verify") or []),
            (ev or {}).get("url", ""), (ev or {}).get("grade", ""),
        ])
    start = 5
    table_rows(ws, start, rows, widths, zebra=True,
               align_map={1: "center", 2: "left", 5: "left", 12: "left", 13: "left", 16: "left", 17: "left",
                          23: "left", 29: "left", 30: "left", 31: "left", 32: "left"})
    end = start + len(rows) - 1
    for rr in range(start, end + 1):
        fit_row(ws, rr, [(ws.cell(row=rr, column=ci).value, widths[ci - 1]) for ci in [2, 5, 12, 23, 29, 30, 31]], min_h=58, max_h=132)
    for rr in range(start, end + 1):
        c = ws.cell(row=rr, column=32)
        if c.value:
            style_link(c, c.value)
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:AG{end}"
    page_setup(ws, "1:4")
    ws.print_area = f"A1:AG{end}"
    ws.conditional_formatting.add(f"F{start}:F{end}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=max(2500, max([int(c.get("distance") or 0) for c in cands] + [1])), color=BLUE, showValue=True))
    ws.conditional_formatting.add(f"K{start}:K{end}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=max([int(c.get("capacity") or 0) for c in cands] + [1]), color=TEAL, showValue=True))
    ws.conditional_formatting.add(f"AB{start}:AB{end}", ColorScaleRule(start_type="num", start_value=1, start_color="C6EFCE", mid_type="num", mid_value=3, mid_color="FFEB9C", end_type="num", end_value=5, end_color="FFC7CE"))
    for val, bg, fg in [("高", "FDE2E1", RED), ("中高", "FFE8CC", AMBER), ("中", "FFF7D6", "8A6A00"), ("低", "E3F4E8", GREEN)]:
        ws.conditional_formatting.add(f"AA{start}:AA{end}", FormulaRule(formula=[f'$AA{start}="{val}"'], fill=fill(bg), font=Font(name=FONT, size=9, bold=True, color=fg)))
    ws.conditional_formatting.add(f"N{start}:N{end}", FormulaRule(formula=[f'$N{start}="是"'], fill=fill("E3F4E8"), font=Font(name=FONT, size=9, bold=True, color=GREEN)))
    ws.conditional_formatting.add(f"Z{start}:Z{end}", FormulaRule(formula=[f'$Z{start}="无"'], fill=fill("E3F4E8")))
    ws.conditional_formatting.add(f"Z{start}:Z{end}", FormulaRule(formula=[f'$Z{start}<>"无"'], fill=fill("FFF5E6")))

    # ------------------------------------------------------------------
    # 02 图表
    # ------------------------------------------------------------------
    ws = wb.create_sheet("02_图表")
    ws.sheet_properties.tabColor = TEAL
    title_block(ws, 6, "02 · 组合视角：规模 × 区位 × 运费", "数据来自 01 表；横条图便于比较现有供应商与备选。")
    section_bar(ws, 5, 6, "图表数据（可编辑，图表自动引用）", color=TEAL)
    headers = ["企业", f"{product_short}产能\n{unit}", "运费中值\n元/吨", "运距\nkm", "工艺", "备注"]
    widths = [26, 16, 14, 12, 12, 38]
    header_row(ws, 6, headers, widths, height=38, fill_color=NAVY2)
    data_rows = []
    for c in [base] + cands:
        data_rows.append([c.get("short"), c.get("capacity"), round((float(c.get("freight_low") or 0) + float(c.get("freight_high") or 0)) / 2),
                          c.get("distance"), c.get("process"), "现有供应商" if c is base else c.get("same_source", "")])
    table_rows(ws, 7, data_rows, widths, zebra=True, align_map={1: "left"})
    for rr in range(7, 7 + len(data_rows)):
        fit_row(ws, rr, [(ws.cell(row=rr, column=1).value, widths[0])], min_h=20)
    ws.freeze_panes = "A7"
    page_setup(ws, "1:6")
    ws.print_area = f"A1:V65"

    ch = BarChart()
    ch.type = "bar"
    ch.style = 10
    ch.title = f"{product_short}产能（{unit}）"
    ch.x_axis.title = unit
    data = Reference(ws, min_col=2, min_row=6, max_row=6 + len(data_rows))
    cats = Reference(ws, min_col=1, min_row=7, max_row=6 + len(data_rows))
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    ch.legend = None
    ch.height = 9.2
    ch.width = 17
    ch.dLbls = DataLabelList()
    ch.dLbls.showVal = True
    ch.series[0].graphicalProperties.solidFill = TEAL
    ws.add_chart(ch, "H7")

    ch2 = BarChart()
    ch2.type = "bar"
    ch2.style = 12
    ch2.title = f"估算运费中值（元/吨，到{buyer_short}）"
    ch2.x_axis.title = "元/吨"
    data = Reference(ws, min_col=3, min_row=6, max_row=6 + len(data_rows))
    ch2.add_data(data, titles_from_data=True)
    ch2.set_categories(cats)
    ch2.legend = None
    ch2.height = 9.2
    ch2.width = 17
    ch2.dLbls = DataLabelList()
    ch2.dLbls.showVal = True
    ch2.series[0].graphicalProperties.solidFill = AMBER
    ws.add_chart(ch2, "H26")

    # process pie
    pie_row = 21
    section_bar(ws, pie_row, 6, "工艺阵营（家数）", color=VIOLET)
    ws.cell(row=pie_row + 1, column=1, value="工艺").font = font(9, True, WHITE)
    ws.cell(row=pie_row + 1, column=2, value="家数").font = font(9, True, WHITE)
    for c in (1, 2):
        ws.cell(row=pie_row + 1, column=c).fill = fill(VIOLET)
        ws.cell(row=pie_row + 1, column=c).border = BORDER
        ws.cell(row=pie_row + 1, column=c).alignment = align("center", "center", True)
    procs = [p["name"] for p in M.get("processes") or []]
    all_s = [base] + cands
    for i, p in enumerate(procs, start=pie_row + 2):
        ws.cell(row=i, column=1, value=p).font = font(9)
        ws.cell(row=i, column=1).border = BORDER
        ws.cell(row=i, column=2, value=sum(1 for x in all_s if x.get("process") == p)).font = font(9, True)
        ws.cell(row=i, column=2).border = BORDER
        ws.cell(row=i, column=2).alignment = align("center", "center", True)
    if procs:
        pie = PieChart()
        pie.title = "工艺阵营（含现有供应商）"
        data = Reference(ws, min_col=2, min_row=pie_row + 1, max_row=pie_row + 1 + len(procs))
        cats = Reference(ws, min_col=1, min_row=pie_row + 2, max_row=pie_row + 1 + len(procs))
        pie.add_data(data, titles_from_data=True)
        pie.set_categories(cats)
        pie.height = 8
        pie.width = 12
        pie.dLbls = DataLabelList()
        pie.dLbls.showPercent = True
        pie.dLbls.showVal = True
        ws.add_chart(pie, "H45")

    # ------------------------------------------------------------------
    # 03 横向对比
    # ------------------------------------------------------------------
    ws = wb.create_sheet("03_横向对比")
    ws.sheet_properties.tabColor = AMBER
    cols = 2 + 1 + len(cands)
    widths = [11, 21, 34] + [26] * len(cands)
    title_block(ws, cols, f"03 · 横向对比：现有供应商 vs {len(cands)} 家备选",
                "绿色=相对更优、红色=相对更弱（仅运距/运费/产能/风险分四项）；其余为事实描述。")
    header = ["类别", "对比维度", f"{base.get('short')}（现有）"] + [c.get("short") for c in cands]
    header_row(ws, 5, header, widths, height=44)
    c = ws.cell(row=6, column=1, value="—")
    c.font = font(8, False, MUTED); c.fill = fill(LIGHT); c.alignment = align("center", "center", True); c.border = BORDER
    c = ws.cell(row=6, column=2, value=f"（工艺 · 产能 · 运距）")
    c.font = font(8, False, MUTED); c.fill = fill(LIGHT); c.alignment = align("center", "center", True); c.border = BORDER
    for j, s in enumerate([base] + cands, start=3):
        c = ws.cell(row=6, column=j, value=f'{s.get("process")} · {s.get("capacity")} {unit} · {s.get("distance")} km')
        c.font = font(8, False, MUTED); c.fill = fill(LIGHT); c.alignment = align("center", "center", True); c.border = BORDER
    ws.row_dimensions[6].height = 20

    cmp = [
        ("身份与区位", [
            ("企业全称", lambda s: s.get("name")),
            ("集团 / 实控人", lambda s: f'{s.get("group","")}\n{s.get("controller","")}'),
            ("股权性质", lambda s: s.get("ownership")),
            ("生产地", lambda s: f'{s.get("province","")} · {s.get("city","")}\n{s.get("site","")}'),
            (f"距{buyer_short}（km）", lambda s: s.get("distance")),
            ("运输方式", lambda s: s.get("transport")),
            ("估算运费中值（元/吨）", lambda s: round((float(s.get("freight_low") or 0) + float(s.get("freight_high") or 0)) / 2)),
        ]),
        ("产能与产品", [
            ("工艺", lambda s: s.get("process")),
            (f"{product_short}产能（{unit}）", lambda s: s.get("capacity")),
            (grade_label, lambda s: f'{s.get("grade","")} / {s.get("brand","")}'),
            ("主营产品", lambda s: s.get("products")),
        ]),
        ("规模与财务", [
            (f"{fy}营收（亿元）", lambda s: s.get("revenue")),
            (f"{fy}净利（亿元）", lambda s: s.get("profit")),
            ("员工（公开）", lambda s: s.get("employees")),
        ]),
        ("资质与规格", [
            (qlabel, lambda s: "是" if s.get("qualified") else "否"),
            ("规格兼容性", lambda s: s.get("spec_fit")),
            ("试样要求", lambda s: s.get("sample")),
            ("风险摘要 / 公开动态", lambda s: s.get("risk_summary")),
        ]),
        ("风险与同源", [
            ("风险等级", lambda s: s.get("risk_level")),
            ("风险分（1–5，低=好）", lambda s: s.get("risk_score")),
            ("同源关系", lambda s: s.get("same_source")),
            ("同源说明", lambda s: s.get("same_note")),
            ("需人工核实", lambda s: "；".join(s.get("verify") or [])),
        ]),
        ("差异判断", [
            (f"与{base.get('short','现有')}差异", lambda s: s.get("diff")),
        ]),
    ]
    center_labels = {f"距{buyer_short}（km）", "估算运费中值（元/吨）", f"{product_short}产能（{unit}）", "风险分（1–5，低=好）",
                     qlabel, "规格兼容性", "试样要求", "风险等级", "工艺", f"{fy}营收（亿元）", f"{fy}净利（亿元）"}
    r = 7
    cat_ranges = []
    numeric_rows = {}
    flag_rows = {}
    for cat, items in cmp:
        cat_start = r
        for label, fn in items:
            c = ws.cell(row=r, column=1, value=cat)
            c.font = font(9, True, WHITE); c.fill = fill(NAVY2); c.alignment = align("center", "center", True); c.border = BORDER
            c = ws.cell(row=r, column=2, value=label)
            c.font = font(9, True, NAVY); c.fill = fill(LIGHT); c.alignment = align("left", "center", True, 1); c.border = BORDER
            for j, s in enumerate([base] + cands, start=3):
                c = ws.cell(row=r, column=j, value=fn(s))
                c.font = font(9); c.border = BORDER
                c.alignment = align("center" if label in center_labels else "left", "top", True)
            fit_row(ws, r, [(ws.cell(row=r, column=j).value, widths[j - 1]) for j in range(2, cols + 1)], min_h=22, max_h=120)
            if label in (f"距{buyer_short}（km）", "估算运费中值（元/吨）", f"{product_short}产能（{unit}）", "风险分（1–5，低=好）"):
                numeric_rows[label] = r
            if label == qlabel:
                flag_rows[label] = r
            r += 1
        cat_ranges.append((cat, cat_start, r - 1))
        if r - 1 > cat_start:
            ws.merge_cells(start_row=cat_start, start_column=1, end_row=r - 1, end_column=1)
    for cat, a, b in cat_ranges:
        c = ws.cell(row=a, column=1, value=cat)
        c.font = font(9.5, True, WHITE); c.fill = fill(NAVY2)
        c.alignment = Alignment(horizontal="center", vertical="center", text_rotation=(90 if b > a else 0), wrap_text=True)
        c.border = BORDER
    ws.freeze_panes = "D7"
    page_setup(ws, "5:6")
    ws.print_area = f"A1:{get_column_letter(cols)}{r-1}"
    for label, rr in numeric_rows.items():
        rng = f"C{rr}:{get_column_letter(cols)}{rr}"
        if label in (f"距{buyer_short}（km）", "估算运费中值（元/吨）", "风险分（1–5，低=好）"):
            rule = ColorScaleRule(start_type="min", start_color="C6EFCE", mid_type="percentile", mid_value=50, mid_color="FFEB9C", end_type="max", end_color="FFC7CE")
        else:
            rule = ColorScaleRule(start_type="min", start_color="FFC7CE", mid_type="percentile", mid_value=50, mid_color="FFEB9C", end_type="max", end_color="C6EFCE")
        ws.conditional_formatting.add(rng, rule)
    if qlabel in flag_rows:
        rr = flag_rows[qlabel]
        ws.conditional_formatting.add(f"C{rr}:{get_column_letter(cols)}{rr}", FormulaRule(formula=[f'C{rr}="是"'], fill=fill("E3F4E8")))

    # ------------------------------------------------------------------
    # 04 试样计划
    # ------------------------------------------------------------------
    ws = wb.create_sheet("04_试样计划")
    ws.sheet_properties.tabColor = VIOLET
    spec_headers = ["企业", "工艺", qshort, "首样量", "检测周期", "判定标准"] + [i["label"] for i in specs] + ["责任", "状态", "备注"]
    spec_widths = [18, 8, 7, 16, 12, 24] + [10] * len(specs) + [9, 10, 28]
    title_block(ws, len(spec_headers), "04 · 规格兼容性与试样计划",
                "首样 20–50 kg × 3 批；判定依据客户内控 + 国标/行标；状态列可下拉更新。")
    header_row(ws, 5, spec_headers, spec_widths, height=46)
    smap = {"COA": "凭COA", "试样": "需试样", "审核": "现场审核"}
    rows = []
    for c in cands:
        st = status.get(c["id"], {})
        note = []
        if c.get("new_plant"):
            note.append("新装置")
        if not c.get("qualified"):
            note.append("非准入：全项检验")
        if c.get("same_source") and c["same_source"] != "无":
            note.append("同源：" + c["same_source"])
        rows.append([c.get("short"), c.get("process"), "是" if c.get("qualified") else "否", "20–50 kg × 3 批", "7–10 个工作日",
                     "客户内控 + 国标/行标"] + [smap.get(st.get(i["id"]), "需试样") for i in specs] + ["质量部", "待取样", "；".join(note)])
    table_rows(ws, 6, rows, spec_widths, zebra=True, align_map={1: "left", 6: "left", len(spec_headers): "left"})
    for rr in range(6, 6 + len(rows)):
        fit_row(ws, rr, [(ws.cell(row=rr, column=1).value, spec_widths[0]), (ws.cell(row=rr, column=len(spec_headers)).value, spec_widths[-1])], min_h=34, max_h=64)
    ws.freeze_panes = "B6"
    ws.auto_filter.ref = f"A5:{get_column_letter(len(spec_headers))}{5 + len(rows)}"
    page_setup(ws, "1:5")
    ws.print_area = f"A1:{get_column_letter(len(spec_headers))}{5 + len(rows)}"
    status_col = get_column_letter(len(spec_headers) - 1)
    ws.conditional_formatting.add(f"{status_col}6:{status_col}{5 + len(rows)}", FormulaRule(formula=[f'${status_col}6="合格"'], fill=fill("E3F4E8"), font=Font(name=FONT, size=9, bold=True, color=GREEN)))
    ws.conditional_formatting.add(f"{status_col}6:{status_col}{5 + len(rows)}", FormulaRule(formula=[f'${status_col}6="不合格"'], fill=fill("FDE2E1"), font=Font(name=FONT, size=9, bold=True, color=RED)))
    ws.conditional_formatting.add(f"{status_col}6:{status_col}{5 + len(rows)}", FormulaRule(formula=[f'${status_col}6="检测中"'], fill=fill("FFF5E6"), font=Font(name=FONT, size=9, color=AMBER)))
    dv = DataValidation(type="list", formula1='"待取样,已取样,检测中,合格,不合格"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{status_col}6:{status_col}{5 + len(rows)}")
    for ci, item in enumerate(specs, start=7):
        col = get_column_letter(ci)
        ws.conditional_formatting.add(f"{col}6:{col}{5 + len(rows)}", FormulaRule(formula=[f'${col}6="需试样"'], fill=fill("FFF5E6")))
        ws.conditional_formatting.add(f"{col}6:{col}{5 + len(rows)}", FormulaRule(formula=[f'${col}6="凭COA"'], fill=fill("E3F4E8")))
        ws.conditional_formatting.add(f"{col}6:{col}{5 + len(rows)}", FormulaRule(formula=[f'${col}6="现场审核"'], fill=fill("E8EFFF")))

    # ------------------------------------------------------------------
    # 05 同源关系
    # ------------------------------------------------------------------
    ws = wb.create_sheet("05_同源关系")
    ws.sheet_properties.tabColor = RED
    headers = ["组", "母公司 / 实控人", "主体", "关系", "影响", "需核实", "标记"]
    widths = [22, 30, 38, 22, 52, 34, 12]
    title_block(ws, len(headers), "05 · 同源与股权关系", "同源不等于不能采购，但“同时入选”会高估第二供应源的独立性。")
    header_row(ws, 5, headers, widths, height=34)
    rows = [[g.get("group"), g.get("parent"), "；".join(g.get("entities") or []), g.get("relation"), g.get("implication"), g.get("verify"), g.get("flag")] for g in same]
    table_rows(ws, 6, rows, widths, zebra=True, align_map={1: "left", 2: "left", 3: "left", 5: "left", 6: "left"})
    for rr in range(6, 6 + len(rows)):
        fit_row(ws, rr, [(ws.cell(row=rr, column=c).value, widths[c - 1]) for c in [1, 2, 3, 5, 6]], min_h=44, max_h=110)
    ws.freeze_panes = "A6"
    page_setup(ws, "1:5")
    ws.print_area = f"A1:G{5 + len(rows)}"
    for val, bg, fg in [("同源", "FDE2E1", RED), ("同集团", "FFE8CC", AMBER), ("纵向竞争", "FFE8CC", AMBER), ("非同源", "E3F4E8", GREEN), ("单一", "EEF3F8", MUTED), ("多厂区", "E8EFFF", BLUE)]:
        ws.conditional_formatting.add(f"G6:G{5 + len(rows)}", FormulaRule(formula=[f'$G6="{val}"'], fill=fill(bg), font=Font(name=FONT, size=9, bold=True, color=fg)))

    # ------------------------------------------------------------------
    # 06 风险台账
    # ------------------------------------------------------------------
    ws = wb.create_sheet("06_风险台账")
    ws.sheet_properties.tabColor = "8A4B08"
    headers = ["企业", "类型", "公开事实", "影响", "来源", "链接", "状态"]
    widths = [16, 9, 56, 36, 24, 46, 9]
    title_block(ws, len(headers), "06 · 公开风险与证据台账", "只列可追溯的公开事实；“待核实”不作为结论使用。")
    header_row(ws, 5, headers, widths, height=32)
    rows = []
    for r in risks:
        s = src.get(r.get("source_id")) or {}
        rows.append([r.get("company"), r.get("type"), r.get("fact"), r.get("impact"), s.get("name") or r.get("source", ""), s.get("url") or r.get("url", ""), r.get("status", "待核实")])
    table_rows(ws, 6, rows, widths, zebra=True, align_map={1: "left", 3: "left", 4: "left", 5: "left", 6: "left"})
    for rr in range(6, 6 + len(rows)):
        c = ws.cell(row=rr, column=6)
        if c.value:
            style_link(c, c.value)
        fit_row(ws, rr, [(ws.cell(row=rr, column=c).value, widths[c - 1]) for c in [1, 3, 4, 5]], min_h=32, max_h=86)
    ws.freeze_panes = "A6"
    ws.auto_filter.ref = f"A5:G{5 + len(rows)}"
    page_setup(ws, "1:5")
    ws.print_area = f"A1:G{5 + len(rows)}"
    for val, bg, fg in [("财务", "FDE2E1", RED), ("装置", "FFF5E6", AMBER), ("资质", "E8EFFF", BLUE), ("同源", "F2ECFF", VIOLET), ("行业", "EEF3F8", MUTED)]:
        ws.conditional_formatting.add(f"B6:B{5 + len(rows)}", FormulaRule(formula=[f'$B6="{val}"'], fill=fill(bg), font=Font(name=FONT, size=9, bold=True, color=fg)))
    ws.conditional_formatting.add(f"G6:G{5 + len(rows)}", FormulaRule(formula=['$G6="已公开"'], fill=fill("E3F4E8"), font=Font(name=FONT, size=9, color=GREEN)))
    ws.conditional_formatting.add(f"G6:G{5 + len(rows)}", FormulaRule(formula=['$G6="待核实"'], fill=fill("FFF5E6"), font=Font(name=FONT, size=9, color=AMBER)))

    # ------------------------------------------------------------------
    # 07 行动清单
    # ------------------------------------------------------------------
    ws = wb.create_sheet("07_行动清单")
    ws.sheet_properties.tabColor = GREEN
    headers = ["阶段", "动作", "说明", "责任", "产出", "状态", "目标完成"]
    widths = [14, 22, 66, 14, 34, 10, 12]
    title_block(ws, len(headers), f"07 · {M.get('action_window')}双源推进清单", "动作是“验证与建源”，不是“替换现有供应商”。")
    header_row(ws, 5, headers, widths, height=32)
    rows = []
    for i, a in enumerate(actions):
        rows.append([a.get("phase"), a.get("action"), a.get("detail"), a.get("owner"), a.get("output"), "未开始", a.get("deadline") or f"T+{30 * (i // 2 + 1)}"])
    table_rows(ws, 6, rows, widths, zebra=True, align_map={1: "center", 3: "left", 5: "left"})
    for rr in range(6, 6 + len(rows)):
        fit_row(ws, rr, [(ws.cell(row=rr, column=3).value, widths[2])], min_h=36, max_h=80)
    ws.freeze_panes = "A6"
    page_setup(ws, "1:5")
    ws.print_area = f"A1:G{5 + len(rows)}"
    dv2 = DataValidation(type="list", formula1='"未开始,进行中,已完成,已搁置"', allow_blank=True)
    ws.add_data_validation(dv2)
    dv2.add(f"F6:F{5 + len(rows)}")
    for val, bg, fg in [("已完成", "E3F4E8", GREEN), ("进行中", "FFF5E6", AMBER), ("已搁置", "EEF3F8", MUTED)]:
        ws.conditional_formatting.add(f"F6:F{5 + len(rows)}", FormulaRule(formula=[f'$F6="{val}"'], fill=fill(bg), font=Font(name=FONT, size=9, bold=True, color=fg)))

    # ------------------------------------------------------------------
    # 08 数据来源
    # ------------------------------------------------------------------
    ws = wb.create_sheet("08_数据来源")
    ws.sheet_properties.tabColor = "64748B"
    headers = ["序号", "来源", "链接", "证据等级", "日期/口径"]
    widths = [6, 66, 62, 9, 14]
    title_block(ws, len(headers), "08 · 数据来源与证据等级", "A = 公告/政府；B = 协会/券商/主流媒体；C = 企业官网/百科。")
    header_row(ws, 5, headers, widths, height=30)
    rows = [[i, s.get("name"), s.get("url"), s.get("grade"), s.get("date")] for i, s in enumerate(sources, start=1)]
    table_rows(ws, 6, rows, widths, zebra=True, align_map={1: "center", 2: "left", 3: "left"})
    for rr in range(6, 6 + len(rows)):
        c = ws.cell(row=rr, column=3)
        if c.value:
            style_link(c, c.value)
        fit_row(ws, rr, [(ws.cell(row=rr, column=2).value, widths[1])], min_h=26, max_h=48)
    ws.freeze_panes = "A6"
    ws.auto_filter.ref = f"A5:E{5 + len(rows)}"
    page_setup(ws, "1:5")
    ws.print_area = f"A1:E{5 + len(rows)}"
    for val, bg, fg in [("A", "E3F4E8", GREEN), ("B", "E8EFFF", BLUE), ("C", "FFF5E6", AMBER)]:
        ws.conditional_formatting.add(f"D6:D{5 + len(rows)}", FormulaRule(formula=[f'$D6="{val}"'], fill=fill(bg), font=Font(name=FONT, size=9, bold=True, color=fg)))

    # ------------------------------------------------------------------
    # 09 尽调清单
    # ------------------------------------------------------------------
    ws = wb.create_sheet("09_尽调清单")
    ws.sheet_properties.tabColor = "0F766E"
    dd = M.get("dd_items") or []
    headers = ["企业", "工艺", qshort, "本次重点尽调"] + [d["label"] for d in dd] + ["备注"]
    widths = [18, 7, 7, 34] + [11] * len(dd) + [42]
    title_block(ws, len(headers), "09 · 供应商尽调 / 准入清单",
                f"{len(dd)} 项证据逐家打勾；状态列可下拉选择。未取得原件或可核验复印件前，不进入合格供方名录。")
    header_row(ws, 5, headers, widths, height=44)
    rows = []
    for c in cands:
        focus = "；".join((c.get("verify") or [])[:2])
        rows.append([c.get("short"), c.get("process"), "是" if c.get("qualified") else "否", focus] + ["未启动"] * len(dd) +
                    [f'{c.get("capacity_note","")}；同源：{c.get("same_source","无")}'])
    table_rows(ws, 6, rows, widths, zebra=True, align_map={1: "left", 4: "left", len(headers): "left"})
    for rr in range(6, 6 + len(rows)):
        fit_row(ws, rr, [(ws.cell(row=rr, column=4).value, widths[3]), (ws.cell(row=rr, column=len(headers)).value, widths[-1])], min_h=42, max_h=90)
    ws.freeze_panes = "E6"
    ws.auto_filter.ref = f"A5:{get_column_letter(len(headers))}{5 + len(rows)}"
    page_setup(ws, "1:5")
    ws.print_area = f"A1:{get_column_letter(len(headers))}{5 + len(rows)}"
    dv3 = DataValidation(type="list", formula1='"未启动,资料中,已核,不适用"', allow_blank=True)
    ws.add_data_validation(dv3)
    dv3.add(f"E6:{get_column_letter(4 + len(dd))}{5 + len(rows)}")
    for val, bg, fg in [("已核", "E3F4E8", GREEN), ("资料中", "FFF5E6", AMBER), ("不适用", "EEF3F8", MUTED), ("未启动", "F8FAFC", MUTED)]:
        ws.conditional_formatting.add(f"E6:{get_column_letter(4 + len(dd))}{5 + len(rows)}", FormulaRule(formula=[f'E6="{val}"'], fill=fill(bg), font=Font(name=FONT, size=9, bold=(val == "已核"), color=fg)))

    # finalize
    wb.active = 0
    os.makedirs(os.path.dirname(os.path.abspath(out_path)) or ".", exist_ok=True)
    wb.save(out_path)
    return out_path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("data")
    ap.add_argument("out")
    a = ap.parse_args()
    p = build(a.data, a.out)
    print(f"OK  XLSX  {p}  {os.path.getsize(p)/1024:.1f} KB")


if __name__ == "__main__":
    main()
