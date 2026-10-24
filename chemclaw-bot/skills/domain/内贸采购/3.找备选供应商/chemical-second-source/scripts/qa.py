#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Post-build QA for the HTML + Excel pair.

Usage:
  python qa.py outdir [--node C:\\path\\node.exe]
"""
from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None


def read(p):
    with open(p, "r", encoding="utf-8") as f:
        return f.read()


def check_html(path, node):
    errors, warnings = [], []
    html = read(path)
    for placeholder in ("__DATA__", "__APP__", "__STYLE__", "__TITLE__", "__XLSX_NAME__"):
        if placeholder in html:
            errors.append(f"HTML 仍有占位符 {placeholder}")
    if "</html>" not in html:
        errors.append("HTML 缺少 </html>")
    for must in ('id="baselineCard"', 'id="pool"', 'id="cmptable"', 'id="scatter"', 'id="sources"'):
        if must not in html:
            errors.append(f"HTML 缺少 {must}")
    m = re.search(r"window\.__DSH_DATA__=(.*?);</script>", html, re.S)
    if not m:
        errors.append("HTML 中找不到 window.__DSH_DATA__")
        return errors, warnings, None
    try:
        data = json.loads(m.group(1))
    except Exception as e:
        errors.append(f"HTML 内嵌 JSON 解析失败: {e}")
        return errors, warnings, None
    if len(data.get("candidates") or []) < 2:
        errors.append("HTML 内嵌数据 candidates < 2")
    if not data.get("meta", {}).get("product"):
        errors.append("HTML 内嵌 meta.product 为空")

    # JS syntax check: write the second <script> block to a temp file
    if node:
        blocks = re.findall(r"<script>(.*?)</script>", html, re.S)
        if len(blocks) < 2:
            errors.append("HTML 脚本块不足 2 个")
        else:
            tmp = tempfile.NamedTemporaryFile("w", suffix=".js", delete=False, encoding="utf-8")
            tmp.write(blocks[-1])
            tmp.close()
            try:
                r = subprocess.run([node, "--check", tmp.name], capture_output=True, text=True, timeout=30)
                if r.returncode != 0:
                    errors.append("node --check 失败: " + (r.stderr or r.stdout).strip()[:600])
            except FileNotFoundError:
                warnings.append("找不到 node，跳过 JS 语法检查")
            except Exception as e:
                warnings.append(f"node 检查异常: {e}")
            finally:
                try:
                    os.unlink(tmp.name)
                except OSError:
                    pass

            # Runtime smoke: execute the inline app against a tiny DOM stub.
            smoke = os.path.join(os.path.dirname(os.path.abspath(__file__)), "smoke_js.js")
            if os.path.exists(smoke):
                try:
                    r = subprocess.run([node, smoke, path], capture_output=True, text=True, timeout=45)
                    if r.returncode != 0:
                        errors.append("运行时 smoke 失败: " + (r.stderr or r.stdout).strip()[:600])
                    else:
                        print("  ·", (r.stdout or "").strip()[:160])
                except Exception as e:
                    warnings.append(f"运行时 smoke 异常: {e}")
    else:
        warnings.append("未提供 --node，跳过 JS 语法检查")
    return errors, warnings, data


def check_xlsx(path, data):
    errors, warnings = [], []
    if load_workbook is None:
        return ["缺少 openpyxl，无法检查工作簿"], warnings
    wb = load_workbook(path, read_only=True, data_only=False)
    expected = ["00_说明", "01_备选池", "02_图表", "03_横向对比", "04_试样计划", "05_同源关系", "06_风险台账", "07_行动清单", "08_数据来源", "09_尽调清单"]
    if wb.sheetnames != expected:
        errors.append(f"工作表顺序/名称不符: {wb.sheetnames}")
    for name in expected:
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        if ws.max_row < 5 or ws.max_column < 3:
            errors.append(f"{name} 内容过少: {ws.max_row} 行 × {ws.max_column} 列")
    if data:
        n = len(data.get("candidates") or [])
        ws = wb["01_备选池"]
        if ws.max_row < n + 4:
            errors.append(f"01_备选池 行数不足：{ws.max_row} < {n + 4}")
        ws = wb["04_试样计划"]
        if ws.max_row < n + 5:
            errors.append(f"04_试样计划 行数不足：{ws.max_row} < {n + 5}")
    wb.close()
    return errors, warnings


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("outdir")
    ap.add_argument("--node", default="")
    a = ap.parse_args()
    node = a.node or shutil.which("node") or ""
    htmls = [x for x in os.listdir(a.outdir) if x.endswith(".html") and "对比工作台" in x]
    xlsxs = [x for x in os.listdir(a.outdir) if x.endswith(".xlsx") and "对比工作簿" in x]
    if not htmls or not xlsxs:
        sys.exit(f"在 {a.outdir} 找不到生成的 HTML/Excel")
    hp = os.path.join(a.outdir, sorted(htmls)[-1])
    xp = os.path.join(a.outdir, sorted(xlsxs)[-1])
    e1, w1, data = check_html(hp, node)
    e2, w2 = check_xlsx(xp, data)
    errors, warnings = e1 + e2, w1 + w2
    for w in warnings:
        print("[WARN ]", w)
    for e in errors:
        print("[ERROR]", e)
    if errors:
        sys.exit(1)
    print(f"OK  QA  HTML {os.path.basename(hp)}  ·  XLSX {os.path.basename(xp)}  ·  {len(data.get('candidates') or [])} candidates")


if __name__ == "__main__":
    main()
