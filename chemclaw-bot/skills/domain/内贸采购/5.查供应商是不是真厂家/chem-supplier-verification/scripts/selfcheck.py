# -*- coding: utf-8 -*-
"""交付前自检：数据 + Excel + HTML 三层一起看。"""
import argparse
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _common import check_doc, load_doc  # noqa: E402

OK, BAD = "[ok]  ", "[FAIL]"


def main():
    ap = argparse.ArgumentParser(description="交付前自检")
    ap.add_argument("input", help="input.json")
    ap.add_argument("--outdir", default=".", help="产物目录")
    a = ap.parse_args()

    doc = load_doc(a.input)
    meta = doc["meta"]
    slug = meta.get("slug", "supplier-verification")
    fails = []

    def chk(cond, msg):
        print((OK if cond else BAD) + " " + msg)
        if not cond:
            fails.append(msg)

    # ---------- 1. 数据
    errors, warnings = check_doc(doc)
    for w in warnings:
        print("[warn] " + w)
    chk(not errors, "数据校验（%d 个 error）" % len(errors))
    if errors:
        for e in errors[:10]:
            print("       - " + e)

    ncase = len(doc["cases"])
    nev = sum(len(c["direct"]) + len(c["indirect"]) + len(c["excluded"]) for c in doc["cases"])
    n_direct = sum(len(c["direct"]) for c in doc["cases"])

    # ---------- 2. HTML
    hp = os.path.join(a.outdir, slug + ".html")
    if not os.path.exists(hp):
        chk(False, "HTML 存在：%s" % hp)
    else:
        h = open(hp, encoding="utf-8").read()
        chk(len(h) > 20000, "HTML 体积（%d 字节）" % len(h.encode("utf-8")))
        chk(h.count('<article class="ev') == nev, "证据卡数量 %d = %d" % (h.count('<article class="ev'), nev))
        chk(h.count('class="casebtn"') + h.count('class="casebtn on"') == ncase, "案例标签 %d 个" % ncase)
        chk("__CASES__" not in h and "__ALIAS__" not in h, "JS 占位符已替换")
        chk("const CASES = [" in h and "const ALIAS = [" in h, "JS 案例表与别名表")
        chk("</script>" in h and "<script>" in h, "内嵌脚本闭合")
        chk(meta.get("product", "")[:2] in h, "标题含产品名")
        miss = [c["id"] for c in doc["cases"] if 'id="case-%s"' % c["id"] not in h]
        chk(not miss, "每个案例都有独立区块%s" % ("（缺 %s）" % miss if miss else ""))
        chk(meta.get("as_of", "") in h, "数据截止日期已写入")
        # 未渲染的占位
        chk(not re.search(r"\{[a-z_]+\(", h), "无未渲染的模板表达式")

    # ---------- 3. Excel
    xp = os.path.join(a.outdir, slug + ".xlsx")
    if not os.path.exists(xp):
        print("[warn] 没找到 Excel（%s）；--html-only 模式可忽略" % xp)
    else:
        try:
            import openpyxl
        except ImportError:
            print("[warn] 没装 openpyxl，跳过 Excel 检查")
        else:
            wb = openpyxl.load_workbook(xp)
            chk(len(wb.sheetnames) >= 9, "工作表 %d 个" % len(wb.sheetnames))
            for need in ("01_核验总表", "05_评分模型", "06_待索取材料"):
                chk(need in wb.sheetnames, "含工作表 %s" % need)
            ws = wb["01_核验总表"]
            names = [ws.cell(row=5 + i, column=2).value for i in range(ncase)]
            chk(all(names), "01 总表 %d 行案例齐全" % ncase)
            f = ws.cell(row=5, column=5).value or ""
            chk(str(f).startswith("='05_评分模型'!"), "01 分数引用 05 表：%s" % f)
            w5 = wb["05_评分模型"]
            total_row = 5 + len(doc["dimensions"])
            chk(str(w5.cell(row=total_row, column=2).value or "").startswith("=SUM"),
                "05 合计行公式在 %d 行" % total_row)
            chk(len([1 for i in range(1, ncase + 1)
                     if w5.cell(row=4, column=2 + i).value]) == ncase, "05 表头含全部案例")
            nmat = len(doc["materials"])
            chk(wb["06_待索取材料"].max_row >= 4 + nmat, "06 材料表 %d 项" % nmat)

    # ---------- 4. 结论
    print("")
    if fails:
        print("[fail] %d 项未通过" % len(fails))
        return 1
    print("[pass] 全部检查通过：%d 案例 / %d 条证据（直接 %d）/ %d 项材料"
          % (ncase, nev, n_direct, len(doc["materials"])))
    return 0


if __name__ == "__main__":
    sys.exit(main())
