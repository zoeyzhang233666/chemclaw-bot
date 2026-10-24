# -*- coding: utf-8 -*-
"""render_xlsx.py — 把 data.json 渲染为 12 页签的 Excel 工作簿。

用法：
    python render_xlsx.py <data.json> [输出.xlsx]

依赖：openpyxl（唯一外部依赖）。数据缺失的页签会自动跳过而不是报错，
便于在调研尚未完成时先生成草稿。
"""
import datetime
import json
import os
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:  # pragma: no cover
    print("需要 openpyxl：pip install openpyxl", file=sys.stderr)
    raise

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import ddlib as L  # noqa: E402

INK = "101C26"
BRAND = "12394A"
BRAND2 = "1D5A70"
BRASS = "9A6C2F"
BRASS_L = "F7F0E3"
PAPER = "FBF9F3"
LINE = "E0DACD"
RED = "A02C2C"
RED_BG = "FBECEB"
AMBER = "A4661A"
AMBER_BG = "FBF0E2"
GREEN = "1C6B4A"
GREEN_BG = "E8F4EE"
GRAY = "6B7178"
GRAY_BG = "EEF0F1"
BLUE = "2B5C8A"
BLUE_BG = "EAF0F7"
FONT = "Microsoft YaHei"

thin = Side(style="thin", color=LINE)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

STATUS_COLOR = {"green": (GREEN, GREEN_BG), "blue": (BLUE, BLUE_BG), "gray": (GRAY, GRAY_BG),
                "amber": (AMBER, AMBER_BG), "red": (RED, RED_BG)}
CERT_COLOR = {"有效": (GREEN, GREEN_BG), "已到期": (RED, RED_BG), "临期": (AMBER, AMBER_BG)}


def fill(c):
    return PatternFill("solid", fgColor=c)


def title_block(ws, ncols, title, sub):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=title)
    c.font = Font(name=FONT, size=15, bold=True, color="FFFFFF")
    c.fill = fill(BRAND)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 34
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    c = ws.cell(row=2, column=1, value=sub)
    c.font = Font(name=FONT, size=9, color="6D7A84")
    c.fill = fill("F1EDE3")
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 22


def header_row(ws, row, headers, widths):
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(name=FONT, size=9.5, bold=True, color="FFFFFF")
        c.fill = fill(BRAND)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30


def write_rows(ws, start_row, rows, widths, aligns=None, wrap_cols=None):
    aligns = aligns or {}
    wrap_cols = wrap_cols or set()
    for r, row in enumerate(rows, start_row):
        maxlines = 1
        for i, v in enumerate(row, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.font = Font(name=FONT, size=9.5, color="243247")
            c.border = BORDER
            c.alignment = Alignment(horizontal=aligns.get(i, "left"), vertical="top", wrap_text=True)
            if r % 2 == 0:
                c.fill = fill(PAPER)
            if i in wrap_cols and isinstance(v, str):
                cw = widths[i - 1] if i - 1 < len(widths) else 20
                import math
                lines = sum(max(1, math.ceil(len(seg) / max(cw * 1.05, 6))) for seg in v.split("\n"))
                maxlines = max(maxlines, lines)
        ws.row_dimensions[r].height = min(max(20, maxlines * 14.5 + 6), 210)


def setup_sheet(ws, ncols, freeze_row, landscape=True):
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.freeze_panes = f"A{freeze_row}"
    ws.print_title_rows = f"{max(1, freeze_row-1)}:{max(1, freeze_row-1)}"
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = "供应商准入尽调底稿 · 数据截止见首页 · 第 &P 页 / 共 &N 页"
    ws.oddFooter.center.size = 8
    ws.oddFooter.center.color = "6D7A84"


def walk(data):
    """收集所有区块，按 name / type 建索引。"""
    by_name, by_type = {}, {}
    for sec in data.get("sections", []):
        for b in sec.get("blocks", []):
            if b.get("name"):
                by_name[b["name"]] = b
            by_type.setdefault(b.get("type"), []).append(b)
            if b.get("type") in ("grid2", "grid3"):
                for c in b.get("cards", []):
                    if c.get("table") and c["table"].get("name"):
                        by_name[c["table"]["name"]] = c["table"]
    return by_name, by_type


def build(data):
    ctx = L.Ctx(data)
    meta = data.get("meta", {})
    verdict = next((b for b in data.get("sections", [{}])[0].get("blocks", [])
                    if b.get("type") == "verdict"), {})
    gates = next((b for b in data.get("sections", [{}])[0].get("blocks", [])
                  if b.get("type") == "gates"), {})
    by_name, by_type = walk(data)
    risks = ctx.risks
    certs = data.get("certificates") or []
    events = (by_type.get("timeline") or [{}])[0].get("items", [])
    materials = next((b for b in by_type.get("checklist", []) if b.get("kind") == "mat"), {})
    sites = next((b for b in by_type.get("checklist", []) if b.get("kind") == "site"), {})
    master = data.get("masters") or {}

    wb = Workbook()

    # ---------------------------------------------------------- 00 使用说明
    ws = wb.active
    ws.title = "00 使用说明"
    ws.sheet_properties.tabColor = BRASS
    setup_sheet(ws, 6, 1, landscape=True)
    for col, w in zip("ABCDEF", [3, 24, 46, 30, 18, 18]):
        ws.column_dimensions[col].width = w

    ws.merge_cells("B2:E2")
    c = ws["B2"]
    c.value = meta.get("title", "供应商准入尽调底稿")
    c.font = Font(name=FONT, size=22, bold=True, color=INK)
    ws.row_dimensions[2].height = 40
    ws.merge_cells("B3:E3")
    c = ws["B3"]
    c.value = f'{meta.get("target", "")} · {meta.get("product", "")} · 内贸采购准入'
    c.font = Font(name=FONT, size=12, color=BRAND2)
    ws.row_dimensions[3].height = 24

    meta_rows = [
        ("底稿编号", meta.get("doc_no", "")),
        ("被调查主体", meta.get("target", "")),
        ("采购品类", meta.get("product", "")),
        ("数据截止", meta.get("as_of", "")),
        ("编制 / 复核", f'{meta.get("author", "")} / {meta.get("reviewer", "")}'),
        ("版本", meta.get("version", "")),
        ("准入结论", verdict.get("headline", "")),
    ] + [[k, v] for k, v in (meta.get("extra_rows") or [])]
    r = 5
    for k, v in meta_rows:
        a = ws.cell(row=r, column=2, value=k)
        a.font = Font(name=FONT, size=10, bold=True, color=BRAND)
        a.fill = fill(BRASS_L)
        a.border = BORDER
        a.alignment = Alignment(vertical="center", indent=1)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
        b = ws.cell(row=r, column=3, value=v)
        b.font = Font(name=FONT, size=10, color=INK)
        b.alignment = Alignment(vertical="center", indent=1, wrap_text=True)
        for cc in range(3, 6):
            ws.cell(row=r, column=cc).border = BORDER
        ws.row_dimensions[r].height = 26 if len(str(v)) < 40 else 40
        r += 1

    r += 1
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    c = ws.cell(row=r, column=2, value="一、状态定义（全文统一口径）")
    c.font = Font(name=FONT, size=11, bold=True, color=INK)
    r += 1
    for i, h in enumerate(["状态", "含义", "证据要求"], 2):
        c = ws.cell(row=r, column=i, value=h)
        c.font = Font(name=FONT, size=9.5, bold=True, color="FFFFFF")
        c.fill = fill(BRAND)
        c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)
    for cc in (4, 5):
        ws.cell(row=r, column=cc).border = BORDER
        ws.cell(row=r, column=cc).fill = fill(BRAND)
    r += 1
    req = {"ok": "年报 / 公告 / 章程 / ESG 报告原文",
           "pub": "至少两个独立公开来源交叉",
           "none": "以供应商材料或监管原始查询替代",
           "todo": "对应材料清单 / 实地核实编号",
           "risk": "具体文件 + 页码 / 条款"}
    for key, st in ctx.status.items():
        c = ws.cell(row=r, column=2, value=st["label"])
        fg, bg = STATUS_COLOR.get(st.get("color", "gray"), (GRAY, GRAY_BG))
        c.font = Font(name=FONT, size=9.5, bold=True, color=fg)
        c.fill = fill(bg)
        c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center")
        for col, val in ((3, st.get("desc", "")), (4, req.get(key, ""))):
            cc = ws.cell(row=r, column=col, value=val)
            cc.font = Font(name=FONT, size=9.5, color="243247")
            cc.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
            cc.border = BORDER
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)
        ws.cell(row=r, column=5).border = BORDER
        ws.row_dimensions[r].height = 30
        r += 1

    r += 1
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    c = ws.cell(row=r, column=2, value="二、使用边界")
    c.font = Font(name=FONT, size=11, bold=True, color=INK)
    r += 1
    notes = meta.get("usage_notes") or [
        "1. 本底稿只做“准入判断”，不预测该供应商未来的交付表现、价格走势或履约概率。",
        "2. 所有事实应能在“11 证据索引”中定位到原文；测算类数据须在正文标注口径。",
        "3. “无公开信息”是查询结论，不是风险结论；该类项目一律转入材料清单与实地核实。",
        f"4. 数据截止 {meta.get('as_of', '')}，此后发生的公告、处罚、诉讼均未纳入。",
        "5. 本工作簿为内部采购准入资料，不得对外提供。",
    ]
    for t in notes:
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        c = ws.cell(row=r, column=2, value=t)
        c.font = Font(name=FONT, size=9.5, color="2B3944")
        c.alignment = Alignment(vertical="top", wrap_text=True, indent=1)
        ws.row_dimensions[r].height = 30
        r += 1

    # ---------------------------------------------------------- 01 结论
    ws = wb.create_sheet("01 尽调结论")
    ws.sheet_properties.tabColor = RED
    widths = [4, 16, 20, 40, 22, 18]
    setup_sheet(ws, 6, 4)
    title_block(ws, 6, "01　尽调结论与准入建议",
                "结论只回答“能不能进、以什么条件进”；不给出履约概率，不预测未来交付表现")
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.merge_cells("B4:E4")
    c = ws["B4"]
    c.value = "准入结论：" + verdict.get("headline", "")
    c.font = Font(name=FONT, size=18, bold=True, color="FFFFFF")
    c.fill = fill(BRAND)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[4].height = 46

    ws.merge_cells("B5:F5")
    c = ws["B5"]
    c.value = verdict.get("text", "").replace("<b>", "").replace("</b>", "")
    c.font = Font(name=FONT, size=10, color=INK)
    c.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
    c.fill = fill(BRASS_L)
    c.border = BORDER
    ws.row_dimensions[5].height = 52

    r = 7
    ws.cell(row=r, column=2, value="准入参数").font = Font(name=FONT, size=11, bold=True, color=INK)
    r += 1
    for row in verdict.get("rows", []):
        ws.cell(row=r, column=2, value=row.get("k")).font = Font(name=FONT, size=9.5, bold=True, color=BRAND)
        ws.cell(row=r, column=2).fill = fill(BRASS_L)
        ws.cell(row=r, column=2).border = BORDER
        ws.cell(row=r, column=2).alignment = Alignment(vertical="center", indent=1)
        c = ws.cell(row=r, column=3, value=row.get("v"))
        c.font = Font(name=FONT, size=9.5, bold=True, color=INK)
        c.border = BORDER
        c.alignment = Alignment(vertical="center", indent=1)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
        ws.cell(row=r, column=4).border = BORDER
        ws.row_dimensions[r].height = 24
        r += 1

    r += 1
    ws.cell(row=r, column=2, value="准入闸门").font = Font(name=FONT, size=11, bold=True, color=INK)
    r += 1
    header_row(ws, r, ["", "闸门", "核心要求", "", "未关闭的后果", ""], widths)
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
    ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=6)
    r += 1
    for g in gates.get("items", []):
        c = ws.cell(row=r, column=2, value=f'{g.get("code")} · {g.get("name")}')
        c.font = Font(name=FONT, size=9.5, bold=True, color=BRAND)
        c.fill = fill(BLUE_BG)
        c.border = BORDER
        c.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
        c = ws.cell(row=r, column=3, value=g.get("text"))
        c.font = Font(name=FONT, size=9.5, color="243247")
        c.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
        for cc in (3, 4):
            ws.cell(row=r, column=cc).border = BORDER
        ws.row_dimensions[r].height = 44
        r += 1

    r += 1
    ws.cell(row=r, column=2, value="维度评分（证据评分，不是履约概率）").font = Font(name=FONT, size=11, bold=True, color=INK)
    r += 1
    header_row(ws, r, ["", "维度", "评分", "依据", "状态", ""], widths)
    r += 1
    for s in verdict.get("scores", []):
        ws.cell(row=r, column=2, value=s.get("name")).font = Font(name=FONT, size=9.5, bold=True, color=INK)
        ws.cell(row=r, column=3, value=s.get("value")).font = Font(name=FONT, size=11, bold=True, color=INK)
        ws.cell(row=r, column=2).alignment = Alignment(vertical="center", indent=1)
        ws.cell(row=r, column=3).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r, column=4, value=s.get("note")).font = Font(name=FONT, size=9.5, color="2B3944")
        ws.cell(row=r, column=4).alignment = Alignment(vertical="center", wrap_text=True, indent=1)
        st = ctx.status.get(s.get("state", "ok"), ctx.status["ok"])
        c = ws.cell(row=r, column=5, value=st["label"])
        fg, bg = STATUS_COLOR.get(st.get("color", "gray"), (GRAY, GRAY_BG))
        c.font = Font(name=FONT, size=9, bold=True, color=fg)
        c.fill = fill(bg)
        c.alignment = Alignment(horizontal="center", vertical="center")
        for cc in range(2, 6):
            ws.cell(row=r, column=cc).border = BORDER
        ws.row_dimensions[r].height = 30
        r += 1

    donts = []
    for blk0 in data.get("sections", [{}])[0].get("blocks", []):
        if blk0.get("type") in ("grid2", "grid3"):
            for card in blk0.get("cards", []):
                donts.extend(card.get("notes") or [])
    if donts:
        ws.cell(row=r, column=2, value="三条明确的“不建议”").font = Font(name=FONT, size=11, bold=True, color=RED)
        r += 1
        for t in donts:
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
            c = ws.cell(row=r, column=2, value="· " + str(t).replace("<b>", "").replace("</b>", ""))
            c.font = Font(name=FONT, size=9.5, color=RED)
            c.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
            c.fill = fill(RED_BG)
            ws.row_dimensions[r].height = 30
            r += 1

    # ---------------------------------------------------------- 02 主体
    blk = by_name.get("subject")
    if blk:
        ws = wb.create_sheet("02 主体信息")
        ws.sheet_properties.tabColor = GREEN
        widths = [6, 18, 62, 11, 8, 26]
        setup_sheet(ws, 6, 5)
        title_block(ws, 6, "02　主体与经营状态", "企业名称、信用代码、注册资本与实缴、注册地与生产地、变更历史")
        header_row(ws, 4, ["序号", "核查项", "披露值 / 结论", "状态", "证据", "核验方式 / 备注"], widths)
        rows = [[i, it.get("k"), it.get("v"), ctx.status.get(it.get("state", "ok"), {}).get("label", ""),
                 it.get("evid", ""), it.get("note", "")]
                for i, it in enumerate(blk["items"], 1)]
        write_rows(ws, 5, rows, widths, aligns={1: "center", 4: "center", 5: "center"}, wrap_cols={3, 6})
        for row in ws[f"D5:D{4+len(rows)}"]:
            for c in row:
                st = next((v for v in ctx.status.values() if v["label"] == c.value), None)
                if st:
                    fg, bg = STATUS_COLOR.get(st.get("color", "gray"), (GRAY, GRAY_BG))
                    c.font = Font(name=FONT, size=9, bold=True, color=fg)
                    c.fill = fill(bg)
        ws.auto_filter.ref = f"A4:F{4+len(rows)}"

    # ---------------------------------------------------------- 03 股权
    sh, rel = by_name.get("shareholders"), by_name.get("related")
    if sh or rel:
        ws = wb.create_sheet("03 股权与关联")
        ws.sheet_properties.tabColor = BRAND2
        widths = [6, 30, 16, 12, 16, 22]
        setup_sheet(ws, 6, 4)
        title_block(ws, 6, "03　股权结构、实际控制人与关联企业",
                    "控制链、持股比例、关联交易及其对本次采购的影响")
        if sh:
            header_row(ws, 4, ["序号"] + [c["h"] for c in sh["columns"]], widths)
            rows = [[i] + list(r) for i, r in enumerate(sh["rows"], 1)]
            write_rows(ws, 5, rows, widths, aligns={1: "center"}, wrap_cols={2})
            r = 5 + len(rows) + 1
        else:
            r = 4
        if rel:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
            c = ws.cell(row=r, column=1, value="主要关联企业及其与本次采购的相关性")
            c.font = Font(name=FONT, size=11, bold=True, color=INK)
            r += 1
            header_row(ws, r, ["序号"] + [c["h"] for c in rel["columns"]], widths)
            r += 1
            rows = [[i] + list(x) for i, x in enumerate(rel["rows"], 1)]
            write_rows(ws, r, rows, widths, aligns={1: "center"}, wrap_cols={2, 3, 4})
            for row in ws[f"F{r}:F{r+len(rows)-1}"]:
                for c in row:
                    fg, bg = {"高": (RED, RED_BG), "中": (AMBER, AMBER_BG), "低": (GRAY, GRAY_BG)}.get(
                        str(c.value), (GRAY, GRAY_BG))
                    c.font = Font(name=FONT, size=9, bold=True, color=fg)
                    c.fill = fill(bg)

    # ---------------------------------------------------------- 04 生产
    cap, ind = by_name.get("capacity"), by_name.get("industry")
    if cap:
        ws = wb.create_sheet("04 生产与供应")
        ws.sheet_properties.tabColor = BRASS
        widths = [6, 20, 46, 11, 9, 34]
        setup_sheet(ws, 6, 4)
        title_block(ws, 6, "04　实际生产与供应能力", "设计产能、实际产量、利用率、工艺路线与改造计划")
        header_row(ws, 4, ["序号", "项目", "数据 / 结论", "状态", "证据", "说明"], widths)
        rows = [[i, it.get("k"), it.get("v"), ctx.status.get(it.get("state", "ok"), {}).get("label", "待核验"),
                 it.get("evid", ""), it.get("note", "")]
                for i, it in enumerate(cap["items"], 1)]
        write_rows(ws, 5, rows, widths, aligns={1: "center", 4: "center", 5: "center"}, wrap_cols={3, 6})
        r = 4 + len(rows) + 2
        if ind:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
            c = ws.cell(row=r, column=1, value="行业与市场背景")
            c.font = Font(name=FONT, size=11, bold=True, color=INK)
            r += 1
            header_row(ws, r, ["序号"] + [c["h"] for c in ind["columns"]] + ["", ""], widths)
            ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
            r += 1
            for i, row in enumerate(ind["rows"], 1):
                ws.cell(row=r, column=1, value=i).alignment = Alignment(horizontal="center")
                ws.cell(row=r, column=2, value=row[0])
                ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
                ws.cell(row=r, column=3, value=row[1])
                ws.cell(row=r, column=5, value=row[2] if len(row) > 2 else "")
                ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=6)
                for cc in range(1, 7):
                    c = ws.cell(row=r, column=cc)
                    c.font = Font(name=FONT, size=9.5, color="243247")
                    c.border = BORDER
                    c.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
                ws.row_dimensions[r].height = 22
                r += 1

    # ---------------------------------------------------------- 05 资质
    if certs:
        ws = wb.create_sheet("05 资质证照")
        ws.sheet_properties.tabColor = AMBER
        widths = [6, 26, 13, 13, 38, 16, 11, 10, 20]
        setup_sheet(ws, 9, 4)
        title_block(ws, 9, f"05　公开可查的资质证照（全量 {len(certs)} 项）",
                    f"状态按 {meta.get('as_of', '')} 计算；临期=180 天内到期；未取得换证凭证")
        header_row(ws, 4, ["序号", "资质名称", "有效期起", "有效期止", "许可范围 / 适用产品",
                           "持有人", "到期状态", "相关性", "备注"], widths)
        rows = []
        for c_ in certs:
            label, cls = L.cert_state(c_.get("end"), ctx.as_of)
            state = {"expired": "已到期", "soon": "临期"}.get(cls, "有效")
            note = {"已到期": "须索取新证", "临期": "90 天内复核"}.get(state, "")
            rows.append([c_.get("no"), c_.get("name"), c_.get("start"), c_.get("end"),
                         c_.get("scope"), c_.get("holder"), state,
                         c_.get("relevance", "相关"), note])
        order = {"已到期": 0, "临期": 1, "有效": 2}
        rows.sort(key=lambda x: order.get(x[6], 3))
        write_rows(ws, 5, rows, widths,
                   aligns={1: "center", 3: "center", 4: "center", 7: "center", 8: "center"},
                   wrap_cols={2, 5, 9})
        for row in ws[f"G5:G{4+len(rows)}"]:
            for c in row:
                fg, bg = CERT_COLOR.get(str(c.value), (GREEN, GREEN_BG))
                c.font = Font(name=FONT, size=9, bold=True, color=fg)
                c.fill = fill(bg)
        for row in ws[f"H5:H{4+len(rows)}"]:
            for c in row:
                fg, bg = {"核心": (RED, RED_BG), "相关": (AMBER, AMBER_BG)}.get(str(c.value), (GRAY, GRAY_BG))
                c.font = Font(name=FONT, size=9, bold=True, color=fg)
                c.fill = fill(bg)
        ws.auto_filter.ref = f"A4:I{4+len(rows)}"

    # ---------------------------------------------------------- 06 风险
    if risks:
        ws = wb.create_sheet("06 风险清单")
        ws.sheet_properties.tabColor = RED
        widths = [7, 8, 11, 10, 34, 60, 40, 40, 14, 12]
        setup_sheet(ws, 10, 4)
        title_block(ws, 10, "06　经营与合规风险清单",
                    "每一项风险均对应具体证据；等级 = 对本次采购的影响程度，不是发生概率")
        header_row(ws, 4, ["编号", "维度", "风险等级", "证据强度", "风险标题",
                           "事实与证据", "采购影响", "缓释动作", "证据 ID", "状态"], widths)
        rows = []
        for x in risks:
            status = "待核验" if x.get("status") == "todo" else "已识别（待缓释）"
            rows.append([x.get("id"), x.get("dim"), L.LEVEL_TEXT.get(x.get("level"), "中风险"),
                         L.STRENGTH_TEXT.get(x.get("strength"), "中"), x.get("title"),
                         x.get("fact"), x.get("impact"), x.get("action"),
                         "、".join(x.get("evid") or []), status])
        write_rows(ws, 5, rows, widths,
                   aligns={1: "center", 2: "center", 3: "center", 4: "center", 10: "center"},
                   wrap_cols={5, 6, 7, 8})
        for row in ws[f"C5:C{4+len(rows)}"]:
            for c in row:
                fg, bg = {"高风险": (RED, RED_BG), "中风险": (AMBER, AMBER_BG),
                          "低风险": (GRAY, GRAY_BG)}.get(str(c.value), (GRAY, GRAY_BG))
                c.font = Font(name=FONT, size=9, bold=True, color=fg)
                c.fill = fill(bg)
        r = 4 + len(rows) + 2
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10)
        c = ws.cell(row=r, column=1, value="风险分布")
        c.font = Font(name=FONT, size=11, bold=True, color=INK)
        r += 1
        stats = [("高风险", sum(1 for x in risks if x.get("level") == "high"), RED),
                 ("中风险", sum(1 for x in risks if x.get("level") == "mid"), AMBER),
                 ("低风险", sum(1 for x in risks if x.get("level") == "low"), GRAY),
                 ("合计", len(risks), BRAND)]
        for i, (label, n, color) in enumerate(stats):
            col = 1 + i * 2
            c = ws.cell(row=r, column=col, value=label)
            c.font = Font(name=FONT, size=9, bold=True, color="FFFFFF")
            c.fill = fill(color)
            c.border = BORDER
            c.alignment = Alignment(horizontal="center", vertical="center")
            ws.merge_cells(start_row=r, start_column=col, end_row=r, end_column=col + 1)
            ws.cell(row=r, column=col + 1).border = BORDER
            ws.cell(row=r, column=col + 1).fill = fill(color)
        r += 1
        for i, (label, n, color) in enumerate(stats):
            col = 1 + i * 2
            c = ws.cell(row=r, column=col, value=n)
            c.font = Font(name=FONT, size=14, bold=True, color=color)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = BORDER
            ws.merge_cells(start_row=r, start_column=col, end_row=r, end_column=col + 1)
            ws.cell(row=r, column=col + 1).border = BORDER

    # ---------------------------------------------------------- 07 事件
    if events:
        ws = wb.create_sheet("07 重大事件")
        ws.sheet_properties.tabColor = BRAND
        widths = [6, 13, 74, 16, 8]
        setup_sheet(ws, 5, 4)
        title_block(ws, 5, "07　近 12 个月重大事件", "按被调查主体公告列表梳理；重大事件对应公告见证据索引")
        header_row(ws, 4, ["序号", "日期", "事件", "类型", "证据"], widths)
        rows = [[i, e.get("date"), e.get("text"), e.get("tag", ""), e.get("evid", "")]
                for i, e in enumerate(events, 1)]
        write_rows(ws, 5, rows, widths, aligns={1: "center", 2: "center", 4: "center", 5: "center"},
                   wrap_cols={3})
        for row in ws[f"D5:D{4+len(rows)}"]:
            for c in row:
                if str(c.value) in ("业绩", "资本开支"):
                    c.font = Font(name=FONT, size=9, bold=True, color=RED)
                    c.fill = fill(RED_BG)

    # ---------------------------------------------------------- 08 / 09 清单
    for sheet_name, blk, kind, cols, statuses, who, tab in (
        ("08 材料清单", materials, "mat",
         [6, 8, 50, 54, 8, 14, 12, 12],
         '"未提交,已提交,已核验,不适用"', "采购员", GREEN),
        ("09 实地核实", sites, "site",
         [6, 18, 44, 56, 8, 14, 12, 12],
         '"未开始,已完成,需复核,不适用"', "采购员 + 质检", BLUE),
    ):
        if not blk:
            continue
        ws = wb.create_sheet(sheet_name)
        ws.sheet_properties.tabColor = tab
        setup_sheet(ws, 8, 4)
        label = "材料名称" if kind == "mat" else "事项"
        title = blk.get("title") or sheet_name[3:]
        title_block(ws, 8, f"{sheet_name[:2]}　{title}", "把“查不到”变成“问得到”，把“问得到”变成“看得见”")
        header_row(ws, 4, ["编号", "类别" if kind == "mat" else "核查区域", label, "核对要点",
                           "优先级", "责任岗位", "计划完成", "状态"], cols)
        rows = [[it.get("code"), it.get("category") if kind == "mat" else it.get("area"),
                 it.get("name"), it.get("note"), it.get("priority", "中"), who, "", "未提交" if kind == "mat" else "未开始"]
                for it in blk["items"]]
        write_rows(ws, 5, rows, cols, aligns={1: "center", 2: "center", 5: "center", 7: "center", 8: "center"},
                   wrap_cols={3, 4})
        for row in ws[f"E5:E{4+len(rows)}"]:
            for c in row:
                fg, bg = {"高": (RED, RED_BG), "中": (AMBER, AMBER_BG), "低": (GRAY, GRAY_BG)}.get(
                    str(c.value), (GRAY, GRAY_BG))
                c.font = Font(name=FONT, size=9, bold=True, color=fg)
                c.fill = fill(bg)
        dv = DataValidation(type="list", formula1=statuses, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"H5:H{4+len(rows)}")
        ws.auto_filter.ref = f"A4:H{4+len(rows)}"

    # ---------------------------------------------------------- 10 交叉核对
    cc = by_name.get("crosscheck")
    if cc:
        ws = wb.create_sheet("10 材料交叉核对")
        ws.sheet_properties.tabColor = BRASS
        widths = [6, 22, 44, 50, 34, 12, 30]
        setup_sheet(ws, 7, 4)
        title_block(ws, 7, "10　供应商上传材料 × 平台公开数据 交叉核对表",
                    "供应商上传后逐字段比对；不一致项进入差异清单，须书面澄清后方可放行")
        header_row(ws, 4, ["序号", "材料", "将核对的字段", "平台侧公开参考值",
                           "供应商上传值", "一致性", "差异说明 / 处理"], widths)
        rows = [[i, r[0], r[1], r[2], "", "待比对", ""] for i, r in enumerate(cc["rows"], 1)]
        write_rows(ws, 5, rows, widths, aligns={1: "center", 6: "center"}, wrap_cols={2, 3, 4, 7})
        dv = DataValidation(type="list", formula1='"待比对,一致,不一致,材料缺失"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"F5:F{4+len(rows)}")

    # ---------------------------------------------------------- 11 证据
    if ctx.ev:
        ws = wb.create_sheet("11 证据索引")
        ws.sheet_properties.tabColor = GRAY
        widths = [7, 40, 14, 12, 62, 40]
        setup_sheet(ws, 6, 4)
        title_block(ws, 6, f"11　证据索引（{len(ctx.ev)} 份一手文件）",
                    "全部来源见原文链接；引用时请标注证据 ID")
        header_row(ws, 4, ["ID", "文件名称", "来源", "披露日期", "原文链接", "引用内容"], widths)
        rows = [[e["id"], e.get("name"), e.get("source"), e.get("date"), e.get("url"), e.get("note")]
                for e in ctx.d.get("evidence", [])]
        write_rows(ws, 5, rows, widths, aligns={1: "center", 3: "center", 4: "center"}, wrap_cols={2, 5, 6})
        for row in ws[f"E5:E{4+len(rows)}"]:
            for c in row:
                if c.value and str(c.value).startswith("http"):
                    c.hyperlink = c.value
                    c.font = Font(name=FONT, size=9, color="2B5C8A", underline="single")
        ws.auto_filter.ref = f"A4:F{4+len(rows)}"

    for w in wb.worksheets:
        w.print_area = f"A1:{get_column_letter(w.max_column)}{w.max_row}"
    return wb


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 1
    src = sys.argv[1]
    data = json.load(open(src, encoding="utf-8"))
    wb = build(data)
    out = sys.argv[2] if len(sys.argv) > 2 else os.path.splitext(src)[0] + ".xlsx"
    wb.save(out)
    print(f"xlsx -> {out}  ({os.path.getsize(out):,} bytes, {len(wb.worksheets)} sheets)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
