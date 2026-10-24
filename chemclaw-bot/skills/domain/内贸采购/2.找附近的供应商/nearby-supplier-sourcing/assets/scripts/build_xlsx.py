# -*- coding: utf-8 -*-
"""build_xlsx.py — 由 dataset.json 生成 8 表交付台账

用法:
    python build_xlsx.py --dataset dataset.json --out 供应商运距清单.xlsx

工作表（固定结构，勿改列序）:
    00 封面与说明 / 01 供应商总表 / 02 匹配度评分 / 03 证据链
    04 距离测算   / 05 运输与合规 / 06 数据缺口   / 07 询价清单
"""
import argparse, json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import DataBarRule, CellIsRule

NAVY = "0D2438"; NAVY2 = "15334C"; TEAL = "0D7A6B"; AMBER = "A85A08"; SLATE = "5C6B7D"
LINE = "E4E9EF"; BAND = "F5F8FA"; WARN = "FBF1E2"; OK = "E7F3F0"; BAD = "FDF1F0"; GREY = "B9C4CF"

def F(sz=10, b=False, c="1F2937", it=False):
    return Font(name="Microsoft YaHei", size=sz, bold=b, color=c, italic=it)

def fill(c):
    return PatternFill("solid", fgColor=c)

THIN = Side(style="thin", color=LINE)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOP = Alignment(vertical="top", wrap_text=True, indent=1)
MID = Alignment(vertical="center", wrap_text=True, indent=1)
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)


def sheet_title(ws, row, cols, text, sub=""):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    c = ws.cell(row=row, column=1, value=text)
    c.font = F(13, True, NAVY); c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[row].height = 26
    if sub:
        ws.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=cols)
        c = ws.cell(row=row + 1, column=1, value=sub)
        c.font = F(8.5, False, "8B98A8"); c.alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[row + 1].height = 17
        return row + 2
    return row + 1


def header_row(ws, row, headers, widths=None):
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = F(9, True, "FFFFFF"); c.fill = fill(NAVY2)
        c.alignment = CEN; c.border = BOX
    ws.row_dimensions[row].height = 30
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
    return row + 1


def data_row(ws, row, vals, widths_fmt=None, band=False, h=None):
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=row, column=j, value=v)
        c.font = F(9, False, "1F2937"); c.alignment = TOP; c.border = BOX
        if band:
            c.fill = fill(BAND)
    if h:
        ws.row_dimensions[row].height = h
    return row + 1


def tier_style(ws, cell, tier):
    col = {"A": OK, "B": WARN, "C": "EEF1F5"}.get(tier, "FFFFFF")
    cell.fill = fill(col)
    cell.font = F(9, True, {"A": "0A5C50", "B": "7A4206", "C": "475569"}.get(tier, "1F2937"))


def short_region(region):
    parts = str(region).replace("　", " ").split()
    prov = parts[0] if parts else ""
    city = parts[1] if len(parts) > 1 else ""
    return prov, city


def gap_impact(g):
    """按关键词推断缺口的影响与动作（通用规则，不依赖具体产品）"""
    rules = [
        (("可供", "供货量", "产量", "库存", "货权", "在库"), "能否按批次稳定供货未定", "索取本月可供量与发货计划（书面）"),
        (("门牌", "地址", "厂址", "坐标", "位置"), "距离与可达性可能偏差", "索取厂区/库房门牌与装卸点定位"),
        (("牌号", "规格", "质检", "等级", "含量", "堆积密度"), "产品是否符合采购规格未定", "索取对应规格质检报告与产品目录"),
        (("运费", "到厂价", "价格", "报价", "口径"), "到厂成本不可比", "要求统一口径报价（吨公里/整车包干）"),
        (("资质", "许可", "危化", "运输", "经营许可"), "合规性与承运能力未定", "核验营业执照、许可与承运资质"),
        (("产能", "装置", "投产", "扩产", "搬迁"), "当前有效产能存疑", "核对最新产量/投产/搬迁公告"),
        (("联系", "电话", "对接", "销售"), "无法触达关键人", "通过总机/园区/同行转介获取对接人"),
    ]
    for kws, impact, action in rules:
        if any(k in g for k in kws):
            return impact, action
    return "信息不完整，影响判断置信度", "首次接触时一并确认"


def build(ds, out):
    wb = Workbook()
    meta = ds["meta"]; prod = meta.get("product", {}); org = meta["origin"]
    items = sorted(ds["items"], key=lambda x: x.get("rank", 99))
    haz = ds.get("hazmat", {})
    dims = ds["match"]["dims"]
    rings = ds["logistics"].get("rings", [100, 300, 500])
    maxr = max(rings)
    gen = meta.get("generated", "")

    # ---------------- 00 封面 ----------------
    ws = wb.active; ws.title = "00 封面与说明"
    ws.sheet_view.showGridLines = False
    for j, w in enumerate([14, 20, 26, 26, 22, 18], 1):
        ws.column_dimensions[get_column_letter(j)].width = w

    ws.merge_cells("A1:F1")
    ws["A1"] = meta.get("title", "供应商寻源清单")
    ws["A1"].font = F(19, True, NAVY); ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 40

    ws.merge_cells("A2:F2")
    ws["A2"] = f'{prod.get("name","")}　|　{prod.get("qty","")}　|　数据基准日 {gen}'
    ws["A2"].font = F(10, False, "66768A"); ws["A2"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[2].height = 22

    r = 4
    info = [
        ("采购标的", prod.get("name", "")),
        ("规格 / 牌号", prod.get("grade", "")),
        ("包装", prod.get("pack", "")),
        ("单次批量", prod.get("qty", "")),
        ("收货地", org.get("label", "")),
        ("坐标", f'{org.get("lat")}, {org.get("lon")}（{org.get("note","")}）'),
        ("运输方式", prod.get("mode", "")),
        ("自提", prod.get("self_pickup", "")),
        ("危化判定", "属于危险化学品" if haz.get("is_hazardous") else "不属于危险化学品（普通货运）"),
    ]
    for k, v in info:
        ws.cell(row=r, column=1, value=k).font = F(9, True, "66768A")
        ws.cell(row=r, column=1).alignment = Alignment(vertical="center", indent=1)
        ws.cell(row=r, column=1).fill = fill("F5F8FA")
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        c = ws.cell(row=r, column=2, value=v)
        c.font = F(9, False, "1F2937"); c.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
        for j in range(1, 7):
            ws.cell(row=r, column=j).border = Border(bottom=Side(style="hair", color=LINE))
        ws.row_dimensions[r].height = 21
        r += 1

    if org.get("is_example"):
        r += 1
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        c = ws.cell(row=r, column=1, value=f'⚠ 示例收货点：{org.get("note","")}　换成真实地址后，全部距离与半径结论需重算。')
        c.font = F(9, True, "7A4206"); c.fill = fill(WARN)
        c.alignment = Alignment(vertical="center", indent=1)
        c.border = BOX
        ws.row_dimensions[r].height = 24
        r += 1

    # KPI
    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    c = ws.cell(row=r, column=1, value="关键数字")
    c.font = F(11, True, NAVY); c.alignment = Alignment(vertical="center", indent=1)
    r += 1
    inmax = [i for i in items if i["road_km"] <= maxr]
    inm = [i for i in items if i["road_km"] <= rings[0]]
    kpis = [
        ("候选企业", f"{len(items)} 家"),
        (f"{maxr} km 内", f"{len(inmax)} 家"),
        (f"{rings[0]} km 内", f"{len(inm)} 家"),
        ("数据缺口", f'{sum(len(i["gaps"]) for i in items)} 项'),
    ]
    for k, v in kpis:
        c = ws.cell(row=r, column=1, value=k); c.font = F(8.5, True, "66768A")
        c.alignment = Alignment(vertical="center", indent=1)
        c2 = ws.cell(row=r, column=2, value=v); c2.font = F(14, True, TEAL)
        c2.alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[r].height = 24
        r += 1

    # 采购路线
    r += 1
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
    c = ws.cell(row=r, column=2, value="建议采购路线")
    c.font = F(11, True, NAVY); c.alignment = Alignment(vertical="center", indent=1)
    r += 1
    for step in ds.get("playbook", []):
        ws.cell(row=r, column=2, value=step.get("step", "")).font = F(9.5, True, TEAL)
        ws.cell(row=r, column=2).alignment = Alignment(vertical="top", indent=1)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
        c = ws.cell(row=r, column=3, value=step.get("text", ""))
        c.font = F(9, False, "1F2937"); c.alignment = Alignment(vertical="top", wrap_text=True, indent=1)
        for j in range(2, 7):
            ws.cell(row=r, column=j).border = Border(bottom=Side(style="hair", color=LINE))
        ws.row_dimensions[r].height = 34
        r += 1

    # 导航
    r += 1
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
    c = ws.cell(row=r, column=2, value="工作表导航")
    c.font = F(11, True, NAVY); c.alignment = Alignment(vertical="center", indent=1)
    r += 1
    nav = [
        ("01 供应商总表", "一企一行：角色、地址、匹配/物流/距离、动作", "01 供应商总表"),
        ("02 匹配度评分", "五个维度的分项打分，可直接改权重复盘", "02 匹配度评分"),
        ("03 证据链", "每条事实的来源、级别、链接（A 交易所/政府 > D 名录）", "03 证据链"),
        ("04 距离测算", "直线/公路/绕行系数/在途时间，含降级标注", "04 距离测算"),
        ("05 运输与合规", "危化判定、资质要求、运费口径（不含金额）", "05 运输与合规"),
        ("06 数据缺口", "每条缺口的影响与下一步动作，带状态下拉", "06 数据缺口"),
        ("07 询价清单", "发起询价用的字段与话术，可打印", "07 询价清单"),
    ]
    for name, desc, link in nav:
        c = ws.cell(row=r, column=2, value=name)
        c.font = F(9, True, TEAL); c.alignment = Alignment(vertical="center", indent=1)
        c.hyperlink = f"#'{link}'!A1"
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
        c2 = ws.cell(row=r, column=3, value=desc)
        c2.font = F(8.8, False, "66768A"); c2.alignment = Alignment(vertical="center", indent=1)
        r += 1

    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    c = ws.cell(row=r, column=1, value="口径与免责")
    c.font = F(11, True, NAVY); c.alignment = Alignment(vertical="center", indent=1)
    r += 1
    for d in ds.get("disclaimer", []):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        c = ws.cell(row=r, column=1, value="· " + d)
        c.font = F(8.5, False, "66768A"); c.alignment = Alignment(vertical="top", wrap_text=True, indent=1)
        ws.row_dimensions[r].height = 30
        r += 1

    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # ---------------- 01 总表 ----------------
    ws = wb.create_sheet("01 供应商总表")
    heads = ["序号", "企业名称", "简称", "角色", "类型", "省份", "城市", "详细地址",
             "工艺/性质", "产能规模", "牌号/品牌", "匹配度", "档位", "公路 km", "直线 km",
             "绕行", "运距等级", "物流分", "证据 A/B", "缺口", "定位精度", "运输建议", "关键动作"]
    widths = [5, 30, 12, 11, 12, 8, 8, 34, 14, 16, 18, 7, 6, 8, 8, 6, 8, 7, 8, 6, 16, 18, 44]
    r = sheet_title(ws, 1, len(heads), "01 供应商总表",
                    f'{prod.get("short","")} · {org.get("short","")} · 按「先档位后距离」排序，匹配与物流分列展示、不加权合并')
    r = header_row(ws, r, heads, widths)
    first = r
    for k, it in enumerate(items, 1):
        prov, city = short_region(it.get("region", ""))
        ev = it.get("evidence", [])
        vals = [k, it["name"], it["short"], it["role"],
                {"factory": "生产厂家", "warehouse": "仓库/交割库", "registered": "仅注册地址", "trader": "贸易商"}.get(it["typeCode"], ""),
                prov, city, it["address"], it.get("process", ""), it.get("capacity", ""), it.get("brand", ""),
                it["matchTotal"], it["tier"], it["road_km"], it["straight_km"], it.get("ratio"),
                it.get("band", ""), it.get("lScore", 0),
                f'A{sum(1 for e in ev if e["level"]=="A")}/B{sum(1 for e in ev if e["level"]=="B")}',
                len(it.get("gaps", [])), it.get("coord_note", ""), it.get("mode", ""), it.get("action", "")]
        r = data_row(ws, r, vals, band=(k % 2 == 0), h=46)
        tier_style(ws, ws.cell(row=r - 1, column=13), it["tier"])
        ws.cell(row=r - 1, column=12).font = F(9, True, NAVY)
    last = r - 1
    ws.freeze_panes = ws.cell(row=first, column=3)
    ws.auto_filter.ref = f"A{first-1}:{get_column_letter(len(heads))}{last}"
    if last >= first:
        ws.conditional_formatting.add(f"N{first}:N{last}",
            DataBarRule(start_type="num", start_value=0, end_type="num", end_value=max(1200, max(i["road_km"] for i in items)),
                        color="1A5F9E", showValue=True))
        ws.conditional_formatting.add(f"R{first}:R{last}",
            DataBarRule(start_type="num", start_value=0, end_type="num", end_value=100, color="0D7A6B", showValue=True))
        ws.conditional_formatting.add(f"T{first}:T{last}",
            CellIsRule(operator="greaterThanOrEqual", formula=["3"], fill=fill(BAD), font=F(9, True, "8C2B23")))
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # ---------------- 02 评分 ----------------
    ws = wb.create_sheet("02 匹配度评分")
    heads = ["序号", "企业", "类型"] + [d["name"] for d in dims] + ["合计", "档位", "说明"]
    widths = [5, 26, 12] + [10] * len(dims) + [7, 6, 52]
    r = sheet_title(ws, 1, len(heads), "02 匹配度评分",
                    "满分 100；仓库/贸易主体在「产品符合」≤30、「供货能力」≤10；档位 A≥85 / B70–84 / C<70")
    r = header_row(ws, r, heads, widths)
    first = r
    for k, it in enumerate(items, 1):
        vals = [k, it["name"], {"factory": "厂家", "warehouse": "仓库", "registered": "注册", "trader": "贸易"}.get(it["typeCode"], "")]
        vals += [it["match"].get(d["name"], 0) for d in dims]
        vals += [it["matchTotal"], it["tier"], it.get("match_note", "")]
        r = data_row(ws, r, vals, band=(k % 2 == 0), h=34)
        tier_style(ws, ws.cell(row=r - 1, column=len(heads) - 1), it["tier"])
        ws.cell(row=r - 1, column=len(heads) - 2).font = F(9, True, NAVY)
    ws.freeze_panes = ws.cell(row=first, column=3)
    ws.auto_filter.ref = f"A{first-1}:{get_column_letter(len(heads))}{r-1}"

    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(heads))
    c = ws.cell(row=r, column=1, value="维度定义")
    c.font = F(10, True, NAVY); c.alignment = Alignment(indent=1)
    r += 1
    for d in dims:
        ws.cell(row=r, column=1, value=d["name"]).font = F(9, True, TEAL)
        ws.cell(row=r, column=2, value=d["max"]).font = F(9, True, NAVY)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=len(heads))
        c = ws.cell(row=r, column=3, value=d.get("desc", ""))
        c.font = F(9, False, "66768A"); c.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
        r += 1

    # ---------------- 03 证据链 ----------------
    ws = wb.create_sheet("03 证据链")
    heads = ["序号", "企业", "级别", "类型", "日期", "事实（claim）", "来源", "链接", "支撑结论"]
    widths = [5, 22, 6, 12, 11, 56, 26, 30, 34]
    r = sheet_title(ws, 1, len(heads), "03 证据链",
                    "A 交易所/政府/上市公司年报　B 企业官网与公告　C 媒体/百科/研报　D 名录与摘要（仅线索）")
    r = header_row(ws, r, heads, widths)
    first = r
    n = 0
    for it in items:
        for e in sorted(it.get("evidence", []), key=lambda x: "ABCD".index(x["level"])):
            n += 1
            r = data_row(ws, r, [n, it["short"], e["level"], e["kind"], e["date"],
                                 e["claim"], e["source"], e["url"], e["supports"]], band=(n % 2 == 0), h=40)
            lv = ws.cell(row=r - 1, column=3)
            lv.fill = fill({"A": OK, "B": "E8F0F9", "C": WARN, "D": BAD}[e["level"]])
            lv.font = F(9, True, {"A": "0A5C50", "B": "1A5F9E", "C": "7A4206", "D": "8C2B23"}[e["level"]])
            lv.alignment = CEN
            lc = ws.cell(row=r - 1, column=8)
            lc.hyperlink = e["url"]; lc.font = F(8.5, False, "1A5F9E")
    ws.freeze_panes = ws.cell(row=first, column=3)
    ws.auto_filter.ref = f"A{first-1}:{get_column_letter(len(heads))}{r-1}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # ---------------- 04 距离 ----------------
    ws = wb.create_sheet("04 距离测算")
    heads = ["序号", "企业", "起点", "终点", "直线 km", "公路 km", "绕行系数", "在途 h",
             "数据来源", "坐标精度", "备注"]
    widths = [5, 22, 26, 34, 9, 9, 9, 8, 12, 22, 26]
    r = sheet_title(ws, 1, len(heads), "04 距离测算",
                    "公路里程为 OSRM 路网规划结果；标“估算”者为直线×系数，不可当路径结果引用")
    r = header_row(ws, r, heads, widths)
    first = r
    for k, it in enumerate(items, 1):
        r = data_row(ws, r, [k, it["short"], org.get("label", ""), it["address"],
                             it["straight_km"], it["road_km"], it.get("ratio"), it.get("hours"),
                             {"osrm": "OSRM 路径", "estimated": "估算（直线×系数）"}.get(it.get("distance_source"), ""),
                             it.get("coord_note", ""), it.get("distance_note", "")], band=(k % 2 == 0), h=32)
    ws.freeze_panes = ws.cell(row=first, column=3)
    ws.auto_filter.ref = f"A{first-1}:{get_column_letter(len(heads))}{r-1}"

    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(heads))
    c = ws.cell(row=r, column=1, value="物流分折算表（公路里程 → 分值，分段线性）")
    c.font = F(10, True, NAVY); c.alignment = Alignment(indent=1)
    r += 1
    ws.cell(row=r, column=1, value="里程 km").font = F(9, True, "66768A")
    ws.cell(row=r, column=2, value="物流分").font = F(9, True, "66768A")
    r += 1
    for km, sc in ds["logistics"].get("anchors", []):
        ws.cell(row=r, column=1, value=km).font = F(9)
        ws.cell(row=r, column=2, value=sc).font = F(9)
        r += 1

    # ---------------- 05 运输与合规 ----------------
    ws = wb.create_sheet("05 运输与合规")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 18
    for j in range(2, 7):
        ws.column_dimensions[get_column_letter(j)].width = 26
    r = sheet_title(ws, 1, 6, "05 运输与合规",
                    "结论 + 依据 + 后果；不含任何运费金额，仅给计费口径")

    def block(title, rows, tint=None):
        nonlocal r
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        c = ws.cell(row=r, column=1, value=title)
        c.font = F(10.5, True, "FFFFFF"); c.fill = fill(NAVY2)
        c.alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[r].height = 24
        r += 1
        for label, text in rows:
            ws.cell(row=r, column=1, value=label).font = F(9, True, "66768A")
            ws.cell(row=r, column=1).alignment = Alignment(vertical="top", indent=1)
            ws.cell(row=r, column=1).fill = fill("F5F8FA")
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
            c = ws.cell(row=r, column=2, value=text)
            c.font = F(9, False, "1F2937"); c.alignment = Alignment(vertical="top", wrap_text=True, indent=1)
            if tint:
                c.fill = fill(tint)
            for j in range(1, 7):
                ws.cell(row=r, column=j).border = Border(bottom=Side(style="hair", color=LINE))
            ws.row_dimensions[r].height = max(22, 16 * (1 + len(str(text)) // 52))
            r += 1

    verdict_tint = BAD if haz.get("is_hazardous") else OK
    block("判定结论", [("结论", haz.get("verdict", "")),
                       ("性质", "危险化学品 / 危险货物" if haz.get("is_hazardous") else "非危险化学品 / 普通货物")], verdict_tint)
    if haz.get("regulatory"):
        block("法规依据", [("依据", x) for x in haz["regulatory"]])
    if haz.get("risks"):
        block("作业风险", [("风险", "；".join(haz["risks"]))])
    if haz.get("transport"):
        block("运输与资质要求", [(f"要求 {i}", x) for i, x in enumerate(haz["transport"], 1)])
    if haz.get("if_spec_changes"):
        block("规格变化提示", [("注意", haz["if_spec_changes"])], WARN)
    block("运费口径（本表不含金额）", [
        ("计费方式", "吨公里 / 整车包干 / 回程车，三者不可混比"),
        ("必须约定", "装卸费、过磅费、雨布苫盖、压车费、超期滞留"),
        ("报价要求", "书面报价须写明计价口径、有效期与油价联动条款"),
        ("到厂价", "拿到运费报价后：到厂价 = 出厂价 + 运费（含税口径需一致）"),
    ])
    if haz.get("sources"):
        block("依据来源", [("来源", "；".join(haz["sources"]))])
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # ---------------- 06 缺口 ----------------
    ws = wb.create_sheet("06 数据缺口")
    heads = ["序号", "企业", "档位", "缺口描述", "影响", "下一步动作", "状态"]
    widths = [5, 22, 6, 52, 24, 30, 12]
    r = sheet_title(ws, 1, len(heads), "06 数据缺口",
                    "未核实的就写未核实：每条缺口都要能落到一个动作，状态列可下拉更新")
    r = header_row(ws, r, heads, widths)
    first = r
    n = 0
    for it in items:
        for g in it.get("gaps", []):
            n += 1
            imp, act = gap_impact(g)
            r = data_row(ws, r, [n, it["short"], it["tier"], g, imp, act, "待确认"], band=(n % 2 == 0), h=34)
            tier_style(ws, ws.cell(row=r - 1, column=3), it["tier"])
    last = r - 1
    if last >= first:
        dv = DataValidation(type="list", formula1='"待确认,已确认,不适用"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"G{first}:G{last}")
        ws.conditional_formatting.add(f"G{first}:G{last}",
            CellIsRule(operator="equal", formula=['"待确认"'], fill=fill(WARN), font=F(9, True, "7A4206")))
    ws.freeze_panes = ws.cell(row=first, column=3)
    ws.auto_filter.ref = f"A{first-1}:{get_column_letter(len(heads))}{last}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # ---------------- 07 询价 ----------------
    ws = wb.create_sheet("07 询价清单")
    heads = ["序号", "企业", "联系人/电话", "匹配/档位", "公路 km", "牌号需求", "可供量（待填）",
             "出厂价（待填）", "到厂价（待填）", "账期（待填）", "状态", "首轮问询要点"]
    widths = [5, 22, 16, 10, 8, 20, 14, 13, 13, 12, 10, 56]
    r = sheet_title(ws, 1, len(heads), "07 询价清单",
                    "按档位排序后发起询价；到厂价口径须一致，未填即空着，不要估算")
    r = header_row(ws, r, heads, widths)
    first = r
    for k, it in enumerate(items, 1):
        r = data_row(ws, r, [k, it["name"], it.get("contact", ""), f'{it["matchTotal"]}/{it["tier"]}',
                             it["road_km"], prod.get("grade", ""), "", "", "", "", "待询价", it.get("ask", "")],
                     band=(k % 2 == 0), h=44)
        tier_style(ws, ws.cell(row=r - 1, column=4), it["tier"])
    last = r - 1
    if last >= first:
        dv = DataValidation(type="list", formula1='"待询价,已报价,已比价,已下单,不适用"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"K{first}:K{last}")
        for col in ("H", "I", "J"):
            for row in range(first, last + 1):
                ws[f"{col}{row}"].number_format = '#,##0'
    ws.freeze_panes = ws.cell(row=first, column=3)
    ws.auto_filter.ref = f"A{first-1}:{get_column_letter(len(heads))}{last}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    wb.save(out)
    print(f"已保存 {out}（8 表，{len(items)} 家企业，{sum(len(i['gaps']) for i in items)} 条缺口）")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dataset", required=True)
    ap.add_argument("--out", default="供应商运距清单.xlsx")
    args = ap.parse_args()
    with open(args.dataset, encoding="utf-8") as f:
        ds = json.load(f)
    build(ds, args.out)


if __name__ == "__main__":
    main()
