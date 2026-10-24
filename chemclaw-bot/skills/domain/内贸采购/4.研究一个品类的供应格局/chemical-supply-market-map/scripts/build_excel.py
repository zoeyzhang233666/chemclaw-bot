#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""build_excel.py —— 由 data bundle 生成多表工作簿（产品无关）

表结构固定，内容全部来自 bundle：
  00_封面  01_核心结论  02_企业产能  03_区域汇总  04_工艺路线  04A_原料与成本
  05_新增项目  05A_净增还原  06_价格历史  06A_价格周期  07_采购关注
  07A_采购策略  08_尽调清单  09_供应商评分  10_数据缺口  11_来源与口径

被 build_all.py 调用，不必单独运行：
  python scripts/build_all.py --dir work/<项目> --excel-only
"""
import datetime
import os

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, LineChart, PieChart, Reference
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import DataPoint, Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.page import PageMargins

import _common as C

MISSING = C.MISSING

# ------------------------------------------------------------------ 主题
NAVY, NAVY2 = "0B1F33", "12324A"
TEAL, TEAL_L, TEAL_M = "0F766E", "ECFDF5", "5EEAD4"
BLUE, BLUE_L = "1D4ED8", "EFF6FF"
AMBER, AMBER_L = "B45309", "FFFBEB"
RED, RED_L = "B91C1C", "FEF2F2"
SLATE, SLATE_L = "475569", "F1F5F9"
LINE, BG, WHITE, INK, BODY, MUTED = "E2E9F1", "F8FAFC", "FFFFFF", "15283D", "40536A", "718096"
FONT = "微软雅黑"

_THIN = Side(style="thin", color=LINE)
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
NO_BORDER = Border()


def fill(hex_):
    return PatternFill("solid", fgColor=hex_)


def fnt(size=10, bold=False, color=BODY, italic=False, underline=None):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic, underline=underline)


def al(h="left", v="center", wrap=True, indent=0):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap, indent=indent)


def page(ws, orientation="landscape", title_rows="1:5"):
    ws.page_setup.orientation = orientation
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = title_rows
    ws.page_margins = PageMargins(left=0.3, right=0.3, top=0.5, bottom=0.5, header=0.2, footer=0.2)
    ws.print_options.horizontalCentered = True
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90


def widths(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w


def band(ws, r1, c1, r2, c2, hex_):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            ws.cell(r, c).fill = fill(hex_)
            ws.cell(r, c).border = NO_BORDER


def banner(ws, title, subtitle, meta, ncols):
    band(ws, 1, 1, 4, ncols, NAVY)
    for r, text, size, color, bold, height in [
            (1, title, 21, WHITE, True, 36), (2, subtitle, 10.5, "C5D2DE", False, 20),
            (3, meta, 9, "8FA6BA", False, 18)]:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
        cell = ws.cell(r, 1, text)
        cell.font = fnt(size, bold, color)
        cell.alignment = al("left", "center", False, 1)
        ws.row_dimensions[r].height = height
    band(ws, 4, 1, 4, ncols, TEAL)
    ws.row_dimensions[4].height = 3


def section(ws, row, text, ncols, note=""):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    cell = ws.cell(row, 1, text)
    cell.font = fnt(12, True, WHITE)
    cell.fill = fill(TEAL)
    cell.alignment = al("left", "center", False, 1)
    ws.row_dimensions[row].height = 24
    if note:
        ws.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=ncols)
        nc = ws.cell(row + 1, 1, note)
        nc.font = fnt(8.5, False, MUTED)
        nc.alignment = al("left", "center", True, 1)
        ws.row_dimensions[row + 1].height = 18
        return row + 2
    return row + 1


def header(ws, row, headers, height=30):
    for i, h in enumerate(headers, 1):
        cell = ws.cell(row, i, h)
        cell.font = fnt(9.5, True, WHITE)
        cell.fill = fill(NAVY2)
        cell.alignment = al("center", "center", True)
        cell.border = BORDER
    ws.row_dimensions[row].height = height


def row_(ws, r, values, formats=None, aligns=None, hex_bg=WHITE, height=30, font=None):
    for i, v in enumerate(values, 1):
        cell = ws.cell(r, i, v)
        cell.font = font or fnt(9.2, False, BODY)
        cell.border = BORDER
        cell.fill = fill(hex_bg)
        if formats and i <= len(formats) and formats[i - 1]:
            cell.number_format = formats[i - 1]
        h = aligns[i - 1] if aligns and i <= len(aligns) and aligns[i - 1] else "left"
        cell.alignment = al(h, "center", True)
    ws.row_dimensions[r].height = height


def note(ws, r, text, ncols, height=34):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
    cell = ws.cell(r, 1, text)
    cell.font = fnt(8.5, False, MUTED)
    cell.alignment = al("left", "top", True, 1)
    cell.fill = fill(BG)
    cell.border = BORDER
    ws.row_dimensions[r].height = height


def zebra(i):
    return WHITE if i % 2 == 0 else BG


def total_row(ws, r, label, ncols, cells):
    """深色合计行：先着色再写值，避免 MergedCell 只读。"""
    band(ws, r, 1, r, ncols, NAVY2)
    cell = ws.cell(r, 1, label)
    cell.font = fnt(9.5, True, WHITE)
    cell.alignment = al("left", "center", False, 1)
    for col, value, fmt, align in cells:
        c2 = ws.cell(r, col, value)
        c2.font = fnt(10, True, TEAL_M)
        c2.alignment = al(align, "center", False)
        if fmt:
            c2.number_format = fmt
    for c in range(1, ncols + 1):
        ws.cell(r, c).border = BORDER
    ws.row_dimensions[r].height = 26


def safe_sheet(name):
    for ch in "\\/:*?[]":
        name = name.replace(ch, "-")
    return name[:31]


# ================================================================== 主入口
def build(bundle, out_path):
    wb = Workbook()
    meta = bundle["meta"]
    product = meta["product"]
    scope = meta["scope"]
    metrics = meta["metrics"]
    short = product.get("short") or product.get("name") or "化工品"
    year = scope.get("base_year") or "—"
    fwd = scope.get("forward_years") or []
    fwd_label = f"{fwd[0]}E" if fwd else "预测年"
    unit = product.get("unit") or "万吨/年"
    price_unit = product.get("price_unit") or "元/吨"
    cutoff = scope.get("data_cutoff") or MISSING

    sheets = []

    def new(title, color):
        ws = wb.create_sheet(safe_sheet(title))
        ws.sheet_properties.tabColor = color
        sheets.append(ws)
        return ws

    # ---------------------------------------------------------- 00 封面
    ws = wb.active
    ws.title = "00_封面"
    ws.sheet_properties.tabColor = NAVY
    sheets.append(ws)
    ws.sheet_view.showGridLines = False
    widths(ws, {"A": 2.5, "B": 20, "C": 18, "D": 18, "E": 18, "F": 18, "G": 18, "H": 2.5})
    band(ws, 1, 1, 4, 8, NAVY)
    ws.row_dimensions[1].height = 18
    ws.row_dimensions[2].height = 46
    ws.row_dimensions[3].height = 26
    ws.row_dimensions[4].height = 22
    for r, text, size, color, bold in [
            (2, f"{short} · 中国供应市场全景", 26, WHITE, True),
            (3, "采购视角 · 产能 / 区域 / 工艺 / 项目 / 价格 / 风险", 12, "C5D2DE", False),
            (4, f"数据截止 {cutoff}  ·  基准年 {year}  ·  覆盖 {scope.get('geography') or '—'}  ·  "
                f"产能：{unit}  ·  价格：{price_unit}", 9.5, "8FA6BA", False)]:
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
        cell = ws.cell(r, 2, text)
        cell.font = fnt(size, bold, color)
        cell.alignment = al("left", "center", False)
    band(ws, 5, 1, 5, 8, TEAL)
    ws.row_dimensions[5].height = 4

    contents = [
        ("01", "核心结论", f"{len(bundle['verdicts'])} 条判断 + 证据 + 采购含义 + 置信度"),
        ("02", "企业产能", f"{len(bundle['companies'])} 集团 / {len(bundle['plants'])} 基地，基准年与 {fwd_label} 产能"),
        ("03", "区域汇总", f"{len(bundle['regions'])} 个生产省份的产能、集中度与贸易活跃度"),
        ("04", "工艺路线", "工艺产能、原料、成本区间与风险"),
        ("04A", "原料与成本", "原料单耗、副产、成本区间与当前价格位置"),
        ("05", "新增项目", f"{len(bundle['projects'])} 个项目：毛产能、置换、净影响与时间"),
        ("05A", "净增还原", "毛产能 ≠ 净新增，逐项还原"),
        ("06", "价格历史", "年度价格、成本带与当前价位"),
        ("06A", "价格周期", "阶段复盘、区域价差与未来判断"),
        ("07", "采购关注", f"{len(bundle['risks'])} 项风险：证据与动作"),
        ("07A", "采购策略", "双源/指数/库存策略、供应商组合与采购量模拟器"),
        ("08", "尽调清单", f"{len(bundle['dd'])} 个问题，可直接用于现场审核"),
        ("09", "供应商评分", f"{len(bundle['scorecard'])} 维度权重模板，可直接打分"),
        ("10", "数据缺口", f"{len(bundle['gaps'])} 项缺口、影响、补齐方式与优先级"),
        ("11", "来源与口径", f"{len(bundle['sources'])} 条来源、口径差异与使用限制"),
    ]
    row = 7
    section(ws, row, "本工作簿包含", 8, "从全景到可执行清单：先看结论，再看企业/区域/项目，最后用尽调与评分模板落地。")
    row += 2
    for i, (numtxt, name, desc) in enumerate(contents):
        r = row + i
        c0 = ws.cell(r, 2, numtxt)
        c0.font = fnt(10, True, TEAL)
        c0.alignment = al("center", "center", False)
        c1 = ws.cell(r, 3, name)
        c1.font = fnt(10.5, True, INK)
        c1.alignment = al("left", "center", False)
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=8)
        c2 = ws.cell(r, 4, desc)
        c2.font = fnt(9, False, MUTED)
        c2.alignment = al("left", "center", True)
        for c in range(2, 9):
            ws.cell(r, c).border = Border(bottom=Side(style="hair", color=LINE))
        ws.row_dimensions[r].height = 22

    # 关键数字
    row += len(contents) + 2
    section(ws, row, "关键数字", 8)
    row += 2
    kpis = [
        (metrics.get("capacity_total_label") or "全国产能", C.fmt_num(metrics.get("capacity_total")),
         f"{unit}；来源 {metrics.get('capacity_total_source') or MISSING}"),
        (metrics.get("output_label") or "产量", C.fmt_num(metrics.get("output")),
         f"开工率 {C.fmt_num(metrics.get('utilization'), 1)}%"),
        ("CR5", f"{C.fmt_num(metrics.get('cr5'), 1)}%" if metrics.get("cr5") else MISSING,
         f"{fwd_label} {metrics.get('cr5_forward', {}).get('low', MISSING)}–{metrics.get('cr5_forward', {}).get('high', MISSING)}%"),
        (metrics.get("capacity_forward_label") or f"{fwd_label} 产能", C.fmt_num(metrics.get("capacity_forward")),
         "名义产能"),
        (metrics.get("price_now_label") or "当前价格", C.fmt_num(metrics.get("price_now")),
         f"{price_unit}"),
        ("样本覆盖", f"{metrics.get('sample_coverage')}%", f"{metrics.get('sample_groups')} 集团 / "
                                                          f"{metrics.get('sample_sites')} 基地 / "
                                                          f"{C.fmt_num(metrics.get('sample_capacity'))} {unit}"),
    ]
    for i, (label, value, sub) in enumerate(kpis):
        r = row + (i // 3) * 3
        c0 = 2 + (i % 3) * 2
        for rr in range(r, r + 3):
            for cc in range(c0, c0 + 2):
                ws.cell(rr, cc).fill = fill(BG)
                ws.cell(rr, cc).border = BORDER
        ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0 + 1)
        c1 = ws.cell(r, c0, label)
        c1.font = fnt(8.5, True, MUTED)
        c1.alignment = al("left", "center", False)
        ws.merge_cells(start_row=r + 1, start_column=c0, end_row=r + 1, end_column=c0 + 1)
        c2 = ws.cell(r + 1, c0, value)
        c2.font = fnt(16, True, NAVY)
        c2.alignment = al("left", "center", False)
        ws.merge_cells(start_row=r + 2, start_column=c0, end_row=r + 2, end_column=c0 + 1)
        c3 = ws.cell(r + 2, c0, sub)
        c3.font = fnt(8, False, MUTED)
        c3.alignment = al("left", "top", True)
        ws.row_dimensions[r].height = 17
        ws.row_dimensions[r + 1].height = 28
        ws.row_dimensions[r + 2].height = 24

    row += 8
    section(ws, row, "口径与使用限制", 8)
    row += 2
    basis = list(scope.get("basis_notes") or []) + [
        "本工作簿为公开信息研究，不含任何企业的非公开成本、合同价、库存与客户数据。",
        "产能为名义/设计能力，实际可用量受检修、环保、能耗、原料与订单影响。",
        "缺失项一律标注「数据缺口」，不做估算；报告不构成投资建议。",
    ]
    for i, text in enumerate(basis):
        ws.merge_cells(start_row=row + i, start_column=2, end_row=row + i, end_column=8)
        cell = ws.cell(row + i, 2, f"{i + 1}. {text}")
        cell.font = fnt(9, False, BODY)
        cell.alignment = al("left", "center", True, 1)
        ws.row_dimensions[row + i].height = 24
    last_cover = row + len(basis) + 1
    ws.print_area = f"A1:H{last_cover}"
    page(ws, "portrait", "1:5")
    ws.page_setup.fitToHeight = 1

    # ---------------------------------------------------------- 01 核心结论
    ws = new("01_核心结论", TEAL)
    banner(ws, "核心结论", f"{len(bundle['verdicts'])} 条最关键的判断：证据、含义与置信度",
           f"数据截止 {cutoff}  ·  置信度：高 / 中高 / 中 / 中低 / 低", 7)
    header(ws, 6, ["#", "结论", "关键证据", "采购含义", "置信度", "来源", "备注"])
    widths(ws, {"A": 5, "B": 30, "C": 54, "D": 52, "E": 9, "F": 20, "G": 22})
    r = 7
    for i, v in enumerate(bundle["verdicts"]):
        row_(ws, r, [v.get("id", f"{i + 1:02d}"), v.get("title", ""), v.get("evidence", ""),
                     v.get("implication", ""), C.norm_conf(v.get("confidence")),
                     "/".join(v.get("sources") or []) or MISSING, v.get("memo", "")],
             aligns=["center", "left", "left", "left", "center", "left", "left"],
             hex_bg=zebra(i), height=62)
        ws.cell(r, 2).font = fnt(10, True, INK)
        conf = C.norm_conf(v.get("confidence"))
        ws.cell(r, 5).font = fnt(9, True, TEAL if conf in ("高", "中高") else AMBER)
        ws.cell(r, 5).fill = fill(TEAL_L if conf in ("高", "中高") else AMBER_L)
        r += 1
    if not bundle["verdicts"]:
        row_(ws, r, ["—", "尚未填写 verdicts", MISSING, MISSING, MISSING, MISSING, MISSING],
             aligns=["center"] * 7, height=30)
        r += 1
    note(ws, r + 1, "读法：置信度「高」= 有年报/公告直接支撑；「中高」= 多来源交叉；"
                    "「中」= 方向明确但时点/净增存在不确定性。证据列中的 S 编号对应「11_来源与口径」。",
         7, 38)
    ws.freeze_panes = "C7"
    if bundle["verdicts"]:
        ws.auto_filter.ref = f"A6:G{r - 1}"
    page(ws, "landscape", "1:6")

    # ---------------------------------------------------------- 02 企业产能
    ws = new("02_企业产能", BLUE)
    banner(ws, "企业产能表", f"{len(bundle['companies'])} 集团 / {len(bundle['plants'])} 基地：集团、基地、工艺与产能",
           f"单位：{unit}  ·  基准年样本合计 {C.fmt_num(metrics.get('sample_capacity'))}（覆盖 {metrics.get('sample_coverage')}%）",
           16)
    header(ws, 6, ["#", "集团", "生产基地", "省份", "城市", "工艺", "基准年产能", f"{fwd_label}产能",
                   "变化", "产品", "主要原料", "状态", "上市代码", "数据来源", "备注", "置信度"])
    widths(ws, {"A": 4.5, "B": 15, "C": 17, "D": 8, "E": 10, "F": 9, "G": 10, "H": 10, "I": 8,
                "J": 22, "K": 22, "L": 8, "M": 9, "N": 22, "O": 46, "P": 7})
    route_color = {r.get("name"): r.get("color", TEAL) for r in bundle["routes"]}
    r = 7
    for i, p in enumerate(bundle["plants"]):
        chg = p["capacity_forward"] - p["capacity_base"]
        row_(ws, r, [i + 1, p.get("company_id"), p.get("site"), p.get("province"), p.get("city", ""),
                     p.get("route", ""), p["capacity_base"], p["capacity_forward"], chg,
                     p.get("product", ""), p.get("raw", ""), p.get("status", ""), p.get("listed", ""),
                     p.get("source", MISSING), p.get("note", ""),
                     "未核实" if p.get("source") in (None, "", "未核实") else C.norm_conf(p.get("confidence", "中"))],
             formats=[None, None, None, None, None, None, "#,##0", "#,##0", "+#,##0;-#,##0;0",
                      None, None, None, None, None, None, None],
             aligns=["center", "left", "left", "center", "center", "center", "right", "right", "right",
                     "left", "left", "center", "center", "left", "left", "center"],
             hex_bg=zebra(i), height=34)
        ws.cell(r, 2).font = fnt(9.4, True, INK)
        rc = str(route_color.get(p.get("route"), SLATE)).lstrip("#")
        ws.cell(r, 6).fill = fill(rc + "" if len(rc) == 6 else SLATE_L)
        ws.cell(r, 6).font = fnt(8.8, True, WHITE)
        r += 1
    last = r - 1
    if last >= 7:
        total_row(ws, r, "样本合计 / 覆盖率", 16, [
            (7, f"=SUM(G7:G{last})", "#,##0", "right"),
            (8, f"=SUM(H7:H{last})", "#,##0", "right"),
            (9, f"=SUM(I7:I{last})", "+#,##0;-#,##0;0", "right"),
            (10, f"覆盖 {metrics.get('sample_coverage')}%", None, "left"),
        ])
        ws.conditional_formatting.add(f"G7:H{last}", DataBarRule(
            start_type="num", start_value=0, end_type="max", color=TEAL_M, showValue=True))
        ws.conditional_formatting.add(f"I7:I{last}", CellIsRule(
            operator="lessThan", formula=["0"], font=Font(name=FONT, size=9, bold=True, color=RED)))
        ws.auto_filter.ref = f"A6:P{last}"
        print_r = r + 2
    else:
        print_r = r + 1
    note(ws, print_r, "口径：① 本表为公开披露的企业/基地样本，上方合计行为实际加总；"
                      f"覆盖率 = 样本合计 ÷ 全国 {C.fmt_num(metrics.get('capacity_total'))} {unit} "
                      f"= {metrics.get('sample_coverage') or C.MISSING}%；未披露的小型企业未分摊。"
                      "② 产能为名义/设计能力，不等于可售量。"
                      "③「变化」= 预测年产能 − 基准年产能，正数为扩产，负数为退出/置换减少。"
                      "④ 来源为空时标「未核实」，不得当作已核实的数字使用。", 16, 56)
    ws.freeze_panes = "C7"
    page(ws, "landscape", "1:6")
    ws.print_area = f"A1:P{print_r + 2}"

    # ---------------------------------------------------------- 03 区域汇总
    ws = new("03_区域汇总", TEAL)
    banner(ws, "区域汇总", f"{len(bundle['regions'])} 个生产省份：产能、集中度、贸易活跃度与下游",
           f"样本口径  ·  基准年合计 {C.fmt_num(metrics.get('sample_capacity'))} {unit}  ·  "
           f"区域 HHI {metrics.get('hhi_region')}", 11)
    header(ws, 6, ["省份", "集团数", "基地数", "基准年产能", "占比", f"{fwd_label}产能", f"{fwd_label}占比",
                   "贸易活跃度", "主要下游 / 流向", "代表企业", "枢纽 / 园区"])
    widths(ws, {"A": 9, "B": 8, "C": 8, "D": 11, "E": 9, "F": 11, "G": 10, "H": 10, "I": 38, "J": 30, "K": 24})
    r = 7
    for i, rg in enumerate(bundle["regions"]):
        meta_r = bundle["region_meta"].get(rg["province"], {})
        row_(ws, r, [rg["province"], rg["group_count"], rg["sites"], rg["capacity_base"], rg["share_base"],
                     rg["capacity_forward"], rg["share_forward"], meta_r.get("activity", MISSING),
                     meta_r.get("downstream", MISSING),
                     "、".join(rg["companies"][:4]) + ("…" if len(rg["companies"]) > 4 else ""),
                     meta_r.get("hub", MISSING)],
             formats=[None, "0", "0", "#,##0", "0.0%", "#,##0", "0.0%", None, None, None, None],
             aligns=["left", "center", "center", "right", "right", "right", "right", "center", "left", "left", "left"],
             hex_bg=zebra(i), height=34)
        ws.cell(r, 1).font = fnt(9.5, True, INK)
        act = meta_r.get("activity")
        ws.cell(r, 8).font = fnt(9, True, {"A": TEAL, "B": BLUE, "C": SLATE}.get(act, SLATE))
        ws.cell(r, 8).fill = fill({"A": TEAL_L, "B": BLUE_L, "C": SLATE_L}.get(act, SLATE_L))
        r += 1
    last = r - 1
    if last >= 7:
        total_row(ws, r, "样本合计", 11, [
            (4, f"=SUM(D7:D{last})", "#,##0", "right"),
            (5, f"=SUM(E7:E{last})", "0.0%", "right"),
            (6, f"=SUM(F7:F{last})", "#,##0", "right"),
            (7, f"=SUM(G7:G{last})", "0.0%", "right"),
            (9, f"全国覆盖率 {metrics.get('sample_coverage')}%", None, "left"),
        ])
        chart = BarChart()
        chart.type = "bar"
        chart.style = 10
        chart.title = f"区域产能：基准年 vs {fwd_label}（{unit}）"
        chart.x_axis.title = None
        chart.height, chart.width = 11, 18
        chart.add_data(Reference(ws, min_col=4, min_row=6, max_row=last), titles_from_data=True)
        chart.add_data(Reference(ws, min_col=6, min_row=6, max_row=last), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=7, max_row=last))
        chart.series[0].graphicalProperties = GraphicalProperties(solidFill=TEAL)
        chart.series[1].graphicalProperties = GraphicalProperties(solidFill="93C5FD")
        chart.legend.position = "b"
        chart.y_axis.majorGridlines = ChartLines()
        chart.y_axis.scaling.orientation = "maxMin"
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showVal = True
        ws.add_chart(chart, "M6")
        ws.conditional_formatting.add(f"D7:D{last}", DataBarRule(
            start_type="num", start_value=0, end_type="max", color=TEAL_M, showValue=True))
        ws.auto_filter.ref = f"A6:K{last}"
        end = r + 3
    else:
        end = r + 1
    note(ws, end, "口径：① 区域产能为样本企业加总，低于全国总产能，差额为未披露的中小企业。"
                  "② 贸易活跃度 A/B/C 为定性分级（A=交易活跃/贸易商密集，B=中等，C=直供为主）；"
                  "按省贸易商数量无公开统计，本表不做估算，缺口见「10_数据缺口」。"
                  "③ 区域 HHI 越高说明越依赖个别省份，采购半径内可替代供应商越少。", 11, 52)
    ws.freeze_panes = "B7"
    page(ws, "landscape", "1:6")
    ws.print_area = f"A1:K{end + 2}"

    # ---------------------------------------------------------- 04 工艺路线
    ws = new("04_工艺路线", AMBER)
    banner(ws, "工艺路线", "各工艺的产能、原料、单耗、成本与风险",
           f"基准年产能结构  ·  单位：{unit}", 10)
    header(ws, 6, ["工艺", "基准年产能", "占比", "主要企业", "主要原料", "单耗 / 副产", "能耗",
                   "成本区间", "优势", "风险"])
    widths(ws, {"A": 12, "B": 10, "C": 8, "D": 44, "E": 26, "F": 44, "G": 18, "H": 14, "I": 40, "J": 40})
    r = 7
    for i, rt in enumerate(bundle["routes"]):
        cost = MISSING
        if rt.get("cost_low") is not None and rt.get("cost_high") is not None:
            cost = f"{C.fmt_num(rt['cost_low'])}–{C.fmt_num(rt['cost_high'])}"
        row_(ws, r, [rt.get("name"), rt.get("cap"), C.num(rt.get("share")) / 100, rt.get("producers", ""),
                     rt.get("raw", ""), rt.get("unit", ""), rt.get("energy", ""), cost,
                     rt.get("advantage", ""), rt.get("risk", "")],
             formats=[None, "#,##0", "0.0%", None, None, None, None, None, None, None],
             aligns=["left", "right", "right", "left", "left", "left", "left", "center", "left", "left"],
             hex_bg=zebra(i), height=58)
        rc = str(rt.get("color", "")).lstrip("#") or SLATE
        ws.cell(r, 1).font = fnt(10, True, rc if len(rc) == 6 else INK)
        r += 1
    last = r - 1
    if last >= 7:
        total_row(ws, r, "合计", 10, [
            (2, f"=SUM(B7:B{last})", "#,##0", "right"),
            (3, f"=SUM(C7:C{last})", "0.0%", "right"),
            (4, f"全国 {C.fmt_num(metrics.get('capacity_total'))}", None, "left"),
        ])
        chart = DoughnutChart()
        chart.title = "基准年工艺结构"
        chart.style = 10
        chart.holeSize = 55
        chart.add_data(Reference(ws, min_col=2, min_row=6, max_row=last), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=7, max_row=last))
        for idx, rt in enumerate(bundle["routes"]):
            pt = DataPoint(idx=idx)
            rc = str(rt.get("color", "")).lstrip("#") or SLATE
            pt.graphicalProperties = GraphicalProperties(solidFill=rc if len(rc) == 6 else SLATE)
            chart.series[0].data_points.append(pt)
        chart.legend.position = "r"
        chart.height, chart.width = 8.5, 12
        ws.add_chart(chart, "L6")
        end = r + 2
    else:
        end = r + 1
    note(ws, end, f"{bundle['meta'].get('routes_note') or ''} 工艺不是「越新越好」：每种路线都受资源、"
                  "副产、环保或能耗约束；采购前应确认供应商的实际路线与原料自给率。", 10, 44)
    ws.freeze_panes = "A7"
    page(ws, "landscape", "1:6")
    ws.print_area = f"A1:J{end + 2}"

    # ---------------------------------------------------------- 04A 原料与成本
    ws = new("04A_原料与成本", TEAL)
    banner(ws, "原料与成本", "原料单耗、成本区间与当前价格位置",
           f"成本单位：{price_unit}  ·  单耗为行业经验值，非企业实测", 5)
    section(ws, 6, "原料单耗与副产（行业经验值）", 5)
    header(ws, 8, ["工艺 / 原料", "主要原料", "单耗", "副产", "能耗与备注"])
    widths(ws, {"A": 14, "B": 26, "C": 40, "D": 22, "E": 44})
    r = 9
    for i, rt in enumerate(bundle["routes"]):
        row_(ws, r, [rt.get("name"), rt.get("raw", MISSING), rt.get("unit", MISSING),
                     rt.get("byproduct", MISSING), rt.get("energy", MISSING)],
             aligns=["left", "left", "left", "left", "left"], hex_bg=zebra(i), height=30)
        ws.cell(r, 1).font = fnt(9.5, True, INK)
        r += 1
    row = r + 1
    section(ws, row, "成本区间与当前价格位置", 5)
    header(ws, row + 2, ["工艺", "成本低", "成本高", "中值", f"当前价 {C.fmt_num(metrics.get('price_now'))}"])
    r = row + 3
    price_now = C.num(metrics.get("price_now"))
    first_cost = r
    for i, rt in enumerate(bundle["routes"]):
        lo, hi = rt.get("cost_low"), rt.get("cost_high")
        if lo is None or hi is None:
            pos = MISSING
        elif price_now >= hi:
            pos = "高于成本上限，仍有利润"
        elif price_now >= lo:
            pos = "位于成本区间内"
        else:
            pos = "低于成本下限，普遍亏损"
        row_(ws, r, [rt.get("name"), lo if lo is not None else MISSING,
                     hi if hi is not None else MISSING, f"=AVERAGE(B{r}:C{r})", pos],
             formats=[None, "#,##0", "#,##0", "#,##0", None],
             aligns=["left", "right", "right", "right", "left"], hex_bg=zebra(i), height=26)
        ws.cell(r, 1).font = fnt(9, True, INK)
        r += 1
    if r - 1 >= first_cost:
        chart = BarChart()
        chart.type = "bar"
        chart.style = 10
        chart.title = "成本区间中值"
        chart.height, chart.width = 7.5, 12
        chart.add_data(Reference(ws, min_col=4, min_row=first_cost, max_row=r - 1), titles_from_data=False)
        chart.set_categories(Reference(ws, min_col=1, min_row=first_cost, max_row=r - 1))
        chart.series[0].graphicalProperties = GraphicalProperties(solidFill=TEAL)
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showVal = True
        chart.legend = None
        ws.add_chart(chart, "G8")
    note(ws, r + 1, "说明：① 单耗为行业经验值，用于理解成本结构，不能用于结算或成本审计。"
                    "② 成本区间来自公开报道/行业研究，不同基地的原料、能源、折旧、副产收益差异很大。"
                    "③「当前价」位置只说明相对盈亏，不代表可以据此压价——还要看交付半径与供应保障。", 5, 46)
    ws.freeze_panes = "A9"
    page(ws, "landscape", "1:6")
    ws.print_area = f"A1:E{r + 3}"

    # ---------------------------------------------------------- 05 新增项目
    ws = new("05_新增项目", "7C3AED")
    banner(ws, "新增项目", "新建 / 扩建 / 技改 / 置换：毛产能、净影响与投产节奏",
           f"单位：{unit}  ·  毛产能 ≠ 净增  ·  置换/搬迁项目需逐项还原", 13)
    header(ws, 6, ["#", "项目", "集团", "省份", "工艺", "毛产能", "项目类型", "状态", "时间",
                   "净影响", "进度", "对供给的影响", "来源"])
    widths(ws, {"A": 4.5, "B": 24, "C": 14, "D": 8, "E": 10, "F": 9, "G": 14, "H": 8, "I": 12,
                "J": 13, "K": 30, "L": 44, "M": 24})
    r = 7
    for i, p in enumerate(bundle["projects"]):
        row_(ws, r, [i + 1, p["name"], p.get("company_id", ""), p.get("province", ""), p.get("route", ""),
                     p["capacity"], p.get("type", ""), p.get("status", ""), p.get("date", ""),
                     p["net_label"], p.get("step", ""), p.get("impact", ""), p.get("source", "")],
             formats=[None, None, None, None, None, "#,##0", None, None, None, None, None, None, None],
             aligns=["center", "left", "left", "center", "center", "right", "center", "center",
                     "center", "center", "left", "left", "left"],
             hex_bg=zebra(i), height=44)
        ws.cell(r, 2).font = fnt(9.5, True, INK)
        ws.cell(r, 8).font = fnt(8.8, True, {"已投产": TEAL, "投产": TEAL, "在建": AMBER,
                                             "拟建": BLUE, "规划": BLUE}.get(p.get("status"), SLATE))
        ws.cell(r, 10).font = fnt(9, True, RED if str(p["net_label"]).startswith("-") else TEAL)
        r += 1
    last = r - 1
    if last >= 7:
        chart = BarChart()
        chart.type = "bar"
        chart.style = 10
        chart.title = f"项目毛产能（{unit}）"
        chart.height, chart.width = 8.5, 17
        chart.add_data(Reference(ws, min_col=6, min_row=6, max_row=last), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=2, min_row=7, max_row=last))
        palette = {"已投产": TEAL, "投产": TEAL, "在建": AMBER, "拟建": BLUE, "规划": BLUE}
        for idx, p in enumerate(bundle["projects"]):
            pt = DataPoint(idx=idx)
            pt.graphicalProperties = GraphicalProperties(solidFill=palette.get(p.get("status"), SLATE))
            chart.series[0].data_points.append(pt)
        chart.legend = None
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showVal = True
        chart.y_axis.scaling.orientation = "maxMin"
        ws.add_chart(chart, "O6")
        ws.auto_filter.ref = f"A6:M{last}"
        end = r + 2
    else:
        end = r + 1
    note(ws, end, "读法：① 毛产能是项目规模，不等于可售量；② 置换/搬迁必须扣减退出产能才是净影响；"
                  "③ 净影响标「待确认」的项目不纳入基准预测，只在上行情景中体现；"
                  "④ 项目延期很常见，签约前要求供应商披露老装置关停时点与过渡期保供方案。", 13, 44)
    ws.freeze_panes = "C7"
    page(ws, "landscape", "1:6")
    ws.print_area = f"A1:M{end + 2}"

    # ---------------------------------------------------------- 05A 净增还原
    ws = new("05A_净增还原", "8B5CF6")
    banner(ws, "净增还原", "毛产能 ≠ 净新增：逐项还原置换、退出与不确定性",
           f"单位：{unit}  ·  「待确认」不计入基准预测  ·  净影响 = 毛产能 − 置换/退出", 5)
    header(ws, 6, ["项目", "毛产能", "置换/退出", "净影响", "判断"])
    widths(ws, {"A": 28, "B": 14, "C": 14, "D": 14, "E": 58})
    r = 7
    for i, p in enumerate(bundle["projects"]):
        net = p["net"]
        judgment = p.get("impact") or MISSING
        if net is None:
            judgment = "净增待确认，不纳入基准预测。" + judgment
        row_(ws, r, [p["name"], p["capacity"], p["capacity_retired"] or "—", p["net_label"], judgment],
             formats=[None, "#,##0", "#,##0", None, None],
             aligns=["left", "right", "right", "center", "left"], hex_bg=zebra(i), height=32)
        ws.cell(r, 1).font = fnt(9.2, True, INK)
        ws.cell(r, 4).font = fnt(9.2, True, RED if str(p["net_label"]).startswith("-")
                                 else TEAL if str(p["net_label"]).startswith("+") else AMBER)
        r += 1
    row = r + 1
    section(ws, row, "判断规则：采购规划只看净增", 5)
    rules = [
        "① 净新增（新建/扩建）：毛产能可加到可用供给，但要按爬坡节奏打折（通常首年 50%–80%）。",
        "② 置换：净影响 = 新产能 − 退出产能；不能只看新产能。",
        "③ 搬迁：新厂投产 ≠ 老厂立即关停；必须确认关停时点、过渡期保供与库存安排。",
        "④ 待确认：不纳入基准预测，只在敏感性/上行情景中体现；签约前要求企业书面确认。",
        "⑤ 没有项目的年份往往比有项目的年份更紧张——供给增量断档会放大检修与补库的影响。",
    ]
    band(ws, row + 2, 1, row + 6, 5, AMBER_L)
    ws.merge_cells(start_row=row + 2, start_column=1, end_row=row + 6, end_column=5)
    cell = ws.cell(row + 2, 1, "\n\n".join(rules))
    cell.font = fnt(9, False, BODY)
    cell.alignment = al("left", "top", True, 1)
    for rr in range(row + 2, row + 7):
        for cc in range(1, 6):
            ws.cell(rr, cc).border = BORDER
    note(ws, row + 8, "用法：把「净影响」列作为供应规划输入；把「待确认」项列成一张跟踪清单，"
                      "每月核对一次项目公告与产能置换公示。", 5, 32)
    ws.freeze_panes = "A7"
    page(ws, "landscape", "1:6")
    ws.print_area = f"A1:E{row + 10}"

    # ---------------------------------------------------------- 06 价格历史
    ws = new("06_价格历史", RED)
    prices = bundle["prices"]
    banner(ws, "价格历史", "年度价格、成本带与当前价位",
           f"单位：{price_unit}  ·  峰值与均价区分标注  ·  当前价 {C.fmt_num(metrics.get('price_now'))}",
           9)
    section(ws, 6, "年度价格", 9)
    header(ws, 8, ["年份", "均价 / 中枢", "低点", "高点", "说明"])
    widths(ws, {"A": 9, "B": 14, "C": 11, "D": 11, "E": 52, "F": 2, "G": 14, "H": 14, "I": 14})
    r = 9
    for i, p in enumerate(prices.get("annual") or []):
        row_(ws, r, [p.get("year"), p.get("avg") if p.get("avg") is not None else (p.get("high") or MISSING),
                     p.get("low") if p.get("low") is not None else MISSING,
                     p.get("high") if p.get("high") is not None else MISSING,
                     p.get("note", "")],
             formats=[None, "#,##0", "#,##0", "#,##0", None],
             aligns=["center", "right", "right", "right", "left"], hex_bg=zebra(i), height=26)
        ws.cell(r, 1).font = fnt(9.5, True, INK)
        r += 1
    if r - 1 >= 9:
        chart = LineChart()
        chart.title = "年度价格中枢与区间"
        chart.style = 12
        chart.height, chart.width = 9, 17
        chart.add_data(Reference(ws, min_col=2, max_col=4, min_row=8, max_row=r - 1), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=9, max_row=r - 1))
        for idx, color in enumerate([TEAL, "94A3B8", "CBD5E1"]):
            if idx < len(chart.series):
                ser = chart.series[idx]
                ser.graphicalProperties = GraphicalProperties()
                ser.graphicalProperties.line.solidFill = color
                ser.graphicalProperties.line.width = 30000 if idx == 0 else 16000
                ser.marker = Marker(symbol="circle", size=7 if idx == 0 else 4)
        chart.legend.position = "b"
        chart.y_axis.majorGridlines = ChartLines()
        ws.add_chart(chart, "G8")
    note(ws, r + 1, prices.get("caveat") or "价格为公开报价，年度均值用于理解周期位置，不作为结算依据。", 9, 36)
    ws.freeze_panes = "A9"
    page(ws, "landscape", "1:8")
    ws.print_area = f"A1:I{r + 3}"

    # ---------------------------------------------------------- 06A 价格周期
    ws = new("06A_价格周期", "EF4444")
    banner(ws, "价格周期", "阶段复盘、区域价差与未来判断",
           f"单位：{price_unit}", 8)
    section(ws, 6, "阶段复盘", 8)
    header(ws, 8, ["阶段", "时间", "主规格", "副规格", "均价 / 区间", "涨跌", "主要驱动", "来源"])
    widths(ws, {"A": 12, "B": 18, "C": 18, "D": 22, "E": 14, "F": 18, "G": 42, "H": 20})
    r = 9
    for i, p in enumerate(prices.get("periods") or []):
        keys = [k for k in p.keys() if k not in ("period", "time", "avg", "change", "driver", "source")]
        spec1 = p.get("light") or (p.get(keys[0]) if keys else MISSING)
        spec2 = p.get("heavy") or (p.get(keys[1]) if len(keys) > 1 else MISSING)
        row_(ws, r, [p.get("period"), p.get("time"), spec1, spec2, p.get("avg", MISSING),
                     p.get("change", MISSING), p.get("driver", ""), p.get("source", "")],
             aligns=["left", "center", "center", "center", "center", "center", "left", "left"],
             hex_bg=zebra(i), height=40)
        ws.cell(r, 1).font = fnt(9.5, True, INK)
        r += 1
    markets = bundle.get("markets") or []
    if markets:
        row = r + 1
        section(ws, row, f"区域价格（{price_unit}）", 8)
        header(ws, row + 2, ["区域", "价格", "备注", "", "", "", "", ""])
        ws.merge_cells(start_row=row + 2, start_column=3, end_row=row + 2, end_column=6)
        r = row + 3
        for i, mk in enumerate(markets):
            ws.cell(r, 1, mk.get("region")).font = fnt(9, True, INK)
            ws.cell(r, 2, mk.get("price")).number_format = "#,##0"
            ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
            ws.cell(r, 3, mk.get("note", ""))
            for cc in range(1, 7):
                cell = ws.cell(r, cc)
                cell.font = fnt(9, False, BODY) if cc != 1 else fnt(9, True, INK)
                cell.border = BORDER
                cell.fill = fill(zebra(i))
                cell.alignment = al("left" if cc in (1, 3) else "right", "center", True)
            ws.row_dimensions[r].height = 26
            r += 1
    outlook = prices.get("outlook") or {}
    row = r + 1
    section(ws, row, "价格判断（未来 6–12 个月）", 8)
    text = "\n\n".join(filter(None, [
        f"基准情景：{outlook.get('base')}" if outlook.get("base") else "",
        f"上行条件：{outlook.get('up')}" if outlook.get("up") else "",
        f"下行风险：{outlook.get('down')}" if outlook.get("down") else "",
        f"采购含义：{outlook.get('procurement')}" if outlook.get("procurement") else "",
    ])) or "尚未填写 prices.outlook：请补充基准/上行/下行情景与采购含义。"
    band(ws, row + 2, 1, row + 6, 6, AMBER_L)
    ws.merge_cells(start_row=row + 2, start_column=1, end_row=row + 6, end_column=6)
    cell = ws.cell(row + 2, 1, text)
    cell.font = fnt(9, False, BODY)
    cell.alignment = al("left", "top", True, 1)
    for rr in range(row + 2, row + 7):
        for cc in range(1, 7):
            ws.cell(rr, cc).border = BORDER
    note(ws, row + 8, "区域价格为公开主流报价，实际成交受批量、账期、运费、质量影响；"
                      "价格判断为情景分析，不是预测承诺。", 8, 30)
    ws.freeze_panes = "A9"
    page(ws, "landscape", "1:8")
    ws.print_area = f"A1:H{row + 10}"

    # ---------------------------------------------------------- 07 采购关注
    ws = new("07_采购关注", AMBER)
    banner(ws, "采购关注", f"{len(bundle['risks'])} 项风险：可能性、影响、证据与动作",
           "风险等级：高 / 中高 / 中 / 低  ·  按等级从高到低排列", 8)
    section(ws, 6, "风险矩阵与动作", 8)
    header(ws, 8, ["风险", "可能性", "影响", "等级", "证据", "采购动作", "建议责任人", "时间"])
    widths(ws, {"A": 22, "B": 8, "C": 8, "D": 8, "E": 50, "F": 50, "G": 12, "H": 10})
    r = 9
    for i, rk in enumerate(C.order_by_priority(bundle["risks"], "level")):
        row_(ws, r, [rk.get("risk"), rk.get("likelihood", MISSING), rk.get("impact", MISSING),
                     rk.get("level", MISSING), rk.get("evidence", ""), rk.get("action", ""),
                     "采购 / 计划", MISSING],
             aligns=["left", "center", "center", "center", "left", "left", "center", "center"],
             hex_bg=zebra(i), height=46)
        ws.cell(r, 1).font = fnt(9.2, True, INK)
        lv = rk.get("level")
        ws.cell(r, 4).font = fnt(9, True, RED if lv == "高" else AMBER if lv == "中高" else SLATE)
        ws.cell(r, 4).fill = fill(RED_L if lv == "高" else AMBER_L if lv == "中高" else SLATE_L)
        r += 1
    note(ws, r + 1, "使用方式：① 先按等级排序，锁定 3 个高/中高风险；② 每个动作落到责任人、时间与验收标准；"
                    "③ 每季度复盘，价格/项目/政策变化时即时更新。策略与模拟器见「07A_采购策略」。", 8, 34)
    ws.freeze_panes = "A9"
    page(ws, "landscape", "1:8")
    ws.print_area = f"A1:H{r + 3}"

    # ---------------------------------------------------------- 07A 采购策略
    ws = new("07A_采购策略", "F59E0B")
    banner(ws, "采购策略", "双源、指数定价、安全库存、质量封样与采购量模拟",
           "示例比例为默认建议，可按企业实际调整", 6)
    widths(ws, {"A": 18, "B": 20, "C": 18, "D": 25, "E": 25, "F": 22})
    section(ws, 6, "采购策略", 6)
    for col, text in [(1, "策略")]:
        cell = ws.cell(8, col, text)
        cell.font = fnt(9.5, True, WHITE)
        cell.fill = fill(NAVY2)
        cell.alignment = al("center", "center", True)
        cell.border = BORDER
    for rng, text in [("B8:C8", "具体做法"), ("D8:E8", "关键条款")]:
        ws.merge_cells(rng)
        first = rng.split(":")[0]
        cell = ws[first]
        cell.value = text
        cell.font = fnt(9.5, True, WHITE)
        cell.fill = fill(NAVY2)
        cell.alignment = al("center", "center", True)
        cell.border = BORDER
    cell = ws.cell(8, 6, "KPI")
    cell.font = fnt(9.5, True, WHITE)
    cell.fill = fill(NAVY2)
    cell.alignment = al("center", "center", True)
    cell.border = BORDER
    ws.row_dimensions[8].height = 28
    r = 9
    strategies = bundle["strategies"] or [{"title": MISSING, "how": "尚未填写 strategies",
                                           "terms": MISSING, "kpi": MISSING}]
    for i, st in enumerate(strategies):
        ws.cell(r, 1, st.get("title")).font = fnt(9.3, True, INK)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        ws.cell(r, 2, st.get("how")).font = fnt(9, False, BODY)
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)
        ws.cell(r, 4, st.get("terms")).font = fnt(9, False, BODY)
        ws.cell(r, 6, st.get("kpi")).font = fnt(9, False, BODY)
        for cc in range(1, 7):
            cell = ws.cell(r, cc)
            cell.border = BORDER
            cell.fill = fill(zebra(i))
            cell.alignment = al("left", "center", True)
        ws.row_dimensions[r].height = 50
        r += 1

    mix_start = r + 2
    section(ws, mix_start, "建议供应商组合（示例）", 6)
    for col, text in [(1, "类型"), (6, "占比")]:
        cell = ws.cell(mix_start + 2, col, text)
        cell.font = fnt(9.5, True, WHITE)
        cell.fill = fill(NAVY2)
        cell.alignment = al("center", "center", True)
        cell.border = BORDER
    for rng, text in [(f"B{mix_start + 2}:C{mix_start + 2}", "候选"),
                      (f"D{mix_start + 2}:E{mix_start + 2}", "作用")]:
        ws.merge_cells(rng)
        cell = ws[rng.split(":")[0]]
        cell.value = text
        cell.font = fnt(9.5, True, WHITE)
        cell.fill = fill(NAVY2)
        cell.alignment = al("center", "center", True)
        cell.border = BORDER
    ws.row_dimensions[mix_start + 2].height = 28
    r = mix_start + 3
    mix_first = r
    mix = bundle["supplier_mix"] or []
    for i, mx in enumerate(mix):
        ws.cell(r, 1, mx.get("type")).font = fnt(9, True, INK)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        ws.cell(r, 2, mx.get("candidates", "")).font = fnt(9, False, BODY)
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)
        ws.cell(r, 4, mx.get("role", "")).font = fnt(9, False, BODY)
        share_cell = ws.cell(r, 6, C.num(mx.get("share")))
        share_cell.number_format = "0%"
        share_cell.font = fnt(10, True, TEAL)
        for cc in range(1, 7):
            cell = ws.cell(r, cc)
            cell.border = BORDER
            cell.fill = fill(zebra(i))
            cell.alignment = al("center" if cc == 6 else "left", "center", True)
        ws.row_dimensions[r].height = 26
        r += 1
    if mix:
        pie = PieChart()
        pie.title = "建议采购组合"
        pie.height, pie.width = 7, 10
        pie.add_data(Reference(ws, min_col=6, min_row=mix_first, max_row=r - 1), titles_from_data=False)
        pie.set_categories(Reference(ws, min_col=1, min_row=mix_first, max_row=r - 1))
        for idx, mx in enumerate(mix):
            pt = DataPoint(idx=idx)
            rc = str(mx.get("color", "")).lstrip("#") or SLATE
            pt.graphicalProperties = GraphicalProperties(solidFill=rc if len(rc) == 6 else SLATE)
            pie.series[0].data_points.append(pt)
        pie.dataLabels = DataLabelList()
        pie.dataLabels.showPercent = True
        ws.add_chart(pie, "H8")

    sim = r + 1
    section(ws, sim, "采购量模拟器（黄底为输入，其余自动计算）", 6)
    for col, text in [(1, "输入 / 输出"), (2, "值"), (3, "单位")]:
        cell = ws.cell(sim + 2, col, text)
        cell.font = fnt(9.5, True, WHITE)
        cell.fill = fill(NAVY2)
        cell.alignment = al("center", "center", True)
        cell.border = BORDER
    ws.merge_cells(start_row=sim + 2, start_column=4, end_row=sim + 2, end_column=6)
    cell = ws.cell(sim + 2, 4, "说明")
    cell.font = fnt(9.5, True, WHITE)
    cell.fill = fill(NAVY2)
    cell.alignment = al("center", "center", True)
    cell.border = BORDER
    buyer = bundle["meta"].get("buyer") or {}
    inputs = [
        ("年采购量", C.num(buyer.get("volume_10kt"), 5), "万吨", "按实际调整；默认 5 万吨/年"),
        ("到厂价", C.num(buyer.get("landed_price"), C.num(metrics.get("price_now"))), price_unit, "含税到厂价"),
        ("运费", C.num(buyer.get("freight"), 150), price_unit, "按采购地调整"),
        ("安全库存", C.num(buyer.get("stock_days"), 30), "天", "消费地库存天数"),
    ]
    r = sim + 3
    for i, (label, value, unit_txt, desc) in enumerate(inputs):
        ws.cell(r, 1, label).font = fnt(9, True, INK)
        vcell = ws.cell(r, 2, value)
        vcell.number_format = "#,##0"
        vcell.font = fnt(9, True, BODY)
        vcell.fill = fill("FEF3C7")
        ws.cell(r, 3, unit_txt).font = fnt(9, False, BODY)
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=6)
        ws.cell(r, 4, desc).font = fnt(9, False, BODY)
        for cc in range(1, 7):
            cell = ws.cell(r, cc)
            cell.border = BORDER
            cell.alignment = al("right" if cc == 2 else "left", "center", True)
        ws.row_dimensions[r].height = 26
        r += 1
    in_vol, in_price, in_freight, in_days = sim + 3, sim + 4, sim + 5, sim + 6
    r += 1
    outputs = [
        ("年度采购金额（含运费）", f"=B{in_vol}*10000*(B{in_price}+B{in_freight})/10000", "万元",
         "年量 × (到厂价 + 运费)"),
        ("月度提货量", f"=B{in_vol}*10000/12", "吨", "年量 / 12"),
        ("安全库存量", f"=B{in_vol}*10000*B{in_days}/365", "吨", "年量 × 天数 / 365"),
        ("价格每涨 100 的影响", f"=B{in_vol}*100", "万元/年", "年量（万吨）× 100"),
    ]
    for i, (label, formula, unit_txt, desc) in enumerate(outputs):
        ws.cell(r, 1, label).font = fnt(9, True, INK)
        cell = ws.cell(r, 2, formula)
        cell.font = fnt(11, True, TEAL)
        cell.number_format = "#,##0"
        ws.cell(r, 3, unit_txt).font = fnt(9, False, BODY)
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=6)
        ws.cell(r, 4, desc).font = fnt(9, False, BODY)
        for cc in range(1, 7):
            c2 = ws.cell(r, cc)
            c2.border = BORDER
            c2.fill = fill(TEAL_L if i == 0 else WHITE)
            c2.alignment = al("right" if cc == 2 else "left", "center", True)
        ws.row_dimensions[r].height = 28
        r += 1
    note(ws, r + 1, "模拟器为量级测算，未含增值税、资金占用、仓储、损耗与账期成本。"
                    "把「年度采购金额」与预算科目对齐，把「价格每涨 100 的影响」作为价格授权与套保额度依据。", 6, 38)
    ws.freeze_panes = "A9"
    page(ws, "landscape", "1:8")
    ws.print_area = f"A1:F{r + 3}"

    # ---------------------------------------------------------- 08 尽调清单
    ws = new("08_尽调清单", BLUE)
    dd = bundle["dd"] or []
    banner(ws, "供应商尽调清单", f"{len(dd)} 个问题：供应、质量、成本、物流、商务、合规、财务、ESG",
           "评分 1–5：1=缺失/不可接受，3=基本满足，5=优秀/可验证  ·  权重 1–10", 8)
    header(ws, 6, ["#", "类别", "尽调问题", "证据 / 文件", "权重", "评分 (1–5)", "加权得分", "备注"])
    widths(ws, {"A": 4.5, "B": 10, "C": 64, "D": 32, "E": 7, "F": 9, "G": 9, "H": 22})
    evidence = {
        "供应能力": "产能表/年报/排产计划/检修记录",
        "质量": "COA/检测报告/封样/小试报告",
        "成本": "成本拆解/原料采购合同/能耗数据",
        "物流": "铁路专用线/港口协议/运费报价",
        "商务": "合同草案/价格公式/授信资料",
        "合规": "排污许可证/能耗指标/处罚记录",
        "财务": "审计报告/银行授信/担保明细",
        "ESG": "碳核算/绿电合同/社区记录",
        "其他": "交割厂库资质/客户证明/审厂报告",
    }
    r = 7
    for i, item in enumerate(dd):
        cat = item[0] if isinstance(item, (list, tuple)) and len(item) > 0 else "其他"
        q = item[1] if isinstance(item, (list, tuple)) and len(item) > 1 else str(item)
        row_(ws, r, [i + 1, cat, q, evidence.get(cat, "—"), 8, None, f"=E{r}*F{r}", ""],
             formats=[None, None, None, None, "0", "0", "0.0", None],
             aligns=["center", "center", "left", "left", "center", "center", "center", "left"],
             hex_bg=zebra(i), height=32)
        ws.cell(r, 2).font = fnt(8.5, True, TEAL)
        ws.cell(r, 6).fill = fill("FEF3C7")
        r += 1
    last = r - 1
    if last >= 7:
        total_row(ws, r, "加权得分合计", 8, [(7, f"=SUM(G7:G{last})", "0.0", "center"),
                                            (8, "满分 = 权重 × 5 之和", None, "left")])
        dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True)
        dv.error = "请输入 1–5 的整数"
        dv.errorTitle = "评分范围"
        ws.add_data_validation(dv)
        dv.add(f"F7:F{last}")
        ws.conditional_formatting.add(f"G7:G{last}", ColorScaleRule(
            start_type="min", start_color="FEE2E2", mid_type="percentile", mid_value=50,
            mid_color="FEF3C7", end_type="max", end_color="CCFBF1"))
        ws.auto_filter.ref = f"A6:H{last}"
        end = r + 2
    else:
        end = r + 1
    note(ws, end, "使用建议：① 先桌面审查（年报/公告/COA/许可证），再现场审核；② 每个问题至少留一份可验证证据；"
                  "③ 评分由采购、质量、工艺、财务四方独立打分后取中位数；④ 加权得分用于横向比较，不能替代试单。",
         8, 38)
    ws.freeze_panes = "C7"
    page(ws, "landscape", "1:6")

    # ---------------------------------------------------------- 09 供应商评分
    ws = new("09_供应商评分", TEAL)
    score = bundle["scorecard"] or []
    banner(ws, "供应商评分模板", f"{len(score)} 维度权重 + 4 家供应商横向比较",
           "评分 1–5  ·  黄底为输入  ·  加权总分自动计算  ·  满分 5 分", 7)
    header(ws, 6, ["评审维度", "权重", "评分要点", "供应商 A", "供应商 B", "供应商 C", "供应商 D"])
    widths(ws, {"A": 20, "B": 8, "C": 58, "D": 13, "E": 13, "F": 13, "G": 13})
    r = 7
    for i, item in enumerate(score):
        dim = item[0] if isinstance(item, (list, tuple)) else str(item)
        weight = item[1] if isinstance(item, (list, tuple)) and len(item) > 1 else 0
        desc = item[2] if isinstance(item, (list, tuple)) and len(item) > 2 else ""
        row_(ws, r, [dim, weight, desc, None, None, None, None],
             formats=[None, "0", None, "0", "0", "0", "0"],
             aligns=["left", "center", "left", "center", "center", "center", "center"],
             hex_bg=zebra(i), height=30)
        ws.cell(r, 1).font = fnt(9.2, True, INK)
        for cc in range(4, 8):
            ws.cell(r, cc).fill = fill("FEF3C7")
        r += 1
    last = r - 1
    if last >= 7:
        band(ws, r, 1, r, 7, NAVY2)
        cell = ws.cell(r, 1, "加权总分（满分 5）")
        cell.font = fnt(9.5, True, WHITE)
        cell.alignment = al("left", "center", False, 1)
        for cc in range(1, 8):
            ws.cell(r, cc).border = BORDER
        wsum = ws.cell(r, 2, f"=SUM(B7:B{last})")
        wsum.font = fnt(9, True, TEAL_M)
        wsum.alignment = al("center", "center", False)
        for cc in range(4, 8):
            col = get_column_letter(cc)
            cell = ws.cell(r, cc,
                           f"=SUMPRODUCT($B$7:$B${last},{col}7:{col}{last})/SUM($B$7:$B${last})")
            cell.font = fnt(11, True, TEAL_M)
            cell.number_format = "0.00"
            cell.alignment = al("center", "center", False)
        ws.row_dimensions[r].height = 28
        dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True)
        dv.error = "请输入 1–5 的整数"
        dv.errorTitle = "评分范围"
        ws.add_data_validation(dv)
        dv.add(f"D7:G{last}")
        end = r + 2
    else:
        end = r + 1
    note(ws, end, "使用说明：① 先给每家供应商打 1–5 分，再乘权重汇总；② 权重可按品类调整；"
                  "③ 建议至少评 3 家，避免只评「已入围」的供应商；④ 评分前先完成「08_尽调清单」。", 7, 34)
    ws.freeze_panes = "A7"
    page(ws, "landscape", "1:6")

    # ---------------------------------------------------------- 10 数据缺口
    ws = new("10_数据缺口", RED)
    gaps = bundle["gaps"] or []
    banner(ws, "数据缺口", f"{len(gaps)} 项：影响、现状、补齐方式与优先级",
           "原则：不估算、不拼凑；缺口本身就是采购决策的风险项", 7)
    header(ws, 6, ["缺口", "影响", "当前状态", "补齐方式", "优先级", "建议责任", "建议截止"])
    widths(ws, {"A": 26, "B": 34, "C": 20, "D": 58, "E": 8, "F": 14, "G": 12})
    r = 7
    for i, g in enumerate(C.order_by_priority(gaps, "priority")):
        row_(ws, r, [g.get("gap"), g.get("impact"), g.get("status"), g.get("fill"),
                     g.get("priority", MISSING), "采购 / 情报", MISSING],
             aligns=["left", "left", "left", "left", "center", "center", "center"],
             hex_bg=zebra(i), height=40)
        ws.cell(r, 1).font = fnt(9.2, True, INK)
        pr = g.get("priority")
        ws.cell(r, 5).font = fnt(9, True, RED if pr == "高" else AMBER if pr == "中高" else SLATE)
        ws.cell(r, 5).fill = fill(RED_L if pr == "高" else AMBER_L if pr == "中高" else SLATE_L)
        r += 1
    note(ws, r + 1, "补齐顺序建议：先补影响「能不能买到」的缺口（项目净增、区域内可替代产能、"
                    "贸易商/渠道），再补影响「买得贵不贵」的缺口（现金成本、库存、区域价差），"
                    "最后补影响「稳不稳」的缺口（质量批次、合规记录）。", 7, 42)
    ws.freeze_panes = "B7"
    if gaps:
        ws.auto_filter.ref = f"A6:G{r - 1}"
    page(ws, "landscape", "1:6")

    # ---------------------------------------------------------- 11 来源与口径
    ws = new("11_来源与口径", SLATE)
    banner(ws, "来源与口径", f"{len(bundle['sources'])} 条来源 + 口径差异 + 使用限制",
           "所有数字都可回溯；不同口径并列，不强行对齐", 7)
    section(ws, 6, "口径说明", 7)
    header(ws, 8, ["项目", "本报告采用口径", "其他公开口径", "处理方式"])
    widths(ws, {"A": 18, "B": 30, "C": 46, "D": 52, "E": 4, "F": 4, "G": 4})
    basis = []
    basis.append(("覆盖范围", scope.get("geography") or MISSING, scope.get("boundary") or MISSING,
                  "只统计口径内的产能；口径外的不混入"))
    basis.append(("基准年 / 预测年", str(year), "、".join(str(y) for y in fwd) or MISSING,
                  "基准年用实测/披露值；预测年只做情景，不做承诺"))
    basis.append(("全国产能", C.fmt_num(metrics.get("capacity_total")),
                  C.fmt_num(metrics.get("capacity_total_note")) if metrics.get("capacity_total_note") else "见来源",
                  "并列展示，不取平均"))
    basis.append(("样本覆盖", f"{metrics.get('sample_coverage')}%",
                  f"{metrics.get('sample_groups')} 集团 / {metrics.get('sample_sites')} 基地",
                  "样本加总除以全国产能；差额为未披露企业"))
    basis.append(("缺失值", MISSING, "不估算、不插值", "缺口单列在「10_数据缺口」"))
    r = 9
    for i, vals in enumerate(basis):
        row_(ws, r, list(vals), aligns=["left", "left", "left", "left"], hex_bg=zebra(i), height=30)
        ws.cell(r, 1).font = fnt(9.5, True, INK)
        r += 1
    for b in (scope.get("basis_notes") or []):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        cell = ws.cell(r, 1, f"· {b}")
        cell.font = fnt(8.5, False, MUTED)
        cell.alignment = al("left", "center", True, 1)
        ws.row_dimensions[r].height = 20
        r += 1
    row = r + 1
    section(ws, row, "来源列表", 7)
    header(ws, row + 2, ["ID", "名称", "发布方", "日期", "URL", "用途", "置信度"])
    widths(ws, {"A": 5, "B": 42, "C": 26, "D": 11, "E": 56, "F": 46, "G": 8})
    r = row + 3
    for i, s in enumerate(bundle["sources"]):
        url = s.get("url", "")
        row_(ws, r, [s.get("id"), s.get("name"), s.get("publisher"), s.get("date"),
                     url or "无链接", s.get("used", ""), C.norm_conf(s.get("confidence"))],
             aligns=["center", "left", "left", "center", "left", "left", "center"],
             hex_bg=zebra(i), height=34)
        ws.cell(r, 1).font = fnt(8.5, True, TEAL)
        ws.cell(r, 2).font = fnt(9, True, INK)
        c_url = ws.cell(r, 5)
        if str(url).startswith("http"):
            c_url.font = Font(name=FONT, size=8.2, color="1D4ED8", underline="single")
            c_url.hyperlink = url
        else:
            c_url.font = fnt(8.2, False, RED)
        r += 1
    last = r - 1
    note(ws, r + 1, "使用限制：① 公开信息研究，不含非公开成本、合同价、库存、客户数据；"
                    "② 产能为名义/设计能力，实际可用量受影响；③ 项目投产存在延期，净增以公告/置换文件为准；"
                    "④ 价格为公开报价，实际成交受批量、账期、运费、质量影响；⑤ 不构成投资建议。", 7, 42)
    ws.freeze_panes = f"B{row + 3}"
    if bundle["sources"]:
        ws.auto_filter.ref = f"A{row + 2}:G{last}"
    page(ws, "landscape", "1:6")

    # ---------------------------------------------------------- 收尾
    wb.properties.title = f"{short}供应格局数据包 {year}"
    wb.properties.subject = "采购视角：产能、区域、工艺、项目、价格与风险"
    wb.properties.creator = "chemical-supply-market-map"
    wb.properties.description = f"数据截止 {cutoff}；公开信息整理，缺口明确标注。"
    wb.active = 0
    for s in wb.worksheets:
        s.sheet_view.tabSelected = (s.title == "00_封面")
    out_dir = os.path.dirname(os.path.abspath(out_path))
    os.makedirs(out_dir, exist_ok=True)
    wb.save(out_path)
    return os.path.getsize(out_path)
