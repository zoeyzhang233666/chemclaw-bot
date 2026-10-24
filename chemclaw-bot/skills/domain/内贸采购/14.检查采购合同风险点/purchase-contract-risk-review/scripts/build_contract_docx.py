"""把案件 JSON 渲染为可阅读的合同 Word 文档（含页码/条款锚点）。

用途：当用户没有真实合同 PDF 时，用本脚本产出一份**明确标注「虚构演示合同」**的
完整合同，作为审查对象与 Skill 示例附件。真实合同请直接使用原件，不要用本脚本改写。

用法:
    python build_contract_docx.py --dir <案例目录> --out <合同.docx>
"""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from common import load_case

FONT = "宋体"
FONT_HEAD = "黑体"


def set_font(run, name=FONT, size=11, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)


def shade(paragraph_or_cell, fill: str):
    el = paragraph_or_cell._element
    pr = el.get_or_add_pPr() if hasattr(el, "get_or_add_pPr") else el.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def add_page_number_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("第 ")
    set_font(run, size=9, color="555555")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    p._p.append(fld)
    run2 = p.add_run(" 页")
    set_font(run2, size=9, color="555555")


def _preamble_lines(c: dict) -> list[str]:
    """首部各句逐行排布，内容与 C01 原文完全一致（仅换行，不改字）。"""
    pre = next((s for s in c["sections"] if s["id"] == "preamble"), None)
    if not pre:
        return []
    text = next(x["text"] for x in c["clauses"] if x["id"] == pre["clause_ids"][0])
    parts = re.split(r"(?<=。)", text)
    return [p for p in parts if p.strip()]


def build(case_dir: Path, out: Path) -> dict:
    case = load_case(case_dir)
    c = case["contract"]
    meta = c.get("contract_meta", {})
    demo = c.get("demo", False)

    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.top_margin, sec.bottom_margin = Cm(2.4), Cm(2.2)
    sec.left_margin, sec.right_margin = Cm(2.8), Cm(2.6)
    add_page_number_footer(sec)

    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    style.paragraph_format.line_spacing = 1.45
    style.paragraph_format.space_after = Pt(5)

    if demo:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        shade(p, "FEF3C7")
        r = p.add_run("【虚构演示合同】本文件为 Skill 示例演示材料，甲乙方名称、金额、条款均为编写，不对应任何真实交易，不得用于真实签约。")
        set_font(r, size=9, bold=True, color="78350F")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(meta.get("title", "采购合同"))
    set_font(r, FONT_HEAD, 18, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"合同编号：{meta.get('contract_no', '')}")
    set_font(r, size=10.5, color="444444")

    doc.add_paragraph()
    # 首部
    # 首部（内容与合同条款 C01 逐字一致，页码标记 第1页）
    for line in _preamble_lines(c):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(line)
        set_font(r, size=11)
    p = doc.add_paragraph()
    r = p.add_run("（原文第 1 页 · 合同第 1 页）")
    set_font(r, size=8.5, color="94A3B8")

    # 条款
    for s in c["sections"]:
        if s["id"] == "preamble":
            continue
        if s["id"].startswith("a"):
            doc.add_page_break()
        ph = doc.add_paragraph()
        ph.paragraph_format.space_before = Pt(10)
        r = ph.add_run(f'{s["no"]}　{s["title"]}')
        set_font(r, FONT_HEAD, 12.5, bold=True)
        ps = doc.add_paragraph()
        ps.paragraph_format.space_after = Pt(2)
        r = ps.add_run(f"（原文第 {s['page']} 页 · 合同第 {s['page']} 页）")
        set_font(r, size=8.5, color="94A3B8")
        for cid in s["clause_ids"]:
            clause = next(x for x in c["clauses"] if x["id"] == cid)
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.left_indent = Cm(0.0)
            if clause.get("kind") == "table_row":
                p.paragraph_format.left_indent = Cm(0.8)
            r = p.add_run(clause["text"])
            set_font(r, size=11)

    # 附件
    for att in c.get("attachments", []):
        doc.add_page_break()
        ph = doc.add_paragraph()
        r = ph.add_run(f'{att["no"]}　{att["title"]}')
        set_font(r, FONT_HEAD, 13, bold=True)
        p = doc.add_paragraph()
        r = p.add_run(f"（原文第 {att['page']} 页 · 合同第 {att['page']} 页）")
        set_font(r, size=8.5, color="94A3B8")
        if att.get("table"):
            t = att["table"]
            tb = doc.add_table(rows=1, cols=len(t["cols"]))
            tb.style = "Table Grid"
            for i, h in enumerate(t["cols"]):
                cell = tb.rows[0].cells[i]
                r = cell.paragraphs[0].add_run(h)
                set_font(r, size=10, bold=True)
                shade(cell, "E8EFF5")
            for row in t["rows"]:
                cells = tb.add_row().cells
                for i, v in enumerate(row):
                    r = cells[i].paragraphs[0].add_run(v)
                    set_font(r, size=10)
            p = doc.add_paragraph()
            r = p.add_run(t.get("note", ""))
            set_font(r, size=10, color="334155")
        if att.get("list"):
            for item in att["list"]:
                p = doc.add_paragraph()
                r = p.add_run(item)
                set_font(r, size=10.5)
                if att.get("kind") == "review_note":
                    shade(p, "F8FAFC")

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("甲方（盖章）：　　　　　　　　　乙方（盖章）：")
    set_font(r, size=11)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("授权代表签字：　　　　　　　　　授权代表签字：")
    set_font(r, size=11)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("日期：　　年　　月　　日　　　　日期：　　年　　月　　日")
    set_font(r, size=11)

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return {"out": str(out), "bytes": out.stat().st_size}


def main() -> None:
    ap = argparse.ArgumentParser(description="由案件 JSON 生成合同 DOCX（演示用）")
    ap.add_argument("--dir", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    print(json.dumps({"ok": True, **build(Path(a.dir), Path(a.out))}, ensure_ascii=False))


if __name__ == "__main__":
    main()
