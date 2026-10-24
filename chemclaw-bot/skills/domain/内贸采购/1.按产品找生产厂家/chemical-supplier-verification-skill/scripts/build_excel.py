#!/usr/bin/env python3
"""Build a 12-sheet Excel workbook from a chemical supplier project.json.

Usage:
    python scripts/build_excel.py path/to/project.json --out path/to/output.xlsx
"""
from __future__ import annotations

import argparse
import json
import math
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NAVY = "0F1C2E"
TEAL = "0F766E"
AMBER = "D97706"
PAPER = "F7F5EF"
LIGHT = "EAF4F1"
WHITE = "FFFFFF"
GRAY = "5B6B7F"
LINE = "D9D5CA"
RED = "B42318"
RED_BG = "FDECEA"
AMBER_BG = "FFF4E5"
GREEN_BG = "E8F5EE"
BLUE_BG = "EAF1F8"

FONT = "Microsoft YaHei"
THIN = Side(style="thin", color=LINE)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

LEVEL_FILL = {
    "直接-A": PatternFill("solid", fgColor=GREEN_BG),
    "直接-B": PatternFill("solid", fgColor=LIGHT),
    "间接-C": PatternFill("solid", fgColor=AMBER_BG),
    "仅工商-D": PatternFill("solid", fgColor=RED_BG),
}
TIER_FILL = {
    "优先谈判": PatternFill("solid", fgColor=TEAL),
    "重点备选": PatternFill("solid", fgColor=AMBER),
}
STATUS_FILL = {
    "部分公开": PatternFill("solid", fgColor=GREEN_BG),
    "未公开": PatternFill("solid", fgColor=AMBER_BG),
    "待验证": PatternFill("solid", fgColor=BLUE_BG),
}


def load(path: Path) -> dict:
    with path.open("r", encoding="utf-8") as f:
        return json.load(f)


def is_num(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v)


def fmt(v, digits=0) -> str:
    if v is None or v == "":
        return "未披露"
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if digits == 0:
            return f"{v:,.0f}"
        return f"{v:,.{digits}f}"
    return str(v)


def pct(v) -> str:
    return f"{v:.0%}" if is_num(v) else "—"


def tier_for(score: float, tiers: list[dict]) -> str:
    for t in sorted(tiers, key=lambda x: x.get("min", 0), reverse=True):
        if score >= t.get("min", 0):
            return t.get("label", "")
    return ""


def calc_composite(c: dict, w: dict) -> float:
    return round(
        c.get("certainty", 0) * w.get("certainty", 0)
        + c.get("spec_score", 0) * w.get("spec", 0)
        + c.get("scale_score", 0) * w.get("scale", 0)
        + c.get("distance_score", 0) * w.get("distance", 0),
        1,
    )


def title(ws, text: str, ncols: int, sub: str | None = None, row: int = 1) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(name=FONT, size=20, bold=True, color=WHITE)
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[row].height = 46
    if sub:
        ws.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=ncols)
        c2 = ws.cell(row=row + 1, column=1, value=sub)
        c2.font = Font(name=FONT, size=10, color="D7E3EA")
        c2.fill = PatternFill("solid", fgColor=NAVY)
        c2.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
        ws.row_dimensions[row + 1].height = 30


def section(ws, row: int, text: str, ncols: int, fill: str = TEAL) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(name=FONT, size=11, bold=True, color=WHITE)
    c.fill = PatternFill("solid", fgColor=fill)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[row].height = 24


def body(ws, row: int, col: int, value, bold=False, size=10, color="1F2937", fill=None,
         align="left", wrap=True, border=True, italic=False):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name=FONT, size=size, bold=bold, color=color, italic=italic)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    if border:
        c.border = BORDER
    return c


def table_header(ws, row: int, headers: list[str], widths: list[int], fill: str = NAVY, height: int = 38) -> None:
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(name=FONT, size=8, bold=True, color=WHITE)
        c.fill = PatternFill("solid", fgColor=fill)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = height


def note_row(ws, row: int, text: str, ncols: int, fill: str = AMBER_BG, color="7A4A00") -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(name=FONT, size=9, color=color)
    c.fill = PatternFill("solid", fgColor=fill)
    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
    ws.row_dimensions[row].height = 34


def set_print(ws, footer: str, landscape=True) -> None:
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = footer
    ws.oddFooter.center.size = 8
    ws.oddFooter.center.color = GRAY


def build(d: dict) -> Workbook:
    M = d["meta"]
    RFQ = d["rfq"]
    SPECS = d["specs"]
    SOURCES = d["sources"]
    COMPANIES = d["companies"]
    TRADERS = d.get("traders", [])
    CONFLICTS = d.get("conflicts", [])
    GAPS = d.get("gaps", [])
    ACTIONS = d.get("actions", [])
    QUESTIONS = d.get("rfq_questions", [])
    RECS = d.get("recommendations", [])
    SCORING = d["scoring"]
    W = SCORING["weights"]
    TIERS = SCORING["tiers"]

    cap_unit = M.get("capacity_unit", "万吨/年")
    qty_unit = M.get("quantity_unit", "吨")
    dest = M.get("destination", "")
    dest_short = M.get("destination_short", dest)
    as_of = M.get("as_of", "")
    product_short = M.get("product_short", M.get("product", "产品"))
    spec_count = len(SPECS)

    plants = [(c, p) for c in COMPANIES for p in c.get("plants", [])]
    capacity_total = sum(c.get("capacity", 0) for c in COMPANIES)
    provinces = {p.get("province", "") for _, p in plants if p.get("province")}
    tiers = Counter(c.get("tier") or tier_for(c.get("composite", 0), TIERS) for c in COMPANIES)
    coa_public = sum(1 for c in COMPANIES if c.get("coa_public"))
    evidence_count = sum(len(c.get("evidence", [])) for c in COMPANIES)
    priority_label = TIERS[0]["label"] if TIERS else "优先"
    priority_min = TIERS[0]["min"] if TIERS else 84

    # normalize
    for c in COMPANIES:
        if c.get("composite") is None:
            c["composite"] = calc_composite(c, W)
        if not c.get("tier"):
            c["tier"] = tier_for(c["composite"], TIERS)
        if c.get("spec_open") is None:
            c["spec_open"] = sum(1 for r in c.get("spec_rows", []) if r.get("status") == "部分公开")
    COMPANIES.sort(key=lambda c: (c.get("rank") or 999, -(c.get("composite") or 0)))

    wb = Workbook()

    # ================================================================ 00 cover
    ws = wb.active
    ws.title = "00_封面"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = NAVY
    for col in "ABCDEFGHIJKLMN":
        ws.column_dimensions[col].width = 13

    title(ws, M.get("title", "生产厂家核验池"), 14,
          f"{M.get('brand', '')} · 按产品找生产厂家 · 示例 RFQ：{RFQ.get('title', '')} · 数据截止 {as_of}", 1)
    ws.row_dimensions[1].height = 54

    kpis = [
        (str(len(COMPANIES)), "生产厂家", "有装置/产能/产量证据"),
        (str(len(plants)), "生产基地", f"覆盖 {len(provinces)} 个省级产区"),
        (fmt(capacity_total), f"样本产能（{cap_unit}）", M.get("sample_note", "")),
        (str(tiers.get(priority_label, 0)), priority_label, f"综合评分 ≥ {priority_min}"),
        (str(len(TRADERS)), "贸易/流通样本", "单独标识，不纳入厂家池"),
        (str(coa_public), "公开批次 COA", f"{spec_count} 项指标无一家公开连续批次数据"),
    ]
    row = 4
    for i, (num, lab, sub) in enumerate(kpis):
        col = 2 + (i % 3) * 4
        r = row + (i // 3) * 3
        ws.merge_cells(start_row=r, start_column=col, end_row=r, end_column=col + 3)
        c = ws.cell(row=r, column=col, value=num)
        c.font = Font(name=FONT, size=24, bold=True, color=TEAL)
        c.fill = PatternFill("solid", fgColor=LIGHT)
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells(start_row=r + 1, start_column=col, end_row=r + 1, end_column=col + 3)
        c = ws.cell(row=r + 1, column=col, value=lab)
        c.font = Font(name=FONT, size=10, bold=True, color=NAVY)
        c.fill = PatternFill("solid", fgColor=LIGHT)
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells(start_row=r + 2, start_column=col, end_row=r + 2, end_column=col + 3)
        c = ws.cell(row=r + 2, column=col, value=sub)
        c.font = Font(name=FONT, size=8, color=GRAY)
        c.fill = PatternFill("solid", fgColor=LIGHT)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[r + 2].height = 34

    row = 11
    section(ws, row, "采购需求卡 / RFQ", 14)
    row += 1
    rfq_rows = [
        ("产品", RFQ.get("product", ""), "CAS / HS", f"{M.get('cas', '—')} / {M.get('hs', '—')}"),
        ("关键规格", RFQ.get("spec_ref", ""), "月需 / 年需", f"{fmt(RFQ.get('qty_month'), 0)} {qty_unit} / {fmt(RFQ.get('qty_year'), 0)} {qty_unit}"),
        ("收货地", RFQ.get("destination", ""), "区域偏好", RFQ.get("region_pref", "")),
        ("下游用途", RFQ.get("application", M.get("application", "")), "包装 / 交付", f"{RFQ.get('packaging', '')}；{RFQ.get('delivery', '')}"),
    ]
    for label1, val1, label2, val2 in rfq_rows:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        body(ws, row, 1, label1, bold=True, fill=PAPER, color=NAVY)
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=7)
        body(ws, row, 3, val1, fill=WHITE)
        ws.merge_cells(start_row=row, start_column=8, end_row=row, end_column=9)
        body(ws, row, 8, label2, bold=True, fill=PAPER, color=NAVY)
        ws.merge_cells(start_row=row, start_column=10, end_row=row, end_column=14)
        body(ws, row, 10, val2, fill=WHITE)
        ws.row_dimensions[row].height = 40
        row += 1

    row += 1
    section(ws, row, "关键规格 / SPECIFICATIONS", 14)
    row += 1
    spec_headers = ["指标", "采购要求", "单位", "检测方法", "为什么关键"]
    for i, h in enumerate(spec_headers, 1):
        body(ws, row, i, h, bold=True, fill=PAPER, color=NAVY, align="center")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=5)
    ws.merge_cells(start_row=row, start_column=6, end_row=row, end_column=8)
    ws.merge_cells(start_row=row, start_column=9, end_row=row, end_column=10)
    ws.merge_cells(start_row=row, start_column=11, end_row=row, end_column=14)
    # Re-place headers after merges
    for col, text in [(1, "指标"), (3, "采购要求"), (6, "单位"), (9, "检测方法"), (11, "为什么关键")]:
        body(ws, row, col, text, bold=True, fill=PAPER, color=NAVY, align="center")
    row += 1
    for i, s in enumerate(SPECS):
        fill = WHITE if i % 2 == 0 else LIGHT
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        body(ws, row, 1, s.get("label", ""), bold=True, fill=fill, color=NAVY)
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=5)
        body(ws, row, 3, s.get("requirement", ""), fill=fill)
        ws.merge_cells(start_row=row, start_column=6, end_row=row, end_column=8)
        body(ws, row, 6, s.get("unit", "—"), fill=fill, align="center")
        ws.merge_cells(start_row=row, start_column=9, end_row=row, end_column=10)
        body(ws, row, 9, s.get("method", "—"), fill=fill)
        ws.merge_cells(start_row=row, start_column=11, end_row=row, end_column=14)
        body(ws, row, 11, s.get("why", "—"), fill=fill)
        ws.row_dimensions[row].height = 30
        row += 1

    row += 1
    section(ws, row, "方法与边界 / METHOD & BOUNDARY", 14)
    row += 1
    method_lines = [
        "生产身份：只有装置/产能/产量/环评/政府项目/官网产品页证据才进入厂家池；经营范围、店铺 SKU、代理宣称不单独作为生产证据。",
        "证据分级：A=定期报告/政府/环评；B=官网/IR/招聘/行业；C=新闻/间接；D=仅工商。每家公司至少一条 A/B 级证据。",
        "产能口径：区分设计/在产/在建/搬迁/集团合并；产能≠产量，在建≠在产，联产装置不能把“双吨”直接当目标产品单吨。",
        "规格匹配：公开标准/官网宣称只写“可对标/部分公开”；没有批次 COA 不写“符合”；关键指标未公开单独标缺口。",
        "排序：综合 = 生产确定性 35% + 规格匹配 30% + 规模 20% + 距离 15%；外销可得性单独展示，不进入分数。",
        "边界：直线距离≠运费；本报告不含报价、账期、运费与最终质量保证；样本覆盖≠市场份额。",
    ]
    for line in method_lines:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=14)
        body(ws, row, 1, "• " + line, fill=WHITE, color="1F2937")
        ws.row_dimensions[row].height = 30
        row += 1

    row += 1
    note_row(ws, row, f"使用顺序：先看 01 候选池与 05 评分 → 再看 02 证据与 03 规格 → 用 06 排除贸易商 → 用 07 缺口和 09 RFQ 行动推进。数据截止 {as_of}；仅作采购初筛。",
             14, fill=AMBER_BG, color="7A4A00")
    set_print(ws, f"{M.get('title', '')} · 数据截止 {as_of} · 仅作采购初筛，不替代现场审核", landscape=True)

    # ================================================================ 01 pool
    ws = wb.create_sheet("01_候选供应商池")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = TEAL
    title(ws, "01 · 候选供应商池", 28,
          f"{len(COMPANIES)} 家生产厂家 / {len(plants)} 个基地；按综合评分排序；贸易商见 06，不混入本表。", 1)
    headers = [
        "排名", "供应商", "代码/属性", "厂家/贸易商", "综合评分", "分级",
        "生产确定性", "规格匹配", "工艺/路线", "主要生产基地", "省/市",
        f"到{dest_short}直线km", "距离带", f"产能（{cap_unit}）", "产能数据年",
        f"产量（{cap_unit}）", "产量数据年", "主营产品", "企业规模/员工",
        "成立年份", "规格公开可对标", "无公开数据指标", "证据等级", "直接证据数",
        "外销可得性", "推荐动作", "数据缺口", "来源ID",
    ]
    widths = [6, 15, 13, 11, 8, 11, 8, 8, 16, 18, 12, 10, 13, 10, 10, 10, 8, 26, 22, 10, 10, 26, 22, 8, 8, 20, 26, 18]
    table_header(ws, 4, headers, widths, height=46)
    r = 5
    for c in COMPANIES:
        levels = " / ".join(sorted({e.get("level", "") for e in c.get("evidence", [])}))
        direct = sum(1 for e in c.get("evidence", []) if str(e.get("level", "")).startswith("直接"))
        spec_gaps = [x.get("label", "") for x in c.get("spec_rows", []) if x.get("status") != "部分公开"]
        nearest = next((p for p in c.get("plants", []) if p.get("name") == c.get("nearest_plant")), (c.get("plants") or [{}])[0])
        province_city = "·".join(x for x in [nearest.get("province", ""), nearest.get("city", "")] if x)
        row_vals = [
            c.get("rank", ""), c.get("name", ""), f"{c.get('ticker', '—')} · {'上市' if c.get('listed') else '非上市'}",
            "生产厂家", c.get("composite", 0), c.get("tier", ""), c.get("certainty", 0), c.get("spec_score", 0),
            c.get("process", ""), c.get("nearest_plant", ""), province_city,
            c.get("distance_km", ""), c.get("distance_band", ""), c.get("capacity", 0), c.get("capacity_year", ""),
            c.get("output") if c.get("output") is not None else "未披露", c.get("output_year", "—"),
            c.get("products", ""), f"{fmt(c.get('employees'))}人" if c.get("employees") else "未披露",
            str(c.get("founded", "")), f"{c.get('spec_open', 0)}/{spec_count}", "、".join(spec_gaps[:4]) + ("…" if len(spec_gaps) > 4 else ""),
            levels, direct, c.get("availability", 0), c.get("action", ""), "；".join(c.get("gaps", [])[:2]),
            "、".join(sorted({sid for p in c.get("plants", []) for sid in p.get("source_ids", [])}))[:60],
        ]
        fill = WHITE if (r % 2) else "FBFCFB"
        for j, v in enumerate(row_vals, 1):
            align = "center" if j in {1, 3, 4, 5, 6, 7, 8, 12, 14, 15, 16, 17, 20, 21, 24, 25} else "left"
            body(ws, r, j, v, size=8, fill=fill, align=align)
        ws.row_dimensions[r].height = 48
        r += 1
    last = r - 1
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:AB{last}"
    ws.conditional_formatting.add(f"E5:E{last}", ColorScaleRule(start_type="num", start_value=60, start_color="FADBD8", mid_type="num", mid_value=80, mid_color="FCF3CF", end_type="num", end_value=100, end_color="D1F2EB"))
    ws.conditional_formatting.add(f"N5:N{last}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=700, color=TEAL))
    ws.conditional_formatting.add(f"F5:F{last}", FormulaRule(formula=['ISNUMBER(SEARCH("优先",F5))'], fill=PatternFill("solid", fgColor=TEAL), font=Font(color=WHITE, bold=True)))
    ws.conditional_formatting.add(f"F5:F{last}", FormulaRule(formula=['ISNUMBER(SEARCH("备选",F5))'], fill=PatternFill("solid", fgColor=AMBER), font=Font(color=WHITE, bold=True)))
    set_print(ws, f"{M.get('title', '')} · 候选池 · 数据截止 {as_of}")

    # ================================================================ 02 evidence
    ws = wb.create_sheet("02_生产证据明细")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = TEAL
    title(ws, "02 · 生产证据明细", 12,
          f"{evidence_count} 条证据；A=定期报告/政府/环评，B=官网/IR/招聘/行业，C=新闻/间接，D=仅工商。", 1)
    headers = ["证据ID", "供应商", "证据等级", "证据类型", "证据标题", "日期", "来源机构", "URL", "关键原文", "证明事项", "来源ID", "可信度"]
    widths = [13, 15, 9, 14, 26, 11, 16, 28, 56, 24, 8, 10]
    table_header(ws, 4, headers, widths, height=36)
    r = 5
    for c in COMPANIES:
        for e in c.get("evidence", []):
            vals = [e.get("id", ""), c.get("name", ""), e.get("level", ""), e.get("type", ""), e.get("title", ""),
                    e.get("date", ""), e.get("org", ""), e.get("url", ""), e.get("quote", ""), e.get("proves", ""),
                    e.get("source_id", ""), "高" if str(e.get("level", "")).startswith("直接-A") else ("中" if str(e.get("level", "")).startswith("直接-B") else "低")]
            for j, v in enumerate(vals, 1):
                cell = body(ws, r, j, v, size=8, fill=WHITE, align="center" if j in {1, 3, 6, 11, 12} else "left")
                if j == 3:
                    cell.fill = LEVEL_FILL.get(str(v), PatternFill("solid", fgColor=WHITE))
                if j == 8 and str(v).startswith("http"):
                    cell.hyperlink = v
                    cell.font = Font(name=FONT, size=8, color="0B63B6", underline="single")
            ws.row_dimensions[r].height = 58
            r += 1
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:L{r-1}"
    set_print(ws, f"{M.get('title', '')} · 生产证据 · 数据截止 {as_of}")

    # ================================================================ 03 wide matrix
    ws = wb.create_sheet("03_规格匹配矩阵")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = TEAL
    ncols = 2 + spec_count * 2
    title(ws, "03 · 规格匹配矩阵", ncols,
          "绿色=部分公开；橙色=未公开；蓝色=待验证。公开标准/宣称不等于批次合格。", 1)
    headers = ["供应商", "综合"]
    widths = [18, 8]
    for s in SPECS:
        headers += [f"{s.get('label', '')}\n公开信息", f"{s.get('label', '')}\n状态/缺口"]
        widths += [30, 34]
    table_header(ws, 4, headers, widths, height=54)
    r = 5
    for c in COMPANIES:
        body(ws, r, 1, c.get("name", ""), bold=True, size=8, fill=WHITE)
        body(ws, r, 2, c.get("composite", 0), size=8, fill=WHITE, align="center")
        for i, s in enumerate(SPECS):
            row = next((x for x in c.get("spec_rows", []) if x.get("id") == s.get("id")), {})
            col = 3 + i * 2
            body(ws, r, col, row.get("public_info", "未公开"), size=8, fill=WHITE)
            cell = body(ws, r, col + 1, f"{row.get('status', '待验证')}｜{row.get('gap', '')}", size=8, fill=WHITE)
            cell.fill = STATUS_FILL.get(row.get("status", ""), PatternFill("solid", fgColor=WHITE))
        ws.row_dimensions[r].height = 42
        r += 1
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:{get_column_letter(ncols)}{r-1}"
    set_print(ws, f"{M.get('title', '')} · 规格矩阵 · 数据截止 {as_of}")

    # ================================================================ 03b long
    ws = wb.create_sheet("03b_规格逐项明细")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = TEAL
    title(ws, "03b · 规格逐项明细", 8,
          "一行一个企业×指标；筛“未公开/待验证”即可生成 RFQ 追问清单。", 1)
    headers = ["供应商", "指标", "采购要求", "状态", "公开信息", "缺口", "下一步动作", "备注"]
    widths = [15, 16, 24, 9, 46, 40, 34, 20]
    table_header(ws, 4, headers, widths, height=34)
    r = 5
    for c in COMPANIES:
        for row in c.get("spec_rows", []):
            vals = [c.get("name", ""), row.get("label", ""), row.get("requirement", ""), row.get("status", ""),
                    row.get("public_info", ""), row.get("gap", ""), row.get("next_action", ""), row.get("note", "")]
            for j, v in enumerate(vals, 1):
                cell = body(ws, r, j, v, size=8, fill=WHITE, align="center" if j in {4} else "left")
                if j == 4:
                    cell.fill = STATUS_FILL.get(str(v), PatternFill("solid", fgColor=WHITE))
            ws.row_dimensions[r].height = 34
            r += 1
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:H{r-1}"
    set_print(ws, f"{M.get('title', '')} · 规格逐项 · 数据截止 {as_of}")

    # ================================================================ 04 plants
    ws = wb.create_sheet("04_工厂与产能")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = TEAL
    title(ws, "04 · 工厂与产能", 15,
          f"{len(plants)} 个基地；产能按基地分列，避免集团产能重复加总；坐标为近似点，仅算直线距离。", 1)
    headers = ["工厂ID", "供应商", "工厂/基地", "省", "市", "详细地址", "经度", "纬度", "工艺", f"产能（{cap_unit}）", "产能数据年", "主要产品/配套", f"到{dest_short}直线km", "距离带", "来源ID"]
    widths = [9, 15, 18, 8, 10, 38, 9, 9, 12, 10, 10, 26, 10, 13, 14]
    table_header(ws, 4, headers, widths, height=40)
    r = 5
    for idx, (c, p) in enumerate(plants, 1):
        vals = [f"P{idx:02d}", c.get("name", ""), p.get("name", ""), p.get("province", ""), p.get("city", ""),
                p.get("address", ""), p.get("lon"), p.get("lat"), p.get("process", ""), p.get("capacity", 0),
                p.get("capacity_year", ""), p.get("products", c.get("products", "")), p.get("distance_km", c.get("distance_km")),
                p.get("distance_band", c.get("distance_band", "")), "、".join(p.get("source_ids", []))]
        for j, v in enumerate(vals, 1):
            body(ws, r, j, v, size=8, fill=WHITE if idx % 2 else "FBFCFB", align="center" if j in {1, 4, 5, 7, 8, 10, 11, 13, 15} else "left")
        ws.row_dimensions[r].height = 36
        r += 1
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:O{r-1}"
    ws.conditional_formatting.add(f"J5:J{r-1}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=max([p.get("capacity", 0) for _, p in plants] or [1]), color=TEAL))
    set_print(ws, f"{M.get('title', '')} · 工厂与产能 · 数据截止 {as_of}")

    # ================================================================ 05 scoring
    ws = wb.create_sheet("05_评分与排序")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = AMBER
    title(ws, "05 · 评分与排序", 15,
          f"综合 = 确定性×{pct(W.get('certainty'))} + 规格×{pct(W.get('spec'))} + 规模×{pct(W.get('scale'))} + 距离×{pct(W.get('distance'))}；外销可得性单列。", 1)
    headers = ["排名", "供应商", "综合评分", "生产确定性", "确定性依据", "规格匹配", "规格依据", "规模分", f"产能（{cap_unit}）", "距离分", f"距离（km）", "加权计算", "分级", "外销可得性", "排序敏感性"]
    widths = [6, 15, 8, 9, 34, 8, 30, 8, 10, 8, 9, 30, 12, 9, 30]
    table_header(ws, 4, headers, widths, height=40)
    r = 5
    for c in COMPANIES:
        calc = f"{c.get('certainty', 0)}×{pct(W.get('certainty'))} + {c.get('spec_score', 0)}×{pct(W.get('spec'))} + {c.get('scale_score', 0)}×{pct(W.get('scale'))} + {c.get('distance_score', 0)}×{pct(W.get('distance'))} = {c.get('composite', 0)}"
        vals = [c.get("rank", ""), c.get("name", ""), c.get("composite", 0), c.get("certainty", 0),
                "；".join(c.get("strengths", [])[:2]), c.get("spec_score", 0), c.get("spec_conclusion", ""),
                c.get("scale_score", 0), c.get("capacity", 0), c.get("distance_score", 0), c.get("distance_km", ""),
                calc, c.get("tier", ""), c.get("availability", 0), c.get("sensitivity", "距离权重↑10pp / 规模权重↑10pp 时需重排")]
        for j, v in enumerate(vals, 1):
            cell = body(ws, r, j, v, size=8, fill=WHITE, align="center" if j in {1, 3, 4, 6, 8, 9, 10, 11, 13, 14} else "left")
            if j == 13:
                cell.fill = TIER_FILL.get(str(v), PatternFill("solid", fgColor="8A94A6"))
                cell.font = Font(name=FONT, size=8, bold=True, color=WHITE)
        ws.row_dimensions[r].height = 48
        r += 1
    ws.auto_filter.ref = f"A4:O{r-1}"
    ws.freeze_panes = "C5"
    ws.conditional_formatting.add(f"C5:C{r-1}", ColorScaleRule(start_type="num", start_value=60, start_color="FADBD8", mid_type="num", mid_value=80, mid_color="FCF3CF", end_type="num", end_value=100, end_color="D1F2EB"))
    set_print(ws, f"{M.get('title', '')} · 评分与排序 · 数据截止 {as_of}")

    # ================================================================ 06 traders
    ws = wb.create_sheet("06_贸易商与排除名单")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = RED
    title(ws, "06 · 贸易商与排除名单", 8,
          "不是“黑名单”，是“生产身份未证实”；单独标识，不进入厂家池。", 1)
    headers = ["名称", "类型", "地区", "证据等级", "公开证据", "为什么不是生产", "处理", "什么情况下可重新考虑"]
    widths = [26, 16, 14, 9, 46, 46, 30, 40]
    table_header(ws, 4, headers, widths, height=36)
    r = 5
    for t in TRADERS:
        vals = [t.get("name", ""), t.get("type", ""), t.get("region", ""), t.get("level", ""), t.get("evidence", ""), t.get("why", ""), t.get("action", ""), t.get("reconsider", "提供生产厂授权链 + 出厂COA + 发货地址 + 产能/环评证明")]
        for j, v in enumerate(vals, 1):
            cell = body(ws, r, j, v, size=8, fill=WHITE, align="center" if j in {3, 4} else "left")
            if j == 4:
                cell.fill = LEVEL_FILL.get(str(v), PatternFill("solid", fgColor=WHITE))
        ws.row_dimensions[r].height = 52
        r += 1
    ws.auto_filter.ref = f"A4:H{r-1}"
    ws.freeze_panes = "B5"
    set_print(ws, f"{M.get('title', '')} · 贸易商与排除 · 数据截止 {as_of}")

    # ================================================================ 07 gaps
    ws = wb.create_sheet("07_数据缺口与复核清单")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = AMBER
    title(ws, "07 · 数据缺口与复核清单", 7,
          "P0=不补不能定主供；P1=影响备选/试单；P2=完善画像。每个缺口必须有复核动作。", 1)
    headers = ["优先级", "数据缺口", "影响企业", "对采购判断的影响", "复核方法", "责任", "截止"]
    widths = [8, 26, 26, 46, 50, 12, 12]
    table_header(ws, 4, headers, widths, height=34)
    r = 5
    for g in GAPS:
        vals = [g.get("priority", ""), g.get("gap", ""), g.get("affected", ""), g.get("impact", ""), g.get("verify", ""), g.get("owner", "采购/质量"), g.get("due", "RFQ后2周")]
        for j, v in enumerate(vals, 1):
            cell = body(ws, r, j, v, size=9, fill=WHITE, align="center" if j in {1, 6, 7} else "left")
            if j == 1:
                cell.fill = PatternFill("solid", fgColor=(RED_BG if v == "P0" else (AMBER_BG if v == "P1" else BLUE_BG)))
                cell.font = Font(name=FONT, size=9, bold=True, color=RED if v == "P0" else "7A4A00")
        ws.row_dimensions[r].height = 54
        r += 1
    ws.auto_filter.ref = f"A4:G{r-1}"
    ws.freeze_panes = "B5"
    set_print(ws, f"{M.get('title', '')} · 数据缺口 · 数据截止 {as_of}")

    # ================================================================ 08 sources
    ws = wb.create_sheet("08_来源库")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = TEAL
    title(ws, "08 · 来源库", 8,
          f"{len(SOURCES)} 个公开来源；A=定期报告/政府/环评；B=官网/IR/招聘/行业；C=新闻/间接；D=仅工商。", 1)
    headers = ["来源ID", "机构", "标题", "日期", "等级", "URL", "用于", "访问日"]
    widths = [8, 22, 50, 12, 6, 44, 56, 12]
    table_header(ws, 4, headers, widths, height=32)
    r = 5
    for s in SOURCES:
        vals = [s.get("id", ""), s.get("org", ""), s.get("title", ""), s.get("date", ""), s.get("grade", ""), s.get("url", ""), s.get("used", ""), s.get("accessed", as_of)]
        for j, v in enumerate(vals, 1):
            cell = body(ws, r, j, v, size=8, fill=WHITE, align="center" if j in {1, 4, 5, 8} else "left")
            if j == 5:
                grade_fill = {"A": GREEN_BG, "B": LIGHT, "C": AMBER_BG, "D": RED_BG}.get(str(v), WHITE)
                cell.fill = PatternFill("solid", fgColor=grade_fill)
            if j == 6 and str(v).startswith("http"):
                cell.hyperlink = v
                cell.font = Font(name=FONT, size=8, color="0B63B6", underline="single")
        ws.row_dimensions[r].height = 46
        r += 1
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:H{r-1}"
    set_print(ws, f"{M.get('title', '')} · 来源库 · 数据截止 {as_of}")

    # ================================================================ 09 RFQ & actions
    ws = wb.create_sheet("09_RFQ与行动清单")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = TEAL
    title(ws, "09 · RFQ 与行动清单", 10,
          "统一规格/数量/收货地询价；先收 COA，再小试，再审厂。", 1)
    row = 4
    section(ws, row, "A. 统一 RFQ 必问", 10)
    row += 1
    qlist = QUESTIONS or [
        "请确认生产工厂全称、地址、是否自有装置。",
        "请提供产能、实际产量、开工率与检修计划。",
        "请确认月度可供量、长协占比、最小起订量。",
        "请提供连续 3 批 COA 与检测方法。",
        "请报出厂价/到厂价、包装、账期与到货周期。",
        "请提供环评/排污许可/安全生产许可。",
        "请提供下游客户案例/第三方检测/审厂报告。",
        "是否接受样品小试与现场审核？",
    ]
    for i, q in enumerate(qlist, 1):
        body(ws, row, 1, f"Q{i:02d}", bold=True, fill=PAPER, color=NAVY, align="center")
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=10)
        body(ws, row, 2, q, fill=WHITE)
        ws.row_dimensions[row].height = 30
        row += 1

    row += 1
    section(ws, row, "B. 规格复核清单", 10)
    row += 1
    spec_headers = ["指标", "采购要求", "检测方法", "为什么关键"]
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=5)
    ws.merge_cells(start_row=row, start_column=6, end_row=row, end_column=7)
    ws.merge_cells(start_row=row, start_column=8, end_row=row, end_column=10)
    for col, text in [(1, "指标"), (3, "采购要求"), (6, "检测方法"), (8, "为什么关键")]:
        body(ws, row, col, text, bold=True, fill=PAPER, color=NAVY, align="center")
    row += 1
    for s in SPECS:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        body(ws, row, 1, s.get("label", ""), bold=True, fill=WHITE, color=NAVY)
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=5)
        body(ws, row, 3, s.get("requirement", ""), fill=WHITE)
        ws.merge_cells(start_row=row, start_column=6, end_row=row, end_column=7)
        body(ws, row, 6, s.get("method", "—"), fill=WHITE)
        ws.merge_cells(start_row=row, start_column=8, end_row=row, end_column=10)
        body(ws, row, 8, s.get("why", "—"), fill=WHITE)
        ws.row_dimensions[row].height = 30
        row += 1

    row += 1
    section(ws, row, "C. 30 天行动", 10)
    row += 1
    action_headers = ["阶段", "动作", "说明"]
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=4)
    ws.merge_cells(start_row=row, start_column=5, end_row=row, end_column=10)
    for col, text in [(1, "阶段"), (3, "动作"), (5, "说明")]:
        body(ws, row, col, text, bold=True, fill=PAPER, color=NAVY, align="center")
    row += 1
    for a in ACTIONS:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        body(ws, row, 1, a.get("phase", ""), bold=True, fill=WHITE, color=NAVY, align="center")
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=4)
        body(ws, row, 3, a.get("action", ""), fill=WHITE)
        ws.merge_cells(start_row=row, start_column=5, end_row=row, end_column=10)
        body(ws, row, 5, a.get("detail", ""), fill=WHITE)
        ws.row_dimensions[row].height = 38
        row += 1

    if RECS:
        row += 1
        section(ws, row, "D. 主供/备供建议", 10)
        row += 1
        for rec in RECS:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
            body(ws, row, 1, rec.get("tier", ""), bold=True, fill=LIGHT, color=NAVY, align="center")
            ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=4)
            body(ws, row, 3, rec.get("company", ""), bold=True, fill=WHITE, color=TEAL)
            ws.merge_cells(start_row=row, start_column=5, end_row=row, end_column=10)
            body(ws, row, 5, rec.get("note", ""), fill=WHITE)
            ws.row_dimensions[row].height = 34
            row += 1

    for col in "ABCDEFGHIJ":
        ws.column_dimensions[col].width = 14
    set_print(ws, f"{M.get('title', '')} · RFQ 与行动 · 数据截止 {as_of}")

    # ================================================================ 10 conflicts
    ws = wb.create_sheet("10_口径与冲突")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = AMBER
    title(ws, "10 · 口径与冲突", 5,
          "同一指标出现两个口径时并列保留，不强行取一个；写清对判断的影响。", 1)
    headers = ["主题", "口径 A", "口径 B", "本报告处理", "对判断的影响"]
    widths = [26, 46, 46, 52, 40]
    table_header(ws, 4, headers, widths, height=32)
    r = 5
    for c in CONFLICTS:
        vals = [c.get("topic", ""), c.get("a", ""), c.get("b", ""), c.get("resolution", ""), c.get("impact", "")]
        for j, v in enumerate(vals, 1):
            body(ws, r, j, v, size=9, fill=WHITE)
        ws.row_dimensions[r].height = 56
        r += 1
    ws.auto_filter.ref = f"A4:E{r-1}"
    ws.freeze_panes = "B5"
    set_print(ws, f"{M.get('title', '')} · 口径与冲突 · 数据截止 {as_of}")

    return wb


def main() -> int:
    ap = argparse.ArgumentParser(description="Build Excel from project.json")
    ap.add_argument("project", type=Path)
    ap.add_argument("--out", type=Path, required=True)
    args = ap.parse_args()
    d = load(args.project)
    wb = build(d)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    print(f"saved {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
