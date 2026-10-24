# -*- coding: utf-8 -*-
"""交付前自检：占位符 / 外链 / 禁用词 / 分位来源 / 口径不匹配标注 / HTML-Excel 同源。

用法：python scripts/check_output.py <model.json> <out_dir>
返回码 0 = 全部通过。
"""
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import common

FORBIDDEN = ['应当还价', '建议还价', '目标价', '底价', '可接受价格', '预计将', '预测价格', '看涨', '看跌',
             '有望上涨', '有望下跌', '见底', '见顶', '抄底', '后市看', '将会上涨', '将会下跌', '建议压价',
             '砍价到', '合理成交价应为']
PLACEHOLDER = ['__TITLE__', '__SECTIONS__', '__DATA__', '__APP__', '__ECHARTS__', '__BADGES__',
               '__SUBTITLE__', '__FOOTER__', '__H1__', '{{', 'TODO', 'FIXME', 'undefined', 'NaN']


def check(model_path, out_dir):
    m = common.load_json(model_path)
    fails, warns = [], []

    htmls = [f for f in os.listdir(out_dir) if f.endswith('.html')]
    xlsxs = [f for f in os.listdir(out_dir) if f.endswith('.xlsx')]
    if len(htmls) != 1:
        fails.append('应当且只应有 1 个 HTML 报告，实际 %d 个' % len(htmls))
    if len(xlsxs) != 1:
        fails.append('应当且只应有 1 个 Excel 工作簿，实际 %d 个' % len(xlsxs))
    if not htmls or not xlsxs:
        return fails, warns

    hp = os.path.join(out_dir, htmls[0])
    html = open(hp, encoding='utf-8').read()
    size_mb = os.path.getsize(hp) / 1024 / 1024

    # 1) 体积（ECharts 应已内联）
    if size_mb < 0.8:
        fails.append('HTML 仅 %.2f MB，ECharts 可能未内联' % size_mb)

    # 2) 占位符（只查可见文本；undefined/NaN 属 JS 字面量，需在正文出现才算异常）
    visible0 = re.sub(r'<script.*?</script>', '', html, flags=re.S)
    visible0 = re.sub(r'<style.*?</style>', '', visible0, flags=re.S)
    for p in PLACEHOLDER:
        hit = p in visible0
        if hit:
            fails.append('HTML 残留占位符/异常值：%s' % p)

    # 3) 外链（离线可用）
    ext = re.findall(r'<(?:script|link)[^>]+(?:src|href)="(https?://[^"]+)"', html)
    if ext:
        fails.append('HTML 存在外部脚本/样式依赖（离线打不开）：%s' % ext[:3])

    # 4) 禁用词：只检查正文，不检查内联 JSON 数据中的"事件事实"原文
    visible = re.sub(r'<script.*?</script>', '', html, flags=re.S)
    visible = re.sub(r'<style.*?</style>', '', visible, flags=re.S)
    for w in FORBIDDEN:
        if w in visible:
            fails.append('正文出现禁用表述（预测/议价建议）：%s' % w)

    # 5) 分位只能来自序列：窗口 ok 的必须有 n 与起止日期
    for p in m['percentiles']:
        if p.get('ok') and not (p.get('n') and p.get('start') and p.get('end')):
            fails.append('分位窗口 %s 缺少样本数或起止日期' % p['window'])
        if p.get('ok') and p.get('quote_pctile') is None:
            fails.append('分位窗口 %s 标记 ok 但无分位值' % p['window'])
        if not p.get('ok') and p.get('quote_pctile') is not None:
            fails.append('分位窗口 %s 为缺口却给出了分位值' % p['window'])

    # 6) 口径不匹配必须在 HTML 中显式提示
    qspec = (m['quote'].get('grade_claimed') or '')
    mspec = (m['primary_anchor'].get('spec') or '')
    if qspec and mspec and qspec.split('（')[0] not in mspec:
        if '口径不匹配' not in visible:
            fails.append('报价规格「%s」与参考规格「%s」不一致，但 HTML 未标注口径不匹配' % (qspec, mspec))

    # 7) 测试输入必须显著标注
    if m['quote'].get('is_test_input') and '测试输入' not in visible:
        fails.append('报价为测试输入，但 HTML 未显著标注')

    # 8) 缺口必须显示
    for g in m.get('gaps') or []:
        if g['id'] not in visible:
            warns.append('缺口 %s 未在 HTML 中出现' % g['id'])

    # 9) 期货不得直接与现货相减
    if m.get('futures_reference'):
        if '不能直接' not in visible and '不得用期货' not in visible:
            fails.append('报告含期货数据但未声明期货与现货不可直接相减')

    # 10) Excel 同源：抽查关键数字出现在 xlsx
    try:
        from openpyxl import load_workbook
        wb = load_workbook(os.path.join(out_dir, xlsxs[0]), data_only=True)
        names = wb.sheetnames
        for must in ('价格历史', '报价比较', '计算公式', '数据来源'):
            if not any(must in n for n in names):
                # 允许语义等价命名
                alias = {'价格历史': ['价格历史', '历史'], '报价比较': ['偏离', '比较'],
                         '计算公式': ['公式'], '数据来源': ['来源']}[must]
                if not any(any(x in n for x in alias) for n in names):
                    fails.append('Excel 缺少必需表：%s（现有：%s）' % (must, '、'.join(names[:14])))
        if len(names) < 10:
            warns.append('Excel 仅 %d 张表' % len(names))
        # 抽查：首个靶心参考中枢应出现在 04 表
        t0 = [x for x in m['targets'] if x['status'] == '可比']
        if t0:
            tgt = [n for n in names if '偏离' in n or '比较' in n][0]
            vals = []
            for row in wb[tgt].iter_rows(values_only=True):
                vals += [v for v in row if isinstance(v, (int, float))]
            if round(float(t0[0]['center']), 2) not in [round(v, 2) for v in vals]:
                fails.append('Excel 04 表中未找到 HTML 显示的参考中枢 %.2f（两文件不同源）' % t0[0]['center'])
        n_formula_sheet = [n for n in names if '公式' in n]
        if n_formula_sheet:
            fr = [c for c in wb[n_formula_sheet[0]].iter_rows(values_only=True)]
            if sum(len([x for x in r if x]) for r in fr) < 10:
                warns.append('计算公式表内容过少')
    except ImportError:
        warns.append('未安装 openpyxl，跳过 Excel 同源检查')

    return fails, warns


def main():
    model_path, out_dir = sys.argv[1], sys.argv[2]
    fails, warns = check(model_path, out_dir)
    print('=' * 60)
    for w in warns:
        print('WARN  %s' % w)
    for f in fails:
        print('FAIL  %s' % f)
    print('=' * 60)
    print('结论：%s（失败 %d，告警 %d）' % ('通过' if not fails else '未通过', len(fails), len(warns)))
    return 0 if not fails else 1


if __name__ == '__main__':
    sys.exit(main())
