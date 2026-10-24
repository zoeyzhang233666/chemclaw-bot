#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""make_quote_template.py — 生成《供应商报价单》标准模板（xlsx），可直接发给供应商填写。

    python templates/make_quote_template.py [输出路径]

设计原则：
  * 一个产品、一份报价单，所有口径字段强制填写，避免“出厂/到厂”“含税/不含税”混着比；
  * 两档量价（长协年量与一次性试单）分开填，防止用试单价冒充长约价；
  * 附带“填写说明”页与下拉校验，供应商不会填错栏位。
"""
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

INK = "1F2D36"
BRAND = "0F3B4C"
PAPER = "F7F5EF"
LINE = "D9D3C6"
FILL_IN = "FFFDF4"
thin = Side(style="thin", color=LINE)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def head(ws, text, sub, ncols=6):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=text)
    c.font = Font(name="微软雅黑", size=14, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=BRAND)
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 30
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    c = ws.cell(row=2, column=1, value=sub)
    c.font = Font(name="微软雅黑", size=9.5, color="5A6670")
    c.fill = PatternFill("solid", fgColor=PAPER)
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[2].height = 24


def row(ws, r, values, widths=None, bold=False, fill=None, height=None):
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=r, column=i, value=v)
        c.font = Font(name="微软雅黑", size=10, bold=bold, color=INK)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = BORDER
        if fill:
            c.fill = PatternFill("solid", fgColor=fill)
    if height:
        ws.row_dimensions[r].height = height
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w


def build(out):
    wb = Workbook()

    # ---------- 报价单 ----------
    ws = wb.active
    ws.title = "供应商报价单"
    head(ws, "供应商报价单（标准模板）", "请逐项填写，空白项视为未报价；每一栏的口径都必须明确，否则无法与其他供应商比较", 6)
    widths = [26, 22, 22, 26, 34, 10]

    r = 3
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    c = ws.cell(row=r, column=1, value="一、基本信息")
    c.font = Font(name="微软雅黑", size=11, bold=True, color=BRAND)
    r = 4
    base = [
        ("供应商全称（工商注册名称）", ""),
        ("统一社会信用代码", ""),
        ("报价主体与生产主体是否一致", "一致 / 不一致（若不一致请说明关联关系）"),
        ("生产企业名称与厂址", ""),
        ("报价日期", ""),
        ("报价有效期（至 YYYY-MM-DD）", ""),
        ("报价人 / 联系电话 / 邮箱", ""),
        ("销售负责人（可加签合同）", ""),
    ]
    for k, v in base:
        row(ws, r, [k, v], widths, height=26)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        ws.cell(row=r, column=2).fill = PatternFill("solid", fgColor=FILL_IN)
        r += 1

    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    c = ws.cell(row=r, column=1, value="二、产品与价格（两档量价请分别填写，不得用试单价代替长协价）")
    c.font = Font(name="微软雅黑", size=11, bold=True, color=BRAND)
    r += 1
    row(ws, r, ["项目", "长协口径（年量）", "试单口径（一次性）", "单位", "备注 / 口径说明", "必填"], widths, bold=True,
        fill="17495C", height=30)
    for i in range(1, 7):
        ws.cell(row=r, column=i).font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
    r += 1
    items = [
        ("产品名称与牌号", "", "", "—", "含等级/执行标准", "是"),
        ("年/批供货量", "", "", "吨", "长协请注明月度均衡量", "是"),
        ("单价（含税）", "", "", "元/吨", "13% 增值税专用发票", "是"),
        ("单价（不含税）", "", "", "元/吨", "用于跨税率比较", "是"),
        ("价格类型", "", "", "—", "出厂价 / 送到价（到厂价）", "是"),
        ("含税口径", "", "", "—", "含税 / 不含税 / 普票不可抵扣", "是"),
        ("运费承担方", "", "", "—", "供方 / 需方 / 含运价（元/吨）", "是"),
        ("包装方式", "", "", "—", "散装罐车 / 吨袋 / 编织袋，注明净重", "是"),
        ("包装费用", "", "", "元/吨", "若含在单价内请写“已含”", "是"),
        ("装卸费 / 港杂费", "", "", "元/吨", "无则填 0", "是"),
        ("结算方式", "", "", "—", "预付款 / 货到付款 / 月结 N 天", "是"),
        ("承兑汇票接受情况", "", "", "—", "接受/不接受，贴息由谁承担", "是"),
        ("交期（下单至到货）", "", "", "天", "常规交期 + 紧急插单最短交期", "是"),
        ("最小起订量 MOQ", "", "", "吨", "按车/按批说明", "是"),
        ("月度可供量上限", "", "", "吨/月", "旺季是否会限量", "是"),
        ("质量指标（主含量/杂质/物性）", "", "", "—", "按我方标准逐项对应，附检测报告", "是"),
        ("检测报告与批次追溯", "", "", "—", "随货检测单是否可查批号", "是"),
        ("质保/质量异议条款", "", "", "—", "异议期与处理方式", "建议"),
        ("价格调整机制", "", "", "—", "固定价 / 月度联动（挂什么指数）", "建议"),
        ("年降/量价优惠", "", "", "—", "达到多少量给什么价", "建议"),
    ]
    for it in items:
        row(ws, r, list(it), widths, height=26)
        for col in (2, 3, 5):
            ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor=FILL_IN)
        r += 1

    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    c = ws.cell(row=r, column=1, value="三、资质与合规（请随报价单一并提交扫描件）")
    c.font = Font(name="微软雅黑", size=11, bold=True, color=BRAND)
    r += 1
    row(ws, r, ["文件", "是否提供", "证号", "有效期至", "覆盖范围是否包含本产品", "必填"], widths, bold=True, fill="17495C", height=28)
    for i in range(1, 7):
        ws.cell(row=r, column=i).font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
    r += 1
    docs = [
        ("营业执照", "是"), ("排污许可证", "是"), ("安全生产许可证（如涉）", "建议"),
        ("危险化学品经营/生产许可（如涉）", "建议"), ("食品生产许可证（食品级产品）", "视产品"),
        ("ISO 9001 质量管理体系", "建议"), ("ISO 14001 环境管理体系", "建议"),
        ("近 6 个月第三方检测报告", "是"), ("信用中国无重大违法记录证明", "是"),
        ("近 3 年重大诉讼与行政处罚清单", "是"),
    ]
    dv_yn = DataValidation(type="list", formula1='"是,否,不适用"', allow_blank=True)
    ws.add_data_validation(dv_yn)
    for name, must in docs:
        row(ws, r, [name, "", "", "", "", must], widths, height=24)
        for col in (2, 3, 4, 5):
            ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor=FILL_IN)
        dv_yn.add(ws.cell(row=r, column=2))
        r += 1

    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    c = ws.cell(row=r, column=1,
                value="声明：以上报价与资料真实、完整。若中标，本报价单构成合同附件。—— 供应商盖章：__________  日期：__________")
    c.font = Font(name="微软雅黑", size=9.5, color="9C2A2A")
    c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[r].height = 40
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"

    # ---------- 填写说明 ----------
    ws2 = wb.create_sheet("填写说明")
    head(ws2, "填写说明", "报价格式不统一，是供应商比较失败最常见的原因。以下 6 条请务必遵守", 3)
    notes = [
        ("1. 口径必须写到栏位里", "“出厂价 1800 元”无法比较：“到厂含税价 2030 元”才可以。价格类型、含税口径、运费承担三栏必须同时填。"),
        ("2. 两档量价分开填", "年量长协与一次性试单的成本结构不同，用试单价冒充长约价会在比价时失真。"),
        ("3. 包装与装卸单列", "散装罐车与吨袋的差价通常在 40–120 元/吨，混在单价里无法比较，也无法优化。"),
        ("4. 账期折算成钱", "月结 30 天、承兑 6 个月、预付 30% 的资金成本不同，比较时应折算为“到厂含税等效价”。"),
        ("5. 交期要区分常规与紧急", "常规交期决定安全库存，紧急插单能力决定断供时的生死。"),
        ("6. 资质随报价提交", "没有排污许可、（食品级）食品生产许可的报价，即使价格最低也不能进入下一步。"),
    ]
    r = 4
    for k, v in notes:
        row(ws2, r, [k, v, ""], [30, 96, 3], height=40)
        ws2.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        r += 1
    ws2.sheet_view.showGridLines = False

    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    wb.save(out)
    print("报价单模板 -> %s" % out)


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else os.path.join(
        os.path.dirname(os.path.abspath(__file__)), "报价单模板.xlsx"))
