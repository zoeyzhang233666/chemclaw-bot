# -*- coding: utf-8 -*-
"""Static QA for the generated HTML + Excel pair. No browser required.

Usage:
    python qa.py data.json out/report.html out/report.xlsx
"""
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

CATS = ["工商", "司法", "处罚", "安全环保", "经营动态"]
EXPECTED_SHEETS = ["00_封面", "01_风险总览", "02_事件时间线", "03_分类汇总",
                   "04_供应商画像", "05_待核实清单", "06_来源清单", "07_方法与口径", "08_行业背景"]

fails, notes = [], []


def check(cond, msg):
    if not cond:
        fails.append(msg)


def main():
    if len(sys.argv) < 4:
        raise SystemExit("usage: python qa.py data.json report.html report.xlsx")
    data_path, html_path, xlsx_path = map(Path, sys.argv[1:4])
    data = json.loads(data_path.read_text(encoding="utf-8"))
    html = html_path.read_text(encoding="utf-8")

    # ---------- HTML ----------
    check(html.lstrip().startswith("<!DOCTYPE html>"), "html: missing doctype")
    check('lang="zh-CN"' in html, "html: missing lang=zh-CN")
    check("%%" not in html, "html: unreplaced %% placeholder")
    check("<script src=" not in html and "<link rel=\"stylesheet\"" not in html,
          "html: external script/stylesheet found; must be self-contained")
    check('id="fullTimeline"' in html and 'id="supplierGrid"' in html and 'id="modal"' in html,
          "html: missing core containers")

    try:
        start = html.index("const DATA = ") + len("const DATA = ")
        end = html.index("\n\nconst CATS =", start)
        payload = json.loads(html[start:end].rstrip().rstrip(";"))
        notes.append(f"html data: {len(payload.get('suppliers', []))} suppliers / {len(payload.get('events', []))} events")
        check(len(payload.get("events", [])) == len(data["events"]), "html: injected event count mismatch")
        check(len(payload.get("suppliers", [])) == len(data["suppliers"]), "html: injected supplier count mismatch")
    except Exception as e:
        fails.append(f"html: cannot parse injected DATA ({e})")

    check(len(html.encode("utf-8")) > 30_000, "html: suspiciously small")
    for e in data["events"]:
        check(e["id"] in html, f"html: event id missing {e['id']}")
    for s in data["suppliers"]:
        check(s["name"] in html, f"html: supplier missing {s['name']}")

    # ---------- Excel ----------
    wb = load_workbook(xlsx_path, data_only=False)
    names = wb.sheetnames
    check(names == EXPECTED_SHEETS, f"xlsx: sheet list mismatch: {names}")
    if "01_风险总览" in names:
        ws = wb["01_风险总览"]
        text = "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
        for s in data["suppliers"]:
            check(s["name"] in text, f"xlsx overview: supplier missing {s['name']}")
    if "02_事件时间线" in names:
        ws = wb["02_事件时间线"]
        ids = {str(ws.cell(row=r, column=1).value) for r in range(5, ws.max_row + 1)}
        for e in data["events"]:
            check(e["id"] in ids, f"xlsx timeline: event missing {e['id']}")
        notes.append(f"xlsx timeline rows: {ws.max_row - 4}")
    if "06_来源清单" in names:
        ws = wb["06_来源清单"]
        urls = {str(c.hyperlink.target) for row in ws.iter_rows() for c in row if c.hyperlink}
        expected = {e["url"] for e in data["events"]}
        check(expected.issubset(urls), f"xlsx sources: missing {len(expected - urls)} links")
        notes.append(f"xlsx source links: {len(urls)}")
    if "07_方法与口径" in names:
        ws = wb["07_方法与口径"]
        text = "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
        check("分类覆盖检查" in text, "xlsx method: coverage section missing")
        for c in data.get("coverage_check", []):
            check(c["cat"] in text, f"xlsx method: coverage row missing {c['cat']}")

    print(f"html: {html_path} ({html_path.stat().st_size:,} bytes)")
    print(f"xlsx: {xlsx_path} ({xlsx_path.stat().st_size:,} bytes, {len(names)} sheets)")
    for n in notes:
        print(" ", n)
    if fails:
        print("\nFAIL")
        for f in fails:
            print("  -", f)
        return 1
    print("\nPASS (static QA: structure, data injection, sheet/link coverage)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
