# -*- coding: utf-8 -*-
"""体检：HTML / Excel 是否真的装下了 data.json 的全部内容，且两边数字一致。"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import common  # noqa: E402

OK, BAD = "PASS", "FAIL"


def check(cond: bool, name: str, detail: str = "") -> bool:
    print(f"[{OK if cond else BAD}] {name}" + (f" — {detail}" if detail else ""))
    return cond


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("data")
    ap.add_argument("html")
    ap.add_argument("xlsx")
    args = ap.parse_args()

    data = common.load(args.data)
    payload = common.build_payload(data)
    n_reg = len(data["regions"])
    n_date = len(data["series"]["dates"])

    html = Path(args.html).read_text(encoding="utf-8")
    ok = True

    ok &= check(len(html) > 300_000, "HTML 体积", f"{len(html)/1024/1024:.2f} MB")
    ok &= check("%%" not in html, "无未替换占位符")
    ok &= check(html.count("<script") >= 3, "脚本块", str(html.count("<script")))
    ok &= check("window.__DATA__" in html, "内嵌数据")
    for cid in ("priceChart", "spreadChart", "decompChart", "drawer", "matrix", "kpis"):
        ok &= check(f'id="{cid}"' in html, f"容器 #{cid}")

    m = re.search(r"window\.__DATA__ = (.*?);</script>", html, re.S)
    ok &= check(bool(m), "数据可解析")
    if m:
        embed = json.loads(m.group(1))
        ok &= check(len(embed["regions"]) == n_reg, "HTML 区域数", f'{len(embed["regions"])}')
        ok &= check(len(embed["series"]["dates"]) == n_date, "HTML 样本数", f'{len(embed["series"]["dates"])}')
        ok &= check("matrix" in embed and len(embed["matrix"]) == n_reg * n_reg, "运距/运费矩阵完整")

    wb = load_workbook(args.xlsx, data_only=True)
    names = wb.sheetnames
    need = ["00 概览", "01 口径与时间轴", "02 价格序列", "03 价差分解", "04 价差矩阵",
            "05 价差统计", "06 运费矩阵", "07 供需与客户", "08 数据缺口", "09 方法与免责"]
    ok &= check(len(names) == 10 and names == need, "Excel 10 张表", " / ".join(names))

    ws2 = wb["02 价格序列"]
    dates_in_book = sum(1 for r in range(3, ws2.max_row + 1) if ws2.cell(row=r, column=1).value)
    ok &= check(dates_in_book == n_date, "价格序列行数", f"{dates_in_book} / {n_date}")

    ws3 = wb["03 价差分解"]
    prices = {x["name"]: x["price"] for x in payload["region_stats"]}
    hits = 0
    for r in range(3, ws3.max_row + 1):
        nm = ws3.cell(row=r, column=1).value
        if nm in prices and ws3.cell(row=r, column=2).value == prices[nm]:
            hits += 1
    ok &= check(hits == n_reg, "Excel 价格与数据一致", f"{hits}/{n_reg}")

    ws7 = wb["07 供需与客户"]
    supply_rows = sum(1 for r in range(3, ws7.max_row + 1) if ws7.cell(row=r, column=1).value)
    ok &= check(supply_rows == n_reg, "供需表区域行数", f"{supply_rows}")

    charts = sum(len(ws._charts) for ws in wb.worksheets)
    ok &= check(charts >= 2, "Excel 图表数量", str(charts))

    ws8 = wb["08 数据缺口"]
    gap_rows = sum(1 for r in range(3, ws8.max_row + 1) if ws8.cell(row=r, column=2).value)
    expect_gaps = sum(len(r.get("gaps", [])) for r in data["regions"])
    ok &= check(gap_rows >= expect_gaps, "缺口逐条入表", f"{gap_rows} ≥ {expect_gaps}")

    print("\n体检结论：" + ("全部通过" if ok else "存在失败项"))
    raise SystemExit(0 if ok else 1)


if __name__ == "__main__":
    main()
