# -*- coding: utf-8 -*-
"""生成 10 表 Excel：结论 / 口径 / 序列 / 价差 / 矩阵 / 统计 / 运费 / 供需 / 缺口 / 方法。"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import common  # noqa: E402

NAVY = "0B2136"
NAVY2 = "12314B"
AMBER = "B8791A"
INK = "101C2C"
INK2 = "31485F"
MUTED = "6B7F93"
LINE = "E3E9F0"
BG = "F4F7FA"
CARDBG = "FBFDFF"
RED = "B42318"
GREEN = "0D7A63"
BLUE = "1D4ED8"

F_TITLE = Font(name="微软雅黑", size=19, bold=True, color="FFFFFF")
F_SUB = Font(name="微软雅黑", size=10.5, color="C6D6E6")
F_H = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
F_B = Font(name="微软雅黑", size=10, color=INK)
F_B2 = Font(name="微软雅黑", size=10, color=INK2)
F_M = Font(name="微软雅黑", size=9, color=MUTED)
F_BIG = Font(name="微软雅黑", size=20, bold=True, color=INK)
F_SEC = Font(name="微软雅黑", size=11.5, bold=True, color=NAVY)
F_LINK = Font(name="微软雅黑", size=9, color=BLUE, underline="single")

THIN = Side(style="thin", color=LINE)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NOB = Border()

FILL_NAVY = PatternFill("solid", fgColor=NAVY)
FILL_HEAD = PatternFill("solid", fgColor=NAVY2)
FILL_LIGHT = PatternFill("solid", fgColor=BG)
FILL_CARD = PatternFill("solid", fgColor=CARDBG)
FILL_AMBER = PatternFill("solid", fgColor="FDF3E1")

NUM = '#,##0'
NUM1 = '#,##0.0'
DELTA = '+#,##0;-#,##0;0'
PCT = '0%'


def col_w(ws, widths: dict) -> None:
    for k, v in widths.items():
        ws.column_dimensions[k].width = v


def sec(ws, row: int, text: str, span: int = 8) -> int:
    ws.cell(row=row, column=1, value=text).font = F_SEC
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    ws.row_dimensions[row].height = 22
    return row + 1


def head(ws, row: int, labels: list[str]) -> int:
    for i, t in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=t)
        c.font = F_H
        c.fill = FILL_HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
    ws.row_dimensions[row].height = 30
    return row + 1


def body(ws, row: int, values: list, fmts: dict | None = None, align: str = "right") -> int:
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = F_B
        c.border = BOX
        c.alignment = Alignment(horizontal=("left" if i == 1 else align), vertical="center")
        if fmts and i in fmts:
            c.number_format = fmts[i]
    return row + 1


def kv_pairs(ws, row: int, pairs: list[tuple[str, str]], span_key=1, span_val=7) -> int:
    for k, v in pairs:
        a = ws.cell(row=row, column=1, value=k)
        a.font = F_B2
        a.alignment = Alignment(vertical="top", wrap_text=True)
        if span_key > 1:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span_key)
        b = ws.cell(row=row, column=2, value=v)
        b.font = F_M
        b.alignment = Alignment(vertical="top", wrap_text=True)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=span_val)
        ws.row_dimensions[row].height = max(18, 15 * (1 + len(str(v)) // 46))
        row += 1
    return row


def cover(ws, data: dict, payload: dict) -> None:
    ws.sheet_view.showGridLines = False
    col_w(ws, {"A": 16, "B": 18, "C": 18, "D": 18, "E": 18, "F": 18, "G": 18, "H": 12})
    ws.merge_cells("A1:H1")
    c = ws["A1"]
    c.value = data["meta"]["title"]
    c.font = F_TITLE
    c.fill = FILL_NAVY
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 46

    ws.merge_cells("A2:H2")
    s = ws["A2"]
    s.value = f'{data["meta"]["product"]} · {data["meta"]["spec"]} · 数据截止 {data["meta"]["asof"]} · {data["meta"]["basis"]}'
    s.font = F_SUB
    s.fill = FILL_NAVY
    s.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
    ws.row_dimensions[2].height = 26

    rows = payload["region_stats"]
    prices = {r["id"]: r["price"] for r in rows if r["price"] is not None}
    hi = max(prices.items(), key=lambda x: x[1])
    lo = min(prices.items(), key=lambda x: x[1])
    ridx = common.region_index(data)
    bench = payload["benchmark"]
    bv = prices[bench]
    pairs = data["freight"]["modes"]
    wmode = next((m for m in pairs if m["id"] == "water"), pairs[0])
    whi = payload["matrix"].get(f"{lo[0]}>{hi[0]}", {}).get("cost", {}).get(wmode["id"])
    cover_ratio = (whi / (hi[1] - lo[1])) if whi and hi[1] > lo[1] else None
    gaps = sum(len(r.get("gaps", [])) for r in data["regions"])

    kpis = [
        ("基准区价", f"{prices[bench]:,}", f"{ridx[bench]['name']} · 元/吨", BLUE),
        ("最高区价", f"{hi[1]:,}", f"{ridx[hi[0]]['name']} · 元/吨", RED),
        ("最低区价", f"{lo[1]:,}", f"{ridx[lo[0]]['name']} · 元/吨", GREEN),
        ("最大价差", f"{hi[1]-lo[1]:,}", f"{ridx[hi[0]]['name']} − {ridx[lo[0]]['name']} · 元/吨", NAVY),
        ("运费/价差", f"{cover_ratio:.0%}" if cover_ratio is not None else "—",
         f"{ridx[lo[0]]['name']}→{ridx[hi[0]]['name']} · {wmode['label']}", AMBER),
        ("数据缺口", str(gaps), "项（已逐条标注）", AMBER),
    ]
    r = 4
    for i, (k, v, s2, color) in enumerate(kpis):
        c0 = ws.cell(row=r, column=i + 1)
        c0.value = k
        c0.font = Font(name="微软雅黑", size=9, color=MUTED)
        c0.fill = FILL_CARD
        c0.border = BOX
        c0.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        v0 = ws.cell(row=r + 1, column=i + 1, value=v)
        v0.font = Font(name="微软雅黑", size=18, bold=True, color=color)
        v0.fill = FILL_CARD
        v0.border = BOX
        v0.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        s0 = ws.cell(row=r + 2, column=i + 1, value=s2)
        s0.font = Font(name="微软雅黑", size=8, color=MUTED)
        s0.fill = FILL_CARD
        s0.border = BOX
        s0.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
    ws.row_dimensions[r].height = 20
    ws.row_dimensions[r + 1].height = 30
    ws.row_dimensions[r + 2].height = 30

    r = 8
    r = sec(ws, r, "一、先看结论（带数字，顺序即优先级）")
    for i, t in enumerate(data["narrative"]["conclusions"], start=1):
        ws.cell(row=r, column=1, value=f"{i:02d}").font = Font(name="微软雅黑", size=10, bold=True, color=AMBER)
        c = ws.cell(row=r, column=2, value=t)
        c.font = F_B2
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
        ws.row_dimensions[r].height = max(20, 15 * (1 + len(t) // 52))
        r += 1

    r += 1
    r = sec(ws, r, "二、判读规则（避免把价差直接读成套利）")
    for t in data["narrative"]["judge_rules"]:
        c = ws.cell(row=r, column=1, value="• " + t)
        c.font = F_B2
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        ws.row_dimensions[r].height = max(18, 15 * (1 + len(t) // 64))
        r += 1

    r += 1
    c = ws.cell(row=r, column=1, value="本页只判断报价环境，不构成套利建议。真实跨区销售还需核对品质、资质、账期、返利与代理协议。")
    c.font = Font(name="微软雅黑", size=9, color=RED)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)


def basis_sheet(ws, data: dict) -> None:
    ws.sheet_view.showGridLines = False
    col_w(ws, {"A": 14, "B": 16, "C": 22, "D": 14, "E": 16, "F": 22, "G": 16})
    r = sec(ws, 1, "口径与方法")
    r = kv_pairs(ws, r, [(m["k"], m["v"]) for m in data["narrative"]["method"]])
    r += 1
    r = sec(ws, r, "样本时间轴（同一口径、逐日可得样本）")
    r = head(ws, r, ["日期", "华东", "华南", "华北", "华中", "西南", "来源"])
    s = data["series"]
    srcmap = {x["date"]: x for x in s["sources"]}
    for i, d in enumerate(s["dates"]):
        vals = [d] + [s["values"][rid][i] for rid in ("east", "south", "north", "central", "southwest")]
        row = r
        r = body(ws, r, vals, fmts={i: NUM for i in range(2, 7)})
        src = srcmap.get(d, {})
        c = ws.cell(row=row, column=7, value=src.get("publisher", ""))
        c.font = F_M
        c.border = BOX
        link = ws.cell(row=row, column=8, value=src.get("url", ""))
        link.font = F_LINK
        link.border = BOX
    ws.column_dimensions["H"].width = 58
    r += 1
    c = ws.cell(row=r, column=1, value="断点说明：" + s.get("gap_note", ""))
    c.font = Font(name="微软雅黑", size=9.5, color=RED)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    ws.row_dimensions[r].height = 34


def series_sheet(ws, data: dict, payload: dict) -> None:
    ws.sheet_view.showGridLines = False
    col_w(ws, {"A": 13, "B": 11, "C": 11, "D": 11, "E": 11, "F": 11, "G": 11, "H": 11, "I": 16})
    r = sec(ws, 1, "区域主流价序列（元/吨，含税）")
    r = head(ws, r, ["日期", "华东", "华南", "华北", "华中", "西南", "最高", "最低", "价差(高-低)"])
    s = data["series"]
    for i, d in enumerate(s["dates"]):
        vals = [s["values"][rid][i] for rid in ("east", "south", "north", "central", "southwest")]
        clean = [v for v in vals if v is not None]
        r = body(ws, r + 0, [d] + vals + [max(clean), min(clean), max(clean) - min(clean)],
                 fmts={i: NUM for i in range(2, 10)})
    first_data, last_data = 3, r - 1
    ws.freeze_panes = "B3"

    ch = LineChart()
    ch.title = "五区重碱主流价"
    ch.style = 2
    ch.height = 8.6
    ch.width = 21
    ch.y_axis.title = "元/吨"
    ch.x_axis.title = "日期"
    dat = Reference(ws, min_col=2, max_col=6, min_row=first_data - 1, max_row=last_data)
    cats = Reference(ws, min_col=1, min_row=first_data, max_row=last_data)
    ch.add_data(dat, titles_from_data=True)
    ch.set_categories(cats)
    for idx, color in enumerate(["1D4ED8", "BE123C", "0E7490", "B45309", "047857"]):
        if idx < len(ch.series):
            ch.series[idx].graphicalProperties.line.solidFill = color
            ch.series[idx].graphicalProperties.line.width = 22000
            ch.series[idx].smooth = False
    ch.y_axis.scaling.min = 1050
    ws.add_chart(ch, f"K2")


def spread_sheet(ws, data: dict, payload: dict) -> None:
    ws.sheet_view.showGridLines = False
    col_w(ws, {"A": 12, "B": 11, "C": 11.5, "D": 9, "E": 10, "F": 10.5, "G": 10, "H": 10.5,
               "I": 10, "J": 10.5, "K": 30, "L": 8})
    bench = payload["benchmark"]
    ridx = common.region_index(data)
    r = sec(ws, 1, f"当前价格与价差分解（基准：{ridx[bench]['name']}；发货地：{ridx[payload['meta']['default_origin']]['name']}）")
    ws.cell(row=r, column=1, value="说明：裸价差 = 该区价 − 基准价；残差 = 裸价差 − 运费。残差是运费解释不了的部分，含供需、库存、品质、账期与口径噪声。").font = F_M
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=12)
    r += 1
    r = head(ws, r, ["区域", "主流价", "裸价差", "运距km", "水运运费", "水运残差",
                     "汽运运费", "汽运残差", "铁运运费", "铁运残差", "判读", "置信度"])
    mx = payload["matrix"]
    prices = {x["id"]: x["price"] for x in payload["region_stats"]}
    origin = payload["meta"]["default_origin"]
    for reg in data["regions"]:
        rid = reg["id"]
        dp = None if rid == bench else (prices[rid] - prices[bench] if prices.get(rid) and prices.get(bench) else None)
        e = mx.get(f"{origin}>{rid}", {})
        wc = e.get("cost", {}).get("water")
        rc = e.get("cost", {}).get("road")
        rl = e.get("cost", {}).get("rail")
        wres = None if (dp is None or wc is None) else round(dp - wc)
        rres = None if (dp is None or rc is None) else round(dp - rc)
        lres = None if (dp is None or rl is None) else round(dp - rl)
        if rid == bench:
            verdict = "基准区"
        elif dp is None:
            verdict = "缺价"
        elif abs(dp) < 30:
            verdict = "价格趋同，运费不可覆盖"
        elif rc is not None and dp < rc and dp < 60:
            verdict = "运费反超，跨区无空间"
        elif wres is not None and abs(wres) <= 50:
            verdict = "运费基本解释"
        elif wres is not None and wres > 50:
            verdict = "残差为正，本地偏紧"
        else:
            verdict = "需逐票询价"
        vals = [reg["name"], prices.get(rid), dp, e.get("km", {}).get("water"),
                wc, wres, rc, rres, rl, lres, verdict, reg.get("price_conf")]
        vals = [("—" if (i in (4, 5, 6, 7, 8, 9, 10) and v is None) else v) for i, v in enumerate(vals, start=1)]
        row = r
        r = body(ws, r, vals, fmts={2: NUM, 3: DELTA, 4: NUM, 5: NUM, 6: DELTA, 7: NUM, 8: DELTA, 9: NUM, 10: DELTA})
        if rid == bench:
            for cc in range(1, 13):
                ws.cell(row=row, column=cc).fill = FILL_LIGHT
        if isinstance(wres, (int, float)) and wres > 60:
            ws.cell(row=row, column=6).font = Font(name="微软雅黑", size=10, bold=True, color=GREEN)
        if isinstance(rres, (int, float)) and rres < -60:
            ws.cell(row=row, column=8).font = Font(name="微软雅黑", size=10, bold=True, color=RED)

    ch = BarChart()
    ch.type = "bar"
    ch.title = "裸价差 vs 运费（元/吨）"
    ch.height = 7.4
    ch.width = 15
    n = len(data["regions"])
    data_ref = Reference(ws, min_col=3, max_col=3, min_row=r - n - 1, max_row=r - 1)
    cats = Reference(ws, min_col=1, min_row=r - n, max_row=r - 1)
    ch.add_data(data_ref, titles_from_data=True)
    ch.set_categories(cats)
    ws.add_chart(ch, "N2")


def matrix_sheet(ws, data: dict, payload: dict) -> None:
    ws.sheet_view.showGridLines = False
    col_w(ws, {"A": 12, "B": 12, "C": 12, "D": 12, "E": 12, "F": 12})
    r = sec(ws, 1, "裸价差矩阵（行 − 列，元/吨；正=行更贵）")
    names = [reg["name"] for reg in data["regions"]]
    r = head(ws, r, ["行 \\ 列"] + names)
    prices = {x["id"]: x["price"] for x in payload["region_stats"]}
    start = r
    for a in data["regions"]:
        row = r
        ws.cell(row=r, column=1, value=a["name"]).font = F_H
        ws.cell(row=r, column=1).fill = FILL_HEAD
        ws.cell(row=r, column=1).border = BOX
        for j, b in enumerate(data["regions"], start=2):
            c = ws.cell(row=r, column=j, value=None if a["id"] == b["id"] else prices[a["id"]] - prices[b["id"]])
            c.number_format = DELTA
            c.border = BOX
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.font = Font(name="微软雅黑", size=10, bold=(a["id"] != b["id"]))
        r += 1
    rng = f"B{start}:F{start + len(names) - 1}"
    ws.conditional_formatting.add(rng, ColorScaleRule(
        start_type="num", start_value=-150, start_color="DBEAFE",
        mid_type="num", mid_value=0, mid_color="FFFFFF",
        end_type="num", end_value=150, end_color="FECDD3"))
    r += 1
    c = ws.cell(row=r, column=1, value="读法：正值说明行区域比列区域贵；颜色越深价差越大。矩阵只描述价格，不含运费，不能直接当成机会。")
    c.font = F_M
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)


def stats_sheet(ws, data: dict, payload: dict) -> None:
    ws.sheet_view.showGridLines = False
    col_w(ws, {"A": 12, "B": 11, "C": 11, "D": 11, "E": 11, "F": 11, "G": 11, "H": 10, "I": 46})
    r = sec(ws, 1, f"价差历史统计（相对 {common.region_index(data)[payload['benchmark']]['name']}，元/吨）")
    r = head(ws, r, ["区域", "当前", "均值", "最小", "最大", "区间", "当前分位", "样本数", "读法"])
    rows = [x for x in payload["region_stats"]]
    for x in rows:
        st = x["stats"]
        bench = payload["benchmark"]
        if x["id"] == bench:
            note = "基准区，恒为 0"
        elif st.get("n", 0) == 0:
            note = "无样本"
        else:
            note = ("价差长期为正，属结构性（销区/运费）" if st["mean"] > 20 else
                    "价差长期为负，基准区偏贵" if st["mean"] < -20 else
                    "价差在 0 附近，报价环境趋同")
            if st["range"] <= 10:
                note += "；区间极窄，价格被运费锁定"
        r = body(ws, r, [x["name"], st.get("cur"), st.get("mean"), st.get("min"), st.get("max"),
                         st.get("range"), (st.get("pctl") / 100 if st.get("pctl") is not None else None),
                         st.get("n"), note],
                 fmts={2: DELTA, 3: DELTA, 4: DELTA, 5: DELTA, 6: NUM, 7: PCT, 8: NUM})
    r += 1
    c = ws.cell(row=r, column=1, value="提示：样本只有 10 个且 9–10 月完全持平，『分位』在此仅表示当前值在样本中的位置，不宜做过强推断。")
    c.font = Font(name="微软雅黑", size=9, color=RED)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)


def freight_sheet(ws, data: dict, payload: dict) -> None:
    ws.sheet_view.showGridLines = False
    col_w(ws, {"A": 12, "B": 12, "C": 11, "D": 11, "E": 11, "F": 11, "G": 11, "H": 11, "I": 26})
    r = sec(ws, 1, "运费矩阵（估算：球面距离 × 绕行系数 × 运价，需询价校准）")
    r = head(ws, r, ["发货地", "目的地", "水运km", "水运成本", "汽运km", "汽运成本", "铁运km", "铁运成本", "备注"])
    for a in data["regions"]:
        for b in data["regions"]:
            if a["id"] == b["id"]:
                continue
            e = payload["matrix"][f"{a['id']}>{b['id']}"]
            kmw, cw = e["km"].get("water"), e["cost"].get("water")
            kmr, cr = e["km"].get("road"), e["cost"].get("road")
            kml, cl = e["km"].get("rail"), e["cost"].get("rail")
            note = ""
            if cw is None:
                note = "内陆枢纽，水运需铁水/江海联运"
            r = body(ws, r, [a["name"], b["name"], kmw, cw, kmr, cr, kml, cl, note],
                     fmts={3: NUM, 4: NUM, 5: NUM, 6: NUM, 7: NUM, 8: NUM})
    ws.freeze_panes = "C3"
    r += 1
    c = ws.cell(row=r, column=1, value="费率：" + "；".join(
        f'{m["label"]} {m["rate"]} 元/吨·km' + (f'+ 固定 {m["base"]} 元/吨' if m.get("base") else "")
        for m in data["freight"]["modes"]))
    c.font = F_M
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
    r += 1
    c = ws.cell(row=r, column=1, value=data["freight"].get("note", ""))
    c.font = F_M
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)


def supply_sheet(ws, data: dict) -> None:
    ws.sheet_view.showGridLines = False
    col_w(ws, {"A": 10, "B": 12, "C": 9, "D": 9, "E": 9, "F": 9, "G": 8, "H": 9, "I": 9,
               "J": 9, "K": 34, "L": 34, "M": 46})
    r = sec(ws, 1, "区域供需、产能与客户格局（产能/需求为万吨/年估算）")
    r = head(ws, r, ["区域", "枢纽", "产能", "产量", "需求", "净平衡", "CR3", "浮法用碱", "光伏用碱", "轻碱用量", "主要生产企业（万吨/年）", "代表客户", "报价环境判断"])
    for reg in data["regions"]:
        s, d = reg["supply"], reg["demand"]
        prods = "；".join(f'{p["n"]} {p["c"]}' for p in s.get("producers", []))
        custs = "；".join(f'{c["n"]}（{c["t"]}）' for c in reg.get("customers", []))
        row = r
        r = body(ws, r, [reg["name"], reg["hub"], s["capacity"], s["output"], d["total"], s["net"],
                         s["cr3"] / 100, d.get("float"), d.get("pv"), d.get("light"), prods, custs, reg["judge"]],
                 fmts={3: NUM, 4: NUM, 5: NUM, 6: DELTA, 7: '0%', 8: NUM, 9: NUM, 10: NUM})
        for cc in (11, 12, 13):
            ws.cell(row=row, column=cc).alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(row=row, column=cc).font = F_B2
        ws.row_dimensions[row].height = 58
        if s["net"] < 0:
            ws.cell(row=row, column=6).font = Font(name="微软雅黑", size=10, bold=True, color=RED)
        else:
            ws.cell(row=row, column=6).font = Font(name="微软雅黑", size=10, bold=True, color=GREEN)


def gaps_sheet(ws, data: dict) -> None:
    ws.sheet_view.showGridLines = False
    col_w(ws, {"A": 12, "B": 26, "C": 8, "D": 56, "E": 40})
    r = sec(ws, 1, "数据缺口清单（先承认缺什么，再谈结论）")
    r = head(ws, r, ["区域", "缺口项", "级别", "说明", "补数建议"])
    for reg in data["regions"]:
        for g in reg.get("gaps", []):
            r = body(ws, r, [reg["name"], g["k"], g["lvl"], g["t"], "电话询价 / 港口周报 / 企业公告"])
    r += 1
    c = ws.cell(row=r, column=1, value="全局缺口")
    c.font = F_SEC
    r += 1
    for t in data["narrative"]["gaps"]:
        c = ws.cell(row=r, column=1, value="• " + t)
        c.font = F_B2
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        ws.row_dimensions[r].height = max(20, 15 * (1 + len(t) // 56))
        r += 1


def method_sheet(ws, data: dict) -> None:
    ws.sheet_view.showGridLines = False
    col_w(ws, {"A": 16, "B": 96})
    r = sec(ws, 1, "方法与免责")
    r = head(ws, r, ["项目", "说明"])
    for m in data["narrative"]["method"]:
        r = body(ws, r, [m["k"], m["v"]])
    for t in data["narrative"]["judge_rules"]:
        r = body(ws, r, ["判读规则", t])
    for g in data["narrative"]["gaps"]:
        r = body(ws, r, ["数据缺口", g])
    for k, v in [("免责声明", data["narrative"]["disclaimer"]),
                 ("样本来源", "；".join(sorted({x["publisher"] for x in data["series"]["sources"]})))]:
        r = body(ws, r, [k, v])
    for row in ws.iter_rows(min_row=3, max_row=r - 1, min_col=2, max_col=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for i in range(3, r):
        ws.row_dimensions[i].height = 46


def build(data_path: Path, out_path: Path) -> Path:
    data = common.load(data_path)
    payload = common.build_payload(data)

    wb = Workbook()
    ws0 = wb.active
    ws0.title = "00 概览"
    cover(ws0, data, payload)

    ws1 = wb.create_sheet("01 口径与时间轴")
    basis_sheet(ws1, data)

    ws2 = wb.create_sheet("02 价格序列")
    series_sheet(ws2, data, payload)

    ws3 = wb.create_sheet("03 价差分解")
    spread_sheet(ws3, data, payload)

    ws4 = wb.create_sheet("04 价差矩阵")
    matrix_sheet(ws4, data, payload)

    ws5 = wb.create_sheet("05 价差统计")
    stats_sheet(ws5, data, payload)

    ws6 = wb.create_sheet("06 运费矩阵")
    freight_sheet(ws6, data, payload)

    ws7 = wb.create_sheet("07 供需与客户")
    supply_sheet(ws7, data)

    ws8 = wb.create_sheet("08 数据缺口")
    gaps_sheet(ws8, data)

    ws9 = wb.create_sheet("09 方法与免责")
    method_sheet(ws9, data)

    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = NAVY if ws.title.startswith("00") else "9FB4C9"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("data")
    ap.add_argument("--out", default="out")
    ap.add_argument("--name", default=None)
    args = ap.parse_args()
    data = common.load(args.data)
    name = args.name or f'{data["meta"]["product"]}多地区价格对比'
    p = build(Path(args.data), Path(args.out) / f"{name}.xlsx")
    print(f"Excel 已生成: {p}  ({p.stat().st_size/1024:.0f} KB)")


if __name__ == "__main__":
    main()
