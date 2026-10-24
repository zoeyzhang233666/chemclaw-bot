#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""check_xlsx.py — Excel 交付自检 + 公式求值（不需要装 Excel）。

检查项：
  1. 表结构：12 张表齐全；04 表关键单元格是公式而不是写死的常数
  2. 与 data.json 同源同数：客户价、市场上下沿、税费/交付口径、历史行数
  3. 公式求值（scripts/xleval.py）：
       默认场景  → 结论应等于 data.json 的结论，偏离金额/比例应一致
       客户价 1100（区间内）      → 接近市场参考区间
       客户价 980（低于下沿且超半宽）→ 明显偏低
       客户交付=出厂且无运费      → 口径不一致，暂时无法直接比较
       补齐运费 80             → 恢复可比较

用法：
  python scripts/check_xlsx.py out/xxx.xlsx build/data.json
"""
from __future__ import annotations

import argparse
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import common as C  # noqa: E402
import xleval  # noqa: E402

from openpyxl import load_workbook  # noqa: E402

REQUIRED = [
    "00_封面与结论", "01_客户报价输入", "02_口径对照", "03_市场报价与历史",
    "04_偏离计算", "05_历史走势数据", "06_区域对比", "07_待确认问题",
    "08_数据缺口", "09_计算过程", "10_沟通事实要点", "11_数据来源清单",
]


def find_label_row(ws, label, col=1):
    for row in ws.iter_rows(max_col=col):
        if isinstance(row[0].value, str) and row[0].value.strip() == label:
            return row[0].row
    return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("data_json", nargs="?")
    a = ap.parse_args()

    wb = load_workbook(a.xlsx)
    data = C.load_json(a.data_json) if a.data_json else None
    fails, warns, notes = [], [], []

    # ---- 1 结构
    for name in REQUIRED:
        if name not in wb.sheetnames:
            fails.append(f"缺表：{name}")
    if len(wb.sheetnames) != 12:
        warns.append(f"表数为 {len(wb.sheetnames)}（应为 12）")

    ws4 = wb["04_偏离计算"] if "04_偏离计算" in wb.sheetnames else None
    ws1 = wb["01_客户报价输入"] if "01_客户报价输入" in wb.sheetnames else None
    if ws4 is None or ws1 is None:
        print("   [FAIL] 关键表缺失，无法继续")
        return 1

    n_formula = sum(1 for row in ws4.iter_rows() for c in row
                    if isinstance(c.value, str) and c.value.startswith("="))
    if n_formula < 15:
        fails.append(f"04 表公式数量过少（{n_formula}），可能存在写死的常数")

    # ---- 2 同源
    if data:
        if ws1["B5"].value != data["customer"].get("stated_price"):
            fails.append(f"01!B5 客户价 {ws1['B5'].value} ≠ {data['customer'].get('stated_price')}")
        m = data["market"]
        if ws1["B23"].value != m.get("low"):
            fails.append(f"01!B23 下沿 {ws1['B23'].value} ≠ {m.get('low')}")
        if ws1["B24"].value != m.get("high"):
            fails.append(f"01!B24 上沿 {ws1['B24'].value} ≠ {m.get('high')}")
        if ws1["B21"].value != C.norm_tax(m.get("tax")):
            fails.append(f"01!B21 市场税费 {ws1['B21'].value} ≠ {C.norm_tax(m.get('tax'))}")
        if ws1["B22"].value != C.norm_delivery(m.get("delivery")):
            fails.append(f"01!B22 市场交付 {ws1['B22'].value} ≠ {C.norm_delivery(m.get('delivery'))}")
        if not (isinstance(ws1["B25"].value, str) and ws1["B25"].value.startswith("=")):
            fails.append("01!B25 市场区间中值应为公式")
        s5 = wb["05_历史走势数据"] if "05_历史走势数据" in wb.sheetnames else None
        if s5:
            n_hist = sum(1 for row in s5.iter_rows(min_row=5, max_col=1)
                         if isinstance(row[0].value, str) and re.match(r"\d{4}-\d{2}-\d{2}", row[0].value))
            if data.get("history") and n_hist != len(data["history"]):
                warns.append(f"05 表历史行数 {n_hist} ≠ data.json {len(data['history'])}")

    # ---- 3 求值
    labels = {
        "客户价": "客户所述价格",
        "偏离额": "① 偏离金额（相对中值）",
        "偏离率": "② 偏离比例（相对中值）",
        "结论": "核验结论",
    }
    rows = {k: find_label_row(ws4, v) for k, v in labels.items()}
    if any(v is None for v in rows.values()):
        fails.append(f"04 表缺少行标签：{[labels[k] for k, v in rows.items() if v is None]}")
    else:
        book = xleval.Book(wb)

        def v(key):
            return book.value("04_偏离计算", f"B{rows[key]}")

        # 默认场景
        if data:
            cust = data["customer"].get("stated_price")
            got = {k: v(k) for k in rows}
            if cust is not None and abs((got["客户价"] or 0) - cust) > 1e-6:
                fails.append(f"求值·客户价 {got['客户价']} ≠ {cust}")
            cmp_ = data["comparison"]
            if cmp_.get("pos") != "undecidable":
                if not isinstance(got["偏离额"], (int, float)) or abs(got["偏离额"] - cmp_["delta_to_dot"]) > 0.51:
                    fails.append(f"求值·偏离金额 {got['偏离额']} ≠ {cmp_['delta_to_dot']}")
                if not isinstance(got["偏离率"], (int, float)) or abs(got["偏离率"] - cmp_["pct_to_dot"]) > 1e-3:
                    fails.append(f"求值·偏离比例 {got['偏离率']} ≠ {cmp_['pct_to_dot']}")
                if got["结论"] != data["verdict"]["label"]:
                    fails.append(f"求值·结论 「{got['结论']}」 ≠ 「{data['verdict']['label']}」")
                else:
                    notes.append(f"默认场景：客户价 {creal(cust)} → 偏离 {got['偏离额']:+,.0f} 元/吨"
                                 f"（{got['偏离率']*100:+.1f}%）→ {got['结论']}")
            else:
                notes.append(f"默认场景：{got['结论']}")

        # 场景 A：客户价落在区间内（仅在有市场区间时有意义）
        lo = book.value("01_客户报价输入", "B23")
        hi = book.value("01_客户报价输入", "B24")
        has_range = isinstance(lo, (int, float)) and isinstance(hi, (int, float))
        if not has_range:
            # 降级案例：没有市场区间，判定必须明确说明「无法判定」，而不是给出高低结论
            r = v("结论")
            if r not in ("无可靠市场数据，无法验证", "无法判定（缺少可比市场价格）",
                         "口径不一致，暂时无法直接比较"):
                fails.append(f"降级案例（无市场区间）结论应为「无法验证」类，实际「{r}」")
            else:
                notes.append(f"降级场景：无市场区间 → {r}")
            print(f"== check_xlsx: {os.path.basename(a.xlsx)} "
                  f"({os.path.getsize(a.xlsx)/1024:.0f} KB, {len(wb.sheetnames)} 张表)")
            for n in notes:
                print(f"   [ok] {n}")
            for w in warns:
                print(f"   [warn] {w}")
            for f_ in fails:
                print(f"   [FAIL] {f_}")
            if not fails:
                print("   PASS  12 表齐全 · 关键单元格为公式 · 降级场景结论正确")
                return 0
            print(f"   共 {len(fails)} 项失败")
            return 1

        mid = (lo + hi) / 2.0
        book.set("01_客户报价输入", "B5", mid)
        r = v("结论")
        if r != "接近市场参考区间":
            fails.append(f"场景A（客户价=区间中值 {mid:g}）应判「接近市场参考区间」，实际「{r}」")
        else:
            notes.append(f"场景A：客户价 {mid:g}（区间内）→ 接近市场参考区间")

        # 场景 B：明显偏低
        book.set("01_客户报价输入", "B5", lo - (hi - lo))
        r = v("结论")
        if r != "明显偏低":
            fails.append(f"场景B（客户价 {lo-(hi-lo):g}）应判「明显偏低」，实际「{r}」")
        else:
            notes.append(f"场景B：客户价 {lo-(hi-lo):g}（低于下沿超半宽）→ 明显偏低")

        # 场景 C：客户口径出厂且无运费 → 无法直接折算
        book = xleval.Book(wb)
        book.set("01_客户报价输入", "B5", 1400)
        book.set("01_客户报价输入", "B9", "出厂")
        book.set("01_客户报价输入", "B15", None)
        r = v("结论")
        if r != "口径不一致，暂时无法直接比较":
            fails.append(f"场景C（出厂·无运费）应判「口径不一致，暂时无法直接比较」，实际「{r}」")
        else:
            notes.append("场景C：客户口径=含税出厂、未给运费 → 口径不一致，暂时无法直接比较")

        # 场景 D：补齐运费后可比较
        book = xleval.Book(wb)
        book.set("01_客户报价输入", "B5", 1400)
        book.set("01_客户报价输入", "B9", "出厂")
        book.set("01_客户报价输入", "B15", 80)
        r = v("结论")
        if r in (None, "", "口径不一致，暂时无法直接比较") or not isinstance(r, str):
            fails.append(f"场景D（补齐运费 80）应恢复可比较，实际「{r}」")
        else:
            notes.append(f"场景D：客户口径=出厂并补齐运费 80 元/吨 → {r}")

    print(f"== check_xlsx: {os.path.basename(a.xlsx)} ({os.path.getsize(a.xlsx)/1024:.0f} KB, {len(wb.sheetnames)} 张表)")
    for n in notes:
        print(f"   [ok] {n}")
    for w in warns:
        print(f"   [warn] {w}")
    for f_ in fails:
        print(f"   [FAIL] {f_}")
    if not fails:
        print("   PASS  12 表齐全 · 关键单元格为公式 · 与 HTML 同源 · 4 组口径场景求值正确")
        return 0
    print(f"   共 {len(fails)} 项失败")
    return 1


def creal(v):
    return f"{v:,.0f}" if isinstance(v, (int, float)) else str(v)


if __name__ == "__main__":
    sys.exit(main())
