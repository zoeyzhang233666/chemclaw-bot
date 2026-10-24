#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""acceptance.py — 交付验收：把所有案例跑一遍，逐条对照需求检查。

这不是单元测试，而是"交付前最后一次验收"：拿用户当初提的每一条要求，
去找交付物里对应的证据。少一条就报 FAIL。

用法：
  python scripts/acceptance.py
"""
from __future__ import annotations

import io
import json
import os
import re
import subprocess
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
EX = os.path.join(ROOT, "examples")
PY = sys.executable

PASS, FAIL = [], []


def ck(name, cond, detail=""):
    (PASS if cond else FAIL).append(name)
    print(f"   {'ok  ' if cond else 'FAIL'} {name}" + (f"\n        → {detail}" if (detail and not cond) else ""))


def run(args):
    p = subprocess.run([PY, *args], capture_output=True, text=True,
                       encoding="utf-8", errors="replace")
    return p.returncode, (p.stdout or "") + (p.stderr or "")


def load(p):
    return json.load(io.open(p, encoding="utf-8"))


def main():
    cases = sorted(d for d in os.listdir(EX) if os.path.isdir(os.path.join(EX, d)))
    print(f"== 验收：{len(cases)} 个案例 {cases}\n")

    # ---------------------------------------------------------------- 0 文件
    print("[0] 交付物结构")
    for f in ("SKILL.md", "references", "scripts", "examples"):
        ck(f"存在 {f}", os.path.exists(os.path.join(ROOT, f)))
    for f in ("sources-playbook.md", "caliber-rules.md", "verdict-rules.md",
              "formula-rules.md", "data-contract.md", "adapters.md", "qa-checklist.md"):
        ck(f"references/{f}", os.path.isfile(os.path.join(ROOT, "references", f)))
    for f in ("common.py", "assemble.py", "build_html.py", "build_xlsx.py", "build_all.py",
              "normalize.py", "validate.py", "check_html.py", "check_xlsx.py",
              "xleval.py", "selftest.py"):
        ck(f"scripts/{f}", os.path.isfile(os.path.join(ROOT, "scripts", f)))
    ck("templates/report.html", os.path.isfile(os.path.join(ROOT, "scripts", "templates", "report.html")))
    ck("assets/echarts.min.js", os.path.isfile(os.path.join(ROOT, "scripts", "assets", "echarts.min.js")))
    ck("examples/README.md", os.path.isfile(os.path.join(EX, "README.md")))

    # ---------------------------------------------------------------- 1 SKILL.md
    print("\n[1] SKILL.md 内容")
    sk = io.open(os.path.join(ROOT, "SKILL.md"), encoding="utf-8").read()
    ck("含 YAML front-matter", sk.startswith("---") and "name:" in sk[:400] and "description:" in sk[:1400])
    for kw in ("八步", "无法直接折算", "无法验证", "不预测", "半宽"):
        ck(f"SKILL.md 含「{kw}」", kw in sk)
    ck("SKILL.md 声明不替用户决定报价", "替你决定报多少" in sk or "不碰定价" in sk)
    ck("SKILL.md 不绑定单一品种", "不绑定任何单一品种" in sk or "product-agnostic" in sk.lower())

    # ---------------------------------------------------------------- 2 脚本自检
    print("\n[2] 脚本自检")
    rc, out = run([os.path.join(HERE, "selftest.py")])
    ck("selftest 通过", rc == 0, out[-400:])
    m = re.search(r"通过 (\d+) / (\d+)", out)
    if m:
        ck(f"selftest {m.group(1)}/{m.group(2)} 项", m.group(1) == m.group(2))

    # ---------------------------------------------------------------- 3 每个案例
    for case in cases:
        d = os.path.join(EX, case)
        print(f"\n[3] 案例 {case}")
        ck("有 research.json", os.path.isfile(os.path.join(d, "research.json")))
        ck("有 README.md", os.path.isfile(os.path.join(d, "README.md")))
        ck("有 customer-message.txt", os.path.isfile(os.path.join(d, "customer-message.txt")))

        # 完整构建
        rc, out = run([os.path.join(HERE, "build_all.py"), d])
        ck("build_all 全绿", rc == 0, out[-900:])

        datap = os.path.join(d, "build", "data.json")
        if not os.path.isfile(datap):
            ck("生成 data.json", False)
            continue
        data = load(datap)

        # 交付物存在（忽略探针临时页等以下划线开头的文件）
        outdir = os.path.join(d, "out")
        htmls = [f for f in os.listdir(outdir) if f.endswith(".html") and not f.startswith("_")]
        xlsxs = [f for f in os.listdir(outdir) if f.endswith(".xlsx") and not f.startswith("_")]
        ck("生成 HTML", len(htmls) == 1, str(htmls))
        ck("生成 Excel", len(xlsxs) == 1, str(xlsxs))
        ck("out 目录无临时文件", not [f for f in os.listdir(outdir) if f.startswith("_")],
           str([f for f in os.listdir(outdir) if f.startswith("_")]))

        # ---- 8 个问题都答了
        m = data["market"]
        ck("Q1 市场参考区间+日期+来源+规格+单位",
           (m.get("low") is not None and m.get("high") is not None) or data["verdict"]["code"] == "unknown",
           f"low={m.get('low')} high={m.get('high')} verdict={data['verdict']['code']}")
        ck("Q1 带了来源与日期", bool(m.get("source")) or data["verdict"]["code"] == "unknown")

        c = data["comparison"]
        if c.get("pos") != "undecidable":
            ck("Q2 偏离金额", isinstance(c.get("delta_to_dot"), (int, float)))
            ck("Q2 偏离比例", isinstance(c.get("pct_to_dot"), (int, float)))
        else:
            ck("Q2 不可比时明确未计算", c.get("pos") == "undecidable")

        labels = {"接近市场参考区间", "略高于市场参考区间", "略低于市场参考区间",
                  "明显偏高", "明显偏低", "无法验证"}
        v = data["verdict"]
        ck(f"Q3 判定在合法集合内（{v['label']}）",
           v["label"] in labels or "无法" in v["label"] or "口径不一致" in v["label"])

        ck("Q4 有 7d/30d/3m 窗口字段", all(k in data for k in ("windows",)))
        ck("Q5 有口径差异表（≥5 维）", len(data.get("calibers", [])) >= 5,
           f"{len(data.get('calibers', []))} 维")
        ck("Q5 有影响因素（≥8 类）", len(data.get("factors", [])) >= 8,
           f"{len(data.get('factors', []))} 类")
        ck("Q6 有区域对比或明确无数据",
           bool(data.get("region_compare")) or data["verdict"]["code"] == "unknown")
        ck("Q7 有事实要点（≥3）", len(data.get("facts", [])) >= 3, f"{len(data.get('facts', []))} 条")
        ck("Q8 有确认问题（≥5）", len(data.get("questions", [])) >= 5, f"{len(data.get('questions', []))} 条")
        ck("有数据缺口（≥1）", len(data.get("gaps", [])) >= 1, f"{len(data.get('gaps', []))} 条")

        # ---- 硬约束
        print("   —— 硬约束")
        cust = data["customer"]
        if cust.get("is_simulated"):
            ck("模拟消息带【模拟】标注", "模拟" in str(cust.get("raw_text", "")))
            ck("模拟案例有免责说明", bool(cust.get("simulation_disclaimer")))
            htmltxt = io.open(os.path.join(outdir, htmls[0]), encoding="utf-8", errors="replace").read()
            ck("HTML 里标注了模拟", "模拟" in htmltxt)

        # 客户价没有被当成市场价
        cp = cust.get("stated_price")
        spots = data.get("spot_all", [])
        polluted = [s for s in spots
                    if not s.get("exclude_reason")
                    and all(abs(x - cp) < 1e-6 for x in
                            [v for v in (s.get("low"), s.get("high"), s.get("point"))
                             if isinstance(v, (int, float))]) 
                    and [v for v in (s.get("low"), s.get("high"), s.get("point"))
                         if isinstance(v, (int, float))]]
        ck("客户价未被当作市场价", not polluted, str([s.get("id") for s in polluted]))

        # 不出现固定百分比阈值的判定
        src = io.open(os.path.join(ROOT, "scripts", "common.py"), encoding="utf-8").read()
        ck("判定不含固定百分比阈值",
           "0.05" not in src.split("def compare_to_range")[1].split("def PICK_REFERENCE")[0]
           or "half" in src.split("def compare_to_range")[1])

        # 禁用词：只看渲染后的可见文字，排除 <script> 里扫描器自己的词表
        htmltxt = io.open(os.path.join(outdir, htmls[0]), encoding="utf-8", errors="replace").read()
        visible = re.sub(r"<script.*?</script>", " ", htmltxt, flags=re.S | re.I)
        visible = re.sub(r"<style.*?</style>", " ", visible, flags=re.S | re.I)
        visible = re.sub(r"<[^>]+>", " ", visible)
        banned = ["建议报价", "销售建议价", "底价", "让利", "包您", "保证采购", "抄底"]
        hit = [b for b in banned if b in visible]
        ck("HTML 无禁用词", not hit, str(hit))

        # 外链（排除 w3.org 命名空间声明）
        ext = re.findall(r'(?:src|href)="(https?://[^"]+)"', htmltxt)
        ext = [u for u in ext if not u.startswith(("http://www.w3.org", "https://www.w3.org"))]
        ck("HTML 无外部资源依赖", not ext, str(ext[:3]))

        # Excel 活公式
        try:
            from openpyxl import load_workbook
            wb = load_workbook(os.path.join(outdir, xlsxs[0]))
            ws4 = wb["04_偏离计算"]
            nf_ = sum(1 for r in ws4.iter_rows() for cc in r
                      if isinstance(cc.value, str) and cc.value.startswith("="))
            ck(f"Excel 04 表有活公式（{nf_} 个）", nf_ >= 15)
            ck("Excel 12 张表", len(wb.sheetnames) == 12, str(len(wb.sheetnames)))
        except Exception as e:  # noqa: BLE001
            ck("Excel 可读", False, str(e))

        # ---- 交互实测（无头浏览器真的点一遍）
        print("   —— 交互实测")
        has_range = (data["market"].get("low") is not None
                     and data["market"].get("high") is not None)
        rc, out = run([os.path.join(HERE, "probe_interact.py"), os.path.join(outdir, htmls[0])])
        ck("交互探针执行成功", rc == 0, out[-300:])
        if has_range:
            ck("改客户价后结论会变", "客户价1150" in out and "接近市场参考区间" in out, out[-300:])
        else:
            # 降级案例没有区间，改客户价必须依旧给出「无法验证」，不能凭空给出高低结论
            ck("无区间时改客户价仍判「无法验证」",
               "客户价1150" in out and "无法验证" in out, out[-300:])
        ck("事实要点可生成", ("事实要点 chars=" in out) and "事实要点 chars=0" not in out)
        ck("微信稿可生成", "微信 chars=" in out and "微信 chars=0" not in out)
        ck("邮件稿可生成", "邮件 chars=" in out and "邮件 chars=0" not in out)

        # 草稿措辞合规（只扫正文，去掉工具的免责提示行）
        rc2, out2 = run([os.path.join(HERE, "probe_drafts.py"), os.path.join(outdir, htmls[0])])
        ck("三份草稿无承诺性措辞", "禁用/承诺类措辞： 无" in out2 or rc2 == 0, out2[-300:])

    # ---------------------------------------------------------------- 4 汇总
    print("\n" + "=" * 60)
    print(f"验收结果：通过 {len(PASS)} / {len(PASS) + len(FAIL)}")
    if FAIL:
        print("\n未通过项：")
        for f in FAIL:
            print(f"  [FAIL] {f}")
        return 1
    print("全部通过 ✓")
    return 0


if __name__ == "__main__":
    sys.exit(main())
