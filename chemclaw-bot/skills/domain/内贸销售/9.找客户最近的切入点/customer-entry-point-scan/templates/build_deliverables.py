# -*- coding: utf-8 -*-
"""
customer-entry-point-scan · 附件生成器
用法:
    python build_deliverables.py cards.json [-o 输出目录] [-n 文件名前缀]
输入: 按 templates/scan_data.schema.json 组织的 cards.json (UTF-8)
输出: <前缀>.xlsx + <前缀>.html
依赖: Python 3.8+ 和 openpyxl (pip install openpyxl)
"""
import argparse, json, re, sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("缺少依赖 openpyxl，请先执行: pip install openpyxl")

TOKEN = "__SCAN_DATA_JSON__"
TEMPLATE = Path(__file__).with_name("report.template.html")

BRAND = "123A5F"
C_HEAD = PatternFill("solid", fgColor=BRAND)
C_ZEBRA = PatternFill("solid", fgColor="F4F8FC")
C_META_F = Font(color="6B7A93", size=10)
C_TITLE_F = Font(color=BRAND, size=15, bold=True)
F_HEAD = Font(color="FFFFFF", size=10.5, bold=True)
F_BODY = Font(color="17233A", size=10)
F_LINK = Font(color="1156C8", size=10, underline="single")
F_BOLD = Font(color="17233A", size=10, bold=True)
F_GOLD = Font(color="B7791F", size=10, bold=True)
THIN = Side(style="thin", color="C9D4E0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
AL_TOP = Alignment(vertical="top", wrap_text=True)
AL_CTR = Alignment(vertical="center", horizontal="center", wrap_text=True)
PRIO_FILL = {
    1: (PatternFill("solid", fgColor="FDECEA"), Font(color="B42318", size=10, bold=True)),
    2: (PatternFill("solid", fgColor="FEF3E2"), Font(color="C2620A", size=10, bold=True)),
    3: (PatternFill("solid", fgColor="FAF5E0"), Font(color="8A6D1A", size=10, bold=True)),
    4: (PatternFill("solid", fgColor="EAF0F7"), Font(color="4F6B8F", size=10, bold=True)),
    5: (PatternFill("solid", fgColor="F0F2F5"), Font(color="7D8CA3", size=10, bold=True)),
    6: (PatternFill("solid", fgColor="F0F2F5"), Font(color="5B7FA6", size=10, bold=True)),
    7: (PatternFill("solid", fgColor="F0F2F5"), Font(color="8B6BA6", size=10, bold=True)),
}
TAG_RE = re.compile(r"<[^>]+>")

def strip_tags(s): return TAG_RE.sub("", s or "")

def die(msg):
    print("[错误] " + msg); sys.exit(1)

def warn(msg): print("[警告] " + msg)

def validate(d):
    for k in ("meta", "snapshot", "cards", "evidence", "market", "inferences", "roles"):
        if k not in d: die(f"cards.json 缺少顶层字段: {k}")
    meta = d["meta"]
    for k in ("title", "task_line", "customer", "window", "generated", "kpis", "file_base"):
        if k not in meta: die(f"meta 缺少字段: {k}")
    if not (1 <= len(d["cards"]) <= 7): die("cards 数量须为1–7（建议3–5）")
    if len(d["cards"]) < 3: warn("卡片少于3张，报告单薄")
    ev_ids = {e.get("id") for e in d["evidence"]}
    for c in d["cards"]:
        for k in ("no", "prio", "title", "event", "relation", "whyNow", "roles", "evidence", "inference", "next_action", "visit", "message"):
            if k not in c: die(f"卡片 {c.get('no','?')} 缺少字段: {k}")
        for eid in c["evidence"]:
            if eid not in ev_ids: die(f"卡片 {c['no']} 引用了不存在的证据 {eid}")
        if not c["inference"]: die(f"卡片 {c['no']} 缺少推断标注（inference）")
    for i, e in enumerate(d["evidence"]):
        for k in ("id", "date", "source", "kind", "tag", "title", "cards", "url"):
            if k not in e: die(f"第{i+1}条证据缺少字段: {k}")
        if not str(e["url"]).startswith("http"): die(f"证据 {e['id']} 的 URL 不合法")
    print(f"[校验] 通过：{len(d['cards'])} 张卡 / {len(d['evidence'])} 条证据 / {len(d['inferences'])} 条推断")

# ---------------- Excel ----------------
def sheet_title(ws, title, subtitle, ncols):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws.cell(row=1, column=1, value=title).font = C_TITLE_F
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    ws.cell(row=2, column=1, value=subtitle).font = C_META_F
    ws.row_dimensions[1].height = 26

def header_row(ws, row, headers):
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = F_HEAD; c.fill = C_HEAD; c.alignment = AL_CTR; c.border = BORDER
    ws.row_dimensions[row].height = 22

def body_cell(ws, row, col, value, bold=False, zebra=False, ctr=False):
    c = ws.cell(row=row, column=col, value=value)
    c.font = F_BOLD if bold else F_BODY
    c.alignment = AL_CTR if ctr else AL_TOP
    c.border = BORDER
    if zebra: c.fill = C_ZEBRA
    return c

def widths(ws, wlist):
    for i, w in enumerate(wlist, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

def page_setup(ws):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

def build_xlsx(d, path):
    meta, cards, ev = d["meta"], d["cards"], d["evidence"]
    cust = meta.get("customer", {})
    win = meta.get("window", {})
    wb = Workbook(); wb.properties.title = meta["title"]; wb.properties.creator = "customer-entry-point-scan"

    # Sheet1 使用说明
    ws = wb.active; ws.title = "使用说明"; ws.sheet_view.showGridLines = False
    widths(ws, [16, 30, 60, 44])
    sheet_title(ws, meta["title"], f"{meta['task_line']} ｜ 生成日 {meta['generated']}", 4)
    rows = [
        ("扫描对象", cust.get("name", ""), (cust.get("codes", "") + " " + cust.get("profile", "")).strip(), ""),
        ("我方产品", meta.get("product", ""), "", ""),
        ("扫描窗口", f"{win.get('start','')} ～ {win.get('end','')}", win.get("label", ""), ""),
        ("扫描范围", meta.get("scope", ""), "", ""),
        ("信息源", meta.get("sources", ""), "全部证据见【证据清单】Sheet", ""),
    ]
    r = 4; header_row(ws, r, ["项目", "内容", "说明", "备注"]); r += 1
    for i, m in enumerate(rows):
        for j, v in enumerate(m, 1):
            body_cell(ws, r, j, v, bold=(j == 1), zebra=(i % 2 == 1))
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="工作表导航").font = Font(color=BRAND, size=12, bold=True); r += 1
    nav = [
        ("事件卡总览", f"{len(cards)}个切入点的完整拆解", "每行=一张事件卡：事件→关系→为什么现在→找谁→证据→推断→下一步", "打印或导入CRM用这一张"),
        ("证据清单", f"{len(ev)}条证据及原文链接", "每条注明支持哪张卡，G列可点开原文", ""),
        ("行业背景数据", f"{len(d['market'].get('rows', []))}组关键数据", "支撑\"为什么是现在\"的市场判断", ""),
        ("推断边界", f"{len(d['inferences'])}条推断及验证方式", "红线：以下内容不得当作事实对客户表述", ""),
        ("接触角色与行动", f"{len(d['roles'])}类角色的目标与开场话题", "对应卡片编号，直接用于约访", ""),
    ]
    header_row(ws, r, ["工作表", "内容", "说明", "使用建议"]); r += 1
    for i, m in enumerate(nav):
        for j, v in enumerate(m, 1):
            body_cell(ws, r, j, v, bold=(j == 1), zebra=(i % 2 == 1))
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="优先级图例").font = Font(color=BRAND, size=12, bold=True); r += 1
    legend_note = {1: "最高：客户事件+时机完全咬合，立即行动", 2: "高：窗口明确，1个月内行动",
                   3: "中高：配合前排卡片推进", 4: "中：时点性强，按节点行动", 5: "辅助：内部使用或次级谈资"}
    for c_ in cards:
        fill, font = PRIO_FILL.get(c_["prio"], PRIO_FILL[5])
        cc = ws.cell(row=r, column=1, value=c_["no"]); cc.fill = fill; cc.font = font; cc.border = BORDER; cc.alignment = AL_CTR
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        c2 = ws.cell(row=r, column=2, value=f"{c_['stars']} {c_['title'][:44]} —— {legend_note.get(c_['prio'], '')}")
        c2.font = F_BODY; c2.alignment = AL_TOP; c2.border = BORDER
        r += 1
    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r + 1, end_column=4)
    c = ws.cell(row=r, column=1, value="免责声明：" + (d.get("footer", {}).get("disclaimer", "本报告由公开信息整理，仅供销售开拓参考；引用内容版权归原出处所有。")))
    c.font = C_META_F; c.alignment = AL_TOP
    page_setup(ws)

    # Sheet2 事件卡总览
    ws = wb.create_sheet("事件卡总览")
    widths(ws, [7, 7, 10, 34, 15, 52, 40, 40, 34, 12, 36, 30])
    sheet_title(ws, "事件卡总览（按建议联系顺序）",
                "结构：事件 → 与产品的关系 → 为什么现在值得联系 → 推荐接触角色 → 证据 → 推断标注 → 建议下一步", 12)
    header_row(ws, 4, ["优先级", "星级", "类型", "切入点主题", "关键日期", "发生了什么（事件）",
                       "与产品的关系", "为什么现在值得联系", "推荐接触角色", "证据", "推断标注（勿当事实）", "建议下一步动作"])
    r = 5
    for i, c_ in enumerate(cards):
        v = [c_["no"], c_["stars"], c_["type"], strip_tags(c_["title"]), c_["dates"],
             strip_tags(c_["event"]), strip_tags(c_["relation"]), strip_tags(c_["whyNow"]),
             "\n".join(f"{x['r']}（{x['d']}）" for x in c_["roles"]),
             " ".join(c_["evidence"]), "\n".join(c_["inference"]), c_["next_action"]]
        fill, font = PRIO_FILL.get(c_["prio"], PRIO_FILL[5])
        for j, val in enumerate(v, 1):
            cc = body_cell(ws, r, j, val, zebra=(i % 2 == 1), ctr=(j in (1, 2, 3, 5, 10)))
            if j == 1: cc.fill = fill; cc.font = font
            if j == 2: cc.font = F_GOLD
        r += 1
    ws.freeze_panes = "E5"; ws.auto_filter.ref = f"A4:L{r-1}"; page_setup(ws)

    # Sheet3 证据清单
    ws = wb.create_sheet("证据清单")
    widths(ws, [7, 12, 24, 12, 72, 12, 60])
    sheet_title(ws, f"证据清单（{len(ev)}条 · 点开原文）", "src1=一手公告（绿） src2=二手来源（蓝）；每条注明支持哪张卡", 7)
    header_row(ws, 4, ["编号", "日期", "来源", "性质", "内容要点", "支持卡片", "原文链接"])
    r = 5
    for i, e in enumerate(ev):
        for j, val in enumerate([e["id"], e["date"], e["source"], e["kind"], e["title"], e["cards"], "打开原文 ↗"], 1):
            cc = body_cell(ws, r, j, val, bold=(j == 1), zebra=(i % 2 == 1), ctr=(j in (1, 2, 4, 6)))
            if j == 7:
                cc.hyperlink = e["url"]; cc.font = F_LINK
        r += 1
    ws.freeze_panes = "A5"; ws.auto_filter.ref = f"A4:G{r-1}"; page_setup(ws)

    # Sheet4 行业背景数据
    ws = wb.create_sheet("行业背景数据")
    widths(ws, [22, 46, 26, 12, 40])
    sheet_title(ws, "行业与市场关键数据", "支撑\"为什么是现在\"的市场判断", 5)
    header_row(ws, 4, ["指标", "数值 / 变化", "时点", "证据", "备注"])
    r = 5
    for i, m in enumerate(d["market"].get("rows", [])):
        for j, val in enumerate(m, 1):
            body_cell(ws, r, j, val, bold=(j == 1), zebra=(i % 2 == 1), ctr=(j == 4))
        r += 1
    ws.freeze_panes = "A5"; page_setup(ws)

    # Sheet5 推断边界
    ws = wb.create_sheet("推断边界")
    widths(ws, [8, 46, 30, 40])
    sheet_title(ws, "推断边界（使用红线）", "以下内容为推断/测算/惯例，不得当作已证实事实对客户或内部表述", 4)
    header_row(ws, 4, ["编号", "推断内容", "依据类型", "如何验证 / 使用限制"])
    r = 5
    for i, m in enumerate(d["inferences"]):
        for j, val in enumerate(m, 1):
            body_cell(ws, r, j, val, bold=(j == 1), zebra=(i % 2 == 1), ctr=(j == 1))
        r += 1
    ws.freeze_panes = "A5"; page_setup(ws)

    # Sheet6 接触角色与行动
    ws = wb.create_sheet("接触角色与行动")
    widths(ws, [8, 40, 14, 12, 34, 40])
    sheet_title(ws, "接触角色地图", "对应事件卡与开场话题，可直接用于约访与CRM录入", 6)
    header_row(ws, 4, ["编号", "角色 / 部门", "层级定位", "对应卡片", "沟通目标", "建议开场 / 话题"])
    r = 5
    for i, m in enumerate(d["roles"]):
        for j, val in enumerate(m, 1):
            body_cell(ws, r, j, val, bold=(j == 1), zebra=(i % 2 == 1), ctr=(j in (1, 3, 4)))
        r += 1
    ws.freeze_panes = "A5"; page_setup(ws)

    wb.save(path)

# ---------------- HTML ----------------
def build_html(d, path):
    if not TEMPLATE.exists(): die(f"找不到模板文件: {TEMPLATE}")
    tpl = TEMPLATE.read_text(encoding="utf-8")
    payload = json.dumps(d, ensure_ascii=False).replace("</", "<\\/")
    if TOKEN not in tpl: die("模板缺少占位符 __SCAN_DATA_JSON__")
    (path).write_text(tpl.replace(TOKEN, payload), encoding="utf-8")

# ---------------- main ----------------
def main():
    ap = argparse.ArgumentParser(description="客户切入点扫描 · 附件生成器")
    ap.add_argument("cards", help="cards.json 路径")
    ap.add_argument("-o", "--outdir", default=None, help="输出目录（默认与 cards.json 同目录）")
    ap.add_argument("-n", "--name", default=None, help="文件名前缀（默认 meta.file_base）")
    a = ap.parse_args()
    src = Path(a.cards)
    if not src.exists(): die(f"找不到输入文件: {src}")
    try:
        d = json.loads(src.read_text(encoding="utf-8"))
    except json.JSONDecodeError as e:
        die(f"cards.json 不是合法 JSON: {e}")
    validate(d)
    base = a.name or d["meta"].get("file_base") or "客户切入点扫描"
    outdir = Path(a.outdir) if a.outdir else src.parent
    outdir.mkdir(parents=True, exist_ok=True)
    xlsx, html = outdir / f"{base}.xlsx", outdir / f"{base}.html"
    build_xlsx(d, xlsx); build_html(d, html)
    print(f"[完成] {xlsx}")
    print(f"[完成] {html}")

if __name__ == "__main__":
    main()
