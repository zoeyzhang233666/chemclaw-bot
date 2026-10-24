# -*- coding: utf-8 -*-
"""交付前校验：HTML 结构、数据完整性、Excel 工作表与关键内容。

用法:
    python verify.py case.json --html out.html --xlsx out.xlsx
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from case_lib import load_case  # noqa: E402

EXPECTED_SHEETS = [
    "00 封面与导读", "01 产品身份确认", "02 危险属性判定卡", "03 判定依据",
    "04 买方资质核查", "05 合规检查清单", "06 需索取证照", "07 留存与备案矩阵",
    "08 缺口与索要话术", "09 同流程产品对照", "10 法规索引", "11 人工确认与免责",
]

RESULT = []


def check(ok: bool, name: str, detail: str = ""):
    RESULT.append((bool(ok), name, detail))
    flag = "PASS" if ok else "FAIL"
    print(f"[{flag}] {name}" + (f" — {detail}" if detail else ""))


def verify_html(case, path: Path):
    html = path.read_text(encoding="utf-8")
    check(path.stat().st_size > 60_000, "HTML 体积 > 60 KB", f"{path.stat().st_size/1024:.1f} KB")
    check("__CASE__" not in html and "__TITLE__" not in html, "无未替换占位符")
    check("const CASE =" in html, "内嵌 CASE 数据")
    check(case["meta"]["case_no"] in html, "案例编号已注入")
    check(case["determination"]["verdict"] in html, "判定结论已注入")
    check(html.count("<section") == html.count("</section>"), "section 标签配对",
          f"{html.count('<section')}/{html.count('</section>')}")
    check(html.count("<script") == html.count("</script>"), "script 标签配对")
    check("localStorage" in html and "download(" in html, "清单可保存/可导出")
    check("buildDraft" in html and "复制草稿" in html, "话术生成与复制")
    check("不受影响" not in html, "无模板残留文案")
    # 每条核查项与义务项都在页面上（由 JS 渲染，检查数据是否内嵌完整）
    for c in case["buyer_checks"]:
        if c["id"] not in html:
            check(False, f"核查项 {c['id']} 缺失")
            return
    check(True, f"核查项 {len(case['buyer_checks'])} 条内嵌完整")
    ids = [o["id"] for k in ("seller", "buyer", "joint") for o in case["obligations"][k]]
    check(all(i in html for i in ids), f"义务项 {len(ids)} 条内嵌完整")


def verify_xlsx(case, path: Path):
    wb = load_workbook(path)
    names = wb.sheetnames
    check(names == EXPECTED_SHEETS, "工作表 12 张且顺序正确", f"{len(names)} 张")
    for name in EXPECTED_SHEETS:
        if name not in names:
            check(False, f"缺少工作表 {name}")
            return
    # 判定卡内容
    ws = wb["02 危险属性判定卡"]
    texts = []
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str):
                texts.append(c.value)
    joined = "\n".join(texts)
    check(case["determination"]["verdict"] in joined, "判定卡含结论")
    check("名录标记" in joined, "判定卡含名录标记")
    check("非危险化学品" in joined, "判定卡明确非危化品")
    # 核查表
    ws = wb["04 买方资质核查"]
    col_e = [ws.cell(row=r, column=5).value for r in range(1, ws.max_row + 1)]
    statuses = [v for v in col_e if v in ("A", "B", "C", "NA")]
    check(len(statuses) == len(case["buyer_checks"]), "核查状态行数一致",
          f"{len(statuses)}/{len(case['buyer_checks'])}")
    # 清单与数据验证
    ws = wb["05 合规检查清单"]
    check(len(ws.data_validations.dataValidation) >= 1, "清单含状态下拉框")
    total = sum(len(case["obligations"][k]) for k in ("seller", "buyer", "joint"))
    ids = set()
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=2).value
        if isinstance(v, str) and re.match(r"^[SPBJ]\d+$", v):
            ids.add(v)
    check(len(ids) == total, "清单条目数一致", f"{len(ids)}/{total}")
    # 证照
    ws = wb["06 需索取证照"]
    doc_ids = {ws.cell(row=r, column=2).value for r in range(1, ws.max_row + 1)}
    check(all(d["id"] in doc_ids for d in case["documents"]), "证照清单完整",
          f"{len(case['documents'])} 项")
    # 链接
    links = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.hyperlink:
                    links += 1
    check(links >= 8, "可点开链接 ≥ 8", f"{links} 条")
    # 空白页
    for name in EXPECTED_SHEETS:
        ws = wb[name]
        nonempty = sum(1 for row in ws.iter_rows() for c in row if c.value not in (None, ""))
        if nonempty < 20:
            check(False, f"{name} 内容过少", f"{nonempty} 个非空单元格")
            return
    check(True, "各表内容量充足")
    check(ws.max_row > 5, "人工确认表有内容")


def main(argv=None):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
    ap = argparse.ArgumentParser()
    ap.add_argument("case")
    ap.add_argument("--html", required=True)
    ap.add_argument("--xlsx", required=True)
    args = ap.parse_args(argv)

    case = load_case(args.case)
    check(True, f"案例加载 {case['meta']['case_no']}")
    verify_html(case, Path(args.html))
    verify_xlsx(case, Path(args.xlsx))

    failed = [r for r in RESULT if not r[0]]
    print("\n" + "=" * 60)
    print(f"RESULT: {'PASS' if not failed else 'FAIL'}  ({len(RESULT)-len(failed)}/{len(RESULT)})")
    for ok, name, detail in failed:
        print(f"  ✗ {name} {detail}")
    return 0 if not failed else 1


if __name__ == "__main__":
    raise SystemExit(main())
