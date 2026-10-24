#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""build_workbook.py —— 把 leads.json 渲染成 5 个工作表的 Excel（需 openpyxl）

表：① 标杆客户画像与复制逻辑 ② 说明与行业地图 ③ 客户线索总表 ④ 重点客户TOP ⑤ 分布统计
产品无关：标题、板块、行业地图、归并规则、相似维度关键词全部来自 project.json。

用法： python build_workbook.py --dir <项目> [--out 文件名.xlsx]
"""
import argparse
import os
import sys
from collections import Counter, OrderedDict

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _common import (bucket_of, load_json, load_project, norm_dims,  # noqa: E402
                     province_of, setup_console)

setup_console()

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    print("!! 需要 openpyxl：pip install openpyxl（HTML 报告不需要它）")
    sys.exit(2)

HEADERS = ["序号", "行业板块", "细分行业", "公司名称", "所在地（省/市）", "用量强度",
           "需求点（用料工序）", "相似维度（相似点落在哪一项）", "为什么会产生相似原料需求",
           "规模线索", "优先级", "置信度", "线索来源"]
WIDTHS = [6, 18, 20, 30, 16, 10, 52, 30, 56, 34, 9, 9, 44]
WRAP_COLS = (3, 7, 8, 9, 10, 13)

THIN = Side(style="thin", color="D0D7E5")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="2F5597")
HEAD_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=16, color="1F3864")
SEC_FONT = Font(bold=True, size=12, color="1F3864")
BODY_FONT = Font(size=10, color="404040")
PRIO_FILL = {"高": PatternFill("solid", fgColor="FCE4D6"),
             "中": PatternFill("solid", fgColor="FFF2CC"),
             "低": PatternFill("solid", fgColor="EDEDED")}


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = 30


def _kv_table(ws, r, headers, rows, widths=None):
    for i, h in enumerate(headers, start=1):
        ws.cell(row=r, column=i, value=h)
    style_header(ws, r, len(headers))
    r += 1
    for row in rows:
        for i, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=i, value=v)
            cell.font = BODY_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = BORDER
        ws.cell(row=r, column=1).font = Font(bold=True, size=10)
        r += 1
    return r


def sheet_profile(wb, profile, leads, cfg, title):
    p = cfg["product"]
    ws = wb.create_sheet("① 标杆客户画像")
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:E1")
    r = 3
    ws.cell(row=r, column=1, value="【标杆客户】").font = SEC_FONT
    ws.cell(row=r, column=2, value=profile.get("标杆客户") or cfg.get("benchmark", {}).get("name", ""))
    r += 1
    for k, v in (profile.get("主体信息") or {}).items():
        ws.cell(row=r, column=1, value=k).font = BODY_FONT
        c = ws.cell(row=r, column=2, value=str(v))
        c.font, c.alignment = BODY_FONT, Alignment(wrap_text=True, vertical="top")
        r += 1
    r += 1
    ws.cell(row=r, column=1, value=f"【关键业务特征（六要素）】产品：{p['name']}"
            + (f"（{p['grades']}）" if p.get("grades") else "")).font = SEC_FONT
    r += 1
    r = _kv_table(ws, r, ["特征项", "内容", "证据来源"],
                  [[f.get("特征项", ""), f.get("内容", ""), f.get("证据", "")]
                   for f in (profile.get("关键业务特征") or [])])
    r += 1
    ws.cell(row=r, column=1,
            value="【可复制的相似性判据：相似点落在哪一项，以及为什么会带来相似原料需求】").font = SEC_FONT
    r += 1
    r = _kv_table(ws, r, ["相似维度", "判据（相似点落在哪一项）", "为什么相似 → 相似原料需求"],
                  [[d.get("维度", ""), d.get("判据", ""), d.get("为什么带来相似原料需求", "")]
                   for d in (profile.get("可复制的相似性判据") or [])])
    r += 1
    ws.cell(row=r, column=1, value="【给销售的复制话术】").font = SEC_FONT
    ws.cell(row=r, column=2, value=profile.get("给销售的复制话术")
            or cfg["report"].get("sales_script", "")).alignment = Alignment(wrap_text=True)
    r += 2
    ws.cell(row=r, column=1, value="【本批潜客按相似维度的命中分布】").font = SEC_FONT
    r += 1
    dims = Counter()
    for ld in leads:
        for x in norm_dims(ld.get("相似维度"), cfg):
            dims[x] += 1
    order = [d.get("name") for d in (cfg.get("similarity_dimensions") or [])] or \
            ["产品结构", "工艺路线", "下游市场", "规模", "产业链位置"]
    r = _kv_table(ws, r, ["相似维度（相似点落在哪一项）", "命中线索数"],
                  [[k, dims.get(k, 0)] for k in order])
    for col, w in zip("ABCDE", [22, 74, 60, 14, 14]):
        ws.column_dimensions[col].width = w


def sheet_guide(wb, leads, cfg, industry_map):
    rep = cfg["report"]
    p = cfg["product"]
    reg = cfg.get("region", {})
    provs = reg.get("provinces") or []
    ws = wb.create_sheet("② 说明与行业地图")
    ws["A1"] = rep["title"]
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:G1")

    n_all = len(leads)
    n_hi = sum(1 for x in leads if x.get("置信度") == "高")
    n_mid = sum(1 for x in leads if x.get("置信度") == "中")
    n_low = sum(1 for x in leads if x.get("置信度") == "低")
    n_cninfo = sum(1 for x in leads if "cninfo" in x.get("来源", ""))
    n_mark = sum(1 for x in leads if "⚠" in x.get("来源", ""))
    n_prov = len({province_of(x.get("所在地", ""), provs) for x in leads} - {"未标注"})
    n_block = len({x.get("板块", "") for x in leads})
    n_prio_hi = sum(1 for x in leads if x.get("优先级") == "高")

    notes = [
        f"本表共 {n_all} 条线索，覆盖 {n_prov} 个省级地区、{n_block} 个行业板块；其中高优先级 {n_prio_hi} 条。",
        f"每条线索都写明了『相似点落在哪一项』（产品结构/工艺路线/下游市场/规模/产业链位置）"
        f"与『为什么会产生相似原料需求』。",
        "用法建议：先在「① 标杆客户画像」确认复制逻辑 → 用「③ 客户线索总表」按板块/省市/相似维度筛选"
        " → 按「需求点」里的用料工序准备开场话术。",
        "用量强度：★★★★★ = 单一客户可达千吨级/年；★★★★ = 百吨级；★★★ = 数十吨级；★★ = 十吨级及以下。",
        "优先级 = 用量强度 × 证据置信度：高=5星且有证据，或4星且证据强；中=4星证据一般，或3星但证据很强；低=证据薄弱或用量偏小。",
        "　　　　优先级只用于排外呼顺序，不代表客户价值绝对高低。某些板块普遍为★★★，故无「高」优先级，但单价高、粘性强，不应低估。",
        "",
        "【数据来源与核验方法】",
        f"证据置信度：高 {n_hi} 条 / 中 {n_mid} 条 / 低 {n_low} 条。",
        f"其中 {n_cninfo} 条的「线索来源」为上市公司公告（如巨潮资讯）原文——年报、招股书、可研报告、"
        "问询函回复里的原文，直接写明该企业采购或使用该原料的环节，电话里可直接复述。",
        "其余来源为企业官网直连核验、交易所/财经站公司资料页。所有 URL 均经实际抓取确认，未凭记忆编造。",
        f"来源列中带 ⚠ 标记的 {n_mark} 条，是链接抽检未通过的条目（官网故障、空页、反爬拦截等），"
        "公司信息可用但请勿直接引用该链接。",
        "",
        "【使用边界（请务必阅读）】",
        "1. 「需求点」列：有公告/官网原文佐证的为实证；披露率低的行业，其用料环节是按品类工艺推导的合理推断，",
        "　　首次接触时建议向对方研发/品控部门确认实际品种与用量。",
        f"2. 已剔除同业：自身也生产{p['name']}的企业属于竞争对手而非下游客户，未计入本表（名单见 Word/HTML 报告第 6 节）。",
        "3. 正式外呼前，建议用企业信息平台核对工商信息、联系方式与经营状态（本表未包含联系人信息）。",
        "4. 未收录经销商与贸易商；仅收录有自有生产实体的工厂（把该产品当原料深加工/复配再转售的厂除外，属合格客户）。",
    ]
    if rep.get("evidence_note"):
        notes.append("5. 取证通道说明：" + str(rep["evidence_note"]))
    if rep.get("boundary_note"):
        notes.append("6. " + str(rep["boundary_note"]))

    r = 3
    for n in notes:
        cell = ws.cell(row=r, column=1, value=n)
        cell.font = Font(size=11, bold=True, color="1F3864") if n.startswith("【") else BODY_FONT
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
        r += 1

    if industry_map:
        r += 1
        for i, h in enumerate(["行业板块", "细分方向", "用料逻辑 / 工序", "用量强度", "采购与开发要点"], start=1):
            ws.cell(row=r, column=i, value=h)
        style_header(ws, r, 5)
        head_row = r
        r += 1
        for row in industry_map:
            for i, v in enumerate(row, start=1):
                cell = ws.cell(row=r, column=i, value=v)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = BORDER
            ws.cell(row=r, column=1).font = Font(bold=True)
            ws.cell(row=r, column=4).alignment = Alignment(horizontal="center", vertical="top")
            r += 1
        ws.freeze_panes = ws.cell(row=head_row + 1, column=1)

    for col, w in zip("ABCDEFG", [16, 34, 62, 11, 64, 12, 12]):
        ws.column_dimensions[col].width = w


def sheet_leads(wb, leads, title):
    ws = wb.create_sheet(title)
    for i, h in enumerate(HEADERS, start=1):
        ws.cell(row=1, column=i, value=h)
    style_header(ws, 1, len(HEADERS))
    prio_rank = {"高": 0, "中": 1, "低": 2}
    ordered = sorted(leads, key=lambda x: (x.get("板块", ""),
                                           prio_rank.get(x.get("优先级", "中"), 3),
                                           {"高": 0, "中": 1, "低": 2}.get(x.get("置信度", "中"), 3),
                                           x.get("公司", "")))
    for idx, item in enumerate(ordered, start=1):
        row = [idx, item.get("板块", ""), item.get("细分行业", ""), item.get("公司", ""),
               item.get("所在地", ""), item.get("用量强度", ""), item.get("需求点", ""),
               item.get("相似维度", ""), item.get("相似逻辑", ""), item.get("规模线索", ""),
               item.get("优先级", "中"), item.get("置信度", "中"), item.get("来源", "")]
        r = idx + 1
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=(c in WRAP_COLS))
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=r, column=4).font = Font(bold=True)
        for c in (5, 6, 11, 12):
            ws.cell(row=r, column=c).alignment = Alignment(horizontal="center", vertical="top")
        fill = PRIO_FILL.get(item.get("优先级", "中"))
        if fill:
            ws.cell(row=r, column=11).fill = fill
    for i, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:M{len(ordered) + 1}"
    return ws


def sheet_stats(wb, leads, buckets, provs):
    ws = wb.create_sheet("⑤ 分布统计")
    ws["A1"] = "线索分布统计"
    ws["A1"].font = Font(bold=True, size=14, color="1F3864")

    def block(title, pairs, hi_counts=None):
        r = ws.max_row + 2
        ws.cell(row=r, column=1, value=title).font = Font(bold=True, size=12)
        r += 1
        heads = ["项目", "线索数", "其中高优先级"] if hi_counts is not None else ["项目", "线索数"]
        for i, h in enumerate(heads, start=1):
            ws.cell(row=r, column=i, value=h)
        style_header(ws, r, len(heads))
        r += 1
        for name, n in pairs:
            ws.cell(row=r, column=1, value=name).border = BORDER
            ws.cell(row=r, column=2, value=n).border = BORDER
            if hi_counts is not None:
                ws.cell(row=r, column=3, value=hi_counts.get(name, 0)).border = BORDER
            r += 1

    blocks = OrderedDict()
    for it in leads:
        b = it.get("板块", "未分类")
        blocks.setdefault(b, [0, 0])
        blocks[b][0] += 1
        if it.get("优先级") == "高":
            blocks[b][1] += 1
    ob = sorted(blocks.items(), key=lambda kv: -kv[1][0])
    block("按行业板块", [(b, v[0]) for b, v in ob], {b: v[1] for b, v in ob})

    prov = Counter(province_of(it.get("所在地", ""), provs) for it in leads)
    block("按所在省份（≥2 条）",
          [(p, n) for p, n in sorted(prov.items(), key=lambda kv: (-kv[1], kv[0])) if n >= 2])

    sub, sub_hi = Counter(), Counter()
    for it in leads:
        b = bucket_of(it.get("细分行业", ""), buckets)
        sub[b] += 1
        if it.get("优先级") == "高":
            sub_hi[b] += 1
    block("按细分品类（已归并）",
          [(s, n) for s, n in sorted(sub.items(), key=lambda kv: (-kv[1], kv[0]))], sub_hi)

    conf = Counter(it.get("置信度", "中") for it in leads)
    block("按置信度", [(k, conf.get(k, 0)) for k in ("高", "中", "低")])
    pr = Counter(it.get("优先级", "中") for it in leads)
    block("按优先级", [(k, pr.get(k, 0)) for k in ("高", "中", "低")])
    for col, w in zip("ABC", [44, 12, 16]):
        ws.column_dimensions[col].width = w


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", default=".")
    ap.add_argument("--out", default=None)
    args = ap.parse_args()
    d = os.path.abspath(args.dir)
    cfg = load_project(d)
    leads = load_json(os.path.join(d, "leads.json"), [])
    if not leads:
        raise SystemExit("leads.json 为空或不存在，先跑 merge_leads.py")
    profile = load_json(os.path.join(d, cfg.get("benchmark", {}).get("profile_file")
                                     or "benchmark_profile.json"), {})
    rep = cfg["report"]
    out = os.path.join(d, args.out or rep.get("workbook_output") or "下游潜客池.xlsx")

    wb = Workbook()
    wb.remove(wb.active)
    sheet_profile(wb, profile, leads, cfg, rep["title"])
    sheet_guide(wb, leads, cfg, cfg.get("industry_map") or [])
    sheet_leads(wb, leads, "③ 客户线索总表")
    top = [x for x in leads if x.get("优先级") == "高"]
    sheet_leads(wb, top, "④ 重点客户TOP")
    sheet_stats(wb, leads, cfg.get("buckets") or {}, cfg.get("region", {}).get("provinces") or [])
    wb.save(out)
    print(f"OK 线索={len(leads)} 高优先级={len(top)} -> {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
