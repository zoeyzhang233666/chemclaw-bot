#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""产物自检：xlsx 与 html 的结构、关键内容、以及和 case.json 的一致性。

用法:
  python check_outputs.py --case case.json [--out out]
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys
import zipfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import caselib as C  # noqa: E402

SHEETS = ["00_使用说明", "01_客户档案", "02_事实台账", "03_主案例版本对照", "04_边界审计",
          "05_场景矩阵", "06_占位符清单", "07_快速调整指令库", "08_发送前检查清单",
          "09_追问应对", "10_复盘看板"]


def check(case_path, outdir):
    case = C.load_case(case_path)
    xlsx, html = C.out_paths(case, outdir)
    bad = []
    ok = []

    # ---------------- xlsx ----------------
    if not os.path.exists(xlsx):
        bad.append("xlsx 不存在：%s" % xlsx)
    else:
        if not zipfile.is_zipfile(xlsx):
            bad.append("xlsx 不是有效的 zip 容器")
        try:
            from openpyxl import load_workbook
            wb = load_workbook(xlsx)
            names = wb.sheetnames
            if names != SHEETS:
                bad.append("工作表不符：%s" % names)
            else:
                ok.append("xlsx 11 张表齐全")
            empty = [n for n in names if sum(1 for r in wb[n].iter_rows() for c in r if c.value) < 3]
            if empty:
                bad.append("疑似空表：%s" % empty)
            txt = " ".join(str(c.value) for n in names for r in wb[n].iter_rows()
                           for c in r if isinstance(c.value, str))
            for kw in ["边界审计", "【待填】"]:
                if kw not in txt:
                    bad.append("xlsx 缺少关键串：%s" % kw)
            n_audit = sum(1 for r in wb["04_边界审计"].iter_rows(min_row=5, max_row=5 + len(case["audit"]) - 1,
                                                                min_col=1, max_col=1) if r[0].value)
            if n_audit != len(case["audit"]):
                bad.append("审计行数不符：xlsx %d vs case %d" % (n_audit, len(case["audit"])))
            if "10_复盘看板" in names:
                f = wb["10_复盘看板"]
                has_formula = any(isinstance(c.value, str) and c.value.startswith("=")
                                  for r in f.iter_rows() for c in r)
                if has_formula:
                    ok.append("看板含公式（打开后按 F9 或自动重算）")
        except ImportError:
            bad.append("缺少 openpyxl，无法校验 xlsx")

    # ---------------- html ----------------
    if not os.path.exists(html):
        bad.append("html 不存在：%s" % html)
    else:
        src = open(html, encoding="utf-8").read()
        size = len(src.encode())
        if size < 20000:
            bad.append("html 偏小（%d 字节），可能未注入内容" % size)
        else:
            ok.append("html 单文件 %d KB" % (size // 1024))
        if re.search(r'(src|href)\s*=\s*"https?://', src, re.I):
            bad.append("html 含外链资源，不能离线使用")
        else:
            ok.append("html 无外链，可离线")
        if src.count('<script src=') or src.count('<link '):
            bad.append("html 引用了外部 script/link")
        m = re.search(r'<script id="case-data" type="application/json">(.*?)</script>', src, re.S)
        if not m:
            bad.append("html 缺少 case-data 注入块")
        else:
            try:
                data = json.loads(m.group(1).replace("<\\/", "</"))
                if data.get("meta", {}).get("file_stem") != case["meta"].get("file_stem"):
                    bad.append("html 注入数据与 case.json 不一致")
                else:
                    ok.append("html 数据注入正确（%d 条事实 / %d 条审计）" %
                              (len(data["facts"]), len(data["audit"])))
            except Exception as exc:
                bad.append("html 注入 JSON 解析失败：%s" % exc)
        ids = re.findall(r'\bid="([^"]+)"', src)
        dup = sorted({i for i in ids if ids.count(i) > 1})
        if dup:
            bad.append("html id 重复：%s" % dup)
        for kw in ["边界审计", "待填", "复制"]:
            if kw not in src:
                bad.append("html 缺少关键串：%s" % kw)

    print("  [自检] " + ("；".join(ok) if ok else "无通过项"))
    for b in bad:
        print("  [问题] " + b)
    return not bad


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--case", required=True)
    ap.add_argument("--out", default="out")
    args = ap.parse_args(argv)
    return 0 if check(args.case, args.out) else 1


if __name__ == "__main__":
    sys.exit(main())
