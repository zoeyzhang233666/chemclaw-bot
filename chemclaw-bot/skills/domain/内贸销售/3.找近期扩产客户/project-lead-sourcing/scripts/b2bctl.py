#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""b2bctl —— 化工品下游「项目级潜客」流水线命令行。

设计目标：**agent 无关、产品无关、可换电脑**。
任何能执行 shell 的 agent（codex / chemclaw / workbuddy / Claude Code / 本机终端）都能用；
数据源能力靠 `probe` 实测，不假设网络。

典型顺序（对应方法论 9 步）：
  1. b2bctl new        建项目（定契约）
  2. b2bctl probe      探通道
  3. 编辑 lanes.json   画应用地图定 lane（用 b2bctl lanes-template 起步）
  4. b2bctl brief      生成每个 lane 的派工单 → 交给子代理/并行执行
  5. b2bctl collect    校验每个 lane 的产出
  6. b2bctl merge      合并去重、peer gate、销售时点推导
  7. b2bctl qa         链接抽检
  8/9. b2bctl build    出 Excel + HTML + 方法与边界说明
  （或直接 b2bctl all 一把跑完 6~9）

用法： python b2bctl.py <command> [项目目录] [选项]
"""
import argparse
import collections
import glob
import json
import os
import re
import shutil
import subprocess
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
if HERE not in sys.path:
    sys.path.insert(0, HERE)

from lib import common as C          # noqa: E402
from lib import merge as M           # noqa: E402
from lib import report_out as R      # noqa: E402
from lib import channels as CH       # noqa: E402
from tools import http as H          # noqa: E402


# ------------------------------------------------------------------ helpers

def die(msg, code=1):
    print('错误：%s' % msg, file=sys.stderr)
    sys.exit(code)


def ok(msg):
    print('[ok] %s' % msg)


# 脚手架模板首行的标记：check 用它判断"判断性内容是否还是空的"
STATUS_UNFILLED = '<!-- STATUS: UNFILLED -->'


def project(path):
    """定位项目根并加载 config。"""
    path = os.path.abspath(path or '.')
    cfgp = os.path.join(path, 'config.json')
    if not os.path.exists(cfgp):
        die('在 %s 找不到 config.json。先用 `b2bctl new <目录>` 建项目。' % path)
    cfg, root = C.load_config(cfgp)
    return cfg, root


def load_leads(cfg, root):
    p = os.path.join(root, cfg['paths']['leads'])
    if not os.path.exists(p):
        die('找不到 %s。先跑 `b2bctl merge`。' % p)
    with open(p, encoding='utf-8-sig') as f:
        return json.load(f)


# ------------------------------------------------------------------ new

LANES_TEMPLATE = {
    "_help": "应用地图。每个 lane = 一个下游应用方向（不是企业名单）。填 logic（用料逻辑）、intensity（1~5 星）、clusters（产业集聚地）、seeds（5~8 家候选企业）。删掉不需要的占位 lane。",
    "product": "",
    "lanes": [],
}

LANE_PLACEHOLDER = {
    "id": "",
    "name": "",
    "logic": "本产品在该方向的工序里起什么作用、大致单耗量级",
    "intensity": 3,
    "intensity_basis": "5★=单点年需求千吨级以上；4★=千吨级；3★=百吨级；2★=十吨级；1★=更小/不确定",
    "clusters": [],
    "seeds": [],
}


def cmd_new(a):
    root = os.path.abspath(a.dir)
    os.makedirs(os.path.join(root, 'leads_raw'), exist_ok=True)
    cfgp = os.path.join(root, 'config.json')
    if os.path.exists(cfgp) and not a.force:
        die('%s 已存在（加 --force 覆盖）' % cfgp)

    ex_path = os.path.join(HERE, 'config.example.json')
    with open(ex_path, encoding='utf-8-sig') as f:
        cfg = json.load(f)
    cfg.pop('_comment', None)
    cfg['product']['name'] = a.product
    if a.formula:
        cfg['product']['formula'] = a.formula
    if a.aliases:
        cfg['product']['aliases'] = [x.strip() for x in a.aliases.split(',') if x.strip()]
    cfg['region']['name'] = a.region
    if a.provinces:
        cfg['region']['provinces'] = [x.strip() for x in a.provinces.split(',') if x.strip()]
    cfg['window_months'] = a.months
    cfg['today'] = a.today or cfg['today']
    if a.peers:
        cfg['peer_keywords'] = [x.strip() for x in a.peers.split(',') if x.strip()]
    with open(cfgp, 'w', encoding='utf-8') as f:
        json.dump(cfg, f, ensure_ascii=False, indent=2)

    lt = dict(LANES_TEMPLATE)
    lt['product'] = a.product
    lt['lanes'] = [dict(LANE_PLACEHOLDER, id='%02d' % i) for i in (1,)]
    lp = os.path.join(root, cfg['paths']['lanes'])
    if not os.path.exists(lp) or a.force:
        with open(lp, 'w', encoding='utf-8') as f:
            json.dump(lt, f, ensure_ascii=False, indent=2)

    gaps = os.path.join(root, 'gaps.md')
    if not os.path.exists(gaps) or a.force:
        with open(gaps, 'w', encoding='utf-8') as f:
            f.write(
                STATUS_UNFILLED + '\n'
                '# 判断性缺口与负结果（由 agent/人填写）\n\n'
                '> 这一节讲**判断**，脚本算不出来。交付前必须写满，否则第四块不完整。\n'
                '> **填完后请删掉第一行的 STATUS 标记**，否则 `b2bctl check` 会判不通过。\n\n'
                '## 1. 哪些通道失效，因此没覆盖到哪些省份/板块\n\n（例：某省生态环境厅栏目为 JS 渲染，'
                '抓不到列表 → 该省非上市项目只能靠公告，系统性低估）\n\n'
                '## 2. 哪些是「负结果」\n\n（例：某板块当前处于下行周期，主体在关停/拆装置而不是新建 → '
                '窗口内不宜作为主攻方向）\n\n'
                '## 3. 有项目但缺关键信息、只能人工补证的\n\n（例：项目确凿但公告未载厂址城市）\n\n'
                '## 4. 非上市企业覆盖程度与原因\n')

    peers_md = os.path.join(root, 'peers_excluded.md')
    if not os.path.exists(peers_md) or a.force:
        with open(peers_md, 'w', encoding='utf-8') as f:
            f.write(
                STATUS_UNFILLED + '\n'
                '# lane 阶段剔除的同行（由 agent/人填写）\n\n'
                '> **为什么必须单独有一份**：lane 在检索时就跳过了同行，所以它们不会出现在\n'
                '> `leads_raw/` 里，`merge` 的 peer 计数必然是 **0**。只依赖自动统计，\n'
                '> 交付四块的第三块会错误地报告"本轮没有同行"。\n'
                '> **填完后请删掉第一行的 STATUS 标记**，否则 `b2bctl check` 会判不通过。\n\n'
                '| 被剔除主体 | 依据（哪份材料证明它自产本产品） | 原以为的机会 |\n'
                '|---|---|---|\n'
                '|  |  |  |\n\n'
                '> 提示：`config.json → peer_keywords` 里预置的关键词是**检索期过滤**，\n'
                '> 这里记录的是**在写作阶段人工判定**、keyword 未覆盖到的那些主体。\n')

    readme = os.path.join(root, 'README.md')
    if not os.path.exists(readme) or a.force:
        with open(readme, 'w', encoding='utf-8') as f:
            f.write('# %s · %s · 未来 %d 个月 潜客项目\n\n'
                    '由 chem-lead-sourcing 生成。执行顺序：\n\n'
                    '```bash\n'
                    'python b2bctl.py probe   .    # 探通道（换电脑必跑）\n'
                    '# 编辑 lanes.json 填应用地图\n'
                    'python b2bctl.py brief   .    # 生成每个 lane 的派工单\n'
                    '# 把 briefs/*.md 分给子代理或自己逐条做，产出写到 leads_raw/\n'
                    'python b2bctl.py collect .    # 校验每个 lane 产出\n'
                    'python b2bctl.py all     .    # merge + qa + build\n'
                    '```\n' % (a.product, a.region, a.months))

    ok('项目已建：%s' % root)
    print('  config.json   产品=%s 地区=%s 窗口=%d个月 today=%s'
          % (a.product, a.region, a.months, cfg['today']))
    print('  lanes.json    请填写应用地图（见 b2bctl lanes-template 的字段说明）')
    print('  gaps.md       第四块的判断性内容占位')
    print('\n下一步： python b2bctl.py probe "%s"' % root)


# ------------------------------------------------------------------ probe

def cmd_probe(a):
    cfg, root = project(a.dir)
    print('通道探测 %s（超时 %ds）\n' % (root, a.timeout))
    res = CH.probe(cfg, root, registry_path=a.channels, timeout=a.timeout)
    print('\n可达通道（status=可达 只说明 URL 能取到内容，不代表能解析出项目条目）：')
    for c in res['channels']:
        if c['status'] == '可达':
            print('  - %-14s %s' % (c['kind'], c['name']))
    bad = [c for c in res['channels'] if c['status'] != '可达']
    if bad:
        print('\n不可达（换网络后可能恢复，建议重跑 probe）：')
        for c in bad:
            print('  - %-40s %s' % (c['status'], c['name']))
    if res.get('needs_adapter_test'):
        print('\n⚠ 下列列表型通道须再用适配器实测「可解析性」，否则可能空手而归：')
        for cid in res['needs_adapter_test']:
            row = next((c for c in res['channels'] if c['id'] == cid), None)
            if row:
                print('  - %s（适配器 %s）' % (row['name'], row.get('adapter') or '未注册'))
        print('  例： python -m tools.eia selftest')
    print('\n提示：若宿主 agent 自带联网搜索，请把它登记进 %s 的 host_tools。'
          % cfg['paths']['channels'])


# ------------------------------------------------------------------ lanes

def cmd_lanes_template(a):
    root = os.path.abspath(a.dir or '.')
    cfgp = os.path.join(root, 'config.json')
    product = a.product
    if not product and os.path.exists(cfgp):
        product = C.load_config(cfgp)[0]['product']['name']
    lt = dict(LANES_TEMPLATE)
    lt['product'] = product or ''
    cats = ['直接原料型（本产品进入产品分子）', '配合料/助熔剂型', 'pH调节/中和/浸出型',
            '反应中间体/催化型', '环保（脱硫脱硝/废水）型', '水处理/清洗型']
    lt['lanes'] = [dict(LANE_PLACEHOLDER, id='%02d' % (i + 1), name=c) for i, c in enumerate(cats)]
    out = os.path.join(root, a.lanes or 'lanes.json')
    with open(out, 'w', encoding='utf-8') as f:
        json.dump(lt, f, ensure_ascii=False, indent=2)
    ok('已生成 %s —— 请把 name/logic/intensity/clusters/seeds 换成该产品真实的 %s 个应用方向'
       % (out, len(cats)))


def cmd_lanes_check(a):
    cfg, root = project(a.dir)
    by_id, lanes = C.load_lanes(cfg, root)
    if not lanes:
        die('lanes.json 不存在或为空（先跑 b2bctl lanes-template）')
    print('应用地图：%d 个 lane，产品=%s\n' % (len(lanes), cfg['product']['name']))
    print('%-4s %-26s %-4s %-22s %s' % ('id', 'name', '★', '集聚地', '候选数'))
    problems = []
    for ln in lanes:
        i = ln.get('intensity', 3)
        if not isinstance(i, int) or not (1 <= i <= 5):
            problems.append('lane %s 的 intensity 必须是 1~5 的整数' % ln['id'])
        if len(ln.get('seeds') or []) < 3:
            problems.append('lane %s 的 seeds 少于 3 家，派工单会太虚' % ln['id'])
        if not (ln.get('logic') or '').strip():
            problems.append('lane %s 缺 logic（用料逻辑）' % ln['id'])
        if ln.get('name', '').startswith(('直接原料型', '配合料', 'pH调节', '反应中间体', '环保', '水处理')):
            problems.append('lane %s 还是占位名「%s」，请换成真实应用方向' % (ln['id'], ln['name']))
        print('%-4s %-26s %-4s %-22s %d'
              % (ln['id'], ln['name'][:26], '★' * int(i if isinstance(i, int) else 3),
                 '/'.join((ln.get('clusters') or [])[:2])[:22], len(ln.get('seeds') or [])))
    total_seeds = sum(len(ln.get('seeds') or []) for ln in lanes)
    print('\n候选企业合计 %d 家' % total_seeds)
    if problems:
        print('\n待修正：')
        for p in problems:
            print('  - %s' % p)
    else:
        ok('lanes.json 结构完整')


# ------------------------------------------------------------------ brief

def _usable_channels_text(cfg, root):
    path = os.path.join(root, cfg['paths']['channels'])
    data = CH.summarize(path)
    if not data:
        return ('（**尚未探测通道**。请先跑 `b2bctl probe`。在探明之前不要假设任何站点可用；'
                '若宿主 agent 自带联网搜索，优先用它，并把命中页面的原始 URL 记进 `来源`。）')
    lines = []
    for c in data.get('channels', []):
        if c['status'] != '可达':
            continue
        tag = ''
        if c.get('needs_adapter_test'):
            tag = '；**列表型：可达 ≠ 能拿到条目**，先用适配器实测再依赖它' \
                  + (('`%s`' % c['adapter']) if c.get('adapter') else '（未注册适配器）')
        elif c.get('adapter'):
            tag = '；适配器 `%s`' % c['adapter']
        lines.append('- **%s** — `%s`%s' % (c['name'], c['url'], tag))
    for t in data.get('host_tools', []):
        if t.get('status', '').startswith('可用'):
            lines.append('- **%s** — %s' % (t['name'], t.get('note', '')))
    if not lines:
        lines.append('（探测结果里没有可达通道：不要硬抓，改为把任务退回给派工方，'
                     '或只使用宿主 agent 自带的联网搜索。）')
    if data.get('needs_adapter_test'):
        lines.append('\n> 上面标了"列表型"的通道，只被证明**可达**。JS 渲染的列表页'
                     '能返回 200 和很长的 HTML，但项目清单是前端注入的、抓不到条目——'
                     '**必须先用适配器实测一次**（如 `python -m tools.eia selftest`），'
                     '确认能抽出条目再依赖它；抽不到就如实报告该通道不可用。')
    return '\n'.join(lines)


def _render_brief(cfg, root, lane, tmpl):
    prod = cfg['product']['name']
    t, _, _, win_end = C.window_bounds(cfg)
    raw_dir = cfg['paths']['raw_dir']
    out_file = os.path.join(root, raw_dir, '%s_%s.json'
                            % (lane['id'], re.sub(r'[\\/:*?"<>|]+', '_', lane['name'])))
    fanout = cfg['lane_fanout']
    vals = {
        'lane_id': lane['id'], 'lane_name': lane['name'],
        'lane_logic': lane.get('logic') or '（未填写——请在 lanes.json 补 logic）',
        'lane_intensity': '★' * int(lane.get('intensity', 3)),
        'intensity_scale': lane.get('intensity_basis') or
                           '5★=单点年需求千吨级以上；4★=千吨级；3★=百吨级；2★=十吨级；1★=更小/不确定',
        'lane_clusters': '、'.join(lane.get('clusters') or []) or '（未填写）',
        'lane_seeds': '、'.join(lane.get('seeds') or []) or '（未填写）',
        'product': prod, 'formula': cfg['product'].get('formula') or '',
        'aliases': '、'.join(cfg['product'].get('aliases') or []) or '（无）',
        'region': cfg['region']['name'],
        'provinces': '、'.join(cfg['region'].get('provinces') or []) or '不限',
        'months': cfg['window_months'], 'today': str(t), 'window_end': str(win_end),
        'stages': ' / '.join(cfg.get('stages_vocab') or C.DEFAULT_STAGES),
        'roles': ' / '.join(cfg['product'].get('roles_vocab') or C.DEFAULT_ROLES),
        'source_types': ' / '.join(cfg.get('source_types_vocab') or C.DEFAULT_SRC_TYPES),
        'peer_keywords': '、'.join(cfg.get('peer_keywords') or []) or '（未配置，请务必补充！）',
        'channels_usable': _usable_channels_text(cfg, root),
        'raw_dir': raw_dir, 'out_file': out_file,
        'min_records': fanout.get('min_records_per_lane', 8),
        'max_records': fanout.get('max_records_per_lane', 15),
        'max_shell_calls': fanout.get('max_shell_calls', 35),
        'product_profile': '',
    }
    pp = os.path.join(root, 'product_profile.md')
    if os.path.exists(pp):
        with open(pp, encoding='utf-8-sig') as f:
            vals['product_profile'] = '## 产品档案（项目根 product_profile.md）\n\n' + f.read().strip()
    text = tmpl
    for k, v in vals.items():
        text = text.replace('{{%s}}' % k, str(v))
    return text, out_file


def cmd_brief(a):
    cfg, root = project(a.dir)
    by_id, lanes = C.load_lanes(cfg, root)
    if not lanes:
        die('lanes.json 不存在或为空')
    lanes = lanes[:a.limit] if a.limit else lanes
    tmpl_path = os.path.join(HERE, 'templates', 'lane_brief.md.tmpl')
    with open(tmpl_path, encoding='utf-8-sig') as f:
        tmpl = f.read()

    bdir = os.path.join(root, 'briefs')
    os.makedirs(bdir, exist_ok=True)
    rows = []
    for lane in lanes:
        text, out_file = _render_brief(cfg, root, lane, tmpl)
        name = '%s_%s.md' % (lane['id'], re.sub(r'[\\/:*?"<>|]+', '_', lane['name']))
        with open(os.path.join(bdir, name), 'w', encoding='utf-8') as f:
            f.write(text)
        rows.append((lane, name, out_file))

    mc = cfg['lane_fanout'].get('max_concurrency', 8)
    wp = ['# 派工单 · %s · %s · 未来 %d 个月\n'
          % (cfg['product']['name'], cfg['region']['name'], cfg['window_months'])]
    wp.append('- 生成日：%s' % C.today_date(cfg))
    wp.append('- lane 数：**%d**' % len(rows))
    wp.append('- 并发上限：**%d**（一次不要超过这个数，否则质量下降）' % mc)
    wp.append('- 每个 lane 的派工单在 `briefs/`，产出写到 `leads_raw/`\n')
    wp.append('| lane | 名称 | 用量强度 | 派工单 | 产出文件 |')
    wp.append('|---|---|---|---|---|')
    for lane, name, out_file in rows:
        wp.append('| %s | %s | %s | `briefs/%s` | `%s` |'
                  % (lane['id'], lane['name'], '★' * int(lane.get('intensity', 3)),
                     name, os.path.relpath(out_file, root)))
    wp.append('\n## 怎么派\n')
    wp.append('把 `briefs/<lane>.md` 的**全文**作为子任务指令交给一个子代理（或你自己按顺序做）。')
    wp.append('每个 brief 已自包含：契约、用料逻辑、候选企业、可用通道、JSON 字段、纪律、回复格式。\n')
    wp.append('**不要**在派工时再做改写或摘要——brief 里的纪律（禁止编造、peer gate、地区闸门、'
              '工具调用上限、增量落盘）是这个方法论的质量来源。\n')
    wp.append('全部 lane 完成后：`python b2bctl.py all .`\n')
    with open(os.path.join(root, cfg['paths']['workplan']), 'w', encoding='utf-8') as f:
        f.write('\n'.join(wp))

    ok('已生成 %d 份派工单 → %s' % (len(rows), bdir))
    for lane, name, out_file in rows:
        print('  %-4s %-28s → %s' % (lane['id'], lane['name'], name))
    print('  派工总表 → %s' % cfg['paths']['workplan'])


# ------------------------------------------------------------------ collect

def cmd_collect(a):
    cfg, root = project(a.dir)
    pattern = a.file or os.path.join(root, cfg['paths']['raw_dir'], '*.json')
    files = sorted(glob.glob(pattern))
    files = [f for f in files if not os.path.basename(f).startswith('_')]
    if not files:
        die('没有匹配的文件：%s' % pattern)

    total_bad = 0
    for f in files:
        base = os.path.basename(f)
        try:
            with open(f, encoding='utf-8-sig') as fh:
                data = json.load(fh)
        except Exception as e:
            print('%-44s JSON 解析失败：%s' % (base, e))
            total_bad += 1
            continue
        if not isinstance(data, list):
            print('%-44s 顶层不是数组' % base)
            total_bad += 1
            continue
        issues = []
        for i, r in enumerate(data):
            if not isinstance(r, dict):
                issues.append('#%d 不是对象' % (i + 1))
                continue
            for k in C.COLLECT_FIELDS:
                if k not in r:
                    issues.append('#%d 缺字段 %s' % (i + 1, k))
            st = C.norm_stage(r.get('建设阶段'), cfg.get('stages_vocab'))
            if st == '未披露' and (r.get('建设阶段') or '').strip():
                issues.append('#%d 建设阶段无法识别：%r' % (i + 1, r.get('建设阶段')))
            role = (r.get('产品角色') or '').strip()
            if role and role not in (cfg['product'].get('roles_vocab') or C.DEFAULT_ROLES):
                issues.append('#%d 产品角色不在词表：%r' % (i + 1, role))
            if (r.get('置信度') or '') not in ('高', '中', '低'):
                issues.append('#%d 置信度必须是 高/中/低：%r' % (i + 1, r.get('置信度')))
            src = (r.get('来源') or '').strip()
            if src and src != '未核实' and not src.lower().startswith(('http://', 'https://')):
                issues.append('#%d 来源不是 URL：%r' % (i + 1, src[:60]))
            if not C.in_region(cfg, r.get('省市')):
                issues.append('#%d 省市不在地区范围：%r' % (i + 1, r.get('省市')))
            if C.is_peer(cfg, r.get('公司')):
                issues.append('#%d 命中 peer（同行，应剔除）：%s' % (i + 1, r.get('公司')))
        conf = collections.Counter((r.get('置信度') or '?') for r in data if isinstance(r, dict))
        srcs = sum(1 for r in data if isinstance(r, dict)
                   and (r.get('来源') or '').lower().startswith('http'))
        flag = 'OK ' if not issues else '!! '
        print('%s%-42s %2d 条 | 高%d 中%d 低%d | 有链接%d'
              % (flag, base, len(data), conf.get('高', 0), conf.get('中', 0),
                 conf.get('低', 0), srcs))
        for s in issues[:12]:
            print('      - %s' % s)
        if len(issues) > 12:
            print('      ... 另有 %d 条' % (len(issues) - 12))
        total_bad += len(issues)
    print()
    if total_bad:
        print('共发现 %d 处待修（修完再 merge，否则会被静默剔除或计入错误分布）' % total_bad)
    else:
        ok('所有 lane 产出结构合规')


# ------------------------------------------------------------------ merge/qa/build

def cmd_merge(a):
    cfg, root = project(a.dir)
    leads, stats = M.merge(cfg, root)
    M.print_summary(stats)
    ok('%d 条 → %s' % (len(leads), cfg['paths']['leads']))


def cmd_qa(a):
    cfg, root = project(a.dir)
    leads = load_leads(cfg, root)
    out = os.path.join(root, cfg['paths']['qa_report'])
    from lib import qa as Q
    rep = Q.sample_and_check(cfg, leads, out, timeout=a.timeout, workers=a.workers)
    if rep['broken_count']:
        print('\n注意：有 %d 条坏链/失败，交付说明里必须如实列出。' % rep['broken_count'])


def cmd_build(a):
    cfg, root = project(a.dir)
    leads = load_leads(cfg, root)
    with open(os.path.join(root, cfg['paths']['stats']), encoding='utf-8-sig') as f:
        stats = json.load(f)
    stem = C.render_stem(cfg)
    outs = cfg['outputs']

    made = []
    if outs.get('xlsx', True):
        from lib import excel_out as X
        p = os.path.join(root, stem + '.xlsx')
        try:
            made.append(X.build_workbook(cfg, leads, p))
        except ImportError:
            print('提示：本机没有 openpyxl，Excel 降级为 CSV。`pip install openpyxl` 后可出 xlsx。',
                  file=sys.stderr)
            from lib import excel_out as X2
            p = os.path.join(root, stem + '.csv')
            made.append(X2.build_csv(cfg, leads, p))
    if outs.get('html', True):
        from lib import html_out as HT
        made.append(HT.build_html(cfg, leads, os.path.join(root, stem + '.html')))
    if outs.get('md', True):
        ch = CH.summarize(os.path.join(root, cfg['paths']['channels']))
        from lib import qa as Q
        qrep = R.load_json(os.path.join(root, cfg['paths']['qa_report']))
        gaps = R.load_gaps(root, cfg)
        excl = R.load_excluded_peers(root, cfg)
        made.append(R.build_report(cfg, leads, stats, os.path.join(root, '方法与边界说明.md'),
                                   channels=ch, qa=qrep, gaps_text=gaps, excluded_text=excl))
    for m in made:
        print('saved: %s' % m)


def cmd_all(a):
    cfg, root = project(a.dir)
    print('== 1/4 merge ==')
    leads, stats = M.merge(cfg, root)
    M.print_summary(stats)
    if not a.skip_qa and leads:
        print('\n== 2/4 链接抽检 ==')
        from lib import qa as Q
        Q.sample_and_check(cfg, leads, os.path.join(root, cfg['paths']['qa_report']),
                           timeout=a.timeout, workers=a.workers)
        print('\n== 3/4 重新 merge（把抽检结论并入）==')
        leads, stats = M.merge(cfg, root)
    print('\n== 4/4 出表 ==')
    cmd_build(a)


# ------------------------------------------------------------------ check

def cmd_check(a):
    cfg, root = project(a.dir)
    leads = load_leads(cfg, root)
    stem = C.render_stem(cfg)
    fails = []

    if not leads:
        fails.append('线索数为 0 —— 这不可能是有效交付（检查 leads_raw/ 是否有产出、'
                     '地区闸门与 peer 是否把所有条目都剔掉了）')
        print('!! 线索数为 0')
    elif len(leads) < 20:
        print('⚠ 只有 %d 条线索，明显偏少。请确认不是通道失效或地区/peer 闸门误剔。' % len(leads))

    xlsx = os.path.join(root, stem + '.xlsx')
    csv = os.path.join(root, stem + '.csv')
    if os.path.exists(xlsx):
        print('Excel: %s (%.1f KB)' % (xlsx, os.path.getsize(xlsx) / 1024))
        try:
            import openpyxl
            wb = openpyxl.load_workbook(xlsx)
            print('  sheets: %s' % wb.sheetnames)
            if len(wb.sheetnames) < 5:
                fails.append('Excel sheet 数偏少')
            ws = wb[wb.sheetnames[1]]
            print('  第二表 %s: %d 行 × %d 列' % (wb.sheetnames[1], ws.max_row, ws.max_column))
            if ws.max_row != len(leads) + 1:
                fails.append('Excel 行数(%d) 与线索数(%d) 不一致' % (ws.max_row - 1, len(leads)))
        except ImportError:
            print('  （本机无 openpyxl，跳过 Excel 深度校验）')
    elif os.path.exists(csv):
        print('CSV（降级）: %s' % csv)
    else:
        fails.append('缺少 Excel/CSV 交付物')

    html = os.path.join(root, stem + '.html')
    if os.path.exists(html):
        with open(html, encoding='utf-8') as f:
            h = f.read()
        print('HTML : %s (%.1f KB bytes / %d chars)'
              % (html, os.path.getsize(html) / 1024, len(h)))
        need = [('时间线', '时间线视图'), ('企业表', '企业表视图'), ('切入点', '切入点话术'),
                ('企业背景', '企业背景'), ('证据与来源', '证据区块'),
                cfg['product']['name'], cfg['region']['name']]
        for token, label in need:
            if token not in h:
                fails.append('HTML 缺 %s' % label)
        js = '\n;\n'.join(re.findall(r'(?is)<script[^>]*>(.*?)</script>', h))
        node = shutil.which('node')
        if node and js.strip():
            tf = os.path.join(root, '_check_tmp.js')
            with open(tf, 'w', encoding='utf-8') as f:
                f.write(js)
            try:
                r = subprocess.run([node, '--check', tf], capture_output=True, text=True)
                print('  内联 JS: %s' % ('语法 OK' if r.returncode == 0 else '语法错误'))
                if r.returncode != 0:
                    fails.append('HTML 内联 JS 语法错误：%s' % (r.stderr or '')[:200])
            finally:
                os.remove(tf)
        else:
            print('  （未找到 node，跳过 JS 语法校验）')
    else:
        fails.append('缺少 HTML 交付物')

    md = os.path.join(root, '方法与边界说明.md')
    if os.path.exists(md):
        with open(md, encoding='utf-8') as f:
            t = f.read()
        print('说明 : %s' % md)
        for token in ('条数与分布', '置信度分级', '被剔除的同行', '未覆盖的缺口'):
            if token not in t:
                fails.append('说明缺第四块的「%s」' % token)
        if 'GAPS:BEGIN' in t:
            print('  ⚠ 第四块的判断性内容（gaps.md）还是空的，交付前必须补写')
            fails.append('gaps.md 未填写（第四块不完整）')
        if 'PEERS:BEGIN' in t:
            print('  ⚠ 第三块的 lane 阶段 peer 名单（peers_excluded.md）还是空的')
            fails.append('peers_excluded.md 未填写（第三块不完整）')
        # 脚手架模板首行的 STATUS 标记还在 = 判断性内容没写
        for fn, block in (('gaps.md', '第四块'), ('peers_excluded.md', '第三块')):
            fp = os.path.join(root, fn)
            if os.path.exists(fp):
                with open(fp, encoding='utf-8-sig') as fh:
                    if STATUS_UNFILLED in fh.read():
                        print('  ⚠ %s 还是脚手架模板（STATUS 标记未删），%s 不完整' % (fn, block))
                        fails.append('%s 未填写（STATUS 标记还在，%s不完整）' % (fn, block))
    else:
        fails.append('缺少《方法与边界说明.md》')

    print()
    if fails:
        print('检查未通过：')
        for f_ in fails:
            print('  - %s' % f_)
        sys.exit(1)
    ok('交付物检查通过（%d 条线索）' % len(leads))


def cmd_stats(a):
    cfg, root = project(a.dir)
    with open(os.path.join(root, cfg['paths']['stats']), encoding='utf-8-sig') as f:
        s = json.load(f)
    print(json.dumps(s, ensure_ascii=False, indent=1))


# ------------------------------------------------------------------ argparse

def build_parser():
    p = argparse.ArgumentParser(
        prog='b2bctl',
        description='化工品下游「项目级潜客」流水线（产品无关，agent 无关）')
    sub = p.add_subparsers(dest='cmd', required=True)

    def add(name, fn, help_, with_dir=True):
        sp = sub.add_parser(name, help=help_)
        if with_dir:
            sp.add_argument('dir', nargs='?', default='.', help='项目目录（含 config.json）')
        sp.set_defaults(func=fn)
        return sp

    sp = add('new', cmd_new, '新建项目（定契约）', with_dir=False)
    sp.add_argument('dir', help='项目目录（不存在会创建）')
    sp.add_argument('--product', required=True, help='产品名，如 纯碱')
    sp.add_argument('--formula', default='', help='分子式，如 Na2CO3')
    sp.add_argument('--aliases', default='', help='别名，逗号分隔')
    sp.add_argument('--region', required=True, help='地区名，如 华东')
    sp.add_argument('--provinces', default='', help='省份，逗号分隔；空=不限')
    sp.add_argument('--months', type=int, default=6, help='未来月数（默认 6）')
    sp.add_argument('--today', default='', help='生成日 YYYY-MM-DD（默认今天）')
    sp.add_argument('--peers', default='', help='同行关键词，逗号分隔（peer gate）')
    sp.add_argument('--force', action='store_true', help='覆盖已有文件')

    sp = add('probe', cmd_probe, '探通道：实测本机能连上哪些数据源')
    sp.add_argument('--timeout', type=int, default=20)
    sp.add_argument('--channels', default=None, help='自定义通道注册表 JSON')

    sp = add('lanes-template', cmd_lanes_template, '生成 lanes.json 骨架（应用地图起步）')
    sp.add_argument('--product', default='')
    sp.add_argument('--lanes', default='lanes.json')

    add('lanes-check', cmd_lanes_check, '校验应用地图')

    sp = add('brief', cmd_brief, '生成每个 lane 的派工单 + workplan.md')
    sp.add_argument('--limit', type=int, default=0, help='只生成前 N 个（调试用）')

    sp = add('collect', cmd_collect, '校验 lane 产出是否符合 schema')
    sp.add_argument('--file', default=None, help='只校验某个文件/通配符')

    add('merge', cmd_merge, '合并去重 + peer gate + 销售时点推导')
    add('stats', cmd_stats, '打印 merge_stats.json')

    sp = add('qa', cmd_qa, '链接抽检（链接闸门）')
    sp.add_argument('--timeout', type=int, default=25)
    sp.add_argument('--workers', type=int, default=8)

    sp = add('build', cmd_build, '生成 Excel + HTML + 方法与边界说明')

    sp = add('all', cmd_all, 'merge + qa + build 一把跑完')
    sp.add_argument('--skip-qa', action='store_true')
    sp.add_argument('--timeout', type=int, default=25)
    sp.add_argument('--workers', type=int, default=8)

    add('check', cmd_check, '交付前自检（Excel/HTML/说明四块/JS 语法）')
    return p


def main():
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass
    args = build_parser().parse_args()
    args.func(args)


if __name__ == '__main__':
    main()