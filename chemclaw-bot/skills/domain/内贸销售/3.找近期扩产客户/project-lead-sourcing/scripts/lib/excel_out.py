# -*- coding: utf-8 -*-
"""Excel 交付物生成模块（产品无关）。

对外契约（docs/INTERFACE.md §5）：
    build_workbook(cfg, leads, out_path) -> str
    build_csv(cfg, leads, out_path) -> str

纪律（docs/INTERFACE.md §6）：
  1. 模块内不硬编码任何产品名 / 地区 / 日期：产品取 cfg['product']['name']，
     地区取 cfg['region']['name']，窗口月数取 cfg['window_months']，生成日取 cfg['today']。
  2. 派生字段（lane / lane_id / 用量强度 / 需求性质 / 销售时点 / 时点理由 / 优先级 / 抽检结论）
     一律 .get(k, 默认值) 读取——新项目在 b2bctl merge 之前可能还没有这些字段，缺字段不能崩。
     `用量强度` 是 `★★★★★` 字符串，只按 `★` 个数比较，绝不当作数字排序。
  3. 环境缺少 openpyxl 时 build_workbook 自动降级为 CSV，stderr 明确告知，并返回实际写入的
     CSV 路径（保证"至少有一份可打开的表格"）。
  4. 单元格里出现 URL 的一律用 openpyxl 挂超链接，整格就是一个 URL 时再设蓝色下划线。

六个 sheet 与源脚本（soda-ash-huadong-projects/build_workbook.py）保持一致：
说明与行业地图 / 项目线索总表 / 项目时间线 / 重点客户TOP / 分布统计 / 剔除与缺口。
"""
import collections
import csv
import datetime as _dt
import os
import re
import sys

try:  # openpyxl 是可选依赖：没有就降级 CSV
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    HAS_OPENPYXL = True
except Exception:  # pragma: no cover - 纯环境问题
    Workbook = None
    HAS_OPENPYXL = False

try:  # 同目录有 common.py（本 skill 的共享工具）时复用，保证口径一致；
    from . import common as _common  # 单独拷走本文件时也能跑（下面退化为自带实现）
except Exception:  # pragma: no cover
    try:
        import common as _common
    except Exception:
        _common = None

# ---------------------------------------------------------------- 小工具

_URL_RE = re.compile(r'https?://[^\s<>"\'，。；、（）()\[\]【】]+')


def _get(d, key, default=''):
    """安全取字段：None / 缺失一律回落默认值（缺字段不崩）。"""
    if not isinstance(d, dict):
        return default
    v = d.get(key, default)
    return default if v is None else v


def _text(d, key, default=''):
    return str(_get(d, key, default))


def _product(cfg):
    return _text(cfg.get('product') or {}, 'name', '本产品')


def _region(cfg):
    return _text(cfg.get('region') or {}, 'name', '未指定地区')


def _provinces(cfg):
    provs = _get(cfg.get('region') or {}, 'provinces', []) or []
    return ' / '.join(str(p) for p in provs) if provs else '未指定'


def _months(cfg):
    try:
        return int(_get(cfg, 'window_months', 6))
    except (TypeError, ValueError):
        return 6


def _today(cfg):
    """生成日：优先 cfg['today']，缺失时退回当天（仅作显示，不编造业务数据）。"""
    if _common is not None:
        try:
            return _common.today_date(cfg).isoformat()
        except Exception:
            pass
    return _text(cfg, 'today') or _dt.date.today().isoformat()


def _add_months(iso, n):
    """在 YYYY-MM-DD 上加 n 个月（不依赖第三方库），返回 YYYY-MM-DD。"""
    m = re.match(r'(\d{4})-(\d{1,2})(?:-(\d{1,2}))?', str(iso or ''))
    if not m:
        return str(iso or '')
    y, mo = int(m.group(1)), int(m.group(2))
    day = int(m.group(3) or 1)
    total = (y * 12 + (mo - 1)) + int(n)
    ny, nm = total // 12, total % 12 + 1
    last = [31, 29 if (ny % 4 == 0 and (ny % 100 != 0 or ny % 400 == 0)) else 28,
            31, 30, 31, 30, 31, 31, 30, 31, 30, 31][nm - 1]
    return f'{ny:04d}-{nm:02d}-{min(day, last):02d}'


def _window_start(cfg):
    return _today(cfg)


def _window_end(cfg):
    """窗口结束日：有 common.window_bounds 就用它（与 merge 口径一致），否则本地推算。"""
    if _common is not None:
        try:
            return str(_common.window_bounds(cfg)[3])
        except Exception:
            pass
    return _add_months(_today(cfg), _months(cfg))


def _month_label(iso):
    return str(iso or '')[:7]


def _stars(rec):
    """按 `★` 个数取星级（字符串比较，不做数字转换）。"""
    s = str(_get(rec, '用量强度', ''))
    return s.count('★')


def _first_url(text):
    m = _URL_RE.search(str(text or ''))
    return m.group(0).rstrip('.,;)]）】。，；') if m else ''


def _lane_of(rec):
    """返回 (lane 全名, lane_id)；缺字段时逐级兜底，绝不抛错。"""
    lane = _text(rec, 'lane')
    lane_id = _text(rec, 'lane_id')
    if not lane_id and lane:
        m = re.match(r'\s*(\d+)', lane)
        lane_id = m.group(1) if m else ''
    if not lane:
        lane = lane_id or _text(rec, '细分行业')   # merge 之前：用细分行业近似下游方向
    return lane or '未分类', lane_id


def _lane_short(lane, lane_id):
    """lane 名去掉前缀编号，得到板块名。"""
    s = str(lane or '')
    if lane_id and s.startswith(str(lane_id)):
        s = s[len(str(lane_id)):].strip(' -_.·')
    return re.sub(r'^\d+\s*', '', s).strip() or '未分类'


def _timing_bucket(rec):
    return _text(rec, '销售时点', '仅需关注') or '仅需关注'


def _confidence(rec):
    return _text(rec, '置信度', '低') or '低'


# ---------------------------------------------------------------- 输入归一

def _resolve_stats(cfg, leads, stats):
    """merge 统计：优先显式参数，其次 cfg['stats']；都没有就如实标注"未提供"。"""
    st = {}
    if isinstance(cfg.get('stats'), dict):
        st.update(cfg['stats'])
    if isinstance(stats, dict):
        st.update(stats)
    st['_provided'] = bool(st)
    st.setdefault('final', len(leads))
    return st


def _resolve_lanes(cfg, leads, lanes_doc, base_dir):
    """lane 地图：优先 lanes.json，其次用 leads 里的真实字段兜底汇总。

    返回 lane_id -> {'name','logic','star','mag','seeds','count'}，按 lane_id 排序输出。
    """
    doc = lanes_doc if isinstance(lanes_doc, dict) else {}
    if not doc:
        path = _text(cfg.get('paths') or {}, 'lanes')
        for cand in ([os.path.join(base_dir, path)] if path and base_dir else []) + \
                    ([path] if path else []):
            if cand and os.path.isfile(cand):
                try:
                    import json
                    with open(cand, encoding='utf-8') as f:
                        doc = json.load(f)
                    break
                except Exception:
                    doc = {}

    out = {}
    for lane in (doc.get('lanes') or []):
        if not isinstance(lane, dict):
            continue
        lid = str(_get(lane, 'id', '') or '')
        if not lid:
            continue
        stars = lane.get('intensity')
        star = ('★' * int(stars)) if isinstance(stars, int) and stars > 0 else ''
        seeds = lane.get('seeds') or []
        out[lid] = {'id': lid,
                    'name': str(_get(lane, 'name', '') or ''),
                    'logic': str(_get(lane, 'logic', '') or ''),
                    'star': star,
                    'mag': str(_get(lane, 'intensity_basis', '') or ''),
                    'seeds': len(seeds) if seeds else '',
                    'count': 0}

    # 用 leads 补全（并统计条数）：缺 lanes.json 时也能出地图，逻辑取自真实"需求点"
    for rec in leads:
        lane, lid = _lane_of(rec)
        name = _lane_short(lane, lid)
        key = str(lid) if lid else f'~{name}'
        slot = out.setdefault(key, {'id': lid, 'name': name, 'logic': '',
                                    'star': '', 'mag': '', 'seeds': '', 'count': 0})
        if not slot['name'] or slot['name'] == '未分类':
            slot['name'] = _lane_short(lane, lid)
        slot['count'] += 1
        if not slot['logic']:
            slot['logic'] = str(_get(rec, '需求点', '') or '')[:120]
        if _stars(rec) > slot['star'].count('★'):
            slot['star'] = '★' * _stars(rec)
    return out


def _channel_rows(cfg, channels, leads):
    """取证通道：有通道探测结果就照实列，没有就如实说明未提供（不编造环境结论）。"""
    ch = channels if channels is not None else cfg.get('channels')
    rows = []
    items = []
    if isinstance(ch, dict):
        items = ch.get('channels') or ch.get('items') or ch.get('results') or []
    elif isinstance(ch, list):
        items = ch
    for it in items if isinstance(items, list) else []:
        if isinstance(it, dict):
            lvl = str(it.get('level') or it.get('name') or it.get('通道') or '通道')
            note = str(it.get('note') or it.get('status') or it.get('说明') or it.get('url') or '')
            ok = it.get('ok', it.get('available'))
            if ok is False:
                lvl = f'{lvl}（不可用）'
            rows.append((lvl, note))
        else:
            rows.append(('通道', str(it)))
    if not rows:
        src_counter = collections.Counter(_text(r, '来源类型', '未披露') or '未披露' for r in leads)
        rows.append(('通道探测结果', '未提供（channels.json 缺失或未传入）；下表仅反映本批线索实际使用的来源类型'))
    src_counter = collections.Counter(_text(r, '来源类型', '未披露') or '未披露' for r in leads)
    rows.append(('本批线索来源类型分布',
                 '、'.join(f'{k} {v} 条' for k, v in src_counter.most_common()) or '无'))
    return rows


def _gap_rows(cfg, leads, stats):
    """缺口与负结果：cfg['gaps'] 优先，否则按真实数据统计（不编造环境结论）。"""
    rows = []
    for g in (cfg.get('gaps') or []):
        if isinstance(g, (list, tuple)) and len(g) >= 2:
            rows.append((str(g[0]), str(g[1])))
        elif isinstance(g, dict):
            rows.append((str(g.get('title') or g.get('项目') or '缺口'),
                         str(g.get('note') or g.get('内容') or '')))
        else:
            rows.append(('缺口', str(g)))
    low = sum(1 for r in leads if _confidence(r) == '低')
    no_est = sum(1 for r in leads if _text(r, '年需求量估算', '') in ('', '无法估算'))
    unverified = sum(1 for r in leads if _text(r, '来源', '') in ('', '未核实'))
    no_time = sum(1 for r in leads if _text(r, '预计投产时间', '') in ('', '未披露'))
    no_source_type = len(leads) - sum(
        1 for r in leads if _text(r, '来源类型', '') in (cfg.get('source_types_vocab') or []))
    rows.append(('低置信度项目', f'{low} 条，主体或项目细节未完全核实，需电话确认'))
    rows.append(('无法估算需求量的项目', f'{no_est} 条（单耗无出处时不估算，宁缺不编）'))
    rows.append(('无来源(未核实)条数', unverified))
    rows.append(('投产时间未披露', f'{no_time} 条，无法进入当期销售时点排程'))
    if no_source_type > 0:
        rows.append(('来源类型超出词表', f'{no_source_type} 条，需复核采集质量'))
    rows.append(('取证通道', '见"说明与行业地图"页的通道小节；未获得通道探测结果时不得声称通道可用'))
    rows.append(('建议人工复核', '反爬拦截类链接（百科/企业查询平台壳页）在浏览器中通常可打开，'
                                '本表标注为待复核而非坏链'))
    return rows


# ---------------------------------------------------------------- 样式

def _styles():
    return {
        'hdr_fill': PatternFill('solid', fgColor='1F4E79'),
        'hdr_font': Font(bold=True, color='FFFFFF', size=10),
        'title_font': Font(bold=True, size=14, color='1F4E79'),
        'sub_font': Font(bold=True, size=11, color='1F4E79'),
        'body_font': Font(size=9),
        'bold_small': Font(bold=True, size=10),
        'border': Border(left=Side(style='thin', color='BFBFBF'),
                         right=Side(style='thin', color='BFBFBF'),
                         top=Side(style='thin', color='BFBFBF'),
                         bottom=Side(style='thin', color='BFBFBF')),
        'wrap': Alignment(wrap_text=True, vertical='top'),
        'center': Alignment(horizontal='center', vertical='center', wrap_text=True),
        'link_font': Font(size=9, color='0563C1', underline='single'),
        'timing_fill': {
            '现在就该联系': PatternFill('solid', fgColor='C6EFCE'),
            '3-6个月后再联系': PatternFill('solid', fgColor='FFEB9C'),
            '仅需关注': PatternFill('solid', fgColor='F2F2F2'),
        },
        'conf_fill': {'高': PatternFill('solid', fgColor='C6EFCE'),
                      '中': PatternFill('solid', fgColor='FFEB9C'),
                      '低': PatternFill('solid', fgColor='FFC7CE')},
    }


def _style_header(ws, st, row=1, ncols=None):
    ncols = ncols or ws.max_column
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = st['hdr_fill']
        cell.font = st['hdr_font']
        cell.alignment = st['center']
        cell.border = st['border']
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _linkify(cell, st):
    """单元格里写了 URL 就挂超链接；整格就是一个 URL 时再设蓝色下划线。"""
    val = cell.value
    if not isinstance(val, str) or 'http' not in val:
        return
    url = _first_url(val)
    if not url:
        return
    cell.hyperlink = url
    if val.strip().rstrip('。，；') == url:
        cell.font = st['link_font']


# ---------------------------------------------------------------- 列定义

FIELD_MAP = ['销售时点', '优先级', 'lane', '公司', '细分行业', '省市', '项目名称', '建设阶段', '项目产品',
             '设计产能', '预计投产时间', '所在园区', '产品角色', '需求性质', '单耗依据', '年需求量估算', '需求点',
             '规模线索', '用量强度', '来源类型', '来源', '抽检结论', '发布日期', '置信度', '时点理由']

TIMELINE_FIELDS = ['预计投产时间', '销售时点', '建设阶段', '公司', '项目名称', '项目产品', '设计产能',
                   '省市', '所在园区', 'lane', '产品角色', '年需求量估算', '置信度', '来源类型']

TOP_FIELDS = ['公司', '项目名称', 'lane', '建设阶段', '销售时点', '预计投产时间', '设计产能',
              '产品角色', '年需求量估算', '省市', '置信度', '优先级', '来源']


def total_columns(product):
    """总表列标题与宽度（顺序、宽度与源脚本一致，仅产品名改为渲染）。"""
    return [
        ('销售时点', 14), ('优先级', 8), ('板块(lane)', 20), ('公司(业主企业)', 30), ('细分行业', 12),
        ('省市', 11), ('项目名称', 40), ('建设阶段', 10), ('项目产品', 22), ('设计产能', 24),
        ('预计投产时间', 13), ('所在园区', 20), ('产品角色', 9), ('需求性质', 9), ('单耗依据', 42),
        ('年需求量估算', 22),
        (f'需求点({product}在哪个工序起什么作用)', 52), ('规模线索', 40), ('用量强度', 8),
        ('来源类型', 9), ('来源', 52), ('来源抽检', 20), ('发布日期', 12), ('置信度', 8), ('时点理由', 46),
    ]


def timeline_columns(product):
    return ['预计投产月', '销售时点', '建设阶段', '公司', '项目名称', '项目产品', '设计产能',
            '省市', '所在园区', '板块', f'{product}角色', '预计需求量级', '置信度', '来源类型']


def top_columns(product):
    return ['排序', '公司', '项目名称', '板块', '建设阶段', '销售时点', '预计投产时间', '设计产能',
            f'{product}角色', '年需求量估算', '省市', '置信度', '优先级', '来源']


def _key_month(s):
    """把"预计投产时间"归一到可比月份键。"""
    txt = str(s or '')
    m = re.search(r'(20\d{2})\s*[-/年\.]\s*(\d{1,2})', txt)
    if m:
        return f'{int(m.group(1))}-{int(m.group(2)):02d}'
    m = re.search(r'(20\d{2})\s*[Qq]\s*([1-4])', txt)
    if m:
        return f'{int(m.group(1))}-Q{m.group(2)}'
    m = re.search(r'(20\d{2})', txt)
    if m:
        return f'{m.group(1)}-年内未明确'
    return '未披露'


def _mkey(k):
    if k == '未披露':
        return (9, 0, 0)
    m = re.match(r'(\d{4})-(\d{2})', k)
    if m:
        return (0, int(m.group(1)), int(m.group(2)))
    m = re.match(r'(\d{4})-Q(\d)', k)
    if m:
        return (1, int(m.group(1)), int(m.group(2)))
    m = re.match(r'(\d{4})', k)
    return (2, int(m.group(1)), 0) if m else (9, 0, 0)


# ---------------------------------------------------------------- 各 sheet 数据（渲染层，供 xlsx/csv 共用）

def sheet1_rows(cfg, leads, stats, lanes, channels):
    """说明与行业地图：行 = (第1列, 第2列, 类型标记)。"""
    product, region = _product(cfg), _region(cfg)
    months = _months(cfg)
    today, wend = _window_start(cfg), _window_end(cfg)
    mid = _add_months(today, 3)     # 与 common.derive_timing 一致：提前 3 个月的销售动作前置期
    stages = ' / '.join(str(x) for x in (cfg.get('stages_vocab') or [])) or '拟建 / 备案 / 在建 / 试车 / 已投产'
    roles = ' / '.join(str(x) for x in (cfg.get('product') or {}).get('roles_vocab') or []) or '主原料 / 辅料 / 助剂'
    grades = (cfg.get('product') or {}).get('grades') or []
    formula = _text(cfg.get('product') or {}, 'formula')
    product_line = f'{product}' + (f' {formula}' if formula else '') + \
        (f'（{" + ".join(str(g) for g in grades)}）' if grades else '')

    if stats.get('_provided'):
        n_lines = (f'{stats.get("final", len(leads))} 条项目级线索'
                   f'（原始 {stats.get("raw_records", len(leads))} 条，去重剔除 {stats.get("dups_removed", 0)} 条）')
    else:
        n_lines = f'{len(leads)} 条项目级线索（本次未传入 merge 统计，去重前后条数未标注）'

    rows = [
        (f'{product} · {region} · 未来 {months} 个月 —— 项目级潜在客户清单', '', 'title'),
        ('', '', ''),
        ('产品', product_line, ''),
        ('地区', f'{region}：{_provinces(cfg)}', ''),
        ('时间窗', f'{today} ~ {wend}（共 {months} 个月；生成日 {today}）', ''),
        ('检索主题', f'在上述地区未来 {months} 个月内新建、扩建、技改或即将投产，'
                     f'并因此新增{product}需求的项目', ''),
        ('线索条数', n_lines, ''),
        ('', '', ''),
        ('—— 怎么读这张表 ——', '', 'sub'),
        ('一条记录 = 一个具体项目，不是一家公司', '', ''),
        (f'建设阶段只取：{stages}', '', ''),
        (f'产品角色只取：{roles}', '', ''),
        ('单耗依据：能估算的给出"单耗数字 + 依据"；缺数据的写"无法估算"，不给假精度', '', ''),
        ('销售时点由"建设阶段 + 预计投产时间"按固定规则统一推导，规则见下', '', ''),
        ('优先级/用量强度/需求性质/抽检结论 均由 merge 阶段派生；新采集项目可能尚未写入，'
         f'本表按"{product}在哪个工序起什么作用"的需求点原文判断', '', ''),
        ('', '', ''),
        ('—— 销售时点分组规则（自动推导，可复算）——', '', 'sub'),
        ('现在就该联系',
         f'试车/试生产阶段；或在建且 {_month_label(mid)} 前投产；或近半年内已投产仍在爬坡。'
         f'理由：装置点火前 1~3 个月即需锁定原料并完成合格供应商入库，采购决策正在发生', ''),
        ('3-6个月后再联系',
         f'在建且预计 {_month_label(_add_months(mid, 1))} ~ {_month_label(wend)} 投产；'
         f'或已备案且投产落在窗口内。理由：采购启动点在未来 3~6 个月，现在留资/送样/技术对接正当其时', ''),
        ('仅需关注',
         f'拟建阶段；或投产时间超出 {_month_label(wend)}；或投产时间未披露。理由：尚未形成当期需求，仅作商机储备', ''),
        ('', '', ''),
        ('—— 本次可用取证通道 ——', '', 'sub'),
    ]
    for lvl, note in _channel_rows(cfg, channels, leads):
        rows.append((lvl, note, ''))
    rows += [
        ('', '', ''),
        (f'—— {product}需求逻辑与用量强度（lane 地图）——', '', 'sub'),
    ]
    rows.append(('lane', '细分板块', f'用料逻辑（{product}在哪个工序、起什么作用）', '用量强度',
                 '量级参考', '候选企业数'))
    for lid in sorted(lanes, key=lambda x: (len(str(x)), str(x))):
        info = lanes[lid]
        rows.append((info.get('id') or '', info['name'], info['logic'],
                     info['star'] or f'{info["count"]} 条线索（强度未提供）',
                     info['mag'] or '未提供',
                     info['seeds'] if info['seeds'] != '' else '未提供'))
    rows.append(('', '', ''))
    rows.append(('—— 重点客户TOP 取数规则 ——', '', 'sub'))
    n_high = sum(1 for r in leads if _text(r, '优先级') == '高')
    rows.append(('TOP 口径', f'优先级 = 高（本项目 {n_high} 条）；'
                             f'若全部暂无优先级派生字段，则按用量强度★数与置信度取前 {_get(cfg, "top_n", 30)} 条', ''))
    return rows


def sheet5_blocks(leads):
    """分布统计：(标题, Counter) 列表，字段缺失不崩。"""
    return [
        ('按销售时点', collections.Counter(_timing_bucket(r) for r in leads)),
        ('按建设阶段', collections.Counter(_text(r, '建设阶段', '未披露') or '未披露' for r in leads)),
        ('按置信度', collections.Counter(_confidence(r) for r in leads)),
        ('按板块', collections.Counter(_lane_of(r)[0] for r in leads)),
        ('按省份', collections.Counter((_text(r, '省市', '/') or '/').split('/')[0] for r in leads)),
        ('按来源类型', collections.Counter(_text(r, '来源类型', '未披露') or '未披露' for r in leads)),
        ('按产品角色', collections.Counter(_text(r, '产品角色', '未披露') or '未披露' for r in leads)),
        ('按用量强度', collections.Counter(_text(r, '用量强度', '未提供') or '未提供' for r in leads)),
    ]


def sheet6_rows(cfg, leads, stats):
    """剔除与缺口：(项目, 内容) 列表。"""
    product = _product(cfg)
    peers = stats.get('peers_removed')
    if peers is None:
        peers_txt = '未提供（未传入 merge 统计，无法列出被剔除的同行名单）'
    else:
        peers_txt = '、'.join(str(p) for p in peers) or '无'
    qa = stats.get('qa')
    rows = [
        ('剔除的同行(peer)名单', peers_txt),
        ('剔除原因', f'命中 peer_keywords 的企业自身生产{product}，是同行而非客户，按 peer gate 剔除'),
        ('无来源(未核实)条数', stats.get('unverified_source',
                                          sum(1 for r in leads if _text(r, '来源', '') in ('', '未核实')))),
        ('原始条数', stats.get('raw_records', '未提供')),
        ('去重剔除条数', stats.get('dups_removed', '未提供')),
        ('最终条数', stats.get('final', len(leads))),
        ('', ''),
        ('—— 覆盖缺口与负结果（必须如实告知）——', ''),
    ]
    if isinstance(qa, dict) and qa:
        checked = _get(qa, 'checked', _get(qa, 'sampled', _get(qa, 'n', '未记录')))
        bad = _get(qa, 'bad', _get(qa, 'broken', '未记录'))
        rows.insert(6, ('链接抽检', f'抽样 {checked} 条，坏链 {bad} 条（详见 leads_raw/_qa_report.json）'))
        rows.insert(7, ('', ''))
    rows += _gap_rows(cfg, leads, stats)
    return rows


# ---------------------------------------------------------------- 主流程（xlsx）

def render_out_stem(cfg):
    """按 cfg['paths']['out_stem'] 模板渲染文件名主干（{product}/{region}/{months}）。

    模块内不拼死任何文件名；调用方直接 build_workbook(cfg, leads, render_out_stem(cfg) + '.xlsx')
    或传 out_path=None 由本模块用当前目录 + 该模板兜底。
    """
    if _common is not None:
        try:
            return _common.render_stem(cfg)      # 与 html_out/report_out 同一套命名口径
        except Exception:
            pass
    tpl = _text(cfg.get('paths') or {}, 'out_stem') or \
        '{product}{region}未来{months}个月项目级潜客表'
    for k, v in (('{product}', _product(cfg)), ('{region}', _region(cfg)),
                 ('{months}', str(_months(cfg)))):
        tpl = tpl.replace(k, v)
    return re.sub(r'[\\/:*?"<>|\r\n\t]+', '_', tpl).strip()   # 文件名安全化


def _default_out_path(cfg, ext):
    """out_path 未给时的兜底路径：当前目录 + 渲染后的 out_stem + 扩展名。"""
    stem = render_out_stem(cfg)
    return os.path.join(os.getcwd(), stem if stem.lower().endswith(ext) else stem + ext)


def _build_xlsx(cfg, leads, out_path, stats, lanes_doc, channels):
    st = _styles()
    product = _product(cfg)
    cols = total_columns(product)
    wb = Workbook()

    # ---------------- Sheet 1: 说明与行业地图 ----------------
    ws = wb.active
    ws.title = '说明与行业地图'
    lanes = _resolve_lanes(cfg, leads, lanes_doc, os.path.dirname(os.path.abspath(out_path)))
    rows = sheet1_rows(cfg, leads, stats, lanes, channels)
    r = 1
    for item in rows:
        first = item[0] if len(item) > 0 else ''
        second = item[1] if len(item) > 1 else ''
        mark = item[2] if len(item) > 2 else ''
        if mark == 'title':
            ws.cell(row=r, column=1, value=first).font = st['title_font']
        elif mark == 'sub':
            ws.cell(row=r, column=1, value=first).font = st['sub_font']
        elif len(item) >= 6 and item[0] == 'lane':
            for c, v in enumerate(item[:6], start=1):
                cell = ws.cell(row=r, column=c, value=v)
                cell.fill = st['hdr_fill']
                cell.font = st['hdr_font']
                cell.alignment = st['center']
        elif len(item) >= 6:
            for c, v in enumerate(item[:6], start=1):
                ws.cell(row=r, column=c, value=v).alignment = st['wrap']
        else:
            ws.cell(row=r, column=1, value=first).font = st['bold_small']
            if second:
                ws.cell(row=r, column=2, value=second).alignment = st['wrap']
        r += 1
    _set_widths(ws, [24, 30, 62, 10, 20, 12])
    ws.column_dimensions['B'].width = 90

    # ---------------- Sheet 2: 项目线索总表 ----------------
    ws2 = wb.create_sheet('项目线索总表')
    ws2.append([c[0] for c in cols])
    for rec in leads:
        ws2.append([_get(rec, f, '') for f in FIELD_MAP])
    _style_header(ws2, st)
    _set_widths(ws2, [c[1] for c in cols])
    for i, rec in enumerate(leads, start=2):
        for c in range(1, len(cols) + 1):
            cell = ws2.cell(row=i, column=c)
            cell.alignment = st['wrap']
            cell.border = st['border']
            cell.font = st['body_font']
        ws2.cell(row=i, column=1).fill = st['timing_fill'].get(_timing_bucket(rec),
                                                               st['timing_fill']['仅需关注'])
        ws2.cell(row=i, column=1).alignment = st['center']
        ws2.cell(row=i, column=24).fill = st['conf_fill'].get(_confidence(rec), st['conf_fill']['低'])
        ws2.cell(row=i, column=24).alignment = st['center']
        for c in (2, 8, 12, 19, 20, 23):
            ws2.cell(row=i, column=c).alignment = st['center']
        for c in range(1, len(cols) + 1):     # 单元格里的 URL 一律可点
            _linkify(ws2.cell(row=i, column=c), st)
    ws2.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{ws2.max_row}'
    ws2.row_dimensions[1].height = 28

    # ---------------- Sheet 3: 项目时间线 ----------------
    ws3 = wb.create_sheet('项目时间线')
    ws3.append(timeline_columns(product))
    mons = {}
    for rec in leads:
        mons.setdefault(_key_month(_text(rec, '预计投产时间')), []).append(rec)
    for k in sorted(mons, key=_mkey):
        for rec in mons[k]:
            ws3.append([k] + [_get(rec, f, '') for f in TIMELINE_FIELDS[1:]])
    _style_header(ws3, st)
    _set_widths(ws3, [16, 16, 10, 30, 42, 22, 24, 12, 20, 20, 10, 22, 8, 10])
    ws3.auto_filter.ref = f'A1:N{ws3.max_row}'
    for i in range(2, ws3.max_row + 1):
        ws3.cell(row=i, column=2).fill = st['timing_fill'].get(
            ws3.cell(row=i, column=2).value, st['timing_fill']['仅需关注'])
        for c in range(1, 15):
            ws3.cell(row=i, column=c).alignment = st['wrap']
            ws3.cell(row=i, column=c).font = st['body_font']
            _linkify(ws3.cell(row=i, column=c), st)

    # ---------------- Sheet 4: 重点客户TOP ----------------
    ws4 = wb.create_sheet('重点客户TOP')
    top = [r for r in leads if _text(r, '优先级') == '高']
    fallback = False
    if not top:
        fallback = True
        limit = int(_get(cfg, 'top_n', 30) or 30)
        top = sorted(leads, key=lambda r: (-_stars(r), _confidence(r)))[:limit] if leads else []
    ws4.append(top_columns(product))
    for i, rec in enumerate(top, start=1):
        ws4.append([i] + [_get(rec, f, '') for f in TOP_FIELDS])
    _style_header(ws4, st)
    _set_widths(ws4, [6, 30, 42, 20, 10, 16, 13, 24, 10, 22, 12, 8, 8, 52])
    for i in range(2, ws4.max_row + 1):
        for c in range(1, 15):
            ws4.cell(row=i, column=c).alignment = st['wrap']
            ws4.cell(row=i, column=c).font = st['body_font']
            _linkify(ws4.cell(row=i, column=c), st)
    if fallback:
        ws4.cell(row=ws4.max_row + 2, column=1,
                 value='注：本批线索尚无"优先级=高"派生字段，上表按用量强度★数与置信度取前若干条，'
                       '仅作排序参考，不得当作高优先级结论').font = st['bold_small']

    # ---------------- Sheet 5: 分布统计 ----------------
    ws5 = wb.create_sheet('分布统计')
    rr = 1
    for title, counter in sheet5_blocks(leads):
        ws5.cell(row=rr, column=1, value=title).font = st['sub_font']
        ws5.cell(row=rr + 1, column=1, value='类别').fill = st['hdr_fill']
        ws5.cell(row=rr + 1, column=1).font = st['hdr_font']
        ws5.cell(row=rr + 1, column=2, value='条数').fill = st['hdr_fill']
        ws5.cell(row=rr + 1, column=2).font = st['hdr_font']
        r2 = rr + 2
        for k, v in counter.most_common():
            ws5.cell(row=r2, column=1, value=k).alignment = st['wrap']
            ws5.cell(row=r2, column=2, value=v)
            r2 += 1
        rr = r2 + 1
    _set_widths(ws5, [34, 10])

    # ---------------- Sheet 6: 剔除与缺口 ----------------
    ws6 = wb.create_sheet('剔除与缺口')
    ws6.append(['项目', '内容'])
    for item in sheet6_rows(cfg, leads, stats):
        ws6.append([item[0], item[1]])
    _style_header(ws6, st)
    _set_widths(ws6, [26, 110])
    for i in range(2, ws6.max_row + 1):
        ws6.cell(row=i, column=1).font = st['bold_small']
        ws6.cell(row=i, column=2).alignment = st['wrap']
        ws6.cell(row=i, column=1).alignment = st['wrap']
        _linkify(ws6.cell(row=i, column=2), st)

    wb.save(out_path)
    return out_path


# ---------------------------------------------------------------- 降级：CSV

def _sheet1_csv_rows(cfg, leads, stats, lanes, channels):
    """说明页转 CSV：普通行取两列，lane 地图行保留六列。"""
    return [list(row[:6]) if len(row) >= 6 else list(row[:2])
            for row in sheet1_rows(cfg, leads, stats, lanes, channels)]


def build_csv(cfg, leads, out_path):
    """CSV 兜底：主文件是"项目线索总表"，其余 sheet 各写一个同名 sidecar CSV。

    用 utf-8-sig 写出，Excel 直接双击不会乱码。返回主 CSV 的实际路径。
    """
    if not out_path:                      # 不拼死文件名：走 cfg['paths']['out_stem'] 模板
        out_path = _default_out_path(cfg, '.csv')
    base = os.path.splitext(out_path)[0]
    main = out_path if str(out_path).lower().endswith('.csv') else f'{base}.csv'
    leads = [r for r in (leads or []) if isinstance(r, dict)]
    d = os.path.dirname(os.path.abspath(main))
    if d and not os.path.isdir(d):
        os.makedirs(d, exist_ok=True)

    product = _product(cfg)
    stats = _resolve_stats(cfg, leads, stats=None)
    lanes = _resolve_lanes(cfg, leads, None, d)

    def _write(path, rows):
        with open(path, 'w', encoding='utf-8-sig', newline='') as f:
            w = csv.writer(f)
            for row in rows:
                w.writerow(row)
        return path

    cols = total_columns(product)
    table = [[c[0] for c in cols]] + [[_get(rec, f, '') for f in FIELD_MAP] for rec in leads]
    _write(main, table)
    _write(f'{base}.说明与行业地图.csv', _sheet1_csv_rows(cfg, leads, stats, lanes, None))

    tl = [timeline_columns(product)]
    mons = {}
    for rec in leads:
        mons.setdefault(_key_month(_text(rec, '预计投产时间')), []).append(rec)
    for k in sorted(mons, key=_mkey):
        for rec in mons[k]:
            tl.append([k] + [_get(rec, f, '') for f in TIMELINE_FIELDS[1:]])
    _write(f'{base}.项目时间线.csv', tl)

    top = [r for r in leads if _text(r, '优先级') == '高']
    if not top:
        top = sorted(leads, key=lambda r: (-_stars(r), _confidence(r)))[:int(_get(cfg, 'top_n', 30) or 30)]
    _write(f'{base}.重点客户TOP.csv',
           [top_columns(product)] + [[i] + [_get(r, f, '') for f in TOP_FIELDS]
                                     for i, r in enumerate(top, start=1)])

    dist = []
    for title, counter in sheet5_blocks(leads):
        dist.append([title, '条数'])
        dist += [[k, v] for k, v in counter.most_common()]
        dist.append(['', ''])
    _write(f'{base}.分布统计.csv', dist)
    _write(f'{base}.剔除与缺口.csv', [['项目', '内容']] + [list(r) for r in sheet6_rows(cfg, leads, stats)])
    print(f'[excel_out] CSV 兜底已写出：{main}（另含 5 个同名 sidecar CSV）', file=sys.stderr)
    return main


# ---------------------------------------------------------------- 对外入口

def build_workbook(cfg: dict, leads: list, out_path: str, stats=None, lanes=None,
                   channels=None, qa=None) -> str:
    """生成 Excel 交付物；返回实际写入的路径。

    cfg / leads 是唯一数据来源；stats / lanes / channels / qa 为可选补充
    （b2bctl 已在内存里有这些结果时可直接传入，缺省则从 cfg 读取或如实标注"未提供"）。
    环境没有 openpyxl 时自动降级为 CSV，并在 stderr 打印明确提示。
    """
    leads = [r for r in (leads or []) if isinstance(r, dict)]
    if not out_path:                      # 不拼死文件名：走 cfg['paths']['out_stem'] 模板
        out_path = _default_out_path(cfg, '.xlsx')
    if not HAS_OPENPYXL:
        csv_path = os.path.splitext(out_path)[0] + '.csv'
        print('[excel_out] 未检测到 openpyxl，无法生成 .xlsx；已自动降级为 CSV：'
              f'{csv_path}（Excel 可直接打开，utf-8-sig 编码）', file=sys.stderr)
        return build_csv(cfg, leads, csv_path)
    if isinstance(cfg, dict) and isinstance(cfg.get('lanes'), list):
        lanes = lanes if lanes is not None else {'lanes': cfg['lanes']}
    if channels is None and isinstance(cfg, dict) and cfg.get('channels'):
        channels = cfg['channels']
    st = _resolve_stats(cfg, leads, stats)
    if qa is not None and str(qa):      # 抽检报告只影响说明，不改变总表结构
        st.setdefault('qa', qa)
    return _build_xlsx(cfg, leads, out_path, st, lanes, channels)