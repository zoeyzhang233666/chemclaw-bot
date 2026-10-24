#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
customer-pool-screening 交付物生成器
用法:
    python build_deliverables.py --input pool.json --outdir out --basename 纯碱_华东_客户池
输入: 按 references/data-schema.md 契约编写的 pool.json
输出: {basename}.xlsx (6表) + {basename}.html (自包含交互报告)
依赖: Python 3.8+, openpyxl (pip install openpyxl)
"""
import argparse
import html
import json
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("[x] 缺少 openpyxl，请先: pip install openpyxl")

# ---------------- helpers ----------------

def fail(msg):
    print("[x] " + msg)
    sys.exit(1)


def norm(v):
    return "" if v is None else str(v)


def as_list(v):
    if v is None:
        return []
    return v if isinstance(v, list) else [v]


HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(bold=True, color="FFFFFF", size=10)
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Border(*[Side(style="thin", color="BFBFBF")] * 4)
FILL_P = {"P0": PatternFill("solid", fgColor="C6EFCE"),
          "P1": PatternFill("solid", fgColor="FFF2CC"),
          "P2": PatternFill("solid", fgColor="E7E6E6")}
FILL_EXCL = PatternFill("solid", fgColor="FCE4E4")

POOL_COLS = [
    ("priority", "优先级", 7), ("score", "综合得分", 9),
    ("raw_name", "原始录入名称", 14), ("std_name", "标准化企业名称", 30),
    ("uscc", "统一社会信用代码", 30), ("status", "经营状态", 12),
    ("region_full", "注册地(省/市)", 14), ("entity_type", "企业类型", 14),
    ("products", "实际主营产品", 34), ("demand_relation", "需求关系", 10),
    ("demand_logic", "需求逻辑", 40),
    ("s_match", "产品匹配", 9), ("s_factory", "工厂真实性", 9),
    ("s_scale", "规模", 8), ("s_region", "地区", 8),
    ("evidence", "证据来源(节选)", 44), ("confidence", "置信度", 6),
    ("dedup", "排重结果", 9), ("next_action", "下一步动作建议", 34),
]
NUM_COLS = {"score", "s_match", "s_factory", "s_scale", "s_region"}

# ---------------- load & validate ----------------

def load_pool(path):
    try:
        data = json.loads(Path(path).read_text(encoding="utf-8"))
    except Exception as e:
        fail("读取/解析 %s 失败: %s" % (path, e))
    meta = data.get("meta")
    if not isinstance(meta, dict):
        fail("pool.json 缺少 meta 对象（产品/区域/评分口径）")
    pool = data.get("pool")
    if not isinstance(pool, list) or not pool:
        fail("pool.json 缺少非空 pool 数组")

    weights = (meta.get("scoring") or {}).get("weights") or {
        "产品匹配": 40, "工厂真实性": 25, "规模": 20, "地区": 15}
    tiers = (meta.get("scoring") or {}).get("tiers") or {"P0": 80, "P1": 60, "P2": 40}

    for i, r in enumerate(pool, 1):
        for key in ("std_name", "entity_type", "demand_relation"):
            if not norm(r.get(key)):
                fail("pool 第 %d 行缺必填字段 %s" % (i, key))
        parts = [r.get(k) for k in ("s_match", "s_factory", "s_scale", "s_region")]
        if all(isinstance(p, (int, float)) for p in parts):
            total = int(sum(parts))
            if "score" in r and int(r["score"]) != total:
                print("[!] 第 %d 行 score(%s) 与分项之和(%d)不一致，已按分项重算"
                      % (i, r.get("score"), total))
            r["score"] = total
        elif not isinstance(r.get("score"), (int, float)):
            fail("pool 第 %d 行缺 score 且分项得分不全" % i)
        if not norm(r.get("priority")):
            p = "P2"
            if r["score"] >= tiers.get("P0", 80):
                p = "P0"
            elif r["score"] >= tiers.get("P1", 60):
                p = "P1"
            r["priority"] = p
        if norm(r.get("province")) == "":
            reg = norm(r.get("region_full"))
            r["province"] = reg[:2] if reg else "未知"
        r["evidence"] = "；".join(as_list(r.get("evidence")))
        if norm(r.get("uscc")) == "":
            r["uscc"] = "未核实"
        if norm(r.get("status")) == "":
            r["status"] = "未核实"
        if norm(r.get("dedup")) == "":
            r["dedup"] = "无"
    pool.sort(key=lambda r: (-r["score"], norm(r["std_name"])))
    data["_weights"], data["_tiers"] = weights, tiers
    return data


# ---------------- xlsx ----------------

def style_header(ws):
    for c in ws[1]:
        c.font = HDR_FONT
        c.fill = HDR_FILL
        c.alignment = WRAP
        c.border = THIN


def style_body(ws, min_row=2):
    for row in ws.iter_rows(min_row=min_row):
        for c in row:
            c.alignment = WRAP
            c.border = THIN


def sheet_kv(ws, pairs, w1=22, w2=110):
    ws.append(["项目", "内容"])
    for r in pairs:
        ws.append(list(r))
    ws.column_dimensions["A"].width = w1
    ws.column_dimensions["B"].width = w2
    style_header(ws)
    style_body(ws)
    for row in ws.iter_rows(min_row=2):
        row[0].font = Font(bold=True, size=10)


def build_xlsx(data, out_path):
    meta = data["meta"]
    weights, tiers = data["_weights"], data["_tiers"]
    wb = Workbook()

    ws = wb.active
    ws.title = "说明与方法"
    pairs = [
        ("任务", "%s · %s 客户名单筛选（偏好：%s）" % (
            meta.get("product", "?"), meta.get("region", "?"), meta.get("preference", "真实工厂优先"))),
        ("输入", "；".join(as_list(meta.get("input_files")))),
        ("流程", "①名称标准化+工商主体匹配 → ②三类分箱(成功/歧义/未找到) → 用户确认歧义主体 → ③生产/贸易判定与主营核实 → ④产品需求匹配分级(强/弱/不相关) → ⑤四维评分排序 → ⑥排除清单(一票否决) → ⑦与现有客户池排重"),
        ("评分口径", "综合分 = " + " + ".join("%s(0-%d)" % (k, v) for k, v in weights.items())
         + "；分级线 P0≥%d / P1≥%d / P2≥%d。需求『不相关』或『同行自产』一票否决进排除清单。"
         % (tiers.get("P0", 80), tiers.get("P1", 60), tiers.get("P2", 40))),
        ("需求逻辑表", "；".join("%s(%s·%s)" % (d.get("行业", "?"), d.get("强度", "?"), d.get("作用", "?"))
                          for d in as_list(meta.get("demand_logic")))),
        ("数据边界", norm(meta.get("notes")) or
         "全部信息来自公开检索，非官方登记系统直查；『未核实』或单一来源字段，签约前以国家企业信用信息公示系统或营业执照终核；需求强度为工艺逻辑推断。"),
        ("生成日期", norm(meta.get("generated"))),
    ]
    sheet_kv(ws, pairs)

    ws2 = wb.create_sheet("重点客户池(排序后)")
    ws2.append([label for _, label, _ in POOL_COLS])
    for r in data["pool"]:
        ws2.append([r.get(k, "") for k, _, _ in POOL_COLS])
    for i, (_, _, w) in enumerate(POOL_COLS, 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    style_header(ws2)
    style_body(ws2)
    for row in ws2.iter_rows(min_row=2):
        key = POOL_COLS[0][0]
        row[0].fill = FILL_P.get(row[0].value, FILL_P["P2"])
        row[0].font = Font(bold=True)
        for idx, (k, _, _) in enumerate(POOL_COLS):
            if k in NUM_COLS:
                row[idx].number_format = "0"
    ws2.freeze_panes = "C2"
    ws2.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(POOL_COLS)), ws2.max_row)

    def simple_sheet(name, headers, rows, widths, mark_col=None):
        s = wb.create_sheet(name)
        s.append(headers)
        for r in rows:
            s.append([norm(x) for x in r])
        for i, w in enumerate(widths, 1):
            s.column_dimensions[get_column_letter(i)].width = w
        style_header(s)
        style_body(s)
        if mark_col is not None:
            for row in s.iter_rows(min_row=2):
                row[mark_col].fill = FILL_EXCL
        return s

    simple_sheet("名称歧义-确认记录",
                 ["原始录入名称", "候选主体及辨识要点", "用户确认结果", "处理结果"],
                 [[a.get("raw_name"), a.get("candidates"), a.get("user_decision"), a.get("outcome")]
                  for a in as_list(data.get("ambiguity"))], [14, 80, 26, 36])

    simple_sheet("未找到",
                 ["原始录入名称", "名单线索", "检索结论", "去向建议/相近名线索"],
                 [[n.get("raw_name"), n.get("hints"), n.get("conclusion"), n.get("next_suggestion")]
                  for n in as_list(data.get("not_found"))], [20, 20, 60, 60])

    simple_sheet("排除企业及原因",
                 ["原始录入名称", "标准化企业名称", "排除类别", "排除原因(含依据)", "关联证据"],
                 [[e.get("raw_name"), e.get("std_name"), e.get("category"), e.get("reason"), e.get("evidence")]
                  for e in as_list(data.get("excluded"))], [16, 34, 20, 62, 52], mark_col=2)

    simple_sheet("现有客户池(输入)",
                 ["客户名称", "客户等级", "合作状态", "备注"],
                 [[c.get("name"), c.get("grade"), c.get("status"), c.get("note")]
                  for c in as_list(data.get("existing_pool"))], [34, 10, 12, 22])

    wb.save(out_path)


# ---------------- html ----------------

HTML_TEMPLATE = r"""<!doctype html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>__TITLE__</title>
<style>
:root{--ink:#1a1a2e;--muted:#6b7280;--line:#e5e7eb;--brand:#1f4e79;}
*{box-sizing:border-box}
body{margin:0;font-family:"Microsoft YaHei",system-ui,sans-serif;color:var(--ink);background:#f6f7f9}
header{background:var(--brand);color:#fff;padding:14px 20px}
header h1{margin:0;font-size:18px}
header .sub{opacity:.85;font-size:12px;margin-top:4px}
.chips{display:flex;gap:8px;flex-wrap:wrap;padding:10px 20px;background:#fff;border-bottom:1px solid var(--line)}
.chip{font-size:12px;padding:3px 10px;border-radius:999px;background:#eef2f7}
.chip.b{background:#dcfce7}.chip.y{background:#fef9c3}.chip.g{background:#e5e7eb}.chip.r{background:#fee2e2}
.tabs{display:flex;gap:4px;padding:10px 20px 0}
.tabs button{border:1px solid var(--line);border-bottom:none;background:#fff;padding:8px 16px;border-radius:8px 8px 0 0;cursor:pointer;font-size:13px}
.tabs button.on{background:var(--brand);color:#fff;font-weight:bold}
.panel{display:none;background:#fff;margin:0 20px 20px;border:1px solid var(--line);border-radius:0 8px 8px 8px;padding:14px}
.panel.on{display:block}
.toolbar{display:flex;gap:8px;flex-wrap:wrap;align-items:center;margin-bottom:10px}
.toolbar input,.toolbar select{padding:6px 8px;border:1px solid var(--line);border-radius:6px;font-size:13px}
.toolbar input{min-width:200px}
table{border-collapse:collapse;width:100%;font-size:12px}
th,td{border:1px solid var(--line);padding:6px 8px;vertical-align:top;text-align:left}
th{background:#f0f4f8;cursor:pointer;user-select:none;white-space:nowrap;position:sticky;top:0}
th.num,td.num{text-align:right}
tbody tr:hover{background:#f8fafc}
.p0{background:#dcfce7!important;font-weight:bold}.p1{background:#fef9c3!important}.p2{background:#e5e7eb!important}
.wrap{max-width:260px;word-break:break-all}
.selbar{position:sticky;bottom:0;background:#111827;color:#fff;padding:8px 14px;border-radius:8px;display:flex;gap:10px;align-items:center;margin-top:10px;font-size:13px}
.selbar button{background:#2563eb;color:#fff;border:none;border-radius:6px;padding:6px 12px;cursor:pointer;font-size:12px}
.selbar button.ghost{background:#374151}
.count{font-size:12px;color:var(--muted);margin:6px 0}
h2{font-size:15px;margin:2px 0 10px}
dl.meta{display:grid;grid-template-columns:120px 1fr;gap:6px 10px;font-size:13px;margin:0}
dl.meta dt{font-weight:bold;color:var(--brand)}
.note{font-size:12px;color:var(--muted);line-height:1.7}
</style>
</head>
<body>
<header>
  <h1 id="title"></h1>
  <div class="sub" id="subtitle"></div>
</header>
<div class="chips" id="chips"></div>
<div class="tabs" id="tabs"></div>
<div id="panels"></div>
<script>
const DATA = __DATA_JSON__;
let state = {q:"",priority:"",province:"",entity_type:"",demand_relation:"",confidence:"",sortKey:"score",sortDir:-1,tab:"pool"};
const sel = new Set();
function esc(s){return String(s==null?"":s).replace(/&/g,"&amp;").replace(/</g,"&lt;").replace(/>/g,"&gt;").replace(/\"/g,"&quot;");}
function norm(v){return v==null?"":String(v);}
function uniq(key){const s=new Set();DATA.pool.forEach(r=>{const v=norm(r[key]);if(v)s.add(v);});return Array.from(s).sort();}
function filtered(){
  let rows=DATA.pool.map((r,i)=>Object.assign({_i:i},r));
  const q=state.q.trim().toLowerCase();
  if(q)rows=rows.filter(r=>JSON.stringify(r).toLowerCase().includes(q));
  ["priority","province","entity_type","demand_relation","confidence"].forEach(k=>{
    if(state[k])rows=rows.filter(r=>norm(r[k])===state[k]);
  });
  const key=state.sortKey;
  rows.sort((a,b)=>{const va=a[key],vb=b[key];
    if(typeof va==="number"&&typeof vb==="number")return (va-vb)*state.sortDir;
    return norm(va).localeCompare(norm(vb),"zh")*state.sortDir;});
  return rows;
}
const POOL_COLS=[
 {k:"sel"},{k:"priority",l:"优先级"},{k:"score",l:"得分",num:1},{k:"std_name",l:"标准化企业名称"},
 {k:"raw_name",l:"原始录入"},{k:"uscc",l:"统一社会信用代码",cls:"wrap"},{k:"status",l:"经营状态"},
 {k:"region_full",l:"注册地"},{k:"entity_type",l:"企业类型"},{k:"products",l:"实际主营产品",cls:"wrap"},
 {k:"demand_relation",l:"需求关系"},{k:"demand_logic",l:"需求逻辑",cls:"wrap"},
 {k:"s_match",l:"匹配",num:1},{k:"s_factory",l:"真实",num:1},{k:"s_scale",l:"规模",num:1},{k:"s_region",l:"地区",num:1},
 {k:"confidence",l:"置信度"},{k:"dedup",l:"排重"},{k:"next_action",l:"下一步动作",cls:"wrap"},{k:"evidence",l:"证据来源",cls:"wrap"}];
function renderPool(){
  const rows=filtered();
  let h="<div class='toolbar'>"+
   "<input id='q' placeholder='搜索企业/主营/逻辑…' value=\""+esc(state.q)+"\">"+
   ["priority","province","entity_type","demand_relation","confidence"].map(k=>{
     const lbl={priority:"优先级",province:"省份",entity_type:"企业类型",demand_relation:"需求关系",confidence:"置信度"}[k];
     const vals=k==="priority"?Object.keys(DATA.meta.scoring?DATA.meta.scoring.tiers||{}:{})||[]:uniq(k);
     return "<select id='f_"+k+"'><option value=''>"+lbl+":全部</option>"+vals.map(v=>"<option"+(state[k]===v?" selected":"")+">"+esc(v)+"</option>").join("")+"</select>";
   }).join("")+
   "<button class='ghost' onclick='resetFilters()'>重置</button>"+
   "<span class='count'>显示 "+rows.length+" / "+DATA.pool.length+" 家 · 已选 "+sel.size+" 家</span></div>";
  h+="<div style='overflow:auto;max-height:62vh'><table><thead><tr>"+
   POOL_COLS.map(c=>{
     if(c.k==="sel")return "<th style='width:34px'><input type='checkbox' onclick='toggleAll(this)' id='ckAll'></th>";
     const active=state.sortKey===c.k;
     return "<th class='"+(c.num?"num":"")+"' onclick='sortBy(\""+c.k+"\")'>"+c.l+(active?(state.sortDir<0?" ▾":" ▴"):"")+"</th>";
   }).join("")+"</tr></thead><tbody>";
  rows.forEach(r=>{
    h+="<tr>";
    h+="<td><input type='checkbox' "+(sel.has(r._i)?"checked":"")+" onclick='toggleSel("+r._i+")'></td>";
    POOL_COLS.forEach(c=>{
      if(c.k==="sel")return;
      let v=norm(r[c.k]);
      let cls=(c.num?"num ":"")+(c.cls||"")+"";
      if(c.k==="priority")cls+=" "+norm(r.priority).toLowerCase();
      h+="<td class='"+cls+"'>"+esc(v)+"</td>";
    });
    h+="</tr>";
  });
  h+="</tbody></table></div>";
  h+="<div class='selbar'><span id='selCount'>已选 "+sel.size+" 家</span>"+
   "<button onclick='exportCSV(true)'>导出选中 CSV</button>"+
   "<button class='ghost' onclick='exportCSV(false)'>导出全部 CSV</button>"+
   "<button class='ghost' onclick='copyNames()'>复制所选名单</button></div>";
  document.getElementById("poolPanel").innerHTML=h;
  const q=document.getElementById("q");
  q.addEventListener("input",e=>{state.q=e.target.value;renderPool();document.getElementById("q").focus();});
  ["priority","province","entity_type","demand_relation","confidence"].forEach(k=>{
    const el=document.getElementById("f_"+k);
    if(el)el.addEventListener("change",e=>{state[k]=e.target.value;renderPool();});
  });
}
function resetFilters(){state={q:"",priority:"",province:"",entity_type:"",demand_relation:"",confidence:"",sortKey:"score",sortDir:-1,tab:"pool"};renderPool();}
function sortBy(k){if(state.sortKey===k)state.sortDir*=-1;else{state.sortKey=k;state.sortDir=1;}renderPool();}
function toggleSel(i){sel.has(i)?sel.delete(i):sel.add(i);renderPool();}
function toggleAll(ck){const rows=filtered();rows.forEach(r=>ck.checked?sel.add(r._i):sel.delete(r._i));renderPool();}
function csvEsc(v){v=norm(v);return /[",\n]/.test(v)?'"'+v.replace(/"/g,'""')+'"':v;}
function download(name,text){const b=new Blob(["\uFEFF"+text],{type:"text/csv;charset=utf-8"});const a=document.createElement("a");a.href=URL.createObjectURL(b);a.download=name;a.click();}
function exportCSV(onlySel){
  const rows=filtered().filter(r=>!onlySel||sel.has(r._i));
  if(!rows.length){alert("没有可导出的行");return;}
  const heads=POOL_COLS.filter(c=>c.k!=="sel").map(c=>c.l);
  const lines=[heads.join(",")].concat(rows.map(r=>POOL_COLS.filter(c=>c.k!=="sel").map(c=>csvEsc(r[c.k])).join(",")));
  download("客户池选中.csv",lines.join("\n"));
}
function copyNames(){
  const rows=DATA.pool.filter((r,i)=>sel.has(i));
  const t=(rows.length?rows:DATA.pool).map(r=>norm(r.std_name)).join("\n");
  navigator.clipboard.writeText(t).then(()=>alert("已复制 "+(rows.length||DATA.pool.length)+" 家企业名称"));
}
function simpleTable(headers,rows){
  return "<div style='overflow:auto'><table><thead><tr>"+headers.map(h=>"<th>"+esc(h)+"</th>").join("")+"</tr></thead><tbody>"+
   rows.map(r=>"<tr>"+r.map(c=>"<td class='wrap'>"+esc(c)+"</td>").join("")+"</tr>").join("")+"</tbody></table></div>";
}
function renderStatic(){
  document.getElementById("exclPanel").innerHTML=simpleTable(
    ["原始录入名称","标准化企业名称","排除类别","排除原因(含依据)","关联证据"],
    DATA.excluded.map(e=>[e.raw_name,e.std_name,e.category,e.reason,e.evidence]));
  document.getElementById("ambPanel").innerHTML=simpleTable(
    ["原始录入名称","候选主体及辨识要点","用户确认结果","处理结果"],
    DATA.ambiguity.map(a=>[a.raw_name,a.candidates,a.user_decision,a.outcome]));
  document.getElementById("nfPanel").innerHTML=simpleTable(
    ["原始录入名称","名单线索","检索结论","去向建议/相近名"],
    DATA.not_found.map(n=>[n.raw_name,n.hints,n.conclusion,n.next_suggestion]));
  let d="<dl class='meta'><dt>产品</dt><dd>"+esc(DATA.meta.product)+"</dd><dt>区域</dt><dd>"+esc(DATA.meta.region)+
   "</dd><dt>偏好</dt><dd>"+esc(DATA.meta.preference||"")+"</dd><dt>生成日期</dt><dd>"+esc(DATA.meta.generated||"")+
   "</dd><dt>评分权重</dt><dd>"+Object.entries(DATA.meta.scoring.weights).map(x=>x[0]+" 0-"+x[1]).join(" + ")+"</dd></dl>";
  d+="<h2 style='margin-top:14px'>需求逻辑表</h2>"+simpleTable(
    ["下游行业","使用工序","作用","强度","判定要点"],
    DATA.meta.demand_logic.map(x=>[x["行业"],x["工序"],x["作用"],x["强度"],x["要点"]]));
  d+="<p class='note'>"+esc(DATA.meta.notes||"")+"</p>";
  document.getElementById("aboutPanel").innerHTML=d;
}
const TABS=[["pool","客户池","poolPanel"],["excluded","排除企业","exclPanel"],["amb","歧义确认","ambPanel"],["nf","未找到","nfPanel"],["about","说明","aboutPanel"]];
function setTab(t){state.tab=t;TABS.forEach(x=>{document.getElementById(x[2]).classList.toggle("on",x[0]===t);
  document.getElementById("tab_"+x[0]).classList.toggle("on",x[0]===t);});}
function init(){
  document.getElementById("title").textContent=DATA.meta.product+" · "+DATA.meta.region+" 客户池";
  document.getElementById("subtitle").textContent="生成于 "+(DATA.meta.generated||"")+" · 偏好: "+(DATA.meta.preference||"真实工厂优先")+
   " · 评分: "+Object.entries(DATA.meta.scoring.weights).map(x=>x[0]+x[1]).join("/");
  const c=DATA.pool;const cnt=p=>c.filter(r=>r.priority===p).length;
  document.getElementById("chips").innerHTML=
   "<span class='chip'>池内 "+c.length+" 家</span><span class='chip b'>P0 "+cnt("P0")+"</span>"+
   "<span class='chip y'>P1 "+cnt("P1")+"</span><span class='chip g'>P2 "+cnt("P2")+"</span>"+
   "<span class='chip r'>排除 "+DATA.excluded.length+"</span><span class='chip'>未找到 "+DATA.not_found.length+"</span>";
  document.getElementById("tabs").innerHTML=TABS.map(x=>"<button id='tab_"+x[0]+"' onclick='setTab(\""+x[0]+"\")'>"+x[1]+"</button>").join("");
  document.getElementById("panels").innerHTML=TABS.map(x=>"<div id='"+x[2]+"' class='panel'></div>").join("");
  renderStatic();setTab("pool");renderPool();
}
init();
</script>
</body>
</html>
"""


def build_html(data, out_path, title):
    payload = json.dumps({
        "meta": {
            "product": data["meta"].get("product", ""),
            "region": data["meta"].get("region", ""),
            "preference": data["meta"].get("preference", ""),
            "generated": data["meta"].get("generated", ""),
            "notes": data["meta"].get("notes", ""),
            "scoring": {"weights": data["_weights"], "tiers": data["_tiers"]},
            "demand_logic": as_list(data["meta"].get("demand_logic")),
        },
        "pool": [{k: r.get(k, "") for k in
                  ("priority", "score", "raw_name", "std_name", "uscc", "status",
                   "province", "region_full", "entity_type", "products",
                   "demand_relation", "demand_logic", "s_match", "s_factory",
                   "s_scale", "s_region", "evidence", "confidence", "dedup",
                   "next_action")} for r in data["pool"]],
        "excluded": as_list(data.get("excluded")),
        "ambiguity": as_list(data.get("ambiguity")),
        "not_found": as_list(data.get("not_found")),
    }, ensure_ascii=False).replace("</", "<\\/")

    out_path.write_text(
        HTML_TEMPLATE.replace("__TITLE__", html.escape(title))
                     .replace("__DATA_JSON__", payload),
        encoding="utf-8")


# ---------------- main ----------------

def main():
    ap = argparse.ArgumentParser(description="客户池筛选交付物生成器 (xlsx + html)")
    ap.add_argument("--input", required=True, help="pool.json 路径")
    ap.add_argument("--outdir", default="out", help="输出目录")
    ap.add_argument("--basename", default="客户池筛选结果", help="输出文件名(不含扩展名)")
    a = ap.parse_args()

    data = load_pool(a.input)
    outdir = Path(a.outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    title = "%s·%s 客户池筛选结果" % (data["meta"].get("product", ""), data["meta"].get("region", ""))

    xlsx_path = outdir / (a.basename + ".xlsx")
    html_path = outdir / (a.basename + ".html")
    build_xlsx(data, xlsx_path)
    build_html(data, html_path, title)

    pool = data["pool"]
    dist = {}
    for r in pool:
        dist[r["priority"]] = dist.get(r["priority"], 0) + 1
    print("[ok] %s" % xlsx_path)
    print("[ok] %s" % html_path)
    print("[i] 池内 %d 家: %s | 排除 %d | 歧义 %d | 未找到 %d | 现有池 %d" % (
        len(pool), dist, len(as_list(data.get("excluded"))),
        len(as_list(data.get("ambiguity"))), len(as_list(data.get("not_found"))),
        len(as_list(data.get("existing_pool")))))


if __name__ == "__main__":
    main()
