#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""产出体检：HTML 渲染与 Excel 结构。

用法: python qa.py data.json out/产品-区域区域市场地图.html out/产品-区域客户分布与走访优先级.xlsx

无 Chrome 时只做静态检查（仍能查出占位符残留、数据未注入、表数不对等问题）；
有 Chrome 时用 headless 打开页面、dump DOM，真实统计渲染出的表格行数与 canvas 数。
"""
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile

CHROME_CANDIDATES = [
    r'C:\Program Files\Google\Chrome\Application\chrome.exe',
    r'C:\Program Files (x86)\Google\Chrome\Application\chrome.exe',
    r'C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe',
    '/Applications/Google Chrome.app/Contents/MacOS/Google Chrome',
    '/usr/bin/google-chrome', '/usr/bin/chromium', '/usr/bin/chromium-browser',
]


def find_chrome():
    for p in CHROME_CANDIDATES:
        if os.path.exists(p):
            return p
    for n in ('chrome', 'google-chrome', 'chromium'):
        p = shutil.which(n)
        if p:
            return p
    return None


def dump_dom(chrome, html_path):
    url = 'file:///' + os.path.abspath(html_path).replace('\\', '/')
    tmp = tempfile.mkdtemp(prefix='crqa_')
    cmd = [chrome, '--headless=new', '--no-sandbox', '--disable-gpu', '--hide-scrollbars',
           '--virtual-time-budget=12000', '--window-size=1560,1400',
           '--user-data-dir=' + tmp, '--dump-dom', url]
    try:
        out = subprocess.run(cmd, capture_output=True, timeout=90).stdout
        return out.decode('utf-8', 'ignore')
    except Exception as e:
        print('  [skip] 无法运行 headless 浏览器：%s' % e)
        return None


def check_html(path, D):
    print('\n== HTML ==')
    ok = True
    if not os.path.exists(path):
        print('  [FAIL] 文件不存在')
        return False
    src = open(path, encoding='utf-8', errors='ignore').read()
    print('  大小 %.2f MB' % (len(src.encode('utf-8')) / 1048576.0))
    for ph in ('%%DATA%%', '%%GEO%%', '__ECHARTS__', '__TITLE__'):
        if ph in src:
            print('  [FAIL] 占位符未替换：%s' % ph)
            ok = False
    if 'window.__DATA__=null' in src or 'window.__DATA__={}' in src:
        print('  [FAIL] 数据对象为空')
        ok = False
    if 'echarts.registerMap' not in src and 'registerMap' not in src:
        print('  [FAIL] 缺少 registerMap（地图不会渲染）')
        ok = False
    for must in (D['meta']['product'].split('（')[0], D['meta']['region'].split('（')[0]):
        if must not in src:
            print('  [WARN] 页面里找不到「%s」' % must)

    chrome = find_chrome()
    if not chrome:
        print('  [skip] 未找到 Chrome，跳过渲染检查')
        return ok
    dom = dump_dom(chrome, path)
    if not dom:
        return ok
    body = re.sub(r'<script[\s\S]*?</script>', '', dom)
    n_city = len(re.findall(r'<tr[^>]*onclick="__selectCity\(', body))
    n_co = len(re.findall(r'<tr[^>]*onclick="__openCompany\(', body))
    n_pj = len(re.findall(r'<tr[^>]*onclick="__openProject\(', body))
    n_canvas = len(re.findall(r'<canvas', dom))
    n_lane = len(re.findall(r'class="lane"', body))
    n_bar = len(re.findall(r'class="bar"', body))
    print('  渲染：城市行 %d / 企业行 %d / 项目行 %d / 下游卡 %d / 数据条 %d / canvas %d'
          % (n_city, n_co, n_pj, n_lane, n_bar, n_canvas))
    if n_city != len(D['cities']):
        print('  [FAIL] 城市行数 %d ≠ 数据 %d' % (n_city, len(D['cities'])))
        ok = False
    if n_co != len(D['companies']):
        print('  [FAIL] 企业行数 %d ≠ 数据 %d' % (n_co, len(D['companies'])))
        ok = False
    if n_pj != len(D['projects']):
        print('  [FAIL] 项目行数 %d ≠ 数据 %d' % (n_pj, len(D['projects'])))
        ok = False
    if n_lane != len(D['lanes']):
        print('  [FAIL] 下游卡片 %d ≠ 数据 %d' % (n_lane, len(D['lanes'])))
        ok = False
    if n_canvas < 4:
        print('  [FAIL] canvas 只有 %d 个（应有地图+3 图）' % n_canvas)
        ok = False
    if n_bar < 1:
        print('  [FAIL] 侧栏数据条未渲染（.bar .fill 需 display:block）')
        ok = False
    for bad in ('NaN', 'undefined<div'):
        if bad in body:
            print('  [FAIL] 渲染结果含 %s' % bad)
            ok = False
    return ok


def check_xlsx(path, D):
    print('\n== Excel ==')
    if not os.path.exists(path):
        print('  [FAIL] 文件不存在')
        return False
    try:
        from openpyxl import load_workbook
    except ImportError:
        print('  [skip] 未安装 openpyxl')
        return True
    wb = load_workbook(path)
    ok = True
    want = 12
    if len(wb.sheetnames) != want:
        print('  [FAIL] 表数 %d ≠ %d' % (len(wb.sheetnames), want))
        ok = False
    print('  表：', ' / '.join(wb.sheetnames))
    expect_city = len(D['cities'])
    expect_co = len(D['companies'])
    ws = wb['03 城市明细']
    rows = sum(1 for r in ws.iter_rows(min_row=5, max_col=2) if r[1].value)
    print('  03 城市明细 数据行 %d' % rows)
    if rows != expect_city:
        print('  [FAIL] 城市行 %d ≠ %d' % (rows, expect_city))
        ok = False
    ws = wb['05 客户线索']
    rows = sum(1 for r in ws.iter_rows(min_row=5, max_col=1) if r[0].value)
    print('  05 客户线索 数据行 %d' % rows)
    if rows != expect_co:
        print('  [FAIL] 客户行 %d ≠ %d' % (rows, expect_co))
        ok = False
    # 合计校验
    ws = wb['02 省级汇总']
    total = None
    for r in ws.iter_rows(min_row=5, max_col=3):
        if r[0].value == '合计':
            total = r[2].value
    print('  02 合计生产企业 %s（数据 %s）' % (total, D['summary']['firms']))
    if total != D['summary']['firms']:
        print('  [FAIL] 省级合计与企业总数不一致')
        ok = False
    # 内容健康度
    longest = 0
    for name in wb.sheetnames:
        for row in wb[name].iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    longest = max(longest, len(c.value))
    print('  最长单元格 %d 字' % longest)
    if longest > 320:
        print('  [WARN] 有超长单元格，Excel 里会很难读，考虑拆到备注区')
    return ok


def main():
    if len(sys.argv) < 4:
        sys.exit(__doc__)
    data, html, xlsx = sys.argv[1:4]
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    import common
    D = common.load(data)
    a = check_html(html, D)
    b = check_xlsx(xlsx, D)
    print('\n结论：%s' % ('通过' if (a and b) else '有问题，见上'))
    sys.exit(0 if (a and b) else 1)


if __name__ == '__main__':
    main()
