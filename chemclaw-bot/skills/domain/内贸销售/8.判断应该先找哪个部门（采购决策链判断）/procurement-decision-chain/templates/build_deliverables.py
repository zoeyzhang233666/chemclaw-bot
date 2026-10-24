#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_deliverables.py — 采购决策链交付物生成器（Excel + HTML）

用法:
    python build_deliverables.py chain.json [输出目录]

输入:
    chain.json —— 按 templates/chain-data.schema.json 填写的结构化分析结果

输出（写入输出目录，默认与 chain.json 同目录）:
    采购决策链-{公司}-{产品}.xlsx   多工作表（需要 openpyxl: pip install openpyxl）
    采购决策链-{公司}-{产品}.html   自包含交互页（部门卡点击展开、按职责筛选），无第三方依赖

无 Python 环境的降级路径（人工或其他工具产出时保持同等结构）:
    xlsx 工作表: 总览 / 企业画像 / 判断逻辑 / 角色链与部门卡 / 卡点 / 场景变体 / 证据与推断 / 联系方式 / 来源
    html 区块: 页首上下文(企业/产品/场景) / 第一入口 / 角色链 / 部门卡(details 交互) /
               判断逻辑 / 企业画像 / 卡点 / 场景变体 / 证据与推断 / 来源
"""
import html
import json
import os
import re
import sys
from datetime import date

DUTIES = ["发起", "评估", "否决", "拍板", "执行"]


def esc(x):
    return html.escape(str(x if x is not None else ""))


def jget(d, *keys, default=""):
    cur = d
    for k in keys:
        if not isinstance(cur, dict):
            return default
        cur = cur.get(k)
        if cur is None:
            return default
    return cur


def joinlist(v, sep="\n"):
    if isinstance(v, list):
        return sep.join(str(x) for x in v if x is not None)
    return str(v or "")


def safe_name(s):
    return re.sub(r'[\\/:*?"<>|\r\n\t]+', "-", str(s)).strip() or "采购决策链"


# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------
def build_xlsx(data, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
    HEAD_FONT = Font(bold=True, color="FFFFFF", size=11)
    THIN = Side(style="thin", color="BFBFBF")
    BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    WRAP = Alignment(vertical="top", wrap_text=True)
    CENTER = Alignment(vertical="center", horizontal="center", wrap_text=True)

    def style_header(ws, ncol):
        for c in range(1, ncol + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill = HEAD_FILL
            cell.font = HEAD_FONT
            cell.alignment = CENTER
            cell.border = BORDER

    def style_body(ws):
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = WRAP
                cell.border = BORDER

    def set_widths(ws, widths):
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    def new_sheet(wb, title, headers, widths):
        ws = wb.create_sheet(title)
        ws.append(headers)
        style_header(ws, len(headers))
        set_widths(ws, widths)
        ws.freeze_panes = "A2"
        return ws

    wb = Workbook()
    meta = data.get("meta", {})

    # --- 总览 ---
    ws = wb.active
    ws.title = "总览"
    ws.append(["项目", "内容"])
    style_header(ws, 2)
    entry = data.get("entry", {})
    rows = [
        ("产品", meta.get("product", "")),
        ("目标企业", meta.get("company", "")),
        ("组织类型", meta.get("company_type", "")),
        ("行业", meta.get("industry", "")),
        ("用途/工序", meta.get("use", "")),
        ("切入场景", meta.get("scenario", "")),
        ("分析日期", meta.get("date", str(date.today()))),
        ("备注", meta.get("note", "")),
        ("第一入口", entry.get("first", "")),
        ("理由", entry.get("why", "")),
        ("次选路径", entry.get("second_path", "")),
    ]
    for k, v in rows:
        ws.append([k, v])
    style_body(ws)
    set_widths(ws, [14, 110])
    ws.freeze_panes = "A2"

    # --- 企业画像 ---
    prof = data.get("company_profile") or []
    ws = new_sheet(wb, "企业画像", ["维度", "事实", "来源", "来源标题", "置信"],
                   [16, 70, 50, 28, 10])
    for p in prof:
        ws.append([p.get("dimension", ""), p.get("fact", ""), p.get("source", ""),
                   p.get("source_title", ""), p.get("confidence", "")])
    style_body(ws)

    # --- 判断逻辑 ---
    logic = data.get("logic_chain") or []
    ws = new_sheet(wb, "判断逻辑", ["步骤", "从（证据/前提）", "推到（结论）", "依据"],
                   [8, 55, 55, 24])
    for l in logic:
        ws.append([l.get("step", ""), l.get("from", ""), l.get("to", ""), l.get("basis", "")])
    style_body(ws)

    # --- 角色链与部门卡 ---
    roles = data.get("roles") or []
    ws = new_sheet(wb, "角色链与部门卡",
                   ["编号", "部门", "岗位", "职责", "第一入口", "最关心的问题", "先聊什么",
                    "最值得确认的问题", "依据", "公开联系方式"],
                   [7, 24, 20, 16, 10, 42, 44, 46, 28, 34])
    for r in roles:
        contacts = joinlist(["{}：{}".format(c.get("channel", ""), c.get("value", ""))
                             for c in (r.get("public_contacts") or [])])
        ws.append([r.get("id", ""), r.get("dept", ""), r.get("title", ""),
                   joinlist(r.get("duties"), "、"),
                   "★" if r.get("is_entry") else "",
                   joinlist(r.get("concerns")), r.get("opener", ""),
                   joinlist(r.get("questions")), r.get("basis", ""), contacts])
    style_body(ws)

    # --- 卡点 ---
    ws = new_sheet(wb, "卡点", ["卡点", "卡在谁", "性质", "破解/窗口"], [40, 20, 14, 52])
    for b in data.get("blocking_points") or []:
        ws.append([b.get("point", ""), b.get("owner", ""), b.get("nature", ""),
                   b.get("workaround", "")])
    style_body(ws)

    # --- 场景变体 ---
    ws = new_sheet(wb, "场景变体", ["客户类型", "决策链变化", "第一入口"], [26, 66, 30])
    for v in data.get("variants") or []:
        ws.append([v.get("customer_type", ""), v.get("change", ""), v.get("first_entry", "")])
    style_body(ws)

    # --- 证据与推断 ---
    ws = new_sheet(wb, "证据与推断", ["类型", "条目"], [10, 120])
    for x in data.get("evidence_list") or []:
        ws.append(["证据", x])
    for x in data.get("inference_list") or []:
        ws.append(["推断", x])
    style_body(ws)

    # --- 联系方式 ---
    ws = new_sheet(wb, "联系方式", ["角色", "渠道", "方式/地址"], [26, 16, 70])
    for r in roles:
        for c in (r.get("public_contacts") or []):
            ws.append([r.get("dept", ""), c.get("channel", ""), c.get("value", "")])
    ws.append(["说明", "", "仅收录企业公开渠道；不提供私人联系方式。"])
    style_body(ws)

    # --- 来源 ---
    ws = new_sheet(wb, "来源", ["标题", "链接", "说明"], [40, 70, 40])
    for s in data.get("sources") or []:
        ws.append([s.get("title", ""), s.get("url", ""), s.get("note", "")])
    style_body(ws)

    wb.save(path)


# --------------------------------------------------------------------------
# HTML
# --------------------------------------------------------------------------
CSS = """
:root{--ink:#1f2937;--line:#e5e7eb;--brand:#1f4e79}
*{box-sizing:border-box}
body{font-family:"Microsoft YaHei",system-ui,-apple-system,"Segoe UI",sans-serif;color:var(--ink);margin:0;background:#f6f7f9}
header{background:linear-gradient(135deg,#1f4e79,#2d6da3);color:#fff;padding:26px 32px}
header h1{margin:0 0 10px;font-size:23px}
.chips span{display:inline-block;background:rgba(255,255,255,.14);border:1px solid rgba(255,255,255,.35);border-radius:14px;padding:3px 12px;margin:3px 6px 0 0;font-size:13px}
main{max-width:1180px;margin:0 auto;padding:18px 24px 40px}
section{background:#fff;border:1px solid var(--line);border-radius:10px;padding:18px 20px;margin-top:16px}
h2{font-size:17px;margin:0 0 12px;border-left:4px solid var(--brand);padding-left:10px}
.entry-box{background:#fffbe6;border:1px solid #f1d888;border-radius:8px;padding:12px 16px;line-height:1.7}
table{width:100%;border-collapse:collapse;font-size:13.5px;margin-top:6px}
th{background:#eef3f8;text-align:left}
th,td{border:1px solid var(--line);padding:7px 9px;vertical-align:top}
tr:nth-child(even) td{background:#fafbfc}
.duty-row{display:flex;align-items:flex-start;gap:8px;margin:6px 0}
.duty{flex:0 0 52px;color:#fff;border-radius:6px;text-align:center;padding:3px 0;font-size:13px}
.d-发起{background:#2e7d32}.d-评估{background:#1c6dd0}.d-否决{background:#b02a37}.d-拍板{background:#6f42c1}.d-执行{background:#5f6b7a}
.chip{border:1px solid var(--line);background:#f2f6fb;border-radius:6px;padding:3px 9px;margin:2px 4px 2px 0;font-size:13px;display:inline-block}
.fbar{margin:2px 0 12px}
.fbtn{border:1px solid var(--line);background:#fff;border-radius:16px;padding:4px 14px;margin-right:6px;cursor:pointer;font-size:13px}
.fbtn.on{background:var(--brand);color:#fff;border-color:var(--brand)}
.cards{display:grid;grid-template-columns:repeat(auto-fill,minmax(360px,1fr));gap:12px}
details.card{border:1px solid var(--line);border-radius:10px;background:#fff;overflow:hidden}
details.card summary{cursor:pointer;list-style:none;padding:12px 14px;background:#f2f6fb;font-weight:600;font-size:14.5px;line-height:1.8}
details.card summary::-webkit-details-marker{display:none}
.bd{display:inline-block;border-radius:10px;padding:1px 8px;margin-left:6px;font-size:12px;background:#e3e8ee;color:#334155;font-weight:400}
.bd.entry{background:#b8860b;color:#fff}
.bd.发起{background:#2e7d32;color:#fff}.bd.评估{background:#1c6dd0;color:#fff}.bd.否决{background:#b02a37;color:#fff}.bd.拍板{background:#6f42c1;color:#fff}.bd.执行{background:#5f6b7a;color:#fff}
.card-body{padding:12px 14px;font-size:13.5px;line-height:1.65}
.card-body h4{margin:10px 0 4px;font-size:13px;color:var(--brand)}
.card-body p{margin:0}
.card-body ul{margin:4px 0;padding-left:18px}
.basis{margin-top:10px;font-size:12.5px;color:#6b7280;border-top:1px dashed var(--line);padding-top:8px}
ol.srcs li{margin:5px 0;font-size:13.5px}
footer{color:#6b7280;font-size:12.5px;text-align:center;margin:24px 0 10px}
"""

JS = """
function flt(k, btn) {
  var bs = document.querySelectorAll('.fbtn');
  for (var i = 0; i < bs.length; i++) { bs[i].classList.remove('on'); }
  btn.classList.add('on');
  var cs = document.querySelectorAll('details.card');
  for (var j = 0; j < cs.length; j++) {
    var c = cs[j], show;
    if (k === 'all') { show = true; }
    else if (k === 'entry') { show = c.getAttribute('data-entry') === '1'; }
    else { show = (c.getAttribute('data-duties') || '').indexOf(k) >= 0; }
    c.style.display = show ? '' : 'none';
  }
}
"""


def src_cell(p):
    s = p.get("source") or ""
    if not s:
        return ""
    label = p.get("source_title") or s
    return "<a href='" + esc(s) + "' target='_blank' rel='noopener'>" + esc(label) + "</a>"


def build_html(data, path):
    meta = data.get("meta", {})
    short = meta.get("short_name") or meta.get("company", "")
    roles = data.get("roles") or []
    entry = data.get("entry", {})
    out = []
    out.append("<!DOCTYPE html><html lang='zh-CN'><head><meta charset='utf-8'>")
    out.append("<meta name='viewport' content='width=device-width, initial-scale=1'>")
    out.append("<title>" + esc(short) + "·" + esc(meta.get("product", "")) + "·采购决策链</title>")
    out.append("<style>" + CSS + "</style></head><body>")

    # 页首：企业/产品/场景上下文（上层平台可在此继承客户上下文）
    out.append("<header><h1>采购决策链 · " + esc(meta.get("product", "")) + " → " + esc(short) + "</h1><div class='chips'>")
    for k in ["product", "company", "company_type", "industry", "use", "scenario"]:
        v = meta.get(k)
        if v:
            out.append("<span><b>" + esc(k) + "</b>：" + esc(v) + "</span>")
    out.append("</div></header><main>")

    # 第一入口
    if entry:
        out.append("<section><h2>★ 第一入口</h2><div class='entry-box'>")
        out.append("<p><b>" + esc(entry.get("first", "")) + "</b></p>")
        out.append("<p>" + esc(entry.get("why", "")) + "</p>")
        if entry.get("second_path"):
            out.append("<p><b>次选路径：</b>" + esc(entry.get("second_path")) + "</p>")
        out.append("</div></section>")

    # 角色链
    out.append("<section><h2>角色链（发起 → 评估 → 否决 → 拍板 → 执行）</h2>")
    for duty in DUTIES:
        rs = [r for r in roles if duty in (r.get("duties") or [])]
        if not rs:
            continue
        out.append("<div class='duty-row'><span class='duty d-" + duty + "'>" + duty + "</span><div>")
        for r in rs:
            star = " ★" if r.get("is_entry") else ""
            name = r.get("dept", "") + (("·" + r["title"]) if r.get("title") else "")
            out.append("<span class='chip'>" + esc(name) + star + "</span>")
        out.append("</div></div>")
    out.append("</section>")

    # 部门卡
    out.append("<section><h2>部门卡（点击展开：先聊什么 / 最值得确认什么）</h2><div class='fbar'>")
    buttons = [("all", "全部")] + [(d, d) for d in DUTIES] + [("entry", "★第一入口")]
    for k, label in buttons:
        cls = "fbtn on" if k == "all" else "fbtn"
        out.append("<button class='" + cls + "' onclick=\"flt('" + k + "',this)\">" + label + "</button>")
    out.append("</div><div class='cards'>")
    for r in roles:
        duties = r.get("duties") or []
        badges = "".join("<span class='bd " + esc(d) + "'>" + esc(d) + "</span>" for d in duties)
        if r.get("is_entry"):
            badges += "<span class='bd entry'>★第一入口</span>"
        name = r.get("dept", "") + ((" · " + r["title"]) if r.get("title") else "")
        out.append("<details class='card' data-entry='" + ("1" if r.get("is_entry") else "0") +
                   "' data-duties='" + esc("、".join(duties)) + "'>")
        out.append("<summary>" + esc(r.get("id", "")) + " " + esc(name) + badges + "</summary><div class='card-body'>")
        if r.get("concerns"):
            out.append("<h4>最关心的三个问题</h4><ul>" +
                       "".join("<li>" + esc(c) + "</li>" for c in r["concerns"]) + "</ul>")
        if r.get("opener"):
            out.append("<h4>先聊什么</h4><p>" + esc(r["opener"]) + "</p>")
        if r.get("questions"):
            out.append("<h4>最值得确认的问题</h4><ul>" +
                       "".join("<li>" + esc(q) + "</li>" for q in r["questions"]) + "</ul>")
        contacts = r.get("public_contacts") or []
        if contacts:
            out.append("<h4>公开联系方式</h4><ul>" +
                       "".join("<li>" + esc(c.get("channel", "")) + "：" + esc(c.get("value", "")) + "</li>"
                               for c in contacts) + "</ul>")
        if r.get("basis"):
            out.append("<div class='basis'>依据：" + esc(r["basis"]) + "</div>")
        out.append("</div></details>")
    out.append("</div></section>")

    # 表格区块
    def table(headers, rows):
        h = "".join("<th>" + esc(x) + "</th>" for x in headers)
        body = "".join(
            "<tr>" + "".join(
                "<td>" + (c if isinstance(c, str) else esc(c)) + "</td>" for c in row
            ) + "</tr>"
            for row in rows
        )
        return "<table><tr>" + h + "</tr>" + body + "</table>"

    logic = data.get("logic_chain") or []
    if logic:
        out.append("<section><h2>判断逻辑（证据 → 结论）</h2>" +
                   table(["步骤", "从（证据/前提）", "推到（结论）", "依据"],
                         [[l.get("step", ""), l.get("from", ""), l.get("to", ""), esc(l.get("basis", ""))] for l in logic]) +
                   "</section>")

    prof = data.get("company_profile") or []
    if prof:
        out.append("<section><h2>企业画像（公开信息）</h2>" +
                   table(["维度", "事实", "来源", "置信"],
                         [[p.get("dimension", ""), p.get("fact", ""), src_cell(p), p.get("confidence", "")] for p in prof]) +
                   "</section>")

    blocks = data.get("blocking_points") or []
    if blocks:
        out.append("<section><h2>常见卡点</h2>" +
                   table(["卡点", "卡在谁", "性质", "破解/窗口"],
                         [[b.get("point", ""), b.get("owner", ""), b.get("nature", ""), b.get("workaround", "")] for b in blocks]) +
                   "</section>")

    variants = data.get("variants") or []
    if variants:
        out.append("<section><h2>客户类型变化</h2>" +
                   table(["客户类型", "决策链变化", "第一入口"],
                         [[v.get("customer_type", ""), v.get("change", ""), v.get("first_entry", "")] for v in variants]) +
                   "</section>")

    ev = data.get("evidence_list") or []
    inf = data.get("inference_list") or []
    if ev or inf:
        rows = [["证据", esc(x)] for x in ev] + [["推断", esc(x)] for x in inf]
        out.append("<section><h2>证据 vs 推断速查</h2>" + table(["类型", "条目"], rows) + "</section>")

    srcs = data.get("sources") or []
    if srcs:
        out.append("<section><h2>来源</h2><ol class='srcs'>" +
                   "".join("<li><a href='" + esc(s.get("url", "")) + "' target='_blank' rel='noopener'>" +
                           esc(s.get("title", "")) + "</a>" +
                           (esc(" — " + s["note"]) if s.get("note") else "") + "</li>" for s in srcs) +
                   "</ol></section>")

    out.append("</main><footer>生成日期：" + str(date.today()) +
               " ｜ 关键判断已标注“证据/推断”；仅展示企业公开联系方式，不包含私人信息。</footer>")
    out.append("<script>" + JS + "</script></body></html>")

    with open(path, "w", encoding="utf-8") as f:
        f.write("".join(out))


# --------------------------------------------------------------------------
def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    src = sys.argv[1]
    with open(src, "r", encoding="utf-8") as f:
        data = json.load(f)
    outdir = sys.argv[2] if len(sys.argv) > 2 else os.path.dirname(os.path.abspath(src))
    os.makedirs(outdir, exist_ok=True)
    meta = data.get("meta", {})
    if meta.get("file_stem"):
        stem = safe_name(meta["file_stem"])
    else:
        name = meta.get("short_name") or meta.get("company", "")
        prod = meta.get("product", "")
        prod = re.split(r"[（(]", prod)[0].strip() or prod  # 去掉括号中的规格说明
        stem = safe_name("采购决策链-{}-{}".format(name, prod))
    xlsx_path = os.path.join(outdir, stem + ".xlsx")
    html_path = os.path.join(outdir, stem + ".html")
    try:
        build_xlsx(data, xlsx_path)
        print("OK  xlsx:", xlsx_path)
    except ImportError:
        print("SKIP xlsx: 未安装 openpyxl（pip install openpyxl 后重跑）；HTML 不受影响")
    build_html(data, html_path)
    print("OK  html:", html_path)


if __name__ == "__main__":
    main()
