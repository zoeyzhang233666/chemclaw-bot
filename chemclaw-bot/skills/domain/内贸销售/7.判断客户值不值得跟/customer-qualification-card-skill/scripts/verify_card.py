# -*- coding: utf-8 -*-
"""verify_card.py —— 校验 data.json 与产物（Excel/HTML），交付前必跑。

用法:
    python verify_card.py <data.json> [--xlsx out.xlsx] [--html out.html]

退出码 0=通过（可有警告），1=存在错误。
"""
import json
import os
import sys

LEVELS = {"高", "中高", "中", "中低", "低"}
ERR, WARN = [], []


def check(data_path):
    with open(data_path, encoding="utf-8") as f:
        d = json.load(f)
    meta = d.get("meta") or {}
    for k in ("title", "product", "seller", "as_of"):
        if not meta.get(k):
            ERR.append("meta.%s 缺失" % k)
    th = meta.get("thresholds") or {"high": 75, "mid": 55}
    dims = d.get("dimensions") or []
    if len(dims) != 7:
        WARN.append("维度数量=%d（标准为 7），确认是否有意调整" % len(dims))
    keys = [x["key"] for x in dims]
    custs = d.get("customers") or []
    if not custs:
        ERR.append("没有客户数据")
    presets = d.get("presets") or []
    if not presets:
        WARN.append("未定义 presets，将使用默认均衡权重")
    for p in presets:
        miss = [k for k in keys if k not in (p.get("weights") or {})]
        if miss:
            ERR.append("preset[%s] 缺权重: %s" % (p.get("key"), ",".join(miss)))

    for c in custs:
        nm = c.get("name", "?")
        for k in ("name", "location", "distance", "demand", "best", "action"):
            if not c.get(k):
                ERR.append("[%s] 缺少字段 %s" % (nm, k))
        if not (c.get("sources") or []):
            WARN.append("[%s] sources 为空（Excel 数据来源表将不完整）" % nm)
        elif len(c["sources"]) < 4:
            WARN.append("[%s] 来源仅 %d 条（建议 ≥4）" % (nm, len(c["sources"])))
        if len(c.get("doubts") or []) < 1:
            WARN.append("[%s] 缺主要疑点" % nm)
        if len(c.get("asks") or []) < 2:
            WARN.append("[%s] 待确认事项 <2 条" % nm)
        cur = c.get("dims") or {}
        for k in keys:
            if k not in cur:
                ERR.append("[%s] 缺维度 %s" % (nm, k))
                continue
            x = cur[k]
            lv, sc, fact = x.get("level"), x.get("score"), x.get("fact") or ""
            if lv not in LEVELS:
                ERR.append("[%s/%s] 评级非法: %s" % (nm, k, lv))
            if not isinstance(sc, (int, float)) or not (0 <= sc <= 100):
                ERR.append("[%s/%s] 分值非法: %s" % (nm, k, sc))
            if len(fact) < 8:
                ERR.append("[%s/%s] 事实过短" % (nm, k))
            if len(fact) > 160:
                WARN.append("[%s/%s] 事实偏长（%d 字），建议 ≤120 字" % (nm, k, len(fact)))
            evs = x.get("evidence") or []
            if not evs:
                ERR.append("[%s/%s] 缺证据链接" % (nm, k))
            for e in evs:
                if not (e.get("url") or "").startswith("http"):
                    ERR.append("[%s/%s] 证据无有效 URL: %s" % (nm, k, e.get("label")))
            if isinstance(sc, (int, float)):
                if k == "risk":
                    pass  # 不确定性维度方向相反：分数越低风险越大
                elif lv and lv.startswith("高") and sc < th["high"]:
                    WARN.append("[%s/%s] 评级高但分值 %s < %d" % (nm, k, sc, th["high"]))
                if k != "risk" and lv and lv.startswith("低") and sc >= th["mid"]:
                    WARN.append("[%s/%s] 评级低但分值 %s ≥ %d" % (nm, k, sc, th["mid"]))
    return d, keys, custs, presets, th


def weighted(c, keys, p):
    w = p["weights"]
    s = sum(w[k] * c["dims"][k]["score"] for k in keys)
    tot = sum(w[k] for k in keys)
    return s / tot if tot else 0


def check_xlsx(path, d, keys, custs, presets, th):
    try:
        from openpyxl import load_workbook
    except Exception as e:
        WARN.append("未安装 openpyxl，跳过 Excel 校验: %s" % e)
        return
    wb = load_workbook(path, data_only=True)
    need = ["客户判断总览", "七维判断明细", "权重与排序", "待确认清单", "数据来源", "市场背景", "阅读说明"]
    miss = [s for s in need if s not in wb.sheetnames]
    if miss:
        ERR.append("Excel 缺工作表: %s" % ",".join(miss))
        return
    ws = wb["客户判断总览"]
    for i, c in enumerate(custs):
        got = ws.cell(row=5 + i, column=23).value
        exp = weighted(c, keys, presets[0])
        if got is None:
            ERR.append("[%s] Excel 排序分无缓存值" % c["name"])
        elif abs(float(got) - exp) > 1e-6:
            ERR.append("[%s] Excel 排序分 %s ≠ 期望 %.4f" % (c["name"], got, exp))
    if ws.cell(row=5, column=1).value is None:
        ERR.append("Excel 排名列无缓存值（RANK 未计算）")


def check_html(path):
    if not os.path.exists(path):
        ERR.append("HTML 不存在: %s" % path)
        return
    size = os.path.getsize(path)
    if size < 50000:
        WARN.append("HTML 仅 %.1f KB，可能未注入数据" % (size / 1024.0))
    with open(path, encoding="utf-8") as f:
        head = f.read(200000)
    if "__CARD_DATA__" in head:
        ERR.append("HTML 仍含占位符 __CARD_DATA__")
    if "const CARD =" not in head:
        ERR.append("HTML 缺少数据注入")


def main():
    args = [a for a in sys.argv[1:]]
    if not args:
        print("用法: python verify_card.py <data.json> [--xlsx out.xlsx] [--html out.html]")
        sys.exit(2)
    data_path = args[0]
    xlsx = args[args.index("--xlsx") + 1] if "--xlsx" in args else None
    html = args[args.index("--html") + 1] if "--html" in args else None

    d, keys, custs, presets, th = check(data_path)
    if xlsx:
        check_xlsx(xlsx, d, keys, custs, presets, th)
    if html:
        check_html(html)

    order = sorted(custs, key=lambda c: -weighted(c, keys, presets[0]))
    print("默认权重排序:")
    for i, c in enumerate(order):
        print("  %d. %-16s %.1f" % (i + 1, c["name"], weighted(c, keys, presets[0])))
    for w in WARN:
        print("WARN:", w)
    for e in ERR:
        print("ERROR:", e)
    print("RESULT:", "FAIL" if ERR else "PASS", "(%d errors, %d warnings)" % (len(ERR), len(WARN)))
    sys.exit(1 if ERR else 0)


if __name__ == "__main__":
    main()
