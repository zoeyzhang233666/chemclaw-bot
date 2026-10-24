#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""build_visit_workbook.py —— 从 visit_data.json 生成 8 表 Excel 拜访工作簿。

用法:
    python build_visit_workbook.py --data visit_data.json [--out 拜访背景卡_客户_产品_日期.xlsx]

仅依赖 openpyxl（pip install openpyxl）。契约见 ../references/data-contract.md。
跨平台（Windows/macOS/Linux）；路径处理只用 pathlib。
"""
import argparse
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# ---------------- 主题（见 references/output-spec.md §A） ----------------
C_DARK, C_LIGHT, C_ALT = "0D4A7A", "D9E7F5", "F4F8FC"
C_AMBER, C_GREEN, C_RED, C_GRAY = "FFF3DE", "E7F5EC", "FBEDEC", "EEF1F4"
F_TITLE = Font(name="微软雅黑", size=14, bold=True, color="FFFFFF")
F_HEAD = Font(name="微软雅黑", size=10.5, bold=True, color="0D4A7A")
F_BODY = Font(name="微软雅黑", size=10.5, color="1C2B36")
F_SMALL = Font(name="微软雅黑", size=9, color="5B6B78")
F_KEY = Font(name="微软雅黑", size=10.5, bold=True, color="0D4A7A")
F_LINK = Font(name="微软雅黑", size=9, color="0E6BA8", underline="single")
F_NOTE = Font(name="微软雅黑", size=10.5, italic=True, color="8A6D1B")
F_EX = Font(name="微软雅黑", size=10.5, italic=True, color="9AA7B2")
FILL_DARK = PatternFill("solid", fgColor=C_DARK)
FILL_LIGHT = PatternFill("solid", fgColor=C_LIGHT)
FILL_ALT = PatternFill("solid", fgColor=C_ALT)
FILL_AMBER = PatternFill("solid", fgColor=C_AMBER)
FILL_GREEN = PatternFill("solid", fgColor=C_GREEN)
FILL_RED = PatternFill("solid", fgColor=C_RED)
FILL_GRAY = PatternFill("solid", fgColor=C_GRAY)
TYPE_FILL = {  # 大事记类型着色
    "点火": FILL_GREEN, "投产": FILL_GREEN, "中标": FILL_GREEN,
    "冷修": FILL_AMBER, "业绩": FILL_RED, "评级": FILL_LIGHT, "行业": FILL_LIGHT,
}
THIN = Side(style="thin", color="C9D4DE")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
AL_TL = Alignment(horizontal="left", vertical="top", wrap_text=True)
AL_C = Alignment(horizontal="center", vertical="center", wrap_text=True)

DEFAULT_USAGE = [
    ("本工作簿是什么", "拜访任何一家 B2B 化工客户前的背景准备材料：与同名 HTML 一页卡内容一致，Excel 更适合现场逐条记录与事后台账化。"),
    ("怎么用（拜访前）", "①「拜访概览」5分钟速览 → ②「客户画像」「大事记」「行情监控」建立背景 → ③ 按「提问清单」选择你面对的角色版本（采购/技术/老板）。"),
    ("怎么用（拜访中）", "在「提问清单」的“现场记录/后续动作/负责人/期限/状态”列直接填写；“状态”列用下拉更新。"),
    ("怎么用（拜访后）", "把当天沟通写入「历史沟通台账」（只写有凭证的内容）；下一张拜访卡的第⑤节即由本表累积。"),
    ("数据政策（重要）", "全卡基于公开信息（见「数据来源」表，均注明日期）；客户需求/痛点均为“待验证假设”；未上传历史沟通记录时，相关章节一律留空，不虚构客户说过的话。"),
]


def sheet_title(ws, text, ncols, sub=None):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=text)
    c.font, c.fill = F_TITLE, FILL_DARK
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 30
    r = 2
    if sub:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
        c = ws.cell(row=2, column=1, value=sub)
        c.font, c.fill = F_SMALL, FILL_GRAY
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
        ws.row_dimensions[2].height = 22
        r = 3
    return r


def table_header(ws, row, headers, widths):
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font, c.fill, c.alignment, c.border = F_HEAD, FILL_LIGHT, AL_C, BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 20
    return row + 1


def data_rows(ws, row, rows, zebra=True, height=None):
    for ridx, vals in enumerate(rows):
        for cidx, v in enumerate(vals, start=1):
            c = ws.cell(row=row, column=cidx, value=v)
            c.font, c.alignment, c.border = F_BODY, AL_TL, BORDER
            if zebra and ridx % 2 == 1:
                c.fill = FILL_ALT
        if height:
            ws.row_dimensions[row].height = height
        row += 1
    return row


def linkify(ws, col, row_start, row_end):
    for row_i in range(row_start, row_end):
        c = ws.cell(row=row_i, column=col)
        if c.value:
            c.hyperlink = str(c.value)
            c.font = F_LINK


def note_block(ws, row, ncols, text, fill=FILL_AMBER, font=F_NOTE, height=None):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1, value=text)
    c.fill, c.font, c.alignment, c.border = fill, font, AL_TL, BORDER
    if height:
        ws.row_dimensions[row].height = height
    return row + 1


def set_landscape(ws):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def empty_state(ws, r, ncols, headers, widths, text):
    r = table_header(ws, r, headers, widths)
    return note_block(ws, r, ncols, text, FILL_GRAY, F_SMALL, 22)


# ---------------- 各表 ----------------

def s_usage(wb, d):
    ws = wb.active
    ws.title = "使用说明"
    ws.sheet_properties.tabColor = C_DARK
    r = sheet_title(ws, "客户拜访背景卡 · 使用说明", 2,
                    f"{d['meta'].get('generator','') } ｜ 客户：{d['meta'].get('customer','')} ｜ 生成：{d['meta'].get('date_generated','')}")
    r = table_header(ws, r, ["条目", "说明"], [26, 96])
    data_rows(ws, r, d.get("usage") or DEFAULT_USAGE, height=42)


def s_overview(wb, d):
    ws = wb.create_sheet("拜访概览")
    ws.sheet_properties.tabColor = "0E6BA8"
    m = d["meta"]
    r = sheet_title(ws, f"拜访概览 · {m.get('customer','')}（{m.get('stock_code','')}）", 4,
                    f"建议先花5分钟读完本页再进入现场 ｜ 数据截至 {m.get('data_cutoff','')}")
    kv = [
        ["我方产品", m.get("our_product", ""), "", ""],
        ["本次拜访目的", m.get("visit_purpose", ""), "", ""],
        ["建议拜访对象", m.get("visit_target", ""), "", ""],
        ["建议时长", m.get("duration", ""), "", ""],
    ]
    r = table_header(ws, r, ["要素", "内容", "要点提示", "备注"], [22, 34, 34, 42])
    r = data_rows(ws, r, kv, height=20)
    if d.get("kpis"):
        r += 1
        r = note_block(ws, r, 4, "▍关键数字看板（KPI）", FILL_LIGHT, F_KEY, 20)
        rows = []
        for k in d["kpis"]:
            arrow = {"up": "▲", "down": "▼", "flat": "→"}.get(k.get("tone", ""), "")
            rows.append([k.get("label", ""), f"{k.get('value','')} {k.get('unit','')}".strip(),
                         k.get("delta", ""), f"{arrow} {k.get('source','')}".strip()])
        r = table_header(ws, r, ["指标", "数值", "口径/说明", "来源（日期）"], [22, 34, 34, 42])
        r = data_rows(ws, r, rows, height=30)
    if d.get("risks_opportunities"):
        r += 1
        r = note_block(ws, r, 4, "▍风险与机会", FILL_LIGHT, F_KEY, 20)
        r = table_header(ws, r, ["类型", "内容", "对我方动作的含义", "状态"], [22, 34, 34, 42])
        rows = [(("机会" if str(v[0]).startswith("机会") else "风险"), v[1], v[2], v[3]) for v in d["risks_opportunities"]]
        r = data_rows(ws, r, rows, height=28)
        for row_i in range(r - len(rows), r):
            c = ws.cell(row=row_i, column=1)
            c.fill = FILL_GREEN if c.value == "机会" else FILL_AMBER
            c.font = F_KEY
    set_landscape(ws)


def s_profile(wb, d):
    ws = wb.create_sheet("客户画像")
    ws.sheet_properties.tabColor = "0E8A8A"
    r = sheet_title(ws, "① 客户实际做什么", 3, "公开信息给来源；推断项标注“推断·待验证”；客户提供项标注“客户确认”")
    if d.get("profile"):
        r = table_header(ws, r, ["维度", "信息", "来源（日期）"], [16, 78, 40])
        data_rows(ws, r, d["profile"], height=40)
    else:
        empty_state(ws, r, 3, ["维度", "信息", "来源（日期）"], [16, 78, 40], "（本卡未包含客户画像数据）")
    set_landscape(ws)


def s_timeline(wb, d):
    ws = wb.create_sheet("大事记")
    ws.sheet_properties.tabColor = "B45309"
    r = sheet_title(ws, "③ 近6–12个月重要事情", 5,
                    "类型：点火/投产/中标=窗口；冷修=收缩；业绩=心态；评级/行业=背景。每条必须写“对本次拜访的意义”")
    r = table_header(ws, r, ["日期", "类型", "事件", "对本次拜访的意义", "来源"], [12, 8, 52, 42, 26])
    if d.get("timeline"):
        start = r
        r = data_rows(ws, r, d["timeline"], height=34)
        for row_i, v in enumerate(d["timeline"], start=start):
            fill = TYPE_FILL.get(str(v[1]).strip(), FILL_GRAY)
            c = ws.cell(row=row_i, column=2)
            c.fill, c.font, c.alignment = fill, F_KEY, AL_C
    else:
        r = note_block(ws, r, 5, "（本卡未包含大事记数据）", FILL_GRAY, F_SMALL, 22)
    set_landscape(ws)


def s_market(wb, d):
    ws = wb.create_sheet("行情监控")
    ws.sheet_properties.tabColor = "B3352C"
    r = sheet_title(ws, f"④ 相关市场行情（截至 {d['meta'].get('data_cutoff','')}）", 8,
                    "▼下跌 ▲上涨 →低位徘徊；每个价格必须带数据日期与来源")
    r = table_header(ws, r, ["品种", "规格/地区", "价格", "单位", "数据日期", "近期变化", "来源（日期）", "链接"],
                     [14, 20, 16, 8, 12, 34, 26, 34])
    if d.get("market"):
        start = r
        r = data_rows(ws, r, d["market"], height=34)
        linkify(ws, 8, start, r)
    if d.get("market_note"):
        r += 1
        r = note_block(ws, r, 8, "▍谈判含义：" + d["market_note"], height=34)
    set_landscape(ws)


def s_questions(wb, d):
    ws = wb.create_sheet("提问清单")
    ws.sheet_properties.tabColor = "1A7F4B"
    r = sheet_title(ws, "⑥ 提问清单（按拜访对象选版本） + 现场记录", 9,
                    "现场只填“现场记录/后续动作/负责人/期限/状态”五列；其余为准备内容")
    r = table_header(ws, r, ["版本", "编号", "问题", "提问目的", "现场记录（客户回答要点）", "后续动作", "负责人", "期限", "状态"],
                     [10, 6, 46, 24, 34, 18, 8, 10, 10])
    rows, first = [], None
    for version, items in (d.get("questions") or {}).items():
        for it in items:
            rows.append([version, it.get("no", ""), it.get("q", ""), it.get("why", ""), "", "", "", "", "未开始"])
    if rows:
        first = r
        r = data_rows(ws, r, rows, height=26)
        dv = DataValidation(type="list", formula1='"未开始,进行中,已确认,待跟进,已关闭"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"I{first}:I{r - 1}")
    else:
        r = note_block(ws, r, 9, "（本卡未生成问题清单）", FILL_GRAY, F_SMALL, 22)
    r = note_block(ws, r, 9,
                   "提示：采购版适合首访；技术版用于对接质量/工艺工程师；老板版用于高层汇报。同一问题跨版本不重复提问。",
                   FILL_GRAY, F_SMALL, 18)
    ws.freeze_panes = "A4"
    set_landscape(ws)


def s_ledger(wb, d):
    ws = wb.create_sheet("历史沟通台账")
    ws.sheet_properties.tabColor = "5B6B78"
    r = sheet_title(ws, "⑤ 历史沟通台账（仅录入有凭证的沟通）", 9,
                    "数据政策：只记录真实发生的沟通；未上传/未发生时保持空白，不虚构客户需求。拜访后当日更新。")
    r = table_header(ws, r, ["日期", "渠道", "对方姓名/角色", "沟通要点", "已确认事项", "未解决问题", "证据文件名", "录入人", "下次跟进动作"],
                     [12, 10, 16, 34, 28, 28, 18, 10, 24])
    if d.get("ledger_example"):
        keys = ["date", "channel", "person", "points", "confirmed", "unresolved", "evidence", "recorder", "next"]
        start = r
        r = data_rows(ws, r, [[d["ledger_example"].get(k, "") for k in keys]], zebra=False, height=30)
        for cidx in range(1, 10):
            c = ws.cell(row=start, column=cidx)
            c.font, c.fill = F_EX, FILL_GRAY
        r = note_block(ws, r, 9, "▲ 上行仅为格式示例（灰色斜体），正式使用时请删除。以下留白行供实际录入：",
                       FILL_LIGHT, F_SMALL, 18)
    for _ in range(12):
        for cidx in range(1, 10):
            ws.cell(row=r, column=cidx).border = BORDER
            ws.cell(row=r, column=cidx).alignment = AL_TL
        ws.row_dimensions[r].height = 24
        r += 1
    r = note_block(ws, r, 9,
                   d.get("ledger_baseline") or "首访现场请建立基线并录入本表：在供供应商与份额 ｜ 招标周期 ｜ 规格红线 ｜ 账期与结算 ｜ 物流与验收方式 ｜ 准入流程。",
                   FILL_GREEN, F_SMALL, 20)
    ws.freeze_panes = "A4"
    set_landscape(ws)


def s_sources(wb, d):
    ws = wb.create_sheet("数据来源")
    ws.sheet_properties.tabColor = "0E6BA8"
    r = sheet_title(ws, f"数据来源（{len(d.get('sources', []))}条 · 全部公开信息）", 6,
                    "与卡内引用一一对应；点击链接可查证")
    r = table_header(ws, r, ["序号", "类别", "要点", "来源名称", "日期", "链接"], [6, 12, 50, 24, 12, 44])
    if d.get("sources"):
        start = r
        r = data_rows(ws, r, d["sources"], height=30)
        linkify(ws, 6, start, r)
    else:
        r = note_block(ws, r, 6, "（本卡未登记来源——这是不合格的交付，请补充研究面A/C/D的出处）", FILL_RED, F_SMALL, 22)
    set_landscape(ws)


def default_out(meta):
    customer = re.sub(r'[\\/:*?"<>|\s]+', "", meta.get("customer", "客户"))[:24]
    date = meta.get("date_generated", "").replace("-", "")
    return f"拜访背景卡_{customer}_{date}.xlsx"


def main():
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
    ap = argparse.ArgumentParser(description="visit_data.json → 8表 Excel 拜访工作簿")
    ap.add_argument("--data", required=True, help="visit_data.json 路径")
    ap.add_argument("--out", default=None, help="输出 xlsx 路径（默认按客户+日期命名）")
    args = ap.parse_args()

    data = json.loads(Path(args.data).read_text(encoding="utf-8-sig"))
    meta = data.get("meta", {})
    for key in ("customer", "our_product", "visit_purpose", "date_generated"):
        if not meta.get(key):
            sys.exit(f"[契约错误] meta.{key} 必填，见 references/data-contract.md")

    wb = Workbook()
    s_usage(wb, data)
    s_overview(wb, data)
    s_profile(wb, data)
    s_timeline(wb, data)
    s_market(wb, data)
    s_questions(wb, data)
    s_ledger(wb, data)
    s_sources(wb, data)

    wb.properties.title = f"客户拜访背景卡·{meta.get('customer', '')}"
    wb.properties.creator = meta.get("generator", "customer-visit-brief")
    out = args.out or default_out(meta)
    Path(out).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"saved: {out}  (sheets: {wb.sheetnames})")


if __name__ == "__main__":
    main()
