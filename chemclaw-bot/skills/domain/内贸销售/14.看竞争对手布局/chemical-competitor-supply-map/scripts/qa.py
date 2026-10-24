#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Hard-gate QA for a built pair (HTML + XLSX).

Usage:
  python qa.py --data data.json --html out.html --xlsx out.xlsx

Exit code 1 on any hard failure. This is intentionally stricter than
validate.py: validate checks the data, qa.py checks the artifacts.
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys

import common
import validate


def _fail(errors, msg):
    errors.append(msg)


def check_html(path, D, errors, warnings):
    if not os.path.exists(path):
        _fail(errors, f"HTML 不存在: {path}")
        return
    size = os.path.getsize(path)
    if size < 500_000:
        _fail(errors, f"HTML 只有 {size/1024:.0f} KB；ECharts/GeoJSON 可能未内联")
    s = open(path, encoding="utf-8").read()
    for token in ("window.__DATA__", "window.__GEO__", "echarts.registerMap", 'id="map"', 'id="matrixBody"', 'id="prospectTable"'):
        if token not in s:
            _fail(errors, f"HTML 缺少关键标记: {token}")
    if re.search(r"<function\s+", s):
        _fail(errors, "HTML 出现未渲染的 Python function 占位符")
    if re.search(r"<script[^>]+src\s*=", s, re.I):
        _fail(errors, "HTML 含外部 <script src>；必须完全离线")
    if re.search(r"<link[^>]+rel=[\"']?stylesheet", s, re.I):
        _fail(errors, "HTML 含外部样式表；CSS 必须内联")
    for sec in ("mapsec", "matrixsec", "densitysec", "newsec", "prospectsec", "methodsec"):
        if f'id="{sec}"' not in s:
            _fail(errors, f"HTML 缺少章节 id={sec}")
    m = re.search(r"window\.__DATA__=(\{.*?\});</script>", s, re.S)
    if not m:
        _fail(errors, "HTML 中无法定位内联 data JSON")
    else:
        try:
            embedded = json.loads(m.group(1))
        except Exception as exc:
            _fail(errors, f"HTML 内联 data JSON 无法解析: {exc}")
        else:
            for k in ("total_companies", "total_plants", "total_cap", "companies", "plants", "sources"):
                if embedded.get(k) != D.get(k):
                    _fail(errors, f"HTML 内联 {k} 与 data.json 不一致")
    if "None" in re.sub(r"<script.*?</script>", "", s, flags=re.S)[:200000]:
        warnings.append("HTML 可见文本出现 None；检查模板空值")
    if "https://" not in s and "http://" not in s:
        warnings.append("HTML 没有可点击来源链接；确认 sources.url 是否填写")
    print(f"  HTML  {size/1024/1024:.2f} MB  sections=6  inline-data=ok")


def check_xlsx(path, D, errors, warnings):
    if not os.path.exists(path):
        _fail(errors, f"XLSX 不存在: {path}")
        return
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=False)
    except Exception as exc:
        _fail(errors, f"XLSX 无法打开: {exc}")
        return
    expected = ["00_封面", "01_企业矩阵", "02_工厂明细", "03_新增项目", "04_区域对照",
                "05_目标区域与潜客", "06_数据来源", "07_口径与冲突", "08_数据缺口与边界"]
    if wb.sheetnames != expected:
        _fail(errors, f"工作表名称/顺序不符: {wb.sheetnames}")
    if len(wb.sheetnames) < 9:
        _fail(errors, f"工作表只有 {len(wb.sheetnames)} 个；至少 9 个")
    for ws in wb.worksheets:
        if ws.max_row < 3:
            _fail(errors, f"{ws.title} 内容过少")
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value in ("#REF!", "#VALUE!", "#NAME?"):
                    _fail(errors, f"{ws.title}!{c.coordinate} 有错误值 {c.value}")
    # merged overlaps
    for ws in wb.worksheets:
        ranges = list(ws.merged_cells.ranges)
        for i, a in enumerate(ranges):
            for b in ranges[i + 1:]:
                if not (a.max_row < b.min_row or b.max_row < a.min_row or a.max_col < b.min_col or b.max_col < a.min_col):
                    _fail(errors, f"{ws.title} 合并单元格重叠: {a} / {b}")
    ws = wb["01_企业矩阵"]
    if "企业对比矩阵" not in str(ws["A1"].value or ""):
        _fail(errors, "01_企业矩阵 A1 标题异常")
    data_rows = sum(1 for r in range(5, ws.max_row + 1) if ws.cell(r, 2).value and str(ws.cell(r, 2).value).startswith("合计") is False)
    if data_rows < D["total_companies"]:
        _fail(errors, f"01_企业矩阵 数据行 {data_rows} < 企业数 {D['total_companies']}")
    ws2 = wb["02_工厂明细"]
    plant_rows = sum(1 for r in range(5, ws2.max_row + 1) if ws2.cell(r, 3).value and str(ws2.cell(r, 3).value).startswith("合计") is False)
    if plant_rows < D["total_plants"]:
        _fail(errors, f"02_工厂明细 数据行 {plant_rows} < 工厂数 {D['total_plants']}")
    ws3 = wb["03_新增项目"]
    proj_rows = sum(1 for r in range(5, ws3.max_row + 1) if ws3.cell(r, 3).value and str(ws3.cell(r, 3).value).startswith("合计") is False)
    if proj_rows < len(D.get("projects") or []):
        _fail(errors, f"03_新增项目 数据行 {proj_rows} < 项目数 {len(D.get('projects') or [])}")
    ws5 = wb["05_目标区域与潜客"]
    prospect_rows = sum(1 for r in range(1, ws5.max_row + 1) if ws5.cell(r, 1).value in ("S", "A", "B", "C"))
    if prospect_rows < len(D.get("prospects") or []):
        _fail(errors, f"05 潜客行 {prospect_rows} < 潜客数 {len(D.get('prospects') or [])}")
    ws6 = wb["06_数据来源"]
    src_rows = sum(1 for r in range(5, ws6.max_row + 1) if ws6.cell(r, 1).value and str(ws6.cell(r, 1).value).startswith("S"))
    if src_rows < len(D.get("sources") or []):
        _fail(errors, f"06 来源行 {src_rows} < 来源数 {len(D.get('sources') or [])}")
    links = sum(1 for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.hyperlink)
    if links < len(D.get("sources") or []):
        _fail(errors, f"超链接只有 {links} 个，少于来源数 {len(D.get('sources') or [])}")
    charts = sum(len(ws._charts) for ws in wb.worksheets)
    if charts < 2:
        warnings.append(f"图表只有 {charts} 个；可接受，但建议至少 2 个")
    formulas = sum(1 for ws in wb.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
    if formulas:
        warnings.append(f"含 {formulas} 个公式；本技能默认不依赖公式，若使用需重算缓存")
    print(f"  XLSX  {os.path.getsize(path)/1024:.1f} KB  sheets={len(wb.sheetnames)}  links={links}  charts={charts}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--data", required=True)
    ap.add_argument("--html", required=True)
    ap.add_argument("--xlsx", required=True)
    a = ap.parse_args()
    D = common.build(common.load_raw(a.data))
    errors, warnings = validate.check(D)
    print("== data ==")
    print(f"  {D['total_companies']} companies / {D['total_plants']} plants / {D['total_cap']:,.0f} sample cap / {len(D.get('sources') or [])} sources")
    print("== artifacts ==")
    check_html(a.html, D, errors, warnings)
    check_xlsx(a.xlsx, D, errors, warnings)
    print("== result ==")
    for w in warnings:
        print("  [WARN ]", w)
    if errors:
        for x in errors:
            print("  [FAIL ]", x)
        print(f"FAIL: {len(errors)} hard error(s)")
        raise SystemExit(1)
    print(f"PASS: 0 hard errors, {len(warnings)} warning(s)")


if __name__ == "__main__":
    main()
