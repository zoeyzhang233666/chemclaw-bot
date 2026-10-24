# -*- coding: utf-8 -*-
"""
build_xlsx.py — 由 followup.json 生成《客户跟进卡》Excel 工作簿

用法：
    python build_xlsx.py followup.json -o 输出.xlsx

产品无关：所有文案来自 JSON，不含任何具体化工品硬编码。
"""
import argparse
import math
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.formatting.rule import FormulaRule
except ImportError:  # pragma: no cover
    sys.exit("缺少依赖 openpyxl。请先执行：pip install openpyxl")

from _common import (
    FONT, INK, INK2, GOLD, GOLD_L, LINE, ZEBRA, BG,
    cat_color, tone_color, est_h, cjk_len,
    parse_as_of, compute_stall_days, promise_state, days_left,
    STAGE_NAMES, validate,
)

thin = Side(style="thin", color=LINE)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


# ───────────────────────── 基础样式 ─────────────────────────
def style_cell(c, size=10, bold=False, color="243247", fill=None,
               wrap=True, ha="left", va="center", border=BOX):
    c.font = Font(name=FONT, size=size, bold=bold, color=color)
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    c.alignment = Alignment(horizontal=ha, vertical=va, wrap_text=wrap)
    if border:
        c.border = border


def title_block(ws, ncols, title, subtitle, tab_color=None):
    last = get_column_letter(ncols)
    ws.merge_cells(f"A1:{last}1")
    c = ws["A1"]
    c.value = title
    for i in range(1, ncols + 1):
        cc = ws.cell(row=1, column=i)
        cc.fill = PatternFill("solid", fgColor=INK)
        cc.border = Border()
    c.font = Font(name=FONT, size=16, bold=True, color="FFFFFF")
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 38

    ws.merge_cells(f"A2:{last}2")
    c = ws["A2"]
    c.value = subtitle
    for i in range(1, ncols + 1):
        cc = ws.cell(row=2, column=i)
        cc.fill = PatternFill("solid", fgColor="F1F5F9")
        cc.border = Border()
    c.font = Font(name=FONT, size=9, color="4A5A6A")
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
    ws.row_dimensions[2].height = 28
    if tab_color:
        ws.sheet_properties.tabColor = tab_color


def header_row(ws, row, headers):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        style_cell(c, size=9.5, bold=True, color="FFFFFF", fill=INK2, ha="center")
    ws.row_dimensions[row].height = 28


def section_bar(ws, row, ncols, text, fill="E8F2F6", color=INK):
    last = get_column_letter(ncols)
    ws.merge_cells(f"A{row}:{last}{row}")
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(name=FONT, size=10.5, bold=True, color=color)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for i in range(1, ncols + 1):
        cc = ws.cell(row=row, column=i)
        cc.fill = PatternFill("solid", fgColor=fill)
        cc.border = Border(bottom=Side(style="thin", color="C7D6DF"))
    ws.row_dimensions[row].height = 24


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def page_setup(ws, ncols, last_row, landscape=True):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = 9
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.oddFooter.center.text = "第 &P 页 / 共 &N 页"
    ws.oddFooter.center.size = 8
    ws.oddFooter.center.color = "808080"


def refit_heights(wb, pad=6.0, factor=0.94):
    for ws in wb.worksheets:
        merged_w = {}
        for rng in ws.merged_cells.ranges:
            w = 0.0
            for ci in range(rng.min_col, rng.max_col + 1):
                w += ws.column_dimensions[get_column_letter(ci)].width or 8.43
            merged_w[(rng.min_row, rng.min_col)] = w
        for row in ws.iter_rows():
            r0 = row[0].row
            if r0 == 1:
                continue
            need = 0.0
            for c in row:
                if c.value is None:
                    continue
                w = merged_w.get((c.row, c.column)) or (ws.column_dimensions[get_column_letter(c.column)].width or 8.43)
                sz = c.font.size or 10
                units = cjk_len(c.value)
                per = max(4.0, w * factor * (10.0 / sz))
                need = max(need, math.ceil(units / per) * sz * 1.62 + pad)
            cur = ws.row_dimensions[r0].height or 0
            ws.row_dimensions[r0].height = round(max(cur, need), 1)


def simple_table(ws, headers, rows, widths, start_row=5, styles=None, heights=None):
    """styles: {col_idx: dict(color=,fill=,bold=,ha=)}"""
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=start_row - 1, column=i, value=h)
        style_cell(c, size=9.5, bold=True, color="FFFFFF", fill=INK2, ha="center")
    ws.row_dimensions[start_row - 1].height = 28

    r = start_row
    for idx, row in enumerate(rows):
        for j, val in enumerate(row, start=1):
            c = ws.cell(row=r, column=j, value=val)
            fill = ZEBRA if idx % 2 == 1 else None
            kw = dict(size=9.5, color="243247", fill=fill, va="top")
            if styles and j in styles:
                s = styles[j]
                kw.update(color=s.get("color", "243247"),
                          fill=s.get("fill", fill),
                          bold=s.get("bold", False),
                          ha=s.get("ha", "left"))
            style_cell(c, **kw)
        ws.row_dimensions[r].height = est_h(row, widths, base=15)
        r += 1
    return r


# ───────────────────────── 主构建 ─────────────────────────
def build(data, out_path):
    meta = data.get("meta") or {}
    as_of = parse_as_of(meta.get("as_of_date"))
    timeline = data.get("timeline") or []
    stage = data.get("stage") or {}
    stall = stage.get("stall") or {}

    computed_days, last_rec = compute_stall_days(timeline, as_of)
    stall_days = stall.get("days")
    if stall_days in (None, ""):
        stall_days = computed_days
    last_date = (last_rec or {}).get("date", "—")
    last_time = (last_rec or {}).get("time", "")

    promises = data.get("promises") or []
    overdue = [p for p in promises if promise_state(p, as_of)[1] == "risk"]
    days_to_expiry = days_left(meta.get("quote_expiry"), as_of)

    wb = Workbook()

    # ══════════════ 00_跟进卡 ══════════════
    ws = wb.active
    ws.title = "00_跟进卡"
    W = [18, 30, 25, 22, 16, 18]
    set_widths(ws, W)
    NC = 6
    ws.sheet_view.showGridLines = False

    sub_bits = []
    if stage.get("current"):
        sub_bits.append(f"阶段：{stage['current']}")
    if stall_days is not None:
        sub_bits.append(f"停滞 {stall_days} 天")
    if overdue:
        sub_bits.append(f"逾期承诺 {len(overdue)} 项")
    if days_to_expiry is not None and days_to_expiry >= 0:
        sub_bits.append(f"报价有效期剩 {days_to_expiry} 天")
    sub_bits.append(f"生成时间 {meta.get('generated_at','—')}")
    sub_bits.append(f"素材 {meta.get('source_count','—')} 份 / {len(timeline)} 条记录")

    title_block(ws, NC, meta.get("title", "客户跟进卡"), "　｜　".join(sub_bits), tab_color=INK)

    r = 4
    ws.merge_cells(f"A{r}:B{r}")
    c = ws.cell(row=r, column=1, value="一句话结论")
    style_cell(c, size=9.5, bold=True, color="6B4D14", fill=GOLD_L, ha="center")
    for i in (1, 2):
        ws.cell(row=r, column=i).fill = PatternFill("solid", fgColor=GOLD_L)
        ws.cell(row=r, column=i).border = BOX
    ws.merge_cells(f"C{r}:F{r}")
    c = ws.cell(row=r, column=3, value=data.get("headline", ""))
    style_cell(c, size=10.5, bold=True, color="3F2A05", fill="FFFDF6")
    for i in range(3, 7):
        ws.cell(row=r, column=i).border = BOX
    ws.row_dimensions[r].height = 46

    r = 6
    section_bar(ws, r, NC, "① 客户档案")
    r += 1
    for item in data.get("profile") or []:
        c = ws.cell(row=r, column=1, value=item.get("label", ""))
        style_cell(c, size=9.5, bold=True, color="4A5A6A", fill=BG, va="top")
        ws.merge_cells(f"B{r}:F{r}")
        val = item.get("value", "")
        if item.get("note"):
            val = f"{val}（{item['note']}）"
        c = ws.cell(row=r, column=2, value=val)
        style_cell(c, size=10, color="243247", va="top")
        for i in range(2, 7):
            ws.cell(row=r, column=i).border = BOX
        ws.row_dimensions[r].height = est_h([val], [sum(W[1:])], base=15)
        r += 1

    r += 1
    section_bar(ws, r, NC, "② 关键指标")
    r += 1
    header_row(ws, r, ["指标", "当前值", "口径 / 说明", "", "", ""])
    ws.merge_cells(f"C{r}:F{r}")
    r += 1
    for k in data.get("kpis") or []:
        c = ws.cell(row=r, column=1, value=k.get("label", ""))
        style_cell(c, size=10, bold=True, color="243247", fill=BG)
        fg, _bg = tone_color(k.get("tone"))
        c = ws.cell(row=r, column=2, value=k.get("value", ""))
        style_cell(c, size=10.5, bold=True, color=fg, ha="center")
        ws.merge_cells(f"C{r}:F{r}")
        c = ws.cell(row=r, column=3, value=k.get("note", ""))
        style_cell(c, size=9.5, color="4A5A6A", va="top")
        for i in range(3, 7):
            ws.cell(row=r, column=i).border = BOX
        ws.row_dimensions[r].height = est_h([k.get("note", "")], [sum(W[2:])], base=15)
        r += 1

    r += 1
    section_bar(ws, r, NC, "③ 卡点（按能否我方单方面推进排序）")
    r += 1
    header_row(ws, r, ["卡点", "现状", "", "状态", "影响 / 我方已给出", ""])
    ws.merge_cells(f"B{r}:C{r}")
    ws.merge_cells(f"E{r}:F{r}")
    r += 1
    for b in data.get("blockers") or []:
        c = ws.cell(row=r, column=1, value=b.get("title", ""))
        style_cell(c, size=9.5, bold=True, color="243247", fill=BG, va="top")
        ws.merge_cells(f"B{r}:C{r}")
        c = ws.cell(row=r, column=2, value=b.get("detail", ""))
        style_cell(c, size=9.5, color="334155", va="top")
        ws.cell(row=r, column=3).border = BOX
        fg, bgc = tone_color(b.get("status_tone"))
        c = ws.cell(row=r, column=4, value=b.get("status", ""))
        style_cell(c, size=9, bold=True, color=fg, fill=bgc, ha="center")
        ws.merge_cells(f"E{r}:F{r}")
        c = ws.cell(row=r, column=5, value=b.get("impact", ""))
        style_cell(c, size=9.5, color="4A5A6A", va="top")
        ws.cell(row=r, column=6).border = BOX
        ws.row_dimensions[r].height = est_h([b.get("detail", ""), b.get("impact", "")],
                                            [W[1] + W[2], W[4] + W[5]], base=15)
        r += 1

    r += 1
    section_bar(ws, r, NC, "④ 下一步动作（详见工作表 06_阶段与下一步）", fill="FDF1E3", color="7A3E05")
    r += 1
    header_row(ws, r, ["优先级", "时间", "动作", "", "闭环标准", "负责人"])
    ws.merge_cells(f"C{r}:D{r}")
    r += 1
    for a in (data.get("next_actions") or [])[:3]:
        p = str(a.get("prio", "P1"))
        pfg, pbg = ("FFFFFF", "C02626") if p.startswith("P0") else \
                   ("FFFFFF", "D97706") if p.startswith("P1") else ("FFFFFF", "94A3B8")
        c = ws.cell(row=r, column=1, value=p)
        style_cell(c, size=9.5, bold=True, color=pfg, fill=pbg, ha="center")
        c = ws.cell(row=r, column=2, value=a.get("when", ""))
        style_cell(c, size=9.5, bold=True, color="243247", va="top")
        ws.merge_cells(f"C{r}:D{r}")
        c = ws.cell(row=r, column=3, value=a.get("action", ""))
        style_cell(c, size=9.5, color="334155", va="top")
        ws.cell(row=r, column=4).border = BOX
        c = ws.cell(row=r, column=5, value=a.get("closure", ""))
        style_cell(c, size=9, color="4A5A6A", va="top")
        c = ws.cell(row=r, column=6, value=a.get("owner", meta.get("owner", "")))
        style_cell(c, size=9, color="4A5A6A", va="top", ha="center")
        ws.row_dimensions[r].height = est_h([a.get("action", ""), a.get("closure", "")],
                                            [W[2] + W[3], W[4]], base=15)
        r += 1
    ws.freeze_panes = "A4"
    page_setup(ws, NC, r)

    # ══════════════ 01_时间线 ══════════════
    ws = wb.create_sheet("01_时间线")
    W = [6, 11, 7, 16, 9, 62, 11, 14]
    set_widths(ws, W)
    title_block(ws, 8, f"跟进时间线 · 共 {len(timeline)} 条记录（按时间正序）",
                meta.get("timeline_note", "只依据上传素材；无法辨认内容不入表，另见 07_待补素材"),
                tab_color="475569")
    rows = [[rec.get("no", i), rec.get("date", ""), rec.get("time", ""),
             rec.get("source", ""), rec.get("actor", ""), rec.get("text", ""),
             rec.get("category", ""), rec.get("flag", "")]
            for i, rec in enumerate(timeline, 1)]
    r = 5
    for idx, row in enumerate(rows):
        for j, val in enumerate(row, start=1):
            c = ws.cell(row=r, column=j, value=val)
            fill = ZEBRA if idx % 2 == 1 else None
            kw = dict(size=9.5, color="243247", fill=fill, va="top")
            if j == 1:
                kw.update(ha="center", color="7A8B9C")
            elif j in (2, 3):
                kw.update(ha="center", color=INK2, bold=(j == 2))
            elif j == 4:
                kw.update(ha="center", color="64748B")
            elif j == 7:
                fg, bgc = cat_color(val)
                kw.update(ha="center", bold=True, color=fg, fill=bgc)
            elif j == 8:
                if "逾期" in str(val) or "无法辨认" in str(val):
                    kw.update(color="C02626", bold=True, ha="center")
                else:
                    kw.update(color="7A8B9C", ha="center")
            style_cell(c, **kw)
        ws.row_dimensions[r].height = est_h([row[5], row[7]], [W[5], W[7]], base=15)
        r += 1
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:H{r-1}"
    page_setup(ws, 8, r - 1)

    # ══════════════ 02_客户明确需求 ══════════════
    ws = wb.create_sheet("02_客户明确需求")
    W = [16, 74, 20, 10, 26]
    set_widths(ws, W)
    reqs = data.get("requirements") or []
    title_block(ws, 5, f"① 客户明确说过的需求与条件（{len(reqs)} 条）",
                data.get("requirements_note",
                         "只收录客户明确表述的条件；客户未提到的内容一律不列入，也不得在沟通中当作已知需求引用。"),
                tab_color="1D4ED8")
    rows = [[x.get("dimension", ""), x.get("content", ""), x.get("source", ""),
             x.get("confidence", ""), x.get("note", "")] for x in reqs]
    simple_table(ws, ["维度", "客户原话 / 纪要内容", "出处", "置信度", "备注"], rows, W,
                 styles={1: dict(bold=True, color=INK),
                         3: dict(color="64748B", ha="center"),
                         5: dict(color="4A5A6A")})
    r = 5 + len(rows)
    for ri in range(5, r):
        conf = ws.cell(row=ri, column=4)
        if conf.value == "高":
            conf.fill = PatternFill("solid", fgColor="E9F8F1")
            conf.font = Font(name=FONT, size=9.5, bold=True, color="0F7A52")
        elif conf.value == "中":
            conf.fill = PatternFill("solid", fgColor="FDF1E3")
            conf.font = Font(name=FONT, size=9.5, bold=True, color="B45309")
        elif conf.value:
            conf.fill = PatternFill("solid", fgColor="FDEEEE")
            conf.font = Font(name=FONT, size=9.5, bold=True, color="C02626")
        conf.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:E{r-1}"
    page_setup(ws, 5, r - 1)

    # ══════════════ 03_客户异议 ══════════════
    ws = wb.create_sheet("03_客户异议")
    W = [5, 58, 16, 16, 30, 34]
    set_widths(ws, W)
    objs = data.get("objections") or []
    title_block(ws, 6, f"② 客户提出的异议（{len(objs)} 条）",
                data.get("objections_note",
                         "异议 = 客户明确表达的不满、拒绝或条件对抗。状态区分未解决/待验证/已解决，避免把「已回应」当成「已解决」。"),
                tab_color="C02626")
    rows = [[x.get("no", i), x.get("content", ""), x.get("source", ""), x.get("status", ""),
             x.get("given", ""), x.get("advice", "")] for i, x in enumerate(objs, 1)]
    r = simple_table(ws, ["#", "客户异议（原话要点）", "出处", "状态", "我方已给出", "应对建议"],
                     rows, W,
                     styles={1: dict(color="7A8B9C", ha="center", bold=True),
                             3: dict(color="64748B", ha="center"),
                             5: dict(color="4A5A6A"), 6: dict(color="4A5A6A")})
    for i, x in enumerate(objs):
        c = ws.cell(row=5 + i, column=4)
        fg, bgc = tone_color(x.get("tone"))
        style_cell(c, size=9.5, bold=True, color=fg, fill=bgc, ha="center", va="top")
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:F{r-1}"
    page_setup(ws, 6, r - 1)

    # ══════════════ 04_我方承诺 ══════════════
    ws = wb.create_sheet("04_我方承诺")
    W = [5, 62, 10, 14, 18, 18, 30]
    set_widths(ws, W)
    title_block(ws, 7, f"③ 我方已做出的承诺（{len(promises)} 条）",
                data.get("promises_note",
                         "承诺必须闭环。状态由截止日与 as_of 日期自动核算：已兑现 / 未到期 / 逾期。"),
                tab_color="0F7A52")
    rows = [[p.get("no", i), p.get("content", ""), p.get("promised_on", ""), p.get("due", ""),
             p.get("source", ""), promise_state(p, as_of)[0], p.get("detail", "")]
            for i, p in enumerate(promises, 1)]
    r = simple_table(ws, ["#", "承诺内容", "承诺日", "截止 / 有效期", "出处", "状态", "备注"],
                     rows, W,
                     styles={1: dict(color="7A8B9C", ha="center", bold=True),
                             3: dict(color=INK2, ha="center"),
                             4: dict(color="4A5A6A", ha="center"),
                             5: dict(color="64748B", ha="center"),
                             7: dict(color="4A5A6A")})
    for i, p in enumerate(promises):
        st, tone = promise_state(p, as_of)
        c = ws.cell(row=5 + i, column=6)
        if tone == "risk":
            style_cell(c, size=9.5, bold=True, color="FFFFFF", fill="C02626", ha="center", va="top")
        else:
            fg, bgc = tone_color(tone)
            style_cell(c, size=9.5, bold=True, color=fg, fill=bgc, ha="center", va="top")
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:G{r-1}"
    page_setup(ws, 7, r - 1)

    # ══════════════ 05_待确认事项 ══════════════
    ws = wb.create_sheet("05_待确认事项")
    W = [5, 62, 14, 22, 32, 16]
    set_widths(ws, W)
    pends = data.get("pending") or []
    title_block(ws, 6, f"④ 尚未解决 / 需确认的问题（{len(pends)} 条）",
                data.get("pending_note",
                         "只放「还没定的事」。涉及我方内部策略的内容不得表述为「客户要求」。"),
                tab_color="6D28D9")
    rows = [[x.get("no", i), x.get("content", ""), x.get("type", ""), x.get("owner", ""),
             x.get("impact", ""), x.get("target", "")] for i, x in enumerate(pends, 1)]
    r = simple_table(ws, ["#", "待确认 / 未解决事项", "类型", "卡在谁", "影响 / 说明", "建议闭环时间"],
                     rows, W,
                     styles={1: dict(color="7A8B9C", ha="center", bold=True),
                             3: dict(color="6D28D9", fill="F3EEFF", ha="center", bold=True),
                             4: dict(color="4A5A6A", ha="center"),
                             5: dict(color="4A5A6A"),
                             6: dict(color=INK2, ha="center", bold=True)})
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:F{r-1}"
    page_setup(ws, 6, r - 1)

    # ══════════════ 06_阶段与下一步 ══════════════
    ws = wb.create_sheet("06_阶段与下一步")
    W = [16, 26, 74, 22, 18]
    set_widths(ws, W)
    title_block(ws, 5, "阶段判定 · 停滞分析 · 下一步动作",
                data.get("stage_note",
                         "阶段六段：初次接触 → 需求确认 → 送样 → 报价谈判 → 合同 → 停滞。停滞天数 = as_of − 最后一次客户实质回复日期。"),
                tab_color=GOLD)
    r = 4
    section_bar(ws, r, 5, "一、阶段进度")
    r += 1
    header_row(ws, r, ["阶段", "状态", "判定依据", "时间", "备注"])
    r += 1
    for s in stage.get("stages") or []:
        state = s.get("state", "todo")
        c = ws.cell(row=r, column=1, value=s.get("name", ""))
        style_cell(c, size=9.5, bold=True, color=INK, fill=BG if (r % 2) else None, va="top")
        c = ws.cell(row=r, column=2, value=s.get("label", ""))
        if state == "done":
            style_cell(c, size=9.5, bold=True, color="0F7A52", fill="E9F8F1", ha="center", va="top")
        elif state == "current":
            style_cell(c, size=9.5, bold=True, color="FFFFFF", fill=GOLD, ha="center", va="top")
        elif state == "stalled":
            style_cell(c, size=9.5, bold=True, color="FFFFFF", fill="C02626", ha="center", va="top")
        elif state == "risk":
            style_cell(c, size=9.5, bold=True, color="C02626", fill="FDEEEE", ha="center", va="top")
        else:
            style_cell(c, size=9.5, color="8A97A6", ha="center", va="top")
        c = ws.cell(row=r, column=3, value=s.get("basis", ""))
        style_cell(c, size=9.5, color="334155", va="top")
        c = ws.cell(row=r, column=4, value=s.get("date", ""))
        style_cell(c, size=9.5, color=INK2, ha="center", va="top")
        c = ws.cell(row=r, column=5, value=s.get("note", ""))
        style_cell(c, size=9, color="4A5A6A", va="top")
        ws.row_dimensions[r].height = est_h([s.get("basis", ""), s.get("note", "")],
                                            [W[2], W[4]], base=15)
        r += 1

    r += 1
    section_bar(ws, r, 5, "二、停滞原因与风险", fill="FDEEEE", color="C02626")
    r += 1
    stall_rows = []
    for i, reason in enumerate(stall.get("reasons") or [], 1):
        stall_rows.append((f"停滞原因 {i}", reason.get("title", ""), reason.get("detail", ""), False))
    for rr in stall.get("risks") or []:
        stall_rows.append(("滑向流失的风险", rr.get("title", ""), rr.get("detail", ""), True))
    if stall.get("days") is not None or computed_days is not None:
        stall_rows.insert(0, ("停滞天数",
                              f"{stall_days if stall_days is not None else computed_days} 天",
                              f"as_of {as_of.isoformat()} − 最后一次客户实质回复（{(last_rec or {}).get('date','—')} {(last_rec or {}).get('time','')}）",
                              False))
    for label, ttl, det, is_risk in stall_rows:
        c = ws.cell(row=r, column=1, value=label)
        style_cell(c, size=9.5, bold=True,
                   color="C02626" if is_risk else "7A3E05",
                   fill="FDEEEE" if is_risk else GOLD_L, va="top")
        ws.merge_cells(f"B{r}:C{r}")
        c = ws.cell(row=r, column=2, value=ttl)
        style_cell(c, size=9.5, bold=True, color="243247", va="top")
        ws.cell(row=r, column=3).border = BOX
        ws.merge_cells(f"D{r}:E{r}")
        c = ws.cell(row=r, column=4, value=det)
        style_cell(c, size=9, color="4A5A6A", va="top")
        ws.cell(row=r, column=5).border = BOX
        ws.row_dimensions[r].height = est_h([det], [W[3] + W[4]], base=15)
        r += 1

    r += 1
    section_bar(ws, r, 5, "三、下一步动作（含联动入口）", fill="FDF1E3", color="7A3E05")
    r += 1
    header_row(ws, r, ["优先级", "时间", "动作", "闭环标准", "联动入口"])
    r += 1
    for a in data.get("next_actions") or []:
        p = str(a.get("prio", "P1"))
        pfg, pbg = ("FFFFFF", "C02626") if p.startswith("P0") else \
                   ("FFFFFF", "D97706") if p.startswith("P1") else ("FFFFFF", "94A3B8")
        c = ws.cell(row=r, column=1, value=p)
        style_cell(c, size=9.5, bold=True, color=pfg, fill=pbg, ha="center", va="top")
        c = ws.cell(row=r, column=2, value=a.get("when", ""))
        style_cell(c, size=9.5, bold=True, color=INK2, ha="center", va="top")
        c = ws.cell(row=r, column=3, value=a.get("action", ""))
        style_cell(c, size=9.5, color="243247", va="top")
        c = ws.cell(row=r, column=4, value=a.get("closure", ""))
        style_cell(c, size=9.5, color="4A5A6A", va="top")
        c = ws.cell(row=r, column=5, value=a.get("link_label", ""))
        style_cell(c, size=9.5, bold=True, color="1D4ED8", fill="EEF3FF", ha="center", va="top")
        ws.row_dimensions[r].height = est_h([a.get("action", ""), a.get("closure", "")],
                                            [W[2], W[3]], base=15)
        r += 1
    page_setup(ws, 5, r - 1)

    # ══════════════ 07_待补素材 ══════════════
    ws = wb.create_sheet("07_待补素材")
    W = [5, 26, 52, 40, 12]
    set_widths(ws, W)
    unks = data.get("unknowns") or []
    title_block(ws, 5, f"无法辨认 / 缺失内容（{len(unks)} 项，需用户确认）",
                data.get("unknowns_note",
                         "以下内容未计入任何事实类清单，避免猜测混入正文。补齐后重新生成即可。"),
                tab_color="B45309")
    rows = [[x.get("no", i), x.get("source", ""), x.get("issue", ""), x.get("need", ""),
             x.get("status", "待补")] for i, x in enumerate(unks, 1)]
    r = simple_table(ws, ["#", "素材", "问题描述", "需要您做什么", "状态"], rows, W,
                     styles={1: dict(color="7A8B9C", ha="center", bold=True),
                             2: dict(bold=True, color=INK)})
    for i, x in enumerate(unks):
        c = ws.cell(row=5 + i, column=5)
        if str(x.get("status", "待补")) == "待补":
            style_cell(c, size=9.5, bold=True, color="FFFFFF", fill="C02626", ha="center", va="top")
        else:
            style_cell(c, size=9.5, bold=True, color="B45309", fill="FDF1E3", ha="center", va="top")
    ws.freeze_panes = "A5"
    page_setup(ws, 5, r - 1)

    # ══════════════ 08_关键字段确认 ══════════════
    ws = wb.create_sheet("08_关键字段确认")
    W = [20, 50, 22, 11, 20, 14]
    set_widths(ws, W)
    fields = data.get("field_confirm") or []
    title_block(ws, 6, f"关键字段确认表（{len(fields)} 项）",
                data.get("field_note",
                         "生成跟进卡前逐条核对。右侧「用户确认」列可下拉选择；标为待确认的字段不进入归档版本。"),
                tab_color=INK2)
    rows = [[f.get("field", ""), f.get("value", ""), f.get("source", ""),
             f.get("confidence", ""), f.get("flag", "—"), "待确认"] for f in fields]
    r = simple_table(ws, ["字段", "提取值", "出处", "置信度", "标记", "用户确认"], rows, W,
                     styles={1: dict(bold=True, color=INK),
                             3: dict(color="64748B", ha="center"),
                             5: dict(color="A0AEC0", ha="center")},
                     heights=None)
    for i, f in enumerate(fields):
        row_i = 5 + i
        ws.row_dimensions[row_i].height = 26
        conf = ws.cell(row=row_i, column=4)
        v = str(conf.value or "")
        if v == "高":
            style_cell(conf, size=9.5, bold=True, color="0F7A52", fill="E9F8F1", ha="center")
        elif v == "中":
            style_cell(conf, size=9.5, bold=True, color="B45309", fill="FDF1E3", ha="center")
        elif v:
            style_cell(conf, size=9.5, bold=True, color="FFFFFF", fill="C02626", ha="center")
        flag = ws.cell(row=row_i, column=5)
        if str(flag.value or "—") != "—":
            style_cell(flag, size=9.5, bold=True, color="C02626", fill="FDEEEE", ha="center")
        else:
            style_cell(flag, size=9.5, color="A0AEC0", ha="center")
        uc = ws.cell(row=row_i, column=6)
        style_cell(uc, size=9.5, bold=True, color="B45309", fill="FDF1E3", ha="center")

    if fields:
        dv = DataValidation(type="list", formula1='"已确认,待确认"', allow_blank=False,
                            showDropDown=False, errorTitle="请选择",
                            error="只能选择「已确认」或「待确认」")
        ws.add_data_validation(dv)
        dv.add(f"F5:F{4 + len(fields)}")
        ws.conditional_formatting.add(
            f"F5:F{4 + len(fields)}",
            FormulaRule(formula=['$F5="已确认"'],
                        fill=PatternFill("solid", fgColor="E9F8F1"),
                        font=Font(name=FONT, size=9.5, bold=True, color="0F7A52")))
        ws.conditional_formatting.add(
            f"A5:F{4 + len(fields)}",
            FormulaRule(formula=['$F5="已确认"'],
                        fill=PatternFill("solid", fgColor="F4FBF7")))
    ws.freeze_panes = "A5"
    page_setup(ws, 6, r - 1)

    # ══════════════ 09_图例与口径 ══════════════
    ws = wb.create_sheet("09_图例与口径")
    W = [18, 16, 74]
    set_widths(ws, W)
    title_block(ws, 3, "图例 · 整理口径 · 免责说明",
                meta.get("disclaimer_note",
                         "本工作簿由「整理客户跟进记录」流程基于上传素材生成。"),
                tab_color="64748B")
    r = 4
    section_bar(ws, r, 3, "一、分类图例")
    r += 1
    header_row(ws, r, ["分类", "色块", "含义 / 使用规则"])
    r += 1
    legend = data.get("legend") or [
        {"name": "客户需求", "desc": "客户明确说过的需求与条件。未说过的一律不补。"},
        {"name": "客户异议", "desc": "客户明确表达的不满、拒绝或条件对抗。"},
        {"name": "我方承诺", "desc": "我方已做出的口头或书面承诺，必须闭环。"},
        {"name": "待确认", "desc": "尚未解决或需确认的问题。"},
        {"name": "客户动作", "desc": "客户发起的行为，不做需求解读。"},
        {"name": "我方动作", "desc": "我方发起的行为，不含承诺。"},
        {"name": "系统提醒", "desc": "系统按规则自动生成的计算项。"},
    ]
    for item in legend:
        fg, bgc = cat_color(item.get("name"))
        c = ws.cell(row=r, column=1, value=item.get("name", ""))
        style_cell(c, size=9.5, bold=True, color=fg, fill=bgc)
        c = ws.cell(row=r, column=2, value="■■■")
        style_cell(c, size=9.5, bold=True, color=fg, fill=bgc, ha="center")
        c = ws.cell(row=r, column=3, value=item.get("desc", ""))
        style_cell(c, size=9.5, color="4A5A6A", va="center")
        ws.row_dimensions[r].height = est_h([item.get("desc", "")], [W[2]], base=15)
        r += 1

    r += 1
    section_bar(ws, r, 3, "二、整理口径")
    r += 1
    for i, rule in enumerate(data.get("rules") or [], 1):
        title = rule.get("title", "") if isinstance(rule, dict) else f"规则 {i}"
        desc = rule.get("desc", "") if isinstance(rule, dict) else str(rule)
        c = ws.cell(row=r, column=1, value=f"{i}. {title}")
        style_cell(c, size=9.5, bold=True, color=INK, fill=BG, va="top")
        ws.merge_cells(f"B{r}:C{r}")
        c = ws.cell(row=r, column=2, value=desc)
        style_cell(c, size=9.5, color="4A5A6A", va="top")
        ws.cell(row=r, column=3).border = BOX
        ws.row_dimensions[r].height = est_h([desc], [W[1] + W[2]], base=15)
        r += 1

    r += 1
    section_bar(ws, r, 3, "三、数据来源与生成信息")
    r += 1
    for s in data.get("sources") or []:
        c = ws.cell(row=r, column=1, value=f"{s.get('no','')} {s.get('name','')}".strip())
        style_cell(c, size=9.5, bold=True, color="4A5A6A", fill=BG)
        c = ws.cell(row=r, column=2, value=s.get("detail", ""))
        style_cell(c, size=9.5, color="243247")
        c = ws.cell(row=r, column=3, value=s.get("count", ""))
        style_cell(c, size=9, color="7A8B9C", ha="center")
        ws.row_dimensions[r].height = est_h([s.get("detail", "")], [W[1]], base=15)
        r += 1
    for k, v in (("生成时间", meta.get("generated_at", "")),
                 ("as_of 日期", as_of.isoformat()),
                 ("记录总数", f"{len(timeline)} 条")):
        c = ws.cell(row=r, column=1, value=k)
        style_cell(c, size=9.5, bold=True, color="4A5A6A", fill=BG)
        c = ws.cell(row=r, column=2, value=v)
        style_cell(c, size=9.5, color="243247")
        ws.cell(row=r, column=3).border = BOX
        r += 1

    r += 1
    ws.merge_cells(f"A{r}:C{r}")
    c = ws.cell(row=r, column=1, value=meta.get("disclaimer", ""))
    style_cell(c, size=9, color="8A97A6", fill=GOLD_L, va="center")
    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
    for i in range(1, 4):
        ws.cell(row=r, column=i).fill = PatternFill("solid", fgColor=GOLD_L)
    ws.row_dimensions[r].height = 40
    page_setup(ws, 3, r)

    refit_heights(wb)

    wb.properties.title = meta.get("title", "客户跟进卡")
    wb.properties.subject = meta.get("subject", "客户跟进记录整理")
    wb.properties.creator = meta.get("creator", "chem-customer-followup")
    wb.properties.description = meta.get("description", "")
    wb.properties.keywords = meta.get("keywords", "")
    wb.save(out_path)
    return out_path


def main():
    ap = argparse.ArgumentParser(description="由 followup.json 生成跟进卡 Excel")
    ap.add_argument("input", help="followup.json 路径")
    ap.add_argument("-o", "--out", required=True, help="输出 .xlsx 路径")
    args = ap.parse_args()

    import json
    with open(args.input, encoding="utf-8-sig") as f:
        data = json.load(f)

    errors, warnings = validate(data)
    for w in warnings:
        print(f"[warn] {w}")
    if errors:
        for e in errors:
            print(f"[error] {e}", file=sys.stderr)
        sys.exit(1)

    build(data, args.out)
    print(f"[ok] xlsx -> {args.out}")


if __name__ == "__main__":
    main()
