#!/usr/bin/env python3
"""build_xlsx.py — research.json → 多 Sheet Excel 工作簿（与 HTML 同源同数）

用法:
    python scripts/build_xlsx.py research.json --out out/

固定结构（12 表，「期现-产业基本面 / 快速看懂一个品种」标准交付）：
  00 封面与阅读指引    01 产品百科与规格    02 品牌牌号与规格对照
  03 工艺路线与单耗    04 产业链节点        05 企业与产能
  06 下游应用与需求    07 贸易与区域市场    08 期货合约与交割（无期货→声明页）
  09 关键研究指标      10 来源清单          11 缺口声明与核验记录
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# ---------------- 样式 ----------------
C_INK = "1F2A3A"
C_ACC = "2F7FE0"
C_ACC2 = "22D3B2"
C_WARN = "F5A524"
C_GAP = "FFF4E5"
C_HEAD = "132033"
C_ZEBRA = "F4F7FB"
C_LINE = "D5DEEA"

F_TITLE = Font(name="微软雅黑", size=18, bold=True, color=C_INK)
F_SUB = Font(name="微软雅黑", size=10, color="5B6B80")
F_SECT = Font(name="微软雅黑", size=12, bold=True, color="FFFFFF")
F_HEAD = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
F_BODY = Font(name="微软雅黑", size=10, color=C_INK)
F_BODY_B = Font(name="微软雅黑", size=10, bold=True, color=C_INK)
F_SRC = Font(name="微软雅黑", size=9, color="6D829D")
F_GAP = Font(name="微软雅黑", size=10, bold=True, color="B45309")
F_NUM = Font(name="Consolas", size=10, color=C_INK)

FILL_SECT = PatternFill("solid", fgColor=C_HEAD)
FILL_HEAD = PatternFill("solid", fgColor=C_ACC)
FILL_ZEBRA = PatternFill("solid", fgColor=C_ZEBRA)
FILL_GAP = PatternFill("solid", fgColor=C_GAP)
FILL_KPI = PatternFill("solid", fgColor="EAF2FE")

THIN = Side(style="thin", color=C_LINE)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
WRAP_C = Alignment(wrap_text=True, vertical="center")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

GAP_TEXT = "缺口"
GAP_MARKS = re.compile(r"^(未查到|无公开|缺口|待核|—$)")


class Sheet:
    """极简表格写入器：标题区 + 分区 + 表头 + 正文，自动斑马纹/缺口底色/行高。"""

    def __init__(self, wb: Workbook, name: str, title: str, subtitle: str = "", widths=None, freeze="A6"):
        self.ws = wb.create_sheet(name)
        self.ws.sheet_view.showGridLines = False
        self.row = 1
        self.ncol = max(len(widths) if widths else 8, 1)
        self.ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=self.ncol)
        c = self.ws.cell(1, 1, title)
        c.font = F_TITLE
        c.alignment = Alignment(vertical="center")
        self.ws.row_dimensions[1].height = 30
        self.ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=self.ncol)
        c = self.ws.cell(2, 1, subtitle)
        c.font = F_SUB
        c.alignment = WRAP_C
        self.ws.row_dimensions[2].height = 30
        self.ws.row_dimensions[3].height = 6
        self.row = 4
        if widths:
            for i, w in enumerate(widths, start=1):
                self.ws.column_dimensions[get_column_letter(i)].width = w
        if freeze:
            self.ws.freeze_panes = freeze

    def section(self, text: str, note: str = ""):
        self.row += 1
        self.ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=self.ncol)
        c = self.ws.cell(self.row, 1, ("  " + text + ("　—　" + note if note else "")))
        c.font = F_SECT
        c.fill = FILL_SECT
        c.alignment = Alignment(vertical="center")
        self.ws.row_dimensions[self.row].height = 24
        self.row += 1

    def header(self, cols):
        for i, h in enumerate(cols, start=1):
            c = self.ws.cell(self.row, i, h)
            c.font = F_HEAD
            c.fill = FILL_HEAD
            c.alignment = CENTER
            c.border = BOX
        self.ws.row_dimensions[self.row].height = 26
        self.row += 1
        self.cols = cols

    def _put(self, r, i, v, num=False):
        c = self.ws.cell(r, i)
        if v is None:
            v = GAP_TEXT
        if isinstance(v, str) and GAP_MARKS.search(v.strip()):
            c.value = v.replace("未查到", "缺口(未查到)") if v.startswith("未查到") else v
            c.font = F_GAP
            c.fill = FILL_GAP
        else:
            c.value = v
            c.font = F_NUM if (num and isinstance(v, (int, float))) else F_BODY
        c.alignment = WRAP
        c.border = BOX
        return c

    def rows(self, rows, numeric=(), height=None):
        for ri, r in enumerate(rows):
            zebra = ri % 2 == 1
            maxlen = 0
            for i, v in enumerate(r, start=1):
                if i > self.ncol:
                    break
                c = self._put(self.row, i, v, num=(i in numeric))
                if zebra and c.fill.fgColor.rgb in (None, "00000000"):
                    c.fill = FILL_ZEBRA
                if isinstance(v, str):
                    maxlen = max(maxlen, len(v))
            self.ws.row_dimensions[self.row].height = height or max(20, min(78, 17 + 9 * max(0, maxlen // 42)))
            self.row += 1

    def kv(self, pairs, kwidth_total=2):
        for k, v in pairs:
            self.ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=kwidth_total)
            c1 = self.ws.cell(self.row, 1, k)
            c1.font = F_BODY_B
            c1.fill = FILL_ZEBRA
            c1.alignment = WRAP_C
            c1.border = BOX
            self.ws.merge_cells(start_row=self.row, start_column=kwidth_total + 1, end_row=self.row, end_column=self.ncol)
            c2 = self.ws.cell(self.row, kwidth_total + 1)
            if v is None or (isinstance(v, str) and GAP_MARKS.search(v.strip())):
                c2.value = v or GAP_TEXT
                c2.font = F_GAP
                c2.fill = FILL_GAP
            else:
                c2.value = v
                c2.font = F_BODY
            c2.alignment = WRAP
            c2.border = BOX
            self.ws.row_dimensions[self.row].height = max(20, min(90, 17 + 12 * (max(len(str(v or "")), len(str(k))) // 60)))
            self.row += 1

    def note(self, text, color="6D829D"):
        self.ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=self.ncol)
        c = self.ws.cell(self.row, 1, text)
        c.font = Font(name="微软雅黑", size=9, color=color, italic=(color == "B45309"))
        c.alignment = WRAP
        self.ws.row_dimensions[self.row].height = max(18, min(70, 16 + 13 * (len(text) // 80)))
        self.row += 1

    def blank(self, n=1):
        self.row += n


# ---------------- 工具 ----------------
def S(data, sid):
    for s in data.get("sources", []):
        if s.get("id") == sid:
            t = s.get("org") or ""
            d = s.get("date") or ""
            u = s.get("url") or ""
            return f"{t}｜{d}" + (f"｜{u}" if u else "")
    return sid


def srcs(data, ids):
    if not ids:
        return ""
    return "\n".join(S(data, i) for i in ids)


def find_block(data, panel_id, btype, idx=0):
    for p in data.get("panels", []):
        if p.get("id") != panel_id:
            continue
        hits = [b for b in p.get("body", []) if b.get("type") == btype]
        if len(hits) > idx:
            return hits[idx]
    return None


def panel(data, pid):
    for p in data.get("panels", []):
        if p.get("id") == pid:
            return p
    return None


def nodes_of(data, col=None, kind=None):
    out = []
    for n in data.get("chain", {}).get("nodes", []):
        if col and n.get("col") != col:
            continue
        if kind and n.get("kind") != kind:
            continue
        out.append(n)
    return out


def walk_nodes(data):
    """展开 nodes 为扁平列表（含 children 摘要，child 仅名字与 cap）。"""
    rows = []

    def rec(n, depth, parent):
        rows.append((n, depth, parent))
        for ch in n.get("children", []) or []:
            cn = {"id": ch.get("id"), "name": ch.get("name"), "meta": ch.get("cap") or "", "_child": True}
            rec(cn, depth + 1, n.get("name"))
            if not ch.get("id"):
                cn["id"] = (n.get("id") or "") + "›" + str(ch.get("name"))

    for n in data.get("chain", {}).get("nodes", []):
        rec(n, 0, None)
    return rows


# ---------------- 各表 ----------------
def sheet_cover(wb, data):
    m = data.get("meta", {})
    p = data.get("product", {})
    name = p.get("name_cn", "品种")
    sh = Sheet(wb, "00 封面与指引", f"{name}（{p.get('name_en','')}）品种研究入门工作簿",
               "面向化工期现研究 / 现货贸易 / 采购销售的品种入门底稿｜事实·工艺参数·公开统计·研究判断 四类信息分层标注",
               widths=[26, 30, 30, 30, 26, 18], freeze=None)
    sh.ws.merge_cells("A4:F4")
    c = sh.ws.cell(4, 1, f"CAS {p.get('cas','—')}　｜　{p.get('main_spec','')}　｜　" +
                   ("有场内期货：" + (data.get("research", {}).get("futures", {}).get("exchange") or "") if p.get("has_futures") else "无场内期货"))
    c.font = Font(name="微软雅黑", size=11, bold=True, color=C_ACC)
    c.alignment = WRAP_C
    sh.ws.row_dimensions[4].height = 22
    sh.row = 5

    sh.section("一句话定位")
    sh.ws.merge_cells(start_row=sh.row, start_column=1, end_row=sh.row, end_column=6)
    c = sh.ws.cell(sh.row, 1, p.get("oneliner", ""))
    c.font = Font(name="微软雅黑", size=11, color=C_INK)
    c.alignment = WRAP
    sh.ws.row_dimensions[sh.row].height = 42
    sh.row += 1

    sh.section("使用前必读")
    for t in [
        "1) 本簿只做「入门结构化」，不做行情预测；所有产能、产量、占比均注明统计时点与来源，缺数一律写「缺口」，不做行业均值推算。",
        "2) 判断与事实分离：工艺路线中的单耗、反应条件多为公开工艺资料/教材典型值，归入「典型工艺参数」；企业与产能多为公开披露或行业统计；两者不可混用。",
        "3) 期货规则以交易所原文为准：合约要素变动（保证金、涨跌停、交割品级调整）需在使用前复核交易所公告。",
        "4) 现货价格必须六项齐全（数值·规格·地区·含税·日期·来源）才可引用；本簿只保留结构性价格锚点，逐日行情请进入「看当前价格处在历史什么位置」场景。",
    ]:
        sh.note(t, color=C_INK)
    sh.blank()

    sh.section("研究指标框架速览")
    sh.header(["维度", "核心指标", "口径/频次", "是否已有数", "用途", ""])
    rows = []
    for f in data.get("frame", []):
        items = f.get("items", [])
        rows.append([f.get("name"), "\n".join(i.get("name", "") for i in items),
                     "\n".join(i.get("caliber", "") for i in items),
                     f"{len(items) - sum(1 for i in items if i.get('gap'))}/{len(items)}",
                     f.get("value_note", ""), ""])
    sh.rows(rows)
    sh.blank()

    sh.section("本簿构成的 12 张表")
    sh.header(["表", "内容", "对应研究问题", "", "", ""])
    sh.rows([
        ["01 产品百科与规格", "名称/别名/CAS/HS/物性/牌号规格/计价单位", "1) 它是什么"],
        ["02 品牌牌号与规格对照", "生产商牌号、执行标准、典型指标、对应交割等级", "1) 规格与等级"],
        ["03 工艺路线与单耗", "各路线流程、原料、典型参数、投资、碳排、差异", "2) 生产工艺路线"],
        ["04 产业链节点", "节点属性、产能/价格摘要、来源（与 HTML 图谱同源）", "产业链图谱"],
        ["05 企业与产能", "企业-基地-产能-工艺-来源；区域结构", "4) 企业与产能"],
        ["06 下游应用与需求", "下游分项占比、单耗系数、驱动变量", "3) 下游与终端用途"],
        ["07 贸易与区域市场", "HS 编码、进出口、口岸、区域价差、集散地", "5) 贸易与区域市场"],
        ["08 期货合约与交割", "合约要素、交割标准品与替代品、期现关系", "6) 期货与交割"],
        ["09 关键研究指标", "供应/需求/库存/价格成本/期现 五类指标与频率", "7) 跟踪哪些指标"],
        ["10 来源清单", "编号、机构、标题、URL、日期、用途、置信度", "来源与时间"],
        ["11 缺口声明与核验", "未覆盖内容、核验记录、口径提醒", "不做数据填充的声明"],
    ])
    sh.blank()
    sh.section("元数据")
    sh.kv([
        ("数据基准日", m.get("data_asof")),
        ("生成日期", m.get("generated")),
        ("Skill 版本", m.get("version")),
        ("来源条数", str(len(data.get("sources", [])))),
        ("缺口条数", str(len(data.get("gaps", [])))),
        ("免责", m.get("footnote")),
    ], kwidth_total=2)
    return sh


def sheet_encyclopedia(wb, data):
    p = data.get("product", {})
    b = find_block(data, "identity_specs", "cards", 0)
    sh = Sheet(wb, "01 产品百科", f"01 产品百科与规格 · {p.get('name_cn','')}",
               "身份字段来自 CAS/UN/GB/HS 等权威口径；牌号与等级区分「国标等级」与「商业牌号」",
               widths=[22, 34, 34, 30, 18])
    sh.section("身份识别")
    sh.header(["字段", "内容", "补充", "", ""])
    rows = [[c.get("k"), c.get("v"), c.get("sub") or "", "", ""] for c in p.get("card", [])]
    sh.rows(rows)
    sh.section("别名与命名")
    sh.kv([
        ("中文通用名", p.get("name_cn")),
        ("英文名", p.get("name_en")),
        ("中文别名", "、".join(p.get("aliases", []))),
        ("英文别名", "、".join(p.get("synonyms_en", []))),
        ("CAS 登记号", p.get("cas")),
        ("HS 编码", "\n".join(f"{h.get('code')} {h.get('desc','')}" for h in p.get("hs", []))),
        ("UN 编号/危险货物", p.get("un")),
        ("危险化学品目录", p.get("hazard_class")),
    ], kwidth_total=1)
    sh.section("物化性质（典型值）")
    sh.kv([
        ("分子式", p.get("formula")), ("分子量", p.get("mw")),
        ("外观", p.get("appearance")), ("沸点", p.get("boiling_point")),
        ("熔点/凝固点", p.get("freezing_point")), ("密度", p.get("density")),
        ("闪点", p.get("flash_point")), ("溶解性", p.get("solubility")),
    ], kwidth_total=1)
    if b:
        sh.section("规格、等级与计价单位")
        sh.header(["分项", "内容", "备注", "", ""])
        rows = []
        for it in b.get("items", []):
            lines = "\n".join(f"{l.get('k')}：{l.get('v')}" for l in it.get("lines", []))
            rows.append([it.get("title"), lines, it.get("lead") or "", "", ""])
        sh.rows(rows)
        for it in b.get("items", []):
            if it.get("src"):
                sh.note("来源：" + "；".join(S(data, s) for s in it["src"]))
    ic = p.get("identity_confidence", {})
    sh.section("身份置信度")
    sh.kv([("等级", ic.get("level")), ("依据", ic.get("basis")), ("来源", srcs(data, ic.get("src")))], kwidth_total=1)
    return sh


def sheet_grades(wb, data):
    p = data.get("product", {})
    sh = Sheet(wb, "02 牌号与规格", "02 品牌牌号与规格对照",
               "牌号=生产商商业名称；等级=标准/交割等级。两者口径不同，不可互相替代。",
               widths=[18, 20, 16, 30, 30, 26])
    b = find_block(data, "identity_specs", "table", 0)
    if b:
        sh.section(b.get("caption", "规格对照表"))
        cols = [c.get("k", str(c)) for c in b.get("columns", [])]
        sh.header(cols + [""] * max(0, 6 - len(cols)))
        rows = []
        for r in b.get("rows", []):
            row = []
            for c in b.get("columns", []):
                k = c.get("k", str(c))
                if k == "src":
                    row.append(srcs(data, r.get("src") or r.get("srcs")))
                else:
                    row.append(r.get(k))
            rows.append(row)
        sh.rows(rows)
    else:
        sh.section("规格对照")
        sh.note("研究稿未提供规格对照表（缺），请以企业质量证明书/交易所交割细则为准。", color="B45309")
    sh.blank()
    sh.section("等级与交割的对应关系（判断项，需逐条核实）")
    sh.note("说明：国标等级 → 交易所允许交割等级 → 现货市场主流报价规格，三者常见差异在纯度、水分、色度与包装方式；"
            "任何“某牌号一定可交割”的结论都必须以交易所交割细则与质检报告为准。", color=C_INK)
    fr = data.get("research", {}).get("futures", {}) or {}
    rows = []
    for it in fr.get("items", []):
        for r in it.get("rows", []):
            if re.search(r"交割品|品级|交割标准|质量", str(r.get("k", ""))):
                rows.append([it.get("title"), r.get("k"), r.get("v"), srcs(data, it.get("src"))])
    if rows:
        sh.header(["合约/场景", "条款", "内容", "来源"])
        sh.rows(rows)
    else:
        sh.note("无期货/未提供交割条款：本品种不做期现等级对应（见 08 表）。", color="B45309")
    return sh


def sheet_routes(wb, data):
    p = data.get("product", {})
    sh = Sheet(wb, "03 工艺路线", f"03 工艺路线与单耗 · {p.get('name_cn','')}",
               "「典型工艺参数」为公开工艺资料/教材值，不代表任一企业实际；「公开统计」需来源；差异判断单列",
               widths=[18, 16, 34, 20, 14, 30])
    b = find_block(data, "routes", "routes", 0)
    if not b:
        sh.section("工艺路线")
        sh.note("研究稿未提供工艺路线（缺）。", color="B45309")
        return sh
    sh.section("路线一览")
    sh.header(["路线", "特点", "流程", "产能占比", "状态", "来源"])
    rows = []
    for r in b.get("items", []):
        rows.append([r.get("name"), r.get("brief"), " → ".join(s.get("name", "") for s in r.get("steps", [])),
                     r.get("share") or "", r.get("status") or "", srcs(data, r.get("src"))])
    sh.rows(rows)

    sh.section("典型参数矩阵（统一口径，便于横向比较）")
    keys, seen = [], set()
    for r in b.get("items", []):
        for x in r.get("params", []):
            if x.get("k") not in seen:
                seen.add(x.get("k"))
                keys.append(x.get("k"))
    if keys:
        for ki, k in enumerate(keys):
            sh.header(["参数", "口径/来源类型"] + [r.get("name") for r in b.get("items", [])] + [""] * 2)
            row = [k, ""]
            for r in b.get("items", []):
                val, note = "", ""
                for x in r.get("params", []):
                    if x.get("k") == k:
                        val = x.get("v") or ""
                        note = x.get("note") or ""
                row.append((str(val) + (f"（{note}）" if note else "")) or None)
            row += ["", ""]
            sh.rows([row])
            sh.blank()

    sh.section("路线说明（原文）")
    for r in b.get("items", []):
        sh.kv([(r.get("name"), (r.get("body") or "") + ("\n来源：" + srcs(data, r.get("src")) if r.get("src") else ""))], kwidth_total=1)
    return sh


def sheet_chain(wb, data):
    p = data.get("product", {})
    sh = Sheet(wb, "04 产业链节点", "04 产业链节点（与 HTML 图谱同源）",
               "层级=距本品种的上游/下游层数；产能与价格摘要仅作索引，逐项数据见对应研究表",
               widths=[16, 22, 30, 24, 22, 26, 18])
    cols = {c.get("id"): c.get("name") for c in data.get("chain", {}).get("cols", [])}
    sh.section("列（产业链骨架顺序）")
    sh.header(["序号", "列 id", "中文名", "", "", "", ""])
    sh.rows([[i + 1, c.get("id"), c.get("name"), "", "", "", ""] for i, c in enumerate(data.get("chain", {}).get("cols", []))])
    sh.section("节点明细")
    sh.header(["层级缩进", "节点", "所属列", "关键数据", "产能/规模", "价格摘要", "期货/标记"])
    rows = []
    for n, depth, parent in walk_nodes(data):
        rows.append([("　" * depth + ("> " if depth else "")), n.get("name"), cols.get(n.get("col"), n.get("col") or ""),
                     n.get("meta") or n.get("summary") or "", n.get("cap") or "", n.get("spot") or "",
                     "有期货" if n.get("has_futures") else (n.get("kind") or "")])
    sh.rows(rows)
    sh.section("节点要点与来源")
    sh.header(["节点", "研究要点", "来源", "", "", "", ""])
    rows = []
    for n in data.get("chain", {}).get("nodes", []):
        r = n.get("research", {}) or {}
        pts = "\n".join("• " + x for x in (r.get("points") or []))
        rows.append([n.get("name"), pts or (n.get("summary") or ""), srcs(data, r.get("sources") or n.get("src")), "", "", "", ""])
    sh.rows(rows)
    return sh


def sheet_capacity(wb, data):
    p = data.get("product", {})
    sh = Sheet(wb, "05 企业与产能", f"05 生产企业与公开产能 · {p.get('name_cn','')}",
               "产能单位为万吨/年（或吨/年，见列头）；含在建/规划须在状态列区分；数据缺失写缺口，不以估算补齐",
               widths=[6, 24, 22, 14, 10, 16, 28, 26])
    b = find_block(data, "capacity", "companies", 0)
    if b:
        sh.section(b.get("caption") or "主要生产企业")
        sh.header(["#", "企业", "生产基地/地区", "公开产能", "单位", "工艺/装置", "备注", "来源"])
        rows = []
        for i, c in enumerate(b.get("items", []), 1):
            rows.append([i, c.get("name"), c.get("site"), c.get("cap"), c.get("unit") or b.get("unit") or "万吨/年",
                         c.get("process") or "", c.get("note") or "", srcs(data, c.get("src"))])
        sh.rows(rows, numeric=(1, 4))
        tot = b.get("total")
        if tot:
            sh.note(f"合计（所列表企业口径）：{tot} {b.get('unit','万吨')}；口径说明：{b.get('caliber') or '见来源'}")

    reg = find_block(data, "capacity", "regions", 0)
    if reg:
        sh.section("区域分布")
        sh.header(["区域", "范围/细分", "数值", "单位", "占比", "说明", "结构", "来源"])
        rows = [[r.get("name"), r.get("scope"), r.get("value"), r.get("unit"), r.get("share"), r.get("note"),
                 "；".join(f"{m.get('n')}{m.get('p')}%" for m in (r.get("mix") or [])), srcs(data, reg.get("src"))]
                for r in reg.get("items", [])]
        sh.rows(rows)

    cs = find_block(data, "capacity", "table", 0)
    if cs:
        sh.section(cs.get("caption") or "产能结构")
        cols = [c.get("k", str(c)) for c in cs.get("columns", [])]
        sh.header(cols + [""] * max(0, 8 - len(cols)))
        rows = []
        for r in cs.get("rows", []):
            row = []
            for c in cs.get("columns", []):
                k = c.get("k", str(c))
                row.append(srcs(data, r.get("src")) if k == "src" else r.get(k))
            rows.append(row)
        sh.rows(rows)

    sh.blank()
    sh.section("产能调查注意（研究判断）")
    for t in [
        "名义产能 vs 有效产能：长期停车、半负荷与置换关停会显著改变有效供给，行业总产能不可直接当供给用。",
        "单套与合计：同一集团多套装置需区分“集团合计”与“单套规模”，联产装置只计目标产品那一环。",
        "投建状态：环评/备案/在建/试车/投产 四个状态对应完全不同的供给含义，不要合并计数。",
    ]:
        sh.note(t, color=C_INK)
    return sh


def sheet_demand(wb, data):
    p = data.get("product", {})
    sh = Sheet(wb, "06 下游与需求", f"06 下游应用与需求 · {p.get('name_cn','')}",
               "分项占比需注明统计机构与时点；单耗系数为工艺典型值，实际随装置与转化率浮动",
               widths=[20, 14, 16, 18, 30, 26, 16])
    b = find_block(data, "demand", "table", 0)
    if b:
        sh.section(b.get("caption") or "下游应用结构")
        cols = [c.get("k", str(c)) for c in b.get("columns", [])]
        sh.header(cols + [""] * max(0, 7 - len(cols)))
        rows = []
        for r in b.get("rows", []):
            row = []
            for c in b.get("columns", []):
                k = c.get("k", str(c))
                row.append(srcs(data, r.get("src")) if k == "src" else r.get(k))
            rows.append(row)
        sh.rows(rows)
    else:
        sh.section("下游结构")
        sh.note("研究稿未提供下游结构表（缺）。", color="B45309")

    b2 = find_block(data, "demand", "table", 1)
    if b2:
        sh.section(b2.get("caption") or "下游单耗/系数")
        cols = [c.get("k", str(c)) for c in b2.get("columns", [])]
        sh.header(cols + [""] * max(0, 7 - len(cols)))
        rows = []
        for r in b2.get("rows", []):
            row = []
            for c in b2.get("columns", []):
                k = c.get("k", str(c))
                row.append(srcs(data, r.get("src")) if k == "src" else r.get(k))
            rows.append(row)
        sh.rows(rows)

    sh.blank()
    sh.section("需求端驱动（研究判断，需以数据验证）")
    for pnl in data.get("panels", []):
        for blk in pnl.get("body", []):
            if blk.get("type") == "logic":
                for lc in blk.get("items", []):
                    if re.search(r"需求|下游", lc.get("name", "")):
                        sh.kv([(lc.get("name"), "\n".join(f"{i+1}. " + re.sub('<[^>]+>', '', s) for i, s in enumerate(lc.get("steps", []))))], kwidth_total=1)
    return sh


def sheet_trade(wb, data):
    p = data.get("product", {})
    sh = Sheet(wb, "07 贸易与区域", f"07 贸易与区域市场 · {p.get('name_cn','')}",
               "进出口须注明 HS 编码与时点；区域价差须同日同规格，否则标“不可比”",
               widths=[20, 30, 22, 20, 20, 26])
    sh.section("HS 编码与监管")
    sh.header(["HS 编码", "描述", "监管条件说明", "", "", "来源"])
    hsrow = [[h.get("code"), h.get("desc"), h.get("note") or "", "", "", srcs(data, h.get("src"))] for h in p.get("hs", [])]
    if hsrow:
        sh.rows(hsrow)
    else:
        sh.note("HS 编码未取得（缺口）：不得在报关/贸易中使用推测编码。", color="B45309")

    bt = find_block(data, "trade_futures", "table", 0)
    if bt:
        sh.section(bt.get("caption") or "贸易与区域市场数据")
        cols = [c.get("k", str(c)) for c in bt.get("columns", [])]
        sh.header(cols + [""] * max(0, 6 - len(cols)))
        rows = []
        for r in bt.get("rows", []):
            row = []
            for c in bt.get("columns", []):
                k = c.get("k", str(c))
                row.append(srcs(data, r.get("src")) if k == "src" else r.get(k))
            rows.append(row)
        sh.rows(rows)

    for bi in (1, 2):
        b = find_block(data, "trade_futures", "table", bi)
        if b:
            sh.section(b.get("caption") or "")
            cols = [c.get("k", str(c)) for c in b.get("columns", [])]
            sh.header(cols + [""] * max(0, 6 - len(cols)))
            rows = []
            for r in b.get("rows", []):
                row = []
                for c in b.get("columns", []):
                    k = c.get("k", str(c))
                    row.append(srcs(data, r.get("src")) if k == "src" else r.get(k))
                rows.append(row)
            sh.rows(rows)
    return sh


def sheet_futures(wb, data):
    p = data.get("product", {})
    name = p.get("name_cn", "")
    fr = data.get("research", {}).get("futures") or {}
    sh = Sheet(wb, "08 期货与交割", f"08 期货合约与交割 · {name}",
               "合约参数以交易所原文为准；本表为研究用摘要，交易与交割前必须复核交易所最新公告",
               widths=[22, 30, 30, 32, 22, 18])
    if not p.get("has_futures"):
        sh.section("无场内期货")
        sh.note(f"{name} 目前没有对应的场内期货/期权合约，按研究规矩不设期现模块，也不做“替代品对冲”的默认假设。", color="B45309")
        gp = find_block(data, "trade_futures", "gapnote", 0)
        if gp:
            sh.kv([(gp.get("title"), gp.get("text"))], kwidth_total=1)
        return sh
    items = fr.get("items", [])
    if not items:
        sh.section("期货模块")
        sh.note("研究稿未提供期货条款（缺口）：不得自行补写合约参数。", color="B45309")
        return sh
    for it in items:
        sh.section(it.get("title", "合约"))
        sh.header(["字段", "内容", "", "", "来源", ""])
        rows = [[r.get("k"), r.get("v"), "", "", srcs(data, it.get("src")), ""] for r in it.get("rows", [])]
        sh.rows(rows)
        sh.blank()
    # 期现关系
    for pnl in data.get("panels", []):
        if pnl.get("id") != "trade_futures":
            continue
        for blk in pnl.get("body", []):
            if blk.get("type") == "table" and re.search(r"对应|交割|现货规格", blk.get("caption", "")):
                sh.section("交割标准品 ↔ 现货主流规格")
                cols = [c.get("k", str(c)) for c in blk.get("columns", [])]
                sh.header(cols + [""] * max(0, 6 - len(cols)))
                rows = []
                for r in blk.get("rows", []):
                    row = []
                    for c in blk.get("columns", []):
                        k = c.get("k", str(c))
                        row.append(srcs(data, r.get("src")) if k == "src" else r.get(k))
                    rows.append(row)
                sh.rows(rows)
    sh.blank()
    sh.section("期现关系说明（研究判断）")
    for pnl in data.get("panels", []):
        if pnl.get("id") != "trade_futures":
            continue
        for blk in pnl.get("body", []):
            if blk.get("type") == "text" and re.search(r"期现|交割", blk.get("title", "")):
                txt = re.sub("<[^>]+>", "\n", blk.get("html", ""))
                for line in [x.strip() for x in txt.split("\n") if x.strip()]:
                    sh.note("• " + line, color=C_INK)
    return sh


def sheet_indicators(wb, data):
    p = data.get("product", {})
    sh = Sheet(wb, "09 关键指标", f"09 关键研究指标 · {p.get('name_cn','')}",
               "每条指标含口径、频次、影响权重与研究要点；权重为研究判断，非统计结论",
               widths=[6, 30, 22, 10, 8, 52, 24])
    b = find_block(data, "indicators", "indicators", 0)
    if not b:
        sh.section("指标")
        sh.note("研究稿未提供指标框架（缺）。", color="B45309")
        return sh
    n = 0
    for g in b.get("groups", []):
        sh.section(g.get("name", "指标组"), g.get("src") and S(data, g["src"][0]) or "")
        sh.header(["#", "指标", "口径", "频次", "影响", "研究要点（为什么看）", "来源"])
        rows = []
        for it in g.get("items", []):
            n += 1
            rows.append([n, it.get("name"), it.get("caliber"), it.get("freq"), it.get("importance"),
                         "\n".join("• " + x for x in it.get("points", [])), srcs(data, it.get("src"))])
        sh.rows(rows, numeric=(1,))
    sh.blank()
    sh.section("首屏框架摘要（五维）")
    sh.header(["维度", "指标数", "已覆盖", "缺口", "用途说明", "", ""])
    rows = []
    for f in data.get("frame", []):
        items = f.get("items", [])
        g = sum(1 for i in items if i.get("gap"))
        rows.append([f.get("name"), len(items), len(items) - g, g, f.get("value_note"), "", ""])
    sh.rows(rows)
    return sh


def sheet_sources(wb, data):
    sh = Sheet(wb, "10 来源清单", "10 来源清单",
               "所有关键事实的来源编号；无 URL 者多为交易所公告/纸质报告/行业综述，使用时注明标题与日期",
               widths=[8, 22, 46, 46, 14, 16, 18, 30])
    sh.header(["编号", "机构", "标题", "URL", "日期", "类型", "用途", "置信度"])
    rows = [[s.get("id"), s.get("org"), s.get("title"), s.get("url") or "（无链接）", s.get("date"), s.get("type"),
             s.get("use"), s.get("confidence") or ""] for s in data.get("sources", [])]
    sh.rows(rows)
    for r in range(1, len(rows) + 1):
        pass
    return sh


def sheet_gaps(wb, data):
    sh = Sheet(wb, "11 缺口与核验", "11 缺口声明与核验记录",
               "“没查到”比“看起来合理”更安全：本表是交付的一部分，不得删除",
               widths=[8, 26, 56, 36, 14])
    sh.section("数据缺口")
    sh.header(["#", "领域", "缺口描述", "补数路径", "影响"])
    rows = [[i, g.get("area"), g.get("desc"), g.get("how"), g.get("impact") or "见描述"] for i, g in enumerate(data.get("gaps", []), 1)]
    sh.rows(rows, numeric=(1,))
    sh.blank()
    sh.section("口径与使用提醒")
    for c in data.get("caveats", []):
        sh.note("• " + c, color=C_INK)
    v = data.get("vcs") or {}
    if v:
        sh.blank()
        sh.section("核验记录（V&V）")
        sh.header(["项", "结果", "", "", ""])
        sh.rows([[k, str(x), "", "", ""] for k, x in v.items()])
    return sh


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("research")
    ap.add_argument("--out", default="out")
    ap.add_argument("--name", default=None)
    a = ap.parse_args(argv)
    data = json.loads(Path(a.research).read_text(encoding="utf-8-sig"))
    wb = Workbook()
    wb.remove(wb.active)
    sheet_cover(wb, data)
    sheet_encyclopedia(wb, data)
    sheet_grades(wb, data)
    sheet_routes(wb, data)
    sheet_chain(wb, data)
    sheet_capacity(wb, data)
    sheet_demand(wb, data)
    sheet_trade(wb, data)
    sheet_futures(wb, data)
    sheet_indicators(wb, data)
    sheet_sources(wb, data)
    sheet_gaps(wb, data)
    name = a.name or data.get("meta", {}).get("output_name") or (data.get("product", {}).get("name_cn", "品种") + "-研究")
    out = Path(a.out) / f"{name}.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.properties.title = data.get("meta", {}).get("title", "")
    wb.properties.creator = "chem-product-quicklook"
    wb.properties.description = "品种研究入门工作簿；事实、工艺参数、公开统计与研究判断分层标注"
    wb.save(out)
    print(f"[build_xlsx] OK → {out}  ({len(wb.sheetnames)} sheets)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
