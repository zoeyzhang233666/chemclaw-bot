#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""交付前自检：Excel 表结构与行数、HTML 区块与内联 JS 语法（用 node --check）、数据契约。"""
import json
import os
import re
import subprocess
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))


def find_node():
    cands = [
        (__import__('shutil').which('node') or 'node'),
    ]
    for c in cands:
        if os.path.exists(c):
            return c
    from shutil import which
    return which("node")


def check_excel(path):
    from openpyxl import load_workbook
    wb = load_workbook(path)
    expect = ["0.说明与口径", "1.项目主表", "2.装置明细", "3.进展记录", "3b.预计投产时间变更史",
              "4.来源登记", "5.年度汇总", "6.隐性供给缺口", "7.口径冲突", "8.剔除与重复披露",
              "9.地区分布", "10.工艺与路线", "11.自选跟踪", "12.待补证清单"]
    errs = []
    for e in expect:
        if e not in wb.sheetnames:
            errs.append("缺少工作表：%s" % e)
    extra = [s for s in wb.sheetnames if s not in expect]
    if extra:
        errs.append("存在未登记的工作表：%s" % extra)
    stats = []
    for ws in wb.worksheets:
        nrow = 0
        for row in ws.iter_rows(values_only=True):
            if any(v not in (None, "") for v in row):
                nrow += 1
        stats.append((ws.title, nrow, ws.max_column))
    # 关键行数
    checks = {"1.项目主表": 3, "3.进展记录": 3, "4.来源登记": 3}
    for name, minrows in checks.items():
        if name in wb.sheetnames:
            ws = wb[name]
            n = sum(1 for row in ws.iter_rows(values_only=True) if any(v not in (None, "") for v in row))
            if n < minrows:
                errs.append("%s 数据行过少：%d" % (name, n))
    # 空表检测
    for ws in wb.worksheets:
        n = sum(1 for row in ws.iter_rows(values_only=True) if any(v not in (None, "") for v in row))
        if n == 0:
            errs.append("工作表为空：%s" % ws.title)
    return errs, stats


def check_html(path, node):
    with open(path, "r", encoding="utf-8") as f:
        html = f.read()
    errs = []
    # 未替换占位符
    for ph in ["@@DATA@@", "@@GEO@@", "@@TITLE@@", "@@ASOF@@", "@@VIEWBOX@@", "@@PROD@@"]:
        if ph in html:
            errs.append("未替换的占位符：%s" % ph)
    # 必需区块
    for sec in ["s-overview", "s-timeline", "s-capacity", "s-map", "s-table",
                "s-delayed", "s-dedup", "s-watch", "s-method"]:
        if 'id="%s"' % sec not in html:
            errs.append("缺少区块：#%s" % sec)
    # 抽取最后一个 <script> 做语法检查
    scripts = re.findall(r"<script>(.*?)</script>", html, re.S)
    if not scripts:
        errs.append("未找到内联脚本")
        return errs, {}
    js = scripts[-1]
    n = find_node()
    if not n:
        errs.append("未找到 node，跳过 JS 语法检查")
        return errs, {}
    with tempfile.NamedTemporaryFile("w", suffix=".js", delete=False, encoding="utf-8") as f:
        f.write(js)
        tmp = f.name
    try:
        p = subprocess.run([n, "--check", tmp], capture_output=True, text=True, timeout=90)
        if p.returncode != 0:
            errs.append("JS 语法错误：%s" % (p.stderr.strip()[:1500] or p.stdout.strip()[:1500]))
    finally:
        try:
            os.unlink(tmp)
        except OSError:
            pass
    # 结构统计
    info = {"size": len(html), "js_size": len(js),
            "scripts": len(scripts), "placeholders_left": html.count("@@")}
    # 必须存在的核心函数
    for fn in ["renderOverview", "renderTimeline", "renderCapacity", "renderMap",
               "renderTable", "renderDelayed", "renderDedup", "renderWatchFull",
               "renderMethod", "openP", "toggleStar", "exportCSV"]:
        if ("function %s(" % fn) not in js:
            errs.append("缺少函数：%s" % fn)
    return errs, info


def check_data(proj_dir):
    import ppctl as C
    d = os.path.join(proj_dir, "data")
    errs = []
    pj = C.load_json(os.path.join(d, "projects.json"))
    pr = C.load_json(os.path.join(d, "progress.json"))["progress"]
    src = {s["id"]: s for s in C.load_json(os.path.join(d, "sources.json"))["sources"]}
    ids = [p["project_id"] for p in pj["projects"]]
    if len(ids) != len(set(ids)):
        errs.append("project_id 重复")
    # 每个项目的来源必须存在
    for p in pj["projects"]:
        for s in p.get("source_ids", []):
            if s not in src:
                errs.append("%s 引用不存在来源 %s" % (p["project_id"], s))
        if not p.get("units"):
            errs.append("%s 无 units" % p["project_id"])
    for r in pr:
        if r["project_id"] not in ids:
            errs.append("%s 指向不存在项目" % r["progress_id"])
        for s in r.get("source_ids", []):
            if s not in src:
                errs.append("%s 引用不存在来源 %s" % (r["progress_id"], s))
        if not r.get("source_ids"):
            errs.append("%s 无来源" % r["progress_id"])
    # 未支撑的事实等级
    for r in pr:
        best = min((src[s]["grade"] for s in r.get("source_ids", []) if s in src), default="D")
        if r.get("evidence_grade") and r["evidence_grade"] < best:
            errs.append("%s 自评等级 %s 高于来源最高等级 %s" %
                        (r["progress_id"], r["evidence_grade"], best))
    # 未被引用的来源（提示而非错误）
    used = set()
    for p in pj["projects"]:
        used.update(p.get("source_ids", []))
    for r in pr:
        used.update(r.get("source_ids", []))
    unused = [s for s in src if s not in used]
    return errs, {"projects": len(ids), "progress": len(pr), "sources": len(src),
                  "unused_sources": unused}


def run_check(proj_dir):
    out = os.path.join(proj_dir, "out")
    node = find_node()
    allerrs = []
    print("=" * 78)
    print("交付自检：%s" % proj_dir)

    xlsx = None
    for f in os.listdir(out):
        if f.endswith(".xlsx"):
            xlsx = os.path.join(out, f)
    html = None
    for f in os.listdir(out):
        if f.endswith(".html"):
            html = os.path.join(out, f)

    if xlsx:
        e, stats = check_excel(xlsx)
        print("\n[Excel] %s" % os.path.basename(xlsx))
        for name, rows, cols in stats:
            print("   %-24s %4d 行 × %2d 列" % (name, rows, cols))
        allerrs += e
    else:
        allerrs.append("未找到 xlsx 产物")

    if html:
        e, info = check_html(html, node)
        print("\n[HTML] %s" % os.path.basename(html))
        print("   大小 %.0f KB ｜ 内联JS %.0f KB ｜ script块 %s ｜ node=%s"
              % (info.get("size", 0) / 1024, info.get("js_size", 0) / 1024,
                 info.get("scripts", "?"), "OK" if node else "未找到"))
        allerrs += e
    else:
        allerrs.append("未找到 html 产物")

    e, info = check_data(proj_dir)
    print("\n[数据] 项目 %d ｜ 进展 %d ｜ 来源 %d" %
          (info["projects"], info["progress"], info["sources"]))
    if info["unused_sources"]:
        print("   未被引用的来源 %d 条（登记但未支撑任何事实）：%s"
              % (len(info["unused_sources"]), ",".join(info["unused_sources"][:20])))
    allerrs += e

    print("\n" + "=" * 78)
    if allerrs:
        for e in allerrs:
            print("[ERR ] " + e)
        print("=" * 78)
        print("自检未通过：%d 项" % len(allerrs))
        sys.exit(1)
    print("自检全部通过")
    print("=" * 78)


if __name__ == "__main__":
    run_check(sys.argv[1] if len(sys.argv) > 1 else ".")
