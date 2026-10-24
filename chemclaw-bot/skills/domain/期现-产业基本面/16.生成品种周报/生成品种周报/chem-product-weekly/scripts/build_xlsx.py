# -*- coding: utf-8 -*-
"""Excel 底稿生成模板（15 张表）。读取 data/*.json，输出 build/<品种>周报_<起>_<止>.xlsx。

依赖：openpyxl。表结构与 assets/weekly_workbook.xlsx 一致，可按品种增删列。
"""
import json, os, sys
sys.stdout.reconfigure(encoding='utf-8')
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, 'data')
BUILD = os.path.join(ROOT, 'build')
os.makedirs(BUILD, exist_ok=True)

M = json.load(open(os.path.join(DATA, 'metrics.json'), encoding='utf-8'))
F = json.load(open(os.path.join(DATA, 'facts.json'), encoding='utf-8'))
TITLE = F.get('title', '品种周报')
OUT = os.path.join(BUILD, '%s.xlsx' % F.get('filename', 'weekly_report'))

HEAD = Font(bold=True, color='FFFFFF', size=10)
HFILL = PatternFill('solid', fgColor='1F4E79')
THIN = Side(style='thin', color='BFBFBF')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')


def sheet(wb, title, head, rows, widths, note=None):
    ws = wb.create_sheet(title)
    ws['A1'] = '%s｜%s' % (TITLE, title)
    ws['A1'].font = Font(bold=True, size=12, color='1F4E79')
    if note:
        ws['A2'] = note
        ws['A2'].font = Font(size=9, color='808080', italic=True)
    r0 = 4 if note else 3
    for j, h in enumerate(head, 1):
        c = ws.cell(row=r0, column=j, value=h)
        c.font = HEAD; c.fill = HFILL; c.border = BOX
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(row=r0 + i, column=j, value=v)
            c.border = BOX; c.alignment = WRAP
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = 'A%d' % (r0 + 1)
    ws.row_dimensions[r0].height = 30
    return ws


def main():
    wb = Workbook(); wb.remove(wb.active)
    # 00 说明
    ws = wb.create_sheet('00_说明与口径')
    ws['A1'] = TITLE; ws['A1'].font = Font(bold=True, size=15, color='1F4E79')
    for i, (k, v) in enumerate(F.get('meta', {}).items(), start=3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True, size=10, color='1F4E79')
        ws.cell(row=i, column=2, value=v).alignment = WRAP
    ws.column_dimensions['A'].width = 20; ws.column_dimensions['B'].width = 110
    # 01 本周vs上周
    rows = [[c, v['prev_close'], v['week_avg'], v['week_high'], v['week_low'], v['week_close'],
             v['chg'], None if v['chg_pct'] is None else round(v['chg_pct'], 2),
             v['oi_start'], v['oi_end'], v['oi_chg'], v['vol_total']] for c, v in M['futures'].items()]
    sheet(wb, '01_本周vs上周', ['合约', '周初收盘', '本周均价', '本周最高', '本周最低', '周末收盘',
                           '周涨跌', '周涨跌幅%', '周初持仓', '周末持仓', '持仓变化', '本周成交'],
          rows, [12, 12, 11, 11, 11, 12, 10, 11, 12, 12, 12, 14], note=M.get('dominant_note'))
    # 02 历史价格
    days = sorted({d for c in M['futures'] for d in M['futures'][c]['daily']})
    rows = [[d] + sum([[M['futures'][c]['daily'].get(d), M['futures'][c]['daily_oi'].get(d)] for c in M['futures']], []) for d in days]
    head = ['日期'] + sum([['%s收' % c, '%s持仓' % c] for c in M['futures']], [])
    sheet(wb, '02_历史价格', head, rows, [12] + [11] * (len(head) - 1))
    # 03 基差
    b = F.get('basis_rows', [])
    if b:
        sheet(wb, '03_基差', list(b[0].keys()), [list(x.values()) for x in b], [14] * len(b[0]))
    # 04 月差
    if M.get('spread'):
        keys = [k for k in M['spread'][0].keys() if k != 'date']
        sheet(wb, '04_月差', ['日期'] + keys, [[x['date']] + [x.get(k) for k in keys] for x in M['spread']],
              [12] + [16] * len(keys))
    # 05 供需
    sheet(wb, '05_供需指标', ['数据日期', '指标', '数值', '单位', '环比/说明', '频率', '来源', '链接', '时点'],
          [[x['date'], x['metric'], x.get('value'), x['unit'], x.get('chg', ''), x.get('period', ''),
            x.get('src', ''), x.get('url', ''), '窗口内' if x.get('in_week') else '背景'] for x in F.get('fund', [])],
          [12, 34, 10, 8, 34, 8, 26, 56, 10])
    # 06 上游
    sheet(wb, '06_上游原料', ['数据日期', '品种/指标', '数值', '单位', '口径', '变动说明', '来源', '链接'],
          [[x['date'], x['item'], x.get('value'), x['unit'], x.get('basis', ''), x.get('chg', ''), x.get('src', ''), x.get('url', '')]
           for x in F.get('upstream', [])], [12, 34, 10, 8, 22, 40, 30, 56])
    # 07 价格锚点
    sheet(wb, '07_价格锚点', ['数据日期', '价格点', '数值', '单位', '口径', '备注', '来源', '链接'],
          [[x['date'], x['name'], x.get('value'), x['unit'], x.get('basis', ''), x.get('note', ''), x.get('src', ''), x.get('url', '')]
           for x in F.get('anchors', [])], [12, 26, 10, 10, 26, 52, 30, 56])
    # 08 事件
    sheet(wb, '08_事件', ['时点', '日期', '时间', '事件', '主体', '机构评述', '来源', '链接'],
          [['窗口外' if x.get('bg') else '窗口内', x['date'], x.get('time', ''), x['what'], x.get('entity', ''),
            x.get('impact', ''), x.get('src', ''), x.get('url', '')] for x in sorted(F.get('events', []), key=lambda y: y['date'])],
          [8, 12, 7, 78, 22, 30, 30, 56])
    # 09 重要变化
    sheet(wb, '09_重要变化', ['序号', '变化点', '数据依据', '对应表'],
          [[i + 1, x['point'], x['evidence'], x['sheet']] for i, x in enumerate(F.get('changes', []))], [6, 46, 86, 22])
    # 10 趋势指标
    sheet(wb, '10_趋势指标', ['日期', '指标值', '短期均线', '长期均线', '说明'],
          [[x['date'], x.get('value'), x.get('m10'), x.get('m20'), x.get('desc', '')] for x in F.get('trend', [])],
          [12, 12, 12, 12, 80])
    # 11 外盘
    sheet(wb, '11_外盘', ['交易日', 'CFR', 'FOB', '其他'], [[x[0], x[1], x[2], x[3] if len(x) > 3 else ''] for x in F.get('foreign', [])],
          [12, 14, 14, 14])
    # 12 历史同期
    sheet(wb, '12_历史同期', ['本年合约', '去年同合约', '去年同周均价', '去年同周收盘'],
          [[c, v['ly_contract'], v['ly_avg'], v['ly_close']] for c, v in M.get('lastyear', {}).items()], [14, 14, 16, 16])
    # 13 缺口
    sheet(wb, '13_缺口清单', ['序号', '缺口说明'], [[i + 1, g] for i, g in enumerate(F.get('gaps', []))], [6, 120])
    # 14 来源
    seen, rows = set(), []
    for key in ('anchors', 'upstream', 'fund', 'events'):
        for x in F.get(key, []):
            if x.get('url') and x['url'] not in seen:
                seen.add(x['url']); rows.append([len(rows) + 1, key, x.get('src', ''), x['url']])
    sheet(wb, '14_来源记录', ['序号', '类别', '来源', 'URL'], rows, [6, 14, 40, 96])
    wb.save(OUT)
    print('OK', OUT)


if __name__ == '__main__':
    main()
