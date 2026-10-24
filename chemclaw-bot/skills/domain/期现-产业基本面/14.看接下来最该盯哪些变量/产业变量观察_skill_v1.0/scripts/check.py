# -*- coding: utf-8 -*-
"""
产物自检：把 references/自检清单.md 中「可机器判定」的条目落成代码。

用法：
    python check.py <index.html> <workbook.xlsx> [--payload my_case.py]

退出码 0 = 全部通过；1 = 有未通过项。
"""
from __future__ import annotations

import argparse
import importlib.util
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import contract  # noqa: E402

# 明确禁止出现的多空预测/交易建议词
# 注意：不含"减仓/增仓"——它们在"持仓量日减仓 3.32 万手"这类市场事实描述中是合规用法。
FORBIDDEN = [
    "建议做多", "建议做空", "建议买入", "建议卖出", "建议逢低", "建议逢高",
    "看涨", "看跌", "买入评级", "卖出评级", "目标价", "建仓", "加仓",
    "必涨", "必跌", "稳赚", "抄底", "梭哈",
]
# 允许出现的合规声明（在免责段落内）
DISCLAIMER_HINTS = ["不构成投资建议", "不构成任何投资建议", "不构成对价格方向的预测"]


class Checker:
    def __init__(self):
        self.items: list[tuple[str, bool, str]] = []

    def add(self, code: str, ok: bool, detail: str = ""):
        self.items.append((code, ok, detail))

    def report(self) -> int:
        passed = sum(1 for _, ok, _ in self.items if ok)
        total = len(self.items)
        print("\n" + "=" * 74)
        print(f"自检结果：{passed}/{total} 通过")
        print("=" * 74)
        for code, ok, detail in self.items:
            mark = "✓" if ok else "✗"
            line = f"  {mark} {code}"
            if detail:
                line += f"  {detail}"
            print(line)
        fails = [c for c, ok, _ in self.items if not ok]
        if fails:
            print("\n未通过项：", ", ".join(fails))
        print("=" * 74)
        return 0 if not fails else 1


def load_payload(module_path: Path) -> dict:
    spec = importlib.util.spec_from_file_location("payload_mod", module_path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)  # type: ignore
    return getattr(mod, "PAYLOAD")


DATE_LIKE = re.compile(r"\d{1,2}[-/月]\d{1,2}|\d{4}-\d{2}-\d{2}|每周|每日|日频|周频|月度|季度|次月")
VAGUE = ["近期", "下个月", "待定", "关注后续", "稍后", "以后再说"]


def check_payload(p: dict, ck: Checker):
    # ---- A 变量筛选 ----
    vs = p["variables"]
    ck.add("A1 入选变量 3-6 个", 3 <= len(vs) <= 6, f"当前 {len(vs)} 个")

    from collections import Counter
    cnt = Counter(v["cat"] for v in vs)
    ck.add("A2a 同分类 ≤2 个", all(n <= 2 for n in cnt.values()), str(dict(cnt)))
    ck.add("A2b 覆盖 ≥2 个分类", len(cnt) >= 2, f"{len(cnt)} 类")

    scored = all(isinstance(v.get(k), int) and 1 <= v[k] <= 5
                 for v in vs for k in ("importance", "verifiability", "urgency"))
    ck.add("A3 三维打分 1-5 整数", scored)

    ranks = [v["rank"] for v in vs]
    ck.add("A4 rank 唯一", len(set(ranks)) == len(ranks))
    # 综合分降序一致性
    def score(v):
        return v["importance"] * 2 + v["verifiability"] + v["urgency"]
    ordered = [score(v) for v in sorted(vs, key=lambda x: x["rank"])]
    ck.add("A4b rank 顺序与综合分一致", ordered == sorted(ordered, reverse=True),
           f"综合分序列 {ordered}")

    ck.add("A5 候选池非空", len(p["candidates"]) >= len(vs),
           f"候选 {len(p['candidates'])} / 入选 {len(vs)}")
    no_reason = [c["name"] for c in p["candidates"]
                 if not c["selected"] and len(str(c.get("reason", ""))) < 8]
    ck.add("A6 落选均有具体理由", not no_reason, str(no_reason[:3]) if no_reason else "")
    ck.add("A7 二级观察非空", len(p["secondary"]) > 0, f"{len(p['secondary'])} 条")

    # ---- B 变量内容七要素 ----
    empty_field = []
    for v in vs:
        for f in ("current", "recent_change", "history_pos", "why", "mechanism",
                  "key_data", "verify_yes", "verify_no", "invalidate", "source", "freq", "next_date"):
            if not str(v.get(f, "")).strip():
                empty_field.append(f"{v['vid']}.{f}")
    ck.add("B1-B7 七要素无空字段", not empty_field, str(empty_field[:5]) if empty_field else "")

    # 阈值：确认/否证信号里应含数字
    no_thresh = [v["vid"] for v in vs
                 if not re.search(r"\d", str(v["verify_yes"]) + str(v["verify_no"]))]
    ck.add("B6 验证条件含数字阈值", not no_thresh, str(no_thresh) if no_thresh else "")

    # 下次可观测：必须是具体日期/周期
    bad_next = [f"{v['vid']}: {v['next_date'][:24]}" for v in vs
                if not DATE_LIKE.search(str(v["next_date"]))
                or any(w in str(v["next_date"]) for w in VAGUE)]
    ck.add("B7/G1 下次可观测具体", not bad_next, str(bad_next) if bad_next else "")

    # 每个变量必须有 gap 说明（口径与局限）
    ck.add("B7 口径局限已写", all(len(str(v.get("gap", ""))) > 5 for v in vs))

    # ---- C 事件 ----
    ev = p["events"]
    ck.add("C1 事件档位齐全", all(e.get("kind") for e in ev), f"{len(ev)} 条")
    vogue_kinds = {"已披露·时点模糊"}
    bad = [e["title"][:18] for e in ev if e["kind"] in vogue_kinds and "以" not in e["certainty"]]
    ck.add("C2 时点模糊写了『以…为准』", not bad, str(bad) if bad else "")

    cond = [e for e in ev if "条件触发" in e["kind"] or "季节性" in e["kind"]]
    bad_cond = [e["title"][:18] for e in cond
                if not any(w in e["certainty"] for w in ("不代表", "不确定", "浮动", "逐年", "窗口"))]
    ck.add("C3 条件触发已声明非必然", not bad_cond, str(bad_cond) if bad_cond else "")

    dir_words = ["利好", "利空"]
    bad_dir = [e["title"][:16] for e in ev if any(w in e.get("impact", "") for w in dir_words)]
    ck.add("C5 影响路径不含方向词", not bad_dir, str(bad_dir) if bad_dir else "")

    # ---- D 数据纪律 ----
    ck.add("D1 快照含口径列", all(s.get("scope") for s in p["snapshot"]),
           f"{len(p['snapshot'])} 行")
    ck.add("D4 快照均有来源", all(s.get("source") for s in p["snapshot"]))
    gaps_ok = all(all(g.get(k) for k in ("item", "why", "workaround", "limit")) for g in p["gaps"])
    ck.add("D6 缺口表四列齐全", gaps_ok, f"{len(p['gaps'])} 项")


def check_html(path: Path, ck: Checker):
    src = path.read_text(encoding="utf-8")
    ck.add("F1 无外部 CDN", not re.search(r'<script[^>]+src=["\']https?://', src)
           and not re.search(r'<link[^>]+href=["\']https?://', src))
    for tab in ("view-cards", "view-timeline", "view-snapshot"):
        ck.add(f"F2 视图 {tab}", tab in src)
    ck.add("F3 交互脚本存在", "<script>" in src and "addEventListener" in src)
    ck.add("E4 含免责声明", any(h in src for h in DISCLAIMER_HINTS))

    hits = [w for w in FORBIDDEN if w in src]
    ck.add("E1-E3 无多空预测词", not hits, f"命中 {hits}" if hits else "")
    ck.add("F7 区块完整",
           all(k in src for k in ("候选池与筛选依据", "数据缺口登记", "固定观察节奏")))


def check_excel(path: Path, ck: Checker):
    import openpyxl
    wb = openpyxl.load_workbook(path)
    need = ["00 封面", "01 变量观察总表", "02 当前产业状态", "03 未来事件时间线",
            "04 候选池与筛选", "05 验证条件与触发器", "06 数据缺口登记", "07 观察日程"]
    titles = [ws.title for ws in wb.worksheets]
    ck.add("F4 8 个 sheet 齐全", titles == need, str(titles))

    ws5 = wb["05 验证条件与触发器"]
    ck.add("F5 回写下拉存在", bool(ws5.data_validations.dataValidation),
           f"{len(ws5.data_validations.dataValidation)} 个校验")

    # 禁止词扫描
    hits = set()
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if isinstance(cell, str):
                    for w in FORBIDDEN:
                        if w in cell:
                            hits.add(w)
    ck.add("E1-E3 Excel 无多空词", not hits, f"命中 {sorted(hits)}" if hits else "")


def cross_consistency(p: dict, html_src: str, xl_path: Path, ck: Checker):
    """抽查关键数字是否在两个产物中一致。"""
    import openpyxl
    wb = openpyxl.load_workbook(xl_path)

    ws1 = wb["01 变量观察总表"]
    names_xl = [str(ws1.cell(r, 2).value) for r in range(4, 4 + len(p["variables"]))]
    names_html = [v["name"] for v in sorted(p["variables"], key=lambda x: x["rank"])]
    ck.add("F6a 变量名单一致", names_xl == names_html,
           "不一致" if names_xl != names_html else f"{len(names_xl)} 个")

    # 抽查快照当前值
    ws2 = wb["02 当前产业状态"]
    mismatch = []
    for i, s in enumerate(p["snapshot"]):
        v = str(ws2.cell(4 + i, 3).value or "")
        if v and v not in html_src:
            mismatch.append(s["indicator"][:14])
    ck.add("F6b 快照值在 HTML 中一致", not mismatch, str(mismatch[:3]) if mismatch else "")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("html")
    ap.add_argument("xlsx")
    ap.add_argument("--payload", default=str(HERE.parent / "examples" / "示例-甲醇-2026Q4" / "payload_甲醇.py"))
    a = ap.parse_args()

    p = load_payload(Path(a.payload))
    ck = Checker()

    probs = contract.validate(p)
    ck.add("E0 契约校验", not probs, "; ".join(probs[:3]) if probs else "")

    print("== payload 检查 ==")
    check_payload(p, ck)
    print("== HTML 检查 ==")
    check_html(Path(a.html), ck)
    print("== Excel 检查 ==")
    check_excel(Path(a.xlsx), ck)
    print("== 一致性检查 ==")
    cross_consistency(p, Path(a.html).read_text(encoding="utf-8"), Path(a.xlsx), ck)

    return ck.report()


if __name__ == "__main__":
    raise SystemExit(main())
