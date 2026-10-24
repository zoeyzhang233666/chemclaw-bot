# -*- coding: utf-8 -*-
"""qa_xlsx.py —— 打开工作簿，核对页签、行数与关键内容是否齐全。"""
import io
import sys

from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

path = sys.argv[1]
wb = load_workbook(path)
expected = ["00_使用说明", "01_证据索引", "02_原始数据", "03_企业产能", "04_项目台账", "05_价格数据",
            "06_产业指标", "07_计算结果", "08_结论", "09_关键变量", "10_资料缺口", "11_章节大纲",
            "12_质量自检"]
print("sheets:", len(wb.sheetnames))
missing = [s for s in expected if s not in wb.sheetnames]
extra = [s for s in wb.sheetnames if s not in expected]
print("missing:", missing or "none", "| unexpected:", extra or "none")
ok = not missing
for name in wb.sheetnames:
    ws = wb[name]
    # 表头行固定在第 4 行
    headers = [ws.cell(row=4, column=c).value for c in range(1, ws.max_column + 1)]
    headers = [h for h in headers if h is not None]
    data_rows = sum(1 for r in range(5, ws.max_row + 1)
                    if any(ws.cell(row=r, column=c).value not in (None, "") for c in range(1, 6)))
    print("  %-12s 行=%-4d 列=%-3d 表头=%d 数据行=%d  %s" %
          (name, ws.max_row, ws.max_column, len(headers), data_rows, "/".join(map(str, headers[:6]))))
    if data_rows == 0 and name not in ("00_使用说明",):
        print("    !! 空表"); ok = False

# 关键内容抽查
ws = wb["08_结论"]
texts = " ".join(str(ws.cell(row=r, column=5).value or "") for r in range(5, ws.max_row + 1))
for kw in ("预计投产", "新增供给", "285"):
    if kw not in texts:
        print("  !! 08_结论 缺少关键词:", kw); ok = False
# 每条判断类结论都必须有反证条件
ws = wb["08_结论"]
for r in range(5, ws.max_row + 1):
    t = ws.cell(row=r, column=3).value
    falsify = ws.cell(row=r, column=9).value
    if t in ("判断", "解释") and not falsify:
        print("  !! 缺少反证条件:", ws.cell(row=r, column=2).value); ok = False
ws = wb["04_项目台账"]
status = [ws.cell(row=r, column=11).value for r in range(5, 25) if ws.cell(row=r, column=11).value]
print("项目状态分布:", {s: status.count(s) for s in set(status)})
ws = wb["12_质量自检"]
res = [ws.cell(row=r, column=1).value for r in range(5, ws.max_row + 1) if ws.cell(row=r, column=1).value]
print("自检结果:", {r: res.count(r) for r in set(res)})
if "FAIL" in res:
    ok = False
print("RESULT", "PASS" if ok else "FAIL")
sys.exit(0 if ok else 1)
