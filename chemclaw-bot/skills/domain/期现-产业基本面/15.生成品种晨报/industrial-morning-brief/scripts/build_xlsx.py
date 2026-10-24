# -*- coding: utf-8 -*-
"""生成 Excel 工作簿（与 HTML 同源数字，逐格可追溯）。

用法：
  python build_xlsx.py build/data.json --out out/晨报.xlsx

工作表：
 0 封面与说明      报告日、基线、口径声明、筛选统计
 1 重要变化总表    全部指标：原始指标/上期值/本期值/变动/幅度/数据日/级别/一句话
 2 价格与价差      按品种的分区价格、区域价差
 3 基差与月差      基差、基差率、月差及其算法
 4 库存与开工      周度/月度基本面指标（含数据日期与陈旧提示）
 5 事件台账        全部事件：日期/主体/量化/影响路径/评分/级别/去重/时效
 6 被排除资讯      排除原因与核对方式
 7 数据缺口        缺什么、为什么、怎么补
 8 数据来源        每个来源的用途
 9 期货日线        主连日线（近 250 日）供复核
10 品种与产业链    上中下游与驱动因素
11 筛选规则        重要性打分与去重规则的说明（可交接给同事）
"""
import argparse
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import common

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

INK = '1F2937'
HEAD_BG = '0F3D3A'
SUB_BG = 'E8F1EF'
A_BG = 'FEF2F2'
B_BG = 'EFF6FF'
C_BG = 'F1F5F9'
WARN_BG = 'FFFBEB'

THIN = Side(style='thin', color='D8DEE7')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
F_TITLE = Font(name='微软雅黑', size=15, bold=True, color='FFFFFF')
F_HEAD = Font(name='微软雅黑', size=10, bold=True, color='FFFFFF')
F_BODY = Font(name='微软雅黑', size=10, color=INK)
F_BOLD = Font(name='微软雅黑', size=10, bold=True, color=INK)
F_SMALL = Font(name='微软雅黑', size=9, color='6B7280')
FILL_HEAD = PatternFill('solid', fgColor=HEAD_BG)
FILL_SUB = PatternFill('solid', fgColor=SUB_BG)
RED = Font(name='微软雅黑', size=10, bold=True, color='C0392B')
GREEN = Font(name='微软雅黑', size=10, bold=True, color='047857')
GREY = Font(name='微软雅黑', size=10, color='94A3B8')
WRAP = Alignment(wrap_text=True, vertical='top')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)


def sheet(wb, name, title, widths, headers, rows, freezes='A4', notes=None):
    ws = wb.create_sheet(name)
    ncol = len(headers)
    # 标题行
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    c = ws.cell(row=1, column=1, value=title)
    c.font = F_TITLE
    c.fill = FILL_HEAD
    c.alignment = Alignment(vertical='center', horizontal='left', indent=1)
    ws.row_dimensions[1].height = 30
    # 表头
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=3, column=j, value=h)
        c.font = F_HEAD
        c.fill = FILL_SUB
        c.border = BORDER
        c.alignment = CENTER
        c.font = Font(name='微软雅黑', size=10, bold=True, color=INK)
    ws.row_dimensions[3].height = 26
    for i, row in enumerate(rows, 4):
        for j, v in enumerate(row, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = F_BODY
            c.border = BORDER
            c.alignment = WRAP
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = freezes
    if notes:
        r = 4 + len(rows) + 1
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
        c = ws.cell(row=r, column=1, value=notes)
        c.font = F_SMALL
        c.alignment = WRAP
        ws.row_dimensions[r].height = 42
    return ws


def pct_of(v):
    return None if v is None else round(v / 100.0, 6)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('data')
    ap.add_argument('--out', required=True)
    a = ap.parse_args()
    d = common.load_json(a.data)
    rep, base = d['report'], d['baseline']
    inds, evs = d['indicators'], d['events']
    wb = Workbook()
    wb.remove(wb.active)

    # ---------------------------------------------------------- 0 封面
    ws = wb.create_sheet('封面与说明')
    ws.merge_cells('A1:F1')
    c = ws['A1']
    c.value = rep['title'] + '　·　' + rep['report_date']
    c.font = F_TITLE
    c.fill = FILL_HEAD
    c.alignment = Alignment(vertical='center', indent=1)
    ws.row_dimensions[1].height = 40
    s = d['summary']
    info = [
        ('报告日', rep['report_date']),
        ('对比基线（上一期晨报）', '%s（%s）' % (base['date'], base.get('prev_label', ''))),
        ('数据截止', rep.get('cutoff', '')),
        ('覆盖品种', ' / '.join(rep.get('products', []))),
        ('指标条数', '%d 条（有基线对比 %d 条 / 本期新增 %d 条）'
         % (s['indicators_total'], s['with_base'], s['is_new'])),
        ('事件条数', '%d 条（A 级 %d / B 级 %d / C 级噪声 %d）'
         % (s['events_total'], s.get('events_A', s['tier_A']),
            s.get('events_B', 0), s.get('events_C', 0))),
        ('被排除资讯', '%d 条（见「被排除资讯」页）' % len(d.get('excluded', []))),
        ('数据缺口', '%d 项（见「数据缺口」页）' % len(d.get('gaps', []))),
    ]
    for i, (k, v) in enumerate(info, 3):
        ws.cell(row=i, column=1, value=k).font = F_BOLD
        ws.cell(row=i, column=2, value=v).font = F_BODY
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=6)
        ws.cell(row=i, column=2).alignment = WRAP
    rules = [
        ('口径声明 1：期货与现货税基不同',
         '期货价格不含税、含交割成本；现货为含税报价。本表基差 = 含税现货 − 不含税期货，'
         '仅为观察期现结构，不能直接作为无风险套利结论。'),
        ('口径声明 2：数据日期 ≠ 事件日期',
         '每条指标与事件都单列「数据日期」与「事件发生日期」。'
         '若数据日期早于报告日，标「未更新」，不称为当天最新变化。'),
        ('口径声明 3：没有的数值写缺口，不写估算',
         '任何未取得的数值，一律在「数据缺口」页列出缺什么、为什么缺、怎么补，不以估算值填充。'),
        ('口径声明 4：不做预测',
         '全文只陈述已发生事实与可核对数字，不含方向性判断、目标价或投资建议。'
         '文本已通过方向性表述扫描（条件句「若…则…」允许，预测性结论禁止）。'),
        ('口径声明 5：同一屏只用一套价格口径',
         '摘要与表格统一用「报告日收盘 vs 上一期基线收盘」；盘中价另标「盘中参考」，不参与涨跌计算。'),
    ]
    r = len(info) + 4
    ws.cell(row=r, column=1, value='使用须知').font = Font(name='微软雅黑', size=11, bold=True, color='0F766E')
    for i, (k, v) in enumerate(rules, r + 1):
        ws.cell(row=i, column=1, value=k).font = F_BOLD
        ws.cell(row=i, column=1).alignment = WRAP
        ws.cell(row=i, column=2, value=v).font = F_BODY
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=6)
        ws.cell(row=i, column=2).alignment = WRAP
        ws.row_dimensions[i].height = 46
    ws.column_dimensions['A'].width = 26
    for col in 'BCDEF':
        ws.column_dimensions[col].width = 24

    # ---------------------------------------------------------- 1 重要变化总表
    rows = []
    for i in sorted(inds, key=lambda x: ({'A': 0, 'B': 1, 'C': 2}[x['significance']],
                                         x['product'], x['module'])):
        rows.append([
            i['significance'], i['product'], i['module'].upper(), i['name'], i['spec'],
            i.get('region', ''), i.get('unit', ''), i.get('tax', ''),
            i['base'], i['base_date'], i['current'], i['current_date'],
            i['abs'], pct_of(i['pct']),
            '新增指标' if i['is_new'] else ('未更新（早 %d 天）' % i['stale_days']
                                            if (i['stale_days'] or 0) > 1 else '已更新'),
            i.get('base_source', ''), i.get('formula', ''),
            i.get('what', ''), i.get('why', ''), i.get('verify', ''),
            i.get('source', ''), i.get('source_url', ''),
        ])
    ws = sheet(wb, '重要变化总表', '全部指标：原始指标 / 上期值 / 本期值 / 变动 / 筛选结果　（基线 %s → 报告 %s）'
               % (base['date'], rep['report_date']),
               [7, 8, 9, 26, 16, 14, 9, 18, 11, 11, 11, 11, 10, 9, 15, 20, 30, 46, 50, 46, 30, 42],
               ['级别', '品种', '模块', '指标名称', '规格/合约', '区域', '单位', '税基',
                '上期值', '上期日期', '本期值', '本期日期', '变动', '幅度', '数据状态',
                '基线来源', '算法', '发生了什么', '为什么重要', '今天继续验证什么', '来源', '来源链接'],
               rows,
               notes='说明：级别 A=必看（≥70 分）、B=记录（45-69 分）、C=噪声（<45 分或纯观点/无量化）。'
                     '「数据状态=未更新」表示该数据发布日期早于报告日，仅代表最新有效值，不得当作当天变化。')
    for i, row in enumerate(rows, 4):
        lvl = row[0]
        fill = {'A': A_BG, 'B': B_BG, 'C': C_BG}[lvl]
        ws.cell(row=i, column=1).fill = PatternFill('solid', fgColor=fill)
        ws.cell(row=i, column=1).alignment = CENTER
        ws.cell(row=i, column=1).font = Font(name='微软雅黑', size=10, bold=True,
                                            color={'A': 'B91C1C', 'B': '1D4ED8', 'C': '64748B'}[lvl])
        v = row[12]
        if isinstance(v, (int, float)):
            ws.cell(row=i, column=13).font = RED if v > 0 else (GREEN if v < 0 else GREY)
        p = row[13]
        if isinstance(p, (int, float)):
            ws.cell(row=i, column=14).number_format = '0.00%'
            ws.cell(row=i, column=14).font = RED if p > 0 else (GREEN if p < 0 else GREY)
        if row[14].startswith('未更新'):
            ws.cell(row=i, column=15).fill = PatternFill('solid', fgColor=WARN_BG)

    # ---------------------------------------------------------- 2 价格与价差
    rows = []
    for i in inds:
        if i['module'] in ('price', 'spread', 'cost'):
            rows.append([i['product'], i['name'], i['spec'], i.get('region', ''), i.get('unit', ''),
                         i.get('tax', ''), i['base'], i['base_date'], i['current'],
                         i['current_date'], i['abs'], pct_of(i['pct']),
                         i['significance'], '新增指标' if i['is_new'] else '',
                         i.get('source', ''), i.get('source_url', '')])
    ws = sheet(wb, '价格与价差', '现货 / 期货 / 区域价差：价格与变动（同口径对比 %s → %s）'
               % (base['date'], rep['report_date']),
               [8, 26, 20, 16, 9, 18, 11, 11, 11, 11, 10, 9, 7, 12, 30, 40],
               ['品种', '指标', '规格/合约', '区域', '单位', '税基', '上期值', '上期日期',
                '本期值', '本期日期', '变动', '幅度', '级别', '备注', '来源', '来源链接'],
               rows)
    for i, row in enumerate(rows, 4):
        if isinstance(row[11], (int, float)):
            ws.cell(row=i, column=12).number_format = '0.00%'

    # ---------------------------------------------------------- 3 基差与月差
    rows = []
    for i in inds:
        if i['module'] == 'basis':
            rows.append([i['product'], i['name'], i.get('formula', ''), i['base'], i['base_date'],
                         i['current'], i['current_date'], i['abs'], pct_of(i['pct']),
                         i.get('tax', ''), i['significance'], i.get('what', ''),
                         i.get('why', ''), i.get('source', '')])
    ws = sheet(wb, '基差与月差', '期现结构与合约间结构（基差 = 含税现货 − 不含税期货，不作套利结论）',
               [8, 24, 34, 11, 11, 11, 11, 10, 9, 22, 7, 44, 50, 34],
               ['品种', '指标', '算法', '上期值', '上期日期', '本期值', '本期日期', '变动',
                '幅度', '税基说明', '级别', '发生了什么', '为什么重要', '来源'],
               rows)
    for i, row in enumerate(rows, 4):
        if isinstance(row[8], (int, float)):
            ws.cell(row=i, column=9).number_format = '0.00%'

    # ---------------------------------------------------------- 4 库存与开工
    rows = []
    for i in inds:
        if i['module'] in ('inventory', 'runrate', 'demand'):
            rows.append([i['product'], i['name'], i.get('spec', ''), i.get('region', ''),
                         i.get('unit', ''), i.get('freq', ''), i['base'], i['base_date'],
                         i['current'], i['current_date'],
                         i['abs'], pct_of(i['pct']), i.get('stale_days'),
                         i['significance'], i.get('what', ''), i.get('verify', ''),
                         i.get('source', ''), i.get('source_url', '')])
    ws = sheet(wb, '库存与开工', '库存 / 开工 / 产量 / 下游需求（周度与月度，注意数据日期）',
               [8, 26, 18, 12, 9, 9, 11, 11, 11, 11, 10, 9, 9, 7, 50, 44, 32, 40],
               ['品种', '指标', '规格', '区域', '单位', '频率', '上期值', '上期日期', '本期值',
                '本期日期', '变动', '幅度', '距报告日', '级别', '发生了什么', '今天继续验证什么',
                '来源', '来源链接'],
               rows,
               notes='「距报告日」>1 天的行表示该数据未在报告日更新，只能作为最新有效值引用，'
                     '严禁在正文中称为「当天变化」。')
    for i, row in enumerate(rows, 4):
        if isinstance(row[11], (int, float)):
            ws.cell(row=i, column=12).number_format = '0.00%'
        if isinstance(row[12], int) and row[12] > 1:
            ws.cell(row=i, column=13).fill = PatternFill('solid', fgColor=WARN_BG)

    # ---------------------------------------------------------- 5 事件台账
    rows = []
    order = {'A': 0, 'B': 1, 'C': 2}
    for e in sorted(evs, key=lambda x: (order.get(x['tier'], 3), -(x.get('score') or 0))):
        rows.append([
            e['tier'], e.get('score'), e.get('id'), e.get('event_date'),
            ' / '.join(e.get('products') or []),
            e.get('event_type_cn') or e.get('event_type'), e.get('entity'),
            e.get('title'), e.get('quant'), e.get('impact_path'),
            {'done': '已兑现', 'ongoing': '持续推进', 'plan': '计划中未兑现',
             'repeat': '旧闻·仅进展', 'comment': '观点类'}.get(e.get('status'), e.get('status')),
            e.get('stale_days'),
            e.get('dup_count', 0), '、'.join(e.get('dup_sources') or []),
            e.get('dedupe_note', ''), e.get('repeat_note', ''), e.get('exclude_reason', ''),
            e.get('source'), e.get('source_url'),
        ])
    ws = sheet(wb, '事件台账', '产业事件台账：发生日 / 量化影响 / 影响路径 / 评分与去重结果',
               [7, 7, 6, 11, 12, 11, 16, 46, 34, 46, 10, 9, 7, 14, 26, 30, 20, 26, 40],
               ['级别', '评分', '编号', '发生日', '涉及品种', '事件类型', '主体', '标题',
                '量化影响', '产业影响路径', '状态', '距报告日', '合并数', '合并来源',
                '去重说明', '时效提示', '排除原因', '来源', '来源链接'],
               rows,
               notes='评分规则：影响路径 30 + 量化 25 + 时效性 20 + 独特性 15 + 覆盖度 10 = 100 分；'
                     '≥70 为 A，45-69 为 B，<45 或纯观点无量化为 C（噪声，不进正文）。'
                     '「状态=旧闻」表示事件发生于上一期基线之前，仅作进展更新，最高只评 B。')
    for i, row in enumerate(rows, 4):
        lvl = row[0]
        ws.cell(row=i, column=1).fill = PatternFill('solid', fgColor={'A': A_BG, 'B': B_BG, 'C': C_BG}[lvl])
        ws.cell(row=i, column=1).alignment = CENTER
        ws.cell(row=i, column=1).font = Font(name='微软雅黑', size=10, bold=True,
                                            color={'A': 'B91C1C', 'B': '1D4ED8', 'C': '64748B'}[lvl])
        if row[10] == 'repeat':
            ws.cell(row=i, column=11).fill = PatternFill('solid', fgColor=WARN_BG)

    # ---------------------------------------------------------- 6 被排除资讯
    rows = [[x.get('id'), x.get('title'), x.get('reason'), x.get('check'),
             x.get('source'), x.get('source_url')] for x in d.get('excluded', [])]
    sheet(wb, '被排除资讯', '为什么这些资讯没有进入晨报（旧闻、转载、噪声、无关）',
          [7, 46, 62, 34, 22, 44],
          ['编号', '标题', '排除原因', '核对方式', '来源', '链接'], rows)

    # ---------------------------------------------------------- 7 数据缺口
    rows = [[x.get('item'), x.get('why'), x.get('how')] for x in d.get('gaps', [])]
    sheet(wb, '数据缺口', '没有拿到的数据：缺什么 / 为什么 / 怎么补（不以估算值填充）',
          [40, 62, 42], ['缺口项', '为什么没有', '怎么补'], rows)

    # ---------------------------------------------------------- 8 数据来源
    rows = [[x.get('name'), x.get('type'), x.get('used'), x.get('url')] for x in d.get('sources', [])]
    sheet(wb, '数据来源', '每个数字都能回到这里',
          [34, 12, 60, 46], ['来源', '类型', '用于', '链接'], rows)

    # ---------------------------------------------------------- 9 期货日线
    rows = []
    for p, f in d.get('futures', {}).items():
        for cd in (f.get('candles_250') or []):
            rows.append([p, f.get('label'), cd[0], cd[1], cd[2], cd[3], cd[4]])
    ws = sheet(wb, '期货日线', '主连日线（近 250 个交易日）供逐日复核　— 开/收/低/高',
               [8, 10, 12, 11, 11, 11, 11],
               ['品种', '序列', '日期', '开盘', '收盘', '最低', '最高'], rows)

    # ---------------------------------------------------------- 10 品种与产业链
    rows = []
    for p, ch in (d.get('chain') or {}).items():
        rows.append([p, ' / '.join(ch.get('up', [])), ' / '.join(ch.get('mid', [])),
                     ' / '.join(ch.get('down', [])), ' / '.join(ch.get('drivers', []))])
    sheet(wb, '品种与产业链', '上中下游与当期核心驱动（用于判断影响路径）',
          [8, 40, 40, 40, 44],
          ['品种', '上游', '中游', '下游', '当期核心驱动'], rows)

    # ---------------------------------------------------------- 11 筛选规则
    rules = [
        ['R1', '基线对比', '本期所有「变动」= 本期值 − 上一期晨报同口径值；不得与更早的旧值相减冒充变动。',
         '重要变化总表·上期值列'],
        ['R2', '同口径', '价格类指标一律用「报告日收盘 vs 上一期基线收盘」；盘中价单独标注，不参与涨跌计算。',
         '价格与价差·备注列'],
        ['R3', '陈旧降级', '数据发布日期距报告日 >1 天 → 标「未更新」，最多降为 B 级，正文不得称当天变化。',
         '库存与开工·距报告日列'],
        ['R4', '新增指标', '上一期未覆盖的指标标「新增指标」，不与空值做百分比。',
         '重要变化总表·数据状态列'],
        ['R5', '事件去重 D1', '同一事件多源转载合并为一条：同事件类型族 + 同主体 + 量化数字指纹有交集，'
         '或标题 2-gram 相似 ≥0.55（纯字符串相似度对中文改写稿无效，实测仅约 0.06）。',
         '事件台账·合并数/合并来源'],
        ['R6', '旧闻识别 D2', '事件发生日早于上一期基线 → 标「旧闻·仅进展更新」，最高评 B，不当作当日新事件。',
         '事件台账·状态/时效提示'],
        ['R7', '泛评剔除 D3', '无量化、无明确影响路径的观点归纳 → 判 C 噪声，不进正文。',
         '事件台账·排除原因'],
        ['R8', '重要性分级', '评分 = 影响路径30 + 量化25 + 时效20 + 独特性15 + 覆盖度10；≥70 A、45-69 B、<45 C。',
         '事件台账·评分'],
        ['R9', '缺口不补', '拿不到的数值一律列入「数据缺口」，不以估算值、同比反推值填表。',
         '数据缺口'],
        ['R10', '禁止预测', '文本经方向性表述扫描：允许「若…则…」条件句，禁止预计涨跌、看多看空、抄底、'
         '底部已现等结论性表述。', '封面·口径声明 4'],
    ]
    sheet(wb, '筛选规则', '判定规则（可交接给同事复算）',
          [7, 18, 88, 30], ['编号', '规则', '做法', '对应本表位置'], rules)

    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    wb.save(a.out)
    common.ok('Excel 已生成：%s（%d 个表，%.1f KB）'
              % (a.out, len(wb.sheetnames), os.path.getsize(a.out) / 1024))


if __name__ == '__main__':
    main()
