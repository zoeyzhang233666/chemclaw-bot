# -*- coding: utf-8 -*-
"""build_xlsx.py —— 由 analysis.json 生成 9 表 Excel（含真实公式，可在 Excel 里改单耗重算）

用法：python scripts/build_xlsx.py build/analysis.json --out-dir out
"""
import argparse, json, os, sys

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

HDR_FILL = PatternFill('solid', fgColor='1F3A5F')
HDR_FONT = Font(color='FFFFFF', bold=True, size=10)
TITLE_FONT = Font(bold=True, size=13, color='1F3A5F')
SUB_FONT = Font(size=9, color='777777')
THIN = Side(style='thin', color='D0D7DE')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NOTE_FILL = PatternFill('solid', fgColor='FFF4CE')
DANGER_FILL = PatternFill('solid', fgColor='FDE7E9')
INPUT_FILL = PatternFill('solid', fgColor='E8F0FE')
NUM = '#,##0.00'
NUM0 = '#,##0.0'


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = BORDER


def autosize(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def add_title(ws, title, sub=None, ncols=8):
    ws.cell(row=1, column=1, value=title).font = TITLE_FONT
    if sub:
        ws.cell(row=2, column=1, value=sub).font = SUB_FONT
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
        ws.cell(row=2, column=1).alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[2].height = 28


def build(analysis, out_path, events=None):
    meta = analysis['meta']
    s = analysis['summary']
    up, dn = meta['upstream'], meta['downstream']
    mats = meta.get('materials') or []
    m0 = mats[0] if mats else {}
    daily = analysis['daily']
    stages = analysis.get('stages') or []
    divs = analysis.get('divergences') or []
    calib = analysis.get('calibration') or []
    steps = analysis.get('steps') or {}

    wb = Workbook()

    # ============ 00 封面与说明 ============
    ws = wb.active
    ws.title = '00_封面与说明'
    add_title(ws, '%s → %s 成本传导研究' % (up['name'], dn['name']),
              sub='观测窗口 %s ~ %s，%d 个交易日。'
                  '本工作簿所有数字与 HTML 报告同源（同一份 analysis.json）。' %
                  (s['window'][0], s['window'][1], s['n_obs']), ncols=6)
    rows = [
        ('', ''),
        ('【必读】加工价差不是利润', '价差 = 下游价 − 单耗×上游价。不含能源、辅料、人工、折旧、'
                                    '财务费用、检维修、运费、库存损益。不得作为利润对外使用。'),
        ('【必读】同涨同跌不等于因果', '上游与下游可能被原油、宏观、政策等共同因素驱动。'
                                      '本工作簿不做因果推断，相关系数仅为观察值。'),
        ('上游', '%s ｜ %s ｜ %s ｜ %s' % (up['name'], up.get('spec'), up.get('region'), up.get('tax'))),
        ('下游', '%s ｜ %s ｜ %s ｜ %s' % (dn['name'], dn.get('spec'), dn.get('region'), dn.get('tax'))),
        ('单耗（主模型）', '%s 吨 %s / 吨 %s（区间 %s ~ %s）' %
         (m0.get('unit'), up['name'], dn['name'], m0.get('low'), m0.get('high'))),
        ('价差定义', (meta.get('params') or {}).get('spread_definition', '')),
        ('', ''),
        ('表目录', '01原料价格 · 02下游价格 · 03工艺单耗 · 04理论成本 · 05加工价差 · '
                   '06分阶段对比 · 07背离区间 · 08计算公式 · 09数据缺口与来源'),
    ]
    r = 4
    for k, v in rows:
        ws.cell(row=r, column=1, value=k).font = Font(bold=True, size=10)
        c = ws.cell(row=r, column=2, value=v)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        if '必读' in k:
            c.fill = DANGER_FILL if '不是利润' in k else NOTE_FILL
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        ws.row_dimensions[r].height = 30
        r += 1
    autosize(ws, [20, 30, 20, 20, 20, 20])

    # ============ 01 原料价格 ============
    ws = wb.create_sheet('01_原料价格')
    add_title(ws, '%s 日度价格（上游）' % up['name'],
              sub='%s ｜ %s ｜ %s ｜ 来源：%s' % (up.get('spec'), up.get('region'),
                                             up.get('tax'), up.get('source')), ncols=5)
    hdr = ['日期', '%s 价格（元/吨）' % up['name'], '日变动', '是否变动', '备注']
    ws.append([]); ws.append(hdr)
    hr = ws.max_row
    style_header(ws, hr, len(hdr))
    prev = None
    for rec in daily:
        v = rec['up']
        chg = None if prev is None or v is None else v - prev
        ws.append([rec['date'], v, chg, ('是' if chg else '否') if chg is not None else '',
                   '台阶（连续不变）' if chg == 0 else ''])
        prev = v
    for row in ws.iter_rows(min_row=hr + 1, max_row=ws.max_row, max_col=5):
        row[1].number_format = NUM
        row[2].number_format = NUM
        for c in row:
            c.border = BORDER
    autosize(ws, [12, 20, 12, 10, 22])
    ws.freeze_panes = 'A%d' % (hr + 1)

    # ============ 02 下游价格 ============
    ws = wb.create_sheet('02_下游价格')
    add_title(ws, '%s 日度价格（下游）' % dn['name'],
              sub='%s ｜ %s ｜ %s ｜ 来源：%s' % (dn.get('spec'), dn.get('region'),
                                             dn.get('tax'), dn.get('source')), ncols=4)
    hdr = ['日期', '%s 价格（元/吨）' % dn['name'], '日变动', '日涨跌%']
    ws.append([]); ws.append(hdr)
    hr = ws.max_row
    style_header(ws, hr, len(hdr))
    prev = None
    for rec in daily:
        v = rec['down']
        chg = None if prev is None or v is None else v - prev
        pc = None if not prev or v is None else (v / prev - 1) * 100
        ws.append([rec['date'], v, chg, pc])
        prev = v
    for row in ws.iter_rows(min_row=hr + 1, max_row=ws.max_row, max_col=4):
        row[1].number_format = NUM
        row[2].number_format = NUM
        row[3].number_format = NUM
        for c in row:
            c.border = BORDER
    autosize(ws, [12, 20, 12, 10])
    ws.freeze_panes = 'A%d' % (hr + 1)

    # ============ 03 工艺单耗 ============
    ws = wb.create_sheet('03_工艺单耗')
    add_title(ws, '工艺与单耗参数（可修改，04/05 表会联动重算）',
              sub='浅蓝底色为可修改输入项。改单耗后，04 表的吨成本与 05 表的价差会按公式自动重算。', ncols=6)
    proc = meta.get('process') or {}
    r = 4
    ws.cell(row=r, column=1, value='工艺路线').font = Font(bold=True)
    ws.cell(row=r, column=2, value=proc.get('route', ''))
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6); r += 1
    ws.cell(row=r, column=1, value='反应式').font = Font(bold=True)
    ws.cell(row=r, column=2, value=proc.get('equation', '')); r += 1
    ws.cell(row=r, column=1, value='核实状态').font = Font(bold=True)
    ws.cell(row=r, column=2, value='已核实' if proc.get('verified') else '待核实'); r += 2

    hdr = ['原料', '单耗（主模型）', '区间下界', '区间上界', '单位口径', '来源']
    ws.append(hdr)
    hr = ws.max_row
    style_header(ws, hr, len(hdr))
    for m in mats:
        ws.append([m['name'], m.get('unit'), m.get('low'), m.get('high'),
                   m.get('unit_basis', ''), m.get('source', '')])
    for row in ws.iter_rows(min_row=hr + 1, max_row=ws.max_row, max_col=6):
        for c in row[1:4]:
            c.number_format = '0.000'
            c.fill = INPUT_FILL
        for c in row:
            c.border = BORDER
    UNIT_CELL = 'B%d' % (hr + 1)
    r = ws.max_row + 2
    ws.cell(row=r, column=1, value='主模型单耗所在单元格：').font = Font(bold=True)
    ws.cell(row=r, column=2, value='=%s' % UNIT_CELL).font = Font(color='1F3A5F', bold=True)
    r += 2
    opt = meta.get('optional_materials') or []
    if opt:
        ws.cell(row=r, column=1, value='未纳入主模型的辅料').font = Font(bold=True); r += 1
        ws.append(['项目', '单耗', '下界', '上界', '状态说明', ''])
        style_header(ws, ws.max_row, 5)
        r += 1
        for m in opt:
            ws.append([m['name'], m.get('unit'), m.get('low'), m.get('high'),
                       m.get('status', ''), ''])
            ws.merge_cells(start_row=ws.max_row, start_column=5, end_row=ws.max_row, end_column=6)
    autosize(ws, [22, 16, 12, 12, 16, 46])

    # ============ 04 理论成本（带公式） ============
    ws = wb.create_sheet('04_理论成本')
    add_title(ws, '理论原料成本（Excel 真公式，可复核）',
              sub='公式：成本 = 单耗 × 上游价。单耗引用 03 表 B 列主模型单元格，改那里即可全表重算。',
              ncols=6)
    hdr = ['日期', '%s 价格（元/吨）' % up['name'], '单耗', '理论原料成本（元/吨）',
           '%s 价格（元/吨）' % dn['name'], '原料成本占下游价%']
    ws.append([]); ws.append(hdr)
    hr = ws.max_row
    style_header(ws, hr, len(hdr))
    n = len(daily)
    for i, rec in enumerate(daily, 1):
        rr = hr + i
        ws.append([rec['date'], rec['up'], None, None, rec['down'], None])
        ws.cell(row=rr, column=3, value='=03_工艺单耗!%s' % UNIT_CELL).number_format = '0.000'
        ws.cell(row=rr, column=4, value='=B%d*C%d' % (rr, rr)).number_format = NUM
        ws.cell(row=rr, column=6, value='=IF(E%d,D%d/E%d*100,"")' % (rr, rr, rr)).number_format = '0.0'
        for c in range(1, 7):
            ws.cell(row=rr, column=c).border = BORDER
        ws.cell(row=rr, column=2).number_format = NUM
        ws.cell(row=rr, column=5).number_format = NUM
    last = hr + n
    # 汇总行
    sr = last + 2
    ws.cell(row=sr, column=1, value='全期均值').font = Font(bold=True)
    for col in (2, 4, 5, 6):
        L = get_column_letter(col)
        ws.cell(row=sr, column=col,
                value='=AVERAGE(%s%d:%s%d)' % (L, hr + 1, L, last)).number_format = NUM
    ws.cell(row=sr + 1, column=1, value='最小 / 最大价差见 05 表').font = SUB_FONT
    autosize(ws, [12, 18, 10, 20, 18, 16])
    ws.freeze_panes = 'A%d' % (hr + 1)

    # ============ 05 加工价差 ============
    ws = wb.create_sheet('05_加工价差')
    add_title(ws, '加工价差（价差口径，不是利润）',
              sub='价差 = 下游价 − 理论原料成本。低/高单耗两列用于敏感性。', ncols=7)
    hdr = ['日期', '%s 价（元/吨）' % dn['name'], '理论原料成本', '加工价差',
           '价差-低单耗情形', '价差-高单耗情形', 'z分数(20日)']
    ws.append([]); ws.append(hdr)
    hr = ws.max_row
    style_header(ws, hr, len(hdr))
    for i, rec in enumerate(daily, 1):
        rr = hr + i
        ws.append([rec['date'], rec['down'], None, None, None, None, rec.get('z')])
        ws.cell(row=rr, column=3, value='=04_理论成本!D%d' % (i + 4)).number_format = NUM
        ws.cell(row=rr, column=4, value='=B%d-C%d' % (rr, rr)).number_format = NUM
        ws.cell(row=rr, column=5, value=rec.get('spread_low_unit')).number_format = NUM
        ws.cell(row=rr, column=6, value=rec.get('spread_high_unit')).number_format = NUM
        for c in range(1, 8):
            ws.cell(row=rr, column=c).border = BORDER
        ws.cell(row=rr, column=2).number_format = NUM
        if ws.cell(row=rr, column=7).value is not None:
            ws.cell(row=rr, column=7).number_format = '0.00'
    last = hr + n
    sr = last + 2
    stats = [('均值', 'AVERAGE'), ('中位数', 'MEDIAN'), ('最小', 'MIN'), ('最大', 'MAX')]
    ws.cell(row=sr, column=1, value='统计').font = Font(bold=True)
    ws.cell(row=sr, column=2, value='加工价差').font = Font(bold=True)
    for j, (lab, fn) in enumerate(stats, 1):
        ws.cell(row=sr + j, column=1, value=lab)
        ws.cell(row=sr + j, column=2,
                value='=%s(D%d:D%d)' % (fn, hr + 1, last)).number_format = NUM
    sp = s['spread']
    ws.cell(row=sr, column=4, value='与 analysis.json 对照（非公式）').font = SUB_FONT
    for j, (lab, val) in enumerate([('均值', sp['mean']), ('中位数', sp['median']),
                                    ('最小', sp['min']), ('最大', sp['max'])], 1):
        ws.cell(row=sr + j, column=4, value=lab)
        ws.cell(row=sr + j, column=5, value=val).number_format = NUM
    autosize(ws, [12, 18, 16, 14, 18, 18, 12])
    ws.freeze_panes = 'A%d' % (hr + 1)

    # ============ 06 分阶段对比 ============
    ws = wb.create_sheet('06_分阶段对比')
    add_title(ws, '分阶段成本传导分解',
              sub='恒等式：Δ下游价 = Δ原料成本 + Δ加工价差（残差应为 0）。'
                  '传导率 = Δ下游价 / Δ原料成本；分母过小（<成本1%）时标记为不稳定。', ncols=11)
    hdr = ['阶段起', '阶段止', '天数', '期初原料成本', '期末原料成本', 'Δ原料成本',
           '期初下游价', '期末下游价', 'Δ下游价', 'Δ加工价差', '传导率']
    ws.append([]); ws.append(hdr)
    hr = ws.max_row
    style_header(ws, hr, len(hdr))
    for st in stages:
        ws.append([st['start'], st['end'], st.get('days'),
                   st.get('cost_start'), st.get('cost_end'), st.get('d_cost'),
                   st.get('down_start'), st.get('down_end'), st.get('d_down'),
                   st.get('d_spread'),
                   st.get('pass_through') if st.get('pass_through_valid', True) else None])
    for row in ws.iter_rows(min_row=hr + 1, max_row=ws.max_row, max_col=len(hdr)):
        for c in row[3:]:
            c.number_format = NUM
        for c in row:
            c.border = BORDER
    r = ws.max_row + 2
    ws.cell(row=r, column=1, value='⚠ 传导率说明').font = Font(bold=True, color='C00000')
    notes = [
        '传导率 ≈1：成本基本全额传导；<1：下游未跟满、价差被压缩；>1：下游涨幅超过成本、价差走阔；<0：反向。',
        '传导率 >1 时必须先检查上游是否为台阶式评估价（见 01 表"是否变动"列），'
        '分母过小会把比值放大成几倍，这是计算假象。',
        '本表只做算术分解，不证明"上游导致下游"。',
    ]
    for t in notes:
        r += 1
        c = ws.cell(row=r, column=1, value=t)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        c.fill = NOTE_FILL
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=11)
        ws.row_dimensions[r].height = 26
    autosize(ws, [12, 12, 7, 15, 15, 13, 13, 13, 12, 13, 10])

    # ============ 07 背离区间 ============
    ws = wb.create_sheet('07_背离区间')
    add_title(ws, '价差背离区间与异常跳变',
              sub='背离：价差相对 20 日滚动均值的 z 分数，连续越阈 ≥5 日。异常跳变：单日变动 ≥400 元/吨。')
    hdr = ['类型', '起', '止', '天数', '期初价差', '期末价差', '区间均值', '极端z', '原因核查']
    ws.append([]); ws.append(hdr)
    hr = ws.max_row
    style_header(ws, hr, len(hdr))
    for d in divs:
        lab = '走阔' if d['type'] == 'spread_widen' else '压缩'
        seg = [e for e in (events or []) if d['start'] <= e['date'] <= d['end']]
        why = '；'.join(e['title'][:26] for e in seg[:3]) if seg else '未检索到可核查事件，原因未确认'
        ws.append([lab, d['start'], d['end'], d['days'], d.get('spread_start'),
                   d.get('spread_end'), d.get('spread_mean'), d.get('z_extreme'), why])
    for row in ws.iter_rows(min_row=hr + 1, max_row=ws.max_row, max_col=len(hdr)):
        for c in row[4:8]:
            c.number_format = NUM
        for c in row:
            c.border = BORDER
            c.alignment = Alignment(wrap_text=True, vertical='top')
    r = ws.max_row + 2
    ws.cell(row=r, column=1, value='单日异常跳变（未剔除）').font = Font(bold=True)
    r += 1
    for j, h in enumerate(['日期', '前一日价差', '当日价差', '跳变'], 1):
        ws.cell(row=r, column=j, value=h)
    style_header(ws, r, 4)
    for o in (analysis.get('outliers') or [])[:30]:
        r += 1
        ws.cell(row=r, column=1, value=o['date'])
        ws.cell(row=r, column=2, value=o['prev']).number_format = NUM
        ws.cell(row=r, column=3, value=o['now']).number_format = NUM
        ws.cell(row=r, column=4, value=o['jump']).number_format = NUM
    autosize(ws, [10, 12, 12, 7, 12, 12, 12, 9, 58])

    # ============ 08 计算公式 ============
    ws = wb.create_sheet('08_计算公式')
    add_title(ws, '全部计算公式与参数', ncols=4)
    ws.append([])
    ws.append(['序号', '项目', '公式', '说明'])
    hr = ws.max_row
    style_header(ws, hr, 4)
    formulas = [
        ('1', '理论原料成本', '成本 = Σ(单耗ᵢ × 原料价ᵢ)',
         '单耗单位：吨原料/吨产品；价格单位：元/吨（已含税、同地区）'),
        ('2', '单耗敏感性', '价差_低 = 下游价 − 下界单耗 × 上游价',
         '下界单耗 → 成本偏低 → 价差偏高'),
        ('3', '加工价差', '价差 = 下游价 − 理论原料成本',
         '★ 价差口径，不是利润。不含能源/辅料/人工/折旧/财务/运费'),
        ('4', '阶段分解', 'Δ下游 = Δ原料成本 + Δ价差',
         '恒等式，残差应恒为 0'),
        ('5', '传导率', '传导率 = Δ下游价 / Δ原料成本',
         '分母 < 成本1% 时不稳定，标记为无效'),
        ('6', '原料成本占比', '占比 = 原料成本 / 下游价 × 100%',
         '正常区间 60%~98%，超出提示单位或单耗有误'),
        ('7', '价差 z 分数', 'z = (价差 − 20日滚动均值) / 20日滚动标准差',
         '|z|≥1.5 且连续 ≥5 日 → 背离区间'),
        ('8', '滞后相关', 'r(k) = corr(Δ上游[t], Δ下游[t+k]), k=0..20',
         '★ 仅观察，不作因果推断'),
        ('9', '美元→人民币', '上游价(元/吨) = CFR × 汇率 × (1+关税) × (1+增值税)',
         '常漏关税项；换完必须检查数量级（化工品通常 6000~11000 元/吨）'),
    ]
    for f_ in formulas:
        ws.append(list(f_))
    for row in ws.iter_rows(min_row=hr + 1, max_row=ws.max_row, max_col=4):
        for c in row:
            c.border = BORDER
            c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.append([])
    r = ws.max_row + 1
    ws.cell(row=r, column=1, value='本期实际参数').font = Font(bold=True)
    r += 1
    for k, v in [('上游', '%s（%s / %s / %s）' % (up['name'], up.get('spec'), up.get('region'), up.get('tax'))),
                 ('下游', '%s（%s / %s / %s）' % (dn['name'], dn.get('spec'), dn.get('region'), dn.get('tax'))),
                 ('单耗', '%s（%s~%s）' % (m0.get('unit'), m0.get('low'), m0.get('high'))),
                 ('观测数', '%d 个交易日' % s['n_obs']),
                 ('上游变动天数', '%s / %d 天' % (steps.get('up_change_days'), steps.get('n'))),
                 ('下游变动天数', '%s / %d 天' % (steps.get('down_change_days'), steps.get('n')))]:
        ws.cell(row=r, column=1, value=k).font = Font(bold=True)
        ws.cell(row=r, column=2, value=v)
        r += 1
    autosize(ws, [7, 20, 46, 52])

    # ============ 09 数据缺口与来源 ============
    ws = wb.create_sheet('09_数据缺口与来源')
    add_title(ws, '数据缺口、口径校验与来源', ncols=5)
    ws.append([])
    ws.append(['类型', '项目', '说明', '影响/偏差', '来源'])
    hr = ws.max_row
    style_header(ws, hr, 5)
    for c in calib:
        ws.append(['锚点校验', c.get('date'), '模型 %s vs 公开 %s' % (c.get('model'), c.get('reported')),
                   c.get('gap'), c.get('source')])
    gaps = [
        ('缺口', '辅料（醋酸等）', '缺同期间日度基准价，未纳入价差计算',
         '价差少扣约 75~160 元/吨', '见 references/data-sources.md'),
        ('缺口', '装置级成本', '无企业级能耗、折旧、财务费用数据',
         '只能做行业价差，不能判断单个装置盈亏', '—'),
        ('口径', '评估价 vs 成交价', '本表用的是资讯机构日度评估价，非成交价',
         '实际成交受品牌/账期/地区影响，可能差百元级', '—'),
        ('口径', '上游台阶式更新', '上游 %s 在 %s 天中仅 %s 天变动' %
         (up['name'], steps.get('n'), steps.get('up_change_days')),
         '阶段传导率在分母极小时失真', '01 表"是否变动"列'),
    ]
    for g in gaps:
        ws.append(list(g))
    for row in ws.iter_rows(min_row=hr + 1, max_row=ws.max_row, max_col=5):
        for c in row:
            c.border = BORDER
            c.alignment = Alignment(wrap_text=True, vertical='top')
    autosize(ws, [12, 22, 46, 34, 40])

    wb.save(out_path)
    return out_path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('analysis')
    ap.add_argument('--out-dir', default='out')
    ap.add_argument('--events', default='')
    args = ap.parse_args()
    a = json.load(open(args.analysis, encoding='utf-8'))
    ev = []
    if args.events and os.path.exists(args.events):
        ev = json.load(open(args.events, encoding='utf-8'))
    os.makedirs(args.out_dir, exist_ok=True)
    name = '%s到%s_成本传导.xlsx' % (a['meta']['upstream']['name'], a['meta']['downstream']['name'])
    p = os.path.join(args.out_dir, name)
    build(a, p, events=ev)
    print('已生成', p, '%.0f KB' % (os.path.getsize(p) / 1024))


if __name__ == '__main__':
    main()
