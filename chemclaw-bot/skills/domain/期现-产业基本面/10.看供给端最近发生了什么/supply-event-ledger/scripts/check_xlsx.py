# -*- coding: utf-8 -*-
"""Excel 自检：sheet 列表、行数、关键合计与引擎结果一致性。"""
import os, sys, json
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
import openpyxl

here = os.path.dirname(os.path.abspath(__file__))
root = os.path.dirname(os.path.dirname(here))
case = os.environ.get('CASE_DIR') or os.path.join(root, 'case')
xlsx = os.environ.get('OUT_XLSX') or os.path.join(case, 'out', '甲醇供给事件台账.xlsx')

sys.path.insert(0, here)
import calc_impact as ci

meta = json.load(open(os.path.join(case, 'datum.json'), encoding='utf-8'))['meta']
ev = json.load(open(os.path.join(case, 'status_events.json'), encoding='utf-8'))['events']
ws_, we_ = ci.parse_date(meta['window_start']), ci.parse_date(meta['window_end'])
S = ci.summarize(ev, ws_, we_, meta['national_capacity_wt'], meta['weekly_output_t'])

wb = openpyxl.load_workbook(xlsx, data_only=True)
print('sheets:', wb.sheetnames)
fails = []

for name in wb.sheetnames:
    sh = wb[name]
    print(f'  {name:<22} rows={sh.max_row:<4} cols={sh.max_column}')

# 校验 02 表合计与引擎一致
sh = wb['02_事件清单_已发生']
tot = 0.0
n_conf = 0
for r in sh.iter_rows(min_row=2, values_only=True):
    if r[1] == '国内·已确认' and isinstance(r[14], (int, float)):
        tot += r[14]
        n_conf += 1
print(f'\n[check] 02表国内已确认合计 = {tot:,.1f} 吨 ; 引擎 = {S["totals"]["domestic_loss_t"]:,.1f} 吨')
if abs(tot - S['totals']['domestic_loss_t']) > 1:
    fails.append('02表合计与引擎不一致')
print(f'[check] 02表国内已确认事件数 = {n_conf} ; 引擎 = {S["totals"]["n_domestic_confirmed"]}')
if n_conf != S['totals']['n_domestic_confirmed']:
    fails.append('事件数不一致')

# 校验 04 公式表合计
sh = wb['04_影响量计算']
t4 = 0.0
for r in sh.iter_rows(min_row=2, values_only=True):
    if r[13] == '是' and isinstance(r[10], (int, float)):
        t4 += r[10]
print(f'[check] 04表确认口径合计 = {t4:,.1f} 吨')
if abs(t4 - S['totals']['domestic_loss_t']) > 1:
    fails.append('04表合计与引擎不一致')

# 校验汇总表（汇总表第4列为“吨”）
for name in ['05_汇总_按企业', '06_汇总_按地区', '07_汇总_按事件类型', '08_汇总_按工艺路线']:
    sh = wb[name]
    last = None
    for r in sh.iter_rows(min_row=2, values_only=True):
        if r[0] == '合计':
            last = r[3]
    print(f'[check] {name} 合计(吨) = {last:,.1f}' if isinstance(last, (int, float)) else f'[check] {name} 合计 = {last}')
    if not isinstance(last, (int, float)) or abs(last - S['totals']['domestic_loss_t']) > 1:
        fails.append(f'{name} 合计不等于引擎结果')

# 校验来源链接完整
sh = wb['10_来源与证据链']
n_src = sh.max_row - 1
n_blank = sum(1 for r in sh.iter_rows(min_row=2, values_only=True) if not r[5])
print(f'[check] 来源条目 = {n_src} ; 缺链接 = {n_blank}')
if n_blank:
    fails.append(f'{n_blank} 条来源缺链接')

print('\nRESULT:', 'PASS' if not fails else 'FAIL -> ' + '; '.join(fails))
sys.exit(0 if not fails else 1)
