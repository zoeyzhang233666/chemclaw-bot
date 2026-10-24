#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""check_all.py — 交付前自检：HTML 与 Excel 双端一致性。

用法:
    python check_all.py report.html workbook.xlsx [research.json]

检查项：
  HTML  零外链 / 无占位符 / 无重复 id / 关键区块存在 / 证据条数一致 / 反误用声明存在
  XLSX  必需工作表齐全 / 证据条数一致 / URL 均为 https / 无脏值
  一致性 同一 evidence id 在两端数值相同
退出码：0 全绿；1 有问题（逐条打印）
"""
import json, os, re, sys

REQUIRED_SHEETS = ['01_下游结构', '02_需求指标库', '03_代理指标分级', '04_历史比较',
                   '05_景气打分', '06_传导链条', '07_政策与事件', '08_价格与库存',
                   '09_缺口清单', '10_来源清单', '11_证据明细', '12_口径规范']


def check_html(path, evs, fails):
    h = open(path, encoding='utf-8').read()
    zone = h[:h.find('<script>const PAGE_DATA')] if '<script>const PAGE_DATA' in h else h
    ids = re.findall(r'\sid="([^"]+)"', h)
    dup = {i for i in ids if ids.count(i) > 1}
    if dup:
        fails.append('HTML 重复 id: %s（会导致渲染函数覆盖整个容器）' % sorted(dup))
    if re.search(r'(?:src|href)="https?://[^"]+\.(?:js|css|woff2?|ttf|png|jpg|svg)"', h):
        fails.append('HTML 存在外链静态资源（必须单文件离线）')
    for pat in ('TODO', 'PLACEHOLDER', '{{', '${', 'undefined</', '>NaN<', '>None<'):
        if pat in zone:
            fails.append('HTML 存在占位/脏值: %r' % pat)
    if '<script>const PAGE_DATA' not in h:
        fails.append('HTML 缺少内联数据 PAGE_DATA')
    for anchor in ('id="kpis"', 'id="structBody"', 'id="cards"', 'id="evBody"', 'id="proxyBody"', 'id="gapsBody"'):
        if anchor not in h:
            fails.append('HTML 缺少区块 %s' % anchor)
    if evs:
        if ('%d 条原始记录' % len(evs)) not in h and ('（%d 条）' % len(evs)) not in h:
            fails.append('HTML 证据条数与 JSON 不一致（应为 %d）' % len(evs))
        miss = [e['id'] for e in evs if e['id'] not in h]
        if miss:
            fails.append('HTML 未引用证据: %s' % miss[:5])
    for kw in ('价格', '不等于', '禁止', '不预测'):
        if kw not in h:
            fails.append('HTML 缺少反误用声明关键词 %r' % kw)
    return h


def check_xlsx(path, evs, doc, fails):
    try:
        from openpyxl import load_workbook
    except ImportError:
        fails.append('未安装 openpyxl，跳过 Excel 检查')
        return
    wb = load_workbook(path)
    for s in REQUIRED_SHEETS:
        if s not in wb.sheetnames:
            fails.append('Excel 缺少工作表 %s' % s)
    if '11_证据明细' in wb.sheetnames and evs:
        ws = wb['11_证据明细']
        got = {}
        for row in ws.iter_rows(min_row=5):
            if row[0].value and re.fullmatch(r'[A-Z]{1,2}\d{2,3}', str(row[0].value)):
                got[str(row[0].value)] = (row[3].value, row[10].value)
        miss = [e['id'] for e in evs if e['id'] not in got]
        if miss:
            fails.append('Excel 11_证据明细 缺 %d 条: %s' % (len(miss), miss[:5]))
        bad = []
        for e in evs:
            if e['id'] in got and e.get('value') is not None and isinstance(got[e['id']][0], (int, float)):
                if abs(float(got[e['id']][0]) - float(e['value'])) > 1e-6:
                    bad.append(e['id'])
        if bad:
            fails.append('Excel 与 JSON 数值不一致: %s' % bad[:5])
        for k, (v, u) in got.items():
            if u and not str(u).startswith('https://'):
                fails.append('Excel 证据 %s 链接非 https' % k)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.strip() in ('None', 'nan', 'undefined'):
                    fails.append('%s!%s 脏值 %r' % (ws.title, c.coordinate, c.value))
    print('Excel sheets:', len(wb.sheetnames))


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(2)
    html, xlsx = sys.argv[1], sys.argv[2]
    js = sys.argv[3] if len(sys.argv) > 3 else None
    evs, doc = [], {}
    if js and os.path.exists(js):
        doc = json.load(open(js, encoding='utf-8'))
        evs = doc.get('evidence', [])
    fails = []
    check_html(html, evs, fails)
    check_xlsx(xlsx, evs, doc, fails)
    print('\n=== 结果 ===')
    if fails:
        for f in fails:
            print(' ✗', f)
        print('\n共 %d 项问题' % len(fails))
        sys.exit(1)
    print(' ✓ 全部通过（HTML + Excel + 一致性）')


if __name__ == '__main__':
    main()
