# -*- coding: utf-8 -*-
"""校验重算后的工作簿：公式值是否与 Python 侧一致"""
import json

from openpyxl import load_workbook

scan = json.load(open('out/scan_result.json', encoding='utf-8'))
wf = load_workbook('out/_recalc.xlsx', data_only=False)
wv = load_workbook('out/_recalc.xlsx', data_only=True)
print('sheets:', wf.sheetnames)

# ---- 04 月差计算：核对 月差(G) / 近1月变化(L) ----
s_f, s_v = wf['04 月差计算'], wv['04 月差计算']
fmap = {}
n_ok = n_bad = n_none = 0
bad = []
for r in range(4, s_f.max_row + 1):
    label = s_f.cell(r, 4).value
    if not label or '-' not in str(label):
        continue
    g_f = s_f.cell(r, 7).value
    g_v = s_v.cell(r, 7).value
    l_f = s_f.cell(r, 12).value
    l_v = s_v.cell(r, 12).value
    fmap[label] = (g_f, g_v, l_f, l_v)
print('\n04 表公式抽查：')
for lab, (g_f, g_v, l_f, l_v) in list(fmap.items())[:12]:
    print(f'  {lab:16s} G公式={str(g_f)[:28]:30s} G值={g_v!s:>10s}  L公式={str(l_f)[:26]:28s} L值={l_v}')

# 与 json 对比
py = {}
for p, b in scan['products'].items():
    for pr in b['pairs']:
        if pr.get('computable'):
            py[pr['label']] = pr
for lab, (g_f, g_v, l_f, l_v) in fmap.items():
    if lab not in py:
        continue
    if g_v is None:
        n_none += 1
        continue
    if abs(float(g_v) - float(py[lab]['spread'])) < 0.01:
        n_ok += 1
    else:
        n_bad += 1
        bad.append((lab, g_v, py[lab]['spread']))
print(f'\n04 月差核对：一致 {n_ok}，不一致 {n_bad}，空值 {n_none}')
for b in bad[:5]:
    print('   MISMATCH', b)

# ---- 06 分位统计 ----
s6f, s6v = wf['06 分位统计'], wv['06 分位统计']
print('\n06 分位统计：')
ok = bad = none = 0
for r in range(4, s6f.max_row + 1):
    lab = s6f.cell(r, 2).value
    if not lab:
        continue
    cnt = s6v.cell(r, 3).value
    mean = s6v.cell(r, 4).value
    pct = s6v.cell(r, 13).value
    grd = s6v.cell(r, 14).value
    pj = py.get(lab)
    if pj is None:
        continue
    if pct is None or pj.get('percentile') is None:
        none += 1
        continue
    if abs(float(pct) - float(pj['percentile'])) < 0.15:
        ok += 1
    else:
        bad += 1
        print(f'   PCT MISMATCH {lab}: xlsx={pct} py={pj["percentile"]}')
    print(f'  {lab:16s} N={cnt} 均值={mean if mean is None else round(mean,1)} 分位={pct} 评级={grd}')
print(f'\n06 分位核对：一致 {ok}，不一致 {bad}，未给分位 {none}')

# ---- 03 斜率 ----
s3v = wv['03 期限结构']
print('\n03 斜率抽查：')
for r in range(1, s3v.max_row + 1):
    if s3v.cell(r, 1).value == '斜率(元/月, OLS)':
        print('  ', round(s3v.cell(r, 2).value, 2) if isinstance(s3v.cell(r, 2).value, (int, float)) else s3v.cell(r, 2).value)

print('\n=== SUMMARY ===')
print('04 ok/bad:', n_ok, n_bad, '| 06 ok/bad:', ok, bad)
