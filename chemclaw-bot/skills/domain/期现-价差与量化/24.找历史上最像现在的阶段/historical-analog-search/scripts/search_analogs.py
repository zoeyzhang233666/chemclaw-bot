# -*- coding: utf-8 -*-
"""
search_analogs.py — 跨品种通用入口：「找历史上最像现在的阶段」
================================================================
只依赖一份宽表 CSV（每行一个观测日，每列一个指标），即可完成：
  数据体检 → 无未来函数标准化 → 相似度检索 → 去重 → 后续演化 → HTML + Excel

用法（最小）
-----------
python scripts/search_analogs.py \
    --csv data/your_product_weekly.csv --date-col date \
    --asof 2026-10-14 --config config_example.json \
    --outdir build

用法（只看数据体检）
-------------------
python scripts/search_analogs.py --csv data/x.csv --date-col date --qc-only

设计约束（不可关闭的硬规则）
---------------------------
1) 任何历史时点 t 的标准化参数只用 t 当日及以前的数据；
2) 缺失一律保留为 NaN：不插值、不前值填充、不补造；
3) 若某指标在历史时点的覆盖率不足 coverage_min，该候选直接剔除
   （除非显式 --allow-partial，此时按有效权重重新归一并在输出中标记）；
4) 历史观察日与当前观察日至少相隔 --lag-days 天，避免同一段行情自比；
5) 去重：按综合分降序取，候选与已选窗口观察日相距 < --min-gap-weeks 周则跳过。
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

import numpy as np
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
import analog_lib as AL  # noqa: E402

sys.stdout.reconfigure(encoding="utf-8")

DEFAULT_STD = {"lookback": 104, "min_obs": 52, "clip": 3.0}


# --------------------------------------------------------------------------- #
# 1. 数据体检
# --------------------------------------------------------------------------- #
def qc(csv: Path, date_col: str, freq_days: int | None = None) -> dict:
    df = pd.read_csv(csv)
    rep = {"file": str(csv), "rows": len(df), "errors": [], "warnings": [], "columns": {}}
    if date_col not in df.columns:
        rep["errors"].append(f"缺少日期列 {date_col}")
        return rep
    d = pd.to_datetime(df[date_col], errors="coerce")
    if d.isna().any():
        rep["errors"].append(f"日期列有 {int(d.isna().sum())} 个无法解析的值")
    df = df.assign(_dt=d).dropna(subset=["_dt"]).sort_values("_dt")
    if df["_dt"].duplicated().any():
        rep["errors"].append(f"日期重复 {int(df['_dt'].duplicated().sum())} 行（将保留最后一条）")
    rep["date_min"] = df["_dt"].min().date().isoformat()
    rep["date_max"] = df["_dt"].max().date().isoformat()
    gaps = df["_dt"].diff().dt.days.dropna()
    if len(gaps):
        rep["gap_median_days"] = float(gaps.median())
        if freq_days:
            off = (gaps != freq_days).sum()
            rep["gap_off_freq"] = int(off)
            if off > len(gaps) * 0.15:
                rep["warnings"].append(
                    f"{off}/{len(gaps)} 个相邻间隔不等于目标频率 {freq_days} 天，"
                    f"中位数实际为 {gaps.median():.0f} 天")
    for c in df.columns:
        if c in (date_col, "_dt"):
            continue
        s = pd.to_numeric(df[c], errors="coerce")
        nn = int(s.notna().sum())
        first = df.loc[s.notna(), "_dt"]
        rep["columns"][c] = {
            "n": nn, "pct_missing": round(1 - nn / max(len(df), 1), 4),
            "first": first.min().date().isoformat() if nn else None,
            "last": first.max().date().isoformat() if nn else None,
            "min": None if not nn else float(s.min()),
            "max": None if not nn else float(s.max()),
        }
        if nn == 0:
            rep["warnings"].append(f"列 {c} 全空，将被忽略")
        elif rep["columns"][c]["pct_missing"] > 0.5:
            rep["warnings"].append(f"列 {c} 缺失 {rep['columns'][c]['pct_missing']:.0%}，"
                                   f"样本内可比时点会显著变少")
    return rep


# --------------------------------------------------------------------------- #
# 2. 标准化 + 检索
# --------------------------------------------------------------------------- #
def build_spec(cfg: dict, df: pd.DataFrame) -> dict:
    inds = []
    for ind in cfg["indicators"]:
        k = ind["key"]
        if k not in df.columns:
            continue
        inds.append({"key": k, "std_method": ind.get("std_method", "rolling_z"),
                     "lookback_days": int(ind.get("lookback", DEFAULT_STD["lookback"])),
                     "min_obs": int(ind.get("min_obs", DEFAULT_STD["min_obs"]))})
    return {"std_defaults": dict(DEFAULT_STD), "indicators": inds}


def run_search(df: pd.DataFrame, cfg: dict, a) -> dict:
    spec = build_spec(cfg, df)
    state = AL.build_state_matrix(df, spec, asof_col=a.date_col)
    raw = df.copy()
    raw[a.date_col] = pd.to_datetime(raw[a.date_col])
    raw = raw.set_index(a.date_col).sort_index()
    raw = raw[~raw.index.duplicated(keep="last")]

    weights = {ind["key"]: float(ind.get("weight", 1.0)) for ind in cfg["indicators"]
               if ind["key"] in df.columns and float(ind.get("weight", 1.0)) > 0}
    asof = pd.Timestamp(a.asof) if a.asof else raw.index.max()
    if asof not in state.index:
        prior = state.index[state.index <= asof]
        if not len(prior):
            raise SystemExit(f"观察日 {asof.date()} 早于数据起点 {raw.index.min().date()}")
        asof = prior[-1]
    hs = asof - pd.Timedelta(days=int(a.history_days))
    hs = max(hs, raw.index.min())

    p = AL.SearchParams(
        asof=asof.date().isoformat(), history_start=hs.date().isoformat(),
        window_weeks=a.window_weeks, top_n=a.top_n, min_gap_weeks=a.min_gap_weeks,
        method=a.method, weights=weights,
        min_effective_weight=0.9 if not a.allow_partial else 0.5,
        hist_z_required=not a.allow_partial,
        min_hist_z_cover=a.coverage_min,
        min_sign_match=a.min_sign_match,
        lag_at_least_days=a.lag_days)
    res = AL.search_analogs(state, raw, p)

    keys = res["keys"]
    sel = res["selected"]
    fwd_cols = a.forward_cols.split(",") if a.forward_cols else [k for k in keys if k in raw.columns]
    fwd_cols = [c.strip() for c in fwd_cols if c.strip() in raw.columns]
    months = [int(m) for m in a.forward_months.split(",")]
    # 主价格列：用于「涨跌%」展示（HTML/Excel 摘要列）
    price_col = a.price_col
    if not price_col:
        for cand in ("fut_close", "close", "price", "fut", "spot"):
            if cand in raw.columns:
                price_col = cand
                break
    if not price_col and fwd_cols:
        price_col = fwd_cols[0]
    # 统一「后续表现」键名：forward 里加一份 price_first / ret_pct，方便模板通用
    def _normalise(fwd: dict, cols) -> dict:
        for m, f in list(fwd.items()):
            if not f:
                continue
            pct = None
            if price_col:
                pct = f.get(price_col + "_pct")
            if pct is None:
                for c in cols:
                    v = f.get(c + "_pct")
                    if v is not None:
                        pct = v
                        break
            f["ret_pct"] = pct
        return fwd

    fees = []
    for r in sel.itertuples():
        fwd = AL.forward_performance(pd.DatetimeIndex(raw.index), raw, r.date,
                                     fwd_cols, tuple(months), res["window_days"])
        fees.append({"rank": int(r.rank), "date": r.date.date().isoformat(),
                     "forward": _normalise(fwd, fwd_cols)})
    now_raw = raw.loc[asof]
    detail = []
    for r in sel.itertuples():
        hraw = raw.loc[r.date]
        for k in keys:
            zv, zh = res["now_std"].get(k), state.loc[r.date, k]
            detail.append({
                "rank": int(r.rank), "date": r.date.date().isoformat(), "indicator": k,
                "weight": weights[k],
                "now_raw": None if pd.isna(now_raw.get(k)) else float(now_raw[k]),
                "hist_raw": None if pd.isna(hraw.get(k)) else float(hraw[k]),
                "now_z": None if pd.isna(zv) else float(zv),
                "hist_z": None if pd.isna(zh) else float(zh)})

    return {
        "asof": asof.date().isoformat(), "history_start": hs.date().isoformat(),
        "history_days": int(a.history_days), "method": a.method,
        "window_weeks": a.window_weeks, "top_n": a.top_n, "min_gap_weeks": a.min_gap_weeks,
        "weights": weights, "keys": keys,
        "dropped_for_now": res.get("dropped_for_now", []),
        "n_candidates": int(res["n_candidates"]),
        "now": {"date": asof.date().isoformat(),
                "raw": {k: (None if pd.isna(now_raw.get(k)) else float(now_raw[k])) for k in keys},
                "z": {k: (None if pd.isna(res["now_std"].get(k)) else float(res["now_std"][k]))
                      for k in keys}},
        "selected": [{"rank": int(r.rank), "date": r.date.date().isoformat(),
                      "similarity": float(r.similarity),
                      "path_similarity": None if pd.isna(r.path_similarity) else float(r.path_similarity),
                      "score": float(r.score), "regime_id": int(r.regime_id),
                      "hist_z_cover": float(r.hist_z_cover),
                      "sign_match": None if pd.isna(r.sign_match) else float(r.sign_match)}
                     for r in sel.itertuples()],
        "candidates": [{"date": c.date.date().isoformat(), "similarity": float(c.similarity),
                        "path_similarity": None if pd.isna(c.path_similarity) else float(c.path_similarity),
                        "score": float(c.score), "hist_z_cover": float(c.hist_z_cover)}
                       for c in res["candidates"].itertuples()],
        "forward": fees, "detail": detail,
        "sensitivity": {}, "drop_one": {},
        "config": {k: v for k, v in cfg.items() if k != "indicators"},
        "indicators": cfg["indicators"],
    }


def sensitivity(df, cfg, a, weights):
    """换相似度方法 / 换样本区间，看 Top 名单是否稳定。"""
    out = {}
    for m in ["mae", "rmse", "cosine"]:
        for hd in [260, 365, 730]:
            aa = argparse.Namespace(**vars(a))
            aa.method = m
            aa.history_days = hd
            try:
                r = run_search(df, cfg, aa)
                out[f"{m}_{hd}"] = [s["date"] for s in r["selected"]]
            except Exception as e:  # noqa: BLE001
                out[f"{m}_{hd}"] = f"ERR {e}"
    return out


def drop_one(df, cfg, a, weights, base_dates):
    out = {}
    for ind in cfg["indicators"]:
        k = ind["key"]
        if float(ind.get("weight", 0)) <= 0:
            continue
        cfg2 = json.loads(json.dumps(cfg))
        for i in cfg2["indicators"]:
            if i["key"] == k:
                i["weight"] = 0
        try:
            r = run_search(df, cfg2, a)
            got = [s["date"] for s in r["selected"]]
            ov = len(set(got) & set(base_dates))
            out[k] = {"removed": k, "top": got, "overlap_with_base": ov,
                      "overlap_pct": round(ov / max(len(base_dates), 1) * 100, 1)}
        except Exception as e:  # noqa: BLE001
            out[k] = {"removed": k, "error": str(e)}
    return out


# --------------------------------------------------------------------------- #
# 3. 产物：JSON / HTML / Excel
# --------------------------------------------------------------------------- #
def build_html(result: dict, cfg: dict, weeks: pd.DataFrame, a, outdir: Path) -> Path:
    tpl_path = Path(__file__).resolve().parent / "templates" / "analog_report_template.html"
    tpl = tpl_path.read_text(encoding="utf-8")
    inds = []
    for ind in cfg["indicators"]:
        if ind["key"] not in weeks.columns:
            continue
        inds.append({"key": ind["key"], "name": ind.get("name", ind["key"]),
                     "unit": ind.get("unit", ""), "family": ind.get("family", "other"),
                     "std_method": ind.get("std_method", "rolling_z"),
                     "default_weight": float(ind.get("weight", 1.0)),
                     "source": ind.get("source", ""),
                     "rationale": ind.get("rationale", "")})
    show = [i["key"] for i in inds]
    for extra in (a.display_cols or "").split(","):
        extra = extra.strip()
        if extra and extra in weeks.columns and extra not in show:
            show.append(extra)
    wk = []
    for _, r in weeks.sort_values(a.date_col).iterrows():
        wk.append({"date": pd.Timestamp(r[a.date_col]).date().isoformat(),
                   "raw": {k: (None if pd.isna(r[k]) else round(float(r[k]), 4)) for k in show}})
    coverage = {}
    for i in inds:
        s = weeks.loc[weeks[i["key"]].notna(), a.date_col]
        coverage[i["key"]] = {"n": int(len(s)),
                              "first": pd.Timestamp(s.min()).date().isoformat() if len(s) else None,
                              "last": pd.Timestamp(s.max()).date().isoformat() if len(s) else None,
                              "note": ""}
    payload = {
        "result": result,
        "app": {
            "product": cfg.get("product", {"key": "?", "name": "品种"}),
            "indicators": inds,
            "default_weights": {i["key"]: i["default_weight"] for i in inds},
            "methods": [
                {"key": "mae", "name": "加权平均绝对差(MAE)", "formula": "1 - Σw|Δz|/Σw/3",
                 "note": "默认，稳健、可解释"},
                {"key": "rmse", "name": "加权RMSE", "formula": "1 - sqrt(Σw·Δz²/Σw)/3",
                 "note": "对大偏差更敏感"},
                {"key": "cosine", "name": "加权余弦", "formula": "cos θ 映射到 [0,1]",
                 "note": "只看方向，不看幅度"}],
            "weeks": wk, "coverage": coverage, "events": [],
            "limits": [
                "<b>结构变化：</b>品种的供给格局、下游结构与统计口径在样本期内会变；"
                "形态相似不等于驱动机制相同。",
                "<b>样本范围：</b>各指标历史长度不同，基本面类指标往往只有近几年，"
                "「指标齐全」的可比窗口可能极少；本工具用覆盖率门槛显式暴露这一点。",
                "<b>相似度方法：</b>距离尺度与权重是人为设定；换权重/换方法名单会变（见敏感性）。",
                "<b>后续表现：</b>每个窗口只有一个结果，N 很小且高度自相关，不能当作期望值。",
                "<b>不做的事：</b>不预测价格、不给交易建议、不填补缺失数据。",
            ],
            "quality_notes": [f"数据文件：{a.csv}", f"频率：约 {a.freq_days or '—'} 天/期",
                              "所有缺失保留为空，不做插值或前值填充。"],
            "latest": wk[-1]["date"] if wk else None,
        },
    }
    html = tpl.replace("/*__DATA__*/",
                       json.dumps(payload, ensure_ascii=False, separators=(",", ":")))
    out = outdir / f"{a.name or 'analog'}_历史相似阶段研究.html"
    out.write_text(html, encoding="utf-8")
    # 把 ECharts 一并复制到输出目录，保证离线双击即可看图表
    asset = Path(__file__).resolve().parent / "templates" / "assets" / "echarts.min.js"
    if asset.exists():
        dst = outdir / "echarts.min.js"
        if not dst.exists() or dst.stat().st_size != asset.stat().st_size:
            import shutil
            shutil.copyfile(asset, dst)
    return out


def build_excel(result: dict, cfg: dict, weeks: pd.DataFrame, a, outdir: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    INK, ACC = "1F3864", "2E75B6"
    wb = Workbook()

    def head(ws, title, sub, width=10):
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
        c = ws.cell(1, 1, title)
        c.font = Font(size=14, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=INK)
        ws.row_dimensions[1].height = 28
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
        ws.cell(2, 1, sub).font = Font(size=9, italic=True, color="595959")

    def block(ws, df, r0=4):
        for j, col in enumerate(df.columns, 1):
            c = ws.cell(r0, j, str(col))
            c.font = Font(bold=True, color="FFFFFF", size=10)
            c.fill = PatternFill("solid", fgColor=ACC)
            c.alignment = Alignment(horizontal="center", wrap_text=True)
        for i, (_, row) in enumerate(df.iterrows(), 1):
            for j, col in enumerate(df.columns, 1):
                v = row[col]
                if isinstance(v, (list, tuple)):
                    v = " | ".join(map(str, v))
                ws.cell(r0 + i, j, v).font = Font(size=10)
        for j, col in enumerate(df.columns, 1):
            mx = max([len(str(col))] + [len(str(x)) for x in df[col].head(300)])
            ws.column_dimensions[get_column_letter(j)].width = min(max(10, mx * 1.4 + 2), 46)

    ws = wb.active
    ws.title = "00 读数先看"
    head(ws, f"{cfg.get('product', {}).get('name', '品种')} · 历史相似阶段",
         "本表为通用输出；结论只能在前几行标注的口径内成立，缺失即空。", 8)
    info = [("观察日", result["asof"]), ("样本起点", result["history_start"]),
            ("样本天数", result["history_days"]), ("状态窗口(周)", result["window_weeks"]),
            ("相似度方法", result["method"]), ("可比时点数", result["n_candidates"]),
            ("命中窗口数", len(result["selected"])),
            ("重要", "历史相似 ≠ 未来预测；后续表现是既成事实，样本少且自相关。")]
    for i, (k, v) in enumerate(info, 4):
        ws.cell(i, 1, k).font = Font(bold=True, size=10)
        ws.cell(i, 2, v).font = Font(size=10)
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=7)
    ws.column_dimensions["A"].width = 18

    ws2 = wb.create_sheet("01 当前状态")
    head(ws2, f"当前状态 {result['asof']}（原始值 + 标准化）",
         "z 值只用当日及以前的数据估计；空值表示该时点无数据。", 6)
    rows = [[k, result["now"]["raw"].get(k), result["now"]["z"].get(k), result["weights"].get(k)]
            for k in result["keys"]]
    block(ws2, pd.DataFrame(rows, columns=["指标", "当前原始值", "标准化z", "权重"]))

    ws3 = wb.create_sheet("02 相似窗口")
    head(ws3, "历史相似窗口（已去重）", "按综合分降序；行情轮次相同=同一段行情。", 8)
    fm = {x["date"]: x["forward"] for x in result["forward"]}
    rows = []
    for s in result["selected"]:
        f1 = (fm.get(s["date"]) or {}).get("1m") or {}
        f3 = (fm.get(s["date"]) or {}).get("3m") or {}
        rows.append([s["rank"], s["date"], round(s["similarity"], 4),
                     None if s["path_similarity"] is None else round(s["path_similarity"], 4),
                     round(s["score"], 4), s["regime_id"], round(s["hist_z_cover"], 3),
                     f1.get("ret_pct"), f3.get("ret_pct")])
    block(ws3, pd.DataFrame(rows, columns=["排名", "历史窗口", "状态相似度", "路径相似度",
                                           "综合得分", "行情轮次", "指标覆盖率", "1M涨跌%", "3M涨跌%"]))

    ws4 = wb.create_sheet("03 逐指标对比")
    head(ws4, "当前 vs 各窗口：逐指标原始值与 z", "「方向」= 标准化偏离符号是否一致。", 9)
    det = pd.DataFrame(result["detail"])
    if not det.empty:
        det["方向"] = det.apply(lambda r: ("无数据" if pd.isna(r["now_z"]) or pd.isna(r["hist_z"])
                                           else ("同向" if (r["now_z"] >= 0) == (r["hist_z"] >= 0) else "反向")),
                                axis=1)
        det.columns = ["排名", "历史窗口", "指标", "权重", "当前原始值", "历史原始值",
                       "当前z", "历史z", "方向"]
    block(ws4, det)

    ws5 = wb.create_sheet("04 后续演化")
    head(ws5, "各窗口随后 1M / 3M 的实际变化", "N 很小，不可当概率使用。", 10)
    rows = []
    for x in result["forward"]:
        for m in ("1m", "3m"):
            f = x["forward"].get(m) or {}
            rows.append([x["rank"], x["date"], m, f.get("start"), f.get("end"),
                         f.get("ret_pct")] + [f.get(c) for c in result["forward_cols"][:4]])
    block(ws5, pd.DataFrame(rows, columns=["排名", "窗口", "期", "起始", "结束", "涨跌%"])
          if not result.get("forward_cols") else
          pd.DataFrame(rows, columns=["排名", "窗口", "期", "起始", "结束", "涨跌%"] +
                       list(result["forward_cols"][:4])))

    ws6 = wb.create_sheet("05 候选全表")
    head(ws6, "样本内全部可比时点", "可自行排序复核。", 6)
    block(ws6, pd.DataFrame(result["candidates"]))

    ws7 = wb.create_sheet("06 敏感性")
    head(ws7, "换方法/换样本区间后的 Top 名单", "名单大幅变化=「最像」本身不稳定。", 6)
    rows = [[k, "、".join(v) if isinstance(v, list) else v] for k, v in result["sensitivity"].items()]
    block(ws7, pd.DataFrame(rows, columns=["设置", "命中的Top窗口"]))

    ws8 = wb.create_sheet("07 权重与指标")
    head(ws8, "指标定义、标准化方法与权重", "权重来自研究者判断，不是拟合结果。", 7)
    block(ws8, pd.DataFrame([{"指标": i.get("name", i["key"]), "代码": i["key"],
                              "单位": i.get("unit", ""), "标准化": i.get("std_method", "rolling_z"),
                              "权重": i.get("weight", 0), "来源": i.get("source", ""),
                              "理由": i.get("rationale", "")}
                             for i in cfg["indicators"]]))

    ws9 = wb.create_sheet("08 原始面板")
    head(ws9, "原始宽表（原样，未加工）", "所有上层结论都可由本表复现。", 12)
    block(ws9, weeks)

    out = outdir / f"{a.name or 'analog'}_历史相似阶段_数据工作簿.xlsx"
    wb.save(out)
    return out


# --------------------------------------------------------------------------- #
def main():
    ap = argparse.ArgumentParser(description="找历史上最像现在的阶段（跨品种通用）")
    ap.add_argument("--csv", required=True, help="宽表 CSV：每行一个观测日，每列一个指标")
    ap.add_argument("--date-col", default="date")
    ap.add_argument("--config", help="JSON 配置：indicators[{key,name,unit,weight,std_method,...}]")
    ap.add_argument("--asof", help="观察日（默认取数据最后一行）")
    ap.add_argument("--history-days", type=int, default=1825)
    ap.add_argument("--window-weeks", type=int, default=4)
    ap.add_argument("--top-n", type=int, default=5)
    ap.add_argument("--min-gap-weeks", type=int, default=8)
    ap.add_argument("--method", default="mae", choices=["mae", "rmse", "cosine", "mahalanobis"])
    ap.add_argument("--lag-days", type=int, default=21,
                    help="历史观察日与当前观察日的最小间隔，避免同一段行情自比")
    ap.add_argument("--coverage-min", type=float, default=0.9,
                    help="历史时点指标覆盖率下限")
    ap.add_argument("--min-sign-match", type=float, default=0.75,
                    help="同向指标权重占比下限")
    ap.add_argument("--allow-partial", action="store_true",
                    help="放宽：允许历史时点缺指标（按有效权重重新归一）")
    ap.add_argument("--freq-days", type=int, default=None, help="期望的观测间隔（天），仅用于体检")
    ap.add_argument("--forward-months", default="1,3")
    ap.add_argument("--forward-cols", default=None, help="后续表现要看哪些列（逗号分隔）")
    ap.add_argument("--price-col", default=None,
                    help="主价格列（用于摘要里的涨跌%%），默认自动识别 fut_close/close/price")
    ap.add_argument("--display-cols", default=None, help="HTML 额外展示列（如价格本身）")
    ap.add_argument("--name", default=None, help="产物文件名前缀")
    ap.add_argument("--outdir", default="build")
    ap.add_argument("--qc-only", action="store_true")
    ap.add_argument("--skip-sensitivity", action="store_true")
    ap.add_argument("--no-html", action="store_true")
    ap.add_argument("--no-excel", action="store_true")
    a = ap.parse_args()

    csv = Path(a.csv)
    if not csv.exists():
        raise SystemExit(f"找不到 {csv}")
    outdir = Path(a.outdir)
    outdir.mkdir(parents=True, exist_ok=True)

    print("=" * 68)
    print("1) 数据体检")
    rep = qc(csv, a.date_col, a.freq_days)
    (outdir / "qc.json").write_text(json.dumps(rep, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"   {rep['rows']} 行 {rep.get('date_min')} ~ {rep.get('date_max')}")
    for w in rep["warnings"]:
        print("   ! " + w)
    for e in rep["errors"]:
        print("   x " + e)
    if rep["errors"]:
        print("   数据体检有 ERROR，请先修复再继续。")
        if a.qc_only:
            return
    if a.qc_only:
        print(f"   体检结果已写入 {outdir/'qc.json'}")
        return

    if not a.config:
        raise SystemExit("需要 --config 指定指标与权重（可参考 data/config_example.json）")
    cfg = json.loads(Path(a.config).read_text(encoding="utf-8"))

    print("\n2) 载入宽表")
    df = pd.read_csv(csv)
    df[a.date_col] = pd.to_datetime(df[a.date_col], errors="coerce")
    df = df.dropna(subset=[a.date_col]).sort_values(a.date_col)
    df = df[~df[a.date_col].duplicated(keep="last")].reset_index(drop=True)

    print("\n3) 标准化 + 相似度检索（无未来函数）")
    result = run_search(df, cfg, a)
    print(f"   观察日 {result['asof']}  样本起点 {result['history_start']}  "
          f"可比时点 {result['n_candidates']}")
    if result["dropped_for_now"]:
        print(f"   ! 当前时点无值、未进入状态向量：{result['dropped_for_now']}")
    for s in result["selected"]:
        f = next((x["forward"] for x in result["forward"] if x["date"] == s["date"]), {})
        f1 = (f.get("1m") or {}).get("ret_pct")
        f3 = (f.get("3m") or {}).get("ret_pct")
        print(f"   #{s['rank']} {s['date']}  综合 {s['score']:.3f}  "
              f"覆盖 {s['hist_z_cover']:.0%}  1M {'—' if f1 is None else f'{f1:+.1f}%'}  "
              f"3M {'—' if f3 is None else f'{f3:+.1f}%'}")
    if not result["selected"]:
        print("   ! 没有满足条件的窗口。可尝试 --allow-partial 或放宽 --coverage-min。")

    if not a.skip_sensitivity and result["selected"]:
        print("\n4) 敏感性与单指标稳健性")
        result["sensitivity"] = sensitivity(df, cfg, a, result["weights"])
        result["drop_one"] = drop_one(df, cfg, a, result["weights"],
                                      [s["date"] for s in result["selected"]])
        for k, v in result["sensitivity"].items():
            print(f"   {k:14s} {v}")
        base = [s["date"] for s in result["selected"]]
        for k, v in result["drop_one"].items():
            print(f"   去掉 {k:18s} 重合 {v.get('overlap_with_base')}/{len(base)}")

    result["forward_cols"] = (a.forward_cols.split(",") if a.forward_cols else
                              [k for k in result["keys"] if k in df.columns])
    (outdir / "result.json").write_text(json.dumps(result, ensure_ascii=False, indent=1,
                                                   default=str), encoding="utf-8")

    print("\n5) 生成产物")
    if not a.no_html:
        p = build_html(result, cfg, df, a, outdir)
        print(f"   HTML  {p}  ({p.stat().st_size/1024:.0f} KB)")
    if not a.no_excel:
        p = build_excel(result, cfg, df, a, outdir)
        print(f"   Excel {p}")
        print(f"   JSON  {outdir/'result.json'}")
    print("\n完成。记得在结论里写明：样本范围、指标覆盖、相似度方法与局限。")


if __name__ == "__main__":
    main()
