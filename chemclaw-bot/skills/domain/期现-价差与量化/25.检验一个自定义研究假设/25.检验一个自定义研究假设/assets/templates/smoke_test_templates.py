# -*- coding: utf-8 -*-
"""smoke_test_templates.py — 两个模板的自检：生成 → 结构检查 → 清理。

模板页是骨架，不含研究内容，所以不能用 check_html.py 的内容词检查；
这里只验证结构（无外链、tab/nav 对应、标签平衡、15 个 sheet 齐全）。
"""
import os
import subprocess
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import html_report_template as H  # noqa: E402

with tempfile.TemporaryDirectory() as td:
    html_path = os.path.join(td, "demo.html")
    H.build(html_path, {
        "title": "模板自检",
        "subtitle": "smoke test",
        "verdict_badge": H.badge("bad", "样本不足"),
        "tabs": [("t1", "① 设计", '<section class="card"><h2>设计</h2></section>'),
                 ("t2", "② 数据", '<section class="card"><h2>数据</h2></section>'),
                 ("t3", "③ 结果", '<section class="card"><h2>结果</h2></section>')],
    })
    s = open(html_path, encoding="utf-8").read()
    probs = []
    if "http://" in s or "https://" in s:
        probs.append("存在外链")
    if s.count('<section class="tab"') != 3 or s.count('data-t="') != 3:
        probs.append("tab 数量不对")
    for tag in ("section", "div", "table"):
        if s.count(f"<{tag}") != s.count(f"</{tag}>"):
            probs.append(f"<{tag}> 不平衡")
    xlsx_path = os.path.join(td, "demo.xlsx")
    r = subprocess.run([sys.executable, os.path.join(HERE, "excel_workbook_template.py"), xlsx_path],
                       capture_output=True, text=True)
    if r.returncode != 0:
        probs.append("xlsx 模板失败: " + r.stderr[:200])
    else:
        import openpyxl
        wb = openpyxl.load_workbook(xlsx_path)
        if len(wb.sheetnames) != 15:
            probs.append(f"sheet 数量 {len(wb.sheetnames)} != 15")
        if not any("SIMULATED" in n for n in wb.sheetnames):
            probs.append("缺少 SIMULATED sheet")
    print("HTML 结构:", "OK" if not probs else probs)
    print("RESULT:", "OK" if not probs else "FAIL")
    sys.exit(0 if not probs else 1)
