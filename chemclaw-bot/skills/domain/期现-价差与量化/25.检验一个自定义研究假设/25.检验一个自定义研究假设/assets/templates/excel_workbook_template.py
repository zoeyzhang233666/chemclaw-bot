# -*- coding: utf-8 -*-
"""excel_workbook_template.py — 15-sheet 研究工作簿模板（改字段即可复用）。

设计约束（不可删减）：
  * 0 号 sheet 必须是「封面与结论」：假设原文、结论、样本量、为什么是这个结论
  * 原始数据 sheet 每一行都要能追到来源（URL / provenance / verified_at）
  * 统计结果 sheet 列出全部检验（含不显著），并带 BH-q
  * 必须有一个 sheet 整表标记 SIMULATED，且标题写明"禁止作为结论引用"
  * 必须有一个 sheet 写局限与风险、一个 sheet 写复现说明（sha256 + 命令）

运行：用工作区自带 Python（含 openpyxl）
  <python> -X utf8 excel_workbook_template.py out.xlsx
"""
import sys
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HDR_FILL = PatternFill("solid", fgColor="134E4A")
HDR_FONT = Font(color="FFFFFF", bold=True, size=10.5)
TITLE_FONT = Font(bold=True, size=15)
SUB_FONT = Font(size=10, color="6B7280")
BOLD = Font(bold=True)
THIN = Side(style="thin", color="E5E7EB")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FILL_SIM = PatternFill("solid", fgColor="EDE9FE")
WRAP = Alignment(wrap_text=True, vertical="top")

SHEETS = [
    "0-封面与结论", "1-研究设计", "2-参数确认", "3-数据来源与口径", "4-原始数据-库存",
    "5-原始数据-价格", "6-因子定义", "7-样本面板", "8-统计结果", "9-稳定性",
    "10-修订口径对照", "11-事件研究", "12-模拟演示(SIMULATED)", "13-局限与风险", "14-复现说明",
]


def head(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = TITLE_FONT
    if sub:
        ws["A2"] = sub
        ws["A2"].font = SUB_FONT


def table(ws, header, rows, start_row=3, widths=None, sim=False):
    for j, h in enumerate(header, 1):
        c = ws.cell(row=start_row, column=j, value=h)
        c.fill = HDR_FILL
        c.font = HDR_FONT
        c.border = BOX
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    for i, row in enumerate(rows, start_row + 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.border = BOX
            c.alignment = WRAP
            if sim:
                c.fill = FILL_SIM
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    return start_row + len(rows)


def kv(ws, pairs, start_row=4, w1=18, w2=110):
    r = start_row
    for k, v in pairs:
        a = ws.cell(row=r, column=1, value=k)
        a.font = BOLD
        a.border = BOX
        a.alignment = Alignment(vertical="top")
        b = ws.cell(row=r, column=2, value=v)
        b.alignment = WRAP
        b.border = BOX
        r += 1
    ws.column_dimensions["A"].width = w1
    ws.column_dimensions["B"].width = w2
    return r


def build(out, ctx):
    """ctx 提供各 sheet 的内容；缺省用占位，保证 15 个 sheet 齐全。"""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEETS[0]
    head(ws, f"研究假设检验报告：{ctx.get('hypothesis','（填写假设原文）')}",
         f"生成于 {datetime.now():%Y-%m-%d %H:%M}｜设计版本 {ctx.get('version','—')}")
    kv(ws, [
        ("原始假设", ctx.get("hypothesis", "")),
        ("检验结论", ctx.get("verdict", "")),
        ("样本量", ctx.get("sample", "")),
        ("为什么是这个结论", ctx.get("why", "")),
        ("数据真实性声明", "原始数据均可回溯原文；模拟演示整表标记 SIMULATED，禁止作为结论引用。"),
        ("免责声明", "观测数据只支持条件相关，不构成因果结论或投资建议。"),
    ])
    for name in SHEETS[1:]:
        s = wb.create_sheet(name)
        head(s, name.replace("-", "、"))
        if "模拟" in name or "SIMULATED" in name:
            s["A2"] = "整表为随机数生成的 SIMULATED 数据，仅用于演示流程，禁止作为结论引用。"
            s["A2"].font = Font(bold=True, color="6D28D9")
    # 各 sheet 的代表性表头
    table(wb[SHEETS[1]], ["项目", "取值"], [["（按 design json 填充）", ""]], widths=[24, 100])
    table(wb[SHEETS[2]], ["问题编号", "问题", "候选口径", "采用"], [], widths=[10, 40, 50, 40])
    table(wb[SHEETS[3]], ["序列", "来源", "URL 规则", "覆盖", "核验方式", "结果"],
          [], widths=[18, 30, 36, 20, 30, 16])
    table(wb[SHEETS[4]], ["观测日", "库存(万吨)", "provenance", "verified_at", "URL"], [],
          widths=[12, 12, 26, 20, 70])
    table(wb[SHEETS[5]], ["交易日", "现货", "近月合约", "近月价", "基差"], [], widths=[12, 12, 11, 11, 11])
    table(wb[SHEETS[6]], ["因子", "定义", "单位", "边界条件"], [], widths=[18, 40, 10, 50])
    table(wb[SHEETS[7]], ["观测日", "状态", "信号", "T0", "Δbasis_10", "ΔP_10", "roll"], [],
          widths=[12, 10, 8, 12, 12, 12, 8])
    table(wb[SHEETS[8]], ["编号", "检验族", "内容", "效应量", "95% CI", "p", "BH-q", "N", "判定"], [],
          widths=[10, 12, 50, 12, 22, 10, 10, 8, 10])
    table(wb[SHEETS[9]], ["分组", "信号 n", "对照 n", "均值差", "p"], [], widths=[14, 10, 10, 14, 10])
    table(wb[SHEETS[10]], ["口径", "说明", "信号 n", "对照 n", "主检验效应量", "p", "结论"], [],
          widths=[10, 40, 10, 10, 15, 10, 20])
    table(wb[SHEETS[11]], ["组别", "n", "T−20…T−1", "T+1…T+10", "T+11…T+20"], [],
          widths=[20, 8, 16, 16, 16])
    table(wb[SHEETS[12]], ["编号", "检验内容", "效应量", "p", "BH-q", "N", "数据状态"], [],
          widths=[10, 50, 12, 10, 10, 8, 14], sim=True)
    table(wb[SHEETS[13]], ["编号", "严重度", "说明"], [], widths=[8, 8, 120])
    table(wb[SHEETS[14]], ["序列", "文件", "行数", "sha256"], [], widths=[14, 46, 10, 22])
    wb.save(out)
    print("xlsx ->", out, len(wb.sheetnames), "sheets")


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else "workbook_template.xlsx",
          {"hypothesis": "（示例）X 连续两周下降时，随后两周 Y 更容易走强",
           "verdict": "样本不足，无法判断"})
