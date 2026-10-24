# -*- coding: utf-8 -*-
"""政策事件研究 · Excel 工作簿构建器

12 张表，与 HTML 同源同数：
 00 封面与结论     01 政策原文与日期   02 影响对象与范围   03 政策事件分类
 04 价格数据与口径 05 事件对齐价格路径 06 窗口变动明细     07 统计结果
 08 逐事件分析     09 混杂因素        10 数据缺口与来源    11 方法说明
"""
import json
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HDR_FILL = PatternFill('solid', fgColor='1F4E79')
HDR_FONT = Font(color='FFFFFF', bold=True, size=10)
TITLE_FONT = Font(size=15, bold=True, color='1F4E79')
SUB_FONT = Font(size=9, color='666666')
WARN_FILL = PatternFill('solid', fgColor='FFF2CC')
BAD_FILL = PatternFill('solid', fgColor='FCE4E4')
OK_FILL = PatternFill('solid', fgColor='E2EFDA')
THIN = Side(style='thin', color='D9D9D9')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _sheet(wb, title, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title[:31]
    return ws


def _header(ws, row, cols):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = 30
    return row + 1


def _widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _body(ws, row, rows, wrap_cols=(), num_cols=()):
    for r in rows:
        for i, v in enumerate(r, 1):
            cell = ws.cell(row=row, column=i, value=v)
            cell.border = BORDER
            cell.font = Font(size=9)
            if i in wrap_cols:
                cell.alignment = Alignment(wrap_text=True, vertical='top')
            if i in num_cols and isinstance(v, (int, float)):
                cell.number_format = '#,##0.00'
        row += 1
    return row


def build(research, study, out_path):
    meta = research.get('meta', {})
    events = research.get('events', [])
    by_id = {e['id']: e for e in study['events']}
    # 统一计算"发布->生效"提前量，避免依赖上游字段
    import datetime as _dt
    for e in events:
        try:
            if e.get('publish_date') and e.get('effective_date'):
                e['lead_days'] = (_dt.date.fromisoformat(e['effective_date']) -
                                  _dt.date.fromisoformat(e['publish_date'])).days
        except Exception:
            e['lead_days'] = None
    wb = Workbook()

    # ============ 00 封面与结论 ============
    ws = _sheet(wb, '00 封面与结论', first=True)
    _widths(ws, [26, 100])
    ws['A1'] = meta.get('title') or '政策事件研究报告'
    ws['A1'].font = TITLE_FONT
    ws['A2'] = meta.get('subtitle') or ''
    ws['A2'].font = SUB_FONT
    r = 4
    info = [('品种', meta.get('product', '')), ('研究口径', meta.get('scope', '')),
            ('时间范围', meta.get('period', '')), ('数据截止', study['calendar']['end']),
            ('交易日数', study['calendar']['n_days']), ('政策事件数', len(events)),
            ('生成日期', meta.get('asof', ''))]
    for k, v in info:
        ws.cell(row=r, column=1, value=k).font = Font(bold=True, size=10)
        ws.cell(row=r, column=2, value=v).font = Font(size=10)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value='核心结论').font = Font(bold=True, size=12, color='1F4E79')
    r += 1
    for line in (research.get('xlsxConclusions') or []):
        c = ws.cell(row=r, column=1, value=line)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        ws.row_dimensions[r].height = max(18, 15 * (len(line) // 90 + 1))
        if line.startswith('【重要】'):
            c.fill = BAD_FILL
        r += 1
    r += 1
    ws.cell(row=r, column=1, value='12 张表导航').font = Font(bold=True, size=11, color='1F4E79')
    r += 1
    for t, d in research.get('xlsxNav') or []:
        ws.cell(row=r, column=1, value=t).font = Font(bold=True, size=9)
        ws.cell(row=r, column=2, value=d).font = Font(size=9)
        ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True)
        r += 1

    # ============ 01 政策原文与日期 ============
    ws = _sheet(wb, '01 政策原文与日期')
    _widths(ws, [10, 40, 22, 20, 13, 13, 13, 10, 20, 12, 12, 60])
    cols = ['ID', '政策文件全称', '发布机关', '文号', '原文发布日', '市场首知日', '生效日',
            '提前量(天)', '政策类型', '力度', '核验状态', '核心内容（原文要点）']
    r = _header(ws, 1, cols)
    rows = []
    for e in events:
        lead = e.get('lead_days')
        rows.append([e['id'], e.get('name'), e.get('issuer'), e.get('doc_no'),
                     e.get('publish_date'), e.get('market_first_aware'), e.get('effective_date'),
                     lead, e.get('policy_type'), e.get('strength'), e.get('status'),
                     e.get('content')])
    r = _body(ws, r, rows, wrap_cols=(2, 12))
    ws.freeze_panes = 'A2'
    r += 1
    ws.cell(row=r, column=1, value='政策原文字段说明').font = Font(bold=True, size=10)
    r += 1
    notes = [
        '「原文发布日」= 文件在发布机关官网/公报上正式对外公布的日期。',
        '「市场首知日」= 市场第一次能获知该政策的时间。若存在征求意见稿、行业提前通知、媒体先曝，此日会早于原文发布日；本表填的是对市场定价生效的最早可交易信号日。',
        '「生效日」= 政策条款开始具有法律/行政约束力的日期。与原文发布日之间通常存在数周至数月的提前量（本表「提前量」列）。',
        '三者必须分开：用原文发布日做事件窗，测的是"信息公开"的一次性冲击；用生效日做事件窗，测的是"约束真正落地"，若市场已提前消化，后者通常反应更弱。',
    ]
    for n in notes:
        ws.cell(row=r, column=1, value=n).font = Font(size=9, color='444444')
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=12)
        ws.cell(row=r, column=1).alignment = Alignment(wrap_text=True)
        r += 1

    # ============ 02 影响对象与范围 ============ #
    ws = _sheet(wb, '02 影响对象与范围')
    _widths(ws, [10, 34, 14, 26, 18, 40, 16, 26])
    cols = ['ID', '政策文件', '覆盖范围', '涉及产品/环节', '涉及区域', '受影响企业与产能', '约束性', '传导渠道']
    r = _header(ws, 1, cols)
    rows = [[e['id'], e.get('name'), e.get('scope'), e.get('objects'), e.get('region_detail'),
             e.get('enterprises'), '强制' if e.get('binding') else '指导', e.get('channel')] for e in events]
    _body(ws, r, rows, wrap_cols=(2, 4, 6, 8))
    ws.freeze_panes = 'A2'

    # ============ 03 政策事件分类 ============ #
    ws = _sheet(wb, '03 政策事件分类')
    _widths(ws, [10, 34, 18, 14, 14, 12, 14, 12, 46])
    cols = ['ID', '政策文件', '政策类型', '力度', '覆盖范围', '约束性', '可比组', '是否纳入统计', '分类理由/可比性判断']
    r = _header(ws, 1, cols)
    rows = [[e['id'], e.get('name'), e.get('policy_type'), e.get('strength'), e.get('scope'),
             '强制' if e.get('binding') else '指导', e.get('comparable_group'),
             '是' if e.get('in_stats') else '否', e.get('comparable_note')] for e in events]
    _body(ws, r, rows, wrap_cols=(2, 9))
    ws.freeze_panes = 'A2'

    # ============ 04 价格数据与口径 ============ #
    ws = _sheet(wb, '04 价格数据与口径')
    _widths(ws, [26, 18, 20, 16, 16, 16, 60])
    cols = ['序列', '张数', '起始日', '截止日', '频率', '单位', '口径与来源说明']
    r = _header(ws, 1, cols)
    dq = study['data_coverage']
    srcmap = {s.get('key'): s for s in (research.get('dataQuality') or {}).get('sources', [])}
    rows = []
    for k, label in [('spot', '现货基准价'), ('future', '期货主力合约'), ('basis', '基差')]:
        d = dq[k]
        rows.append([label, d['n'], d['first'], d['last'], '日频(交易日)', '元/吨',
                     (srcmap.get(k) or {}).get('note') or ''])
    _body(ws, r, rows, wrap_cols=(7,))

    # ============ 05 事件对齐价格路径 ============ #
    ws = _sheet(wb, '05 事件对齐价格路径')
    _widths(ws, [10, 30, 12, 8, 14, 12, 12, 12, 12])
    cols = ['ID', '事件', '锚点', 'T+k', '日期', '现货', '期货', '基差', '备注']
    r = _header(ws, 1, cols)
    rows = []
    for e in events:
        s = by_id.get(e['id'])
        if not s:
            continue
        for ank, av in s['anchors'].items():
            if not av.get('ok'):
                continue
            for p in av.get('path') or []:
                rows.append([e['id'], e.get('short') or e.get('name'), ank, p['k'], p['date'],
                             p['spot'], p['future'], p['basis'],
                             '缺失(未插值)' if (p['spot'] is None or p['future'] is None) else ''])
    _body(ws, r, rows, num_cols=(6, 7, 8))
    ws.freeze_panes = 'A2'

    # ============ 06 窗口变动明细 ============ #
    ws = _sheet(wb, '06 窗口变动明细')
    _widths(ws, [10, 28, 11, 11, 8, 12, 11, 11, 10, 11, 11, 10, 11, 11, 11, 10])
    cols = ['ID', '事件', '锚点', '锚点日', '窗口', '终点日', '现货', 'Δ现货', 'Δ现货%',
            '期货', 'Δ期货', 'Δ期货%', '基差', 'Δ基差', '基准日', '可比组']
    r = _header(ws, 1, cols)
    rows = []
    for e in events:
        s = by_id.get(e['id'])
        if not s:
            continue
        for ank, av in s['anchors'].items():
            if not av.get('ok'):
                continue
            base = av['windows']['base']
            for w, wd in av['windows'].items():
                if w == 'base' or not wd.get('ok'):
                    continue
                rows.append([e['id'], e.get('short') or e.get('name'), ank, av['align']['anchor'], w,
                             wd['date'], wd['spot'], wd['spot_chg'], wd['spot_pct'],
                             wd['future'], wd['future_chg'], wd['future_pct'],
                             wd['basis'], wd['basis_chg'], base['date'],
                             e.get('comparable_group')])
    _body(ws, r, rows, num_cols=(7, 8, 9, 10, 11, 12, 13, 14))
    ws.freeze_panes = 'A2'
    if rows:
        r += 1
        ws.cell(row=r, column=1, value='Δ 为相对"基准日"（事件前最后一个交易日）的绝对变动；% 仅对现货/期货给出，基差可正可负不给百分比。').font = SUB_FONT

    # ============ 07 统计结果 ============ #
    ws = _sheet(wb, '07 统计结果')
    _widths(ws, [10, 12, 8, 12, 12, 12, 12, 12, 12, 12, 12, 14, 30])
    cols = ['指标', '窗口', '组', 'n', '均值', '中位数', '标准差', 'P25', 'P75', '最小', '最大',
            '正比例', '显著性', '成员']
    r = _header(ws, 1, cols)
    titles = {'future_chg': '期货变动(元/吨)', 'spot_chg': '现货变动(元/吨)', 'basis_chg': '基差变动(元/吨)'}
    rows = []
    any_allow = False
    for mk, mt in titles.items():
        for w in study['windows']:
            gmap = study['stats']['market_first_aware'].get(str(w), study['stats']['market_first_aware'].get(w, {})).get(mk, {})
            for g, s in gmap.items():
                q = s.get('quantiles') or {}
                any_allow = any_allow or s['significance_allowed']
                rows.append([mt, w, g, s['n'], s['mean'], s['median'], s['std'],
                             q.get('q0.25'), q.get('q0.75'), s['min'], s['max'], s['positive_ratio'],
                             ('p=%.3f' % s['p_value']) if (s['significance_allowed'] and s.get('p_value') is not None)
                             else ('样本不足' if not s['significance_allowed'] else '—'),
                             ', '.join(m['id'] for m in s.get('members', []))])
    r = _body(ws, r, rows, num_cols=(5, 6, 7, 8, 9, 10, 11, 12))
    ws.freeze_panes = 'A2'
    r += 1
    msg = ('【重要】各组样本量均 < 5，按研究规范不计算任何显著性检验，上述均为描述性统计。'
           if not any_allow else
           '仅 n≥5 的组计算了单样本 t 检验；事件不可重复、混杂因素不可控，p 值只能读作窗口内方向一致性。')
    ws.cell(row=r, column=1, value=msg).font = Font(size=9, bold=True, color='C00000')
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=14)
    ws.cell(row=r, column=1).alignment = Alignment(wrap_text=True)
    ws.cell(row=r, column=1).fill = WARN_FILL

    # ============ 08 逐事件分析 ============ #
    ws = _sheet(wb, '08 逐事件分析')
    _widths(ws, [10, 34, 40, 46, 40])
    cols = ['ID', '事件', '窗口内发生了什么（事实）', '可能的归因与边界', '同期混杂因素']
    r = _header(ws, 1, cols)
    rows = [[e['id'], e.get('name'), e.get('fact_summary'), e.get('interpretation'),
             '\n'.join(e.get('confounders') or [])] for e in events]
    _body(ws, r, rows, wrap_cols=(2, 3, 4, 5))
    for i in range(len(rows)):
        ws.row_dimensions[r - len(rows) + i].height = 120
    ws.freeze_panes = 'A2'

    # ============ 09 混杂因素 ============ #
    ws = _sheet(wb, '09 混杂因素')
    _widths(ws, [16, 30, 14, 52, 44])
    cols = ['时间节点', '因素', '类型', '为何可能影响结果', '对结论的影响']
    r = _header(ws, 1, cols)
    rows = [[c.get('when'), c.get('factor'), c.get('type'), c.get('why'), c.get('impact')]
            for c in research.get('confounders') or []]
    _body(ws, r, rows, wrap_cols=(2, 4, 5))

    # ============ 10 数据缺口与来源 ============ #
    ws = _sheet(wb, '10 数据缺口与来源')
    _widths(ws, [30, 70, 30])
    r = 1
    ws.cell(row=r, column=1, value='数据缺口').font = Font(bold=True, size=11, color='C00000')
    r += 1
    r = _header(ws, r, ['缺口', '说明', '是否影响结论'])
    gaps = (research.get('dataQuality') or {}).get('gaps') or []
    rows = [[g, '', '见说明'] if isinstance(g, str) else [g.get('item'), g.get('why'), g.get('impact')] for g in gaps]
    r = _body(ws, r, rows, wrap_cols=(1, 2, 3))
    r += 1
    ws.cell(row=r, column=1, value='数据来源清单').font = Font(bold=True, size=11, color='1F4E79')
    r += 1
    r = _header(ws, r, ['来源名称', 'URL', '抓取/日期'])
    srcs = (research.get('dataQuality') or {}).get('sources') or []
    rows = [[s.get('name'), s.get('url'), s.get('fetched')] for s in srcs]
    r = _body(ws, r, rows, wrap_cols=(1, 2, 3))
    r += 1
    ws.cell(row=r, column=1, value='政策原文来源').font = Font(bold=True, size=11, color='1F4E79')
    r += 1
    r = _header(ws, r, ['事件', '来源标题', 'URL'])
    rows = []
    for e in events:
        for s in e.get('sources') or []:
            rows.append([e['id'] + ' ' + (e.get('short') or ''), s.get('title'), s.get('url')])
    _body(ws, r, rows, wrap_cols=(2,))

    # ============ 11 方法说明 ============ #
    ws = _sheet(wb, '11 方法说明')
    _widths(ws, [30, 110])
    r = _header(ws, 1, ['项目', '说明'])
    rows = [(m.get('k'), m.get('v')) for m in research.get('method') or []]
    _body(ws, r, rows, wrap_cols=(2,))
    for i in range(len(rows)):
        ws.row_dimensions[2 + i].height = 46

    wb.save(out_path)
    return out_path


if __name__ == '__main__':
    import argparse
    ap = argparse.ArgumentParser()
    ap.add_argument('--research', required=True)
    ap.add_argument('--study', required=True)
    ap.add_argument('--out', required=True)
    a = ap.parse_args()
    def L(p):
        with open(p, 'r', encoding='utf-8-sig') as f:
            return json.load(f)
    p = build(L(a.research), L(a.study), a.out)
    print('已生成 Excel：%s' % p)
