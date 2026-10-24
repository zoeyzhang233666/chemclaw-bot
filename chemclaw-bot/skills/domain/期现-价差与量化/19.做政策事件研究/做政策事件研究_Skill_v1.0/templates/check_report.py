# -*- coding: utf-8 -*-
"""政策事件研究报告 · 交付前自动检查

检查 HTML + event_stats.json，覆盖：
  A. HTML 静态：占位符残留、外链、禁用词
  B. 统计口径：样本不足时是否真的没出现 p 值（防止"编造显著性"）
  C. 浏览器实测（需 playwright，缺失则跳过）：JS 报错、图表像素、模块渲染、三档分辨率溢出
  D. Excel（可选）：12 表齐全、关键内容非空

用法：
    python check_report.py out/report.html build/event_stats.json [--xlsx out/book.xlsx]

退出码：0 全部通过；1 存在失败项。
"""
import argparse
import json
import os
import re
import sys

FAILS = []
WARNS = []
PASSES = []


def ok(msg):
    PASSES.append(msg)
    print('  [PASS] %s' % msg)


def bad(msg):
    FAILS.append(msg)
    print('  [FAIL] %s' % msg)


def warn(msg):
    WARNS.append(msg)
    print('  [WARN] %s' % msg)


FORBIDDEN = ['预计', '看涨', '看跌', '抄底', '底部已现', '必然上涨', '一定会上涨',
             '稳赚', '必涨', '必跌', '推荐买入']


# ---------------------------------------------------------------- A 静态
def check_static(html_path):
    print('\n[A] HTML 静态检查')
    with open(html_path, 'r', encoding='utf-8') as f:
        html = f.read()

    leftovers = [p for p in ('__TITLE__', '__SUBTITLE__', '__ECHARTS__', '__DATA__') if p in html]
    if leftovers:
        bad('存在未替换占位符：%s' % leftovers)
    else:
        ok('无占位符残留')

    ext = re.findall(r'(?:src|href)="https?://[^"]+"', html)
    # 报告内链到政策原文是允许的（<a href>），只禁 <script src> / <link href> 外链
    bad_ext = re.findall(r'<script[^>]+src="https?://[^"]+"', html) + \
              re.findall(r'<link[^>]+href="https?://[^"]+"', html)
    if bad_ext:
        bad('存在外部 script/link 引用（离线打开会失效）：%s' % bad_ext[:3])
    else:
        ok('零外部 script/link（ECharts 已内联）')

    size = len(html.encode('utf-8')) / 1e6
    if size < 0.5:
        bad('HTML 仅 %.2f MB，ECharts 可能未内联成功' % size)
    else:
        ok('HTML 体积 %.2f MB（ECharts 内联正常）' % size)

    hit = [(w, html.count(w)) for w in FORBIDDEN if w in html]
    # 允许出现在"不要写 XX"这类否定语境中，做一次粗筛
    real = []
    for w, c in hit:
        for m in re.finditer(re.escape(w), html):
            seg = html[max(0, m.start() - 12):m.start()]
            if not re.search(r'(禁止|不要|不得|不应|勿)', seg):
                real.append(w)
                break
    if real:
        bad('出现预测性/诱导性禁用词：%s' % set(real))
    else:
        ok('禁用词检查通过')


# ---------------------------------------------------------------- B 统计口径
def check_stats(study_path):
    print('\n[B] 统计口径检查')
    with open(study_path, 'r', encoding='utf-8') as f:
        st = json.load(f)

    min_n = st.get('min_n_for_significance', 5)
    violations, allowed_groups, all_n = [], [], []
    for anchor, by_w in st['stats'].items():
        for w, by_m in by_w.items():
            for metric, by_g in by_m.items():
                for g, s in by_g.items():
                    all_n.append(s['n'])
                    if s['n'] < min_n and s.get('p_value') is not None:
                        violations.append('%s/+%s/%s/%s n=%d 却给出了 p=%s'
                                          % (anchor, w, metric, g, s['n'], s['p_value']))
                    if s.get('significance_allowed'):
                        allowed_groups.append((g, s['n']))

    if violations:
        bad('样本不足却计算了显著性（严重造假风险）：%s' % violations[:3])
    else:
        ok('n < %d 的组全部未计算 p 值（p_value=None）' % min_n)

    if allowed_groups:
        gs = sorted(set(allowed_groups))
        ok('达到统计门槛的组：%s' % gs)
    else:
        warn('没有任何组达到 n≥%d，报告应只做逐事件分析（若非预期请检查分组）' % min_n)

    if not all_n:
        bad('统计结果为空，没有任何组')
    else:
        ok('统计覆盖 %d 个 组×窗口×指标 组合，最大样本量 n=%d' % (len(all_n), max(all_n)))

    # 缺数据是否被插值
    interp = []
    for ev in st.get('events', []):
        for ank, av in ev['anchors'].items():
            if not av.get('ok'):
                continue
            for p in av.get('path') or []:
                if p.get('spot') is not None and p.get('future') is None and p.get('basis') is not None:
                    interp.append('%s/%s/%s' % (ev['id'], ank, p['date']))
    if interp:
        bad('发现只用现货就凭空得出基差的点（跨日拼凑）：%s' % interp[:3])
    else:
        ok('基差仅在现货与期货同日齐备时计算，无跨日拼凑')

    # 三锚点是否都跑了
    anchors = set()
    for ev in st.get('events', []):
        anchors.update(k for k, v in ev['anchors'].items() if v.get('ok'))
    if len(anchors) >= 2:
        ok('已使用多个时间锚点：%s' % sorted(anchors))
    else:
        bad('只跑了一个时间锚点（%s），必须区分"市场首知/原文发布/政策生效"'
            % sorted(anchors))

    groups = st.get('groups') or {}
    if len(groups) >= 2:
        ok('已按可比性分 %d 组，未机械合并' % len(groups))
    else:
        warn('只有 1 个可比组，请确认是否把所有事件强行合并了')


# ---------------------------------------------------------------- C 浏览器实测
def check_render(html_path):
    print('\n[C] 浏览器实测')
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        warn('未安装 playwright，跳过浏览器实测（python -m pip install playwright && python -m playwright install chromium）')
        return

    errs = []
    url = 'file:///' + os.path.abspath(html_path).replace('\\', '/')
    with sync_playwright() as pw:
        b = pw.chromium.launch()
        pg = b.new_page(viewport={'width': 1600, 'height': 1000})
        pg.on('console', lambda m: errs.append(m.text) if m.type == 'error' else None)
        pg.on('pageerror', lambda e: errs.append(str(e)))
        pg.goto(url)
        pg.wait_for_timeout(2500)

        r = pg.evaluate("""() => {
          const q = s => document.querySelectorAll(s).length;
          const charts = {};
          ['multiPath','singlePath','boxChart'].forEach(id=>{
            const c=document.getElementById(id);
            if(!c) return;
            const cv=c.querySelector('canvas');
            const inst=echarts.getInstanceByDom(c);
            charts[id]={canvas: cv?[cv.width,cv.height]:null,
                        series: inst?inst.getOption().series.length:0};
          });
          return {kpis:q('.kpi'), callouts:q('#verdict .callout'),
                  cand:q('#candTable tbody tr'), tl:q('.tl-item'),
                  cards:q('.ev-card'), statRows:q('#statsBlock tbody tr'),
                  conf:q('#confounders tbody tr'), limits:q('#limits .callout'),
                  gaps:q('#dataQuality ul li'), charts};
        }""")

        if errs:
            bad('浏览器 JS 错误 %d 条：%s' % (len(errs), errs[:3]))
        else:
            ok('浏览器无 JS 错误')

        for name, need in [('kpis', 1), ('callouts', 1), ('cand', 1), ('tl', 1),
                           ('cards', 1), ('statRows', 1), ('conf', 1),
                           ('limits', 1), ('gaps', 1)]:
            if r[name] >= need:
                ok('模块 %s 渲染 %d 项' % (name, r[name]))
            else:
                bad('模块 %s 未渲染（%d 项）' % (name, r[name]))

        for cid, c in r['charts'].items():
            if not c['canvas']:
                bad('图表 %s 未画出 canvas' % cid)
            elif c['canvas'][0] < 100 or c['canvas'][1] < 100:
                bad('图表 %s 尺寸异常 %s' % (cid, c['canvas']))
            elif c['series'] < 1:
                bad('图表 %s 没有 series 数据' % cid)
            else:
                ok('图表 %s 已绘制 %dx%d，%d 条序列' % (cid, c['canvas'][0], c['canvas'][1], c['series']))

        for w in (1600, 1280, 420):
            pg.set_viewport_size({'width': w, 'height': 900})
            pg.wait_for_timeout(600)
            ov = pg.evaluate("() => document.documentElement.scrollWidth - document.documentElement.clientWidth")
            if ov > 2:
                bad('视口 %dpx 横向溢出 %dpx' % (w, ov))
            else:
                ok('视口 %dpx 无横向溢出' % w)
        b.close()


# ---------------------------------------------------------------- D Excel
def check_xlsx(path):
    print('\n[D] Excel 检查')
    try:
        import openpyxl
    except ImportError:
        warn('未安装 openpyxl，跳过 Excel 检查')
        return
    wb = openpyxl.load_workbook(path)
    need = ['00 封面与结论', '01 政策原文与日期', '02 影响对象与范围', '03 政策事件分类',
            '04 价格数据与口径', '05 事件对齐价格路径', '06 窗口变动明细', '07 统计结果',
            '08 逐事件分析', '09 混杂因素', '10 数据缺口与来源', '11 方法说明']
    miss = [n for n in need if n not in wb.sheetnames]
    if miss:
        bad('缺少工作表：%s' % miss)
    else:
        ok('12 张表齐全')

    empty = [n for n in wb.sheetnames if wb[n].max_row < 2]
    if empty:
        bad('工作表无数据：%s' % empty)
    else:
        ok('所有工作表均有数据')

    # 规模检查
    for n, minrow in [('05 事件对齐价格路径', 20), ('06 窗口变动明细', 10), ('07 统计结果', 5)]:
        if n in wb.sheetnames and wb[n].max_row < minrow:
            bad('%s 仅 %d 行，内容可能缺失' % (n, wb[n].max_row))
    ok('核心表行数达标')


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('html')
    ap.add_argument('study')
    ap.add_argument('--xlsx', default=None)
    ap.add_argument('--no-render', action='store_true')
    a = ap.parse_args()

    print('=' * 70)
    print('政策事件研究报告 · 交付前检查')
    print('=' * 70)
    check_static(a.html)
    check_stats(a.study)
    if not a.no_render:
        check_render(a.html)
    if a.xlsx:
        check_xlsx(a.xlsx)

    print('\n' + '=' * 70)
    print('通过 %d 项 ｜ 警告 %d 项 ｜ 失败 %d 项' % (len(PASSES), len(WARNS), len(FAILS)))
    if FAILS:
        print('\n失败项：')
        for f in FAILS:
            print('  - ' + f)
        print('\n结论：未通过，请修复后再交付。')
        sys.exit(1)
    print('\n结论：全部通过。')
    sys.exit(0)


if __name__ == '__main__':
    main()
