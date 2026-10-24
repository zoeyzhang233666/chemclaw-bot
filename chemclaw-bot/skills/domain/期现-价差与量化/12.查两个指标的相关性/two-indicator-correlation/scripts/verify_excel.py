# -*- coding: utf-8 -*-
"""核对重算后工作簿的关键数值（含分阶段、滚动、矩阵），并与 Python 端逐项比对。"""
import csv
import os
import sys
from datetime import datetime

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
import corrlib as C  # noqa: E402

XLSX = os.path.join(ROOT, 'docs', 'recalc.xlsx')
wb = load_workbook(XLSX, data_only=True)
rows = list(csv.DictReader(open(os.path.join(ROOT, 'data', 'demo_methanol_weekly.csv'), encoding='utf-8-sig')))
dates = [r['week_ending'] for r in rows]
inv = [float(r['inventory_total']) for r in rows]
price = [float(r['spot_taicang']) for r in rows]
dI = C.transform(inv, 'change')
pP = C.transform(price, 'pct')
a, b, dts = C.align(dI, pP, dates)

errors = []


def cmp(name, got, want, tol=1e-6):
    ok = (got is not None and want is not None and abs(got - want) <= tol)
    print(f'  {"OK " if ok else "BAD"} {name}: excel={got} python={want if want is None else round(want, 6)}')
    if not ok:
        errors.append(name)


def P(xv, yv):
    """先按位置配对剔除缺失，再求相关（与 Excel 的 CORREL 行为一致）。"""
    px, py, _ = C.align(xv, yv, dates)
    return C.pearson(px, py)


print('== 05 总体相关（与 Python 对照）')
ws5 = wb['05 总体相关']
pairs = [
    ('水平×水平', P(inv, price)),
    ('水平×环比', P(inv, pP)),
    ('差分×环比', P(dI, pP)),
    ('差分×差分', P(dI, C.transform(price, 'change'))),
    ('环比×环比', P(C.transform(inv, 'pct'), pP)),
]
for i, (nm, want) in enumerate(pairs):
    r = 5 + i
    cmp(nm, ws5.cell(row=r, column=3).value, want)
    pv = ws5.cell(row=r, column=7).value
    print(f'      p={pv} 判定={ws5.cell(row=r, column=8).value}')

print('\n== 06 滚动相关（26 周）')
ws6 = wb['06 滚动相关']
py_roll = C.rolling_corr(dI, pP, 26, min_pairs=21)
checked = 0
for i in range(len(rows) - 1, -1, -1):
    xl = ws6.cell(row=5 + i, column=5).value
    want = py_roll[i]
    if isinstance(xl, (int, float)) and want is not None:
        cmp(f'26w r @{dates[i]}', xl, want, 1e-6)
        checked += 1
    if checked >= 4:
        break
for i in range(0, 60):
    xl = ws6.cell(row=5 + i, column=5).value
    if xl in ('样本不足', None, ''):
        print(f'  OK  26w 前段 @{dates[i]} 双方均不计算（excel={xl!r}, python={py_roll[i]!r}）')
        break

print('\n== 07 分阶段')
ws7 = wb['07 分阶段结果']
breaks = ['2023-10-11', '2024-09-01', '2025-09-01', '2026-09-30']
for k in range(3):
    lo, hi = breaks[k], breaks[k + 1]
    idx = [i for i, d in enumerate(dates) if lo <= d <= hi]
    want = P([dI[i] for i in idx], [pP[i] for i in idx])
    cmp(f'阶段{k+1} r', ws7.cell(row=8 + k, column=5).value, want, 1e-6)
    print(f'      N excel={ws7.cell(row=8+k,column=4).value} python={sum(1 for i in idx if dI[i] is not None and pP[i] is not None)}')

print('\n== 08 形态矩阵 3x3')
ws8 = wb['08 形态组合矩阵']
forms = [('level', inv), ('change', dI), ('pct', C.transform(inv, 'pct'))]
pforms = [('level', price), ('change', C.transform(price, 'change')), ('pct', pP)]
for i, (_, xv) in enumerate(forms):
    for j, (_, yv) in enumerate(pforms):
        want = P(xv, yv)
        got = ws8.cell(row=5 + i, column=3 + j).value
        cmp(f'矩阵[{i+1},{j+1}]', got, want, 1e-6)

print('\n== 09b 去趋势')
ws9 = wb['09b 辅助计算(残差)']
t = list(range(len(inv)))
want_dt = C.pearson(C.detrend(t, inv), C.detrend(t, price))
cmp('去趋势 r', ws9.cell(row=len(rows) + 6, column=4).value, want_dt, 1e-6)

print('\n结果：', '全部一致' if not errors else f'{len(errors)} 项不一致: {errors}')
sys.exit(0 if not errors else 1)

