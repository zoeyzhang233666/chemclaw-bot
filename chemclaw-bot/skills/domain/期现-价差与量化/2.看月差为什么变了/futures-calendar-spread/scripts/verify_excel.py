# -*- coding: utf-8 -*-
"""校验 Excel 工作簿：结构完整、无空表、公式可算、关键数字与 result.json 一致。

用法：python scripts/verify_excel.py <workbook.xlsx> <enriched.json>
"""
from __future__ import annotations

import json
import sys

from openpyxl import load_workbook

EXPECT_SHEETS = [
    '01-使用说明', '02-结论摘要', '03-合约价格', '04-月差历史', '05-近期变化',
    '06-历史分位', '07-可比样本', '08-期限结构', '09-贡献分解', '10-近端结构演变',
    '11-仓单库存', '12-事件台账', '13-计算公式', '14-数据来源', '15-数据质量',
]
FAILS = []


def ck(name, cond, detail=''):
    print(f"  [{'PASS' if cond else 'FAIL'}] {name}" + ('' if cond else f' {detail}'))
    if not cond:
        FAILS.append(name)


def main() -> int:
    xlsx, js = sys.argv[1], sys.argv[2]
    D = json.load(open(js, encoding='utf-8'))
    wb = load_workbook(xlsx)  # 不带 data_only，用于读公式
    print('1) 工作表结构')
    ck('工作表数量 15', len(wb.sheetnames) == 15, wb.sheetnames)
    ck('工作表名称与顺序正确', wb.sheetnames == EXPECT_SHEETS, wb.sheetnames)

    print('2) 每张表都有实质内容')
    for name in wb.sheetnames:
        ws = wb[name]
        nonempty = sum(1 for row in ws.iter_rows() for c in row if c.value not in (None, ''))
        ck(f'{name} 非空单元格 ≥ 12', nonempty >= 12, f'实际 {nonempty}')
        ck(f'{name} 有标题', ws['A1'].value not in (None, ''), ws['A1'].value)

    print('3) 关键数字与 result.json 一致')
    ws = wb['03-合约价格']
    # 找出最后一行有效数据
    last = None
    for r in range(5, ws.max_row + 1):
        if isinstance(ws.cell(row=r, column=1).value, str) and ws.cell(row=r, column=1).value[:2] == '20':
            last = r
    ck('价格表有数据行', last is not None and last > 100, last)
    ck('末行近月价格与当期一致',
       abs(float(ws.cell(row=last, column=2).value) - D['spread']['a_price']) < 1e-6,
       (ws.cell(row=last, column=2).value, D['spread']['a_price']))
    ck('末行远月价格与当期一致',
       abs(float(ws.cell(row=last, column=5).value) - D['spread']['b_price']) < 1e-6,
       (ws.cell(row=last, column=5).value, D['spread']['b_price']))
    ck('月差用公式而非硬编码', str(ws.cell(row=last, column=7).value).startswith('='),
       ws.cell(row=last, column=7).value)
    ck('行数 = 月差样本数', last - 5 + 1 == len(D['spread']['series']),
       (last - 5 + 1, len(D['spread']['series'])))

    ws4 = wb['04-月差历史']
    ck('月差历史表首行日期正确',
       ws4.cell(row=5, column=1).value == D['spread']['series'][0]['date'],
       ws4.cell(row=5, column=1).value)

    print('4) 分位页如实标注不可用区间')
    ws6 = wb['06-历史分位']
    txt = ' '.join(str(c.value) for row in ws6.iter_rows() for c in row if c.value)
    ck('含“不可称为该区间分位”类提示', ('不足' in txt) or ('重合' in txt) or ('不可' in txt), txt[:160])

    print('5) 事件台账带链接与等级说明')
    ws12 = wb['12-事件台账']
    n_ev = len(D['events'])
    got = sum(1 for r in range(5, 5 + n_ev) if ws12.cell(row=r, column=1).value)
    ck('事件行数一致', got == n_ev, (got, n_ev))
    links = sum(1 for r in range(5, 5 + n_ev) if ws12.cell(row=r, column=7).hyperlink)
    ck('每条事件都有超链接', links == n_ev, (links, n_ev))

    print('6) 可比样本与期限结构')
    ws7 = wb['07-可比样本']
    ck('可比样本组合数一致',
       sum(1 for r in range(5, 5 + len(D['comparable']['pairs'])) if ws7.cell(row=r, column=1).value)
       == len(D['comparable']['pairs']))
    ws8 = wb['08-期限结构']
    n_ts = len(D['term_structure']['rows'])
    ck('期限结构行数一致',
       sum(1 for r in range(5, 5 + n_ts) if ws8.cell(row=r, column=1).value) == n_ts)
    ck('期限结构首行 TS 归零',
       abs(float(ws8.cell(row=5, column=6).value)) < 1e-9, ws8.cell(row=5, column=6).value)

    print('7) 公式页可自洽')
    ws13 = wb['13-计算公式']
    ck('对数贡献用 LN', any('LN(' in str(c.value) for row in ws13.iter_rows() for c in row if c.value))
    ck('有残差校验公式', any('B14+B15' in str(c.value) for row in ws13.iter_rows() for c in row if c.value))

    print('8) 数据质量页')
    ws15 = wb['15-数据质量']
    ck('数据质量表有结论行', ws15.cell(row=5, column=1).value is not None)

    print()
    if FAILS:
        print(f'FAILED {len(FAILS)}: {FAILS}')
        return 1
    print('ALL PASSED')
    return 0


if __name__ == '__main__':
    sys.exit(main())
